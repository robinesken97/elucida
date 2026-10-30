"""elucida/selftest.py - Selbsttest (CI und gebaute App: Elucida --selftest ORDNER [--selftest-isnet]).

Neue Funktion -> eigener Abschnitt mit report.append("…: OK"); Meldungen nur ASCII/cp1252 (Windows-Konsole).
"""
from __future__ import annotations

import json
import os
import re
import shutil
import sys
import time
import traceback
from pathlib import Path

from elucida.app import main_window as MW  # zuerst: setzt das matplotlib-Backend (QtAgg)

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure as MplFigure
from PySide6.QtCore import QEvent, QSettings, Qt
from PySide6.QtGui import QDesktopServices
from PySide6.QtWidgets import (QAbstractItemView, QAbstractSpinBox, QApplication, QFrame, QGroupBox, QLabel,
                               QPushButton, QScrollArea, QSpinBox, QSplitter, QVBoxLayout, QWidget)

from elucida import APP_NAME, APP_VERSION, ROOT  # noqa: E402
from elucida.app.main_window import (DEFAULT_GLOB, OLD_APP_NAMES, PREVIEW_DPI, MainWindow, TextFormatDialog,
                                     compound_from_structure, migrate_settings)
from elucida.app import project as PRJ
from elucida.app import style as STYLE
from elucida.app import update as UPD
from elucida.app.i18n import T
from elucida.core import color as COLOR
from elucida.core import images as IMG
from elucida.core import uvvis as core
from elucida.si import lines as SI
from elucida.tabs import base as TABS
from elucida.tabs import cv as CVMOD
from elucida.tabs import dft as DFT
from elucida.tabs import epr as EPRMOD
from elucida.tabs import hrms as HRMS
from elucida.tabs import ir_ea as IRMS
from elucida.tabs import nmr as NMR
from elucida.tabs import sec as SECMOD
from elucida.tabs import synthesis as SYN
from elucida.tabs import xray as XRAY


# ----------------------------------------------------------------------------
# Selbsttest (für CI und nach der Installation)
# ----------------------------------------------------------------------------
def make_selftest_inputs(d: Path):
    """Synthetische Cary-CSV (2 Serien), CDXML und Küvettenfoto (HEIC, falls möglich)."""
    from PIL import Image, ImageDraw
    d.mkdir(parents=True, exist_ok=True)
    wl = np.arange(1100.0, 199.0, -1.0)
    bands = {"TEST-A": [(262, 6.8, 20), (484, 1.6, 30), (732, 0.25, 40), (816, 0.24, 40)],
             "TEST-B": [(255, 7.3, 18), (351, 2.7, 20), (444, 3.5, 12), (719, 0.42, 35), (804, 0.41, 35)]}
    concs = [("1", 1.0), ("0p5", 0.5), ("0p25", 0.25), ("0p125", 0.125), ("0p0625", 0.0625)]
    rng = np.random.default_rng(0)
    cols, names = [], []
    for ser, bl in bands.items():
        a = sum(h * np.exp(-0.5 * ((wl - l) / w) ** 2) for l, h, w in bl)
        for tag, c in concs:
            names.append(f"{ser}-{tag}mgml_THF")
            y = np.minimum(a * c, 4.0) - 0.03 + rng.normal(0, 0.002, wl.size)
            cols.append(y)
    lines = ["".join(f"{n},," for n in names), "".join("Wavelength (nm),Abs," for _ in names)]
    for i, x in enumerate(wl):
        lines.append("".join(f"{x:.6f},{y[i]:.8f}," for y in cols))
    lines += ["", "TEST-A-1mgml_THF", "Instrument  Cary 60"]
    (d / "selftest.csv").write_text("\r\n".join(lines), encoding="utf-8", newline="")
    nodes = [(1, 100, 100, 6), (2, 112.5, 92.8, 6), (3, 125, 100, 6), (4, 125, 114.4, 6),
             (5, 112.5, 121.6, 6), (6, 100, 114.4, 7), (7, 112.5, 78.4, 8)]
    bonds = [(1, 2, 2), (2, 3, 1), (3, 4, 2), (4, 5, 1), (5, 6, 2), (6, 1, 1), (2, 7, 1)]
    xml = ['<?xml version="1.0" encoding="UTF-8" ?><CDXML BondLength="14.40"><page id="1"><fragment id="99">']
    xml += [f'<n id="{i}" p="{x} {y}"' + (f' Element="{e}"' if e != 6 else "") + "/>" for i, x, y, e in nodes]
    xml += [f'<b id="{20 + k}" B="{b}" E="{e}"' + (' Order="2"' if o == 2 else "") + "/>"
            for k, (b, e, o) in enumerate(bonds)]
    xml += ["</fragment></page></CDXML>"]
    (d / "pyridinol.cdxml").write_text("".join(xml), encoding="utf-8")
    im = Image.new("RGB", (600, 900), (190, 185, 175))
    dr = ImageDraw.Draw(im)
    dr.rectangle([240, 100, 360, 200], fill=(25, 25, 25))
    dr.rectangle([220, 200, 380, 820], fill=(215, 220, 228))
    dr.rectangle([235, 420, 365, 805], fill=(195, 90, 35))
    photo = d / "kuevette.png"
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
        im.save(d / "kuevette.heic", quality=90)
        photo = d / "kuevette.heic"
    except Exception:
        im.save(photo)
    return d / "selftest.csv", d / "pyridinol.cdxml", photo


def selftest(outdir, with_isnet=False):
    out = Path(outdir).resolve()
    out.mkdir(parents=True, exist_ok=True)
    MW.SETTINGS_FILE = out / "selftest_settings.ini"  # zuletzt geöffnet, Sprache usw. nicht beim Nutzer ändern
    MW.SETTINGS_FILE.unlink(missing_ok=True)
    report = [f"{APP_NAME} {APP_VERSION} selftest", f"python {sys.version.split()[0]}, {sys.platform}"]
    ok = True
    app = QApplication.instance() or QApplication(sys.argv[:1])  # noqa: F841
    QApplication.setApplicationVersion(APP_VERSION)
    try:
        csv, cdxml, photo = make_selftest_inputs(out / "input")
        report.append(f"photo format: {photo.suffix}")
        w = MainWindow()
        w._allow_download = with_isnet
        assert w.center.currentIndex() == 0, "ohne Projekt: Startbildschirm"
        proj = out / "projekt_test"
        assert not w.open_csv(csv), "ohne Projekt darf nichts geladen werden"
        assert w.create_project(proj)
        QApplication.processEvents()
        w.open_csv(csv)
        assert set(w.groups) == {"TEST-A", "TEST-B"}, w.groups
        assert Path(w.csv_path).parent == proj / "data" / "uvvis", w.csv_path          # ins Projekt kopiert
        w.set_structure(str(cdxml), "TEST-A", render=False, sync=True)
        st = w.series["TEST-A"]
        report.append(f"CDXML: {st['chem_info']!r}")
        assert st["mw"] and abs(st["mw"] - 95.10) < 0.05, st["mw"]      # 3-Hydroxypyridin C5H5NO
        hint_ = lambda: T("cdxml_drawn", name="pyridinol")                # noqa: E731  Hinweis folgt der Sprache
        assert hint_() in st["chem_info"], st["chem_info"]
        for lang in ("en", "de"):
            w.switch_lang(lang)
            assert hint_() in w.series["TEST-A"]["chem_info"], (lang, w.series["TEST-A"]["chem_info"])
        report.append("structure hints follow language switch: OK")
        w.series["TEST-A"]["bg"] = "isnet" if with_isnet else "grabcut"
        w.set_photo(str(photo), "TEST-A", render=False, sync=True)
        bgm = st["photo_asset"]["bg_method"]
        report.append(f"background: {bgm}")
        if with_isnet:
            assert bgm == "isnet", bgm
        w.render()
        files = w.export_all(out)
        for f in files:
            assert f.exists() and f.stat().st_size > 200, f
        import pymupdf
        page = pymupdf.open(str(out / "selftest_TEST-A.pdf"))[0]
        report.append(f"pdf: {len(page.get_drawings())} vector paths, {len(page.get_images())} images")
        assert page.get_images(), "photo missing in PDF"
        svg = (out / "selftest_TEST-A.svg").read_text(encoding="utf-8")
        assert "structure_0" in svg and "<image" in svg, "SVG ohne Struktur/Foto"
        report.append(f"pdf size: {page.rect.width / 72 * 2.54:.2f} x {page.rect.height / 72 * 2.54:.2f} cm")
        # Probennamen ohne Konzentration (z. B. c0..c4) -> Overrides wie aus dem Dialog
        txt = csv.read_text(encoding="utf-8").splitlines()
        n_samples = len([n for n in txt[0].split(",") if n])
        txt[0] = "".join(f"c{i},," for i in range(n_samples))
        plain = out / "input" / "unnamed.csv"
        plain.write_text("\r\n".join(txt), encoding="utf-8", newline="")
        assert w.open_csv(plain, ask=False) is False                 # ohne Angaben: nichts
        w.conc = {f"c{i}": {"conc": 1.0 / 2 ** i, "unit": "mg/mL", "series": "unbenannt"}
                  for i in range(5)}
        assert w.rebuild_groups() and "unbenannt" in w.groups, w.groups
        assert len(w.groups["unbenannt"]) == 5
        w.render()
        report.append("unnamed samples: OK")
        # Einzelmessungen in getrennten Dateien + Cary-Binärdatei (.DSW)
        import struct
        singles = []
        for i, s_ in enumerate(core.read_samples([csv])[:5]):
            f = out / "input" / f"single_{i}.csv"
            f.write_text("\r\n".join([f"{s_['name']},,", "Wavelength (nm),Abs,"] +
                                      [f"{a:.4f},{b:.6f}," for a, b in zip(s_["x"], s_["y"])]),
                         encoding="utf-8", newline="")
            singles.append(f)
        assert w.open_files(singles, ask=False) and len(w.groups["TEST-A"]) == 5, w.groups
        magic = b"Varian UV-VIS Spectrophotometer"
        blob = bytearray([len(magic)]) + magic
        blob += bytes(0x400 - len(blob))
        for k in range(3):
            blk = bytearray(256)
            t = f"DSW-T-{0.5 / 2 ** k:g}".replace(".", "p").encode() + b"mgml"
            blk[:len(t)] = t
            blob += blk
            for x in np.arange(800.0, 249.0, -1.0):
                blob += struct.pack("<ff", x, 0.4 / 2 ** k * np.exp(-((x - 450) / 40) ** 2))
            blob += bytes(64)
        dsw = out / "input" / "test.DSW"
        dsw.write_bytes(bytes(blob))
        assert w.open_files([dsw], ask=False) and len(w.groups.get("DSW-T", [])) == 3, w.groups
        report.append("multiple files + DSW: OK")
        # Overlay (Messung + TD-DFT) und Fluoreszenz
        orca = out / "input" / "orca.out"
        orca.write_text("\n".join([
            "-" * 77, "         ABSORPTION SPECTRUM VIA TRANSITION ELECTRIC DIPOLE MOMENTS", "-" * 77,
            "State   Energy    Wavelength  fosc         T2        TX        TY        TZ",
            "        (cm-1)      (nm)                 (au**2)    (au)      (au)      (au)", "-" * 77,
            "   1   20660.0    484.0   0.150000000   1.00000   1.00000   0.00000   0.00000",
            "   2   38168.0    262.0   0.600000000   3.00000   1.70000   0.00000   0.00000", ""]),
            encoding="utf-8")
        assert len(__import__("elucida.core.tddft", fromlist=["_"]).parse_tddft(orca)["E"]) == 2
        w.ov_state["entries"] = [
            {"kind": "exp", "file": str(csv), "sample": "TEST-A-0p25mgml_THF", "label": "exp", "color": "#000000",
             "visible": True, "own_norm": False, "nlo": 450, "nhi": 520, "conc": 0.25, "unit": "mg/mL",
             "mw": 95.1, "d": 1.0},
            {"kind": "calc", "file": str(orca), "label": "calc", "color": "#1764e8", "visible": True,
             "own_norm": False, "nlo": 450, "nhi": 520, "fwhm": 0.3, "shift": 0.0}]
        w.ov_state.update(nlo=450, nhi=520)
        for mode in ("norm", "eps"):
            w.ov_state["mode"] = mode
            files = w.ov_tab.export_to(out / f"overlay_{mode}")
            assert all(f.exists() for f in files), files
        # Hinweis im Tab statt nur im Protokoll: ε ohne Konzentration -> Messung ausgelassen, Grund unter der Vorschau
        w.ov_state["entries"][0]["conc"] = None
        w.ov_tab.render()
        assert not w.ov_tab.lbl_note.isHidden() and "„exp“" in w.ov_tab.lbl_note.text(), w.ov_tab.lbl_note.text()
        w.ov_state["entries"][0]["conc"] = 0.25
        w.ov_state["mode"] = "norm"
        w.ov_tab.render()
        assert w.ov_tab.lbl_note.isHidden(), w.ov_tab.lbl_note.text()
        report.append("overlay notes in the tab: OK")
        em = out / "input" / "emission.csv"
        xs = np.arange(400.0, 701.0, 1.0)
        em.write_text("\r\n".join(["EM ex352,,", "Wavelength (nm),Intensity (a.u.),"] +
                                    [f"{x:.1f},{500 * np.exp(-((x - 520) / 30) ** 2):.3f}," for x in xs]),
                      encoding="utf-8", newline="")
        w.fl_state.update(abs_file=str(csv), abs_sample="TEST-A-0p25mgml_THF", em_file=str(em),
                          em_sample="EM ex352", ex=352.0, alo=450, ahi=520, eps="1234",
                          photo_day=str(photo), bg_day="grabcut", photo_uv=str(photo), bg_uv="none")
        files = w.fl_tab.export_to(out / "fluorescence")
        assert all(f.exists() for f in files), files
        # Cary-Eclipse-CSV (Name in Zeile 2, λex aus den Metadaten)
        ecl = out / "input" / "eclipse.csv"
        ecl.write_text("\r\n".join(["Wavelength (nm),Intensity (a.u.),Z Axis,", ",PROBE-1-45p3µM", ",1"] +
                                     [f"{x:g},{300 * np.exp(-((x - 480) / 25) ** 2):.4f}" for x in xs] +
                                     ["", "PROBE-1-45p3µM", "Ex. Wavelength (nm)               387.00"]),
                       encoding="utf-8", newline="")
        es = core.read_samples([ecl])
        assert len(es) == 1 and es[0]["name"] == "PROBE-1-45p3µM" and es[0]["ex"] == 387.0, es[0]["name"]
        report.append("overlay + fluorescence + Eclipse CSV: OK")

        # Fluoreszenz: Stokes-Verschiebung, A(λex), λabs-Label
        w.fl_state.update(stokes_label=True)
        files = w.fl_tab.export_to(out / "fluorescence_stokes")
        assert all(f.exists() for f in files), files
        inf = w.fl_tab.info
        exp_cm = 1e7 / inf["lam_abs"] - 1e7 / inf["lam_em"]
        assert inf["lam_em"] == 520 and abs(inf["stokes_cm"] - exp_cm) < 1e-6 and inf["stokes_cm"] > 0, inf
        assert inf["a_ex"] is not None and np.isfinite(inf["a_ex"]), inf
        assert TABS.fmt_stokes(1473.4) == "1470"
        figf = MplFigure(figsize=(8, 6), dpi=PREVIEW_DPI)
        FigureCanvasAgg(figf)
        ltxt = " ".join(a.get_text() for a in w.fl_tab.build(figf).ann_items.values())
        assert "abs" in ltxt and "\\tilde" in ltxt, ltxt
        w.fl_state.update(stokes_label=False)
        report.append("fluorescence Stokes shift + A(ex): OK")
        # Schultern, Zusatzbande, Cutoff und Ausblenden pro Bande
        w.open_csv(csv, ask=False)
        w.glob.update(shoulders=True, bands_add=True, bands="600")
        w.series["TEST-A"]["band_cutoff"] = {"484": 2.0}
        w.series["TEST-A"]["band_hidden"] = [600]
        cfg_t = w.make_cfg("TEST-A")
        _, _, res_t = core.prepare_series(cfg_t, w.groups["TEST-A"], "TEST-A")
        lams = {int(round(r["lam_requested"])): r for r in res_t}
        assert 600 in lams and lams[600].get("hidden") and lams[484]["cutoff"] == 2.0, sorted(lams)
        report.append("bands/shoulders/cutoff: OK")
        # einstellbare R²-Schwelle: 1.0 blendet alles aus, 0 zeigt alles mit Fit
        w.series["TEST-A"]["band_hidden"] = []
        for thr, expect_hidden in ((1.0, True), (0.0, False)):
            w.glob["min_r2"] = thr
            w.apply_globals()
            assert abs(w.sp_r2.value() - thr) < 1e-9
            w.collect_globals()
            _, _, res_r = core.prepare_series(w.make_cfg("TEST-A"), w.groups["TEST-A"], "TEST-A")
            fitted = [r for r in res_r if r["fit"]]
            assert fitted and all(bool(r.get("hidden")) == expect_hidden for r in fitted), (thr, res_r)
        w.glob["min_r2"] = 0.98
        w.apply_globals()
        report.append("min. R² threshold: OK")
        # Audit P1: Grenzfälle der ε-Auswertung (synthetisch)
        cfg_a = core.deep_merge(core.DEFAULTS, dict(conc_unit="mM", path_length_cm=1.0))
        assert core.parse_sample_name("probe_0,1mM") == ("probe", 0.1, "mM")         # Dezimalkomma
        assert core.parse_sample_name("probe_1e-3mM") == ("probe", 0.001, "mM")
        assert core.parse_sample_name("S-0p25mgml2") == ("S", 0.25, "mg/mL")         # Cary-Laufnummer dahinter
        x_a = np.arange(490.0, 511.0)
        sp_a = [dict(x=x_a, y=0.2 * c * np.exp(-((x_a - 503) / 2) ** 2), c=c) for c in (1.0, 2.0, 3.0, 4.0)]
        r_a = core.evaluate(dict(cfg_a, wavelengths_nm=[500], snap_window_nm=5), sp_a)[0]
        assert r_a["lam"] == 503 and abs(r_a["eps"] - 200) < 1e-6, (r_a["lam"], r_a.get("eps"))  # Maximum-Snap
        x_a = np.arange(400.0, 601.0)
        flat = [dict(x=x_a, y=np.full_like(x_a, 0.1), c=c) for c in (1.0, 2.0, 3.0, 4.0)]
        r_a = core.evaluate(dict(cfg_a, wavelengths_nm=[500]), flat)[0]
        assert r_a.get("hidden"), r_a                                    # R² undefiniert -> nicht ins Bild/SI
        prop = [dict(x=x_a, y=0.125 * c * np.ones_like(x_a), c=c) for c in (1.0, 2.0, 3.0, 4.0)]
        r_a = core.evaluate(dict(cfg_a, wavelengths_nm=[500], fit_intercept=False), prop)[0]
        assert core.fmt_ve(r_a["eps"], r_a["eps_err"], max_decimals=0)[1] != "0"   # nie „± 0“
        assert core.fmt_ve(9525.4, 455.15, max_decimals=0) == ("9525", "455")         # ε ganzzahlig wie publiziert
        assert core.fmt_ve(9.4802, 0.0123) == ("9.480", "0.012")                       # a nach Unsicherheit
        assert core.fmt_ve(12345.0, 0.0, max_decimals=0) == ("12345", "1")              # SE 0: eine Einheit
        assert core.parse_sample_name("probe-p5mM") is None                            # 0,5 oder 5? nicht raten
        same_c = [dict(x=x_a, y=np.full_like(x_a, 0.2), c=1.0) for _ in range(3)]
        r_a = core.evaluate(dict(cfg_a, wavelengths_nm=[500]), same_c)[0]
        assert r_a["fit"] is None and "eps" not in r_a, r_a              # gleiche c: kein ε statt Absturz
        x_a = np.arange(300.0, 801.0)
        edge = [dict(x=x_a, y=0.2 * c * np.exp(-0.5 * ((x_a - 600) / 20) ** 2), c=c) for c in (1.0, 2.0, 3.0, 4.0)]
        _, _, r_e = core.prepare_series(cfg_a, edge, "edge")             # Basislinienfenster am Messbeginn
        assert any(abs(r_["lam"] - 600) <= 2 for r_ in r_e), [r_["lam"] for r_ in r_e]
        two = [dict(x=x_a, y=0.2 * c * (np.exp(-0.5 * ((x_a - 300) / 15) ** 2) + np.exp(-0.5 * ((x_a - 600) / 60) ** 2)),
                    c=c) for c in (1.0, 2.0, 3.0, 4.0)]                  # Basislinienfenster 350-370 nm dazwischen
        lam_two = [r_["lam"] for r_ in core.prepare_series(cfg_a, two, "two")[2]]
        assert any(abs(l_ - 600) <= 2 for l_ in lam_two), lam_two      # Banden auf beiden Seiten des Fensters
        _, _, r_s = core.prepare_series(dict(cfg_a, peak_range_nm=[900, 950], find_shoulders=True), edge, "out")
        assert r_s == [], r_s                                             # Suchbereich ausserhalb: kein Absturz
        semi = out / "input" / "semicolon.csv"
        semi.write_text("\r\n".join(["S_0,1mM;;S_0,2mM;", "Wavelength (nm);Abs;Wavelength (nm);Abs;"] +
                                    [f"{lam};{0.1 * (lam - 395):.1f};{lam};{0.2 * (lam - 395):.1f};".replace(".", ",")
                                     for lam in range(400, 406)]), encoding="utf-8", newline="")
        quoted = out / "input" / "quoted.csv"                              # Namen mit Komma, korrekt zitiert
        quoted.write_text("\r\n".join(['"S_0,1mM",,"S_0,2mM",', "Wavelength (nm),Abs,Wavelength (nm),Abs,"] +
                                      [f"{lam},{0.01 * (lam - 395):.2f},{lam},{0.02 * (lam - 395):.2f},"
                                       for lam in range(400, 406)]), encoding="utf-8", newline="")
        assert [s_["parsed"] for s_ in core.read_samples([quoted])] == [("S", 0.1, "mM"), ("S", 0.2, "mM")]
        smp_ = core.read_samples([semi])
        assert [s_["parsed"] for s_ in smp_] == [("S", 0.1, "mM"), ("S", 0.2, "mM")], [s_["name"] for s_ in smp_]
        assert core.parse_sample_name("p_.5mM")[1] == 0.5 and core.parse_sample_name("p_5μM")[2] == "uM"
        ansi = out / "input" / "ansi.csv"                                 # Windows-Export (cp1252): µ = 0xB5
        ansi.write_bytes(("A-10 \xb5M,,A-5 \xb5M,\r\n" + "".join(f"{lam},0.{lam - 390},{lam},0.0{lam - 390},\r\n"
                                                           for lam in range(400, 406))).encode("cp1252"))
        assert [s_["parsed"] for s_ in core.read_samples([ansi])] == [("A", 10.0, "uM"), ("A", 5.0, "uM")]
        x_a = np.arange(200.0, 801.0)
        part = [dict(x=x_a, y=1.2 * c / 0.04 * np.exp(-((x_a - 400) / 30) ** 2), c=c) for c in (0.04, 0.02, 0.01, 0.005)]
        part[0] = dict(part[0], x=part[0]["x"][301:], y=part[0]["y"][301:])  # stärkste Probe erst ab 501 nm
        assert any(abs(p_ - 400) <= 2 for p_ in core.detect_peaks(part, dict(cfg_a, peak_range_nm=[230, None]))), \
            "Bande der verduennten Proben verloren"
        nan_bl = [dict(s_) for s_ in part[1:]]
        assert core.apply_baseline(nan_bl, dict(cfg_a, baseline_nm=[900, 950])) is None \
            and not any(np.isnan(s_["y"]).any() for s_ in nan_bl)        # Fenster ausserhalb: nicht korrigieren
        icpt = [dict(x=x_a, y=np.full_like(x_a, v), c=c) for c, v in zip((0.1, 0.2, 0.3, 0.4), (0.20, 0.28, 0.45, 0.50))]
        r_a = core.evaluate(dict(cfg_a, wavelengths_nm=[500]), icpt)[0]
        assert not r_a["warnings"], r_a["warnings"]                       # |b|/SE = 2.1 < t(0.975; 2) = 4.3
        res_x = [dict(lam=280.0, fit={"r2": 1.0}, eps=5000.0, eps_err=10.0),
                 dict(lam=450.0, fit={"r2": 1.0}, eps=3000.0, eps_err=10.0)]
        line_x = SI.uvvis_line(res_x, dict(cfg_a, xlim=[300, 600]))[0]
        assert "450 nm" in line_x and "280" not in line_x, line_x         # nur Banden des Bildes
        nir = bytearray([len(magic)]) + magic                            # Cary-NIR bis 1200 nm
        nir += bytes(0x400 - len(nir))
        blk = bytearray(256)
        blk[:10] = b"NIR-1mgml\x00"
        nir += blk + b"".join(struct.pack("<ff", x, 0.3) for x in np.arange(1200.0, 999.0, -1.0)) + bytes(64)
        (out / "input" / "nir.DSW").write_bytes(bytes(nir))
        s_nir = core.read_samples([out / "input" / "nir.DSW"])
        assert s_nir and s_nir[0]["x"].max() > 1199 and len(s_nir[0]["x"]) == 201, [len(s_["x"]) for s_ in s_nir]
        spc = bytearray(600)                                              # unbekannte SPC-Variante (Version 0)
        spc[4:8] = struct.pack("<i", 6)
        spc[8:24] = struct.pack("<dd", 400.0, 405.0)
        spc[544:568] = struct.pack("<6f", *[0.1] * 6)
        (out / "input" / "bad.spc").write_bytes(bytes(spc))
        try:
            core.read_samples([out / "input" / "bad.spc"])
            raise AssertionError("SPC ohne Versionskennung angenommen")
        except ValueError:
            pass
        spc[0:4] = bytes([0, 0x4B, 0, 0x80])                              # gültig: neues Format, y als float
        (out / "input" / "good.spc").write_bytes(bytes(spc))
        assert len(core.read_samples([out / "input" / "good.spc"])[0]["x"]) == 6
        report.append("eps edge cases (names, snap, undefined R2, SE 0, equal c, baseline at start, ';' CSV): OK")

        def pump(sec=0.4):
            import time
            t0 = time.time()
            while time.time() - t0 < sec:
                QApplication.processEvents()
                time.sleep(0.02)

        # TD-DFT mit Spin-Bahn-Kopplung (ORCA 6): Zustandsnamen wie "0-1.0A" dürfen nicht als Zahl gelten
        soc = out / "input" / "orca_soc.out"
        hdr = ["-" * 104, "      SOC CORRECTED ABSORPTION SPECTRUM VIA TRANSITION ELECTRIC DIPOLE MOMENTS", "-" * 104,
               "      Transition         Energy     Energy  Wavelength fosc(D2)      D2       |DX|      |DY|      |DZ|",
               "                          (eV)      (cm-1)    (nm)                 (au**2)    (au)      (au)      (au)",
               "-" * 104,
               "  0-1.0A  ->  1-3.0A    2.898266   23376.1   427.8   0.000000000   0.00000   0.00001   0.00000   0.00000",
               "  0-1.0A  -> 10-1.0A    4.004892   32301.6   309.6   0.120701528   1.23017   0.69013   0.85513   0.15048",
               ""]
        soc.write_text("\n".join(orca.read_text(encoding="utf-8").splitlines() + hdr), encoding="utf-8")
        from elucida.core import tddft as XX
        r_soc = XX.parse_tddft(soc)
        assert r_soc["variant"] == "soc" and abs(r_soc["E"][1] - 4.004892) < 1e-6 and abs(r_soc["f"][1] - 0.1207) < 1e-3, r_soc
        assert XX.parse_tddft(soc, "nosoc")["variant"] == "nosoc"
        report.append("TD-DFT SOC: OK")

        # Zuschnitt von Strukturen: senkrechte/waagerechte Endbindungen dürfen nicht abgeschnitten werden
        import pymupdf
        tdoc = pymupdf.open()
        pg = tdoc.new_page(width=200, height=200)
        pg.draw_line((50, 50), (150, 50), width=1)        # waagerecht
        pg.draw_line((100, 50), (100, 160), width=1)      # senkrecht, reicht am weitesten nach unten
        cr = IMG._content_rect(pg)
        assert cr.y1 >= 160 and cr.x0 <= 50 and cr.x1 >= 150, cr
        report.append("structure trimming: OK")

        # Einfärben von Strukturen (Overlay: Struktur pro Spektrum in Kurvenfarbe)
        ca = IMG.load_structure(cdxml, color="#e8231b")
        opaque = ca["rgba"][..., 3] > 200
        assert ca["pdf"] and abs(int(np.median(ca["rgba"][opaque][:, 0])) - 232) < 3, "Einfärben"
        report.append("structure colour: OK")

        # Neu platzierte Labels dürfen nicht auf gespeicherten landen (zwei Banden, eine gespeichert)
        cfg_l = w.make_cfg("TEST-A")
        cfg_l2, sp_l, res_l = core.prepare_series(cfg_l, w.groups["TEST-A"], "TEST-A")
        figl = MplFigure(figsize=cfg_l2["figsize_in"], dpi=PREVIEW_DPI)
        FigureCanvasAgg(figl)
        auto = core.build_figure(plt, cfg_l2, sp_l, res_l, {}, fig=figl).current_layout()
        lk = [k for k in auto if k.startswith("label_")]
        assert len(lk) >= 2, lk
        Fl = core.build_figure(plt, cfg_l2, sp_l, res_l, {lk[-1]: auto[lk[0]]}, fig=figl)
        rl = Fl._renderer()
        b0, b1 = (Fl.ann_items[k].get_window_extent(rl) for k in (lk[0], lk[-1]))
        assert not b0.overlaps(b1), "Label-Kollision"
        report.append("label collision: OK")
        # Inset nie über den Kurven: füllen die Spektren das ganze Bild, kommt es rechts neben die (schmalere) Achse
        sp_f = [dict(s_, y=np.full_like(np.asarray(s_["y"], float), 5.0)) for s_ in sp_l]
        Ff = core.build_figure(plt, cfg_l2, sp_f, res_l, {}, fig=figl)
        if "inset" in Ff.axes_items:
            assert Ff._inset_beside and Ff.axes_items["inset"].get_position().x0 >= Ff.ax.get_position().x1, \
                (Ff.axes_items["inset"].get_position(), Ff.ax.get_position())
            assert "inset" not in Ff.current_layout()
        report.append("inset beside the plot when no free space: OK")

        # Overlay-Optionen: Skalierung, Versatz, Linienart, Strich-Modi
        w.ov_state["entries"][1].update(scale=0.8, offset=0.2, ls="dotted", lw=2.0)
        for mode in ("max", "band", "faxis"):
            w.ov_state.update(mode="norm", stick_mode=mode, sticks=True)
            files = w.ov_tab.export_to(out / f"overlay_sticks_{mode}")
            assert all(f.exists() for f in files), files
        report.append("overlay options: OK")
        # Audit P1 (Feedback: Spektren ausserhalb des Bildes, Einzelnormierung): Overlay-Grenzfälle
        ov, ents0 = w.ov_tab, w.ov_state["entries"]
        calc_e = dict(ents0[1], scale=1.0, offset=0.0, ls=None, lw=None, own_norm=False, nlo=450, nhi=520)
        w.ov_state.update(entries=[calc_e], mode="norm", nlo=450, nhi=520, ymax=0.0, x_auto=True, sticks=True,
                          stick_mode="band", stack=False)
        cfg_o, curves_o = ov.figure_spec()[:2]
        top_o = max(np.nanmax(c_["y"][(c_["x"] >= cfg_o["xlim"][0]) & (c_["x"] <= cfg_o["xlim"][1])]) for c_ in curves_o)
        assert cfg_o["ylim"][1] >= top_o, (cfg_o["ylim"], top_o)          # 262-nm-Bande ~4x: nicht abschneiden
        w.ov_state["sticks"] = False
        assert ov.figure_spec()[0]["xlim"] == cfg_o["xlim"], "x-Bereich haengt an den Strichen"
        w.ov_state.update(sticks=True, nlo=700, nhi=800)                  # keine Bande im Fenster
        spec_o = ov.figure_spec()
        assert spec_o is None or max(np.nanmax(c_["y"]) for c_ in spec_o[1]) < 10, "Gauss-Auslaeufer normiert"
        orca0 = out / "input" / "orca_f0.out"                               # nur f = 0: keine NaN-Kurve
        orca0.write_text(orca.read_text(encoding="utf-8").replace("0.150000000", "0.000000000")
                         .replace("0.600000000", "0.000000000"), encoding="utf-8")
        w.ov_state["entries"] = [dict(calc_e, file=str(orca0))]
        spec_o = ov.figure_spec()
        assert spec_o is None or all(np.isfinite(c_["y"]).all() for c_ in spec_o[1]), "f = 0 -> NaN"
        w.ov_state.update(nlo=450, nhi=520)
        w.ov_state["entries"] = [dict(calc_e, label="a"), dict(calc_e, label="b", own_norm=True, nlo=400, nhi=300),
                                 dict(calc_e, label="c", sources=[{"file": calc_e["file"], "weight": 1.0},
                                                                  {"file": calc_e["file"], "weight": 1.0}])]
        ov.fill_table()
        ov.table.setCurrentCell(0, 1)
        ov.remove_selected()                                               # Details muessen nachziehen
        ov.sp_lw.setValue(1.5)
        e_b = w.ov_state["entries"][0]
        assert e_b["label"] == "b" and e_b["own_norm"] and (e_b["nlo"], e_b["nhi"]) == (400, 300), e_b
        ov.table.setCurrentCell(1, 1)
        ov.ed_weights.setText("0.7, 0.3")
        ov.detail_changed()
        assert [s_["weight"] for s_ in w.ov_state["entries"][1]["sources"]] == [0.7, 0.3]
        ov.table.setCurrentCell(0, 1)
        w.ov_state["entries"][0]["own_norm"] = False
        ov.show_details()
        ov.sp_elo.setValue(420)                                            # eigenes Fenster = eigene Normierung
        assert w.ov_state["entries"][0]["own_norm"] and w.ov_state["entries"][0]["nlo"] == 420
        ov.cmb_ymode.setCurrentIndex(ov.cmb_ymode.findData("norm"))
        ov.sp_ymax.setValue(1.2)
        ov.cmb_ymode.setCurrentIndex(ov.cmb_ymode.findData("eps"))       # ε-Skala: alte Obergrenze passt nicht
        assert ov.sp_ymax.value() == 0 and not ov.cb_own.isEnabled() and w.ov_state["mode"] == "eps"
        ov.cmb_ymode.setCurrentIndex(ov.cmb_ymode.findData("norm"))
        assert ov.cb_own.isEnabled()
        # f-Achse: ein Faktor – jeder Strich so hoch wie seine Bande, f rechts nur als zweite Skala
        w.ov_state.update(entries=[dict(calc_e)], stick_mode="faxis", mode="norm", nlo=450, nhi=520)
        cfg_f, curves_f, _, sticks_f = ov.figure_spec()
        lam_f, y_f = curves_f[0]["x"], curves_f[0]["y"]
        k_ = int(np.argmin(np.abs(sticks_f[0]["x"] - 484.0)))
        h_ = sticks_f[0]["h"][k_]
        assert abs(h_ * cfg_f["f_axis"] - 0.15) < 1e-6, h_ * cfg_f["f_axis"]          # f = 0.15 auf der f-Skala
        assert h_ <= np.interp(484.0, lam_f, y_f) + 1e-6 and h_ > 0.9 * np.interp(484.0, lam_f, y_f)
        w.ov_state["entries"] = [dict(calc_e), dict(calc_e, fwhm=0.5, label="breit")]
        assert "f_axis" in ov.figure_spec()[0]                            # normiert: beide Banden = 1, ein f-Maßstab
        w.ov_state["mode"] = "eps"
        assert "f_axis" not in ov.figure_spec()[0]                        # ε: verschiedene Maßstäbe, keine f-Achse
        w.ov_state["mode"] = "norm"
        # Absorbanz wie gemessen (y-Achse folgt den Daten, nicht immer ε)
        exp_e = dict(ents0[0], conc=None, own_norm=False, scale=1.0, offset=0.0, ls=None, lw=None)
        w.ov_state.update(entries=[exp_e, dict(calc_e)], mode="abs", stick_mode="band", baseline=False)
        cfg_a2, curves_a2 = ov.figure_spec()[:2]
        s_ab = ov.sample(exp_e["file"], exp_e["sample"])
        assert len(curves_a2) == 1 and np.allclose(curves_a2[0]["y"], s_ab["y"]) and "absorbance" in cfg_a2["ylabel"] \
            and "normalized" not in cfg_a2["ylabel"], cfg_a2["ylabel"]
        w.ov_state.update(mode="norm", baseline=True, stick_mode="band")
        w.ov_state["entries"] = ents0
        ov.fill_table()
        report.append("overlay: y range, x range without sticks, tail normalisation, own window, remove, weights: OK")

        # Overlay gestapelt: Ebenen in Listenreihenfolge, TD-DFT auf der Ebene der Messung, y ohne Zahlen
        ents = w.ov_state["entries"]
        ents.append(dict(ents[0], uid="stack3", label="exp 2", color="#e8231b"))
        ents[1]["stack_with_prev"] = True

        def ov_fig():
            figo = MplFigure(figsize=(8, 6), dpi=PREVIEW_DPI)
            FigureCanvasAgg(figo)
            Fo = w.ov_tab.build(figo)
            figo.canvas.draw()
            return Fo
        w.ov_state.update(mode="norm", stick_mode="faxis", stack=False, curve_labels=False, zero_lines=False)
        flat = ov_fig()
        w.ov_state.update(stack=True, stack_gap=1.2, y_axis="auto", curve_labels=True, zero_lines=True)
        stk = ov_fig()
        for i, dy in enumerate((1.2, 1.2, 0.0)):
            d = np.nanmedian(stk.curves[i]["y"] - flat.curves[i]["y"])
            assert abs(d - dy) < 1e-9, (i, d)
        assert stk.ax2 is None and not stk.ax.yaxis.get_major_ticks()[0].label1.get_visible()
        assert len(stk.curve_texts) == 2 and stk.legend is None, len(stk.curve_texts)
        files = w.ov_tab.export_to(out / "overlay_stacked")
        assert all(f.exists() for f in files), files
        ents.pop()
        ents[1].pop("stack_with_prev")
        w.ov_state.update(stack=False, curve_labels=False, zero_lines=False, y_axis="auto")
        report.append("overlay stacked: OK")

        # TD-DFT: Triplett-Zustände (ORCA 6 Bezeichnung, ORCA 4/5 eigener Block, Gaussian) und Kombinieren
        from elucida.core import tddft as XT
        sep = "-" * 100
        tab_head = [sep, "      ABSORPTION SPECTRUM VIA TRANSITION ELECTRIC DIPOLE MOMENTS", sep,
                    "   Transition    Energy   Energy  Wavelength fosc(D2)   D2   DX   DY   DZ",
                    "                  (eV)    (cm-1)    (nm)", sep]
        o6 = out / "input" / "orca6_trip.out"
        o6.write_text("\n".join(tab_head + [
            "  0-1A  ->  1-3A    2.500000   20164.1   495.9   0.000000000   0.00000   0.00000   0.00000   0.00000",
            "  0-1A  ->  1-1A    3.000000   24196.9   413.3   0.200000000   1.00000   1.00000   0.00000   0.00000",
            ""]), encoding="utf-8")
        o5 = out / "input" / "orca5_trip.out"
        o5.write_text("\n".join([
            "TD-DFT/TDA EXCITED STATES (TRIPLETS)", "------------------------------------", "",
            "STATE  1:  E=   0.091874 au      2.600 eV    20970.5 cm**-1 <S**2> =   2.000000", "",
            "STATE  2:  E=   0.110000 au      2.993 eV    24142.0 cm**-1 <S**2> =   2.000000", ""]
            + tab_head + ["   1   26617.0    375.7   0.300000000   1.00000   1.00000   0.00000   0.00000", ""]),
            encoding="utf-8")
        gx = out / "input" / "gauss_trip.log"
        gx.write_text("\n".join([
            " Excited State   1:     Triplet-A      2.7000 eV  459.20 nm  f=0.0000  <S**2>=2.000",
            " Excited State   2:     Singlet-A      3.2000 eV  387.45 nm  f=0.1500  <S**2>=0.000", ""]),
            encoding="utf-8")
        for pth, n_t, e_t in ((o6, 1, 2.5), (o5, 2, 2.6), (gx, 1, 2.7)):
            rt = XT.parse_tddft(pth)
            trips = [E for E, m in zip(rt["E"], rt["mult"]) if m == 3]
            assert len(trips) == n_t and abs(min(trips) - e_t) < 1e-6, (pth.name, rt)
            assert all(fv == 0 for fv, m in zip(rt["f"], rt["mult"]) if m == 3), pth.name
        n_before = len(ents)
        ents += [{"kind": "calc", "file": str(p_), "label": p_.stem, "color": "#1f9e4a", "visible": True,
                  "own_norm": False, "nlo": 450, "nhi": 520, "fwhm": 0.3, "shift": 0.0} for p_ in (o6, gx)]
        w.ov_tab.fill_table()
        for r_ in range(n_before):                      # nur die beiden neuen markieren
            ents[r_]["visible"] = r_ == 0
        w.ov_tab.table.clearSelection()
        w.ov_tab.table.setSelectionMode(QAbstractItemView.MultiSelection)
        assert not w.ov_tab.btn_combine.isEnabled()                 # nichts markiert: „Kombinieren“ aus
        for r_ in (n_before, n_before + 1):
            w.ov_tab.table.selectRow(r_)
        assert w.ov_tab.btn_combine.isEnabled()                     # zwei TD-DFT-Einträge markiert
        w.ov_tab.combine_selected()
        comb = ents[-1]
        assert len(comb["sources"]) == 2 and not ents[n_before]["visible"], comb
        E_c, f_c, m_c = w.ov_tab.calc_data(comb)
        assert len(E_c) == 4 and m_c.count(3) == 2 and abs(f_c.sum() - 0.35) < 1e-9, (E_c, f_c, m_c)
        comb["sources"][1]["weight"] = 0.5
        assert abs(w.ov_tab.calc_data(comb)[1].sum() - 0.275) < 1e-9
        w.ov_state.update(mode="norm", triplets=True, sticks=True, stick_mode="band")
        _, _, _, stx = w.ov_tab.figure_spec()
        dashed = np.concatenate([s_["x"] for s_ in stx if s_.get("ls") == "dashed"])
        assert np.allclose(sorted(dashed), sorted(1e7 / (np.array([2.5, 2.7]) * XT.EV_TO_CM))), dashed
        files = w.ov_tab.export_to(out / "overlay_triplets")
        assert all(f.exists() for f in files), files
        assert any(c_ is comb["sources"][0] for c_, _, _ in PRJ._path_slots({"overlay": w.ov_state})), \
            "kombinierte Quellen fehlen im Projekt"
        del ents[n_before:]
        ents[1]["visible"] = True
        w.ov_state.update(triplets=False)
        w.ov_tab.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        w.ov_tab.fill_table()
        e_cc = next(e_ for e_ in w.ov_state["entries"] if e_["kind"] == "calc")
        e_cc.update(color_c=1.0, color_d=1.0)                  # Lösungsfarbe aus dem TD-DFT-Spektrum (1 mM, 1 cm)
        hx1_, _ = w.ov_tab.calc_color(e_cc)
        e_cc["color_c"] = 1e4
        hx2_, _ = w.ov_tab.calc_color(e_cc)
        lum_ = lambda h_: sum(int(h_[i_:i_ + 2], 16) for i_ in (1, 3, 5))      # noqa: E731
        assert re.fullmatch(r"#[0-9A-F]{6}", hx1_) and lum_(hx2_) <= lum_(hx1_), (hx1_, hx2_)
        e_cc["color_c"] = 1.0
        w.ov_tab.update_calc_color(e_cc)
        assert hx1_ in w.ov_tab.lbl_ccol_info.text(), w.ov_tab.lbl_ccol_info.text()
        report.append("TD-DFT triplets + combine: OK")

        # Export ins Projekt inkl. Excel-Report
        w.open_csv(csv, ask=False)
        w.series["TEST-A"]["mw"] = 95.1
        w.set_structure(str(cdxml), "TEST-A", render=False, sync=True, autofill_mw=False)
        w.render()
        pump()
        files = w.export_all()
        rep = proj / "tables" / "TEST-A_UV-Vis_report.xlsx"
        assert rep in files and rep.exists(), files
        assert (proj / "figures" / "TEST-A_UV-Vis.pdf").exists() and not (proj / "exports").exists()
        meth_ = (proj / "SI" / "TEST-A_UV-Vis_methods.txt").read_text(encoding="utf-8")
        assert "ordinary least squares" in meth_ and "95% confidence interval" in meth_ and " nm: " in meth_ \
            and "M = 95.10 g mol-1" in meth_, meth_
        assert meth_.split("\n\n")[0] in w.eps_methods_text(), "SI-Text Export/Knopf verschieden"
        from openpyxl import load_workbook
        wbk = load_workbook(rep)
        assert len(wbk.sheetnames) >= 3 and any("nm" in n for n in wbk.sheetnames), wbk.sheetnames
        fmts_x = {c_.number_format for ws_ in wbk.worksheets for row_ in ws_.iter_rows() for c_ in row_}
        assert "0.0" not in fmts_x, "Excel: SE/KI von ε mit anderer Rundung als Tabelle/SI"
        report.append("export into project + Excel report: OK")

        # Textformat: λ/ε in Symbolschrift, Labels fett/farbig, Rahmen um Plot und Struktur
        w.text_fmt = dict(core.TEXT_DEFAULTS, bold=True, color="#1764e8", label_pt=12.0,
                          axes_frame=True, struct_frame=True)
        w.render()
        assert w.F.ax.spines["top"].get_visible(), "Rahmen oben fehlt"
        ann = next(a for k, a in w.F.ann_items.items() if k.startswith("label_"))
        assert "\\mathsf{λ}" in ann.get_text() and ann.get_fontweight() == "bold" and ann.get_fontsize() == 12.0
        sa = w.F.axes_items["image_struktur"]
        assert sa._frame is not None
        rf = w.F._renderer()
        fb, ib = sa._frame.get_window_extent(rf), sa.get_window_extent(rf)
        assert fb.x0 < ib.x0 and fb.y1 > ib.y1, "Rahmen liegt nicht außen um die Struktur"
        for k, b in [(k, a.get_window_extent(rf)) for k, a in w.F.ann_items.items()]:
            assert not fb.overlaps(b), f"Strukturrahmen überdeckt {k}"
        files = w.export_all(out / "textformat")
        pdf = next(f for f in files if f.suffix == ".pdf")
        fonts = " ".join(f[3] for f in pymupdf.open(str(pdf))[0].get_fonts())
        sym_font = core.first_font(core.DEFAULTS["symbol_font"], "DejaVu Serif").replace(" ", "")
        assert sym_font.lower() in fonts.lower().replace("-", ""), fonts
        svg = next(f for f in files if f.suffix == ".svg").read_text(encoding="utf-8")
        assert "structure_0" in svg, "SVG ohne Struktur"
        w.fl_tab.export_to(out / "textformat_fluo")
        w.text_fmt = dict(core.TEXT_DEFAULTS)
        w.render()
        report.append("text format + frames: OK")

        # Hover: Maus über einem Messpunkt -> λ und A in der Statusleiste; Doppelklick -> Bande
        from matplotlib.backend_bases import MouseEvent
        w.glob["bands"] = ""
        w.apply_globals()
        w.render()
        w.canvas.draw()
        hv, ax = w.hover, w.F.ax
        def peak(c):                                   # Bande bei 484 nm, nur Kurven mit Peak im Bild
            xs = np.asarray(c["x"])
            i = int(np.nanargmax(np.where((xs > 400) & (xs < 560), c["y"], -np.inf)))
            return i, float(c["y"][i])
        top = ax.get_ylim()[1]
        sp0 = max((c for c in hv.curves if peak(c)[1] < 0.9 * top), key=lambda c: peak(c)[1])
        i0 = peak(sp0)[0]
        px, py = ax.transData.transform((sp0["x"][i0], sp0["y"][i0]))
        w.canvas.callbacks.process("motion_notify_event", MouseEvent("motion_notify_event", w.canvas, px, py))
        msg = w.statusBar().currentMessage()
        # Mausereignisse haben ganzzahlige Pixel, Messpunkte liegen in der Vorschau < 1 px auseinander:
        # Treffer darf der Nachbarpunkt sein. Geprüft wird, dass die Anzeige genau den markierten Punkt zeigt.
        assert hv.hit is not None, msg
        hc, hi = hv.hit
        # erlaubt: ein Messpunkt Abstand plus ein Pixel in Datenkoordinaten (Vorschaugröße je nach System/Schrift)
        step = float(np.median(np.abs(np.diff(np.asarray(sp0["x"], float)))))
        nm_per_px = abs(ax.transData.inverted().transform((px + 1, py))[0] - sp0["x"][i0])
        assert abs(hc["x"][hi] - sp0["x"][i0]) <= step + nm_per_px + 1e-9, (msg, hc["x"][hi], step, nm_per_px)
        assert f"λ = {hc['x'][hi]:.1f} nm" in msg and f"A = {hc['y'][hi]:.4f}" in msg, msg
        assert list(hv.marker.get_xdata()) == [hc["x"][hi]], hv.marker.get_xdata()
        assert hv.marker.get_animated(), "Marker darf nicht exportiert werden"
        w.canvas.callbacks.process("button_press_event",
                                   MouseEvent("button_press_event", w.canvas, px, py, button=1, dblclick=True))
        lam0 = round(float(hc["x"][hi]))
        assert w.glob["bands"] == f"{lam0:g}", w.glob["bands"]
        w.pick_band(lam0 + 1)                       # erneut (±2 nm) -> wieder entfernt
        assert w.glob["bands"] == "", w.glob["bands"]
        w.render()
        report.append("hover readout + band picking: OK")

        # Ziehen mit Blitting: Inset verschieben, beim Loslassen wieder normal gezeichnet und gespeichert
        w.render()
        w.canvas.draw()
        key_, a_ = next((k, a) for k, a in w.F.axes_items.items() if a.get_visible())
        p0 = a_.get_position().frozen()
        cx, cy = a_.get_window_extent().x0 + 5, a_.get_window_extent().y0 + 5
        proc = w.canvas.callbacks.process
        proc("button_press_event", MouseEvent("button_press_event", w.canvas, cx, cy, button=1))
        assert a_.get_animated() and w.dragger._bg is not None, "Blitting nicht aktiv"
        proc("motion_notify_event", MouseEvent("motion_notify_event", w.canvas, cx + 30, cy + 20, button=1))
        # nach dem Loslassen muss ganz neu gezeichnet werden (Blitting-Hintergründe, auch der der Hover-Anzeige, sind
        # ohne das gezogene Element entstanden); geprüft wird die Anforderung, nicht ihr Zeitpunkt (CI ohne Fenster)
        idle_ = []
        w.canvas.draw_idle = lambda *a_k: idle_.append(1)
        proc("button_release_event", MouseEvent("button_release_event", w.canvas, cx + 30, cy + 20, button=1))
        del w.canvas.draw_idle
        assert idle_, "nach dem Ziehen kein Neuzeichnen angefordert"
        assert w.hover.bg is None, "Hover-Hintergrund aus der Zeit des Ziehens noch gültig"
        pump()
        p1 = a_.get_position()
        assert p1.x0 > p0.x0 and p1.y0 > p0.y0 and not a_.get_animated() and w.dragger._bg is None, (p0, p1)
        assert all(getattr(ann, "_draggable", None) is None or ann._draggable._use_blit
                   for ann in w.F.ann_items.values()), "Labels ohne Blitting"
        w.series[w.current]["layout"] = {}
        w.render()
        report.append("dragging with blitting: OK")
        w.glob["label_index"] = True                            # λ₁/ε₁ … aufsteigend nach λ
        w.apply_globals()
        w.render()
        lt_ = sorted((float(k_.split("_")[1]), a_.get_text()) for k_, a_ in w.F.ann_items.items())
        assert lt_ and all(t_.count(f"_{{{i_}}}") == 2 for i_, (_, t_) in enumerate(lt_, 1)), lt_
        w.glob["label_index"] = False
        w.apply_globals()
        w.render()
        assert not any("_{" in a_.get_text() for a_ in w.F.ann_items.values())
        report.append("label indices: OK")

        # Cyclovoltammetrie: Gamry-Dateien (beide Varianten), Redoxpaare, interne Referenz, Export
        def cv_sweep(e0, ip, lo, hi, step=0.002, ref=None):
            """Modell-CV: Hinlauf lo->hi, Rücklauf hi->lo; Peaks bei E0 ± 30 mV (+ optional Referenzpaar)."""
            up = np.arange(lo, hi + step / 2, step)
            E_ = np.concatenate([up, up[::-1]])
            fw = np.r_[np.ones(len(up)), np.zeros(len(up))].astype(bool)

            def wave(e_c, i_p):
                pa = i_p * np.exp(-((E_ - (e_c + 0.03)) / 0.045) ** 2) + 0.3 * i_p / (1 + np.exp(-(E_ - e_c - 0.03) / 0.02))
                pc = -0.9 * i_p * np.exp(-((E_ - (e_c - 0.03)) / 0.045) ** 2)
                return np.where(fw, pa, pc + 0.3 * i_p * np.exp(-(hi - E_) / 0.1))
            I_ = wave(e0, ip) + (wave(ref, 0.8 * ip) if ref is not None else 0)
            return E_, I_

        def write_dta(path, cycles, rate, cycle_col=False, times=None):
            head = ["EXPLAIN", "TAG\tCV", "DATE\tLABEL\t1/1/2026\tDate", "TIME\tLABEL\t12:00:00\tTime",
                    f"SCANRATE\tQUANT\t{rate:.5E}\t&Scan Rate (mV/s)", "AREA\tQUANT\t7.10000E-002\tElectrode &Area (cm^2)",
                    "CONVENTION\tIQUANT\t1\tCurrent Convention"]
            cols = "\tPt\tT\tVf\tIm\tVu" + ("\tCycle" if cycle_col else "")
            units = "\t#\ts\tV vs. Ref.\tA\tV" + ("\t#" if cycle_col else "")
            body = []
            if cycle_col:
                body += [f"CURVE\tTABLE\t{sum(len(c[0]) for c in cycles)}", cols, units]
            pt = 0
            for n, (E_, I_) in enumerate(cycles):
                if not cycle_col:
                    body += [f"CURVE{n + 1}\tTABLE", cols, units]
                for e_, i_ in zip(E_, I_):
                    t_ = times[pt] if times is not None else pt * 0.02
                    body.append(f"\t{pt}\t{t_:.2f}\t{e_:.5E}\t{i_ * 1e-6:.5E}\t0" + (f"\t{n}" if cycle_col else ""))
                    pt += 1
            path.write_text("\r\n".join(head + body) + "\r\n", encoding="latin-1", newline="")
        cvdir = out / "input" / "cv"
        cvdir.mkdir(parents=True, exist_ok=True)
        cv_files = []
        for rate, ip in ((50, 40.0), (100, 56.0), (200, 80.0)):
            p_ = cvdir / f"cv_{rate}.DTA"
            write_dta(p_, [cv_sweep(-1.02, ip, -1.4, -0.5)], rate)
            cv_files.append(p_)
        ref_file = cvdir / "cv_ref.DTA"                  # zwei Zyklen, Variante mit Spalte 'Cycle'
        write_dta(ref_file, [cv_sweep(-1.02, 80.0, -1.4, 0.9, ref=0.55)] * 2, 200, cycle_col=True)
        d_ = CVMOD.read_gamry(ref_file)
        assert len(d_["cycles"]) == 2 and d_["meta"]["scanrate"] == 200 and d_["meta"]["area"] == 0.071, d_["meta"]
        # weites Fenster, Start innerhalb (Feedback „abgeschnitten“): Gamry legt den Rückweg zu VFINAL in eine eigene
        # Tabelle; mit CYCLES 1 gehört sie zum Zyklus (beide Varianten)
        up_w = np.arange(-1.2, 2.0 + 1e-9, 0.004)
        leg1_ = np.concatenate([up_w, np.arange(2.0, -3.0 - 1e-9, -0.004)[1:]])
        leg2_ = np.arange(-3.0 + 0.004, -1.2 + 1e-9, 0.004)
        for cc_ in (False, True):
            p_ = cvdir / f"cv_wide_{int(cc_)}.DTA"
            write_dta(p_, [(leg1_, np.sin(leg1_)), (leg2_, np.sin(leg2_))], 200, cycle_col=cc_)
            p_.write_text(p_.read_text(encoding="latin-1").replace(
                "EXPLAIN", "EXPLAIN\r\nCYCLES\tIQUANT\t1\tC&ycles (#)", 1), encoding="latin-1", newline="")
            cw_ = CVMOD.read_gamry(p_)["cycles"]
            assert len(cw_) == 1 and len(cw_[0]["E"]) == len(leg1_) + len(leg2_), [len(c_["E"]) for c_ in cw_]
            assert abs(cw_[0]["E"].min() + 3.0) < 1e-6 and abs(cw_[0]["E"].max() - 2.0) < 1e-6
            assert abs(cw_[0]["E"][-1] + 1.2) < 1e-6, cw_[0]["E"][-1]        # endet wieder am Start
        report.append("CV: weites Fenster, Rueckweg zu VFINAL im Zyklus (Gamry CURVE2 / Cycle 1): OK")
        c_ = CVMOD.read_gamry(cv_files[0])["cycles"][0]
        sug = CVMOD.suggest_couples(c_["E"], c_["I"])
        assert len(sug) == 1 and sug[0][0] < -1.02 < sug[0][1], sug
        r_ = CVMOD.analyze_couple(c_["E"], c_["I"], *sug[0])
        assert abs(r_["e12"] + 1.02) < 0.003 and abs(r_["dep"] - 0.06) < 0.006 and not r_["edge"], r_
        st_cv = w.cv_state
        st_cv["files"] = [{"file": str(p_), "uid": f"cv{i}", "visible": True, "color": "#000000",
                           "label": f"{r} mV/s"} for i, (p_, r) in enumerate(zip(cv_files, (50, 100, 200)))]
        st_cv["files"].append({"file": str(ref_file), "uid": "cvref", "visible": False, "color": "#000000",
                               "label": "ref"})
        st_cv.update(couples=[list(sug[0])], ref_mode="internal", ref_uid="cvref", ref_lo=0.3, ref_hi=0.8,
                     ref_compound="ac2fc", ref_value=0.49, cycle=2)
        w.cv_tab.fill_table()
        files = w.cv_tab.export_to(out / "cv_test")
        assert all(f.exists() for f in files), files
        res = w.cv_tab.results
        assert len(res) == 3 and all(abs(r["e12"] - (-1.02 - (0.55 - 0.49))) < 0.004 for r in res), res
        assert "Ac₂Fc" in w.cv_tab.note and "+0.49 V vs. Fc" in w.cv_tab.note, w.cv_tab.note
        b_cv = w.cv_tab.table.cellWidget(0, 2)                # Farbverlauf: Feld zeigt die Farbe im Bild, gesperrt
        assert st_cv["gradient"] and not b_cv.isEnabled() and \
            w.cv_tab.grad_colors[st_cv["files"][0]["uid"]] in b_cv.styleSheet(), b_cv.styleSheet()
        rs = CVMOD.randles_sevcik([50, 100, 200], [r["ipa"] * 1e6 for r in res])
        assert rs and rs[2] > 0.99, rs
        csv_txt = files[-1].read_text(encoding="utf-8-sig")
        assert csv_txt.count("\n") == 4 and "E_half_V" in csv_txt and "charging_step_uA" in csv_txt, csv_txt
        meth = files[-2].read_text(encoding="utf-8")
        assert "Nicholson" in meth and "diacetylferrocene" in meth and "95% confidence" in meth \
            and "Couple 1: E1/2 = −1.08 V vs. Fc0/+" in meth, meth
        assert any(c_ is st_cv["files"][0] for c_, _, _ in PRJ._path_slots({"cv": st_cv})), "CV-Dateien fehlen im Projekt"
        # „Alles auf Standard“: Einstellungen zurück; Messungen, Redoxpaare, Referenz, Unterschrift bleiben
        snap_cv = json.loads(json.dumps({k_: st_cv[k_] for k_ in ("files", "couples", "ref_mode", "ref_value", "cycle",
                                                                  "gradient", "baseline")}))
        st_cv.update(gradient=False, baseline=False, stack=True, cap_temp="−78 °C")
        w.cv_tab.reset_defaults()
        assert st_cv["gradient"] and st_cv["baseline"] and not st_cv["stack"] and st_cv["cycle"] == 1
        assert st_cv["cap_temp"] == "−78 °C" and all(st_cv[k_] == snap_cv[k_] for k_ in ("files", "couples", "ref_mode",
                                                                                     "ref_value"))
        assert w.cv_tab.cb_gradient.isChecked() and w.main_tabs.currentWidget() is not None
        st_cv.update(cycle=snap_cv["cycle"], cap_temp="rt")
        w.cv_tab.fill_table()
        report.append("cyclic voltammetry: OK")

        # SI-Export: Abbildungen aller Tabs mit Daten (ε je Serie, CV …), zwei je Seite
        fb_ = w.si_figure_blocks(out / "si_figs", "TEST-1")
        caps_ = [b_[1] for b_ in fb_ if b_[0] == "figcap"]
        imgs_ = [b_ for b_ in fb_ if b_[0] == "image"]
        assert any(c_.startswith("UV-Vis spectra of <b>TEST-1</b>") and "cm quartz cuvettes." in c_ for c_ in caps_), caps_
        assert any(c_.startswith("Cyclic voltammograms of <b>TEST-1</b>") for c_ in caps_) and len(imgs_) == len(caps_), caps_
        assert all(Path(b_[1]).exists() and Path(b_[4]).exists() and 0 < b_[2] <= 16.0 for b_ in imgs_), imgs_
        assert fb_.count(("pagebreak",)) == (len(imgs_) - 1) // 2 and w.si_available()["figs"]
        # Methodik der Fits im Word-Export: Texte wie „SI-Text kopieren“, je Technik ein Absatz mit fettem Titel
        mb_ = w.si_methods_blocks()
        assert mb_[0] == ("heading", "Data processing and fitting") and w.si_available()["methods"], mb_[:1]
        mtx_ = [b_[1] for b_ in mb_[1:]]
        assert any(t_.startswith("<b>UV-Vis absorption spectroscopy.</b> ") for t_ in mtx_), mtx_
        assert any(t_.startswith("<b>Cyclic voltammetry.</b> ") and "Nicholson" in t_ for t_ in mtx_), mtx_
        assert any("E<sub>1/2</sub> = −1.08 V" in t_ for t_ in mtx_) and "cm<sup>-1</sup>" in "".join(mtx_), mtx_
        assert SI.methods_html("A_red, g_lit, E1/2,app, M-1 cm-1") == \
            "A<sub>red</sub>, g<sub>lit</sub>, E<sub>1/2</sub>,app, M<sup>-1</sup> cm<sup>-1</sup>"
        b_m, _ = w.si_blocks(["methods"])
        assert ("heading", "Data processing and fitting") in b_m and not any(b_[0] == "image" for b_ in b_m)
        mdoc_ = SI.write_docx(out / "si_methods.docx", b_m)
        with __import__("zipfile").ZipFile(mdoc_) as z_:
            mx_ = z_.read("word/document.xml").decode("utf-8")
            assert "Data processing and fitting" in mx_ and 'w:ascii="Arial"' in mx_
        SI.set_font("Times New Roman")                                   # Schriftart wählbar, Größe bleibt 11 pt
        with __import__("zipfile").ZipFile(SI.write_docx(out / "si_font.docx", b_m)) as z_:
            mx_ = z_.read("word/document.xml").decode("utf-8")
            assert 'w:ascii="Times New Roman"' in mx_ and 'w:ascii="Arial"' not in mx_ and 'w:sz w:val="22"' in mx_
            assert 'w:ascii="Times New Roman"' in z_.read("word/styles.xml").decode("utf-8")
        SI.set_font("Arial")
        cap2_ = "Cyclic voltammograms of <b>X</b> (1.8±0.1 mg mL-1) in MeCN (0.1 M n-Bu4NPF6) at rt; " * 2
        tp_ = SI.two_per_page([("char", "Setup text. " * 40)] + [("image", "a.png", 16.0, 11.0, None),
                                                                   ("figcap", cap2_)] * 3)
        assert [b_[0] for b_ in tp_] == ["char", "image", "figcap", "image", "figcap", "pagebreak", "image", "figcap"]
        h_ = [b_[3] for b_ in tp_ if b_[0] == "image"]
        assert h_[0] == h_[1] < h_[2] == 11.0 and abs(tp_[1][2] / tp_[1][3] - 16 / 11) < 1e-9, h_   # zwei je Seite
        w.cv_state.update(cap_conc="1.8±0.1 mg mL-1", cap_el="0.1 M n-Bu4NPF6")        # CV-Unterschrift wie in der SI
        cap_cv = w.cv_tab.si_caption("TEST-1")
        assert cap_cv.startswith("Cyclic voltammograms of <b>TEST-1</b> (1.8±0.1 mg mL<sup>-1</sup>)") \
            and "(0.1 M n-Bu<sub>4</sub>NPF<sub>6</sub>) at rt" in cap_cv and cap_cv.endswith(".") \
            and "potentials referenced against Fc<sup>0/+</sup> using diacetylferrocene (E<sub>1/2</sub> = +0.49 V vs. " \
            "Fc<sup>0/+</sup>) as internal standard" in cap_cv, cap_cv
        assert SI.chem_html("CH2Cl2, AgNO3 (0.01 M), 7.1 mm², 3 min") == \
            "CH<sub>2</sub>Cl<sub>2</sub>, AgNO<sub>3</sub> (0.01 M), 7.1 mm², 3 min"
        w.settings.setValue("si_setup/cv", "CVs were measured in CH3CN.")              # Aufbau-Absatz vor den CVs
        fb_ = w.si_figure_blocks(out / "si_figs", "TEST-1")
        assert ("char", "CVs were measured in CH<sub>3</sub>CN.") in fb_ and \
            ("figcap", cap_cv) in fb_, [b_ for b_ in fb_ if b_[0] != "image"]
        w.settings.remove("si_setup/cv")
        w.cv_state.update(cap_conc="", cap_el="")
        report.append("SI figures from all tabs (CV caption, setup text): OK")

        # Tab-Reihenfolge wie die SI; offener Tab über den Namen (alte Projekte: Position in der alten Reihenfolge)
        assert [w.main_tabs.widget(i) for i in range(5)] == [w.syn_tab, w.nmr_tab, w.irms_tab, w.hrms_tab, w.eps_page]
        st_ = w.get_state()
        st_["tab"] = 3                                                   # 1.23: Position 3 = CV
        w.set_state(st_)
        assert w.main_tabs.currentWidget() is w.cv_tab and w.get_state()["tab"] == "cv"
        # Textformat je Tab: eigenes Format nur für CV, der Rest folgt dem Standard
        w.tab_fmt["cv"] = dict(core.TEXT_DEFAULTS, bold=True)
        assert w.cv_tab.base_cfg()["labels"]["bold"] and not w.sec_tab.base_cfg()["labels"]["bold"]
        assert w.get_state()["text_tabs"]["cv"]["bold"] and w.text_tab()[0] == "cv"
        dlg_ = TextFormatDialog(w, w.text_fmt, w.tab_fmt["cv"], w.cv_tab.TEXT_OPTS, "cv", "CV")
        assert dlg_.cmb_scope.currentData() == "own" and dlg_.cb_bold.isChecked() and dlg_.cb_greek.isHidden()
        dlg_.cmb_scope.setCurrentIndex(0)                                  # zurück auf „alle Tabs“
        assert not dlg_.cb_bold.isChecked() and dlg_.results()[1] is None
        w.tab_fmt.clear()
        w.main_tabs.setCurrentWidget(w.eps_page)
        report.append("tab order + open tab by name, text format per tab: OK")

        # CV-Basislinie: Ladestrom ±C·v + linearer Untergrund dürfen i_p und i_r/i_f nicht verändern
        E_, I_ = cv_sweep(-1.02, 50.0, -1.4, -0.5)
        dEs_ = np.gradient(E_)
        I_c = I_ + np.where(dEs_ > 0, 6.0, -6.0) + 15.0 * (E_ + 1.0)
        win_ = sug[0]
        r0_ = CVMOD.analyze_couple(E_, I_, *win_, use_baseline=True)
        r1_ = CVMOD.analyze_couple(E_, I_c, *win_, use_baseline=True)
        r2_ = CVMOD.analyze_couple(E_, I_c, *win_)
        assert r0_["bl"] and r1_["bl"] and not r1_["bl_fail"], r1_
        assert abs(r1_["bl"]["step"] - 12.0) < 0.5, r1_["bl"]
        for k_ in ("ipa", "ipc", "ratio"):
            assert abs(r1_[k_] - r0_[k_]) < 0.01 * abs(r0_[k_]) + 0.05, (k_, r0_[k_], r1_[k_])
        assert abs(r2_["ipa"] - r1_["ipa"]) > 4, (r2_["ipa"], r1_["ipa"])        # ohne Basislinie verfälscht
        assert abs(r1_["e12"] - r2_["e12"]) < 1e-12                               # Potentiale unverändert
        # kathodisch zuerst (Spiegelbild): gleiche Korrektur mit umgekehrtem Vorzeichen
        r3_ = CVMOD.analyze_couple(-E_, -I_c, -win_[1], -win_[0], use_baseline=True)
        assert abs(r3_["ratio"] - r1_["ratio"]) < 1e-6 and abs(r3_["bl"]["step"] - r1_["bl"]["step"]) < 1e-6, r3_
        # zu wenig Vorlauf im Fenster: Fallback von null, markiert
        r4_ = CVMOD.analyze_couple(E_, I_c, -1.08, -0.8, use_baseline=True)
        assert r4_ is not None and r4_["bl_fail"] and r4_["ipa"] == r4_["ipa_raw"], r4_
        # breite Welle (iR-Abfall, wie THF mit ΔE_p ≈ 240 mV): Basislinie endet vor dem Anstieg, nicht im Fuß
        upw_ = np.arange(-1.6, -0.3, 0.002)
        Ew_ = np.concatenate([upw_, upw_[::-1]])
        fww_ = np.r_[np.ones(len(upw_)), np.zeros(len(upw_))].astype(bool)
        Iw_ = np.where(fww_, 100 * np.exp(-((Ew_ + 0.80) / 0.12) ** 2), -80 * np.exp(-((Ew_ + 1.04) / 0.12) ** 2))
        rw_ = CVMOD.analyze_couple(Ew_, Iw_, -1.5, -0.5, use_baseline=True)
        assert rw_["bl"] and abs(rw_["ipa"] - rw_["ipa_raw"]) < 1.0 and rw_["bl"]["e1"] < -1.0, rw_["bl"]   # alte Regel: -0.95 (21 % im Fuß)
        # Nicholson nur ohne weiteren Prozess zwischen Hinpeak und Umkehrpotential (hier: Referenzpaar bei +0.55 V)
        cr_ = CVMOD.read_gamry(ref_file)["cycles"][0]
        rr1_ = CVMOD.analyze_couple(cr_["E"], cr_["I"], *win_)
        rr2_ = CVMOD.analyze_couple(cr_["E"], cr_["I"], 0.3, 0.8)
        assert not rr1_["ratio_ok"] and not np.isfinite(rr1_["ratio"]) and rr2_["ratio_ok"], (rr1_, rr2_)
        assert r1_["ratio_ok"] and np.isfinite(r1_["ratio"])
        f_ = CVMOD.rate_fit([1, 2, 3, 4], [2.1, 3.9, 6.2, 7.9])
        # von Hand: RSS = 0.063, SE = √(0.063/2/5) = 0.079373, t(0.975; 2) = 4.302653
        assert abs(f_["slope"] - 1.97) < 1e-9 and abs(f_["ci"] - 4.302653 * 0.0793725) < 1e-5, f_
        assert CVMOD.fmt_ci(f_["slope"], f_["ci"]) == "2.0 ± 0.3", CVMOD.fmt_ci(f_["slope"], f_["ci"])
        ll_ = CVMOD.loglog_fit([50, 100, 200, 400], [10 * (v / 50) ** 0.5 for v in (50, 100, 200, 400)])
        assert abs(ll_["slope"] - 0.5) < 1e-9, ll_
        report.append("CV baseline, Nicholson with baseline, fits with CI: OK")

        # CV-Ausbau: Gruppe + eigene Referenz je Messung (zweites "Lösungsmittel"), Stapeln, E½ je Kurve,
        # Auftragungen gegen die Scanrate (ΔE_p, Randles-Ševčík, Trumpet)
        t_cv = w.cv_tab
        for e_ in st_cv["files"]:
            e_["group"] = "A"
        for r_ in (50, 100):                            # Gruppe B: E½ −0.92 V, eigenes Ac₂Fc bei +0.75 V (gemessen)
            p_ = cvdir / f"cvB_{r_}.DTA"
            write_dta(p_, [cv_sweep(-0.92, 60.0 * (r_ / 50) ** 0.5, -1.4, 1.1, ref=0.75)], r_)
            st_cv["files"].append({"file": str(p_), "uid": f"cvB{r_}", "visible": True, "color": "#000000",
                                   "label": f"B {r_} mV/s", "group": "B"})
        for e_ in st_cv["files"][-2:]:
            e_["ref"] = {"mode": "internal", "uid": e_["uid"], "lo": 0.5, "hi": 1.0, "compound": "ac2fc", "value": 0.49}
        st_cv.update(plot="cv", stack=True, e12_each=True, show_e12=True)
        t_cv.fill_table()
        spec_ = t_cv.figure_spec()
        resB = [r_ for r_ in t_cv.results if r_["group"] == "B"]
        assert len(resB) == 2 and all(abs(r_["e12"] - (-0.92 - (0.75 - 0.49))) < 0.004 for r_ in resB), resB
        resA = [r_ for r_ in t_cv.results if r_["group"] == "A"]
        assert len(resA) == 3 and all(abs(r_["e12"] - (-1.02 - (0.55 - 0.49))) < 0.004 for r_ in resA), resA
        cfg_, curves_, labels_, _ = spec_
        offs_ = [c_.get("off", 0) for c_ in curves_]
        assert offs_ == sorted(offs_, reverse=True) and offs_[0] > 0 == offs_[-1] and cfg_["y_axis"] == "no_numbers"
        assert len(labels_) == 5 and all(l_.get("y") is not None and "1/2" in l_["text"] for l_ in labels_), labels_
        t_cv.render()
        assert all(getattr(a_, "_anchor_y", None) is not None for a_ in t_cv.F.ann_items.values())
        # Peakpotentiale an einzelnen Peaks, y-Achse aus, Struktur je Messung (in Kurvenfarbe)
        st_cv.update(peak_labels="both", y_axis="off")
        st_cv["files"][0].update(structure=str(cdxml), struct_color=True)
        cfg_, curves_, labels_, _ = t_cv.figure_spec()
        assert len(labels_) == 10 and {l_["side"] for l_ in labels_} == {"above", "below"}, len(labels_)
        assert cfg_["y_axis"] == "off" and [i_["id"] for i_ in cfg_["images"]] == ["struktur_cv0"], cfg_["images"]
        t_cv.render()
        assert "image_struktur_cv0" in t_cv.F.axes_items
        assert any(c_ is st_cv["files"][0] and k_ == "structure" for c_, k_, _ in PRJ._path_slots({"cv": st_cv}))
        st_cv["files"][0].pop("structure")
        st_cv.update(peak_labels="e12", y_axis="auto")
        st_cv["show_bl"] = True                          # Basislinien zur Kontrolle: Hilfslinien, nicht in CSV/Hover
        cfg_, curves_, _, _ = t_cv.figure_spec()
        aux_ = [c_ for c_ in curves_ if c_.get("aux")]
        assert len(aux_) >= 5 * 3 and all(c_.get("label") is None for c_ in aux_), len(aux_)
        t_cv.render()
        assert len(t_cv.hover.curves) == len(curves_) - len(aux_), (len(t_cv.hover.curves), len(curves_))
        st_cv["show_bl"] = False
        st_cv.update(stack=False, e12_each=False)          # E½ ohne Stapeln: je Gruppe, nie über Gruppen gemittelt
        _, _, labels_, _ = t_cv.figure_spec()
        assert sorted(l_["text"].split(" =")[0] for l_ in labels_) == ["E$_{1/2}$ (A)", "E$_{1/2}$ (B)"], labels_
        assert {round(l_["lam"], 2) for l_ in labels_} == {-1.08, -1.18}, labels_
        st_cv.update(stack=True, e12_each=True)
        for plot_, n_series in (("dep", 2), ("rs", 4), ("loglog", 4), ("ratio", 1), ("trumpet", 4)):
            st_cv["plot"] = plot_
            cfg_, curves_, _, _ = t_cv.figure_spec()
            pts_ = [c_ for c_ in curves_ if c_.get("marker")]
            assert len(pts_) == n_series, (plot_, len(pts_))
            if plot_ == "dep":
                assert all(abs(y_ - 60) < 8 for c_ in pts_ for y_ in c_["y"]), [list(c_["y"]) for c_ in pts_]
                assert sorted(len(c_["x"]) for c_ in pts_) == [2, 3]
            if plot_ == "rs":
                assert len([c_ for c_ in curves_ if c_.get("ls") == "--"]) == 2      # Gerade nur ab 3 Scanraten (A)
            if plot_ == "trumpet":
                assert abs(min(c_["x"].min() for c_ in pts_) - np.log10(0.05)) < 1e-9
            if plot_ == "ratio":
                assert all(0.8 < y_ < 1.3 for c_ in pts_ for y_ in c_["y"]), [list(c_["y"]) for c_ in pts_]
            if plot_ == "loglog":
                fits_ = t_cv.rate_fits()
                assert list(fits_) == [("A", 1)] and abs(fits_[("A", 1)]["ll_a"]["slope"] - 0.5) < 0.05, fits_
                assert "log-log" in t_cv.lbl_res.text() or "log–log" in t_cv.lbl_res.text(), t_cv.lbl_res.text()
        files = t_cv.export_to(out / "cv_dep")
        assert all(f.exists() for f in files) and "group" in files[-1].read_text(encoding="utf-8-sig")
        assert "Couple 1 (A)" in files[-2].read_text(encoding="utf-8")
        st_cv["files"] = st_cv["files"][:1]
        assert t_cv.figure_spec() is None and "zwei Scanraten" in t_cv.note or "two scan rates" in t_cv.note
        st_cv.update(files=[], couples=[], ref_mode="none", plot="cv", stack=False, e12_each=False)
        w.cv_tab.fill_table()
        report.append("CV stacking, reference per file, rate plots: OK")

        # Spektroelektrochemie: BWSpec-Spektren + Gamry-Potential über die Uhrzeit, Nernst-Fit, Export
        import datetime as _dt
        secdir = out / "input" / "sec"
        secdir.mkdir(parents=True, exist_ok=True)
        t_start = _dt.datetime(2026, 1, 1, 12, 0, 0)
        T_ = np.arange(0.0, 1150.0, 2.0)
        E_run = -1.3 + 0.002 * T_                      # 2 mV/s
        write_dta(secdir / "sec.DTA", [(E_run, 1.0 + 0 * E_run)], 2.0, cycle_col=True, times=T_)  # Start 12:00:00
        pix = np.arange(0, 2048)
        lam_ = 200.0 + 0.35 * pix
        sec_files = []
        for j in range(0, 70):
            ts = j * 15.0 + 5.0
            e_ = -1.3 + 0.002 * ts
            ox = 1 / (1 + np.exp(-SECMOD.F_RT * (e_ + 0.80)))          # E½ = −0.80 V, n = 1
            A_ = (1 - ox) * 0.3 * np.exp(-((lam_ - 432) / 25) ** 2) + ox * 0.35 * np.exp(-((lam_ - 373) / 15) ** 2)
            hdr = ["File Version;BWSpec4.11_1", f"Date;{(t_start + _dt.timedelta(seconds=ts)):%Y-%m-%d %H:%M:%S}",
                   "coefs_a0;200", "coefs_a1;0.35", "coefs_a2;0", "coefs_a3;0", "sel_pixel_start;0",
                   "Pixel;Wavelength;Wavenumber;Raman Shift;Dark;Reference;Raw data #1;Dark Subtracted #1;%TR #1;"
                   "Absorbance #1;Irradiance (W/m^2/nm) #1"]
            rows_ = [f"{p_};   ;   ;   ;0;0;0;0;0;{a_:.5f};0;" for p_, a_ in zip(pix, A_)]
            f_ = secdir / f"s{j:03d}.txt"
            f_.write_text("\r\n".join(hdr + rows_), encoding="latin-1", newline="")
            sec_files.append(str(f_))
        (secdir / "LOG_s.txt").write_text("File Format;TXTR\r\nFile Description;Peak Monitor File\r\n",
                                         encoding="latin-1")
        assert not SECMOD.is_bwspec(secdir / "LOG_s.txt") and SECMOD.is_bwspec(sec_files[0])
        sp_ = [SECMOD.read_bwspec(f_) for f_ in sec_files]
        E_map = SECMOD.potentials(sp_, CVMOD.read_gamry(secdir / "sec.DTA"))
        assert np.allclose(E_map, -1.3 + 0.002 * (np.arange(70) * 15.0 + 5.0), atol=1e-3), E_map[:3]
        a432 = np.array([np.interp(432, s_["lam"], s_["A"]) for s_ in sp_])
        fit_ = SECMOD.nernst_fit(E_map, a432)
        assert abs(fit_["e12"] + 0.80) < 0.005 and abs(fit_["n"] - 1.0) < 0.1, fit_
        st_sec = w.sec_state
        st_sec.update(spec_files=sec_files, dta=str(secdir / "sec.DTA"), bands="373, 432", mode="spectra")
        files = w.sec_tab.export_to(out / "sec_spectra")
        assert all(f.exists() for f in files), files
        assert not any(f.name.endswith("_methods.txt") for f in files), files           # kein Fit, kein SI-Text
        st_sec.update(mode="trace", trace_lams="432")
        files = w.sec_tab.export_to(out / "sec_trace")
        assert all(f.exists() for f in files) and abs(w.sec_tab.info[432.0]["e12"] + 0.80) < 0.005
        meth_ = files[-1].read_text(encoding="utf-8")
        assert files[-1].name.endswith("_methods.txt") and "apparent" in meth_ and "grid search" in meth_ \
            and "λ = 432 nm: E1/2,app = −0.80" in meth_, meth_
        slots_ = PRJ._path_slots({"sec": st_sec})
        assert sum(1 for c_, k_, _ in slots_ if c_ is st_sec["spec_files"]) == 70 and \
            any(c_ is st_sec and k_ == "dta" for c_, k_, _ in slots_), "SEC-Dateien fehlen im Projekt"
        st_sec.update(mode="spectra")
        w.sec_state.update(cap_solvent="THF", cap_start="neutral", cap_end="dication")    # Unterschrift wie in der SI
        cap_s = w.sec_tab.si_caption("TEST-1")
        assert cap_s.startswith("UV-vis-SEC of <b>TEST-1</b> in THF.") and "Starting potential (vs. " in cap_s \
            and "End potential" in cap_s, cap_s
        # Zwischenspezies: Potential -> nächstliegendes Spektrum grün (Nr. 17: −0.81 V), über den grauen, mit Label
        assert "green" not in cap_s and cap_s.count("<b>") == 3, cap_s                      # ohne Eingabe wie bisher
        cv0_ = w.sec_tab.figure_spec()[1]
        assert [c_["color"] for c_ in cv0_ if c_["lw"] > 1] == [SECMOD.START_COLOR, SECMOD.END_COLOR]
        w.sec_state.update(mid="-0.80", cap_mid="radical cation")
        _, cv_, lb_, _ = w.sec_tab.figure_spec()
        hi_ = [c_ for c_ in cv_ if c_["lw"] > 1]
        assert [c_["color"] for c_ in hi_] == [SECMOD.START_COLOR, SECMOD.MID_COLORS[0][0], SECMOD.END_COLOR], hi_
        assert all(c_["z"] == 3 for c_ in cv_[-3:]) and cv_[-3:] == hi_, "hervorgehobene nicht oben"
        assert "0.81" in hi_[1]["label"] and np.allclose(hi_[1]["y"], np.interp(
            cv_[0]["x"], sp_[16]["lam"], sp_[16]["A"])), hi_[1]["label"]
        assert {l_["lam"]: l_["color"] for l_ in lb_} == {373.0: SECMOD.END_COLOR, 432.0: SECMOD.START_COLOR}
        cap_s = w.sec_tab.si_caption("TEST-1")
        assert ("Color coding: <b>neutral</b> (blue), <b>radical cation</b> (green), <b>dication</b> (red).") \
            in cap_s, cap_s
        w.sec_state.update(mid="-0.80, 0.30, -0.81", cap_mid="radical cation; X")      # Doppelte fallen weg
        hi_ = [c_ for c_ in w.sec_tab.figure_spec()[1] if c_["lw"] > 1]
        assert [c_["color"] for c_ in hi_].count(SECMOD.MID_COLORS[1][0]) == 1 and len(hi_) == 4
        assert "(green), <b>X</b> (orange), <b>dication</b> (red)" in w.sec_tab.si_caption("TEST-1")
        w.sec_state.update(dta=None, mid="35")                       # ohne Potentiale: Nummer des Spektrums
        hi_ = [c_ for c_ in w.sec_tab.figure_spec()[1] if c_["lw"] > 1]
        assert hi_[1]["color"] == SECMOD.MID_COLORS[0][0] and np.allclose(hi_[1]["y"], np.interp(
            hi_[1]["x"], sp_[34]["lam"], sp_[34]["A"]))
        w.sec_state.update(highlight=False)
        assert "green" not in w.sec_tab.si_caption("TEST-1")
        w.sec_state.update(dta=str(secdir / "sec.DTA"), highlight=True, mid="", cap_mid="")
        w.sec_state.update(cap_solvent="", cap_start="", cap_end="")
        report.append("spectroelectrochemistry: OK")

        # Audit P2 (CV + SEC): Grenzfälle mit synthetischen Daten
        E_s = np.concatenate([np.arange(0, 1, 0.01), np.arange(1, -1, -0.01), np.arange(-1, 1e-9, 0.01)])
        dE_s = np.gradient(E_s)                          # Start innerhalb: 0 -> +1 -> −1 -> 0, Paar bei −0,5 V

        def g_s(c_, w_):
            return np.exp(-((E_s - c_) / w_) ** 2)
        I_s = np.where(dE_s < 0, -g_s(-0.55, 0.05), 0) + np.where(dE_s > 0, 0.8 * g_s(-0.45, 0.05), 0)
        r_ = CVMOD.analyze_couple(E_s, I_s, -0.75, -0.25)
        assert not r_["anodic_first"] and r_["ratio_ok"] and abs(r_["e12"] + 0.5) < 0.005, r_
        I_s2 = I_s + np.where(dE_s < 0, -0.4 * g_s(-0.85, 0.04), 0)   # zweiter Prozess vor dem Umkehrpunkt (−1 V)
        r_ = CVMOD.analyze_couple(E_s, I_s2, -0.75, -0.25)
        assert not r_["anodic_first"] and not r_["ratio_ok"] and not np.isfinite(r_["ratio"]), r_
        assert CVMOD.rate_fit([100, 100, 100], [1, 2, 3]) is None and \
            CVMOD.randles_sevcik([100, 100, 100], [1, 2, 3]) is None
        assert SECMOD.nernst_fit([0.0] * 5, [0.1, 0.2, 0.3, 0.4, 0.5]) is None
        E_n = np.linspace(-1.3, -0.5, 60)
        for e0_, n0_, bound_ in ((-0.4, 1.0, True), (-0.6, 4.0, True), (-0.9, 0.2, True), (-0.9, 1.0, False)):
            f_ = SECMOD.nernst_fit(E_n, 0.1 + 0.5 / (1 + np.exp(-n0_ * SECMOD.F_RT * (E_n - e0_))))
            assert f_ is not None and bool(f_.get("bound")) == bound_, (e0_, n0_, f_)
        bw_ = secdir / "nocal.txt"                       # ohne Wellenlängen und Kalibrierung: ablehnen
        bw_.write_text("File Version;BWSpec4\r\nDate;2026-01-01 12:00:00\r\nPixel;Wavelength;Absorbance #1\r\n"
                       + "".join(f"{k_};;0.{k_ + 1}\r\n" for k_ in range(3)), encoding="latin-1", newline="")
        try:
            SECMOD.read_bwspec(bw_)
            raise AssertionError("BWSpec ohne Kalibrierung angenommen")
        except ValueError:
            pass
        bw_.write_text(bw_.read_text(encoding="latin-1").replace("Pixel;", "coefs_a0;400\nPixel;"),
                       encoding="latin-1")                # nur a0: konstante λ, ebenso ablehnen
        try:
            SECMOD.read_bwspec(bw_)
            raise AssertionError("BWSpec mit unvollständiger Kalibrierung angenommen")
        except ValueError:
            pass
        nt_ = secdir / "noT.DTA"                         # ohne Zeitspalte: keine Zuordnung über die Uhrzeit
        nt_.write_text("DATE\tLABEL\t1/1/2026\r\nTIME\tLABEL\t12:00:00\r\nCURVE1\tTABLE\r\n\tPt\tVf\tIm\r\n"
                       "\t#\tV\tA\r\n" + "".join(f"\t{k_}\t{k_ * 0.1}\t1e-6\r\n" for k_ in range(10)),
                       encoding="latin-1", newline="")
        assert SECMOD.potentials(sp_[:3], CVMOD.read_gamry(nt_)) is None
        # Stromdichte nur, wenn alle Messungen eine Fläche haben; sonst Strom mit Hinweis
        na_ = cvdir / "cv_noarea.DTA"
        na_.write_text(cv_files[0].read_text(encoding="latin-1").replace(
            "AREA\tQUANT\t7.10000E-002\tElectrode &Area (cm^2)", ""), encoding="latin-1", newline="")
        st_cv.update(files=[{"file": str(p_), "uid": f"d{k_}", "visible": True, "color": "#000000",
                             "label": f"d{k_}"} for k_, p_ in enumerate((cv_files[0], na_))],
                     density=True, couples=[[-1.25, -0.8]], ref_mode="none", cycle=1, baseline=True)
        cfg_, curves_, _, _ = t_cv.figure_spec()
        assert cfg_["ylabel"] == "I [µA]" and np.allclose(curves_[0]["y"], curves_[1]["y"]), cfg_["ylabel"]
        assert T("cv_area_missing", files="d1") in t_cv.note, t_cv.note
        st_cv["files"] = st_cv["files"][:1]
        cfg_, curves_, _, _ = t_cv.figure_spec()
        assert cfg_["ylabel"].startswith("j [mA") and T("cv_area_missing", files="d0") not in t_cv.note
        st_cv["density"] = False
        # gewählter Zyklus fehlt in einer Messung -> Hinweis (nicht still den letzten nehmen)
        st_cv["cycle"] = 2
        t_cv.figure_spec()
        assert T("cv_cycle_missing", files="d0", n=2) in t_cv.note, t_cv.note
        assert "For d0 (fewer cycles), the last cycle was used." in t_cv.methods_text(), t_cv.methods_text()
        st_cv["cycle"] = 1
        # Basislinie verlangt, aber nicht möglich -> Methodentext sagt es
        st_cv["couples"] = [[-1.13, -0.8]]
        t_cv.figure_spec()
        assert t_cv.results[0]["bl_fail"] and "from the ratios given below" in t_cv.methods_text(), \
            t_cv.methods_text()
        # interne Referenz nicht gefunden: Achse wie gemessen -> Text und Unterschrift nicht „vs. Fc0/+“
        st_cv.update(couples=[[-1.25, -0.8]], ref_mode="internal", ref_uid="missing", ref_compound="ac2fc")
        t_cv.figure_spec()
        assert "Fc0/+" not in t_cv.methods_text() and "Fc" not in t_cv.si_caption("X"), t_cv.methods_text()
        # Vorschlag bei eigener Referenz der ersten Messung: Fenster in der Skala der globalen Referenz speichern
        p_ = cvdir / "cvB_50.DTA"
        st_cv.update(ref_mode="none", couples=[], files=[{
            "file": str(p_), "uid": "sB", "visible": True, "color": "#000000", "label": "B", "ref": {
                "mode": "internal", "uid": "sB", "lo": 0.5, "hi": 1.0, "compound": "ac2fc", "value": 0.49}}])
        t_cv.suggest()
        t_cv.figure_spec()
        assert t_cv.results and not any(r_["edge"] for r_ in t_cv.results) and \
            abs(t_cv.results[0]["e12"] - (-0.92 - (0.75 - 0.49))) < 0.004, t_cv.results
        # eine Messung mit gefundener, eine mit fehlender Referenz in einer Gruppe: E½ nie über Skalen mitteln
        st_cv.update(couples=[[-1.5, -0.6]], stack=False, e12_each=False, show_e12=True, files=[
            dict(st_cv["files"][0], group="G"),
            {"file": str(cvdir / "cvB_100.DTA"), "uid": "sB2", "visible": True, "color": "#000000", "label": "B2",
             "group": "G", "ref": {"mode": "internal", "uid": "fehlt", "lo": 0.5, "hi": 1.0, "compound": "ac2fc",
                                   "value": 0.49}}])
        _, _, labels_, _ = t_cv.figure_spec()
        assert {r_["ref"] for r_ in t_cv.results} == {"Fc$^{0/+}$", st_cv["ref_electrode"]}, t_cv.results
        assert not labels_ and "not averaged" in t_cv.methods_text(), (labels_, t_cv.methods_text())
        st_cv.update(files=[], couples=[], ref_mode="none", baseline=True)
        t_cv.fill_table()
        # SEC: Namen der Zwischenstufen nach Eingabeposition (weggefallene Einträge verschieben nichts)
        w.sec_state.update(mid="-1.3; -0.80", cap_mid="startlike; radical", cap_start="neutral", cap_end="dication")
        cap_s = w.sec_tab.si_caption("TEST-1")
        assert "<b>radical</b> (green)" in cap_s and "startlike" not in cap_s, cap_s
        w.sec_state.update(mid=";-0.80", cap_mid=";radical")         # leerer Eintrag behält seine Position
        assert "<b>radical</b> (green)" in w.sec_tab.si_caption("TEST-1"), w.sec_tab.si_caption("TEST-1")
        w.sec_state.update(first=3, last=3, mid="", cap_mid="")      # ein Spektrum: nur Rot im Bild
        cap_s = w.sec_tab.si_caption("TEST-1")
        assert "(blue)" not in cap_s and "<b>dication</b> (red)" in cap_s, cap_s
        w.sec_state.update(first=1, last=0, cap_start="", cap_end="")
        w.sec_state["offset_s"] = -400.0                 # Spektren vor Scanbeginn: Hinweis mit Anzahl
        w.sec_tab.figure_spec()
        assert T("sec_outside", n=25) in w.sec_tab.lbl_info.text(), w.sec_tab.lbl_info.text()   # > 30 s davor
        w.sec_state["offset_s"] = -20.0                  # ein Spektrum kurz vor Scanbeginn: normal, kein Hinweis
        w.sec_tab.figure_spec()
        assert "⚠" not in w.sec_tab.lbl_info.text(), w.sec_tab.lbl_info.text()
        w.sec_state.update(offset_s=0.0, mode="trace", trace_lams="432", baseline=True, diff=True)
        w.sec_tab.figure_spec()
        meth_ = w.sec_tab.methods_text()
        assert "780" in meth_ and "Difference spectra" in meth_, meth_
        short_ = []                                      # Spektren nur bis 700 nm: keine Basislinie 780–820 nm
        for f_ in sec_files:
            g_ = secdir / "short" / Path(f_).name
            g_.parent.mkdir(exist_ok=True)
            ln_ = Path(f_).read_text(encoding="latin-1").splitlines()
            g_.write_text("\r\n".join(ln_[:7 + 1428]), encoding="latin-1", newline="")
            short_.append(str(g_))
        w.sec_state.update(spec_files=short_, diff=False)
        w.sec_tab.figure_spec()
        assert w.sec_tab.fits and "780" not in w.sec_tab.methods_text(), w.sec_tab.methods_text()
        w.sec_state.update(spec_files=sec_files, mode="spectra", trace_lams="", baseline=False, diff=False)
        nat_ = []                                        # ohne Zeit: natürliche Reihenfolge der Namen
        for k_ in (1, 2, 10):
            f_ = secdir / "nat" / f"s_{k_}.txt"
            f_.parent.mkdir(exist_ok=True)
            f_.write_text(Path(sec_files[0]).read_text(encoding="latin-1").replace("Date;", "Datum;"),
                          encoding="latin-1", newline="")
            nat_.append(str(f_))
        w.sec_state["spec_files"] = nat_[::-1]
        assert [s_["name"] for s_ in w.sec_tab.series()[0]] == ["s_1", "s_2", "s_10"]
        w.sec_state["spec_files"] = sec_files
        report.append("audit P2 CV/SEC edge cases: OK")

        # SEC im Overlay: Spektren nach Potential übernehmen, nur normiert (kein ε aus der Dünnschichtzelle)
        ov_, st_ov = w.ov_tab, w.ov_state
        saved_ov = {k: json.loads(json.dumps(st_ov[k])) for k in ("entries", "mode", "nlo", "nhi", "baseline")}
        st_ov.update(entries=[], mode="norm", nlo=340.0, nhi=460.0)
        ov_.add_sec_entries([0, 69])
        e0, e1 = st_ov["entries"]
        assert e0["kind"] == "sec" and abs(e0["E"] - (-1.3 + 0.002 * 5.0)) < 1e-3, e0
        assert "V vs. Ag/AgNO" in e0["label"] and e1["E"] > e0["E"], (e0["label"], e1["E"])
        ov_.render()
        cs_ = [c_ for c_ in ov_.F.curves if c_.get("label") in (e0["label"], e1["label"]) or "SEC" in str(c_.get("label"))]
        assert len(cs_) == 2, [c_.get("label") for c_ in ov_.F.curves]
        for c_, peak_ in zip(cs_, (432, 373)):          # reduziert: Bande bei 432 nm, oxidiert: 373 nm
            win_ = (c_["x"] >= 340) & (c_["x"] <= 460)
            assert abs(float(np.max(c_["val"][win_])) - 1.0) < 1e-6, "SEC nicht normiert"
            assert abs(float(c_["x"][win_][np.argmax(c_["val"][win_])]) - peak_) < 2, (peak_, c_["label"])
        assert c_["x"].min() >= SECMOD.LAM_RANGE[0] and c_["x"].max() <= SECMOD.LAM_RANGE[1]
        assert any(c_ is e0 and k_ == "file" for c_, k_, _ in PRJ._path_slots({"overlay": st_ov}))
        # Serie: EIN Eintrag, gemeinsamer Faktor (auf das Startspektrum) -> Verhältnisse bleiben erhalten
        st_ov.update(entries=[], baseline=False)
        ov_.add_sec_series(0, 69)
        ov_.render()
        assert len(ov_.F.curves) == 70 and len(st_ov["entries"]) == 1, len(ov_.F.curves)
        named_ = [c_["label"] for c_ in ov_.F.curves if c_.get("label")]
        assert len(named_) == 2 and all("V vs. Ag/AgNO" in n_ for n_ in named_), named_
        raw0, raw1 = (SECMOD.read_bwspec(sec_files[i_]) for i_ in (0, 69))
        start_ = next(c_ for c_ in ov_.F.curves if c_["color"] == SECMOD.START_COLOR)
        end_ = next(c_ for c_ in ov_.F.curves if c_["color"] == SECMOD.END_COLOR)
        win_ = (start_["x"] >= 340) & (start_["x"] <= 460)
        assert abs(float(start_["val"][win_].max()) - 1.0) < 1e-6
        r_raw = np.interp(373, raw1["lam"], raw1["A"]) / np.interp(432, raw0["lam"], raw0["A"])
        r_plot = np.interp(373, end_["x"], end_["val"]) / np.interp(432, start_["x"], start_["val"])
        assert abs(r_plot / r_raw - 1) < 0.02, (r_plot, r_raw)      # nicht einzeln normiert
        assert sum(1 for c_, k_, _ in PRJ._path_slots({"overlay": st_ov}) if c_ is st_ov["entries"][0]["files"]) == 70
        st_ov["entries"][0]["ends_only"] = True
        ov_.render()
        assert len(ov_.F.curves) == 2
        st_ov["mode"] = "eps"                           # ε-Modus: SEC wird ausgelassen, nicht falsch skaliert
        assert ov_.figure_spec() is None or all("SEC" not in str(c_.get("label")) for c_ in ov_.figure_spec()[1])
        st_ov.update(saved_ov)
        st_sec.update(spec_files=[], dta=None, mode="spectra", trace_lams="", bands="")
        report.append("SEC spectra in overlay: OK")

        # EPR: BES3T lesen (G -> mT), Dunkelspektrum, Basislinie, Ausblenden, Simulation anpassen, EasySpin-Parameter
        eprdir = out / "input" / "epr"
        eprdir.mkdir(parents=True, exist_ok=True)
        Bg = np.linspace(4400.0, 5100.0, 7001)                       # G
        signal = -np.gradient(np.exp(-((Bg - 4700) / 30) ** 2)) * 1000  # Ableitungslinie bei 470 mT
        drift = 0.05 * (Bg - 4400)                                     # lineare Drift der Basislinie

        def write_bes3t(stem, y, time="11:30:15", temp=100.0):
            (eprdir / f"{stem}.dsc").write_text("\n".join([
                "#DESC\t1.2", "DSRC\tEXP", "BSEQ\tLIT", "IKKF\tREAL", "XTYP\tIDX", "YTYP\tNODATA", "IRFMT\tD",
                f"XPTS\t{len(Bg)}", f"XMIN\t{Bg[0]}", f"XWID\t{Bg[-1] - Bg[0]}", "XUNI\t'G'", "#SPL\t1.2",
                "MWFQ\t9338553208.5", f"STMP\t{temp}", "B0MA\t0.00095", "MWPW\t0.01", "DATE\t08/22/25",
                f"TIME\t{time}"]),
                encoding="latin-1")
            y.astype("<f8").tofile(eprdir / f"{stem}.DTA")
        dark_bg = 3.0 * np.sin((Bg - 4400) / 200)                      # Untergrund, der auch dunkel da ist
        write_bes3t("dark", drift + dark_bg)
        write_bes3t("light", drift + dark_bg + signal)
        sim_B = np.linspace(440.0, 510.0, 1401)                        # mT, wie EasySpin
        sim_I = 0.5 * np.interp(sim_B, Bg / 10, signal)
        np.savetxt(eprdir / "sim.txt", np.column_stack([sim_B, sim_I]), delimiter="\t")
        (eprdir / "sim_params.txt").write_text("Sys.S = 1;\nSys.g = 2.0023;\nSys.D = [11100 450]; % MHz\n"
                                               "Exp.mwFreq = 9.3386;\n", encoding="utf-8")
        sp_e = EPRMOD.read_any(eprdir / "light.DTA")[0]
        assert abs(sp_e["B"][0] - 440.0) < 1e-9 and abs(sp_e["meta"]["mw_GHz"] - 9.3385535) < 1e-6, sp_e["meta"]
        assert abs(sp_e["meta"]["mod_mT"] - 0.95) < 1e-9 and abs(sp_e["meta"]["power_mW"] - 10) < 1e-9
        pr_ = EPRMOD.parse_easyspin((eprdir / "sim_params.txt").read_text(encoding="utf-8"))
        assert pr_["Sys.D"] == [11100.0, 450.0] and "0.3703 cm" in EPRMOD.zfs_summary(pr_)
        assert abs(EPRMOD.g_value(9.3386, 333.2) - 2.0024) < 1e-3
        t_e = w.epr_tab
        ed = t_e._new("exp", str(eprdir / "dark.dsc"))
        el = t_e._new("exp", str(eprdir / "light.dsc"))
        es = t_e._new("sim", str(eprdir / "sim.txt"))
        ed["visible"] = False
        el["dark"] = ed["uid"]
        es.update(fit_to=el["uid"], params=str(eprdir / "sim_params.txt"))
        w.epr_state.update(entries=[ed, el, es], bl_order=1, exclude="499-501", smooth=0.0, normalize=False,
                           stack=False, labels="Z@452, X@490", show_params=True)
        cache_ = {}
        Bp, Ip = t_e.processed(el, cache_)
        cleaned = Ip[np.isfinite(Ip)]
        ref_ = (signal / 1.0)[np.isfinite(Ip)]
        assert np.max(np.abs(cleaned - ref_)) < 0.02 * np.ptp(signal), "Dunkelspektrum/Basislinie"
        assert np.isnan(Ip[(Bp > 499.2) & (Bp < 500.8)]).all(), "Bereich nicht ausgeblendet"
        data_, _ = t_e.components([el, es], {}, [])
        Is = data_[es["uid"]][1]
        assert abs(np.nanmax(Is) / np.nanmax(Ip) - 1) < 0.05, "Simulation nicht an Messung angepasst"
        t_e.fill_table()
        files = w.epr_tab.export_to(out / "epr_test")
        assert all(f.exists() for f in files), files
        F_e = w.epr_tab.build(MplFigure(figsize=(8, 6), dpi=PREVIEW_DPI))
        assert any("cm" in a.get_text() for a in F_e.ann_items.values()), "D/E-Label fehlt"
        slots_ = PRJ._path_slots({"epr": w.epr_state})
        assert any(c_ is es and k_ == "params" for c_, k_, _ in slots_), "EPR-Parameter fehlen im Projekt"
        # Intensität gegen Zeit: Amplituden 25/50/75/100 % im Minutentakt nach dem Dunkelspektrum (11:30:15)
        series = [ed]
        for k_, amp in enumerate((0.25, 0.5, 0.75, 1.0), 1):
            write_bes3t(f"kin{k_}", drift + dark_bg + amp * signal, time=f"11:3{k_}:15")
            ek = t_e._new("exp", str(eprdir / f"kin{k_}.dsc"))
            ek["label"] = f"kin{k_}"
            series.append(ek)
        w.epr_state.update(entries=series, exclude="", labels="", show_params=False, mode="intensity",
                           int_window="455-485", int_method="double", int_x="time", normalize=True)
        t_e.fill_table()
        t_e.table.selectRow(1)
        series[1]["dark"] = ed["uid"]
        t_e.dark_for_all()
        files = w.epr_tab.export_to(out / "epr_kin")
        xs_ = [v[0] for v in t_e.series_values]
        ys_ = [v[1] for v in t_e.series_values]
        assert np.allclose(xs_, [1, 2, 3, 4]) and np.allclose(ys_, [0.25, 0.5, 0.75, 1.0], atol=0.02), (xs_, ys_)
        w.epr_state["int_method"] = "pp"
        w.epr_tab.export_to(out / "epr_kin_pp")
        assert np.allclose([v[1] for v in t_e.series_values], [0.25, 0.5, 0.75, 1.0], atol=0.02)
        txt_ = next(f for f in files if f.name.endswith("_kin4.txt"))
        arr_ = np.loadtxt(txt_, comments="%")
        assert arr_.shape[1] == 2 and abs(arr_[0, 0] - 440.0) < 1e-6 and "mwFreq" in txt_.read_text(encoding="utf-8")
        assert np.max(np.abs(arr_[:, 1] - signal)) < 0.02 * np.ptp(signal), "Export nicht verarbeitet"
        assert any(f.name.endswith("_intensity.csv") for f in files) and any(f.name.endswith("_data.csv") for f in files)
        w.epr_state.update(entries=[], labels="", exclude="", show_params=False, mode="spectra")
        report.append("EPR (BES3T, dark spectrum, simulation, intensity, export): OK")

        # EPR S = ½: TEMPO-artiges Spektrum mit Feldfehler +0.25 mT fitten -> Kalibrierung; zwei Spezies 70:30
        mw_ = 9.40
        Bt = np.linspace(330.0, 345.0, 6001)
        tempo_I = EPRMOD.iso_spectrum(Bt, mw_, 2.0055, [(1, 1.0, 1.55)], 0.15)
        noise = np.random.default_rng(0).standard_normal(len(Bt)) * 0.01 * np.ptp(tempo_I)
        np.savetxt(eprdir / "tempo-9p400000.txt", np.column_stack([Bt + 0.25, tempo_I + noise]),
                   header="B [mT]\tI [a.U.]", comments="", delimiter="\t")
        rad_I = EPRMOD.iso_spectrum(Bt, mw_, 2.0030, [(2, 0.5, 0.45)], 0.12)
        np.savetxt(eprdir / "radikal-9p400000.txt", np.column_stack([Bt + 0.25, rad_I]),
                   header="B [mT]\tI [a.U.]", comments="", delimiter="\t")
        et = t_e._new("exp", str(eprdir / "tempo-9p400000.txt"))
        er = t_e._new("exp", str(eprdir / "radikal-9p400000.txt"))
        assert abs(t_e.spectrum(et)["meta"]["mw_GHz"] - 9.4) < 1e-9
        w.epr_state.update(entries=[et, er], bl_order=-1, smooth=0.0, normalize=True, stack=False,
                           show_sum=True, show_residual=True, calib_uid=None, calib_g=2.0055)
        t_e.fill_table()
        t_e.table.selectRow(0)
        t_e.add_iso()
        ti = w.epr_state["entries"][-1]
        ti["nuclei"] = "1 14N 1.4"
        t_e.table.selectRow(2)
        t_e.fit_current()
        assert abs(ti["g"] - EPRMOD.g_value(mw_, EPRMOD.b_res(mw_, 2.0055) + 0.25)) < 2e-5, ti
        assert "1.55" in ti["nuclei"] and ti["fit"]["r2"] > 0.98, ti     # 1 % Rauschen
        w.epr_state["calib_uid"] = ti["uid"]
        t_e.fill_table()
        dB_, _ = t_e.calibration_shift()
        assert abs(dB_ - 0.25) < 0.003, dB_
        meth_ = t_e.methods_text()
        assert "Levenberg–Marquardt" in meth_ and "Field calibration: g_fit" in meth_ and "ΔB = +0.25" in meth_ \
            and "a(1 × I = 1) = 1.5" in meth_, meth_
        t_e.table.selectRow(1)
        t_e.add_iso()
        ri = w.epr_state["entries"][-1]
        ri["nuclei"] = "2 1H 0.4"
        t_e.table.selectRow(3)
        t_e.fit_current()
        assert abs(ri["g"] - 2.0030) < 5e-5 and abs(EPRMOD.parse_nuclei(ri["nuclei"])[0][2] - 0.45) < 0.005, ri
        w.epr_state.update(cap_species="TEST-1^{•+}", cap_solvent="THF", cap_conc="1 mM")   # Unterschrift wie in der SI
        cap_e = t_e.si_caption("TEST-1")
        assert cap_e.startswith("X-Band EPR spectr") and "<b>TEST-1<sup>•+</sup></b> in THF (1 mM)." in cap_e \
            and "2×<sup>1</sup>H: 12." in cap_e and "g = 2.003" in cap_e and "referenced to a standard" in cap_e, cap_e
        w.epr_state.update(cap_species="", cap_solvent="", cap_conc="")
        # zwei Spezies: Messung = 0.7·A + 0.3·B, beide als EasySpin-ASCII importiert
        A_ = EPRMOD.iso_spectrum(Bt, mw_, 2.0030, [(1, 1.0, 1.4)], 0.2)
        B_ = EPRMOD.iso_spectrum(Bt, mw_, 2.0100, [], 0.5)
        np.savetxt(eprdir / "mix-9p400000.txt", np.column_stack([Bt, 0.7 * A_ + 0.3 * B_]), delimiter="\t")
        np.savetxt(eprdir / "simA.txt", np.column_stack([Bt, A_]), delimiter="\t")
        np.savetxt(eprdir / "simB.txt", np.column_stack([Bt, B_]), delimiter="\t")
        em = t_e._new("exp", str(eprdir / "mix-9p400000.txt"))
        sa, sb = t_e._new("sim", str(eprdir / "simA.txt")), t_e._new("sim", str(eprdir / "simB.txt"))
        sa["fit_to"] = sb["fit_to"] = em["uid"]
        w.epr_state.update(entries=[em, sa, sb], calib_uid=None, bl_order=-1)
        info_ = []
        data_, extra_ = t_e.components([em, sa, sb], {}, info_)
        wa = np.nanmax(data_[sa["uid"]][1]) / np.nanmax(A_)
        wb = np.nanmax(data_[sb["uid"]][1]) / np.nanmax(B_)
        assert abs(wa - 0.7) < 1e-3 and abs(wb - 0.3) < 1e-3, (wa, wb)
        assert np.nanmax(np.abs(extra_[em["uid"]]["res"][1])) < 1e-6 * np.ptp(A_) and "sum" in extra_[em["uid"]]
        assert info_ and "%" in info_[0], info_
        files = w.epr_tab.export_to(out / "epr_species")
        assert all(f.exists() for f in files), files
        w.epr_state.update(entries=[], calib_uid=None, show_residual=False, bl_order=1)
        report.append("EPR S=1/2 fit, TEMPO calibration, multiple species: OK")

        # EPR-Zerfall bei konstanter Heizrate: Gauß-Intensität eines Merkmals, Arrhenius-Fit (E_a = 50 kJ/mol, ln A = 45)
        Ea_true, lnA_true = 50e3, 45.0
        temps = np.arange(100.0, 135.01, 0.5)
        beta_ = 0.5 / 120.0                                         # 0,5 K alle 2 min
        I_true = EPRMOD.decay_model(temps, np.exp(lnA_true), Ea_true, beta_)
        Bk = Bg                                                       # G, 440–510 mT
        kin_entries = []
        for j, (Tk, Ik) in enumerate(zip(temps, I_true)):
            y = -2000.0 * Ik * np.exp(-((Bk / 10 - 487.0) ** 2) / 9.0) / Tk + 0.3 * (Bk - 4400) / 10
            stem = f"temp{j:03d}"
            hh, mm = divmod(2 * j, 60)
            write_bes3t(stem, y, time=f"{12 + hh:02d}:{mm:02d}:00", temp=Tk)
            ek = t_e._new("exp", str(eprdir / f"{stem}.dsc"))
            kin_entries.append(ek)
        w.epr_state.update(entries=kin_entries, mode="intensity", int_method="gauss", int_window="470-510",
                           int_center=487.0, int_x="T", int_curie=True, ref_freq=0.0, T_offset=0.0, kin_fit=True,
                           kin_norm="lsq", bl_order=-1, smooth=0.0, normalize=True, calib_uid=None)
        t_e.fill_table()
        files = w.epr_tab.export_to(out / "epr_decay")
        r_ = t_e.kin_result
        assert r_ and abs(r_["Ea"] - Ea_true) / Ea_true < 0.02 and abs(r_["lnA"] - lnA_true) < 1.0, r_
        # rauschfrei: Fit exakt, Intervall nur Rundungsbreite
        assert r_["ci_Ea"][0] - 1e-3 < Ea_true < r_["ci_Ea"][1] + 1e-3 and abs(r_["beta"] - beta_) < 1e-9, r_
        meth_ = next(f for f in files if f.name.endswith("_methods.txt")).read_text(encoding="utf-8")
        assert "Gaussian on a linear baseline" in meth_ and "t(0.975; n − 3)" in meth_ and "Ea = 5" in meth_ \
            and "multiplied by the temperature (I·T) before fitting" in meth_ and "± 0.01 kJ" in meth_ \
            and " × 10^" in meth_, meth_
        w.epr_state["kin_norm"] = "script"
        w.epr_tab.export_to(out / "epr_decay_script")
        assert abs(t_e.kin_result["Ea"] - Ea_true) / Ea_true < 0.02, t_e.kin_result
        assert "Arrhenius" in next(f for f in files if f.name.endswith("_intensity.csv")).read_text(encoding="utf-8-sig")
        w.epr_state.update(entries=[], mode="spectra", kin_fit=False, int_method="double", int_window="", int_center=0.0)
        report.append("EPR decay at constant heating rate (Arrhenius): OK")

        # Audit P6 (EPR): Funde mit synthetischen Daten
        p6 = eprdir / "p6"
        p6.mkdir(exist_ok=True)
        B6 = np.linspace(330.0, 345.0, 1501)
        S6 = EPRMOD.iso_spectrum(B6, 9.4, 2.003, [], 0.15)
        MINUS = "−"

        def asc6(name, I, B=B6):
            np.savetxt(p6 / name, np.column_stack([B, I]), header="B [mT]\tI [a.U.]", comments="", delimiter="\t")
            return t_e._new("exp", str(p6 / name))
        # Fit, Simulation und Hover auf der Achse B·ν_ref/ν: g bleibt gleich (vorher mit ν der Datei gerechnet)
        e6 = asc6("s-9p400000.txt", S6)
        w.epr_state.update(entries=[e6], bl_order=-1, smooth=0.0, ref_freq=9.41, exclude="", calib_uid=None,
                           normalize=False, mode="spectra", unit="mT", kin_fit=False)
        t_e.fill_table()
        t_e.table.selectRow(0)
        t_e.add_iso()
        i6 = w.epr_state["entries"][-1]
        t_e.table.selectRow(1)
        t_e.fit_current()
        assert abs(i6["g"] - 2.003) < 2e-5 and abs(i6["fit"]["mw"] - 9.41) < 1e-12, i6
        t_e.figure_spec()
        assert abs(float(t_e.hover_text(EPRMOD.b_res(9.41, 2.003), None, "").split("g = ")[1]) - 2.003) < 1e-4
        # Handänderung nach dem Fit: kein „Fitting parameters“ mehr; Lorentz-Häkchen allein ändert keinen Wert
        t_e.cb_fit_lor.setChecked(True)
        assert i6["fit"], "fit_lorentz allein verwirft den Fit"
        t_e.cb_fit_lor.setChecked(False)
        t_e.sp_g.setValue(2.0040)
        assert i6["fit"] is None and "Fitting parameters" not in t_e.si_caption("X"), i6
        # unbekannter Kern: Fit verweigern statt Kopplung still weglassen; kompakte Schreibweise bleibt erhalten
        i6["nuclei"] = "2 14n 0.65"
        t_e.fit_current()
        assert i6["fit"] is None and i6["nuclei"] == "2 14n 0.65", i6
        assert EPRMOD.unknown_nuclei("2 14n 0.65; 4 1H 0.3") == ["2 14n 0.65"]
        assert EPRMOD.nuclei_with_a("2x14N 0.65", [0.7]) == "2 14N 0.7000", EPRMOD.nuclei_with_a("2x14N 0.65", [0.7])
        assert EPRMOD.nuclei_names("2x14N 0.65") == ["14N"], EPRMOD.nuclei_names("2x14N 0.65")
        # Dunkelspektrum vor dem Skalieren/Verschieben abziehen (hell = dunkel -> 0, keine Randlücken)
        el6, ed6 = asc6("l-9p400000.txt", B6 - 330), asc6("d-9p400000.txt", B6 - 330)
        el6["dark"] = ed6["uid"]
        w.epr_state.update(entries=[el6, ed6])
        I6 = t_e.processed(el6, {})[1]
        assert np.isfinite(I6).all() and np.max(np.abs(I6)) < 1e-9, (np.isnan(I6).sum(), np.nanmax(np.abs(I6)))
        # ausgeblendete Messung (7·S + 3·Bn): Simulationen behalten Faktor 7 und den Normierungsfaktor der Messung
        Bn6 = EPRMOD.iso_spectrum(B6, 9.4, 2.010, [], 0.5)
        np.savetxt(p6 / "simn.txt", np.column_stack([B6, Bn6]), delimiter="\t")
        em6 = asc6("m-9p400000.txt", 7 * S6 + 3 * Bn6)
        np.savetxt(p6 / "sim6.txt", np.column_stack([B6, S6]), delimiter="\t")
        es6, eb6 = t_e._new("sim", str(p6 / "sim6.txt")), t_e._new("sim", str(p6 / "simn.txt"))
        es6["fit_to"] = eb6["fit_to"] = em6["uid"]
        w.epr_state.update(entries=[em6, es6, eb6], ref_freq=0.0)
        em6["visible"] = False
        cur6 = t_e.figure_spec()[1]
        assert abs(np.nanmax(np.abs(cur6[0]["y"])) / np.nanmax(np.abs(7 * S6)) - 1) < 0.01, np.nanmax(cur6[0]["y"])
        w.epr_state["normalize"] = True
        cur6 = t_e.figure_spec()[1]
        want6 = 7 * np.nanmax(np.abs(S6)) / np.nanmax(np.abs(7 * S6 + 3 * Bn6))
        assert abs(np.nanmax(np.abs(cur6[0]["y"])) / want6 - 1) < 0.01, (np.nanmax(cur6[0]["y"]), want6)
        em6["visible"] = True
        assert "normalized" in t_e.methods_text(), t_e.methods_text()
        # negatives Gewicht: keine Anteile in %, sondern Hinweis
        en6, sn6 = asc6("n-9p400000.txt", S6 - 0.2 * Bn6), t_e._new("sim", str(p6 / "simn.txt"))
        sa6 = t_e._new("sim", str(p6 / "sim6.txt"))
        sa6["fit_to"] = sn6["fit_to"] = en6["uid"]
        info6 = []
        t_e.components([en6, sa6, sn6], {}, info6)
        assert info6 and T("epr_negative_weight", name=en6["label"]) in info6 and "%" not in info6[0], info6
        # Export: gleiche bereinigte Beschriftung überschreibt nicht
        e2_6, e3_6 = asc6("m2-9p400000.txt", 2 * S6), asc6("m3-9p400000.txt", 3 * S6)
        em6["label"], e2_6["label"], e3_6["label"] = f"sample_A_{e3_6['uid']}", "sample A", "sample/A"
        w.epr_state.update(entries=[em6, e2_6, e3_6], normalize=False)
        txt6 = [f for f in t_e.export_to(out / "epr_p6") if f.suffix == ".txt" and "sample" in f.name]
        assert len(set(txt6)) == 3, txt6
        # I·T ohne Temperatur: Spektrum auslassen statt I und I·T zu mischen
        w.epr_state.update(entries=[e6, kin_entries[0]], mode="intensity", int_method="pp", int_x="index",
                           int_curie=True, int_window="")
        assert len(t_e.intensity_series([e6, kin_entries[0]])[0]) == 1
        # Zerfall: Kurve = gefittetes Modell an den Messpunkten, Band aus der Kovarianz (nicht aus Randwerten)
        w.epr_state.update(entries=kin_entries[::10], int_method="gauss", int_window="470-510", int_x="T",
                           kin_fit=True, kin_norm="lsq")
        sp6 = t_e.intensity_spec(kin_entries[::10])
        r6 = t_e.kin_result
        x6 = sp6[1][0]["x"]
        fit6 = sp6[1][1]
        assert np.allclose(np.interp(x6, fit6["x"], fit6["y"]),
                           EPRMOD.decay_model(x6, r6["A"], r6["Ea"], r6["beta"]), atol=1e-9), fit6
        band6 = EPRMOD.decay_band(x6, r6, r6["beta"], (x6 - x6[0]) / r6["beta"])   # gezeichnetes Band = Kovarianz
        assert np.allclose(sp6[1][2]["y"], band6[0], atol=1e-9) and np.allclose(sp6[1][3]["y"], band6[1], atol=1e-9)
        rng6 = np.random.default_rng(6)
        Tn6 = np.arange(100.0, 135.01, 0.5)
        yn6 = EPRMOD.decay_model(Tn6, np.exp(45.0), 50e3, 0.5 / 120) + rng6.normal(0, 0.01, len(Tn6))
        rn6 = EPRMOD.fit_decay(Tn6, yn6, 0.5 / 120, norm="lsq")
        lo6, hi6 = EPRMOD.decay_band(Tn6, rn6, 0.5 / 120)
        smp6 = rng6.multivariate_normal([rn6["Ea"], rn6["lnA"]], rn6["cov"], 20000)
        k6 = int(np.argmin(np.abs(Tn6 - 116.5)))
        mc6 = np.std([EPRMOD.decay_model(Tn6[:k6 + 1], np.exp(b_), a_, 0.5 / 120)[-1] for a_, b_ in smp6])
        assert abs((hi6[k6] - lo6[k6]) / 2 / (rn6["t"] * mc6) - 1) < 0.1, ((hi6[k6] - lo6[k6]) / 2, rn6["t"] * mc6)
        # kein „± None“ im SI-Text, wenn das Intervall nicht bestimmbar ist
        t_e.kin_result = dict(Ea=4e4, lnA=34.5, A=1e15, ci_Ea=(np.nan, np.nan), ci_lnA=(np.nan, np.nan),
                              corr=np.nan, beta=0.01)
        assert "None" not in t_e.methods_text(), t_e.methods_text()
        # R1: ohne Min-Max mit freier Startintensität (unvollständiger Zerfall) und Integral über die Zeit (Haltephase)
        bR = 1 / 60

        def truth6(T_, t_):                         # dichtes Zeitraster, unabhängig vom Messraster
            tf_ = np.linspace(t_[0], t_[-1], 200001)
            return np.interp(t_, tf_, EPRMOD.decay_model(np.interp(tf_, t_, T_), np.exp(45.0), 100e3, bR, tf_))
        TR = np.arange(150.0, 300.0, 2.0)
        IR = truth6(TR, (TR - 150) / bR)
        kR = int(np.flatnonzero(IR <= 0.57)[0])
        rR = EPRMOD.fit_decay(TR[:kR + 1], 5 * IR[:kR + 1], bR, norm="lsq", time=(TR[:kR + 1] - 150) / bR)
        assert abs(rR["Ea"] / 100e3 - 1) < 0.005 and abs(rR["I0"] / 5 - 1) < 1e-3, rR
        t1R = np.arange(0, 4800.1, 120.0)
        t2R = t1R[-1] + np.arange(120, 1800.1, 120.0)
        t3R = t2R[-1] + np.arange(120, 4000.1, 120.0)
        tR = np.r_[t1R, t2R, t3R]
        TRh = np.r_[150 + bR * t1R, np.full(len(t2R), 150 + bR * t1R[-1]), 150 + bR * (t1R[-1] + t3R - t2R[-1])]
        rR = EPRMOD.fit_decay(TRh, truth6(TRh, tR), bR, norm="lsq", time=tR)
        assert abs(rR["Ea"] / 100e3 - 1) < 0.01, rR
        w.epr_state.update(entries=[], mode="spectra", kin_fit=False, int_method="double", int_window="",
                           int_curie=False, normalize=True)
        # Leser: Dezimalkomma, BOM, echtes Minus, ν aus dem Dateinamen, YFMT, fallende Achse, unbekannte Varianten
        (p6 / "komma.txt").write_text("B [mT];I\n" + "\n".join(f"{330 + k / 2:.1f};{1.25 + k:.2f}".replace(".", ",")
                                                              for k in range(12)), encoding="utf-8")
        rk6 = EPRMOD.read_ascii(p6 / "komma.txt")
        assert abs(rk6["B"][1] - 330.5) < 1e-9 and abs(rk6["I"][0] - 1.25) < 1e-9, (rk6["B"][:2], rk6["I"][:2])
        pm6 = p6 / "Cpd1p5_ox-9p404368-339p0230-10p0708.txt"
        pm6.write_text("﻿" + "\n".join(f"{330 + k:.1f}\t{MINUS if k % 2 else ''}{k}" for k in range(12)),
                       encoding="utf-8")
        rm6 = EPRMOD.read_ascii(pm6)
        assert len(rm6["B"]) == 12 and rm6["I"][1] == -1 and abs(rm6["meta"]["mw_GHz"] - 9.404368) < 1e-9, rm6

        def bes6(stem, extra=(), xwid=99.0, ny=1, yv=None, yfmt="D"):
            h = ["BSEQ\tLIT", "IKKF\tREAL", "XTYP\tIDX", "IRFMT\tD", "XPTS\t100", f"YPTS\t{ny}", "XMIN\t3300",
                 f"XWID\t{xwid}", "XUNI\t'G'", "MWFQ\t9.4e9"] + list(extra)
            if yv is not None:
                h += ["YTYP\tIGD", f"YFMT\t{yfmt}"]
                np.asarray(yv, {"D": "<f8", "F": "<f4"}[yfmt]).tofile(p6 / f"{stem}.YGF")
            (p6 / f"{stem}.dsc").write_text("\n".join(h), encoding="latin-1")
            np.arange(100 * ny, dtype="<f8").tofile(p6 / f"{stem}.DTA")
            return p6 / f"{stem}.dsc"
        assert [s["meta"]["y"] for s in EPRMOD.read_bes3t(bes6("yf", ny=2, yv=[100, 110], yfmt="F"))] == [100.0, 110.0]
        sd6 = EPRMOD.read_bes3t(bes6("down", xwid=-99.0))[0]
        assert sd6["B"][0] < sd6["B"][-1] and sd6["I"][0] == 99, (sd6["B"][:2], sd6["I"][:2])
        for ex6 in (["XUNI\t'MHz'"], ["BSEQ\tVAX"], ["IRFMT\tZ"]):
            try:
                EPRMOD.read_bes3t(bes6("bad", ex6))
            except ValueError:
                continue
            raise AssertionError(f"BES3T-Variante angenommen: {ex6}")
        # BES3T-Partner gemeinsam übernehmen: gleiche .dsc, andere .DTA -> eigener Stamm für beide
        for sub6, v6 in (("a", 1.0), ("b", 2.0)):
            (p6 / sub6).mkdir(exist_ok=True)
            (p6 / sub6 / "same.dsc").write_text("IKKF\tREAL\nXPTS\t10\nXMIN\t3300\nXWID\t9\nBSEQ\tLIT\nIRFMT\tD",
                                                encoding="latin-1")
            np.full(10, v6, "<f8").tofile(p6 / sub6 / "same.DTA")
        got6 = [PRJ.import_group(p6 / "proj", [p6 / s6 / "same.dsc", p6 / s6 / "same.DTA"], "data/epr")
                for s6 in ("a", "b")]
        assert [float(EPRMOD.read_bes3t(g_[0])[0]["I"][0]) for g_ in got6] == [1.0, 2.0], got6
        assert PRJ.import_group(p6 / "proj", [p6 / "b" / "same.dsc", p6 / "b" / "same.DTA"], "data/epr") == got6[1]
        # Bereiche mit Dezimalkomma; Glätten breiter als der Sweep; EasySpin: Ausdrücke, mehrere Anweisungen, [Dx Dy Dz]
        assert EPRMOD.ranges("330,5-340,5") == [(330.5, 340.5)], EPRMOD.ranges("330,5-340,5")
        assert EPRMOD.ranges("330-340,480-485") == [(330.0, 340.0), (480.0, 485.0)]
        assert EPRMOD.ranges("330,5-340,5; 480-485") == [(330.5, 340.5), (480.0, 485.0)]
        assert len(EPRMOD.smooth(np.linspace(330, 331, 101), np.ones(101), 2)) == 101
        pe6 = EPRMOD.parse_easyspin("Sys.D = [1.1 0.16]*1000;\nSys.g = 2.0023; Sys.lwpp = 1.5; % Kommentar")
        assert pe6["Sys.D"] == "[1.1 0.16]*1000" and pe6["Sys.g"] == [2.0023] and pe6["Sys.lwpp"] == [1.5], pe6
        z6 = EPRMOD.zfs_summary({"Sys.D": [-100.0, -100.0, 200.0]})
        assert "(300 MHz)" in z6 and "(0 MHz)" in z6, z6
        # Codex-Gegenlesen: ΔB folgt einer später geänderten ν_ref; Toleranz der Handänderung; Leser/Import/Glätten
        mwC = 9.4
        BC = np.linspace(330.0, 345.0, 3001)
        erC = asc6("refC-9p400000.txt", EPRMOD.iso_spectrum(BC, mwC, 2.0055, [], 0.15), B=BC + 0.25)
        esC = asc6("smpC-9p400000.txt", EPRMOD.iso_spectrum(BC, mwC, 2.003, [], 0.15), B=BC + 0.25)
        w.epr_state.update(entries=[erC, esC], bl_order=-1, smooth=0.0, ref_freq=0.0, calib_uid=None, calib_g=2.0055,
                           exclude="", mode="spectra")
        t_e.fill_table()
        t_e.table.selectRow(0)
        t_e.add_iso()
        riC = w.epr_state["entries"][-1]
        t_e.table.selectRow(2)
        t_e.fit_current()
        w.epr_state.update(calib_uid=riC["uid"], ref_freq=9.8)
        t_e.fill_table()
        t_e.table.selectRow(1)
        t_e.add_iso()
        siC = w.epr_state["entries"][-1]
        t_e.table.selectRow(3)
        t_e.fit_current()
        assert abs(siC["g"] - 2.003) < 2e-5, siC
        siC["lw"] = 0.10049                                   # angezeigt 0.100; Handänderung auf 0.101 zählt
        t_e.table.selectRow(2)
        t_e.table.selectRow(3)
        t_e.sp_lw.setValue(0.101)
        assert siC["fit"] is None, siC
        w.epr_state.update(entries=[], calib_uid=None, ref_freq=0.0)
        assert EPRMOD.parse_easyspin("Sys.D=[1000-100 20]")["Sys.D"] == "[1000-100 20]"
        assert EPRMOD.parse_easyspin("Sys.D = [1000 -100]")["Sys.D"] == [1000.0, -100.0]
        (p6 / "kp.txt").write_text("\n".join(f"{330 + k},1.25, 99" for k in range(12)), encoding="utf-8")
        rkp = EPRMOD.read_ascii(p6 / "kp.txt")
        assert len(rkp["B"]) == 12 and rkp["B"][0] == 330 and rkp["I"][0] == 1.25, rkp
        np.savetxt(p6 / "s-9p40.txt", np.column_stack([B6[:20], S6[:20]]))
        assert EPRMOD.read_ascii(p6 / "s-9p40.txt")["meta"]["mw_GHz"] == 9.4
        for sub6, nm6, x6m in (("c", "same.DSC", 3300), ("d", "same.dsc", 3400)):   # nur die Beschreibung anders
            (p6 / sub6).mkdir(exist_ok=True)
            (p6 / sub6 / nm6).write_text(f"IKKF\tREAL\nXPTS\t10\nXMIN\t{x6m}\nXWID\t9\nBSEQ\tLIT\nIRFMT\tD",
                                         encoding="latin-1")
            np.ones(10, "<f8").tofile(p6 / sub6 / "same.DTA")
        gC = [PRJ.import_group(p6 / "projC", [p6 / s6 / n6, p6 / s6 / "same.DTA"], "data/epr")
              for s6, n6 in (("c", "same.DSC"), ("d", "same.dsc"))]
        assert [EPRMOD.read_bes3t(g_[0])[0]["B"][0] for g_ in gC] == [330.0, 340.0], gC
        eS = asc6("sm-9p400000.txt", S6)                       # Fenster < 3 Punkte: keine Glättung, kein Satz dazu
        w.epr_state.update(entries=[eS], smooth=0.01, normalize=False)
        assert "smoothed" not in t_e.methods_text(), t_e.methods_text()
        w.epr_state["smooth"] = 0.05
        assert "moving average over 0.05 mT" in t_e.methods_text(), t_e.methods_text()
        w.epr_state.update(entries=[], smooth=0.0, normalize=True)
        report.append("EPR audit P6 (reference frequency, dark axis, decay band, readers, import): OK")

        # Farbe der Lösung: Farbmetrik (Weiß, Grau, Gelb bei Blauabsorption) und Anzeige im ε-Tab
        L_ = COLOR.LAM
        assert COLOR.to_hex(COLOR.absorbance_to_rgb(0 * L_)[0]) == "#FFFFFF"
        assert COLOR.to_hex(COLOR.absorbance_to_rgb(0 * L_ + 0.3)[0]) == "#BCBCBC"     # T = 50 % -> sRGB 188
        r_, g_, b_ = COLOR.absorbance_to_rgb(np.exp(-0.5 * ((L_ - 450) / 30) ** 2))[0]
        assert r_ > 0.9 and g_ > 0.9 and b_ < 0.5, (r_, g_, b_)                        # absorbiert Blau -> Gelb
        w.glob.update(color_show=True, color_d=1.0)
        w.series["TEST-A"].pop("color_c", None)
        w.apply_globals()
        w.render()
        assert w.lbl_color.isVisibleTo(w) and "#" in w.lbl_color_info.text(), w.lbl_color_info.text()
        hx1 = w.lbl_color_info.text().split()[0]
        c_max = max(s_["c"] for s_ in w.groups["TEST-A"])
        assert f"c = {c_max:g}" in w.lbl_color_info.text(), w.lbl_color_info.text()
        w.sp_color_c.setValue(4 * c_max)                    # konzentrierter -> dunkler
        hx2 = w.lbl_color_info.text().split()[0]
        lum = lambda h: sum(int(h[i:i + 2], 16) for i in (1, 3, 5))
        assert lum(hx2) < lum(hx1) and w.series["TEST-A"]["color_c"] == 4 * c_max, (hx1, hx2)
        w.glob["color_show"] = False
        w.series["TEST-A"].pop("color_c", None)
        w.apply_globals()
        assert not w.lbl_color.isVisibleTo(w)
        report.append("solution colour: OK")

        # Layout ε-Tab: Vorschau statt Leer-Zustand, eingeklappte Bereiche melden Abweichungen vom Standard
        assert w.preview_stack.currentIndex() == 1
        w.sec_eval.btn.setChecked(False)
        assert not w.sec_eval.body.isVisibleTo(w)
        r2_old = w.sp_r2.value()
        w.sp_r2.setValue(DEFAULT_GLOB["min_r2"])
        w.sp_min.setValue(DEFAULT_GLOB["min_points"])
        w.cb_shoulders.setChecked(DEFAULT_GLOB["shoulders"])
        assert not w.sec_eval.changed
        w.sp_r2.setValue(0.95)
        assert w.sec_eval.changed and T("adv_changed") in w.sec_eval.btn.text()
        w.sp_r2.setValue(r2_old)
        w.sec_eval.btn.setChecked(True)
        assert w.sec_eval.body.isVisibleTo(w)
        labels = [lb for lb in w.eps_page.findChildren(QLabel)
                  if lb.text() in (T("molar_mass_short"), T("ymax"))]
        assert len({lb.minimumWidth() for lb in labels}) == 1 and labels[0].minimumWidth() > 0
        # „Alles auf Standard“ im ε-Tab: Einstellungen zurück, Schichtdicke und Exportgröße bleiben
        if w.current:
            w.glob.update(cutoff=2.0, ymax=2.5, path_length=0.1, shoulders=True, size="custom")
            w.series[w.current]["band_hidden"] = [600]
            w.reset_defaults()
            assert w.glob["cutoff"] == 1.0 and w.glob["ymax"] == 1.0 and not w.glob["shoulders"]
            assert w.glob["path_length"] == 0.1 and w.glob["size"] == "custom" and abs(w.sp_cut.value() - 1.0) < 1e-9
            assert w.series[w.current]["band_hidden"] == []
            w.glob.update(path_length=1.0, size="half_a4")
            w.apply_globals()
        report.append("eps tab layout (sections, empty state): OK")
        # Anleitung in leeren Tabs: Texte in beiden Sprachen, Link zeigt auf eine vorhandene README-Überschrift
        from elucida.app.i18n import TEXTS as TEXTS_
        heads_ = {lang_: [ln_[3:].strip() for ln_ in (ROOT / f_).read_text(encoding="utf-8").splitlines()
                          if ln_.startswith("## ")]
                  for lang_, f_ in (("de", "README.de.md"), ("en", "README.md"))
                  if (ROOT / f_).exists()}           # im gebauten Programm fehlt die README
        for key_, (de_, en_) in TABS.GUIDE_HEAD.items():
            assert not heads_ or (de_ in heads_["de"] and en_ in heads_["en"]), (key_, de_, en_)
            for part_ in ("title", "files", "steps"):
                assert T(f"guide_{key_}_{part_}") != f"guide_{key_}_{part_}", (key_, part_)
        for lang_ in heads_:                                       # Startbildschirm/Hilfe: Schnellstart-Link
            q_ = TEXTS_["readme_quickstart"][0 if lang_ == "de" else 1]
            assert q_ in heads_[lang_], (lang_, q_)
        assert TABS.gh_anchor("IR + EA tab") == "ir--ea-tab"
        assert TABS.gh_anchor("CV-Tab (Cyclovoltammetrie)") == "cv-tab-cyclovoltammetrie"
        guide_ = w.cv_tab.preview_stack.widget(0).findChild(TABS.EmptyGuide)
        assert guide_ is not None and T("guide_cv_title") in guide_.label.text()
        w.resize(1440, 900)
        w.show()
        for tab_, key_ in ((w.syn_tab, "syn"), (w.nmr_tab, "nmr"), (w.cv_tab, "cv")):   # überall gleich und mittig
            w.main_tabs.setCurrentWidget(tab_)
            pump(0.1)
            if hasattr(tab_, "ensure_rendered"):                # erst zeichnen (verzögert beim Anzeigen), dann messen
                tab_.ensure_rendered()
            pump(0.3)
            assert tab_.preview_stack.currentIndex() == 0, key_
            g_ = tab_.preview_stack.widget(0)
            g_ = g_ if isinstance(g_, TABS.EmptyGuide) else g_.findChild(TABS.EmptyGuide)
            g_.fit_width()                                       # Breite nach der Schrift des Systems, Layout fertig
            g_.layout().activate()
            pump(0.1)
            assert g_ is not None and T(f"guide_{key_}_title") in g_.label.text(), key_
            b_ = g_.findChild(QPushButton)
            assert b_ is not None, key_
            top_, bottom_ = g_.label.geometry().top(), g_.height() - b_.geometry().bottom()
            assert abs(top_ - bottom_) <= 12, (key_, top_, bottom_, g_.size())
            mid_ = (g_.label.geometry().center().x(), b_.geometry().center().x())
            assert all(abs(m_ - g_.width() / 2) <= 12 for m_ in mid_), (key_, mid_, g_.width())
        # Meldungen als Einblendung: keine Statusleiste unter den Fußleisten, Knöpfe oben und unten gleich weit weg
        w.statusBar().showMessage("selftest", 5000)
        assert w._toast.isVisible() and w._toast.text() == "selftest"     # sofort (andere Meldungen folgen beim pump)
        pump(0.1)
        bar_ = w.cv_tab.findChild(QFrame, "bottomBar")
        y_ = bar_.mapTo(w, bar_.rect().bottomLeft()).y()
        assert y_ >= w.height() - 3, (y_, w.height())
        btn_ = max((c_ for c_ in bar_.findChildren(QPushButton) if c_.isVisible()), key=lambda c_: c_.height())
        gap_ = (btn_.geometry().top(), bar_.height() - btn_.geometry().bottom() - 1)
        assert abs(gap_[0] - gap_[1]) <= 2, gap_
        w.statusBar().clearMessage()                                      # Signal kommt sofort (ohne pump: Autosave)
        assert not w._toast.isVisible(), (w._toast.text(), w.statusBar().currentMessage())
        report.append("empty-tab guides: OK")
        # Navigation: Gruppen oben statt Reiter; Knopf wechselt den Tab, Tabwechsel markiert den Knopf
        assert len(w.nav_buttons) == 12 and w.main_tabs.tabBar().isHidden()
        i_cv_ = w._tab_keys.index("cv")
        w.nav_buttons[i_cv_].click()
        assert w.main_tabs.currentWidget() is w.cv_tab
        w.main_tabs.setCurrentWidget(w.ov_tab)
        assert w.nav_buttons[w._tab_keys.index("overlay")].isChecked()
        # Einstellungen des gewählten Eintrags: ohne Auswahl ausgeblendet, stattdessen ein Hinweis
        w.ov_tab.table.clearSelection()
        w.ov_tab.table.setCurrentCell(-1, -1)
        w.ov_tab.show_details()
        assert w.ov_tab.det.isHidden() == (not w.ov_tab.det.isEnabled())
        assert w.toolbar.widgetForAction(w.act_si).property("accent")
        # Mausrad über einem Zahlenfeld in der linken Spalte: ohne Fokus scrollt die Spalte, der Wert bleibt
        from PySide6.QtCore import QPoint, QPointF
        from PySide6.QtGui import QWheelEvent
        sa_ = QScrollArea()
        inner_ = QWidget()
        lv_ = QVBoxLayout(inner_)
        spins_ = [QSpinBox(maximum=100, value=50) for _ in range(40)]
        for sp_ in spins_:
            lv_.addWidget(sp_)
        sa_.setWidget(inner_)
        sa_.setWidgetResizable(True)
        sa_.resize(200, 200)
        sa_.show()
        pump(0.2)
        sp0_ = spins_[2]

        def wheel_(target_):
            return QWheelEvent(QPointF(5, 5), QPointF(target_.mapToGlobal(QPoint(5, 5))), QPoint(0, 0), QPoint(0, -120),
                               Qt.NoButton, Qt.NoModifier, Qt.NoScrollPhase, False)
        QApplication.sendEvent(sp0_, wheel_(sp0_))
        assert sp0_.value() == 50 and sa_.verticalScrollBar().value() > 0, (sp0_.value(), sa_.verticalScrollBar().value())
        assert sp0_.focusPolicy() == Qt.StrongFocus
        sp0_.setFocus()
        QApplication.sendEvent(sp0_, wheel_(sp0_))
        assert sp0_.value() != 50 or not sp0_.hasFocus()      # mit Fokus wie gewohnt (offscreen evtl. ohne Fokus)
        sa_.close()
        report.append("navigation, detail hint, SI button, wheel guard: OK")
        # Synthese/NMR mit Fußleiste wie die anderen Tabs; „Aus Schema“ nur mit Schema; Exportmeldung mit Ordner
        assert w.syn_tab.btn_export.property("accent") and w.nmr_tab.btn_export.property("accent")
        assert w.syn_tab.btn_from_scheme.isEnabled() == (w.syn_tab.scheme_parts() is not None)
        w.exported([proj / "figures"])
        assert w.last_export == [str(proj / "figures")]
        # linke Spalte: Karten eingeklappt mit „(geändert)“, Ergebnisse als Reiter neben der Abbildung
        fc_ = next(f_ for f_ in w.cv_tab.folds if f_.group is w.cv_tab.grp_common)
        fc_.set_open(False, remember=False)
        assert fc_.body.isHidden() and fc_.group.title().startswith("▶")
        w.cv_state["size"] = "custom"
        fc_.refresh()
        assert T("adv_changed") in fc_.group.title(), fc_.group.title()
        w.cv_state["size"] = "half_a4"
        fc_.refresh()
        assert T("adv_changed") not in fc_.group.title()
        fc_.set_open(True, remember=False)
        assert not fc_.body.isHidden()
        assert w.cv_tab.right_tabs is not None and w.cv_tab.right_tabs.tabText(0) == T("tab_figure")
        assert [w.xray_tab.right_tabs.tabText(i_) for i_ in range(3)] == [T("tab_figure"), T("xr_caption"), T("xr_tables")]
        assert w.dft_tab.right_tabs.count() == 2 and w.hrms_tab.right_tabs is None
        # Ergebnisseiten überall gleich: ε neben der Abbildung, Karte füllt die Seite, Hauptknopf blau rechts
        assert [w.tabs.tabText(i_) for i_ in range(w.tabs.count())] == [T("tab_figure"), T("tab_results"),
                                                                       T("tab_si_line")]
        for rt_, j_ in ((w.cv_tab.right_tabs, 1), (w.xray_tab.right_tabs, 2), (w.dft_tab.right_tabs, 1), (w.tabs, 2)):
            card_ = rt_.widget(j_).widget().findChild(QGroupBox, "resultCard")
            assert card_ is not None and card_.title() == "", (rt_.tabText(j_),)
            acc_ = [b_ for b_ in card_.findChildren(QPushButton) if b_.property("accent")]
            assert len(acc_) == 1 and not acc_[0].icon().isNull(), (rt_.tabText(j_), len(acc_))
        # Raster: jedes sichtbare Zahlenfeld der linken Spalten genau STYLE.FIELD_W breit (das Stylesheet darf die
        # feste Breite nicht nachträglich aufheben – so liefen die Felder früher auseinander)
        w.show()
        for i_ in range(w.main_tabs.count()):
            w.main_tabs.setCurrentIndex(i_)
            pump(0.05)
            page_ = w.main_tabs.widget(i_)
            odd_ = [sp_.width() for sp_ in page_.findChildren(QAbstractSpinBox)
                    if sp_.isVisibleTo(page_) and sp_.property("fieldw") and sp_.width() != STYLE.FIELD_W]
            assert not odd_, (w._tab_keys[i_], odd_)
        # linke Spalte nie abgeschnitten: Inhalt passt in die sichtbare Breite, auch auf kleinerem Bildschirm
        for size_ in ((1600, 1000), (1280, 800)):
            w.resize(*size_)
            for i_ in range(w.main_tabs.count()):
                w.main_tabs.setCurrentIndex(i_)
                pump(0.1)
                page_ = w.main_tabs.widget(i_)
                cf_ = getattr(page_, "column_fit", None) or (w.column_fit if w._tab_keys[i_] == "eps" else None)
                assert cf_ is not None, w._tab_keys[i_]
                cf_.fit()
                pump(0.05)
                need_ = cf_.content.minimumSizeHint().width()
                have_ = cf_.scroll.viewport().width()
                assert need_ <= have_ + 1, (w._tab_keys[i_], size_, need_, have_)
        # Neuaufbau (Projekt öffnen): ausstehende Anpassung einer gelöschten Spalte darf nicht abstürzen
        sp_ = QSplitter()
        col_, sc_ = QWidget(), QScrollArea()
        cont_ = QWidget()
        sc_.setWidget(cont_)
        QVBoxLayout(col_).addWidget(sc_)
        sp_.addWidget(col_)
        sp_.addWidget(QWidget())
        cf_ = STYLE.fit_column(cont_, sc_, None, col_, sp_)
        QApplication.sendEvent(cont_, QEvent(QEvent.LayoutRequest))       # Anpassung steht aus
        hook_, errs_ = sys.excepthook, []
        sys.excepthook = lambda *a: errs_.append(a[1])
        try:
            import shiboken6 as _shb
            _shb.delete(sp_)                                                  # wie der Neuaufbau der Tabs
            pump(0.1)
            cf_.fit()
            cf_._moved(10, 1)
        finally:
            sys.excepthook = hook_
        assert not errs_, errs_
        report.append("left column (folds, changed mark, result tabs, field grid, never clipped): OK")

        # Erscheinungsbild: Fusion, Palette je Modus, Akzentbuttons, Symbole; Ergebnisse bleiben erhalten
        w.render()
        n_rows = w.table.rowCount()
        for mode in ("dark", "light"):
            w.switch_theme(mode)
            pal = QApplication.instance().palette()
            assert STYLE.base_style == "fusion"
            assert pal.color(pal.ColorRole.Window).name().upper() == STYLE.COLORS[mode]["window"], mode
            assert w.btn_export.property("accent") and w.btn_open.property("accent")
            assert not w.btn_export.icon().isNull(), "icon"
            assert w.table.rowCount() == n_rows, (w.table.rowCount(), n_rows)
            css_ = QApplication.instance().styleSheet()       # Checkboxen mit eigenem Rand (dunkel sonst unsichtbar)
            assert "QCheckBox::indicator" in css_ and STYLE.COLORS[mode]["check_border"] in css_, mode
            ic_ = w.app_icon()                                # App-Symbol hell/dunkel passend zum Theme
            assert ic_.exists() and ic_.name == ("icon-dark.png" if mode == "dark" else "icon.png"), (mode, ic_)
            assert not w.windowIcon().isNull(), mode
        assert w.findChildren(QFrame, "card") and not [g_ for g_ in w.eps_page.findChildren(QGroupBox)
                                                        if g_.objectName() != "resultCard"]
        w.switch_theme("system")
        # Werkzeugleiste (nach Neuaufbau nur eine; Projekt, Undo, Verbindung, SI) und Listenknöpfe mit Symbol
        bars = w.findChildren(TABS.QWidget, "mainToolbar")
        acts = [a for a in w.toolbar.actions() if not a.isSeparator()]
        assert len(bars) == 1 and len(acts) == 8 and all(not a.icon().isNull() for a in acts), (len(bars), len(acts))
        opened_ = []                                         # Projektordner im Dateimanager (hier nur abgefangen)
        orig_open_ = QDesktopServices.openUrl
        QDesktopServices.openUrl = staticmethod(lambda u: opened_.append(u.toLocalFile()) or True)
        try:
            assert w.act_folder in acts and w.act_folder.isEnabled()
            w.act_folder.trigger()
        finally:
            QDesktopServices.openUrl = orig_open_
        assert opened_ and Path(opened_[0]) == Path(w.project_dir), opened_
        arrows = [b for b in w.epr_tab.findChildren(QPushButton) if b.toolTip() == T("move_up")]
        assert arrows and arrows[0].text() == "" and not arrows[0].icon().isNull()
        report.append("appearance (Fusion, light/dark, toolbar): OK")

        # Umbenennungen UVVisTool -> Spectravolt (1.16) -> Elucida (2.0): Einstellungen/Datenordner einmalig übernehmen
        old_s = QSettings(str(out / "old_settings.ini"), QSettings.IniFormat)
        old_s.setValue("lang", "en")
        old_s.setValue("recent", [str(out / "p1"), str(out / "p2")])
        new_s = QSettings(str(out / "new_settings.ini"), QSettings.IniFormat)
        new_s.clear()
        assert migrate_settings(old_s, new_s) and new_s.value("lang") == "en" and len(new_s.value("recent")) == 2
        new_s.setValue("lang", "de")
        assert not migrate_settings(old_s, new_s) and new_s.value("lang") == "de"   # nur einmal, nichts überschreiben
        base_ = out / "appdata"
        shutil.rmtree(base_, ignore_errors=True)
        (base_ / "UVVisTool" / "models").mkdir(parents=True)
        (base_ / "UVVisTool" / "models" / "m.onnx").write_bytes(b"x")
        d_ = IMG.app_data_dir(base_)
        assert d_.name == "Elucida" and (d_ / "models" / "m.onnx").exists() and not (base_ / "UVVisTool").exists()
        shutil.rmtree(base_, ignore_errors=True)                 # Spectravolt hat Vorrang vor UVVisTool
        for n_, f_ in (("Spectravolt", "neu.onnx"), ("UVVisTool", "alt.onnx")):
            (base_ / n_ / "models").mkdir(parents=True)
            (base_ / n_ / "models" / f_).write_bytes(b"x")
        d_ = IMG.app_data_dir(base_)
        assert (d_ / "models" / "neu.onnx").exists() and not (base_ / "Spectravolt").exists(), list(d_.rglob("*"))
        assert APP_NAME == "Elucida" and OLD_APP_NAMES[0] == "Spectravolt" and UPD.REPO.endswith("/elucida")
        report.append("rename (settings + data folder migration): OK")

        # Verbindung auf Projektebene + Tab IR + MS (künstliche OPUS-, PDF- und Word-Dateien)
        import struct as _st
        import zipfile as _zf
        w.set_compound({"name": "TEST-1", "structure": str(cdxml), **compound_from_structure(cdxml)})
        assert w.compound["formula"] and w.compound["mw"] and "TEST-1" in w.windowTitle(), w.compound
        import time as _time
        t_end_ = _time.monotonic() + 60                     # Strukturen im Hintergrund (macOS-Intel-Runner langsam)
        while not all(st_.get("structure") for st_ in w.series.values()) and _time.monotonic() < t_end_:
            pump(0.1)
        assert all(st_.get("structure") for st_ in w.series.values()), ("Verbindung nicht als Vorgabe übernommen", {g_: st_.get("structure") for g_, st_ in w.series.items()}, w.compound)
        irdir = out / "input" / "irms"
        irdir.mkdir(parents=True, exist_ok=True)
        wn_ = np.linspace(3999.0, 400.0, 1800)
        bands_ = [(2961.2, 25, 12), (1678.4, 60, 8), (1501.3, 40, 6), (1105.0, 15, 10), (758.1, 55, 6), (3350.0, 20, 160)]
        T_ = 1.0 - sum(dp / 100 * np.exp(-((wn_ - c) / wd) ** 2) for c, dp, wd in bands_)
        def opus_par(pairs):
            out_ = b""
            for nm, typ, val in pairs:
                raw = _st.pack("<i", val) if typ == 0 else _st.pack("<d", val)
                out_ += nm.encode() + b"\0" + _st.pack("<hh", typ, len(raw) // 2) + raw
            return out_ + b"END\0" + _st.pack("<hh", 0, 0)
        par = opus_par([("NPT", 0, len(wn_)), ("FXV", 1, 3999.0), ("LXV", 1, 400.0), ("CSF", 1, 1.0)])
        dat = T_.astype("<f4").tobytes()
        off_par, off_dat = 24 + 12 * 2, 24 + 12 * 2 + len(par)
        head = b"\x0a\x0a\xfe\xfe" + _st.pack("<d", 920622.0) + _st.pack("<iii", 24, 2, 2)   # Version, Verz., max, n
        dirs = _st.pack("<BBBBii", 31, 20, 0, 0, len(par) // 4, off_par) + _st.pack("<BBBBii", 15, 20, 0, 0, len(wn_), off_dat)
        (irdir / "test.0").write_bytes(head + dirs + par + dat)
        d_ir = IRMS.read_opus(irdir / "test.0")
        assert d_ir["kind"] == "T" and abs(d_ir["y"].max() - 100) < 1 and IRMS.is_opus(irdir / "test.0")
        pk_ = IRMS.pick_peaks(d_ir["wn"], d_ir["y"], "T")
        got_ = [round(p_["wn"]) for p_ in pk_]
        assert all(any(abs(g_ - c) <= 1 for g_ in got_) for c, _, _ in bands_), got_
        br_ = next(p_ for p_ in pk_ if abs(p_["wn"] - 3350) < 5)
        assert IRMS.intensity_tag(br_).startswith("br"), IRMS.intensity_tag(br_)
        doc_ = pymupdf.open()
        pg_ = doc_.new_page()
        for i_, v_ in enumerate(["2961.20", "1678.41", "1501.33", "500", "1000", "Transmittance [%]", "80"]):
            pg_.insert_text((50, 60 + 20 * i_), v_)
        doc_.save(str(irdir / "test_ir.pdf"))
        assert IRMS.read_ir_pdf(irdir / "test_ir.pdf") == [2961.2, 1678.41, 1501.33]
        def emf(texts):
            out_ = b""
            for t_ in texts:
                body = t_.encode("utf-16-le")
                size = 76 + len(body) + (-(76 + len(body)) % 4)
                rec = bytearray(size)
                _st.pack_into("<II", rec, 0, 84, size)
                _st.pack_into("<II", rec, 44, len(t_), 76)
                rec[76:76 + len(body)] = body
                out_ += bytes(rec)
            return out_
        xml_ = "".join(f"<w:p><w:r><w:t>{t_}</w:t></w:r></w:p>" for t_ in
                       ("sample name:", "TEST-1", "molecular formula:", "C18H20N2", "exact mass:", "264.1626",
                        "method of Ionisation:", "ESI ☒ APCI ☐ HESI ☐"))
        with _zf.ZipFile(irdir / "ms.docx", "w") as z_:
            z_.writestr("word/document.xml", f"<w:document><w:body>{xml_}</w:body></w:document>")
            z_.writestr("word/media/image1.emf", emf(["Chemical Formula:", "C", "18", "Exact Mass:", "264,1626"]))
            z_.writestr("word/media/image2.emf", emf(["265.1703", "266.1735", "x.d: +MS, 0.1-0.4min", "m/z"]))
        rep_ = IRMS.read_ms_docx(irdir / "ms.docx")
        assert rep_["formula"] == "C18H20N2" and rep_["mz"] == [265.1703, 266.1735] and rep_["polarity"] == "+", rep_
        plain_, html_, calc_ = IRMS.ms_line("C18H20N2", "M+H", 265.1703)
        assert plain_ == "HR-MS-ESI(+) calc. C18H21N2+ [M+H]+ 265.1699; found 265.1703." and "<sub>21</sub>" in html_
        assert abs(IRMS.ion("C18H20N2", "M+Na")[2] - 287.1519) < 1e-4 and IRMS.ion("C18H20N2", "M-H")[1] == -1
        t_ir = w.irms_tab
        t_ms = w.hrms_tab
        assert w.main_tabs.indexOf(t_ir) + 1 == w.main_tabs.indexOf(t_ms)      # IR + EA vor HRMS (wie die SI)
        t_ir.load_ir(str(irdir / "test.0"))
        t_ms.load_ms(str(irdir / "ms.docx"))
        w.compound["formula"] = "C18H20N2"
        t_ir.render()
        t_ms.render()
        assert t_ir._ir_line[0].startswith("IR [cm-1]: ṽ = 3350, 2961, 1678, 1501, 1105, 758"), t_ir._ir_line[0]
        assert t_ms._ms_line[0].endswith("265.1699; found 265.1703."), t_ms._ms_line[0]
        assert t_ms.F is None                                    # ohne Rohdaten kein Spektrum
        w.irms_state["ir_off"] = [1105]
        w.irms_state["with_int"] = True
        t_ir.render()
        assert "1105" not in t_ir._ir_line[0] and "(br " in t_ir._ir_line[0], t_ir._ir_line[0]
        # IR: Absorbanz-Umschalter, y-Bereich (%T 0-100, exakt), Labels senkrecht unter dem Peakminimum
        a_ = IRMS.to_absorbance([100.0, 10.0, 1.0, 0.0, -1.0])
        assert np.allclose(a_[:3], [0, 1, 2]) and np.isnan(a_[3:]).all(), a_
        assert np.allclose(IRMS.to_transmission(a_[:3]), [100, 10, 1])
        yt_ = np.linspace(40.0, 100.0, 50)
        lim_, rej_ = IRMS.y_limits(yt_, "T")                 # auto: oben immer 100
        assert not rej_ and lim_[1] == 100.0 and lim_[0] < 40, lim_
        assert IRMS.y_limits(yt_, "T", (20.0, 75.0)) == ([20.0, 75.0], False)       # exakt, keine Erweiterung
        assert IRMS.y_limits(yt_, "T", (-5.0, 150.0)) == ([0.0, 100.0], False)      # %T strikt 0-100
        lim_, rej_ = IRMS.y_limits(yt_, "T", (80.0, 50.0))  # min >= max -> auto (früher verschwand das Spektrum)
        assert rej_ and lim_[1] == 100.0, lim_
        assert IRMS.y_limits([0.1, 0.8], "A")[0][0] == 0.0
        assert t_ir.sp_ymax.isHidden() or not t_ir.sp_ymax.isVisibleTo(t_ir)
        assert (t_ir.sp_ymin.minimum(), t_ir.sp_ymax2.maximum(), t_ir.sp_ymin.decimals()) == (0.0, 100.0, 0)
        assert t_ir.mode() == "T" and t_ir.cb_yauto.isChecked()
        w.irms_state.update(ymax=75.0, y_auto=True)         # altes Projekt mit „y-Achse max“
        t_ir.migrate_ymax()
        assert (w.irms_state["ymax"], w.irms_state["y_auto"], w.irms_state["y_max"]) == (0.0, False, 75.0), w.irms_state
        assert 0.0 <= w.irms_state["y_min"] < 75.0
        w.irms_state.update(y_auto=True, y_min=0.0, y_max=100.0)
        line_t_ = t_ir._ir_line
        w.irms_state["ir_labels"] = True
        t_ir.render()
        rend_ = t_ir.F._renderer()
        bbs_ = {k_: a_.get_window_extent(rend_) for k_, a_ in t_ir.F.ann_items.items()}
        assert bbs_ and all(abs(a_.get_rotation() - 90) < 1e-6 for a_ in t_ir.F.ann_items.values())
        ks_ = list(bbs_)
        assert not any(bbs_[k1].overlaps(bbs_[k2]) for i_, k1 in enumerate(ks_) for k2 in ks_[i_ + 1:]), "IR-Labels überlappen"
        ab_ = t_ir.F.ax.get_window_extent(rend_)
        for k_, ann_ in t_ir.F.ann_items.items():
            bb_ = bbs_[k_]
            yd_ = t_ir.F.ax.transData.transform([ann_.xy])[0][1]     # Peakminimum in Pixeln
            assert bb_.y1 < yd_ and bb_.y0 >= ab_.y0 - 1, (k_, bb_, yd_)
            xd0_, xd1_ = t_ir.F.ax.transData.inverted().transform([[bb_.x0, 0], [bb_.x1, 0]])[:, 0]
            m_ = (d_ir["wn"] <= max(xd0_, xd1_)) & (d_ir["wn"] >= min(xd0_, xd1_))
            if m_.any():                                   # Linie über die ganze Labelbreite oberhalb des Labels
                assert t_ir.F.ax.transData.transform([[0, d_ir["y"][m_].min()]])[0][1] > bb_.y1, k_
        t_ir.cmb_mode.setCurrentIndex(t_ir.cmb_mode.findData("A"))
        t_ir.render()
        assert w.irms_state["ir_mode"] == "A" and t_ir._ir_line == line_t_, t_ir._ir_line   # Peaks wie aus %T
        assert t_ir.F.cfg["ylabel"] == "absorbance" and t_ir.sp_ymax2.maximum() == 10.0
        assert t_ir.F.ax.get_ylim()[0] == 0.0                  # Absorbanz: unten 0
        rend_ = t_ir.F._renderer()
        for ann_ in t_ir.F.ann_items.values():                # Absorbanz: Labels über dem Maximum
            assert ann_.get_window_extent(rend_).y0 > t_ir.F.ax.transData.transform([ann_.xy])[0][1]
        t_ir.cmb_mode.setCurrentIndex(t_ir.cmb_mode.findData("T"))
        t_ir.render()
        assert t_ir.F.ax.get_ylim()[1] == 100.0                # %T auto: oben fest 100, unten Platz für Labels
        t_ir.cb_yauto.setChecked(False)                         # Felder übernehmen den gezeigten Bereich
        assert not w.irms_state["y_auto"] and w.irms_state["y_max"] == 100.0
        t_ir.sp_ymin.setValue(20.0)
        t_ir.render()
        assert np.allclose(t_ir.F.ax.get_ylim(), (20.0, 100.0)), t_ir.F.ax.get_ylim()   # exakt, nicht erweitert
        rend_ = t_ir.F._renderer()
        assert all(t_ir.F.ax.get_window_extent(rend_).y0 - 1 <= a_.get_window_extent(rend_).y0
                   for a_ in t_ir.F.ann_items.values())          # Labels innerhalb geklemmt
        t_ir.sp_ymax2.setValue(10.0)                             # max unter min -> min weicht aus
        assert w.irms_state["y_min"] < w.irms_state["y_max"] == 10.0, w.irms_state
        w.irms_state.update(ir_labels=False, ir_mode=None, ymax=0.0, y_auto=True, y_min=0.0, y_max=100.0)
        t_ir.sync_widgets()
        t_ir.render()
        report.append("IR: Absorbanz-Umschalter, y-Bereich (0-100 %T, exakt), Labels unter Peakminimum: OK")
        sp_real_ = t_ir.spectrum                                # Absorbanzdatei: Standard Transmission (Feedback)
        t_ir.spectrum = lambda: {**sp_real_(), "kind": "A"}
        try:
            assert IRMS.IRMSTab.DEFAULT_STATE["ir_mode"] == "T" and w.irms_state.get("ir_mode") is None
            assert t_ir.mode() == "A"                           # altes Projekt (None = wie die Datei) bleibt
            w.irms_state["ir_mode"] = IRMS.IRMSTab.DEFAULT_STATE["ir_mode"]
            assert t_ir.mode() == "T"
        finally:
            del t_ir.spectrum
            w.irms_state["ir_mode"] = None
        report.append("IR: Standard Transmission auch fuer Absorbanzdateien, alte Projekte unveraendert: OK")
        w.compound["formula"] = "C18H22N2"                 # falsche Verbindung -> Warnung
        t_ms.render()
        assert "⚠" in t_ms.lbl_ms_info.text(), t_ms.lbl_ms_info.text()
        files = t_ir.export_to(out / "irms") + t_ms.export_to(out / "hrms")
        assert sum(f.name.endswith("_SI.txt") for f in files) == 2 and all(f.exists() for f in files)
        slots_ = PRJ._path_slots(w.get_state())
        assert any(k_ == "ir_file" for _, k_, _ in slots_) and any(k_ == "ms_file" for _, k_, _ in slots_)
        # ChemDraw-Objekt im Bericht: Abgleich und Übernahme der Verbindung (künstliche Struktur, kein CDX-Schreiber)
        from rdkit import Chem as _Chem
        rep_ = t_ms.ms_report()
        rep_["structure"] = {"formula": "C18H20N2", "mw": 264.37, "n_fragments": 1,
                             "cdxml": _Chem.MolToCDXMLBlock(_Chem.MolFromSmiles("CC(C)(C)c1ccc2nc(-c3ccccc3)n(C)c2c1"))
                             if hasattr(_Chem, "MolToCDXMLBlock") else None}          # macOS Intel: RDKit 2025.9
        w.compound = {}
        t_ms.render()
        assert t_ms.formula() == ("C18H20N2", "report_struct") == t_ir.formula(), t_ms.formula()
        assert not t_ms.btn_take.isHidden()
        rep_["formula"] = "C18H21N2"                       # Formelfeld vertippt -> Hinweis
        t_ms.render()
        assert "⚠" in t_ms.lbl_ms_info.text(), t_ms.lbl_ms_info.text()
        t_ms.take_compound()
        assert w.compound["formula"] == "C18H20N2" and w.compound["name"] == "TEST-1", w.compound
        if rep_["structure"]["cdxml"]:
            assert Path(w.compound["structure"]).exists() and "aus_MS-Bericht" in Path(w.compound["structure"]).name
        w.irms_state.update(ir_file=None, ir_pdf=None, ir_off=[], with_int=False)
        w.hrms_state.update(ms_file=None)
        w.compound = {}
        # alte Projekte (bis 1.20): MS-Teil im Zustand von IR + MS wird in den Tab HRMS übernommen
        st_old = w.get_state()
        st_old.pop("hrms")
        st_old["irms"].update(ms_file=str(irdir / "ms.docx"), adduct="M+Na")
        w.set_state(st_old)
        assert w.hrms_state["ms_file"] and w.hrms_state["adduct"] == "M+Na", w.hrms_state
        t_ir, t_ms = w.irms_tab, w.hrms_tab
        w.hrms_state.update(ms_file=None, adduct="M+H")
        w.irms_state.pop("ms_file", None)
        report.append("project compound + IR/HRMS lines: OK")

        # Bruker-Rohdaten (.d): künstlicher LC-Lauf, m/z über die Umkehrung der Kalibrierung in Flugzeit-Indizes
        import sqlite3 as _sq
        from matplotlib.backend_bases import MouseEvent
        ddir = irdir / "test.d"
        ddir.mkdir(exist_ok=True)
        cal_ = (24790.0, 0.2, 309.0, 158509.1, 2.54e-4, 0.0, -0.0011)   # Delay, Takt, c0, c1, c2, –, c6
        def to_idx(m):
            m1 = np.asarray(m, float) + cal_[6]
            return (cal_[2] + np.sqrt(1e12 * m1 / cal_[3]) + cal_[4] * m1 - cal_[0]) / cal_[1]
        calc_ = IRMS.ion("C18H20N2", "M+H")[2]
        rts_ = np.arange(0.0, 120.0, 10.0)
        offs_ = 1e-4 * np.array([3, -2, 4, -3, 2, 1, -1, 3, -4, 2, 0, 1])            # Streuung je Scan
        amp_ = 1e5 * np.exp(-0.5 * ((rts_ - 60) / 12) ** 2)
        baf_, rows_ = bytearray(b"\0" * 64), []
        for i_, rt_ in enumerate(rts_):
            mz_ = np.array([180.0634, calc_ + offs_[i_], calc_ + 1.00335, 400.1234])
            in_ = np.array([5e3, amp_[i_], 0.2 * amp_[i_], 8e3], "<f4")
            n_ = len(mz_)
            body = (_st.pack("<4I", n_, 16, 16 + 8 * n_, 16 + 12 * n_) + to_idx(mz_).astype("<f8").tobytes()
                    + in_.tobytes() + np.full(n_, 0.01, "<f4").tobytes())
            head_ = bytearray(52)
            _st.pack_into("<III", head_, 0, 52 + len(body), IRMS.BAF_LINE_MAGIC, 52)
            rows_.append((i_ + 1, float(rt_), 1, (0x10 << 56) | len(baf_), 0))
            baf_ += head_ + body
        rows_.append((99, 61.0, 1, (0x10 << 56) | 64, 0))                            # MS2 -> wird übersprungen
        (ddir / "analysis.baf").write_bytes(bytes(baf_))
        (ddir / "readme.txt").write_text("nicht ins Projekt")
        tf_ = _st.pack("<i7d", 10, *cal_)
        tf_ += b"\0" * (112 - len(tf_))
        (ddir / "analysis.sqlite").unlink(missing_ok=True)       # Selbsttest-Ordner kann wiederverwendet werden
        con_ = _sq.connect(ddir / "analysis.sqlite")
        con_.executescript("create table Properties (Key text, Value text); create table AcquisitionKeys "
                           "(Id integer, Polarity integer); create table Spectra (Id integer, Rt real, TransformatorId "
                           "integer, LineIndexId integer, AcquisitionKey integer); create table Steps (TargetSpectrum "
                           "integer); create table Transformators (Id integer, Blob blob);")
        con_.executemany("insert into Properties values (?, ?)", [("SchemaType", "Baf2Sql"), ("SampleName", "TEST-1"),
                                                                   ("InstrumentName", "timsTOF test")])
        con_.execute("insert into AcquisitionKeys values (0, 0)")
        con_.executemany("insert into Spectra values (?, ?, 1, ?, ?)", [(a, b, c, d) for a, b, _, c, d in rows_])
        con_.execute("insert into Steps values (99)")
        con_.execute("insert into Transformators values (1, ?)", (tf_,))
        con_.commit()
        con_.close()
        raw_ = IRMS.read_bruker_d(ddir)
        assert IRMS.is_bruker_d(ddir) and len(raw_["spectra"]) == 12 and raw_["polarity"] == "+", raw_["polarity"]
        assert np.abs(raw_["spectra"][6][1][1] - (calc_ + offs_[6])) < 1e-7          # Kalibrierung zurückgerechnet
        lc_ = IRMS.lc_found(raw_["spectra"], calc_)
        sel_ = slice(5, 8)                                                         # 50/60/70 s: >= 50 % des Maximums
        exp_ = float(np.average(calc_ + offs_[sel_], weights=amp_[sel_]))
        assert lc_["n"] == 3 and abs(lc_["mz"] - exp_) < 1e-6 and lc_["rt_lo"] == 50 and lc_["rt_hi"] == 70, lc_
        tf_bad = _st.pack("<i7d", 12, *cal_) + b"\0" * 80
        try:
            IRMS._baf_mz([1.0], tf_bad)
            raise AssertionError("unbekannte Kalibrier-Variante nicht abgelehnt")
        except ValueError:
            pass
        w.compound = {"formula": "C18H20N2"}
        t_ms.load_raw(str(ddir))
        t_ms.render()
        assert w.hrms_state["ms_raw"] and t_ms._ms_line[0].endswith(f"found {exp_:.4f}."), t_ms._ms_line[0]
        assert "mDa" in t_ms.lbl_ms_info.text() and t_ms.cmb_found.currentData() == "raw", t_ms.lbl_ms_info.text()
        assert abs(t_ms.sp_rt_lo.value() - 50 / 60) < 0.01 and abs(t_ms.sp_rt_hi.value() - 70 / 60) < 0.01
        if w.project_dir:                                  # aus dem .d nur Rohdaten + Index übernommen
            assert sorted(x.name for x in Path(w.hrms_state["ms_raw"]).iterdir()) == ["analysis.baf", "analysis.sqlite"]
        # Spektrum: gemittelt über den LC-Bereich, Isotopenmuster aus der Formel, Peak per Klick
        pat_ = HRMS.isotope_pattern("C18H21N2", 1)
        assert abs(pat_[0][0] - calc_) < 2e-5 and pat_[0][1] == 100 and 19.5 < pat_[1][1] < 22.5, pat_[:3]
        assert t_ms.F is not None and len(t_ms._sticks[0]) >= 2, t_ms._sticks
        lo_, hi_ = t_ms.view_range()
        assert lo_ < calc_ < hi_ and hi_ - lo_ < 6, (lo_, hi_)
        assert "M+1" in t_ms.lbl_ms_info.text(), t_ms.lbl_ms_info.text()
        mzs_, rel_ = t_ms._sticks
        j_ = int(np.argmin(np.abs(mzs_ - calc_)))
        px_, py_ = t_ms.F.ax.transData.transform((mzs_[j_], rel_[j_] / 2))
        t_ms.spectrum_clicked(MouseEvent("button_press_event", t_ms.canvas, px_, py_, button=1))
        pump(0.05)
        assert w.hrms_state["found"] == "pick" and abs(w.hrms_state["found_pick"] - mzs_[j_]) < 1e-5, w.hrms_state
        assert t_ms._ms_line[0].endswith(f"found {mzs_[j_]:.4f}."), t_ms._ms_line[0]
        t_ms.rt_selected(0.0, 2.0)                         # Ziehen im Chromatogramm: alle Scans mit Signal
        assert w.hrms_state["ms_rt"] == [0.0, 2.0] and t_ms.raw_found(calc_)["n"] == 12
        assert not t_ms.cb_rt_auto.isChecked()
        files = t_ms.export_to(out / "hrms" / "raw")
        assert any(f.name.endswith("_sticks.csv") for f in files) and all(f.exists() for f in files)
        w.hrms_state["found"] = None
        t_ms.load_ms(str(irdir / "ms.docx"))               # Bericht + Rohdaten: Bericht vorgewählt, Rohdaten wählbar
        t_ms.render()
        assert t_ms._ms_line[0].endswith("found 265.1703.") and t_ms.cmb_found.findData("raw") >= 0
        assert any(k_ == "ms_raw" and c_ is w.hrms_state for c_, k_, _ in PRJ._path_slots(w.get_state()))
        w.hrms_state.update(ms_file=None, ms_raw=None, ms_rt=None, found=None, found_pick=None)
        w.compound = {}
        report.append("Bruker .d raw data (LC-peak mean): OK")

        # Audit P3 (HRMS + IR + EA): jeder Fall scheiterte vor den Fixes; Fälle gesammelt, Fehler alle auf einmal
        fails_ = []
        def chk_(name, fn):
            try:
                fn()
            except Exception as e_:
                fails_.append(f"{name}: {type(e_).__name__} {e_}"[:160])
        def p3_formula():
            assert IRMS.parse_formula("Ca(OH)2") == {"Ca": 1, "O": 2, "H": 2}
            assert IRMS.parse_formula("C18H20N2·2H2O") == {"C": 18, "H": 24, "N": 2, "O": 2}
            for bad_ in ("C18H20n2", "C18H20N2(", "C18H20N2 (264.16)"):
                try:
                    IRMS.parse_formula(bad_)
                    raise AssertionError(f"{bad_} angenommen")
                except ValueError:
                    pass
        chk_("formula parentheses/hydrate/strict", p3_formula)
        def p3_heavy():                                  # Ir: mittlere Masse vorhanden, HRMS über RDKit (193Ir)
            assert abs(sum(SI.ea_composition("C20H20ClIrN2").values()) - 100) < 1e-9
            m_ = IRMS.ion("C20H20ClIrN2", "M+H")[2]
            assert abs(m_ - (240 + 21 * 1.00782503207 + 192.9629264 + 2 * 14.0030740048 + 34.96885268
                             - 0.000548579909)) < 1e-4, m_
        chk_("EA/HRMS heavy elements", p3_heavy)
        def p3_isotope():                                # RDKit schreibt Isotope als [13C]
            f_, z_, m_, _ = IRMS.ion("[13C]H4", "M+H")
            assert abs(m_ - 18.0419) < 1e-4 and f_ == "[13C]H5", (f_, m_)
            assert "<sup>13</sup>C" in IRMS.ms_line("[13C]H4", "M+H", None)[1]
            assert abs(HRMS.isotope_pattern(f_, 1)[0][0] - m_) < 2e-5
        chk_("isotope labels", p3_isotope)
        def p3_ea_nan():
            try:
                SI.ea_line("C18H20N2", "", {"C": "NaN"})
                raise AssertionError("NaN angenommen")
            except ValueError:
                pass
        chk_("EA found NaN", p3_ea_nan)
        def opus_file_(path, y, kind=(15, 20), par_kind=(31, 20)):
            par_ = opus_par([("NPT", 0, len(y)), ("FXV", 1, 4000.0), ("LXV", 1, 400.0), ("CSF", 1, 1.0)])
            hd_ = b"\x0a\x0a\xfe\xfe" + _st.pack("<d", 920622.0) + _st.pack("<iii", 24, 2, 2)
            dr_ = (_st.pack("<BBBBii", par_kind[0], par_kind[1], 0, 0, len(par_) // 4, 48)
                   + _st.pack("<BBBBii", kind[0], kind[1], 0, 0, len(y), 48 + len(par_)))
            path.write_bytes(hd_ + dr_ + par_ + np.asarray(y, "<f4").tobytes())
            return path
        def p3_opus_outlier():
            y_ = np.full(128, 0.8)
            y_[5] = 1.6
            r_ = IRMS.read_opus(opus_file_(irdir / "outlier.0", y_))
            assert abs(np.median(r_["y"]) - 80) < 1e-3, np.median(r_["y"])
        chk_("OPUS fraction with outlier", p3_opus_outlier)
        def p3_opus_single():                            # nur Einkanalspektrum (2. Kennbyte 4): ablehnen
            try:
                IRMS.read_opus(opus_file_(irdir / "single.0", np.full(128, 0.05), (7, 4), (23, 4)))
                raise AssertionError("Einkanal angenommen")
            except ValueError:
                pass
        chk_("OPUS single channel rejected", p3_opus_single)
        def p3_plateau():                                # gesättigte Bande: ein Peak in der Mitte
            x_ = np.arange(4000.0, 399.0, -2.0)
            t_ = 100 - 40 * np.exp(-((x_ - 2900) / 10) ** 2) - 300 * np.exp(-((x_ - 1700) / 25) ** 2)
            pk_ = IRMS.pick_peaks(x_, np.clip(t_, 0, None), "T")
            w_ = [p_["wn"] for p_ in pk_]
            assert len(w_) == 2 and abs(w_[1] - 1700) < 1.5, w_
        chk_("IR saturated band", p3_plateau)
        def p3_round():                                  # round() rundet x.5 zur geraden Zahl
            l_ = IRMS.ir_line([{"wn": 2960.5}, {"wn": 1000.5}])[0]
            assert l_.endswith("2961, 1001"), l_
        chk_("IR x.5 rounded up", p3_round)
        def p3_ir_pdf():
            for name_, vals_ in (("p3a.pdf", ["1678.41"]), ("p3b.pdf", ["4500.10", "1700.20", "380.30"])):
                d_ = pymupdf.open()
                pg_ = d_.new_page()
                for i_, v_ in enumerate(vals_):
                    pg_.insert_text((50, 60 + 20 * i_), v_)
                d_.save(str(irdir / name_))
            t_ir.load_ir(str(irdir / "test.0"))
            t_ir.load_ir(str(irdir / "p3a.pdf"))
            (irdir / "other.0").write_bytes((irdir / "test.0").read_bytes())
            t_ir.load_ir(str(irdir / "other.0"))         # neues Spektrum: altes PDF gilt nicht mehr
            assert not w.irms_state.get("ir_pdf"), w.irms_state.get("ir_pdf")
            t_ir.clear_ir()
            t_ir.load_ir(str(irdir / "p3b.pdf"))         # nur PDF: alle Peakwerte, nicht 400-4000
            assert [p_["wn"] for p_ in t_ir.peaks()] == [4500.1, 1700.2, 380.3], t_ir.peaks()
        chk_("IR PDF replace/range", p3_ir_pdf)
        t_ir.clear_ir()
        def ms_docx_(path, paras, mzs, emf_extra=()):
            with _zf.ZipFile(path, "w") as z_:
                z_.writestr("word/document.xml", "<w:document><w:body>" + "".join(paras) + "</w:body></w:document>")
                z_.writestr("word/media/image1.emf", emf(list(emf_extra) + [*mzs, "m/z"]))
            return path
        par_ = lambda t_: f"<w:p><w:r><w:t>{t_}</w:t></w:r></w:p>"
        def p3_report():
            rep_ = IRMS.read_ms_docx(ms_docx_(irdir / "p3_del.docx", [
                par_("molecular formula:"), '<w:p><w:del w:id="1"><w:r><w:delText>C18H22N2</w:delText></w:r></w:del>'
                '<w:ins w:id="2"><w:r><w:t>C18H20N2</w:t></w:r></w:ins></w:p>',
                par_("method of Ionisation:"), par_("ESI ☐ APCI ☑ HESI ☐")], ["265.1703"], ["HR-MS"]))
            assert rep_["formula"] == "C18H20N2" and rep_["ionisation"] == "APCI" and rep_["polarity"] == "+", rep_
        chk_("MS report tracked changes/checkbox/HR-MS", p3_report)
        def set_adduct_(a_):                             # Combobox kann nach Zustandstests abweichen
            t_ms.cmb_adduct.setCurrentIndex(t_ms.cmb_adduct.findData(a_))
            w.hrms_state["adduct"] = a_
        def p3_adduct():                                 # Wechsel M+H -> M+Na: found des neuen Ions
            w.compound = {"formula": "C18H20N2"}
            set_adduct_("M+H")
            t_ms.load_ms(str(ms_docx_(irdir / "p3_na.docx", [par_("molecular formula:"), par_("C18H20N2")],
                                      ["265.1703", "287.1519"], ["+MS"])))
            t_ms.render()
            assert t_ms._ms_line[0].endswith("found 265.1703."), t_ms._ms_line[0]
            t_ms.cmb_adduct.setCurrentIndex(t_ms.cmb_adduct.findData("M+Na"))
            t_ms.render()
            assert t_ms._ms_line[0].endswith("found 287.1519."), t_ms._ms_line[0]
        chk_("HRMS adduct change", p3_adduct)
        def p3_nomatch():                                # kein passender Wert: found bleibt leer
            set_adduct_("M+H")
            t_ms.load_ms(str(ms_docx_(irdir / "p3_none.docx", [par_("molecular formula:"), par_("C18H20N2")],
                                      ["100.0000", "200.0000"], ["+MS"])))
            t_ms.render()
            t_ms.cb_calc.toggle()
            t_ms.render()
            t_ms.cb_calc.toggle()
            t_ms.render()
            assert t_ms._ms_line[0].endswith("found –."), t_ms._ms_line[0]
        chk_("HRMS no match stays empty", p3_nomatch)
        w.hrms_state.update(ms_file=None, found=None)
        def make_d_(dir_, scans):                        # scans: (rt, [m/z], [Intensität], Polarität 0/1)
            dir_.mkdir(parents=True, exist_ok=True)
            baf_, rows_ = bytearray(b"\0" * 64), []
            for i_, (rt_, mz_, in_, key_) in enumerate(scans):
                n_ = len(mz_)
                body = (_st.pack("<4I", n_, 16, 16 + 8 * n_, 16 + 12 * n_) + to_idx(mz_).astype("<f8").tobytes()
                        + np.asarray(in_, "<f4").tobytes() + np.full(n_, 0.01, "<f4").tobytes())
                head_ = bytearray(52)
                _st.pack_into("<III", head_, 0, 52 + len(body), IRMS.BAF_LINE_MAGIC, 52)
                rows_.append((i_ + 1, float(rt_), (0x10 << 56) | len(baf_), key_))
                baf_ += head_ + body
            (dir_ / "analysis.baf").write_bytes(bytes(baf_))
            tf_ = _st.pack("<i7d", 10, *cal_)
            (dir_ / "analysis.sqlite").unlink(missing_ok=True)
            con_ = _sq.connect(str(dir_ / "analysis.sqlite"))
            con_.executescript("create table Properties (Key text, Value text); create table AcquisitionKeys "
                               "(Id integer, Polarity integer); create table Spectra (Id integer, Rt real, "
                               "TransformatorId integer, LineIndexId integer, AcquisitionKey integer); create table "
                               "Steps (TargetSpectrum integer); create table Transformators (Id integer, Blob blob);")
            con_.executemany("insert into Properties values (?, ?)", [("SchemaType", "Baf2Sql"), ("SampleName", "P3")])
            con_.executemany("insert into AcquisitionKeys values (?, ?)", [(0, 0), (1, 1)])
            con_.executemany("insert into Spectra values (?, ?, 1, ?, ?)", rows_)
            con_.execute("insert into Transformators values (1, ?)", (tf_ + b"\0" * (112 - len(tf_)),))
            con_.commit()
            con_.close()
            return dir_
        amps_ = [1e4, 8e4, 1e5, 8e4, 1e4]
        scans_ = [(41.5 + 10 * i_, [calc_], [a_], 0) for i_, a_ in enumerate(amps_)]
        scans_ += [(46.5 + 10 * i_, [calc_ + 0.004], [10 * a_], 1) for i_, a_ in enumerate(amps_)]   # negativ
        def p3_uri():
            r_ = IRMS.read_bruker_d(make_d_(irdir / "P3#1.d", scans_[:2]))
            assert len(r_["spectra"]) == 2 and not (irdir / "P3").exists()
        chk_("Bruker path with #", p3_uri)
        def p3_polarity():                               # Polaritätswechsel: nur Scans des Addukts
            d_ = make_d_(irdir / "p3mixed.d", sorted(scans_))
            assert IRMS.read_bruker_d(d_)["polarity"] == "±"
            w.compound = {"formula": "C18H20N2"}
            set_adduct_("M+H")
            t_ms.load_raw(str(d_))
            t_ms.render()
            assert abs(t_ms._found - calc_) < 1e-6 and t_ms._lc["n"] == 3, (t_ms._found, t_ms._lc)
            t_ms.cb_rt_auto.setChecked(False)            # Felder (0.01 min) dürfen Randscans nicht abschneiden
            t_ms.render()
            assert t_ms._lc["n"] == 3 and abs(w.hrms_state["ms_rt"][0] * 60 - 51.5) < 1e-6, w.hrms_state["ms_rt"]
            w.hrms_state.update(found="pick", found_pick=calc_)
            t_ms.clear_raw()
            assert w.hrms_state["found"] is None and w.hrms_state["found_pick"] is None, w.hrms_state
        chk_("Bruker polarity/RT fields/pick", p3_polarity)
        def p3_view():                                   # Isotope unter calc (Sn: 116 bei calc - 4) im Bild
            w.compound = {"formula": "C4H10Sn"}
            w.hrms_state["x_auto"] = True
            t_ms.update_ms()
            lo_, _ = t_ms.view_range()
            assert lo_ < t_ms._ion[2] - 4, (lo_, t_ms._ion)
        chk_("HRMS view range lower isotopes", p3_view)
        w.hrms_state.update(ms_file=None, ms_raw=None, ms_rt=None, found=None, found_pick=None, adduct="M+H")
        w.compound = {}
        assert not fails_, "\n".join(fails_)
        report.append("audit P3 HRMS/IR/EA edge cases: OK")

        # Audit P3, Codex-Gegenlesen + Entscheidungen R1-R3 (jeder Fall scheiterte an b076a23)
        fails_ = []
        def r3_chemdraw():                               # ChemDraw: 13C beim C mit Isotopenmasse (Nutzerbeispiel)
            pl_ = SI.ea_line("C20[13C]H24N4", "", {})[0]
            assert "C 75.94, H 7.26, N 16.80" in pl_, pl_
            assert abs(IRMS.ion("C20[13C]H24N4", "M")[2] + 0.000548579909 - 333.2035) < 5e-5
        chk_("EA isotopes like ChemDraw", r3_chemdraw)
        def r2_printed():                                # Warnung nach den gedruckten Werten (calc 81.78)
            assert SI.ea_line("C18H20N2", "", {"C": "82.18"})[2] == []
            assert SI.ea_line("C18H20N2", "", {"C": "82.19"})[2]
        chk_("EA warning on printed values", r2_printed)
        def r1_tags():                                   # w/m/s: Tiefe in T, ganzes Spektrum, auch für A-Dateien
            x_ = np.arange(4000.0, 399.0, -2.0)
            t_ = 100 - 95 * np.exp(-((x_ - 1700) / 10) ** 2) - 40 * np.exp(-((x_ - 2900) / 10) ** 2)
            def tag_(wn_, y_, kind_, lo_=None, hi_=None):
                return next(IRMS.intensity_tag(p_) for p_ in IRMS.pick_peaks(wn_, y_, kind_, lo_, hi_)
                            if abs(p_["wn"] - 2900) < 2)
            tags_ = (tag_(x_, t_, "T"), tag_(x_, IRMS.to_absorbance(t_), "A"), tag_(x_, t_, "T", 2000, 3500))
            assert tags_ == ("m", "m", "m"), tags_
        chk_("IR intensity class", r1_tags)
        def c1_del():                                    # gelöschte Absatzmarke <w:del/>: Text danach bleibt
            rep_ = IRMS.read_ms_docx(ms_docx_(irdir / "p3_delmark.docx", [
                par_("molecular formula:"), '<w:p><w:pPr><w:rPr><w:del w:id="3" w:author="x"/></w:rPr></w:pPr>'
                '<w:r><w:t>C18H20N2</w:t></w:r></w:p>', '<w:p><w:del w:id="4"><w:r><w:delText>alt</w:delText>'
                '</w:r></w:del></w:p>'], ["265.1703"], ["+MS"]))
            assert rep_["formula"] == "C18H20N2", rep_["formula"]
        chk_("MS report deleted paragraph mark", c1_del)
        def c2_batch():                                  # A.0, A.pdf, B.0: B ohne das PDF von A
            (irdir / "A.0").write_bytes((irdir / "test.0").read_bytes())
            (irdir / "B.0").write_bytes((irdir / "test.0").read_bytes())
            (irdir / "A.pdf").write_bytes((irdir / "p3a.pdf").read_bytes())
            t_ir.load_ir_files([str(irdir / "A.0"), str(irdir / "A.pdf"), str(irdir / "B.0")])
            assert Path(w.irms_state["ir_file"]).stem.startswith("B") and not w.irms_state.get("ir_pdf"), w.irms_state
            t_ir.load_ir_files([str(irdir / "A.0"), str(irdir / "A.pdf")])    # ein Spektrum: PDF gehört dazu
            assert Path(w.irms_state["ir_pdf"]).stem.startswith("A")
            t_ir.clear_ir()
        chk_("IR batch load PDF pairing", c2_batch)
        def c3_unit():                                   # Prozentdaten mit vielen Werten < 1.5: nicht raten
            y_ = np.where(np.arange(128) % 2, 80.0, 0.8)
            try:
                IRMS.read_opus(opus_file_(irdir / "mixed_unit.0", y_))
                raise AssertionError("Einheit geraten")
            except ValueError:
                pass
        chk_("OPUS ambiguous unit rejected", c3_unit)
        def c5_tc():                                     # ohne natürliche Isotope nur mit Massenzahl
            try:
                IRMS.parse_formula("C10H10Tc")
                raise AssertionError("Tc angenommen")
            except ValueError:
                pass
            assert HRMS.isotope_pattern("C10H11[99Tc]", 1)[0][1] == 100
        chk_("element without natural isotopes", c5_tc)
        def c4_edges():                                  # Grenzen s -> min -> s: Randscan bleibt
            sp_ = [(t_, np.array([calc_]), np.array([1e4])) for t_ in (13.0, 14.0, 15.002)]
            assert IRMS.lc_found(sp_, calc_, rt=[60 * (13.0 / 60), 60 * (15.002 / 60)])["n"] == 3   # 15.0019999…
            d_ = make_d_(irdir / "p3mixed2.d", sorted(scans_))
            w.compound = {"formula": "C18H20N2"}
            set_adduct_("M+H")
            t_ms.load_raw(str(d_))
            t_ms.render()
            mz_ = t_ms._sticks[0]                        # Spektrum ohne die negativen Scans
            assert len(mz_) and not (np.abs(mz_ - (calc_ + 0.004)) < 1e-3).any(), mz_
            w.hrms_state["ms_rt"] = [0.865, 1.2]         # Feld zeigt 0.86; bewusste Änderung auf 0.87 gilt
            t_ms.sync_widgets()
            t_ms.sp_rt_lo.setValue(0.87)
            assert abs(w.hrms_state["ms_rt"][0] - 0.87) < 1e-9, w.hrms_state["ms_rt"]
            t_ms.clear_raw()
        chk_("HRMS RT edges/spectrum polarity", c4_edges)
        w.hrms_state.update(ms_file=None, ms_raw=None, ms_rt=None, found=None, found_pick=None, adduct="M+H")
        w.compound = {}
        assert not fails_, "\n".join(fails_)
        report.append("audit P3 review fixes (Codex) + EA/IR decisions: OK")

        # X-Ray: künstliche CIF (Ring auf Inversionszentrum, PART 1/2, Wasser als zweites Fragment, ohne Torsionen)
        from matplotlib.backend_bases import MouseEvent
        xdir = out / "input" / "xray"
        xdir.mkdir(parents=True, exist_ok=True)
        cif_, xyz_ = XRAY.make_test_cif(xdir / "xtest.cif")
        assert abs(XRAY.chi2_scale(0.5) - 1.5382) < 1e-4, XRAY.chi2_scale(0.5)
        S_ = XRAY.Structure(XRAY.structure_blocks(cif_)[0])
        assert len(S_.ops) == 4 and len(S_.atoms) == 9 and S_.parts() == [1, 2]
        c1_ = S_.cart(S_.atoms[S_.by_label["C1"]]["frac"])
        assert np.abs(c1_ - xyz_["C1"]).max() < 1e-3, c1_                         # fraktional -> kartesisch
        mol_ = XRAY.Molecule(S_, "main")
        names_ = sorted(XRAY.atom_name(a_).encode("ascii", "replace").decode() for a_ in mol_.atoms)
        assert len(mol_.atoms) == 13 and mol_.ops == {"\u2212x, \u2212y, \u2212z": "i"}, names_
        assert mol_.find_key("Cl1B") is None and mol_.find_key("Cl1|\u2212x, \u2212y, \u2212z") is not None
        assert [f_["formula"] for f_ in mol_.fragments] == ["C6H2Cl2", "H2O"], mol_.fragments
        assert XRAY.Molecule(S_, "all").find_key("Cl1B") is not None
        mi_ = XRAY.Molecule(S_, "minor")
        assert mi_.find_key("Cl1B") is not None and mi_.find_key("Cl1") is None
        inv_ = "|\u2212x, \u2212y, \u2212z"
        m_ = XRAY.measure(mol_, ["C1", "C2"])
        d_ = float(np.linalg.norm(mol_.atoms[mol_.find_key("C2")]["xyz"] - mol_.atoms[mol_.find_key("C1")]["xyz"]))
        assert m_["text"] == "1.390(3)" and abs(d_ - 1.390) < 2e-3, (m_, d_)      # Wert aus _geom_bond
        assert XRAY.measure(mol_, ["C3", "C1" + inv_])["text"] == "1.390(3)"       # Symmetriecode 3_555
        assert XRAY.measure(mol_, ["C3" + inv_, "C1", "C2"])["text"] == "120.0(2)"  # umgekehrte Reihenfolge
        tor_ = XRAY.measure(mol_, ["C1", "C2", "C3", "C1" + inv_])
        assert tor_["source"] == "calc" and not tor_["esd"] and tor_["text"] == "0.0", tor_
        sg_ = XRAY.space_group_markup("P 21/c")
        assert sg_ == "*P*2_{1}/*c*" and "<sub>1</sub>" in XRAY.to_html(sg_), sg_
        assert XRAY.to_plain(XRAY.space_group_markup("P -1")) == "P1\u0305"
        cv_ = XRAY.crystal_values(S_.block)
        assert cv_["diff"] == "0.42/\u22120.31" and cv_["theta"] == "5.000 to 55.000", cv_["diff"]
        assert cv_["index"].startswith("\u22129 \u2264 *h*") and cv_["formula"] == "C_{6}H_{4}Cl_{2}O"
        t_x = w.xray_tab
        w.main_tabs.setCurrentWidget(t_x)
        t_x.load_cif(str(cif_))
        assert len(w.xray_state["entries"]) == 2 and Path(w.xray_state["entries"][0]["file"]).parent == proj / "data" / "xray"
        assert any(k_ == "file" and c_ is w.xray_state["entries"][0] for c_, k_, _ in PRJ._path_slots(w.get_state()))
        w.xray_state["current"] = w.xray_state["entries"][0]["uid"]
        t_x.sync_widgets()
        t_x.render()
        pump()
        F_x = t_x.F
        assert F_x is not None and len(F_x.scene["atoms"]) == 13 and F_x.foot is None   # keine Symmetrie-Fußnote im Bild
        # simulierter Klick auf C1 und C2 (Atommitte; Pixel gerundet -> nur das getroffene Atom prüfen)
        proc_ = t_x.canvas.callbacks.process

        def click_(key):
            it_ = t_x.F.scene["by_k"][t_x.F.mol.find_key(key)]
            px_, py_ = t_x.F.ax.transData.transform((it_["x"], it_["y"]))
            proc_("button_press_event", MouseEvent("button_press_event", t_x.canvas, px_, py_, button=1))
            proc_("button_release_event", MouseEvent("button_release_event", t_x.canvas, px_, py_, button=1))
            pump(0.05)
        t_x.sel = []
        click_("C1")
        assert t_x.sel == ["C1"], t_x.sel
        click_("C2")
        assert t_x.sel == ["C1", "C2"] and t_x._meas["text"] == "1.390(3)", (t_x.sel, t_x._meas)
        t_x.take_param()
        click_("C1")
        click_("C2")
        click_("C3")
        click_("C1" + inv_)
        assert t_x._meas and t_x._meas["kind"] == "torsion" and t_x._meas["source"] == "calc", t_x._meas
        t_x.take_param()
        ent_ = w.xray_state["entries"]
        ent_[1]["params"] = [{"atoms": ["C2", "C1"]}]                      # gleiche Bindung, andere Richtung
        hdr_, rows_, notes_ = t_x.params_table()
        flat_ = [r_[0] for r_ in rows_]
        assert flat_ == ["Bond lengths [\u00c5]", "C1\u2013C2", "Torsion angles [\u00b0]",
                         "C1\u2013C2\u2013C3\u2013C1"], flat_
        assert rows_[1][1:] == ["1.390(3)", "1.390(3)"] and rows_[3][1].startswith("0.0^{a}"), rows_
        assert any("no s.u." in n_ for n_ in notes_) and any(n_.startswith("Symmetry") for n_ in notes_), notes_
        hc_, rc_ = t_x.crystal_table()
        assert len(hc_) == 3 and [r_ for r_ in rc_ if r_[0] == "Space group"][0][1] == "*P*2_{1}/*c*"
        # Drehen per Ziehen: Ansicht im Zustand, Projektion ändert sich
        x0_ = [(a_["x"], a_["y"]) for a_ in t_x.F.scene["atoms"]]
        cx_, cy_ = t_x.F.ax.transData.transform((0.0, 0.0))
        cy_ += 0.3 * t_x.F.ax.bbox.height                               # neben das Molekül (kein Label, kein Atom)
        proc_("button_press_event", MouseEvent("button_press_event", t_x.canvas, cx_, cy_, button=1))
        proc_("motion_notify_event", MouseEvent("motion_notify_event", t_x.canvas, cx_ + 40, cy_ + 10, button=1))
        proc_("button_release_event", MouseEvent("button_release_event", t_x.canvas, cx_ + 40, cy_ + 10, button=1))
        pump(0.1)
        assert ent_[0]["view"] and len(ent_[0]["view"]) == 9, ent_[0].get("view")
        x1_ = [(a_["x"], a_["y"]) for a_ in t_x.F.scene["atoms"]]
        assert max(abs(a_[0] - b_[0]) + abs(a_[1] - b_[1]) for a_, b_ in zip(x0_, x1_)) > 0.05
        # Mercury-Stil (Standard): schattierte Schichten je Atom; je Atom ausblenden / Label per Doppelklick
        assert w.xray_state["style"] == "mercury"
        pc_ = t_x.F.atom_art[t_x.F.mol.find_key("C1")][0]
        assert len(pc_.get_paths()) == 13, len(pc_.get_paths())                 # 12 Schichten + Umriss
        assert "lbl_C1" in t_x.F.ann_items and t_x.F.ann_items["lbl_C1"].get_fontweight() == "bold"   # wie Mercury-Bilder
        w.xray_state["label_pt"] = 12.5                                     # Labelgröße je Tab
        t_x.render(push=False)
        assert t_x.F.ann_items["lbl_C1"].get_fontsize() == 12.5
        w.xray_state["label_pt"] = 0.0
        t_x.render(push=False)
        it_ = t_x.F.scene["by_k"][t_x.F.mol.find_key("C2")]
        px_, py_ = t_x.F.ax.transData.transform((it_["x"], it_["y"]))
        t_x.sel = []
        proc_("button_press_event", MouseEvent("button_press_event", t_x.canvas, px_, py_, button=1, dblclick=True))
        pump(0.6)
        assert ent_[0]["label_on"] == {"C2": False} and "lbl_C2" not in t_x.F.ann_items, ent_[0]["label_on"]
        assert not t_x.sel
        # Shift + Rahmen: Atome im Rahmen -> Menü (hier abgefangen), dann Labels mehrerer Atome aus und wieder an
        got_, boxed_ = [], t_x.atoms_boxed
        t_x.atoms_boxed = got_.extend
        try:
            p1_, p2_ = (t_x.F.ax.transData.transform((t_x.F.scene["by_k"][t_x.F.mol.find_key(l_)]["x"],
                                                      t_x.F.scene["by_k"][t_x.F.mol.find_key(l_)]["y"]))
                        for l_ in ("C1", "C2"))
            x0_, y0_ = min(p1_[0], p2_[0]) - 3, min(p1_[1], p2_[1]) - 3
            x1_, y1_ = max(p1_[0], p2_[0]) + 3, max(p1_[1], p2_[1]) + 3
            proc_("button_press_event", MouseEvent("button_press_event", t_x.canvas, x0_, y0_, button=1, key="shift"))
            proc_("motion_notify_event", MouseEvent("motion_notify_event", t_x.canvas, x1_, y1_, button=1,
                                                    key="shift"))
            proc_("button_release_event", MouseEvent("button_release_event", t_x.canvas, x1_, y1_, button=1,
                                                     key="shift"))
            pump(0.1)
        finally:
            t_x.atoms_boxed = boxed_
        keys_ = {t_x.F.mol.atoms[k_]["key"] for k_ in got_}
        assert {"C1", "C2"} <= keys_, keys_
        assert not any(p_.get_linestyle() == "--" for p_ in t_x.F.ax.patches), "Rahmen bleibt stehen"
        t_x.set_labels(got_, False)
        assert not any(f"lbl_{k_}" in t_x.F.ann_items for k_ in keys_), keys_
        t_x.set_labels(got_, True)
        assert all(f"lbl_{k_}" in t_x.F.ann_items for k_ in keys_ if not k_.startswith("H")), keys_
        ent_[0]["label_on"] = {}
        t_x.set_label("H3", True)                                           # H-Label einzeln an
        ent_[0]["hidden_atoms"] = ["H3" + inv_]                             # ein einzelnes H ausblenden
        t_x.render()
        assert "lbl_H3" in t_x.F.ann_items and t_x.F.mol.find_key("H3" + inv_) not in t_x.F.scene["by_k"]
        w.xray_state["style"] = "flat"
        t_x.render()
        assert len(t_x.F.atom_art[t_x.F.mol.find_key("C1")][0].get_paths()) == 2
        assert t_x.F.ann_items["lbl_C1"].get_fontweight() == "bold"
        w.xray_state["style"] = "mercury"
        ent_[0].update(label_on={}, hidden_atoms=[])
        ent_[0]["hidden_frags"] = [1]                                   # Wasser ausblenden
        w.xray_state["show_h"] = False
        t_x.sync_widgets()
        t_x.render()
        assert len(t_x.F.scene["atoms"]) == 8, len(t_x.F.scene["atoms"])
        files = t_x.export_to(out / "xray" / "xray_test")
        for f in files:
            assert f.exists() and f.stat().st_size > 200, f
        import pymupdf
        npath_ = len(pymupdf.open(str(out / "xray" / "xray_test_xtest_a.pdf"))[0].get_drawings())
        assert npath_ > 20, npath_
        csv_ = (out / "xray" / "xray_test_crystal_data.csv").read_text(encoding="utf-8-sig")
        assert "P2\u2081/c" in csv_ and "0.42/\u22120.31" in csv_
        from openpyxl import load_workbook
        wb_ = load_workbook(out / "xray" / "xray_test_tables.xlsx")
        assert wb_.sheetnames == ["Crystal data", "Selected parameters"]
        assert (out / "xray" / "xray_test_captions.txt").read_text(encoding="utf-8").count("X-ray solid state structure") == 2
        # Lösungsmittel-Schalter, Startansicht, Label-Modus, Drahtmodell, Bildunterschrift
        mol_ = t_x.molecule(ent_[0])
        assert [f_["kind"] for f_ in mol_.fragments] == ["main", "solvent"] and mol_.fragments[1]["name"] == "water"
        assert XRAY.cif_order(mol_, ["C2", "C1"]) == ["C1", "C2"]                # Reihenfolge wie in der CIF
        R_ = XRAY.default_view(mol_)
        assert np.allclose(R_ @ R_.T, np.eye(3), atol=1e-9) and abs(np.linalg.det(R_) - 1) < 1e-9
        # Ausrichten: Bindung nach unten, Ebene von vorn, Drehen, ChemDraw als Vorlage
        xyz_k = lambda k_: mol_.atoms[mol_.find_key(k_)]["xyz"]  # noqa: E731
        Rb_ = XRAY.align_bond(R_, mol_, ["C1", "Cl1"], "down")
        v_ = Rb_ @ (xyz_k("Cl1") - xyz_k("C1"))
        assert abs(v_[0]) < 1e-6 and abs(v_[2]) < 1e-6 and v_[1] < 0, v_
        Rp_ = XRAY.align_plane(R_, mol_, ["C1", "C2", "C3"])
        n_ = np.cross(xyz_k("C2") - xyz_k("C1"), xyz_k("C3") - xyz_k("C1"))
        assert abs(abs((Rp_ @ (n_ / np.linalg.norm(n_)))[2]) - 1) < 1e-6
        Rpb_ = XRAY.align_bond(Rp_, mol_, ["C1", "C2"], "right")           # Bindung in der Ebene: Ebene bleibt
        assert abs(abs((Rpb_ @ (n_ / np.linalg.norm(n_)))[2]) - 1) < 1e-6
        from rdkit import Chem as _Chem
        if hasattr(_Chem, "MolToCDXMLBlock"):
            from rdkit.Chem import AllChem as _AC
            tm_ = _Chem.MolFromSmiles("Clc1ccc(Cl)cc1")
            _AC.Compute2DCoords(tm_)
            tpl_ = xdir / "dcb.cdxml"
            tpl_.write_text(_Chem.MolToCDXMLBlock(tm_), encoding="utf-8")
            Rt_, inf_ = XRAY.align_to_template(mol_, tpl_)
            assert inf_["n"] == 8 and inf_["rmsd"] < 0.3, inf_
            cls_ = [a_.GetIdx() for a_ in tm_.GetAtoms() if a_.GetSymbol() == "Cl"]
            cf_ = tm_.GetConformer()
            dt_ = np.array(cf_.GetAtomPosition(cls_[1])) - np.array(cf_.GetAtomPosition(cls_[0]))
            dc_ = (Rt_ @ (xyz_k("Cl1|" + [o_ for o_ in mol_.ops][0]) - xyz_k("Cl1")))[:2]
            cosang_ = abs(np.dot(dt_[:2], dc_)) / np.linalg.norm(dt_[:2]) / np.linalg.norm(dc_)
            assert cosang_ > 0.98, cosang_                                   # Cl–Cl wie in der Zeichnung
        ent_[0].update(hidden_frags=[], view=None)
        w.xray_state.update(show_h=True, hide_solv=True)
        t_x.sync_widgets()
        t_x.render()
        assert t_x.F.scene["atoms"] and all(mol_.atoms[it_["k"]]["frag"] == 0 for it_ in t_x.F.scene["atoms"])
        assert "<sub>1</sub>" in t_x.lbl_info.text(), t_x.lbl_info.text()          # Raumgruppe formatiert
        cap_ = t_x.caption()[0]
        assert cap_.startswith("X-ray solid state structure of xtest_a (70% (PART 1) and 30% (B) occupancy; major part "
                               "shown). Thermal ellipsoids are shown with 50% probability."), cap_
        assert "Co-crystallised water is omitted for clarity." in cap_, cap_
        # Audit P4 R1/R2 (wissenschaftliche Praxis, Transparenz): Symmetriecode der genannten erzeugten Atome und
        # Kennzeichnung gerechneter Werte in der Unterschrift (vorher: weder noch)
        assert ("Selected bond lengths and torsion angles in [Å] and [°]: C1–C2: 1.390(3); C1–C2–C3–C1ⁱ: 0.0\u1d43. "
                "Symmetry code: (i) \u2212x, \u2212y, \u2212z. \u1d43 Calculated from the atomic coordinates; no s.u. "
                "in the CIF.") in cap_, cap_
        assert "0.0<sup>a</sup>." in t_x.caption()[1] and "(i) \u2212x" in t_x.caption()[1], t_x.caption()[1]
        assert "<b>xtest_a</b>" in t_x.caption()[1] and not t_x.param_tables(ent_[0])   # wenige Werte: in der Unterschrift
        c1_, cl1_ = mol_.find_key("C1"), mol_.find_key("Cl1")
        assert mol_.substituent(c1_, mol_.find_key("C2")) is None                  # Ringbindung im Kern
        assert mol_.beyond(c1_, cl1_) == {cl1_}
        assert mol_.substituent(c1_, cl1_) == {cl1_} and mol_.substituent(cl1_, c1_) == {cl1_}
        t_x.sel = ["C1", "Cl1"]
        t_x.wire_selected()
        assert ent_[0]["wire"] == ["Cl1"] and not t_x.sel, ent_[0]["wire"]
        t_x.render()                                                              # Bindung Kern–Draht dünn
        mixed_ = [ln_ for a_, b_, ln_, h_ in t_x.F.bond_art if {a_, b_} == {c1_, cl1_}]
        assert len(mixed_) == 2 and all(ln_.get_linewidth() < t_x.F.lw_bond for ln_ in mixed_), mixed_
        assert "lbl_C1" in t_x.F.ann_items and t_x.F.scene["by_k"][c1_]["kind"] != "wire"
        ent_[0]["wire"] = ["C1", "Cl1"]                                           # beide Enden -> Drahtbindung
        ent_[0]["wire_name"] = "the C-Cl group"
        t_x.render()
        assert t_x.F.scene["by_k"][cl1_]["kind"] == "wire" and "lbl_Cl1" not in t_x.F.ann_items
        halves_ = [ln_ for a_, b_, ln_, h_ in t_x.F.bond_art if {a_, b_} == {c1_, cl1_}]
        assert len(halves_) == 2 and halves_[0].get_linewidth() < t_x.F.lw_bond, halves_
        assert "the C-Cl group is shown as wireframe" in t_x.caption()[0], t_x.caption()[0]
        t_x.sel = []
        t_x.unwire_selected()
        t_x.layout()["lbl_C1"] = [0.0, 0.0]                  # gespeichert, liegt aber auf dem Atom -> neu platziert
        t_x.render()
        assert tuple(float(v_) for v_ in t_x.F.ann_items["lbl_C1"].xyann) != (0.0, 0.0)
        t_x.layout().pop("lbl_C1")
        w.xray_state["label_mode"] = "hetero"
        t_x.sync_widgets()
        t_x.render()
        assert not ent_[0]["wire"] and "lbl_Cl1" in t_x.F.ann_items and "lbl_C1" not in t_x.F.ann_items
        # Darstellung auf Standard: Einstellungen + gewählte Struktur zurück, CIFs/Namen/Parameter bleiben
        e_x = t_x.entry()
        w.xray_state.update(bond_w=0.2, prob=30, style="flat", label_pt=14.0, show_h=False, disorder="all")
        e_x.update(hidden_atoms=[e_x["hidden_atoms"] or "C2"], label_on={"C1": False}, wire=["C1"], zoom=2.0,
                   name="umbenannt")
        n_ent, params_ = len(w.xray_state["entries"]), list(e_x.get("params") or [])
        t_x.reset_display()
        d_x = XRAY.XrayTab.DEFAULT_STATE
        assert all(w.xray_state[k_] == d_x[k_] for k_ in ("bond_w", "prob", "style", "show_h", "disorder"))
        assert w.xray_state["label_pt"] == 0.0 and abs(t_x.sp_bond.value() - 0.12) < 1e-9 and t_x.sp_prob.value() == 50
        assert not e_x["hidden_atoms"] and not e_x["label_on"] and not e_x["wire"] and e_x["zoom"] == 1.0
        assert len(w.xray_state["entries"]) == n_ent and e_x["name"] == "umbenannt" and e_x["params"] == params_
        assert t_x.F is not None
        # Stile: Mercury (C weiß, schwarze Bindungen, fette Labels), Olex2, Kugel-Stab glänzend, Kugel-Stab mit Tiefe,
        # Strichzeichnung. Im Tab X-Ray wählbar nur Mercury/Olex2 (alte Projekte -> Mercury); gezeichnet werden alle
        assert [t_x.cmb_style.itemData(i_) for i_ in range(t_x.cmb_style.count())] == list(XRAY.XRAY_STYLES)
        w.xray_state["style"] = "ballstick"
        assert t_x.opts(t_x.entry())["style"] == "mercury"
        styles0_, XRAY.XRAY_STYLES = XRAY.XRAY_STYLES, XRAY.STYLES
        for sty_ in XRAY.STYLES:
            w.xray_state["style"] = sty_
            t_x.render()
            F_s = t_x.F
            assert F_s.style == sty_
            c_it = next(i_ for i_ in F_s.scene["atoms"] if i_["el"] == "C")
            if sty_ in XRAY.BALL_STYLES:
                assert (sty_ == "cartoon" or abs(c_it["w"] - 2 * XRAY.ball_radius("C", sty_)) < 1e-9) \
                    and c_it["w"] == c_it["h"] and not c_it["curves"]
                assert F_s.bond_w == XRAY.ROD_W[sty_] and (sty_ == "cartoon" or sty_ == "iboview" and
                                                           len(F_s.rods["seg"]) > 10 or
                                                           len(F_s.rods["frac"]) > 10 and F_s.rods["frac"].max() == 1.0)
                # schnelles Drehen: Segmente im gemeinsamen Puffer (verjüngte IboView-Stäbe: Trapeze)
                assert sty_ in ("iboview", "cartoon") or F_s.rods.get("inplace"), "Stabsegmente nicht im gemeinsamen Puffer"
                pc_s = F_s.atom_art[c_it["k"]][0]
                nv_ = [len(q_.vertices) for q_ in pc_s.get_paths()]
                F_s.set_drag(True)                                     # Ziehen: gröbere Kugelpolygone, gleiche Schichten
                nd_ = [len(q_.vertices) for q_ in pc_s.get_paths()]
                F_s.set_drag(False)
                assert len(nd_) == len(nv_) and (sum(nd_) < sum(nv_) or sty_ == "iboview"), (sum(nd_), sum(nv_))
                assert [len(q_.vertices) for q_ in pc_s.get_paths()] == nv_
                fc_ = F_s.atom_art[c_it["k"]][0].get_facecolors()                         # Phong: dunkel … Glanz
                if sty_ == "cartoon":                           # weiß mit schwarzem Umriss, hohle Stäbe, Perspektive
                    assert all(tuple(f_[:3]) == (1, 1, 1) for f_ in fc_[:-1]), fc_
                    assert tuple(pc_s.get_edgecolors()[-1][:3]) == (0, 0, 0)
                    rods_ = [(ln_, h_) for _, _, ln_, h_ in F_s.bond_art]          # je Bindung Rand + Kern
                    outl_ = [ln_ for ln_, h_ in rods_ if h_ is None]
                    core_ = [ln_ for ln_, h_ in rods_ if h_ == "core"]
                    assert outl_ and len(outl_) == len(core_) and all(c_.get_color() == "white" for c_ in core_)
                    assert all(c_.get_zorder() > o_.get_zorder() for o_, c_ in zip(outl_, core_))
                    edge_ = (outl_[0].get_linewidth() - core_[0].get_linewidth()) / 2   # Rand = Kugelumriss
                    assert abs(edge_ - F_s._cartoon_edge()) < 1e-6 or core_[0].get_linewidth() == 0.2, edge_
                    toward_ = 0                                  # Stab zum Betrachter beginnt im Kugelumriss
                    for a_, b_, ln_, h_ in F_s.bond_art:
                        A_, B_ = F_s.scene["by_k"][a_], F_s.scene["by_k"][b_]
                        xs_, ys_ = ln_.get_data()
                        if h_ is None and np.isfinite(xs_).all() and B_["z"] - A_["z"] > 0.3:
                            assert float(np.hypot(xs_[0] - A_["x"], ys_[0] - A_["y"])) < A_["w"] / 2, (a_, b_)
                            toward_ += 1
                    assert toward_ > 0
                    ws_ = {round(i_["w"], 6) for i_ in F_s.scene["atoms"] if i_["el"] == "C"}
                    assert len(ws_) > 1, ws_                     # vordere C größer als hintere
                    continue
                assert len(fc_) > 10 and sum(fc_[0][:3]) < sum(fc_[len(fc_) // 2][:3])    # Rand dunkler als Mitte
                if sty_ == "ballstick":
                    assert min(max(f_[:3]) for f_ in fc_[-3:-1]) > 0.95, fc_[-3:]          # Glanzpunkt fast weiß
            if sty_ == "mercury":
                assert XRAY.default_color("C", "mercury") == XRAY.OLEX2_COLORS["C"]               # C grau wie bisher
                assert all(a_.get_fontweight() == "bold" for a_ in F_s.ann_items.values())
                assert all(ln_.get_color() == "black" for _, _, ln_, h_ in F_s.bond_art if h_ is None)
            if sty_ == "iboview":                               # wie IboView: Radien/Farben, Stab zur Mitte 72 %
                from elucida.core import orbitals as ORB
                assert abs(c_it["w"] / 2 - 0.4 * 1.43 * ORB.BOHR) < 1e-9 and XRAY.default_color("C", sty_) == "#999999"
                sp_ = F_s.rods["specs"]
                assert all(abs(x_["r1"] / x_["r0"] - ORB.IBO_TAPER) < 1e-9 for x_ in sp_)          # Einfachbindungen
                assert len(F_s.rods["seg"]) == len(sp_) * ORB.IBO_ROD_STRIPS
                assert F_s._fog(-20.0) == 1.0 and F_s._fog(0.0) == 0.0
                assert set(F_s.elem_lab) == {i_["k"] for i_ in F_s.scene["atoms"] if i_["el"] not in ("H", "D")}
                assert F_s.elem_lab[c_it["k"]].get_text() == "C"
            if sty_ == "depth":                                 # wie ORCA-Handbuch: ohne Aufhellung, Jmol-Farben
                assert F_s._fog(F_s._zmax - XRAY.DEPTH_A - 1) == 0.0
                assert XRAY.default_color("O", "depth") == "#FF0D0D" and XRAY.default_color("O") != "#FF0D0D"
            f_s = t_x.export_to(out / f"xray_style_{sty_}")
            assert all(f_.exists() and f_.stat().st_size > 500 for f_ in f_s if f_.suffix in (".pdf", ".svg", ".png"))
        XRAY.XRAY_STYLES = styles0_
        w.xray_state["style"] = "mercury"
        # Labelfarbe: schwarz (Standard) oder wie das Atom, helle Farben abgedunkelt (lesbar)
        w.xray_state["label_color"] = "atom"
        t_x.sync_widgets()
        t_x.render()
        assert t_x.cmb_lcol.currentData() == "atom"
        from matplotlib.colors import to_rgb
        n_lab_ = next(a_ for k_, a_ in t_x.F.ann_items.items()
                      if k_.startswith("lbl_") and t_x.F.mol.atoms[a_._atom]["el"] not in ("C", "H"))
        el_ = t_x.F.mol.atoms[n_lab_._atom]["el"]
        assert to_rgb(n_lab_.get_color()) == to_rgb(XRAY.label_color(el_, {}, "atom", "mercury")), (el_, n_lab_.get_color())
        assert to_rgb(n_lab_.get_color()) != (0, 0, 0)
        hc_ = XRAY.label_color("H", {}, "atom")
        assert 0.299 * hc_[0] + 0.587 * hc_[1] + 0.114 * hc_[2] <= 0.45 + 1e-9, hc_
        w.xray_state["label_color"] = "black"
        t_x.sync_widgets()
        t_x.render()
        assert all(to_rgb(a_.get_color()) == (0, 0, 0) for a_ in t_x.F.ann_items.values())
        # eigene Labelfarbe (Hex) für alle Labels, auch im Export
        w.xray_state["label_color"] = "#1a7f37"
        t_x.sync_widgets()
        t_x.render(push=False)
        assert t_x.cmb_lcol.currentData() == "color" and t_x.cmb_lcol.value() == "#1a7f37"
        assert t_x.F.ann_items and all(to_rgb(a_.get_color()) == to_rgb("#1a7f37") for a_ in t_x.F.ann_items.values())
        Fe_ = t_x.build(MplFigure(), interactive=False)
        assert all(to_rgb(a_.get_color()) == to_rgb("#1a7f37") for a_ in Fe_.ann_items.values())
        w.xray_state["label_color"] = "black"
        t_x.sync_widgets()
        t_x.render(push=False)
        report.append("X-Ray custom label colour: OK")
        # Drehen (GUI-Pfad über die Canvas-Ereignisse): das Molekül behält seine Größe – beim Ziehen, nach dem Loslassen
        # (render -> neue XrayFigure), bei mehreren Drehungen, mit verschobenem Label und nach Zoom per Rad;
        # „Ansicht zurücksetzen“ passt wieder an
        def rot_keeps_scale_(tab_):
            e_r = tab_.entry()
            keep_ = {k_: json.loads(json.dumps(e_r.get(k_))) for k_ in ("view", "layout", "zoom", "pan", "ext")}
            proc_r = tab_.canvas.callbacks.process

            def drag_(key_=None, dx_=24, dy_=8):
                F_ = tab_.F
                cx_, cy_ = F_.ax.bbox.x0 + 6, F_.ax.bbox.y0 + 6                  # Ecke: kein Atom, kein Label
                proc_r("button_press_event", MouseEvent("button_press_event", tab_.canvas, cx_, cy_, button=1))
                got_ = []
                for i_ in range(1, 4):
                    proc_r("motion_notify_event", MouseEvent("motion_notify_event", tab_.canvas, cx_ + dx_ * i_ / 3,
                                                             cy_ + dy_ * i_ / 3, button=1))
                    off_ = tuple(float(v_) for v_ in tab_.F.ann_items[key_].xyann) if key_ else None
                    got_.append((tab_.F.ax.get_xlim(), tab_.F.ax.get_ylim(), off_))
                proc_r("button_release_event", MouseEvent("button_release_event", tab_.canvas, cx_ + dx_, cy_ + dy_,
                                                          button=1))
                pump(0.1)
                return got_

            def lims_():
                return np.array([tab_.F.ax.get_xlim(), tab_.F.ax.get_ylim()])
            e_r.update(layout={}, zoom=1.0, pan=[0.0, 0.0])
            e_r.pop("ext", None)
            tab_.render(push=False)
            lim0_ = lims_()
            v0_ = list(e_r.get("view") or [])
            for dx_, dy_ in ((24, 8), (-18, 20), (30, -10)):                 # mehrere Drehungen ohne Labels
                got_ = drag_(dx_=dx_, dy_=dy_)
                assert all(np.allclose(g_[:2], lim0_) for g_ in got_), (got_, lim0_)
                assert np.allclose(lims_(), lim0_, atol=1e-4), (lims_(), lim0_)
            assert e_r.get("ext") and e_r.get("view") != v0_, e_r.get("ext")
            lk_ = next(iter(tab_.F.ann_items))                                # Label „verschoben“
            off0_ = tuple(float(v_) for v_ in tab_.F.ann_items[lk_].xyann)
            tab_.layout()[lk_] = list(off0_)
            tab_.render(push=False)
            got_ = drag_(lk_)
            assert all(np.allclose(g_[:2], lim0_) and g_[2] == off0_ for g_ in got_), (got_, lim0_)
            assert np.allclose(lims_(), lim0_, atol=1e-4), (lims_(), lim0_)
            F_ = tab_.F                                                     # Zoom per Rad, dann drehen
            proc_r("scroll_event", MouseEvent("scroll_event", tab_.canvas, F_.ax.bbox.x0 + 40, F_.ax.bbox.y0 + 40,
                                              button="up", step=1))
            pump(0.1)
            lim1_ = lims_()
            assert not np.allclose(lim1_, lim0_)
            drag_()
            assert np.allclose(lims_(), lim1_, atol=1e-4), (lims_(), lim1_)
            tab_.reset_view()
            assert "ext" not in e_r
            e_r.update({k_: v_ for k_, v_ in keep_.items() if v_ is not None})
            if keep_["ext"] is None:
                e_r.pop("ext", None)
            tab_.render(push=False)
        rot_keeps_scale_(t_x)
        report.append("X-Ray rotate keeps size (GUI path, labels moved, zoom): OK")
        # Audit P4 (X-Ray): jeder Fall scheiterte vor den Fixes; künstliche CIFs (kubische 10-Å-Zelle, Å-Koordinaten)
        fails_ = []
        def chk_(name, fn):
            try:
                fn()
            except Exception as e_:
                fails_.append(f"{name}: {type(e_).__name__} {e_}"[:160])
        import math
        from types import SimpleNamespace
        from elucida.formats.cif import parse_cif_text as _pcif
        def xcif_(atoms, ops=("x, y, z",), ids=None, bonds=(), angles=(), tors=(), extra=(), aniso=()):
            L_ = ["data_t"] + [f"_cell_length_{x_} 10" for x_ in "abc"] + \
                 [f"_cell_angle_{x_} 90" for x_ in ("alpha", "beta", "gamma")] + list(extra) + ["loop_"]
            L_ += (["_space_group_symop_id"] if ids else []) + ["_space_group_symop_operation_xyz"]
            L_ += [(f"{ids[k_]} " if ids else "") + f"'{o_}'" for k_, o_ in enumerate(ops)]
            L_ += ["loop_"] + [f"_atom_site_{t_}" for t_ in ("label", "type_symbol", "fract_x", "fract_y", "fract_z",
                                                             "U_iso_or_equiv", "occupancy", "disorder_group")]
            for at_ in atoms:
                lab_, el_, p_ = at_[:3]
                L_.append(f"{lab_} {el_} " + " ".join(f"{v_ / 10:.6f}" for v_ in p_) +
                          f" 0.02 {at_[4] if len(at_) > 4 else 1} {at_[3] if len(at_) > 3 else '.'}")
            if aniso:
                L_ += ["loop_", "_atom_site_aniso_label"] + [f"_atom_site_aniso_U_{k_}" for k_ in
                                                             ("11", "22", "33", "23", "13", "12")] + list(aniso)
            for tag_, n_, rows_ in (("bond", 2, bonds), ("angle", 3, angles), ("torsion", 4, tors)):
                if rows_:
                    L_ += ["loop_"] + [f"_geom_{tag_}_atom_site_label_{k_ + 1}" for k_ in range(n_)] + \
                          [f"_geom_{tag_}" + ("_distance" if tag_ == "bond" else "")] + \
                          [f"_geom_{tag_}_site_symmetry_{k_ + 1}" for k_ in range(n_)] + list(rows_)
            return "\n".join(L_) + "\n"
        def xmol_(txt, **kw):
            S_x = XRAY.Structure(_pcif(txt)[0])
            return S_x, XRAY.Molecule(S_x, **kw)
        inv2_ = ("x, y, z", "-x, -y, -z")
        def p4_torsion_sign():                       # Inversion kehrt das Vorzeichen der CIF-Torsion um
            _, m_ = xmol_(xcif_([("C1", "C", (0.75, 1.5, 0)), ("C2", "C", (0.75, 0, 0)), ("C3", "C", (2.25, 0, 0)),
                                 ("C4", "C", (2.25, 0, 1.5))], ops=inv2_,
                                bonds=["C1 C2 1.500(2) . .", "C2 C3 1.500(2) . .", "C3 C4 1.500(2) . .",
                                       "C2 C2 1.500(2) . 2"], tors=["C1 C2 C3 C4 90.0(2) . . . ."]))
            g_ = {a_["label"]: a_["key"] for a_ in m_.atoms if a_["sym"]}
            r_ = XRAY.measure(m_, [g_[x_] for x_ in ("C1", "C2", "C3", "C4")])
            assert r_["text"] == "-90.0(2)" and abs(r_["value"] + 90) < 1e-6, r_
            assert XRAY.measure(m_, ["C1", "C2", "C3", "C4"])["text"] == "90.0(2)"
        chk_("torsion sign under inversion", p4_torsion_sign)
        def p4_symop_id():                           # Codes beziehen sich auf _space_group_symop_id, nicht die Zeile
            _, m_ = xmol_(xcif_([("C1", "C", (1, 0, 0))], ops=("-x, -y, -z", "x, y, z"), ids=(2, 1),
                                bonds=["C1 C1 2.000(3) . 2_555"]))
            assert len(m_.atoms) == 2 and m_.bonds == [(0, 1)], (len(m_.atoms), m_.bonds)
        chk_("symop ids", p4_symop_id)
        def p4_truncated():                          # abgeschnittener Atom-Loop: Fehler statt still ein Atom weniger
            txt_ = xcif_([("C1", "C", (1, 1, 1)), ("C2", "C", (2, 1, 1))]).rstrip("\n")
            try:
                XRAY.Structure(_pcif(txt_.rsplit(" ", 4)[0] + "\n")[0])
                raise AssertionError("unvollständige Schleife angenommen")
            except ValueError:
                pass
        chk_("truncated atom loop", p4_truncated)
        def p4_special_position():                   # Zn im Inversionszentrum: CIF-esd auch für die invertierten N
            _, m_ = xmol_(xcif_([("Zn1", "Zn", (0, 0, 0)), ("N1", "N", (2, 0, 0)), ("N2", "N", (0, 2, 0))], ops=inv2_,
                                bonds=["Zn1 N1 2.000(2) . .", "Zn1 N2 2.000(2) . .", "Zn1 N1 2.000(2) . 2",
                                       "Zn1 N2 2.000(2) . 2"], angles=["N1 Zn1 N2 90.00(5) . . ."]))
            g_ = {a_["label"]: a_["key"] for a_ in m_.atoms if a_["sym"]}
            assert XRAY.measure(m_, [g_["N1"], "Zn1", g_["N2"]])["text"] == "90.00(5)"
        chk_("special position esd", p4_special_position)
        def p4_zprime():                             # zwei C6-Ringe ohne H: Z′ = 2 vor „Lösungsmittel Benzol“
            ring_ = lambda c_, n0_: [(f"C{n0_ + i_}", "C", (c_ + 1.39 * math.cos(math.radians(60 * i_)),
                                                            c_ + 1.39 * math.sin(math.radians(60 * i_)), c_))
                                     for i_ in range(6)]
            _, m_ = xmol_(xcif_(ring_(2, 1) + ring_(6, 7)))
            assert [f_["kind"] for f_ in m_.fragments] == ["main", "molecule"], m_.fragments
        chk_("Z' before solvent", p4_zprime)
        def p4_antiparallel():
            R_ = XRAY.rot_between(np.array([1e-9, 1.0, 0.0]), (0, -1, 0))
            assert np.isfinite(R_).all() and np.allclose(R_ @ [0, 1, 0], [0, -1, 0], atol=1e-6), R_
        chk_("rot_between near antiparallel", p4_antiparallel)
        def p4_threshold():                          # R-Werte unter der Schwelle der CIF beschriften
            b_ = _pcif(xcif_([("C1", "C", (1, 1, 1))], extra=["_reflns_threshold_expression 'I > 3sigma(I)'",
                                                               "_refine_ls_R_factor_gt 0.035",
                                                               "_refine_ls_wR_factor_gt 0.087"]))[0]
            v_ = XRAY.crystal_values(b_)
            assert v_["thr"] == "*I*>3*σ*(I)" and v_["method"] == "", v_
            from elucida.si.publication_si import merge_crystal
            t1_ = (["Compound", "A"], [["Final *R* indices [*I*>2*σ*(I)]", "R1"]])
            t2_ = (["Compound", "B"], [["Final *R* indices [*I*>3*σ*(I)]", "R2"]])
            assert merge_crystal([t1_, t2_])[0][1][0] == ["Final *R* indices [*I*>2*σ*(I)]", "R1",
                                                           "R2 [*I*>3*σ*(I)]"], merge_crystal([t1_, t2_])
            b2_ = _pcif(xcif_([("C1", "C", (1, 1, 1))], extra=["_refine_diff_density_max 0.42",
                                                                "_refine_diff_density_min 0.31"]))[0]
            assert XRAY.crystal_values(b2_)["diff"] == "0.42/0.31", XRAY.crystal_values(b2_)["diff"]
        chk_("R threshold, method, hole", p4_threshold)
        def p4_markup_names():                       # „Cp*“ im Namen: nicht kursiv/verschluckt, kein Absturz
            for nm_ in ("Cp*2Fe", "A_{1", "x^2~y"):
                assert XRAY.to_plain(XRAY.esc_markup(nm_)) == nm_ and XRAY.to_html(XRAY.esc_markup(nm_)) == nm_
        chk_("markup escaping", p4_markup_names)
        def p4_disorder_sites():                     # zwei unabhängige Fehlordnungen ohne assembly ('.')
            S_x, _ = xmol_(xcif_([("C1", "C", (5, 5, 5)), ("C2A", "C", (6.5, 5, 5), 1, 0.7),
                                  ("C2B", "C", (6.3, 5.6, 5), 2, 0.3), ("C3A", "C", (3.5, 5, 5), 1, 0.4),
                                  ("C3B", "C", (3.7, 5.6, 5), 2, 0.6)]))
            assert [S_x.atoms[i_]["label"] for i_ in S_x.visible_asu("main")] == ["C1", "C2A", "C3B"]
            assert XRAY.disorder_occupancy(S_x) is None
            S_x, _ = xmol_(xcif_([("C1", "C", (5, 5, 5)), ("C2A", "C", (6.5, 5, 5), 1, 0.545),
                                  ("C2B", "C", (6.3, 5.6, 5), 2, 0.455)]))
            assert XRAY.disorder_occupancy(S_x) == "55% (A) and 45% (B)", XRAY.disorder_occupancy(S_x)
        chk_("disorder sites", p4_disorder_sites)
        def p4_reader_misc():                        # Typ '?' -> aus dem Label; ohne Symmetrie-Loop Hinweis
            S_x, _ = xmol_(xcif_([("H1A", "?", (1, 1, 1))]))
            assert S_x.atoms[0]["el"] == "H", S_x.atoms[0]["el"]
            txt_ = "\n".join(l_ for l_ in xcif_([("C1", "C", (1, 1, 1))]).splitlines()
                             if "symop" not in l_ and "x, y, z" not in l_).replace("loop_\nloop_", "loop_")
            S_x = XRAY.Structure(_pcif(txt_ + "\n_space_group_name_H-M_alt 'P 21/c'\n")[0])
            assert S_x.no_symops, "kein Hinweis"
        chk_("reader type '?' / no symops", p4_reader_misc)
        def p4_plane_collinear():
            _, m_ = xmol_(xcif_([("C1", "C", (1, 1, 1)), ("C2", "C", (2, 1, 1)), ("C3", "C", (3, 1, 1))]))
            try:
                XRAY.align_plane(np.eye(3), m_, ["C1", "C2", "C3"])
                raise AssertionError("Ebene aus drei Atomen auf einer Linie")
            except ValueError:
                pass
        chk_("plane from collinear atoms", p4_plane_collinear)
        class _XS:                                   # Tab-Attrappe für Tabellen (nur Modell)
            _sym_ops = XRAY.XrayTab._sym_ops
            def __init__(s_, S_x, m_, params):
                s_.e_ = dict(name="t", block="t", params=params)
                s_.S_x, s_.m_ = S_x, m_
            def entries(s_):
                return [s_.e_]
            def structure(s_, e_):
                return s_.S_x
            def molecule(s_, e_):
                return s_.m_
        def p4_params_rows():                        # C1–C2 und C1–C2ⁱ derselben Struktur: beide Zeilen
            S_x, m_ = xmol_(xcif_([("C1", "C", (1, 0, 0)), ("C2", "C", (2.5, 0, 0))], ops=inv2_,
                                  bonds=["C1 C2 1.500(2) . .", "C1 C1 2.000(2) . 2"]))
            g_ = [a_["key"] for a_ in m_.atoms if a_["sym"] and a_["label"] == "C2"]
            assert g_, [a_["key"] for a_ in m_.atoms]
            st_ = _XS(S_x, m_, [{"atoms": ["C1", "C2"]}, {"atoms": ["C1", g_[0]]}])
            _, rows_, _ = XRAY.XrayTab.params_table(st_)
            vals_ = [r_[1] for r_ in rows_[1:]]
            assert len(vals_) == 2 and vals_[0] == "1.500(2)" and vals_[1].startswith("3.500"), rows_
            assert XRAY.XrayTab.sym_note(st_, st_.e_).startswith("Symmetry codes: i: −x"), "Word-Fußnote fehlt"
        chk_("params table rows / Word symmetry note", p4_params_rows)
        def p4_template_core():                      # tBu-reiche Vorlage: beste Ansicht trotz vieler Permutationen
            from rdkit import Chem as _Ch
            from rdkit.Chem import AllChem as _AC
            m3_ = _Ch.AddHs(_Ch.MolFromSmiles("CC(C)(C)c1cc(C(C)(C)C)c(Oc2c(C(C)(C)C)cc(C(C)(C)C)cc2C(C)(C)C)"
                                              "c(C(C)(C)C)c1"))
            _AC.EmbedMolecule(m3_, randomSeed=7)
            _AC.MMFFOptimizeMolecule(m3_)
            cf_ = m3_.GetConformer()
            ats_ = [dict(el=a_.GetSymbol(), xyz=np.array(list(cf_.GetAtomPosition(a_.GetIdx()))), frag=0,
                         key=str(a_.GetIdx())) for a_ in m3_.GetAtoms()]
            mm_ = SimpleNamespace(atoms=ats_, bonds={(b_.GetBeginAtomIdx(), b_.GetEndAtomIdx())
                                                    for b_ in m3_.GetBonds()})
            t2_ = _Ch.RemoveHs(_Ch.Mol(m3_))
            t2_.RemoveAllConformers()
            _AC.Compute2DCoords(t2_)
            tp_ = xdir / "tbu_template.cdxml"
            tp_.write_text(_Ch.MolToCDXMLBlock(t2_), encoding="utf-8")
            R_, inf_ = XRAY.align_to_template(mm_, tp_)
            assert inf_["rmsd"] < 0.8 and abs(np.linalg.det(R_) - 1) < 1e-9, inf_
        chk_("template with many tBu", p4_template_core)
        assert not fails_, "\n".join(fails_)
        report.append("audit P4 X-Ray edge cases: OK")
        # Audit P4, Codex-Gegenlesen: jeder Fall scheiterte an e4c5324
        fails_ = []
        def p4r_params_order():                      # Symmetrie-Variante zuerst gewählt: beide Zeilen bleiben
            S_x, m_ = xmol_(xcif_([("C1", "C", (1, 0, 0)), ("C2", "C", (2.5, 0, 0))], ops=inv2_,
                                  bonds=["C1 C2 1.500(2) . .", "C1 C1 2.000(2) . 2"]))
            g_ = [a_["key"] for a_ in m_.atoms if a_["sym"] and a_["label"] == "C2"][0]
            _, rows_, _ = XRAY.XrayTab.params_table(_XS(S_x, m_, [{"atoms": ["C1", g_]}, {"atoms": ["C1", "C2"]}]))
            vals_ = sorted(r_[1] for r_ in rows_[1:])
            assert len(vals_) == 2 and vals_[0] == "1.500(2)" and vals_[1].startswith("3.500"), rows_
        chk_("params table, symmetric variant first", p4r_params_order)
        def p4r_near_mirror():                       # fast auf einer Spiegelebene: kein CIF-Text für einen anderen Satz
            _, m_ = xmol_(xcif_([("C1", "C", (1, 1, 0.004)), ("C2", "C", (2.5, 1, 0.004)), ("C3", "C", (3.2, 2.3, 0.004)),
                                 ("C4", "C", (4.7, 2.3, 0.006))], ops=("x, y, -z", "x, y, z"),
                                tors=["C1 C2 C3 C4 179.9(3) . . . ."]))
            ks_ = ["C1", "C2|x, y, −z", "C3", "C4|x, y, −z"]
            assert all(m_.find_key(k_) is not None for k_ in ks_), [a_["key"] for a_ in m_.atoms]
            r_ = XRAY.measure(m_, ks_)
            assert r_["source"] == "calc", r_
        chk_("near-mirror torsion not matched", p4r_near_mirror)
        def p4r_merge_mixed():                       # gemischte Schwellen in der Gesamt-SI: keine doppelte Angabe
            from elucida.si.publication_si import merge_crystal
            L2_, L3_ = "Final *R* indices [*I*>2*σ*(I)]", "Final *R* indices [*I*>3*σ*(I)]"
            row_ = merge_crystal([(["Compound", "A"], [[L3_, "RA"]]),
                                  (["Compound", "B", "C"], [[L2_, "RB", "RC [*I*>3*σ*(I)]"]])])[0][1][0]
            assert row_ == [L3_, "RA", "RB [*I*>2*σ*(I)]", "RC"], row_
        chk_("merged mixed thresholds", p4r_merge_mixed)
        def p4r_template_small():                    # kleines Molekül: nie schlechter als die vollständige Suche
            from rdkit import Chem as _Ch
            from rdkit.Chem import AllChem as _AC
            m3_ = _Ch.AddHs(_Ch.MolFromSmiles("OC(=O)CC(C)C"))
            _AC.EmbedMolecule(m3_, randomSeed=1)
            _AC.MMFFOptimizeMolecule(m3_)
            cf_ = m3_.GetConformer()
            mm_ = SimpleNamespace(atoms=[dict(el=a_.GetSymbol(), xyz=np.array(list(cf_.GetAtomPosition(a_.GetIdx()))),
                                              frag=0, key=str(a_.GetIdx())) for a_ in m3_.GetAtoms()],
                                  bonds={(b_.GetBeginAtomIdx(), b_.GetEndAtomIdx()) for b_ in m3_.GetBonds()})
            t2_ = _Ch.RemoveHs(_Ch.Mol(m3_))
            t2_.RemoveAllConformers()
            _AC.Compute2DCoords(t2_)
            tp_ = xdir / "small_template.cdxml"
            tp_.write_text(_Ch.MolToCDXMLBlock(t2_), encoding="utf-8")
            inf_ = XRAY.align_to_template(mm_, tp_)[1]
            assert inf_["rmsd"] < 0.7, inf_          # vollständige Suche 0.57, nur Gerüst 0.85
        chk_("template small molecule", p4r_template_small)
        assert not fails_, "\n".join(fails_)
        report.append("audit P4 review fixes (Codex): OK")
        # Zeichenfläche darf durch Leisten nicht schrumpfen: bei 1440 × 900 mindestens die halbe Fensterbreite
        # (gemessen ~55 %; eine Seitenleiste läge bei ~44 %). Relativ, damit andere Schriften in CI nicht stören.
        size0_ = w.size()
        w.show()
        w.resize(1440, 900)
        w.main_tabs.setCurrentWidget(t_x)
        pump(0.5)
        t_x.host.fit()
        pump(0.2)
        if t_x.canvas.width() < 0.5 * w.width():                  # Messwerte für die Fehlersuche (andere Systeme)
            left_ = t_x.lv.parentWidget()
            cards_ = sorted(((c_.minimumSizeHint().width(), getattr(c_, "title", lambda: type(c_).__name__)()[:28])
                             for c_ in left_.findChildren(QWidget) if c_.parentWidget() is left_ and c_.isVisibleTo(left_)),
                            reverse=True)[:4]
            wide_ = sorted(((x_.minimumSizeHint().width(), type(x_).__name__, (getattr(x_, "text", lambda: "")() or "")[:30])
                            for x_ in left_.findChildren(QWidget) if x_.isVisibleTo(left_)), reverse=True)[:6]
            raise AssertionError(f"Abbildung {t_x.canvas.width()} px bei Fenster {w.width()} px; linke Spalte "
                                 f"{left_.width()} px (min {left_.minimumSizeHint().width()}), Karten {cards_}, "
                                 f"breiteste Elemente {wide_}, Schrift {QApplication.font().family()} "
                                 f"{QApplication.font().pointSizeF()} pt, DPI {w.logicalDpiX()}")
        w.resize(size0_)
        report.append(f"X-Ray (CIF, symmetry, disorder, ORTEP, picking, tables, wireframe, caption): OK "
                      f"({npath_} pdf paths)")
        w.xray_state.update(entries=[], current=None, show_h=True, hide_solv=False, label_mode="all")
        t_x.sel = []
        t_x.sync_widgets()
        t_x.render()

        # DFT: xyz (zwei Bilder, CRLF), Bindungen, Umnummerieren per Klick, xyz speichern (Zeilen unverändert)
        ddir = out / "input" / "dft"
        xyzf_ = DFT.make_test_xyz(ddir / "dtest.xyz")
        fr_ = DFT.read_xyz(xyzf_)
        assert len(fr_) == 2 and fr_[-1]["els"] == ["H", "C", "Cl", "C", "H", "H", "H", "H"], fr_[-1]["els"]
        assert fr_[-1]["comment"] == "E -539.123456789 (test)" and not fr_[0]["lines"][0].endswith("\r")
        assert DFT.parse_xyz("1\n\n6 0 0 0\n")[0]["els"] == ["C"]                 # Ordnungszahl statt Symbol
        for bad_ in ("C 0 0 0\n", "3\n\nC 0 0 0\n"):
            try:
                DFT.parse_xyz(bad_)
                raise AssertionError(bad_)
            except ValueError:
                pass
        bo_ = DFT.bonds_of(fr_[-1]["els"], fr_[-1]["xyz"])
        assert sorted(bo_) == [(0, 3), (1, 2), (1, 3), (1, 4), (1, 5), (3, 6), (3, 7)], bo_
        nb_ = {k_: [] for k_ in range(8)}
        for a_, b_ in bo_:
            nb_[a_].append(b_)
            nb_[b_].append(a_)
        assert DFT.renumber(range(8), [1, 3, 2]) == [1, 3, 2, 0, 4, 5, 6, 7]
        # H folgen ihrem Atom: angeklickte behalten 1–3, H danach nach ihrem Atom (C1: 4, 5; C3: 0, 6, 7)
        assert DFT.renumber(range(8), [1, 3, 2], fr_[-1]["els"], nb_) == [1, 3, 2, 4, 5, 0, 6, 7]
        assert DFT.renumber([2, 1, 0, 3, 4, 5, 6, 7], [3]) == [3, 2, 1, 0, 4, 5, 6, 7]   # Rest: bisherige Reihenfolge
        t_d = w.dft_tab
        w.main_tabs.setCurrentWidget(t_d)
        e_d = t_d.load_xyz(str(xyzf_))
        assert e_d and Path(e_d["file"]).parent == proj / "data" / "dft"
        assert any(k_ == "file" and c_ is e_d for c_, k_, _ in PRJ._path_slots(w.get_state()))
        pump()
        assert t_d.F is not None and len(t_d.F.scene["atoms"]) == 8 and len(t_d.F.scene["bonds"]) == 7
        assert t_d.F.ann_items["lbl_a2"].get_text() == "Cl3"
        proc_d = t_d.canvas.callbacks.process

        def click_d(k_):
            it_ = t_d.F.scene["by_k"][k_]
            px_, py_ = t_d.F.ax.transData.transform((it_["x"], it_["y"]))
            proc_d("button_press_event", MouseEvent("button_press_event", t_d.canvas, px_, py_, button=1))
            proc_d("button_release_event", MouseEvent("button_release_event", t_d.canvas, px_, py_, button=1))
            pump(0.05)
        click_d(1)
        assert t_d.picks == [], t_d.picks                                     # ohne Modus: Klick nummeriert nicht
        t_d.btn_num.setChecked(True)
        for k_ in (1, 3, 2):
            click_d(k_)
        assert t_d.picks == [1, 3, 2], t_d.picks
        click_d(3)                                                             # erneuter Klick nimmt es heraus
        assert t_d.picks == [1, 2], t_d.picks
        click_d(3)
        assert t_d.picks == [1, 2, 3], t_d.picks
        assert t_d.F.ann_items["lbl_a3"].get_text() == "C3" and t_d.F.ann_items["lbl_a3"].get_fontweight() == "bold"
        t_d.picks = [1, 3, 2]
        t_d.apply_numbering()
        assert e_d["order"] == [1, 3, 2, 0, 4, 5, 6, 7] and not t_d.numbering and not t_d.btn_num.isChecked()
        assert t_d.F.ann_items["lbl_a2"].get_text() == "Cl3" and t_d.F.ann_items["lbl_a3"].get_text() == "C2"
        w.dft_state["label_pt"] = 14.0                                       # eigene Labelgröße des Tabs
        t_d.render(push=False)
        assert t_d.F.ann_items["lbl_a1"].get_fontsize() == 14.0
        w.dft_state["label_pt"] = 0.0
        w.dft_state["zero_based"] = True
        t_d.render(push=False)
        assert t_d.F.ann_items["lbl_a1"].get_text() == "C0"
        w.dft_state["zero_based"] = False
        pump(0.3)
        w.undo()                                                               # Umnummerieren rückgängig
        pump(0.4)
        assert w.dft_state["entries"][0]["order"] is None, w.dft_state["entries"][0]["order"]
        w.redo()
        pump(0.4)
        t_d = w.dft_tab
        assert w.dft_state["entries"][0]["order"] == [1, 3, 2, 0, 4, 5, 6, 7]
        sx_ = t_d.save_xyz(out / "dft" / "dtest_new.xyz")
        fr2_ = DFT.read_xyz(sx_)
        assert len(fr2_) == 2 and fr2_[-1]["els"] == ["C", "C", "Cl", "H", "H", "H", "H", "H"]
        assert fr2_[0]["lines"] == [fr_[0]["lines"][k_] for k_ in (1, 3, 2, 0, 4, 5, 6, 7)]    # Zeilen unverändert
        src_ = Path(w.dft_state["entries"][0]["file"])
        assert t_d.save_xyz(src_) is None and DFT.read_xyz(src_)[0]["lines"] == fr_[0]["lines"]  # nie überschreiben
        files = t_d.export_to(out / "dft" / "dft_test")
        assert {f_.suffix for f_ in files} == {".pdf", ".svg", ".png", ".xyz"}, files
        assert all(f_.exists() and f_.stat().st_size > 100 for f_ in files)
        t_d.reset_order()
        assert w.dft_state["entries"][0]["order"] is None
        # Rechnungen: ORCA 6 (UTF-16/CRLF, CPCM, eine imaginäre Frequenz) und Gaussian (Format geraten); SI-Block, Word
        orf_, gf_ = DFT.make_test_orca(ddir / "otest.out"), DFT.make_test_gaussian(ddir / "gtest.log")
        assert DFT.is_calc(orf_) and DFT.is_calc(gf_) and not DFT.is_calc(xyzf_)
        co_ = DFT.read_calc(orf_)
        assert (co_["version"], co_["charge"], co_["mult"], co_["opt_converged"], co_["normal"]) == ("6.1.1", 0, 1, True, True)
        assert DFT.imaginary(co_) == ["-31.20"] and co_["thermo"]["G"] == "-539.94683332", co_["thermo"]
        assert co_["solvation"]["e_diel"] == "-0.00412345678901" and co_["solvation"]["solvent"] == "TETRAHYDROFURAN"
        assert co_["conv"][0] == ("Energy change", "-0.0000002000", "0.0000010000", "YES") and len(co_["conv"]) == 5
        cg_ = DFT.read_calc(gf_)
        assert cg_["solvation"]["model"] == "SMD" and cg_["thermo"]["H"] == "-539.916710" and len(cg_["conv"]) == 4
        assert cg_["energy"] == "-539.987654321" and not DFT.imaginary(cg_) and cg_["geometry"][1][1][0] == "Cl"
        # gleichnamige xyz (ORCA, 14 Stellen) wird mitkopiert und benutzt, wenn sie zur Endgeometrie passt
        with open(ddir / "otest.xyz", "w", encoding="utf-8", newline="\n") as fh_:
            fh_.write("8\ntest\n" + "".join(f"{el_}  {float(x_):.14f}  {float(y_):.14f}  {float(z_):.14f}\n"
                                            for el_, x_, y_, z_ in DFT._TEST_GEOM))
        fr_p, used_ = DFT.calc_frame(co_, partner=ddir / "otest.xyz")
        assert used_ and fr_p["lines"][1].startswith("Cl  0.00000000000000"), fr_p["lines"][1]
        assert not DFT.calc_frame(co_, partner=xyzf_)[1]                      # andere Geometrie -> ignoriert
        assert co_["point_group"] == "C1" and cg_["point_group"] == "CS" and DFT.pg_html("CS") == "C<sub>s</sub>"
        e_o = t_d.load_xyz(str(orf_))
        assert (Path(e_o["file"]).with_suffix(".xyz")).is_file()
        t_d.load_xyz(str(gf_))
        e_o["order"] = [1, 0, 2, 3, 4, 5, 6, 7]
        w.dft_state["current"] = e_o["uid"]
        t_d.sync_widgets()
        t_d.render()
        assert t_d.F is not None and len(t_d.F.scene["atoms"]) == 8
        lab_ = [k_ for k_ in range(8) if t_d.F.mol.atoms[k_]["el"] != "H"]    # Labels mehrerer Atome (Rahmen)
        t_d.set_labels(lab_, False)
        assert all(e_o["label_on"].get(t_d.F.mol.atoms[k_]["key"]) is False for k_ in lab_), e_o["label_on"]
        t_d.set_labels(lab_, True)
        assert not e_o["label_on"], e_o["label_on"]                            # Vorgabe: Ausnahmen fallen weg
        pl_ = DFT.si_plain(t_d.si_of(e_o))
        for frag_ in ("ORCA 6.1.1: ! B3LYP D4 def2-SVP TightOpt Freq CPCM(thf)",
                      "Charge: 0\tMultiplicity: 1\tPoint group: C1", "1 imaginary frequency (\u221231.20 cm-1)",
                      "thermochemistry at 298.15 K and 1.00 atm (quasi-RRHO)",
                      "CPCM dielectric energy (included in E(el)): \u22120.00412345678901 Eh",
                      "Energy change\t\u22120.0000002000\t0.0000010000\tyes", "\nCl  0.00000000000000",
                      "E_el [Eh]\tH [Eh]\tG [Eh]\n\u2212539.98765432\t\u2212539.91670987\t\u2212539.94683332"):
            assert frag_ in pl_, (frag_, pl_)
        assert pl_.index("Energy change") < pl_.index("Frequency calculation")      # Tabelle unter der Optimierung
        assert pl_.index("\nCl  ") < pl_.index("\nC  ")                              # Koordinaten umnummeriert
        assert T("dft_partner_used") in t_d.lbl_info.text()
        w.dft_state["si_conv"] = False                                             # Teile abschaltbar
        assert "Energy change" not in DFT.si_plain(t_d.si_of(e_o))
        w.dft_state["si_conv"] = True
        assert T("dft_w_solv_cpcm") in t_d.lbl_warn.text() and t_d.btn_si_copy.isEnabled()
        files = t_d.export_to(out / "dft" / "dft_calc")
        names_ = {f_.name for f_ in files}
        assert {"dft_calc_SI.docx", "dft_calc_SI.txt"} <= names_, names_
        import zipfile as _zf
        with _zf.ZipFile(out / "dft" / "dft_calc_SI.docx") as z_:
            doc_ = z_.read("word/document.xml").decode("utf-8")
            assert "word/media/image1.png" in z_.namelist() and "word/media/image2.png" in z_.namelist()
        for frag_ in ("Point group", "Courier New", 'w:cols w:num="2"', "Max. gradient", "otest", "Gaussian 16",
                      "quasi-RRHO", "<w:tab/>"):
            assert frag_ in doc_, frag_
        # Messen und Ausrichten wie im Tab X-Ray (gemeinsame Klasse XRAY.MeasureAlign), Werte aus der Geometrie
        t_d.set_numbering(False)
        g_d = t_d.geometry(t_d.entry())
        xyz_d = np.asarray(g_d[2], float)
        for keys_, kind_ in ((["a1", "a3"], "bond"), (["a2", "a1", "a3"], "angle"), (["a2", "a1", "a3", "a0"], "torsion")):
            t_d.sel = []
            for k_ in keys_:
                t_d.atom_clicked(int(k_[1:]))
            assert t_d.sel == keys_, t_d.sel
            m_d = t_d._meas
            idx_ = [int(k_[1:]) for k_ in keys_]
            ref_ = {"bond": XRAY.distance, "angle": XRAY.angle, "torsion": XRAY.torsion}[kind_]([xyz_d[i_] for i_ in idx_])
            assert m_d["kind"] == kind_ and abs(m_d["value"] - ref_) < 1e-9 and m_d["source"] == "calc", m_d
            assert "\n" not in t_d.lbl_meas.text() and t_d.lbl_meas.styleSheet() == "", t_d.lbl_meas.text()  # keine esd-Warnung
        t_d.sel = ["a1", "a2"]
        t_d.align_bond("down")                                 # Bindung C–Cl zeigt im Bild nach unten
        R_d = np.array(t_d.entry()["view"]).reshape(3, 3)
        v_d = R_d @ (xyz_d[2] - xyz_d[1])
        assert v_d[1] < -0.99 * np.linalg.norm(v_d) and t_d.sel == [], v_d
        t_d.sel = ["a1", "a2", "a3"]
        t_d.align_plane()
        R_d = np.array(t_d.entry()["view"]).reshape(3, 3)
        n_d = np.cross(xyz_d[2] - xyz_d[1], xyz_d[3] - xyz_d[1])
        assert abs(abs((R_d @ n_d)[2]) - np.linalg.norm(n_d)) < 1e-6           # Ebene von vorn
        t_d.sel = ["a1"]
        t_d.set_numbering(True)                                # Nummerieren hebt die Messauswahl auf
        assert t_d.sel == [] and t_d.picks == []
        t_d.set_numbering(False)
        # Kugel-Stab: Atomradius/Bindungsdicke relativ zum Standard; Hintergrund wählbar (auch im Export)
        # (alte Projekte mit „ballstick“/IboView zeichnen als Kugel-Stab „depth“; Auswahl ohne IboView, mit Strichzeichnung)
        assert [t_d.cmb_style.itemData(i_) for i_ in range(t_d.cmb_style.count())] == list(XRAY.DFT_STYLES)
        w.dft_state.update(style="ballstick", radius=0.5, bond_w=0.16, bg="#203040")
        t_d.sync_widgets()
        t_d.render()
        assert t_d.F.style == "depth" and t_d.cmb_style.currentData() == "depth"
        w.dft_state.update(style="cartoon", cartoon_color=True)          # Strichzeichnung, Atome in Elementfarben
        t_d.sync_widgets()
        t_d.render()
        cc_ = next(i_ for i_ in t_d.F.scene["atoms"] if i_["el"] == "C")
        fcc_ = tuple(t_d.F.atom_art[cc_["k"]][0].get_facecolors()[0][:3])
        assert fcc_ != (1, 1, 1) and t_d.cb_ccol.isChecked() and not t_d.cb_ccol.isHidden(), fcc_
        w.dft_state.update(style="ballstick", cartoon_color=False)
        t_d.sync_widgets()
        t_d.render()
        c_d = next(i_ for i_ in t_d.F.scene["atoms"] if i_["el"] == "C")
        assert abs(c_d["w"] - 4 * XRAY.ball_radius("C", "depth")) < 1e-9, c_d["w"]
        assert abs(t_d.F.bond_w - 2 * XRAY.ROD_W["depth"]) < 1e-9 and t_d.cmb_bg.currentData() == "color"
        from PIL import Image as PILImage
        png_bg = next(f_ for f_ in t_d.export_to(out / "dft_bg") if f_.suffix == ".png")
        with PILImage.open(png_bg) as im_:
            assert im_.convert("RGB").getpixel((1, 1)) == (0x20, 0x30, 0x40), im_.convert("RGB").getpixel((1, 1))
        w.dft_state["bg"] = "transparent"
        t_d.sync_widgets()
        png_bg = next(f_ for f_ in t_d.export_to(out / "dft_bg") if f_.suffix == ".png")
        with PILImage.open(png_bg) as im_:
            assert im_.convert("RGBA").getpixel((1, 1))[3] == 0
        w.dft_state["style"] = "mercury"
        e_dr = t_d.entry()                                     # Darstellung auf Standard, Nummerierung bleibt
        ord_ = list(e_dr.get("order") or [])
        w.dft_state.update(bond_w=0.3, radius=0.5, show_h=False)
        e_dr["zoom"] = 3.0
        t_d.reset_display()
        assert w.dft_state["bond_w"] == 0.08 and w.dft_state["radius"] == 0.25 and w.dft_state["show_h"]
        assert w.dft_state["bg"] == "white" and t_d.cmb_bg.currentData() == "white"
        assert e_dr["zoom"] == 1.0 and list(e_dr.get("order") or []) == ord_ and abs(t_d.sp_bond.value() - 0.08) < 1e-9
        report.append("DFT xyz (reader, bonds, click renumbering, save, export): OK")
        report.append("DFT outputs (ORCA 6 UTF-16, Gaussian, SI block, Word): OK")
        w.main_tabs.setCurrentWidget(t_d)
        pump(0.2)
        assert not t_d.numbering
        rot_keeps_scale_(t_d)
        report.append("DFT rotate keeps size (GUI path, labels moved, zoom): OK")
        # Vergleichsmodus: Nummerierung einer anderen Geometrie bzw. einer CIF übernehmen
        w.dft_state.update(entries=[], current=None)
        e_ref = t_d.load_xyz(str(orf_))                                   # Vorlage: C, Cl, C, H…
        e_tg = t_d.load_xyz(str(xyzf_))                                   # gleiche Geometrie, andere Reihenfolge
        t_d.fill_refs()
        t_d.cmb_ref.setCurrentIndex(t_d.cmb_ref.findData("dft:" + e_ref["uid"]))
        t_d.apply_reference()
        assert e_tg["order"] == [1, 2, 3, 4, 5, 6, 7, 0], e_tg["order"]
        assert "0.00" in t_d.lbl_ref.text() and not e_tg["ref_labels"], t_d.lbl_ref.text()
        w.xray_tab.load_cif(str(cif_))                                    # CIF als Vorlage: Labels der Kristallstruktur
        mol_c = w.xray_tab.molecule(w.xray_state["entries"][0])
        main_c = [k_ for k_, a_ in enumerate(mol_c.atoms) if a_["frag"] == 0]
        perm_c = main_c[::-1]                                              # umgekehrte Reihenfolge, leicht verschoben
        cxyz_ = ddir / "from_cif.xyz"
        cxyz_.write_text(f"{len(perm_c)}\nfrom cif\n" + "".join(
            f"{mol_c.atoms[k_]['el']} {x_ + 0.01:.6f} {y_:.6f} {z_:.6f}\n" for k_ in perm_c
            for x_, y_, z_ in [mol_c.atoms[k_]["xyz"]]), encoding="utf-8")
        e_cx = t_d.load_xyz(str(cxyz_))
        w.main_tabs.setCurrentWidget(w.xray_tab)
        w.main_tabs.setCurrentWidget(t_d)                                 # Tabwechsel füllt die Vorlagen neu
        t_d.cmb_ref.setCurrentIndex(t_d.cmb_ref.findData("xray:" + w.xray_state["entries"][0]["uid"]))
        t_d.apply_reference()
        labs_ = e_cx["ref_labels"]
        assert sorted(labs_.values()) == sorted(XRAY.atom_name(mol_c.atoms[k_]) for k_ in main_c), labs_
        assert all(labs_[str(i_)].rstrip("\u2071").startswith(mol_c.atoms[k_]["el"]) for i_, k_ in enumerate(perm_c))
        assert w.dft_state["label_style"] == "ref" and t_d.F.ann_items["lbl_a0"].get_text() in labs_.values()
        assert "\u26a0" in t_d.lbl_ref.text()        # zentrosymmetrisch: C1/C1(i) gleichwertig -> gemeldet
        g_cx = t_d.geometry(e_cx)
        ok_ = DFT.match_reference(dict(els=g_cx[1], xyz=g_cx[2], bonds=g_cx[3], order=None),
                                  dict(els=g_cx[1], xyz=g_cx[2], bonds=g_cx[3], order=None))
        assert ok_["order"] == list(range(len(perm_c))) or ok_["ambiguous"]
        t_d.reset_order()
        assert e_cx["order"] is None and not e_cx["ref_labels"]
        w.xray_state.update(entries=[], current=None)
        w.dft_state.update(label_style="el_num")
        report.append("DFT compare mode (number like xyz / CIF template): OK")

        # Orbitale (Cube): Leser, Isowert wie IboView (80 % der Dichte), Raycasting, Verdeckung, Tab, Export, Word
        from elucida.core import orbitals as ORB
        cub_ = ORB.make_test_cube(ddir / "eth.mo7a.cube")
        cc_ = ORB.read_cube(cub_)
        assert cc_["els"] == ["C", "C", "H", "H", "H", "H"] and cc_["data"][0].shape == (48, 48, 48)
        assert abs(cc_["xyz"][1][0] - 0.667) < 1e-5 and ORB.cube_mo(cub_, cc_) == (7, "a")
        v_ = cc_["data"][0].astype(float)
        assert abs(ORB.grid_norm(v_, cc_["axes"]) - 1) < 1e-3
        iso_ = ORB.iso_for_fraction(v_, 80)
        frac_ = float(np.sum(v_[np.abs(v_) > iso_] ** 2) / np.sum(v_ ** 2))
        assert abs(frac_ - 0.8) < 0.01, frac_
        txt_ = ("ORBITAL ENERGIES\n----------------\n\n  NO   OCC          E(Eh)            E(eV) \n"
                "   6   2.0000      -0.400000       -10.8846 \n   7   2.0000      -0.300000        -8.1634 \n"
                "   8   0.0000       0.050000         1.3606 \n   9   0.0000       0.100000         2.7211 \n\n"
                "------------\nMULLIKEN\n")
        tab_ = ORB.orbital_energies(txt_)
        assert ORB.frontier_name(7, "", tab_) == "HOMO" and ORB.frontier_name(6, "a", tab_) == "HOMO\u22121"
        assert ORB.frontier_name(9, "", tab_) == "LUMO+1" and abs(ORB.orbital_energy(8, "", tab_) - 1.3606) < 1e-9
        gtab_ = ORB.orbital_energies(" Population analysis using the SCF Density.\n Alpha  occ. eigenvalues --"
                                     "  -10.20000  -0.50000\n Alpha virt. eigenvalues --    0.10000\n")
        assert ORB.frontier_name(2, "", gtab_) == "HOMO" and ORB.frontier_name(3, "", gtab_) == "LUMO", gtab_
        # zwei Orbitale in einem Cube (Gaussian: Atomzahl negativ, MO-Liste)
        mc_ = ddir / "two.cube"
        mc_.write_text("t\nc\n   -1 0 0 0\n    2 1 0 0\n    2 0 1 0\n    2 0 0 1\n    1 1.0 0 0 0\n    2 5 6\n"
                       + " ".join(str(float(i_)) for i_ in range(16)) + "\n", encoding="utf-8")
        m2_ = ORB.read_cube(mc_)
        assert m2_["mos"] == [5, 6] and m2_["data"][1][1, 1, 1] == 15.0 and m2_["data"][0][0, 0, 1] == 2.0
        # Raycasting: beide Phasen in ihren Farben; Kugel davor verdeckt alles
        vol_ = ORB.Volume(cc_["data"][0], cc_["origin"], cc_["axes"])
        cp_, cm_, al_ = ORB.orbital_colors(0)
        Rv_ = XRAY.rot("x", -1.2)
        im_ = ORB.render(vol_, iso_, Rv_, cc_["xyz"].mean(0), 0.0, (-4, 4), (-3, 3), 160, 120, cp_, cm_, al_)
        hit_ = im_[..., 3] > 0.3
        assert hit_.sum() > 2000, hit_.sum()
        up_, dn_ = im_[80:][hit_[80:]][:, :3].mean(0), im_[:40][hit_[:40]][:, :3].mean(0)
        assert up_[0] > up_[1] and dn_[1] > dn_[0], (up_, dn_)          # oben violett (+), unten hellblau (−)
        occ_ = ORB.occluder_depth([(0.0, 0.0, 50.0, 10.0)], [], (-4, 4), (-3, 3), 160, 120)
        assert not (ORB.render(vol_, iso_, Rv_, cc_["xyz"].mean(0), 0.0, (-4, 4), (-3, 3), 160, 120, cp_, cm_, al_,
                               occ=occ_)[..., 3] > 0).any()
        # Tab: Cube ohne passende Geometrie -> neue Geometrie, Orbital gezeigt, Atome wie IboView
        w.dft_state.update(entries=[], current=None)
        w.main_tabs.setCurrentWidget(t_d)
        e_orb = t_d.load_cube(str(cub_))
        assert e_orb["orbitals"] and t_d.F.orb is not None and t_d.F.style == "iboview", t_d.F.style
        assert "HOMO" not in t_d.cmb_orb.currentText() and "MO 7" in t_d.cmb_orb.currentText()
        assert not t_d.cmb_style.isEnabled()
        t_d.canvas.draw()
        assert t_d.F.orb_art.get_array().shape[0] > 50 and t_d.F.orb_art.get_array()[..., 3].max() > 0.5
        t_d.F.set_drag(True)
        t_d.F.set_view(XRAY.rot("y", 0.3) @ t_d.F.R)
        t_d.canvas.draw()
        small_ = t_d.F.orb_art.get_array().shape[1]
        t_d.F.set_drag(False)
        t_d.canvas.draw()
        assert t_d.F.orb_art.get_array().shape[1] > 2 * small_                 # beim Ziehen gröber gerechnet
        # passender zweiter Cube hängt an derselben Geometrie, Farben je Orbital wie IboView
        cub2_ = ddir / "eth.mo8a.cube"
        cub2_.write_bytes(Path(cub_).read_bytes())
        t_d.load_cube(str(cub2_))
        assert len(w.dft_state["entries"]) == 1 and len(e_orb["orbitals"]) == 2 and e_orb["orb_cur"] == 1
        assert t_d.F.orb["plus"] == ORB.orbital_colors(1)[0]
        t_d.cmb_iso.setCurrentIndex(t_d.cmb_iso.findData("abs"))
        t_d.sp_iso.setValue(0.03)
        t_d.render(push=False)
        assert w.dft_state["orb_abs"] == 0.03 and t_d.F.orb["iso"] == 0.03 and "0.0300" in t_d.lbl_orb.text()
        t_d.cmb_iso.setCurrentIndex(t_d.cmb_iso.findData("frac"))
        files_o = t_d.export_to(out / "dft_orb" / "orb")
        names_o = sorted(Path(f_).name for f_ in files_o)
        assert "orb_MO7a.pdf" in names_o and "orb_MO8a.pdf" in names_o and "orb.pdf" in names_o, names_o
        assert b"/Subtype /Image" in (out / "dft_orb" / "orb_MO7a.pdf").read_bytes()
        assert b"/Subtype /Image" not in (out / "dft_orb" / "orb.pdf").read_bytes()   # Struktur ohne Orbital
        cap_ = t_d.orbital_caption(e_orb, e_orb["orbitals"][0], "X")
        assert cap_ == "MO 7 \u03b1 of <b>X</b>, isosurface enclosing 80% of the orbital density.", cap_
        blk_ = t_d.orbital_blocks(e_orb, "X", out / "dft_orb")
        assert [b_[0] for b_ in blk_] == ["image", "text", "image", "text"]
        t_d.cmb_orb.setCurrentIndex(0)
        assert t_d.F.orb is None and t_d.cmb_style.isEnabled()
        w.dft_state.update(entries=[], current=None)
        t_d.sync_widgets()
        t_d.render(push=False)
        report.append("DFT orbitals (cube, 80% isosurface like IboView, ray casting, occlusion, export): OK")
        # Molden: Basis bis g, Konvention über CᵀSC = 1 (Standard / ORCA: rohe Primitive), Werte, HOMO/LUMO, Tab
        for orca_ in (False, True):
            mf_ = ORB.make_test_molden(ddir / f"t{int(orca_)}.molden", orca=orca_)
            mm_ = ORB.parse_molden(Path(mf_).read_text())
            bb_ = ORB.Basis(mm_)
            assert bb_.nao == 27 and bb_.setup() < 1e-6 and bb_.variant[0] == ("raw" if orca_ else "norm"), bb_.errors
            assert min(v_ for k_, v_ in bb_.errors.items() if k_[0] != bb_.variant[0]) > 0.1
            phi_ = sum(c_ * ORB._prim_norm(a_, 0) * np.exp(-a_ * 0.49) for a_, c_ in
                       zip([3.42525091, 0.62391373, 0.16885540], [0.15432897, 0.53532814, 0.44463454]))
            val_ = bb_.evaluate(mm_["mos"][0]["c"], np.array([[0.0, 0.0, 0.7]]))[0]
            assert abs(val_ - 2 * phi_ * mm_["mos"][0]["c"][0]) < 1e-6, (val_, phi_)
            tb_ = ORB.molden_table(mm_)
            assert ORB.frontier_name(tb_[""][0][0], "", tb_) == "HOMO" and tb_[""][0][0] == (0 if orca_ else 1)
        dat_, org_, axs_ = ORB.molden_grid(bb_, mm_["mos"][0]["c"], 0.15)
        assert abs(ORB.grid_norm(dat_, axs_) - 1) < 0.02 and org_[2] < -1.0 / 1.9 and dat_.shape[0] < 60
        w.dft_state.update(entries=[], current=None)
        e_m = t_d.load_molden(str(mf_))
        assert [it_["mo"] for it_ in e_m["orbitals"]] == [0, 1] and e_m["orb_cur"] == 0, e_m["orbitals"]
        assert t_d.cmb_orb.currentText().startswith("HOMO (MO 0") and t_d.row_mo.isVisibleTo(t_d)
        assert t_d.F.orb is not None and t_d.F.orb_art is not None and "Molden" in t_d.lbl_orb.text()
        t_d.cmb_mo.setCurrentIndex(5)
        t_d.add_molden_mo()
        assert e_m["orb_cur"] == 2 and e_m["orbitals"][2]["mo"] == 5 and "LUMO+4" in t_d.cmb_orb.currentText()
        bad_ = ddir / "bad.molden"
        bad_.write_text(Path(mf_).read_text().replace("    1     0.", "    1     5.", 1), encoding="utf-8")
        assert ORB.Basis(ORB.parse_molden(bad_.read_text())).setup() > ORB.Basis.MAX_ORTH_ERR   # wird abgelehnt
        w.dft_state.update(entries=[], current=None)
        t_d.sync_widgets()
        t_d.render(push=False)
        # Bindungsordnungen wie IboView (IAO-Wiberg): H2 Einfachbindung, Ne ohne Bindung; Doppelbindung -> zwei Stäbe
        bo_ = ORB.iao_bond_orders(bb_)
        assert list(bo_) == [(0, 1)] and abs(bo_[(0, 1)] - 1) < 0.02, bo_
        e_m = t_d.load_molden(str(mf_))
        assert t_d.bond_orders(e_m) is not None and t_d.F.opts.get("bond_orders") and t_d.F.elem_lab
        w.dft_state.update(entries=[], current=None)
        t_d.sync_widgets()
        t_d.render(push=False)
        Xd_ = np.array([[0.0, 0, 0], [1.2, 0, 0], [-0.6, 0.9, 0], [-0.6, -0.9, 0]])
        md_ = DFT.XyzMolecule(["C", "O", "H", "H"], Xd_, [(0, 1), (0, 2), (0, 3)], ["C1", "O2", "H3", "H4"])
        assert np.allclose(np.abs(ORB.ibo_tan_u(Xd_, ["C", "O", "H", "H"], 0, 1)), [0, 1, 0])   # in der Molekülebene
        from matplotlib.figure import Figure as Fig_
        fg_ = Fig_(figsize=(6, 4), dpi=80)
        FigureCanvasAgg(fg_)
        cfg_d = core.deep_merge(core.DEFAULTS, {})
        Fd_ = XRAY.XrayFigure(cfg_d, md_, np.eye(3), dict(style="iboview", labels=False,
                                                         bond_orders={(0, 1): 1.95, (0, 2): 0.93, (0, 3): 1.45}), {}, fg_)
        sp_ = Fd_.rods["specs"]
        assert sum(1 for x_ in sp_ if x_["bond"] == 0) == 4 and sum(1 for x_ in sp_ if x_["bond"] == 1) == 2
        assert sum(1 for x_ in sp_ if x_["bond"] == 2 and x_["r0"] == x_["r1"]) > 2       # 1.45: zweiter Stab gestrichelt
        offs_ = {round(float(x_["off"][1]), 4) for x_ in sp_ if x_["bond"] == 0}
        assert len(offs_) == 2 and abs(sum(offs_)) < 1e-9                                 # symmetrisch neben der Achse
        fg_.canvas.draw()
        Fd_.set_view(XRAY.rot("y", 1.2))                                   # C=O zeigt zum Betrachter: Stab vor dem C
        tw_ = Fd_.atom_rods_over[0].get_paths()
        assert any(np.ptp(p_.vertices[:, 0]) > 0 for p_ in tw_)
        st_p_ = {"dft": {"entries": [{"file": "g.xyz", "molden": "g.molden", "orbitals": [{"file": "g.cube"}]}]}}
        assert {PRJ._get(c_, k_) for c_, k_, _ in PRJ._path_slots(st_p_)} >= {"g.xyz", "g.molden", "g.cube"}   # im Projekt
        report.append("DFT orbitals from Molden (basis up to g, convention check, HOMO/LUMO, IAO bond orders): OK")

        # Audit P5 (DFT, TD-DFT, Orbitale): je Fall ein Fund, der vor dem Fix scheiterte
        from elucida.formats import calc as CALC
        from elucida.formats.elements import element_of as elem_of_
        from elucida.core import tddft as XT5
        fails_ = []
        def chk_(name, fn):
            try:
                fn()
            except Exception as e_:
                fails_.append(f"{name}: {type(e_).__name__} {e_}"[:160])
        def g_or_(title, x):                     # Gaussian-Orientierungstabelle (H2O)
            return [f"                         {title}", " " + "-" * 69, " Center  Atomic  Atomic  Coordinates (Angstroms)",
                    " Number  Number  Type  X  Y  Z", " " + "-" * 69,
                    f"      1          8           0     {x:>9s}    0.000000    0.000000",
                    "      2          1           0     0.000000    0.757000    0.586000",
                    "      3          1           0     0.000000   -0.757000    0.586000", " " + "-" * 69]
        g_head_ = [" Entering Gaussian System, Link 0=g16", " ---", " #p B3LYP/def2SVP opt", " ---",
                   " Charge =  0 Multiplicity = 1"]
        g_link_ = [" Normal termination of Gaussian 16", " Link1:  Proceeding to internal job step number  2."]
        def p5_g_orient():                       # spätere Input orientation nach Link1 gilt
            c_ = CALC.parse_gaussian("\n".join(g_head_ + g_or_("Standard orientation:", "0.000000") + g_link_ +
                                               g_or_("Input orientation:", "2.000000")))
            assert c_["geometry"][0][1][1] == "2.000000", c_["geometry"][0]
        chk_("gaussian orientation after Link1", p5_g_orient)
        def p5_g_freq_jobs():                    # Frequenzen des letzten Jobs; Single Point danach: E(el) des Freq-Jobs
            c_ = CALC.parse_gaussian("\n".join(
                g_head_ + g_or_("Standard orientation:", "0.0") + [" Frequencies --   -100.0   200.0   300.0"] + g_link_ +
                g_or_("Standard orientation:", "0.0") +
                [" SCF Done:  E(RB3LYP) =  -76.400000000     A.U. after 1 cycles",
                 " Frequencies ---   100.1234   250.2345   350.3456",          # hpmodes: Kopie mit „---“
                 " Frequencies --   100.1234   250.2345   350.3456",
                 " Sum of electronic and thermal Free Energies=        -76.390000"] + g_link_ +
                [" SCF Done:  E(RwB97XD) =  -76.500000000     A.U. after 1 cycles", " Normal termination of Gaussian 16"]))
            assert c_["freqs"] == ["100.1234", "250.2345", "350.3456"], c_["freqs"]
            assert c_["thermo"]["E_el"] == "-76.400000000", c_["thermo"]
        chk_("gaussian frequencies of the last job", p5_g_freq_jobs)
        def p5_g_route():                        # Route am Leerzeichen umbrochen: Wörter bleiben getrennt
            c_ = CALC.parse_gaussian("\n".join([" Entering Gaussian System", " ---", " #p B3LYP/def2SVP opt",
                                                "  freq scrf=(smd,solv", " ent=thf)", " ---"] + g_or_("Standard orientation:", "0.0")))
            assert c_["keywords"] == ["#p B3LYP/def2SVP opt freq scrf=(smd,solvent=thf)"] and c_["freq_requested"], c_
        chk_("gaussian wrapped route", p5_g_route)
        o_geo_ = ["CARTESIAN COORDINATES (ANGSTROEM)", "---------------------------------", "  O  0.0 0.0 0.0",
                  "  H  0.0 0.757 0.586", "  H  0.0 -0.757 0.586", ""]
        def p5_o_empty_freq():                   # abgebrochener Frequenzblock: keine Aussage im SI-Text
            c_ = CALC.parse_orca("\n".join(["Program Version 6.1.1"] + o_geo_ +
                                           ["-----------------------", "VIBRATIONAL FREQUENCIES", "-----------------------", ""]))
            assert "Frequency" not in DFT.si_plain(DFT.si_parts(c_, "x", [r_[0] for r_ in c_["geometry"]])), c_["freqs"]
        chk_("orca empty frequency block", p5_o_empty_freq)
        def p5_o_energy_comment():               # Text nach der Energie, Kommentar „#“ in der Eingabezeile
            c_ = CALC.parse_orca("\n".join(["Program Version 6.1.1", "INPUT FILE", "|  1> ! B3LYP def2-SVP Opt  # test",
                                            "END OF INPUT"] + o_geo_ +
                                           ["FINAL SINGLE POINT ENERGY      -76.123456789 (Wavefunction not fully converged!)"]))
            assert c_["energy"] == "-76.123456789" and c_["keywords"] == ["! B3LYP def2-SVP Opt"], c_
        chk_("orca energy suffix and comment", p5_o_energy_comment)
        def p5_o_not_conv():                     # nicht konvergiert / ORCA 5: Hinweis nur in der Anzeige
            c_ = CALC.parse_orca("\n".join(["Program Version 5.0.4", "INPUT FILE", "|  1> ! B3LYP Opt", "END OF INPUT",
                "          ----------------------|Geometry convergence|-------------------------",
                "          Energy change      -0.0010000000            0.0000010000      NO",
                "          ---------------------------------------------------------------------"] + o_geo_))
            assert c_["opt_requested"] and not c_["opt_converged"] and len(c_["conv"]) == 1, c_
            assert "NOT" not in DFT.si_plain(DFT.si_parts(c_, "x", [r_[0] for r_ in c_["geometry"]]))
            assert any("5.0.4" in w_ for w_ in DFT.calc_warnings(c_)), DFT.calc_warnings(c_)
        chk_("not converged / ORCA 5 only as warning", p5_o_not_conv)
        def p5_kabsch_planar():                  # ebene Schweratome mit Rauschen: immer echte Drehung
            P_ = np.array([[1.25 * i_, 0.4 * (i_ % 2), 0.0] for i_ in range(9)])
            dets_ = []
            for s_ in range(40):
                r_ = np.random.default_rng(s_)
                dets_.append(np.linalg.det(DFT.kabsch(P_ + r_.normal(0, 0.01, P_.shape),
                                                      P_ + r_.normal(0, 0.01, P_.shape))[1]))
            assert min(dets_) > 0, sum(d_ < 0 for d_ in dets_)
            Q_ = np.array([[0, 0, 0], [1.5, 0, 0], [0, 1.6, 0], [0, 0, 1.7], [-0.9, -0.8, -0.7]])
            assert np.linalg.det(DFT.kabsch(Q_, Q_ * [-1, 1, 1])[1]) < 0     # Enantiomer: Spiegelung bleibt
        chk_("kabsch mirror only if clearly better", p5_kabsch_planar)
        def p5_ambig_union():                    # zwei unabhängige symmetrische Gabeln: beide gemeldet
            X_ = [[1.25 * i_, 0.4 * (i_ % 2), 0.6 * np.sin(0.5 * i_) + 0.02 * i_] for i_ in range(40)]   # chiral
            els_, bonds_ = ["N"] + ["C"] * 38 + ["O"], [(i_, i_ + 1) for i_ in range(39)]
            for k_ in (10, 30):
                for z_ in (0.14, -0.04):
                    X_.append([1.25 * k_, 0.4 * (k_ % 2) + (1.0 if k_ % 2 else -1.0), X_[k_][2] + z_])
                    els_.append("C")
                    bonds_.append((k_, len(X_) - 1))
            g_ = dict(els=els_, xyz=np.array(X_), bonds=bonds_, order=None)
            r_ = DFT.match_reference(g_, dict(g_))
            assert len(r_["ambiguous"]) == 4, r_["ambiguous"]
        chk_("ambiguity of independent groups", p5_ambig_union)
        def p5_element():                        # nur Unbekanntes auf den ersten Buchstaben kürzen
            got_ = [elem_of_(x_, x_) for x_ in ("Ho", "Nd", "CL")]
            assert got_ == ["Ho", "Nd", "Cl"], got_
        chk_("element Ho stays Ho", p5_element)
        def p5_orb_spin():                       # ROHF (nur α, Besetzung 1) und UHF (α/β je orthonormal)
            def mold_(mos_):
                t_ = ["[Molden Format]", "[Atoms] AU", "H 1 1 0.0 0.0 0.0", "H 2 1 0.0 0.0 1.4", "[GTO]"]
                for a_ in (1, 2):
                    t_ += [f"{a_} 0", "s 1 1.0", "1.0 1.0", ""]
                t_ += ["[MO]"]
                for sp_, oc_, c_ in mos_:
                    t_ += [" Sym= A", " Ene= 0.0", f" Spin= {sp_}", f" Occup= {oc_}"] + \
                          [f"{i_ + 1} {v_:.15f}" for i_, v_ in enumerate(c_)]
                return ORB.Basis(ORB.parse_molden("\n".join(t_) + "\n"))
            s_ = np.exp(-1.4 ** 2 / 2)
            b_, a_ = 1 / np.sqrt(2 + 2 * s_), 1 / np.sqrt(2 - 2 * s_)
            B_ = mold_([("Alpha", 1.0, [b_, b_]), ("Alpha", 0.0, [a_, -a_])])
            assert B_.setup() < 1e-9 and abs(ORB.iao_bond_orders(B_, 0.0)[(0, 1)] - 0.5) < 1e-6
            B_ = mold_([("Alpha", 1.0, [b_, b_]), ("Alpha", 0.0, [a_, -a_]), ("Beta", 0.0, [b_, b_]), ("Beta", 0.0, [a_, -a_])])
            assert B_.setup() < 1e-9 and abs(ORB.iao_bond_orders(B_, 0.0)[(0, 1)] - 0.5) < 1e-6
            bad_h_ = "[Molden Format]\n[Atoms] AU\nH 1 1 0 0 0\n[GTO]\n1 0\nh 1 1.0\n1.0 1.0\n\n[MO]\n"
            try:
                ORB.parse_molden(bad_h_)
                raise AssertionError("h shell accepted")
            except ValueError as ex_:
                assert "'h'" in str(ex_) and "not supported" in str(ex_), ex_
        chk_("molden ROHF/UHF bond orders, h shell", p5_orb_spin)
        def p5_orb_names():                      # Dateiname, Gaussian-Sterne, gebrochene Besetzung
            assert ORB.cube_mo("demo12.cube") == (None, "") and ORB.cube_mo("NHO.mo67a.cube") == (67, "a")
            tb_ = ORB.orbital_energies(" Population analysis using the SCF density.\n"
                                       " Alpha  occ. eigenvalues -- **********-100.12345 -10.19134 -10.18734  -0.80000\n"
                                       " Alpha virt. eigenvalues --    0.10000   0.20000\n")
            assert [n_ for n_, o_, _ in tb_[""] if o_ > 0] == [1, 2, 3, 4, 5], tb_
            assert ORB.frontier_name(6, "", tb_) == "LUMO"
            assert ORB.frontier_name(3, "", {"": [(1, 2.0, -9), (2, 1.56, -7), (3, 0.43, -1)]}) is None
        chk_("orbital names", p5_orb_names)
        def p5_td_utf16_sf():                    # UTF-16; Spin-Flip-Labels sind keine Multiplizität
            u_ = ddir / "td_u16.log"
            u_.write_bytes(" Excited State   1:      Singlet-A      3.0000 eV  413.28 nm  f=0.1000\n".encode("utf-16"))
            assert XT5.tddft_variants(u_) == ["nosoc"] and XT5.parse_tddft(u_)["f"][0] == 0.1
            sf_ = ddir / "td_sf.out"
            sf_.write_text("\n".join(["SPIN FLIP ABSORPTION SPECTRUM VIA TRANSITION ELECTRIC DIPOLE MOMENTS", "-" * 100,
                                      "  Transition  Energy", "   (eV)", "-" * 100,
                                      "  1-3A  ->  2-3A    3.269852   26373.1   379.2   0.124943095   1.5  0.1  0.1  0.1", ""]),
                           encoding="utf-8")
            assert XT5.parse_tddft(sf_)["mult"] == [None], XT5.parse_tddft(sf_)["mult"]
        chk_("TD-DFT UTF-16 and spin flip", p5_td_utf16_sf)
        def ov_spec_(name_, rows_, **st_):           # Overlay mit einer TD-DFT-Rechnung (ORCA-6-Tabelle)
            td_ = ddir / name_
            td_.write_text("\n".join(["ABSORPTION SPECTRUM VIA TRANSITION ELECTRIC DIPOLE MOMENTS", "-" * 100,
                                      "  Transition  Energy", "   (eV)", "-" * 100] +
                                     [f"  0-1A  ->  {lab_}    {e_:.6f}   1.0   1.0   {f_:.9f}   0 0 0 0"
                                      for lab_, e_, f_ in rows_] + [""]), encoding="utf-8")
            ent_ = {"kind": "calc", "file": str(td_), "label": "t", "color": "#000000", "visible": True,
                    "own_norm": False, "nlo": 300, "nhi": 600, "fwhm": 0.3, "shift": 0.0, "offset": st_.pop("offset", 0)}
            st0_ = dict(entries=[ent_], triplets=True, x_auto=True, mode="norm", sticks=True, stick_mode="band",
                        nlo=300, nhi=600)
            old_ = {k_: w.ov_state.get(k_) for k_ in st0_}
            try:
                w.ov_state.update(st0_, **st_)
                return w.ov_tab.figure_spec()
            finally:
                w.ov_state.update(old_)
        chk_("overlay triplets only", lambda: ov_spec_("td_trip.out", [("1-3A", 2.5, 0.0), ("2-3A", 3.0, 0.0)])[0])
        def p5_ov_xrange():                      # ausgeblendetes Triplett bei 1771 nm bestimmt nicht den x-Bereich
            xl_ = ov_spec_("td_mixed.out", [("1-3A", 0.7, 0.0), ("1-1A", 3.5, 0.2)], triplets=False)[0]["xlim"]
            assert xl_[1] < 1000, xl_
        chk_("overlay x range without hidden triplets", p5_ov_xrange)
        def p5_ov_faxis():                       # f-Achse beginnt bei y = 0: mit Versatz keine f-Achse
            assert "f_axis" not in ov_spec_("td_f.out", [("1-1A", 3.5, 0.2)], stick_mode="faxis", offset=0.5)[0]
        chk_("overlay f axis with offset", p5_ov_faxis)
        def p5_dft_dupes():                      # gleich benannte Geometrien: eigene Dateien; HOMO nach erneutem Laden
            w.dft_state.update(entries=[], current=None)
            e1_ = t_d.load_xyz(str(xyzf_))
            t_d.entries().append(dict(e1_, name=e1_.get("name") or Path(e1_["file"]).stem))
            e1_["name"] = t_d.entries()[1]["name"]
            fs_ = [str(f_) for f_ in t_d.export_to(out / "p5_dupes" / "d")]
            assert len(fs_) == len(set(fs_)), fs_
            w.dft_state.update(entries=[], current=None)
            e_m_ = t_d.load_molden(str(mf_))
            e_m_["orbitals"].append(dict(file=e_m_["orbitals"][0]["file"], mo=5))
            e_m_ = t_d.load_molden(str(mf_))
            assert e_m_["orbitals"][e_m_["orb_cur"]]["mo"] == 0, (e_m_["orb_cur"], e_m_["orbitals"])
            w.dft_state.update(entries=[], current=None)
            t_d.sync_widgets()
            t_d.render(push=False)
        chk_("DFT duplicate names, HOMO after reload", p5_dft_dupes)
        assert not fails_, "\n".join(fails_)
        report.append("audit P5 DFT/TD-DFT/orbitals edge cases: OK")
        fails_ = []
        def p5r_names():                         # opt, opt_3, opt: keine Endlosschleife; opt.1/opt.2 nicht beide „opt“
            got_ = DFT.export_names(["opt", "opt_3", "opt", "opt.1", "opt.2", "opt"])
            assert len(set(got_)) == 6 and not any("." in g_ for g_ in got_), got_
        chk_("export names unique without dots", p5r_names)
        chk_("overlay triplets only, eps, hidden triplets",
             lambda: ov_spec_("td_t2.out", [("1-3A", 2.5, 0.0)], mode="eps", triplets=False)[0])
        def p5r_thermo():                        # zweiter Freq-Job bricht nach H ab: kein altes G/E(el) daneben
            c_ = CALC.parse_gaussian("\n".join(
                g_head_ + g_or_("Standard orientation:", "0.0") +
                [" SCF Done:  E(RB3LYP) =  -76.400000000     A.U. after 1 cycles", " Frequencies --   100.0   250.0   350.0",
                 " Sum of electronic and thermal Enthalpies=           -76.370000",
                 " Sum of electronic and thermal Free Energies=        -76.420000"] + g_link_ + g_or_("Standard orientation:", "0.0") +
                [" SCF Done:  E(RB3LYP) =  -77.410000000     A.U. after 1 cycles", " Frequencies --   110.0   260.0   360.0",
                 " Sum of electronic and thermal Enthalpies=           -77.390000"]))
            assert c_["thermo"].get("H") == "-77.390000" and "G" not in c_["thermo"], c_["thermo"]
            assert c_["thermo"]["E_el"] == "-77.410000000", c_["thermo"]
        chk_("gaussian thermo of an aborted second job", p5r_thermo)
        def p5r_nan():                           # „**********“: kein „nan eV“ in Beschriftungen
            tb_ = ORB.orbital_energies(" Population analysis using the SCF density.\n"
                                       " Alpha  occ. eigenvalues -- **********  -0.80000\n Alpha virt. eigenvalues --  0.10000\n")
            assert ORB.orbital_energy(1, "", tb_) is None and ORB.orbital_energy(2, "", tb_) is not None, tb_
        chk_("no nan orbital energy", p5r_nan)
        assert not fails_, "\n".join(fails_)
        report.append("audit P5 review fixes (Codex): OK")

        # NMR: Mnova-SVGs (Kern aus der Achse, sonst von Hand), Sortierung, Mnova-Text + Zuordnung, Bruker acqus, Word
        ndir = out / "input" / "nmr"
        sv_h = NMR.make_test_svg(ndir / "t-1H.svg", ("1H",))
        sv_c = NMR.make_test_svg(ndir / "t-13C.svg", ("13C",), peaks=("141.20", "77.16 CDCl3", "21.05"))
        sv_q = NMR.make_test_svg(ndir / "t-HSQC.svg", ("1H", "13C"))
        sv_f = NMR.make_test_svg(ndir / "t-fluor.svg", ("19F",), axis=False)
        assert NMR.svg_info(sv_q)["nuclei"] == ["1H", "13C"] and NMR.svg_info(sv_q)["is2d"]
        assert NMR.svg_info(sv_f)["nuclei"] == [] and NMR.is_nmr_svg(sv_h)
        t_n = w.nmr_tab
        w.main_tabs.setCurrentWidget(t_n)
        for sv_ in (sv_q, sv_f, sv_c, sv_h):
            t_n.load_svg(str(sv_))
        e_f = [s_ for s_ in w.nmr_state["spectra"] if Path(s_["file"]).stem == "t-fluor"][0]
        w.nmr_state["current"] = e_f["uid"]
        t_n.sync_widgets()
        t_n.ed_nuc.setText("19F")
        t_n.spec_edited()
        assert e_f["nuclei"] == ["19F"] and not e_f["nuc_auto"]
        assert [(s_["kind"], s_["nuclei"]) for s_ in t_n.spectra()] == [("1D", ["1H"]), ("1D", ["13C"]), ("1D", ["19F"]),
                                                                        ("HSQC", ["1H", "13C"])], t_n.spectra()
        t_n.ed_text.setPlainText("1H NMR (501 MHz, CDCl3) \u03b4 7.52 \u2013 7.45 (m, 2H), 6.91 (dd, J = 7.0, 1.5 Hz, 1H), "
                                 "2.47 (s, 3H).\n13C NMR (126 MHz, CDCl3) \u03b4 141.2, 21.1.\n"
                                 "19F NMR (471 MHz, CDCl3) \u03b4 -63.2 (s).")
        t_n.text_changed()
        assert t_n.tbl_sig.rowCount() == 6
        t_n.tbl_sig.item(1, 5).setText("Ar-H")                           # Zuordnung je Signal
        t_n.tbl_sig.item(3, 5).setText("Ar-*C*_{q}")
        ln_ = t_n.nmr_lines()
        assert ln_[0][0] == ("\u00b9H NMR (500 MHz, CDCl\u2083): \u03b4 [ppm] = 7.52 \u2013 7.45 (m, 2H), "
                             "6.91 (dd, J = 7.0, 1.5 Hz, 1H, Ar-H), 2.47 (s, 3H)"), ln_[0][0]
        assert "<i>J</i>" in ln_[0][1] and "Ar-<i>C</i><sub>q</sub>" in ln_[1][1] and "\u221263.2" in ln_[2][0]
        assert NMR.svg_peaks(sv_c) == (["141.20", "21.05"], "CDCl3"), NMR.svg_peaks(sv_c)
        w.nmr_state["current"] = [s_ for s_ in t_n.spectra() if s_["nuclei"] == ["13C"]][0]["uid"]
        t_n.peaks_from_svg()                                   # 13C-Zeile aus den Peaklabels (ersetzt die alte)
        assert "13C NMR (126 MHz, CDCl3) \u03b4 141.20, 21.05." in w.nmr_state["text"] and \
            w.nmr_state["text"].count("13C NMR") == 1, w.nmr_state["text"]
        acq_ = NMR.make_test_acqus(ndir / "exp" / "10", nuc="1H", te=300.2)
        (acq_ / "fid").write_bytes(b"\0" * 1000)                          # ganzer Ordner ins Projekt (seit 2.1)
        e_h = t_n.spectra()[0]
        assert t_n.set_bruker(e_h, str(acq_ / "acqus"))
        assert Path(e_h["bruker"]) == proj / "data" / "nmr" / "1H" / "exp_10" and (Path(e_h["bruker"]) / "fid").exists()
        assert abs(e_h["params"]["o1p"] - 6.175) < 1e-6 and e_h["params"]["d1"] == 1.0
        w.nmr_state["current"] = e_h["uid"]
        t_n.temp_from_bruker()
        assert e_h["temp"] == 300
        cap_q = NMR.caption_html(t_n.spectra()[3], t_n.blocks(), "TEST-1")
        assert cap_q == ("<sup>1</sup>H/<sup>13</sup>C HSQC (500/126 MHz, CDCl<sub>3</sub>, 298 K) of "
                         "<b>TEST-1</b>."), cap_q
        assert any(k_ == "bruker" and c_ is e_h for c_, k_, _ in PRJ._path_slots(w.get_state()))
        from rdkit import Chem as _Chem
        if hasattr(_Chem, "MolToCDXMLBlock"):                               # Struktur als Vektor ins Spektrum
            from rdkit.Chem import AllChem as _AC
            sm_ = _Chem.MolFromSmiles("Cc1ccc(C#C)nc1")             # erfundenes Testmolekül
            _AC.Compute2DCoords(sm_)
            (ndir / "struct.cdxml").write_text(_Chem.MolToCDXMLBlock(sm_), encoding="utf-8")
            w.nmr_state["structure"] = str(ndir / "struct.cdxml")
            fs_ = t_n.figure_svg(t_n.spectra()[0])
            assert fs_ != t_n.spectra()[0]["file"] and 'id="structure_0"' in Path(fs_).read_text(encoding="utf-8")
            t_n.spectra()[1]["struct_on"] = False
            assert t_n.figure_svg(t_n.spectra()[1]) == t_n.spectra()[1]["file"]
            e0_ = t_n.spectra()[0]                                         # gezogen: Lage im Dateinamen
            e0_.update(struct_pos="manual", struct_xy=[0.5, 0.3])
            assert Path(t_n.figure_svg(e0_)).stem.endswith("_500_300_200"), t_n.figure_svg(e0_)
            w.nmr_state["current"] = e0_["uid"]                            # Mausrad über der Struktur: Größe
            t_n.sync_widgets()
            t_n.update_preview()
            if t_n._drag_rect is not None:
                from PySide6.QtCore import QPoint, QPointF
                from PySide6.QtGui import QWheelEvent
                x_, y_, wr_, hr_ = t_n._drag_rect
                pw_, ph_ = t_n._pm_size
                pos_ = QPointF((t_n.lbl_prev.width() - pw_) / 2 + (x_ + wr_ / 2) * pw_,
                               (t_n.lbl_prev.height() - ph_) / 2 + (y_ + hr_ / 2) * ph_)
                ev_ = QWheelEvent(pos_, pos_, QPoint(), QPoint(0, 240), Qt.NoButton, Qt.NoModifier,
                                  Qt.NoScrollPhase, False)
                assert t_n.eventFilter(t_n.lbl_prev, ev_) and w.nmr_state["struct_w"] == 22   # gilt für alle
            e0_.update(struct_pos="auto", struct_xy=None)
            w.nmr_state["struct_w"] = 20.0
        m_ = NMR.occupancy(sv_h)                                           # freie Stelle: verdeckt nichts
        fx_, fy_ = NMR.auto_place(m_, 0.2, 0.5)
        H_, W_ = m_.shape
        x0_, y0_ = int(fx_ * W_), int(fy_ * H_)
        assert m_.any() and not m_[y0_:y0_ + int(0.1 * W_), x0_:x0_ + int(0.2 * W_)].any(), (fx_, fy_)
        nd_ = t_n.export_dialog(out / "nmr" / "nmr_test.docx")
        with _zf.ZipFile(nd_) as z_:
            xml_ = z_.read("word/document.xml").decode("utf-8")
            svgs_ = [n_ for n_ in z_.namelist() if n_.endswith(".svg")]
        assert len(svgs_) == 4 and xml_.count("<w:pageBreakBefore/>") == 1 and xml_.count("SEQ Figure") == 4, len(svgs_)
        assert "HSQC" in xml_ and ">Figure S<" in xml_ and "projekt_test" in xml_
        # Rohdaten: TopSpin-Datensatz (ganz, nach Art sortiert), automatische Zuordnung, Mnova-Dokument, EMF-Spektrum
        ds_ = ndir / "topspin" / "TS1"
        for no_, nuc_, pp_, n2_ in (("1", "1H", "zg30", None), ("2", "13C", "zgpg30", None),
                                    ("3", "13C", "dept135", None), ("4", "1H", "hsqcedetgpsisp2.3", "13C")):
            e_ = NMR.make_test_acqus(ds_ / no_, nuc=nuc_, nuc2=n2_, pulprog=pp_, te=299.0)
            (e_ / ("ser" if n2_ else "fid")).write_bytes(b"\1" * 64)
            (e_ / "pdata" / "1").mkdir(parents=True)
            (e_ / "pdata" / "1" / "procs").write_text("##$SI= 1024\n", encoding="latin-1")
        assert [p_.name for p_ in NMR.find_experiments(ds_.parent)] == ["1", "2", "3", "4"]
        assert [NMR.experiment_kind(NMR.acq_params(ds_ / n_)) for n_ in "1234"] == ["1H", "13C", "DEPT135", "HSQC"]
        w.nmr_state.update(raw=[])
        for s_ in t_n.spectra():
            s_.update(bruker=None, params=None)
        assert t_n.import_topspin(str(ds_)) == 4 and t_n.import_topspin(str(ds_)) == 0     # zweimal: nichts doppelt
        for kind_, no_ in (("1H", "1"), ("13C", "2"), ("DEPT135", "3"), ("HSQC", "4")):
            d_ = proj / "data" / "nmr" / kind_ / f"TS1_{no_}"
            assert (d_ / "acqus").exists() and (d_ / "pdata" / "1" / "procs").exists(), d_
        sp_ = {NMR.spectrum_kind(s_): s_ for s_ in t_n.spectra()}
        assert sp_["1H"]["bruker"] == str(proj / "data" / "nmr" / "1H" / "TS1_1") and sp_["1H"]["temp"] == 299
        assert sp_["HSQC"]["bruker"].endswith("TS1_4") and sp_["13C"]["bruker"].endswith("TS1_2")
        t_n.sync_widgets()
        assert t_n.tbl_raw.rowCount() == 4 and t_n.tbl_raw.item(3, 2).text() == "hsqcedetgpsisp2.3"
        assert t_n.right_tabs.tabText(1) == T("nmr_raw")
        (ndir / "TS1.mnova").write_bytes(b"MNOVA-TEST")
        mn_ = t_n.add_mnova_file(str(ndir / "TS1.mnova"))
        assert Path(mn_) == proj / "data" / "nmr" / "mnova" / "TS1.mnova" and t_n.tbl_raw.rowCount() == 5
        slots_ = {str(PRJ._get(c_, k_)) for c_, k_, _ in PRJ._path_slots(w.get_state())}
        assert mn_ in slots_ and sp_["HSQC"]["bruker"] in slots_
        def emf_spec(path_, texts_):                       # minimale EMF: Header (rclFrame 16 × 11 cm) + Texte
            hdr_ = bytearray(108)
            _st.pack_into("<II", hdr_, 0, 1, 108)
            _st.pack_into("<iiii", hdr_, 24, 0, 0, 16000, 11000)
            body_ = b""
            for t_ in texts_:
                b_ = t_.encode("utf-16-le")
                size_ = 76 + len(b_) + (-(76 + len(b_)) % 4)
                rec_ = bytearray(size_)
                _st.pack_into("<II", rec_, 0, 84, size_)
                _st.pack_into("<II", rec_, 44, len(t_), 76)
                rec_[76:76 + len(b_)] = b_
                body_ += bytes(rec_)
            Path(path_).write_bytes(bytes(hdr_) + body_ + _st.pack("<III", 14, 20, 0) + b"\0" * 8)
            return Path(path_)
        em_ = emf_spec(ndir / "S2_19F.emf", ["10", "\u03b4 19F (ppm)", "0"])
        assert NMR.emf_info(em_) == dict(nuclei=["19F"], is2d=False, size_cm=(16.0, 11.0))
        # Vorschau füllt die Fläche (Breite oder Höhe, ganzes Bild sichtbar), auch nach Größenänderung; die Größe der
        # Vorschau kommt nie vom Bild (früher: Mindestgröße = altes Pixmap, Bild nur nach der Breite)
        w.nmr_state["current"] = [s_ for s_ in t_n.spectra() if s_["kind"] == "1D" and s_["nuclei"] == ["1H"]][0]["uid"]
        w.show()
        w.main_tabs.setCurrentWidget(t_n)
        t_n.right_tabs.setCurrentIndex(0)
        t_n.sync_widgets()
        for size_ in ((1440, 900), (1440, 600), (1000, 700), (1440, 900)):
            w.resize(*size_)
            pump(0.6)                                           # Layout setzt die Breite verzögert
            aw_, ah_ = t_n.lbl_prev.width() - 10, t_n.lbl_prev.height() - 10
            pw_, ph_ = t_n._pm_size
            assert pw_ <= aw_ + 1 and ph_ <= ah_ + 1 and (pw_ >= aw_ - 2 or ph_ >= ah_ - 2), (size_, aw_, ah_, pw_, ph_)
        s_em_ = t_n.load_svg(str(em_))
        assert s_em_["nuclei"] == ["19F"] and Path(s_em_["file"]).suffix == ".emf"
        w.nmr_state["current"] = s_em_["uid"]
        t_n.update_preview()
        assert "EMF" in t_n.lbl_prev.text()
        nd2_ = t_n.export_dialog(out / "nmr" / "nmr_test_emf.docx")
        with _zf.ZipFile(nd2_) as z_:
            assert any(n_.endswith(".emf") for n_ in z_.namelist())
            assert z_.read([n_ for n_ in z_.namelist() if n_.endswith(".emf")][0]) == em_.read_bytes()
        w.nmr_state.update(spectra=[], current=None, text="", assign={}, structure=None, raw=[], mnova=[])
        t_n.sync_widgets()
        report.append("NMR (Mnova SVG/EMF, nucleus from axis/manual, order, Mnova text + assignment, TopSpin raw data "
                      "sorted by kind, Mnova file, Word, preview fits area): OK")

        # Dialog Verbindung: umbrechender Hinweis unter „Literatur“ in voller Höhe (früher letzte Zeile abgeschnitten)
        cdlg_ = MW.CompoundDialog(w, dict(w.compound or {}, known=True, lit="J. Org. Chem. 2020, 85, 1234"))
        cdlg_.show()
        pump(0.1)
        hint_lb_ = [lb_ for lb_ in cdlg_.findChildren(QLabel) if lb_.text() == T("compound_hint")][0]
        assert hint_lb_.height() >= hint_lb_.heightForWidth(hint_lb_.width()), \
            (hint_lb_.height(), hint_lb_.heightForWidth(hint_lb_.width()))
        cdlg_.reject()
        # Schema-Hinweis beschreibt den aktuellen Ablauf (CDXML + optionale SVG/PDF, .cdx, .emf) in beiden Sprachen
        from elucida.app.i18n import TEXTS as TX_
        for k_ in ("syn_scheme_hint", "guide_syn_files"):
            assert all(all(x_ in t_ for x_ in ("CDXML", "SVG/PDF", ".cdx", ".emf")) for t_ in TX_[k_]), k_
        report.append("compound dialog hint not clipped, scheme hint (CDXML/SVG/PDF/.cdx/.emf): OK")

        # Synthese: Mengenprüfung mit Rundung, Reagenzientabelle, Schema als bearbeitbares ChemDraw-Objekt (CDX)
        proc_ = ("In a flask <b>S1</b> (1.03 g, 5.00 mmol, 1.0 eq.) was dissolved in MeOH (10 mL) and "
                 "K<sub>2</sub>CO<sub>3</sub> (829 mg, 6.00 mmol, 1.2 eq.) was added to give <b>S2</b> "
                 "(525 mg, 4.95 mmol, 99%).")                     # erfundene Werte
        tab_ = {"S1": {"M": 206.36}, "K2CO3": {"M": SYN.formula_mw("K2CO3")}, "S2": {"M": 106.12}}
        res_ = SYN.check_amounts(SYN.parse_amounts(proc_), tab_)
        assert [l_ for l_, _ in res_] == ["ok", "ok", "ok", "ok", "ok"], res_        # 1.03 g -> 4.99 vs. 5.00 mmol: Rundung
        bad_ = SYN.check_amounts(SYN.parse_amounts(proc_.replace("6.00", "6.60").replace("99%", "95%")), tab_)
        assert [l_ for l_, _ in bad_].count("warn") == 3, bad_
        assert abs(SYN.formula_mw("K2CO3") - 138.204) < 0.01 and SYN.molar_mass("206.36") == 206.36
        t_s = w.syn_tab
        w.main_tabs.setCurrentWidget(t_s)
        w.syn_state["reagents"] = [dict(name="S1", M="C12H18OSi", m_mg="1032", role="ref"),
                                   dict(name="K2CO3", M="K2CO3", m_mg="829", role="reagent"),
                                   dict(name="S2", M="C7H6O", m_mg="525", role="product")]
        t_s.sync_widgets()
        assert t_s.tbl.item(1, len(SYN.COLS) + 1).text() == "1.20" and t_s.tbl.item(2, len(SYN.COLS) + 1).text() == "99%"
        r0_ = w.syn_state["reagents"][0]
        assert SYN.amount_text(r0_, dict(SYN.row_amounts(r0_), eq=1.0)) == "(1.032 g, 5.00 mmol, 1.0 eq.)"
        t_s.ed_proc.setHtml(proc_)
        t_s.proc_changed()
        assert "<b>S1</b>" in w.syn_state["proc"] and "<sub>2</sub>" in w.syn_state["proc"], w.syn_state["proc"]
        assert t_s.lbl_check.text().count("\u2713") == 5, t_s.lbl_check.text()
        sdir = out / "input" / "syn"
        sdir.mkdir(parents=True, exist_ok=True)
        if hasattr(_Chem, "MolToCDXMLBlock"):
            (sdir / "scheme.cdxml").write_text(_Chem.MolToCDXMLBlock(sm_), encoding="utf-8")
            from rdkit.Chem import Draw as _Draw
            _Draw.MolToFile(sm_, str(sdir / "scheme.svg"))                   # steht für den ChemDraw-SVG-Export
            (sdir / "scheme.cdx").write_bytes(b"VjCD0100\x04\x03\x02\x01" + b"\0" * 200)   # Platzhalter-CDX
            t_s.set_scheme(str(sdir / "scheme.cdxml"))
            sp_ = t_s.scheme_parts()
            assert sp_["picture"].suffix == ".svg" and sp_["cdx"] is not None and sp_["emf"] is None
            assert Path(w.syn_state["scheme"]).parent == proj / "structures"
            sb_ = t_s.si_blocks(out / "syn")
            assert sb_[0] == ("synthesis", "TEST-1" if w.compound.get("name") == "TEST-1" else t_s.compound_name())
            assert [b_[0] for b_ in sb_] == ["synthesis", "ole", "char"], [b_[0] for b_ in sb_]
            sd_ = t_s.export_dialog()
            import olefile
            with _zf.ZipFile(sd_) as z_:
                xml_ = z_.read("word/document.xml").decode("utf-8")
                ole_ = olefile.OleFileIO(z_.read("word/embeddings/oleObject1.bin"))
            assert 'ProgID="ChemDraw.Document.6.0"' in xml_ and ole_.openstream("CONTENTS").read().startswith(b"VjCD0100")
            assert any(k_ == "scheme" and c_ is w.syn_state for c_, k_, _ in PRJ._path_slots(w.get_state()))
            (Path(w.syn_state["scheme"]).with_suffix(".cdx")).unlink()           # ohne CDX: Vektorbild
            assert [b_[0] for b_ in t_s.si_blocks(out / "syn")][1] == "image"
            # Schema ohne ChemDraw-Export: selbst gezeichnet (Pfeil, Bedingungen), Reagenzien aus den Pfeiltexten
            import xml.etree.ElementTree as _ET
            from elucida.formats import cdxml as CDXD
            fr_ = _ET.fromstring(_Chem.MolToCDXMLBlock(sm_)).find(".//fragment")
            fr2_ = _ET.fromstring(_Chem.MolToCDXMLBlock(sm_)).find(".//fragment")
            for k_, n_ in enumerate(fr2_.iter()):                              # Produkt: 150 pt weiter rechts
                if n_.get("id"):
                    n_.set("id", str(90000 + k_))
                if n_.tag == "n":
                    x_, y_ = map(float, n_.get("p").split())
                    n_.set("p", f"{x_ + 150} {y_}")
            ids_ = {o_.get("id"): str(90000 + k_) for k_, o_ in enumerate(
                _ET.fromstring(_Chem.MolToCDXMLBlock(sm_)).find(".//fragment").iter()) if o_.get("id")}
            for b_ in fr2_.iter("b"):
                b_.set("B", ids_[b_.get("B")])
                b_.set("E", ids_[b_.get("E")])
            ys_ = [float(n_.get("p").split()[1]) for n_ in fr_.iter("n")]
            xs_ = [float(n_.get("p").split()[0]) for n_ in fr_.iter("n")]
            ym_, xr_ = sum(ys_) / len(ys_), max(xs_)
            sch_ = (f'<?xml version="1.0"?><CDXML LineWidth="0.6" BondSpacing="18" LabelSize="10"><page>'
                    + _ET.tostring(fr_, encoding="unicode") + _ET.tostring(fr2_, encoding="unicode")
                    + f'<t id="501" p="{xr_ + 30} {ym_ - 6}"><s face="96" size="10">K2CO3 (1.2 eq.)</s></t>'
                    + f'<t id="502" p="{xr_ + 30} {ym_ + 12}"><s size="9.7">MeOH, rt, 1 h</s></t>'
                    + f'<arrow id="503" Head3D="{xr_ + 120} {ym_} 0" Tail3D="{xr_ + 20} {ym_} 0"/>'
                    + f'<scheme><step ReactionStepReactants="{fr_.get("id")}" ReactionStepProducts="{fr2_.get("id")}" '
                    + 'ReactionStepObjectsAboveArrow="501" ReactionStepObjectsBelowArrow="502"/></scheme></page></CDXML>')
            (sdir / "rxn.cdxml").write_text(sch_, encoding="utf-8")
            svg_, notes_ = CDXD.render_svg(sdir / "rxn.cdxml")
            assert svg_.count(">N</text>") == 2, svg_.count(">N</text>")      # Labels auch ohne <t> (RDKit-CDXML)
            assert not notes_ and "<polygon" in svg_ and ">2</text>" in svg_ and "1.2\u00a0eq.)" in svg_, notes_
            nest_ = sch_.replace('<t id="501"', '<group id="601"><group id="602"><t id="501"').replace(
                '<arrow id="503"', '</group><arrow id="503"').replace('"/><scheme>', '"/></group><scheme>')
            assert nest_.count("<group") == 2 and nest_.count("</group>") == 2
            (sdir / "rxn_grp.cdxml").write_text(nest_, encoding="utf-8")     # Bedingungen + Pfeil gruppiert (Feedback)
            assert CDXD.render_svg(sdir / "rxn_grp.cdxml") == (svg_, notes_), "Pfeil/Texte in Gruppe fehlen"
            assert CDXD.scheme_reagents(sdir / "rxn.cdxml") == [("K2CO3", "K2CO3", "reagent"), ("MeOH", "CH4O", "other")]
            t_s.set_scheme(str(sdir / "rxn.cdxml"))
            assert t_s.scheme_asset()["drawing_from"] == "cdxml" and t_s.scheme_asset()["pdf"]
            w.syn_state["reagents"] = []
            t_s.rows_from_scheme()
            assert [(r_["name"], r_["role"]) for r_ in w.syn_state["reagents"]] == \
                [("", "ref"), ("K2CO3", "reagent"), ("MeOH", "other"), ("", "product")], w.syn_state["reagents"]
        b_all, _ = w.si_blocks(["char"])
        assert b_all[0][0] == "synthesis" and any(b_[0] == "char" and "S1" in b_[1] for b_ in b_all)
        w.syn_state.update(name="", scheme=None, proc="", reagents=[])
        t_s.sync_widgets()
        report.append("Synthesis (amount check with rounding, reagent table, ChemDraw object, SI): OK")
        w.dft_state.update(entries=[], current=None)
        t_d.sync_widgets()
        t_d.render()

        # SI-Zeilen: Elementaranalyse, UV-Vis, SI-Abschnitt als Word-Datei
        comp_ = SI.ea_composition("C18H20N2")
        assert [round(comp_[e_], 2) for e_ in "CHN"] == [81.78, 7.63, 10.60], comp_
        assert SI.ea_formula_label("C18H20N2", "0.5 CH2Cl2") == "C18H20N2·0.5CH2Cl2"
        assert abs(SI.ea_composition("C18H20N2", "0.5 CH2Cl2")["C"] - 72.42) < 0.01
        pl_, ht_, nt_ = SI.ea_line("C18H20N2", "", {"C": "81.52", "H": "7.70", "N": "10.10"})
        assert pl_ == "Elemental analysis (%) calc. C18H20N2: C 81.78, H 7.63, N 10.60; found C 81.52, " \
                      "H 7.70, N 10.10", pl_
        assert [e_ for e_, _ in nt_] == ["N"] and "<sub>18</sub>" in ht_       # N weicht um 0.50 % ab
        cfg_uv = dict(core.DEFAULTS, conc_unit="mM")
        res_uv = [{"lam": 412.0, "fit": {"n": 5}, "eps": 12345.6, "eps_err": 120.0},
                  {"lam": 298.0, "fit": {"n": 5}, "eps": 24100.2, "eps_err": 0.4},
                  {"lam": 350.0, "fit": {"n": 5}, "eps": 5000.0, "eps_err": 50.0, "hidden": True},
                  {"lam": 500.0, "fit": None}]
        uvp_, uvh_ = SI.uvvis_line(res_uv, cfg_uv, "THF")
        assert uvp_ == "UV-Vis (THF): λmax: 298 nm (ε = 24100 cm-1 M-1), 412 nm (ε = 12346 cm-1 M-1)", uvp_
        assert "<b>λ<sub>max</sub></b>" in uvh_ and "cm<sup>-1</sup>" in uvh_, uvh_
        assert SI.join_lines([("m.p. 67 °C", "m.p. 67 °C"), (uvp_, uvh_)])[0].endswith("M-1).")
        w.compound = {"name": "TEST-1", "formula": "C18H20N2", "mp": "67 °C"}
        w.irms_state.update(ea_found={"C": "81.70", "H": "7.60", "N": "10.55"}, ea_solvate="")
        w.irms_tab.sync_widgets()
        blocks_, missing_ = w.si_blocks()
        para_ = [b_[1] for b_ in blocks_ if b_[0] == "char"][0]
        assert blocks_[0] == ("synthesis", "TEST-1"), blocks_[0]
        assert para_.startswith("<b>m.p.</b> 67 °C; ") and "<b>Elemental analysis</b>" in para_, para_
        assert "IR" in missing_ and "m.p." not in missing_, missing_
        w.export_si()
        docx_ = Path(w.project_dir) / "SI" / "TEST-1_SI.docx"
        with _zf.ZipFile(docx_) as z_:
            xml_ = z_.read("word/document.xml").decode("utf-8")
            assert 'w:val="15"' in z_.read("word/settings.xml").decode("utf-8")   # kein Kompatibilitätsmodus
        assert "Elemental analysis" in xml_ and 'w:val="subscript"' in xml_ and "TEST-1" in xml_
        assert "67\u00a0°C" in xml_ and "C\u00a081." in xml_                  # geschützte Leerzeichen
        assert SI.nbsp("298 nm (ε = 5 cm-1 M-1)") == "298\u00a0nm (ε\u00a0=\u00a05\u00a0cm-1\u00a0M-1)"
        assert "Courier New" not in xml_                                   # ohne DFT-Rechnung kein DFT-Teil
        # Teile wählbar: DFT zweispaltig hinter der Charakterisierung, ohne Charakterisierung nur Überschrift
        w.dft_tab.load_xyz(str(orf_))
        assert {k_: v_ for k_, v_ in w.si_available().items() if k_ not in ("figs", "methods")} == \
            {"char": True, "nmr": False, "xray": False, "dft": True}
        b_dft, _ = w.si_blocks(["dft"])
        assert b_dft[0] == ("heading", "TEST-1") and ("columns", 2) in b_dft and not any(b_[0] == "char" for b_ in b_dft)
        w.export_si(["char", "dft"])
        with _zf.ZipFile(docx_) as z_:
            xml_ = z_.read("word/document.xml").decode("utf-8")
        assert "Elemental analysis" in xml_ and "Courier New" in xml_ and xml_.count("<w:sectPr>") == 2
        assert xml_.index("Elemental analysis") < xml_.index('w:cols w:num="2"')   # einspaltig, dann zweispaltig
        w.dft_state.update(entries=[], current=None)
        w.dft_tab.sync_widgets()
        w.irms_state.update(ea_found={}, ea_solvate="")
        w.compound = {}
        report.append("SI lines (EA, UV-Vis) + SI section as Word (parts selectable): OK")

        # Release Notes aus den Changelogs (nur im Quellcode; das Skript liegt nicht in der gebauten App)
        rn_path = ROOT / "tools" / "release_notes.py"
        if rn_path.exists() and not getattr(sys, "frozen", False):
            import importlib.util
            spec_ = importlib.util.spec_from_file_location("release_notes", rn_path)
            RN = importlib.util.module_from_spec(spec_)
            spec_.loader.exec_module(RN)
            en_ = (ROOT / "CHANGELOG.md").read_text(encoding="utf-8")
            de_ = (ROOT / "CHANGELOG.de.md").read_text(encoding="utf-8")
            txt_ = RN.notes(APP_VERSION, en_, de_)
            assert txt_.startswith("## English") and "## Deutsch" in txt_ and f"{APP_NAME}-windows-x64.zip" in txt_
            assert f"## {APP_VERSION}" not in txt_ and "\n## 1." not in txt_          # nur dieser Abschnitt
            try:
                RN.notes("0.0.1", en_, de_)
                raise AssertionError("fehlende Version nicht erkannt")
            except ValueError:
                pass
            report.append("release notes from changelogs: OK")

        # Update-Prüfung (ohne Netz: Versionsvergleich, API-Antwort, Meldungstext)
        assert UPD.parse_version("v1.10.0") > UPD.parse_version("1.9.9") and UPD.parse_version("x") == ()
        assert UPD.is_newer("1.4.1", "1.4.0") and not UPD.is_newer("1.4", "1.4.0") and not UPD.is_newer("", "1.0")
        lbl = UPD.platform_label()
        fake = {"tag_name": "v99.0.0", "html_url": "https://example.invalid/r",
                "assets": [{"name": f"Elucida-{x}.zip"} for x in ("windows-x64", "macos-arm64", "macos-x64")]
                + [{"name": "Elucida-linux-x64.tar.gz"}]}
        info = UPD.parse_release(fake)
        assert info["version"] == "99.0.0" and lbl in info["asset"], info
        assert "99.0.0" in w.update_text(info, manual=False)
        assert w.update_text(dict(info, version=APP_VERSION), manual=False) is None
        assert APP_VERSION in w.update_text(dict(info, version=APP_VERSION), manual=True)
        report.append("update check: OK")

        # Projekt automatisch gespeichert -> in neuem Fenster öffnen (Dateien als Kopie, Pfade relativ)
        w.ov_state["entries"][0]["file"] = w.import_file(w.ov_state["entries"][0]["file"], PRJ.data_sub("overlay"))
        w.push_undo()
        pump(1.5)                                      # automatisches Speichern abwarten
        w.maybe_save()
        txt = (proj / PRJ.PROJECT_FILE).read_text(encoding="utf-8")
        assert str(out / "input") not in txt.replace(str(proj), ""), "absolute Pfade im Projekt"
        assert (proj / "data" / "uvvis" / csv.name).exists() and (proj / "structures" / cdxml.name).exists()
        w2 = MainWindow()
        assert w2.open_project(proj, ask=False)
        pump()
        assert set(w2.groups) == {"TEST-A", "TEST-B"} and w2.series["TEST-A"]["mw"] == 95.1, w2.groups
        assert len(w2.ov_state.get("entries", [])) == 2 and w2.fl_state.get("em_file"), "Tabs nicht geladen"
        report.append("project save/load: OK")

        # Projektordner (2.1): Exporte nach Art sortiert, Name der Verbindung, Zwischendateien in .elucida
        w2.compound["name"] = "TEST-1"
        ov_out_ = w2.ov_tab.export_dialog()
        assert proj / "figures" / "TEST-1_Overlay.pdf" in ov_out_, ov_out_
        assert proj / "tables" / "TEST-1_Overlay_data.csv" in ov_out_, ov_out_
        assert not (proj / PRJ.WORK / "export" / "overlay").exists()      # Zwischenordner wieder weg
        assert (proj / PRJ.WORK / "project.elucida.bak").exists() and not (proj / "project.elucida.bak").exists()
        assert (proj / "project.elucida").exists() and not (proj / "project.uvvis").exists()
        assert not (proj / "exports").exists() and not (proj / "images").exists()
        assert [PRJ.export_kind(n_) for n_ in ("a.svg", "a_SI.txt", "a.docx", "a_methods.txt", "a.csv", "a.xyz")] \
            == ["figures", "SI", "SI", "SI", "tables", "tables"]
        assert PRJ.data_sub("irms") == "data/ir" and PRJ.export_stem("S/2", "xray") == "S_2_X-ray"
        old_ = out / "old_layout"                          # Projekt im alten Aufbau (bis 2.0): Pfade gelten weiter
        (old_ / "data").mkdir(parents=True, exist_ok=True)
        shutil.copy2(csv, old_ / "data" / csv.name)
        (old_ / "project.uvvis").write_text(f"eps:\n  files: [data/{csv.name}]\nformat_version: 1\n", encoding="utf-8")
        assert PRJ.is_project(old_) == old_ and PRJ.is_project(old_ / "project.uvvis") == old_   # alter Name wird erkannt
        st_old_ = PRJ.load(old_)
        assert Path(st_old_["eps"]["files"][0]) == (old_ / "data" / csv.name).resolve(), st_old_
        PRJ.backup(old_)
        assert (old_ / PRJ.WORK / "project.uvvis.bak").exists()
        PRJ.save(old_, st_old_)                            # Speichern: neuer Name, alter als Sicherung
        assert "data/" + csv.name in (old_ / "project.elucida").read_text(encoding="utf-8")
        assert not (old_ / "project.uvvis").exists() and (old_ / PRJ.WORK / "project.uvvis.old").exists()
        assert PRJ.load(old_)["eps"]["files"] == st_old_["eps"]["files"]
        assert not (old_ / "data" / "uvvis").exists()      # alte Dateien bleiben, wo sie sind
        report.append("project folder layout (SI/figures/tables/data/structures/.elucida, old projects): OK")

        # elucida.json für Characterization: atomar, UTF-8 ohne BOM, Pfade relativ, Werte als Text
        from elucida.exchange import characterization as EXCH
        w2.compound.update(mp="67 °C", structure=str(proj / "structures" / cdxml.name))
        jx_ = w2.write_exchange()
        assert jx_ == proj / "elucida.json" and not (proj / "elucida.json.tmp").exists(), jx_
        raw_ = jx_.read_bytes()
        assert not raw_.startswith(b"\xef\xbb\xbf")
        ex_ = json.loads(raw_.decode("utf-8"))
        assert ex_["format"] == "elucida-exchange" and ex_["schema_version"] == 1 and ex_["app"]["version"] == APP_VERSION
        assert re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", ex_["written_at"]), ex_["written_at"]
        assert ex_["compound"]["name"] == "TEST-1" and ex_["compound"]["structure"] == "structures/" + cdxml.name
        assert ex_["analytics"]["mp"] == {"text": "67 °C"} and [s_["name"] for s_ in ex_["analytics"]["uvvis"]] \
            == ["TEST-A", "TEST-B"], ex_["analytics"]
        assert all(not Path(f_).is_absolute() and "\\" not in f_ for f_ in ex_["analytics"]["overlay"]["files"])
        kinds_ = [l_["kind"] for l_ in ex_["si"]["lines"]]
        assert [l_["series"] for l_ in ex_["si"]["lines"] if l_["kind"] == "uvvis"] == ["TEST-A", "TEST-B"]
        assert kinds_[0] == "mp" and "uvvis" in kinds_ and ex_["si"]["paragraph_html"].startswith("<b>m.p.</b> 67"), kinds_
        fig_ = next(f_ for f_ in ex_["figures"] if f_["tab"] == "overlay")
        assert fig_["pdf"] == "figures/TEST-1_Overlay.pdf" and fig_.get("caption_html"), fig_
        assert any(f_["tab"] == "eps" for f_ in ex_["figures"]), ex_["figures"]
        t0_ = time.perf_counter()
        EXCH.build(w2)
        report.append(f"exchange file elucida.json (schema 1, {time.perf_counter() - t0_:.2f} s): OK")

        # Datenpaket (Zip) je Verbindung wie die Datensätze der Gruppe; Laborkürzel ersetzt, Binärdateien unverändert
        from elucida.exchange import datapack as DP
        import zipfile as _zfp
        dp_ = out / "datapack" / "proj"
        ex_ = NMR.make_test_acqus(dp_ / "data" / "nmr" / "1H" / "RE-3-150_10", nuc="1H")
        with open(ex_ / "acqus", "a", encoding="latin-1", newline="") as fh_:   # Windows: \n bleibt \n
            fh_.write("$$ 2024-03-12 14:35:12.123 +0100  user@nmr\n$$ /opt/topspin/data/RE-3-150/nmr/RE-3-150/10/acqus\n"
                      "##$DATE= 1710251712\n")
        fid_ = bytes(range(256)) * 4
        (ex_ / "fid").write_bytes(fid_)
        (ex_ / "pdata" / "1").mkdir(parents=True)
        (ex_ / "pdata" / "1" / "title").write_text("RE-3-150 in CDCl3 12.03.2024\n", encoding="latin-1", newline="")
        hs_ = NMR.make_test_acqus(dp_ / "data" / "nmr" / "HSQC" / "RE-3-150_12", nuc="1H", nuc2="13C",
                                  pulprog="hsqcedetgpsisp2.3")
        hm_ = NMR.make_test_acqus(dp_ / "data" / "nmr" / "HMBC" / "RE-3-150_13", nuc="1H", nuc2="13C",
                                  pulprog="hmbcgplpndqf")
        mn_ = dp_ / "data" / "nmr" / "mnova" / "RE-3-150.mnova"
        mn_.parent.mkdir(parents=True)
        mn_.write_bytes(b"\0MNOVA\0RE-3-150\0")                        # binär: bleibt, wird gemeldet
        ir0_ = dp_ / "data" / "ir" / "RE-3-150_IR.0"
        ir0_.parent.mkdir(parents=True)
        ir0_bytes_ = b"\x0a\x0a\x00OPUS\x00SNM\x00RE-3-150\x00DAT\x0012/03/2024\x00\xff\x01"
        ir0_.write_bytes(ir0_bytes_)
        msr_ = dp_ / "data" / "hrms" / "MS-Bericht.docx"
        msr_.parent.mkdir(parents=True)
        with _zfp.ZipFile(msr_, "w") as z_:                               # Bericht: Kürzel in Runs geteilt, Datum
            z_.writestr("word/document.xml", '<w:document><w:body><w:p><w:r><w:t>sample name: RE-3-</w:t></w:r>'
                        '<w:r><w:t>150</w:t></w:r></w:p><w:p><w:r><w:t>date: 12.03.2024</w:t></w:r></w:p>'
                        '<w:p><w:r><w:t>exact mass: 264.1626</w:t></w:r></w:p></w:body></w:document>')
            z_.writestr("docProps/core.xml", "<cp:coreProperties><dc:creator>Jane Doe</dc:creator>"
                        "<dcterms:created>2024-03-12T10:00:00Z</dcterms:created><dc:title>MS</dc:title>"
                        "</cp:coreProperties>")
            z_.writestr("word/media/image1.emf", b"\x01\0\0\0" + "RE-3-150 12.03.2024".encode("utf-16-le"))
        cvf_ = dp_ / "data" / "cv" / "RE-3-150-thf-200mVs.DTA"
        cvf_.parent.mkdir(parents=True)
        cvf_.write_text("EXPLAIN\nTITLE\tLABEL\tRE-3-150 CV\tTest\n", encoding="latin-1", newline="")
        epr_ = dp_ / "data" / "epr" / "RE-3-150_epr.txt"
        epr_.parent.mkdir(parents=True)
        epr_.write_text("B [G]\tRE-3-150\tmeasured 2024-03-12 09:15\n3380\t0.1\n", encoding="utf-8", newline="")
        cif_ = dp_ / "data" / "xray" / "s.cif"
        cif_.parent.mkdir(parents=True)
        cif_.write_text("data_x\n", encoding="utf-8")
        st_dp_ = {"nmr": {"raw": [dict(path=str(ex_), kind="1H"), dict(path=str(hs_), kind="HSQC"),
                                  dict(path=str(hm_), kind="HMBC")], "mnova": [str(mn_)]},
                  "irms": {"ir_file": str(ir0_)}, "hrms": {"ms_file": str(msr_)},
                  "cv": {"files": [dict(file=str(cvf_), group="THF")]}, "epr": {"entries": [dict(file=str(epr_))]},
                  "xray": {"entries": [dict(file=str(cif_))]}, "syn": {"scheme": str(cif_)}}
        its_ = DP.collect(st_dp_, dp_)
        arcs_ = {i_.arc: i_.on for i_ in its_}
        assert arcs_ == {"NMR/1H/10": True, "NMR/HSQC + HMBC/12": True, "NMR/HSQC + HMBC/13": True,
                         "NMR/RE-3-150.mnova": True, "IR/RE-3-150_IR.0": True, "MS/MS-Bericht.docx": True,
                         "CV/THF/RE-3-150-thf-200mVs.DTA": True, "EPR/RE-3-150_epr.txt": True}, arcs_   # ohne X-Ray
        assert DP.find_codes(its_) == ["RE-3-150"]
        import zipfile as _zfp
        zp_ = out / "datapack" / "6m.zip"
        rep_ = DP.write_zip(its_, zp_, "6m", {"RE-3-150": "6m"}, {"RE-3-150": "Compound 6m"})
        with _zfp.ZipFile(zp_) as z_:
            names_ = set(z_.namelist())
            assert {"6m/NMR/1H/10/acqus", "6m/NMR/1H/10/fid", "6m/NMR/1H/10/pdata/1/title", "6m/NMR/6m.mnova",
                    "6m/IR/6m_IR.0", "6m/MS/MS-Bericht.docx", "6m/CV/THF/6m-thf-200mVs.DTA",
                    "6m/EPR/6m_epr.txt", "6m/NMR/HSQC + HMBC/12/acqu2s"} <= names_, sorted(names_)
            assert not any("RE-3-150" in n_ for n_ in names_)
            assert z_.read("6m/NMR/1H/10/fid") == fid_                        # Rohdaten bytegleich
            assert b"Compound 6m in CDCl3" in z_.read("6m/NMR/1H/10/pdata/1/title")
            acq_txt_ = z_.read("6m/NMR/1H/10/acqus").decode("latin-1")
            assert "RE-3-150" not in acq_txt_ and "/Compound 6m/10/acqus" in acq_txt_ and "##$PULPROG= <zg30>" in acq_txt_
            assert "2024" not in acq_txt_ and "##$DATE= 0" in acq_txt_ and "user@nmr" in acq_txt_     # nur Datum weg
            assert z_.read("6m/NMR/1H/10/pdata/1/title") == b"Compound 6m in CDCl3 \n"
            assert "2024" not in z_.read("6m/EPR/6m_epr.txt").decode("utf-8")
            ir_z_ = z_.read("6m/IR/6m_IR.0")                                # OPUS: gleiche Länge, Rest bytegleich
            assert len(ir_z_) == len(ir0_bytes_) and b"RE-3-150" not in ir_z_ and b"2024" not in ir_z_
            assert ir_z_.startswith(b"\x0a\x0a\0OPUS\0SNM\x006m      \0") and ir_z_.endswith(b"\0\xff\x01")
            ms_txt_ = DP.docx_text(z_.read("6m/MS/MS-Bericht.docx"))
            assert "sample name: Compound 6m" in ms_txt_ and "2024" not in ms_txt_ and "264.1626" in ms_txt_
            with _zfp.ZipFile(__import__("io").BytesIO(z_.read("6m/MS/MS-Bericht.docx"))) as zm_:
                core_ = zm_.read("docProps/core.xml").decode("utf-8")
                assert "Jane Doe" not in core_ and "dcterms:created" not in core_ and "<dc:title>MS</dc:title>" in core_
            assert "Compound 6m CV" in z_.read("6m/CV/THF/6m-thf-200mVs.DTA").decode("latin-1")
            assert z_.read("6m/NMR/6m.mnova") == mn_.read_bytes()
        assert rep_["kept"] == ["6m/NMR/6m.mnova"], rep_["kept"]
        assert not (out / "datapack" / "6m.zip.tmp").exists()
        rep2_ = DP.write_zip(its_, out / "datapack" / "6m_raw.zip", "6m", {"RE-3-150": "6m"}, None)   # Inhalt unverändert
        with _zfp.ZipFile(out / "datapack" / "6m_raw.zip") as z_:
            assert b"RE-3-150 in CDCl3" in z_.read("6m/NMR/1H/10/pdata/1/title") and not rep2_["text_replaced"]
        from elucida.app.datapack_dialog import DatapackDialog          # Dialog mit dem Projekt des Selbsttests
        dd_ = DatapackDialog(w2)
        assert dd_.ed_name.text() == w2.export_name() and dd_.tree.topLevelItemCount() >= 1
        rep3_ = dd_.export(out / "datapack" / "dialog.zip")
        assert rep3_ and Path(rep3_["path"]).exists() and rep3_["files"] > 0
        report.append("data package zip (layout like the group's datasets, lab codes replaced, binaries unchanged): OK")

        # Undo/Redo
        before = w2.glob["cutoff"]
        w2.sp_cut.setValue(before + 0.5)
        pump(0.8)
        assert w2.glob["cutoff"] == before + 0.5
        w2.undo()
        pump()
        assert abs(w2.glob["cutoff"] - before) < 1e-9, w2.glob["cutoff"]
        w2.redo()
        pump()
        assert abs(w2.glob["cutoff"] - (before + 0.5)) < 1e-9, w2.glob["cutoff"]
        # Rückgängig baut nur, was sich geändert hat: ε ohne Tab-Neuaufbau, ein Tab -> nur dieser Tab neu
        tabs_ = lambda: {k_: getattr(w2, w2._tab_spec(k_)[0]) for k_ in w2._TAB_BUILD_ORDER}   # noqa: E731
        pump(0.4)
        w2._init_undo()                                    # frischer Verlauf (vorher Verbindung ohne Eintrag gesetzt)
        w2.sp_cut.setValue(before + 1.0)
        pump(0.8)
        objs_ = tabs_()
        w2.undo()
        pump(0.4)
        diff_ = [k_ for k_, v_ in tabs_().items() if v_ is not objs_[k_]]
        assert not diff_ and abs(w2.glob["cutoff"] - (before + 0.5)) < 1e-9, (
            diff_, w2.glob["cutoff"], before, len(w2._undo))
        assert abs(w2.sp_cut.value() - (before + 0.5)) < 1e-9          # Bedienelement nachgezogen, nicht neu
        cap_ = w2.cv_state.get("cap_temp", "rt")
        w2.cv_state["cap_temp"] = "-30 \u00b0C"
        w2.push_undo()
        pump(0.4)
        objs_ = tabs_()
        w2.undo()
        pump(0.4)
        now_ = tabs_()
        assert w2.cv_state["cap_temp"] == cap_ and now_["cv"] is not objs_["cv"], w2.cv_state.get("cap_temp")
        assert all(now_[k_] is objs_[k_] for k_ in now_ if k_ != "cv")                  # übrige Tabs unverändert
        assert w2.main_tabs.widget(w2._tab_keys.index("cv")) is w2.cv_tab and w2.cv_tab.st is w2.cv_state
        assert w2.cv_tab.ed_temp.text() == cap_ if hasattr(w2.cv_tab, "ed_temp") else True
        w2.redo()
        pump(0.4)
        assert w2.cv_state["cap_temp"] == "-30 \u00b0C"
        w2.undo()
        pump(0.4)
        # Öffnen: Fenster bleibt, nur der sichtbare Tab wird gebaut; übrige beim ersten Zugriff/Anzeigen
        w3_ = MainWindow()
        w3_.resize(1200, 800)
        w3_.show()
        assert w3_.open_project(proj, ask=False)
        pump(0.4)
        built_ = [k_ for k_ in w3_._TAB_BUILD_ORDER if w3_._built(k_)]
        assert len(built_) <= 1, built_
        key_ = next(k_ for k_ in ("cv", "sec", "epr") if not w3_._built(k_))
        tab3_ = getattr(w3_, w3_._tab_spec(key_)[0])                       # Zugriff baut ihn
        assert w3_._built(key_) and w3_.main_tabs.widget(w3_._tab_keys.index(key_)) is tab3_
        w3_.main_tabs.setCurrentIndex(w3_._tab_keys.index("xray"))         # Anzeigen baut ihn
        pump(0.2)
        assert w3_._built("xray") and w3_.main_tabs.currentWidget() is w3_.xray_tab
        w3_.close()
        report.append("undo/redo (only changed tabs rebuilt), lazy tabs on open: OK")
        # Startbildschirm/Tabs je nach Projekt; Beispielprojekt (künstliche Daten) füllt die Tabs
        assert w.center.currentIndex() == 1
        d_demo = w.open_demo(out / "demo")
        assert d_demo and Path(w.project_dir) == out / "demo" and w.center.currentIndex() == 1, (
            d_demo, w.project_dir, w.center.currentIndex(), w.log.toPlainText()[-3000:])
        assert w.groups and w.cv_state["files"] and w.xray_state["entries"] and w.dft_state["entries"]
        assert w.syn_state["proc"] and w.nmr_tab.spectra()
        st_file_ = PRJ.load(out / "demo")                  # gespeichert, nicht nur im Fenster (war leer: _restoring)
        assert st_file_.get("cv", {}).get("files") and st_file_.get("eps", {}).get("files"), sorted(st_file_)
        from elucida.app import demo as DEMO_                 # Beispiel-CDXML in ChemDraw auf der ersten Seite
        import xml.etree.ElementTree as ET_c
        cdx_demo_ = [p_ for p_ in (out / "demo").rglob("*.cdxml") if ET_c.parse(p_).getroot().find(".//n") is not None]
        assert any(p_.name == "rxn.cdxml" for p_ in cdx_demo_), cdx_demo_
        for p_ in cdx_demo_:
            x0_, y0_, x1_, y1_ = DEMO_.cdxml_extent(ET_c.parse(p_).getroot())
            m_ = DEMO_.PAGE_MARGIN                            # innerhalb der Druckränder (ACS-1996-Größe)
            assert m_ <= x0_ and x1_ <= DEMO_.PAGE_W - m_ and m_ <= y0_ and y1_ <= DEMO_.PAGE_H - m_, \
                (p_.name, x0_, y0_, x1_, y1_)
            if p_.name == "rxn.cdxml":                            # ein Pfeil, K2CO3 mit tiefgestellten Ziffern
                r_ = ET_c.parse(p_).getroot()
                assert len(r_.findall(".//arrow")) == 1 and not r_.findall(".//graphic")
                assert [s_.text for s_ in r_.iter("s") if s_.get("face") == "32"] == ["2", "3"]
                assert r_.get("BondLength") == "14.40" and r_.find(".//fonttable") is not None
        report.append("Example CDXML on the ChemDraw page (ACS 1996 size, one arrow): OK")
        from elucida.app import welcome as WEL                 # Startbildschirm: Projektliste mit Struktur und Analytik
        w.compound["note"] = "Radikalkation (Test)"           # Notiz zur Verbindung: Titel, Startbildschirm, elucida.json
        w.update_title()
        pump(0.4)                                              # Ladevorgang des Beispielprojekts abgeschlossen
        w.save_project()
        w.write_exchange()
        assert "(Radikalkation (Test))" in w.windowTitle(), w.windowTitle()
        import json as _json
        assert _json.loads((out / "demo" / "elucida.json").read_text(encoding="utf-8"))["compound"]["note"] == \
            "Radikalkation (Test)"
        w.add_recent(out / "demo")
        w.fill_recent_list()
        inf_ = WEL.project_info(out / "demo")
        assert inf_ and {"nmr", "cv", "uvvis", "xray", "dft"} <= set(inf_["badges"]), inf_
        assert WEL.thumb_path(out / "demo").exists() and inf_["thumb"], "Strukturbild fehlt"
        assert any(r_.info["path"] == str((out / "demo")) for r_ in w.recent_rows)
        assert inf_["note"] == "Radikalkation (Test)"
        w._filter_recent("radikalkation")
        assert any(r_.isVisibleTo(r_.parentWidget()) for r_ in w.recent_rows)
        w._filter_recent("gibtsnicht")
        assert not any(r_.isVisibleTo(r_.parentWidget()) for r_ in w.recent_rows)
        w._filter_recent("")
        import datetime as _dt
        now_ = _dt.datetime(2026, 10, 24, 15, 0)
        assert WEL.relative_time(_dt.datetime(2026, 10, 24, 9, 5), now_) == T("time_today", t="09:05")
        assert WEL.relative_time(_dt.datetime(2026, 10, 21, 9, 5), now_) == T("time_days", n=3)
        w.remove_recent(out / "demo")
        assert str(out / "demo") not in w.recent()
        report.append("welcome page + example project: OK")
        # macOS-Titelleiste: offscreen nie AppKit (winId kein NSView); Knöpfe der Leiste bekommen ihre Klicks
        from elucida.app import macos as MAC
        from PySide6.QtCore import QPointF
        from PySide6.QtGui import QMouseEvent
        assert not MAC.enabled() and not w._mac_title and getattr(w, "_title_inset", None) is None
        btn_ = w.toolbar.widgetForAction(w.act_si)
        ev_ = QMouseEvent(QEvent.MouseButtonPress, QPointF(btn_.geometry().center()), QPointF(),
                          Qt.LeftButton, Qt.LeftButton, Qt.NoModifier)
        assert MAC.TitleDrag(w).eventFilter(w.toolbar, ev_) is False
        report.append("macOS title bar (off offscreen, toolbar buttons keep clicks): OK")
        pump()                                             # Laden abgeschlossen (_restoring)
        t0_ = time.perf_counter()
        exd_ = json.loads(w.write_exchange().read_text(encoding="utf-8"))   # alle Tabs gefüllt
        dt_ = time.perf_counter() - t0_
        an_ = exd_["analytics"]
        assert {"synthesis", "nmr", "xray", "dft"} <= set(an_), sorted(an_)   # CV: Demo-Dateien liegen noch außerhalb
        assert an_["nmr"][0]["svg"].startswith("data/nmr/") and an_["xray"][0]["cif"].startswith("data/xray/"), an_
        assert {l_["kind"] for l_ in exd_["si"]["lines"]} >= {"nmr_1H", "uvvis"}, exd_["si"]["lines"]
        if "hrms" in an_ and "found" in an_["hrms"]:
            assert re.fullmatch(r"\d+\.\d{4}", an_["hrms"]["found"]) and an_["hrms"]["ion"].endswith(("+", "\u2212"))
        report.append(f"elucida.json for example project ({dt_:.2f} s, {len(an_)} techniques): OK")
        # Publikation: übergeordneter Ordner, Verbindungen in Reihenfolge, Datensatz (Zips + README), Gesamt-SI
        from elucida.app import publication as PUB
        from elucida.si import publication_si as PSI
        import shutil as _sh
        pub_ = out / "pub"
        _sh.rmtree(pub_, ignore_errors=True)
        pub_.mkdir()
        w.save_project()
        _sh.copytree(out / "demo", pub_ / "6m")
        PUB.create(pub_, "Test publication")
        assert [d_.name for d_ in PUB.sync(pub_)[0]] == ["6m"], PUB.load(pub_)
        w.open_publication(pub_)
        assert w.center.currentIndex() == 2 and w.project_dir is None and len(w.pub_page.rows) == 1
        assert w.new_compound_in_publication("2a") and Path(w.project_dir) == pub_ / "2a"
        assert Path(w.publication_dir) == pub_ and "Test publication" in w.windowTitle(), w.windowTitle()
        assert w.act_to_pub.isEnabled()
        PUB.move_compound(pub_, pub_ / "2a", -1)
        assert [d_.name for d_ in PUB.sync(pub_)[0]] == ["2a", "6m"]
        assert PUB.compound_name(pub_ / "2a") == "2a"
        w.go_publication()
        assert w.project_dir is None and w.center.currentIndex() == 2 and len(w.pub_page.rows) == 2
        PUB.remove_compound(pub_, pub_ / "2a")                     # Ordner bleibt, kommt nicht von selbst zurück
        assert [d_.name for d_ in PUB.sync(pub_)[0]] == ["6m"] and (pub_ / "2a").is_dir()
        assert PUB.publication_of(pub_ / "2a") is None and PUB.publication_of(pub_ / "6m") == pub_
        PUB.add_compound(pub_, pub_ / "2a")
        PUB.move_compound(pub_, pub_ / "2a", -1)
        d_ = PUB.load(pub_)
        d_.update(authors="A. Author and B. Author", doi="10.1234/test", license="CC-BY-4.0",
                  si_texts={"methods": "All reactions were carried out under argon.", "cv": "CV intro CH2Cl2."})
        PUB.save(pub_, d_)
        assert PUB.load(pub_)["si_texts"]["cv"] == "CV intro CH2Cl2."
        w.close_publication()
        w.add_recent_publication(pub_)
        w.fill_recent_list()
        assert w.center.currentIndex() == 0 and any(r_.info["path"] == str(pub_) for r_ in w.pub_rows)
        n6_ = PUB.compound_name(pub_ / "6m")                       # Verbindung des kopierten Beispielprojekts
        assert w.pub_rows[0].info["compounds"] == ["2a", n6_], w.pub_rows[0].info["compounds"]
        w.add_recent(pub_ / "6m")
        w.fill_recent_list()
        assert any(r_.info.get("publication") == "Test publication" for r_ in w.recent_rows)
        rd_ = PUB.readme_text(PUB.load(pub_), {"nmr", "cv", "dft"}, ["2a", "6m"])
        assert rd_.startswith("TITLE:\nTest publication") and "https://doi.org/10.1234/test" in rd_
        assert "CC-BY 4.0" in rd_ and "CompChem.zip" in rd_ and "NMR -> TopSpin" in rd_, rd_
        assert PUB.dataset_title({"paper_title": "Imidazo[1,5-<i>a</i>]pyridines"}) == \
            "Electronic supplement to Imidazo[1,5-a]pyridines"
        assert PUB.paper_title({"title": "Electronic supplement to X"}) == "X"
        mc_ = PSI.merge_crystal([(["Compound", "a", "b"], [["Formula", "C1", "C2"], ["Z", "4", "2"]]),
                                 (["Compound", "c", "d"], [["Formula", "C3", "C4"], ["Z", "8", "1"]])])
        assert [h_ for h_, _ in mc_] == [["Compound", "a", "b", "c"], ["Compound", "d"]] and mc_[1][1][1] == ["Z", "1"]
        from elucida.app.publication_view import PublicationPackDialog
        pdl_ = PublicationPackDialog(w, pub_)
        rep_ = pdl_.export(out / "pubpack")
        names_ = sorted(Path(z_).name for z_ in rep_["zips"])
        assert f"{n6_}.zip" in names_ and "CompChem.zip" in names_ and "2a.zip" not in names_, names_   # 2a ohne Daten
        assert (out / "pubpack" / "README.txt").exists() and (out / "pubpack" / "LICENSE-CC_BY_4.0_en.txt").exists()
        with _zfp.ZipFile(out / "pubpack" / "CompChem.zip") as z_:
            assert all(n_.startswith(f"CompChem/{n6_}/") for n_ in z_.namelist()), z_.namelist()[:5]
        with _zfp.ZipFile(out / "pubpack" / f"{n6_}.zip") as z_:
            assert not any("/DFT/" in n_ for n_ in z_.namelist())
        PUB.set_group(pub_, pub_ / "6m", "NHOs")                            # Gruppen im Projekt
        PUB.set_group(pub_, pub_ / "2a", "Precursor")
        PUB.move_group(pub_, "Precursor", -1)
        assert [(g_, [d_.name for d_ in ds_]) for g_, ds_ in PUB.grouped(pub_)] == [("Precursor", ["2a"]),
                                                                                    ("NHOs", ["6m"])]
        assert [d_.name for d_ in PUB.sync(pub_)[0]] == ["2a", "6m"]
        w.open_publication(pub_)
        assert [g_ for _, g_ in w.pub_page.headers] == ["Precursor", "NHOs"] and \
            [r_.group for r_ in w.pub_page.rows] == ["Precursor", "NHOs"]
        w.pub_page.ed_filter.setText("NHOs")                                 # Filter findet auch die Gruppe
        assert [r_.isVisibleTo(r_.parentWidget()) for r_ in w.pub_page.rows] == [False, True]
        w.pub_page.ed_filter.setText("")
        w.open_project(pub_ / "6m")
        pump(0.3)
        psi_ = w.export_publication_si(pub_, PUB.sync(pub_)[0], [k_ for k_, _, _ in PSI.SECTIONS],
                                       path=out / "pub_SI.docx")
        assert Path(w.project_dir) == pub_ / "6m", w.project_dir          # vorher offene Verbindung wieder da
        with _zfp.ZipFile(psi_) as z_:
            doc_ = z_.read("word/document.xml").decode("utf-8")
            assert "updateFields" in z_.read("word/settings.xml").decode("utf-8")
        heads_ = re.findall(r'w:val="Heading1".*?<w:t[^>]*>([^<]+)</w:t>', doc_)
        assert heads_[0] == "1. Materials and Methods" and "All reactions were carried out" in doc_, heads_
        order_ = [h_.split(". ", 1)[1] for h_ in heads_]
        ref_ = [t_ for _, t_, _ in PSI.SECTIONS]
        assert order_ == sorted(order_, key=ref_.index) and {"NMR Spectra", "Cyclic Voltammetry",
                                                             "X-Ray Characterization Data"} <= set(order_), order_
        assert doc_.index("CV intro") > doc_.index("Cyclic Voltammetry") and "TOC \\o" in doc_
        assert "Supplementary Materials for" in doc_ and "Synthesis of " in doc_
        subs_ = re.findall(r'w:val="Heading2".*?<w:t[^>]*>([^<]+)</w:t>', doc_)        # Unterkapitel je Gruppe
        assert subs_[:2] == ["2.1 Precursor", "2.2 NHOs"] and "3.1 NHOs" in subs_, subs_
        assert 'w:val="TOC2"' in doc_
        assert "PAGEREF _Toc1" in doc_ and 'w:name="_Toc1"' in doc_            # Verzeichnis schon gefüllt
        toc_ = SI.toc_entries([("title", "T"), ("toc",), ("section", "1. A"), ("char", "x"), ("section", "2. B")])
        assert toc_ == [("1. A", 3, 1), ("2. B", 4, 1)], toc_
        tp_ = SI.two_per_page([("char", "x " * 3000), ("image", "a.png", 16.0, 10.9), ("figcap", "c")])
        assert [b_[0] for b_ in tp_] == ["char", "pagebreak", "image", "figcap"] and tp_[2][3] == 10.9, \
            [b_[0] for b_ in tp_]                                            # langer Text: Bild auf neue Seite
        psi2_ = w.export_publication_si(pub_, PUB.sync(pub_)[0], ["char", "nmr"], path=out / "pub_SI_nature.docx",
                                        template="nature")
        with _zfp.ZipFile(psi2_) as z_:
            doc2_ = z_.read("word/document.xml").decode("utf-8")
        assert "Supplementary Information for" in doc2_ and "Supplementary Figure " in doc2_ and "Figure S" not in doc2_
        psi3_ = w.export_publication_si(pub_, PUB.sync(pub_)[0], ["char", "nmr"], path=out / "pub_SI_acs.docx",
                                        template="acs")
        with _zfp.ZipFile(psi3_) as z_:
            assert "PAGE" in z_.read("word/footer1.xml").decode("utf-8") and ">S<" in z_.read(
                "word/footer1.xml").decode("utf-8")                         # ACS: Seiten S1, S2 …
        assert SI.TEMPLATES and SI._TEMPLATE == "group"                    # danach wieder Standard
        orph_ = out / "orphan"                                            # Verbindung ohne Projekt zuordnen,
        _sh.rmtree(orph_, ignore_errors=True)                             # verschieben und duplizieren
        assert w.create_project(orph_)
        w.set_compound({"name": "9z"})
        w.save_project()
        assert w.publication_dir is None and w.btn_pub.text().startswith("\u26a0")
        new_ = w.move_compound(target=pub_)                               # offene Verbindung: schließen, verschieben,
        assert new_ == (pub_ / "orphan").resolve() and not orph_.exists()  # am neuen Ort wieder öffnen
        assert Path(w.project_dir).resolve() == new_ and Path(w.publication_dir).resolve() == pub_.resolve()
        assert PUB.compound_name(new_) == "9z" and new_ in [d_.resolve() for d_ in PUB.sync(pub_)[0]]
        pub2_ = out / "pub2"
        _sh.rmtree(pub2_, ignore_errors=True)
        PUB.create(pub2_, "Second paper")
        cp_ = w.move_compound(new_, copy=True, target=pub2_)               # dieselbe Verbindung im zweiten Paper
        assert cp_ == (pub2_ / "orphan").resolve() and new_.exists() and PRJ.is_project(cp_)
        assert not (cp_ / PRJ.WORK).exists() and [d_.name for d_ in PUB.sync(pub2_)[0]] == ["orphan"]
        mv_ = w.move_compound(new_, target=pub2_)                          # zweites Mal: freier Name
        assert mv_ == (pub2_ / "orphan_2").resolve() and not new_.exists()
        assert [d_.name for d_ in PUB.sync(pub2_)[0]] == ["orphan", "orphan_2"]
        assert not any(d_.name == "orphan" for d_ in PUB.sync(pub_)[0])
        # ChemDraw: leere Vorlagen je Verbindung, Übernahme beim Speichern, veraltete Exporte nicht verwenden
        s_t_, sch_t_ = PRJ.template_paths(w.project_dir)
        assert s_t_.exists() and sch_t_.exists() and not PRJ.cdxml_has_content(s_t_)
        assert 'BondLength="14.40"' in s_t_.read_text(encoding="utf-8")            # ACS 1996
        assert w.edit_in_chemdraw("structure").resolve() == s_t_.resolve()
        assert not w.chemdraw_saved(str(s_t_))                                    # leer: nichts übernehmen
        _sh.copy(cdxml, s_t_)
        assert w.chemdraw_saved(str(s_t_)) and Path(w.compound["structure"]).resolve() == s_t_.resolve()
        assert w.compound.get("formula") and w.compound.get("name") == "9z", w.compound
        assert w.chemdraw_file("structure").resolve() == s_t_.resolve()
        _sh.copy(cdxml, sch_t_)
        assert w.chemdraw_saved(str(sch_t_)) and Path(w.syn_state["scheme"]).resolve() == sch_t_.resolve()
        old_svg_ = s_t_.with_suffix(".svg")
        old_svg_.write_text("<svg xmlns='http://www.w3.org/2000/svg'/>", encoding="utf-8")
        t_old_ = s_t_.stat().st_mtime - 100
        os.utime(old_svg_, (t_old_, t_old_))
        stale_ = []
        assert IMG.find_sibling_drawing(s_t_, stale_) is None and stale_ == [old_svg_]
        os.utime(old_svg_, None)                                                  # frisch exportiert: gilt wieder
        assert IMG.find_sibling_drawing(s_t_) == old_svg_
        old_svg_.unlink()
        w.go_home()
        w.add_recent_publication(pub2_)                                   # Suche über alle bekannten Projekte
        w.ed_recent.setText("9z")
        assert [r_.info["name"] for r_ in w.recent_rows] == ["9z", "9z"], [r_.info["path"] for r_ in w.recent_rows]
        assert all(r_.info.get("publication") for r_ in w.recent_rows)
        w.ed_recent.setText("C5H5N 9z")                                   # mehrere Wörter: alle müssen passen
        w.ed_recent.setText("RE-9-999")
        assert not w.recent_rows and w.lbl_recent_empty.isVisibleTo(w.lbl_recent_empty.parentWidget())
        w.ed_recent.setText("")
        assert not getattr(w, "_searching", False)
        dpub_ = w.open_demo_publication(out / "demo_pub")                  # Beispielpublikation
        assert dpub_ and Path(w.publication_dir) == out / "demo_pub" and w.center.currentIndex() == 2
        assert [PUB.compound_name(d_) for d_ in PUB.sync(dpub_)[0]] == ["S2", "S3"] and len(w.pub_page.rows) == 2
        assert PUB.load(dpub_)["si_texts"]["methods"].startswith("All reactions")
        ck_ = PUB.checklist(dpub_ / "S2")                                 # Checkliste: neue Verbindung
        assert {"1H", "13C", "2D", "ms", "ir", "mp"} <= ck_["have"] and not ck_["missing"], ck_   # Beispiel vollständig
        w.open_project(dpub_ / "S3")
        assert w.btn_pub.isVisibleTo(w.btn_pub.parentWidget()) and w.btn_pub.text().startswith("Example")
        assert w.btn_home.isVisibleTo(w.btn_home.parentWidget()) and w.btn_comp.text() == "S3"   # Pfadleiste
        w._fill_comp_menu()
        assert [a_.text() for a_ in w.comp_menu.actions() if a_.isCheckable()] == ["S2", "S3"]
        assert [a_.isChecked() for a_ in w.comp_menu.actions() if a_.isCheckable()] == [False, True]
        w.set_compound({**w.compound, "known": True, "lit": "J. Test. Chem. 2020, 1, 1"})
        pm_ = w.si_char_parts({**w.compound, "oil": True})[0]                 # Öl: m.p. <25 °C
        assert pm_[0][:2] == ("mp", "m.p. <25 \u00b0C"), pm_[0]
        for in_, out_ in (("145", "145 \u00b0C"), ("145-147", "145-147 \u00b0C"), ("145.5 \u2013 147", "145.5 \u2013 147 \u00b0C"),
                          ("m.p. >300", ">300 \u00b0C"), ("145 \u00b0C", "145 \u00b0C"), ("145 \u00b0C (dec.)", "145 \u00b0C (dec.)"),
                          ("decomp.", "decomp.")):                      # m.p. ohne Einheit: °C ergänzen (Feedback)
            pm_ = w.si_char_parts({**w.compound, "mp": in_, "oil": False})[0]
            assert pm_[0][:2] == ("mp", "m.p. " + out_), (in_, pm_[0])
        w.save_project()
        ck_ = PUB.checklist(dpub_ / "S3")                                 # literaturbekannt: nur 1H/13C Pflicht
        assert ck_["known"] and ck_["missing"] == [] and ck_["lit"].startswith("J. Test"), ck_
        w.btn_pub.click()
        assert w.project_dir is None and w.center.currentIndex() == 2 and not w.btn_pub.isVisibleTo(
            w.btn_pub.parentWidget())
        w.pub_page.ed_filter.setText("S3")                                   # Filter auf der Projektseite
        assert [r_.isVisibleTo(r_.parentWidget()) for r_ in w.pub_page.rows] == [False, True]
        w.pub_page.ed_filter.setText("")
        assert "2 von 2" in w.pub_page.lbl_head.text() or "2 of 2" in w.pub_page.lbl_head.text(), \
            w.pub_page.lbl_head.text()
        w.open_project(dpub_ / "S2")                                      # Pfadleiste: ⌂ Übersicht
        w.btn_home.click()
        assert w.project_dir is None and w.publication_dir is None and w.center.currentIndex() == 0
        from elucida.app import demo as DEMO                              # Beispiele nie in der Übersicht
        real_dir_ = DEMO.IMG.app_data_dir
        DEMO.IMG.app_data_dir = lambda: out
        try:
            w.add_recent(out / "Beispielverbindung")
            w.add_recent_publication(out / "Beispielprojekt")
            assert not any(DEMO.is_demo(r_) for r_ in w.recent() + w.recent_publications())
            assert DEMO.is_demo(out / "Beispielprojekt" / "S2") and not DEMO.is_demo(out / "demo")
        finally:
            DEMO.IMG.app_data_dir = real_dir_
        from PySide6.QtCore import QFile as _QF                           # Papierkorb (Attrappe, nicht echt)
        trashed_ = []
        real_trash_ = _QF.moveToTrash
        _QF.moveToTrash = staticmethod(lambda p_: trashed_.append(p_) or True)
        try:
            w.add_recent(pub_ / "2a")
            w.add_recent_publication(pub_)
            assert w.trash_folder(pub_ / "2a", confirm=False) and trashed_ == [str(pub_ / "2a")]
            assert str(pub_ / "2a") not in w.recent()
            assert w.trash_folder(pub_, publication=True, confirm=False) and str(pub_) not in w.recent_publications()
            assert not any(pub_.resolve() in Path(r_).resolve().parents for r_ in w.recent())   # auch die darin
        finally:
            _QF.moveToTrash = real_trash_
        w.close_publication()
        report.append(f"publication (order, datapack {len(names_)} zips, complete SI {len(heads_)} chapters): OK")
        for lang in ("en", "de"):
            w.switch_lang(lang)
        import shiboken6 as _shb_end                                     # Beenden: alle Widgets vor dem Interpreter-
        from PySide6.QtWidgets import QWidget as _QW_end                 # Ende löschen (macOS-Absturz 2.3.7)
        MW.release_widgets(app)
        left_ = [type(o_).__name__ for o_ in _shb_end.getAllValidWrappers() if isinstance(o_, _QW_end)]
        assert not app.topLevelWidgets() and not left_, left_[:10]
        report.append("Quit: no widgets left for PySide's atexit cleanup: OK")
        report.append("OK")
    except Exception:
        ok = False
        report.append("FAILED\n" + traceback.format_exc())
    (out / "selftest_report.txt").write_text("\n".join(report), encoding="utf-8")
    try:                                   # Windows-Konsole (cp1252) kann z. B. kein λ: ersetzen statt abstürzen
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass
    print("\n".join(report))
    return 0 if ok else 1
