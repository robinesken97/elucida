"""elucida/tabs/xray.py - Kristallstrukturen aus CIF: ORTEP-Abbildung, Bindungen/Winkel/Diederwinkel, SI-Tabellen.

CIF:   eigener Leser (CIF 1.1-Syntax: data_-Blöcke, loop_, Strings in '…'/"…", Textfelder ;…;), Zahlen bleiben
       zusätzlich als Text erhalten, damit Werte mit Standardabweichung exakt wie in der CIF erscheinen.
Modell: fraktionale -> kartesische Koordinaten (a entlang x, b in der xy-Ebene), U_ij -> U_kart = M·N·U·N·Mᵀ.
       Bindungen aus _geom_bond (mit Symmetriecode n_klm), sonst aus kovalenten Radien. Molekül über die Bindungen
       und die Symmetrieoperationen vervollständigen; erzeugte Atome tragen einen Zusatz (C1ⁱ) mit Fußnote.
ORTEP: Ellipsoid mit Wahrscheinlichkeit p: xᵀU⁻¹x = k², k² = χ²-Quantil (3 Freiheitsgrade); 50 % -> k = 1.5382.
       Projektion auf den Bildschirm = 2×2-Block von R·U·Rᵀ (exakter Umriss), dazu die drei Hauptachsen-Ellipsen
       (vordere Hälften) wie in Olex2.
esd:   Standardabweichungen nur aus der CIF (_geom_bond/_geom_angle/_geom_torsion), nie geschätzt.
"""
from __future__ import annotations

import csv
import math
import re
from pathlib import Path

import numpy as np

from elucida.core import orbitals as ORB
from elucida.formats.elements import (COV, element_of)  # noqa: E402,F401  (Leser in elucida/formats)
from elucida.formats.cif import (Block, parse_cif_text, read_cif, is_cif, _NUM, num)  # noqa: E402,F401  (Leser in elucida/formats)

# ---------------------------------------------------------------------------- CIF lesen


def structure_blocks(path) -> list[Block]:
    """Nur Blöcke mit Atomen (Olex2/SHELXL: ein Block je Struktur)."""
    return [b for b in read_cif(path) if b.column("_atom_site_fract_x")]


def esd_parts(s):
    """'1.392(3)' -> ('1.392', '3'); ohne esd -> ('1.392', None)."""
    m = _NUM.match(str(s).strip())
    if not m:
        return str(s), None
    return m.group(1), m.group(2)


def minus(s):
    """Echtes Minuszeichen statt Bindestrich (Tabellen, Diff.-Dichte)."""
    return re.sub(r"(^|[\s(/\[])-(?=[\d.])", "\\1\u2212", str(s))


# ---------------------------------------------------------------------------- Symmetrie
def parse_symop(s):
    """'-x+1/2, y+1/2, -z+1/2' -> (R 3×3 int, t 3)."""
    parts = [p.strip().lower().replace(" ", "") for p in str(s).split(",")]
    if len(parts) != 3:
        raise ValueError(f"Symmetrieoperation unlesbar: {s}")
    R = np.zeros((3, 3), int)
    t = np.zeros(3)
    for i, p in enumerate(parts):
        for term in re.findall(r"[+-]?[^+-]+", p):
            sign = -1 if term.startswith("-") else 1
            term = term.lstrip("+-")
            m = re.match(r"^(\d*\.?\d*(?:/\d+)?)\*?([xyz])$", term)
            if m:
                coef = m.group(1)
                c = float(eval_frac(coef)) if coef else 1.0
                R[i, "xyz".index(m.group(2))] += int(round(sign * c))
            else:
                t[i] += sign * eval_frac(term)
    return R, t


def eval_frac(s):
    if "/" in s:
        a, b = s.split("/")
        return float(a) / float(b)
    return float(s)


IDENT = (np.eye(3, dtype=int), np.zeros(3))


def tr_key(T):
    R, t = T
    return tuple(int(v) for v in R.ravel()) + tuple(round(float(v), 4) + 0.0 for v in t)


def compose(A, B):
    """(A∘B)(x) = A(B(x))."""
    return A[0] @ B[0], A[0] @ B[1] + A[1]


def inverse(A):
    Ri = np.rint(np.linalg.inv(A[0])).astype(int)
    return Ri, -Ri @ A[1]


def apply(T, frac):
    return T[0] @ frac + T[1]


def is_identity(T):
    return tr_key(T) == tr_key(IDENT)


def op_string(T):
    """(R, t) -> '1−x, 1−y, 1−z' (wie Fußnoten in SI-Tabellen)."""
    out = []
    for i in range(3):
        terms = []
        tv = float(T[1][i])
        if abs(tv) > 1e-6:
            fr = _frac_str(tv)
            terms.append(fr)
        for j, ax in enumerate("xyz"):
            c = int(T[0][i, j])
            if c:
                s = ("−" if c < 0 else ("+" if terms else "")) + (str(abs(c)) if abs(c) != 1 else "") + ax
                terms.append(s)
        out.append("".join(terms) or "0")
    return ", ".join(out)


def _frac_str(v):
    for den in (1, 2, 3, 4, 6, 8, 12):
        if abs(v * den - round(v * den)) < 1e-4:
            nume = int(round(v * den))
            s = str(abs(nume)) if den == 1 else f"{abs(nume)}/{den}"
            return ("−" if nume < 0 else "") + s
    return f"{v:g}".replace("-", "−")


def symcode(code, ops):
    """SHELXL-Code '3_665' -> Transformation (Operation 3, Translation +1,+1,0); '.' -> Identität."""
    if code in (None, ".", "?", "") or code == "1_555":
        return IDENT
    m = re.match(r"^(\d+)(?:_(\d+))?$", str(code))
    if not m:
        raise ValueError(f"Symmetriecode unlesbar: {code}")
    n, tt = int(m.group(1)), m.group(2) or "555"            # '3' = Operation 3 ohne Translation
    if len(tt) % 3:
        raise ValueError(f"Symmetriecode unlesbar: {code}")
    k = len(tt) // 3                             # 555 = keine Translation; zweistellig (z. B. 1_105555) -> 55
    shift = np.array([int(tt[i * k:(i + 1) * k]) for i in range(3)], float) - (5 if k == 1 else 55)
    R, t = ops[n - 1]
    return R, t + shift


ROMAN = ["i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x", "xi", "xii", "xiii", "xiv", "xv", "xvi"]
SUP = str.maketrans("0123456789+−-=()nivx", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁻⁼⁽⁾ⁿⁱᵛˣ")
SUB = str.maketrans("0123456789+−-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₋₌₍₎")


def roman(i):
    return ROMAN[i] if i < len(ROMAN) else str(i + 1)


# ---------------------------------------------------------------------------- Elemente
# Olex2-ähnliche Standardfarben (Beispielbilder), sonst Jmol/CPK
OLEX2_COLORS = {"C": "#8C8C8C", "H": "#FFFFFF", "N": "#1F3FE0", "O": "#E01010", "F": "#C8E632", "Cl": "#3CD23C",
                "S": "#E6A03C", "B": "#FFB5B5", "Br": "#A62929", "P": "#FF8000", "I": "#940094", "Si": "#F0C8A0"}
METAL_COLOR = "#0A7D8C"
JMOL = {"He": "#D9FFFF", "Li": "#CC80FF", "Be": "#C2FF00", "Ne": "#B3E3F5", "Na": "#AB5CF2", "Mg": "#8AFF00",
        "Al": "#BFA6A6", "Ar": "#80D1E3", "K": "#8F40D4", "Ca": "#3DFF00", "Se": "#FFA100", "Kr": "#5CB8D1",
        "Rb": "#702EB0", "Sr": "#00FF00", "Te": "#D47A00", "Xe": "#429EB0", "Cs": "#57178F", "Ba": "#00C900",
        "As": "#BD80E3", "Ge": "#668F8F", "Sb": "#9E63B5", "Ga": "#C28F8F", "In": "#A67573", "Sn": "#668080",
        "Tl": "#A6544D", "Pb": "#575961", "Bi": "#9E4FB5", "D": "#FFFFC0"}
METALS = set("Sc Ti V Cr Mn Fe Co Ni Cu Zn Y Zr Nb Mo Tc Ru Rh Pd Ag Cd Hf Ta W Re Os Ir Pt Au Hg "
             "La Ce Pr Nd Sm Eu Gd Tb Dy Ho Er Tm Yb Lu U Th".split())
BOND_TOL = 0.45


# Kugel-Stab wie im ORCA-Handbuch (Avogadro/Jmol-Farben, gesättigt): nur für den Stil „depth“
ORCA_COLORS = {"C": "#909090", "H": "#FFFFFF", "N": "#3050F8", "O": "#FF0D0D", "F": "#90E050", "Cl": "#1FF01F",
               "Br": "#A62929", "I": "#940094", "S": "#FFFF30", "P": "#FF8000", "Si": "#F0C8A0", "B": "#FFB5B5"}


def default_color(el, style=None):
    if style == "iboview":
        return ORB.ibo_color(el)
    if style == "depth" and el in ORCA_COLORS:
        return ORCA_COLORS[el]
    if el in OLEX2_COLORS:
        return OLEX2_COLORS[el]
    if el in METALS:
        return METAL_COLOR
    return JMOL.get(el, "#FF1493")


# ---------------------------------------------------------------------------- Zelle / ADP
def cell_of(b: Block):
    vals = [b.num(f"_cell_length_{x}") for x in "abc"] + [b.num(f"_cell_angle_{x}") for x in ("alpha", "beta", "gamma")]
    if any(v is None for v in vals):
        raise ValueError("Zellparameter fehlen")
    return vals


def ortho_matrix(a, b, c, al, be, ga):
    """Fraktional -> kartesisch (Å), a entlang x, b in der xy-Ebene."""
    ca, cb, cg = (math.cos(math.radians(v)) for v in (al, be, ga))
    sg = math.sin(math.radians(ga))
    v = a * b * c * math.sqrt(max(1e-12, 1 - ca * ca - cb * cb - cg * cg + 2 * ca * cb * cg))
    return np.array([[a, b * cg, c * cb],
                     [0, b * sg, c * (ca - cb * cg) / sg],
                     [0, 0, v / (a * b * sg)]])


def u_cart(M, U):
    """U_ij (CIF, bezogen auf a*, b*, c*) -> kartesische U-Matrix."""
    Minv = np.linalg.inv(M)
    nstar = np.diag(np.linalg.norm(Minv, axis=1))   # Zeilen von M⁻¹ = reziproke Vektoren -> |a*|, |b*|, |c*|
    A = M @ nstar
    return A @ U @ A.T


def chi2_scale(p):
    """k mit P(χ²₃ ≤ k²) = p (Ellipsoid mit Aufenthaltswahrscheinlichkeit p); 0.5 -> 1.5382."""
    p = min(max(float(p), 0.01), 0.999)

    def cdf(k):        # χ²-Verteilung mit 3 Freiheitsgraden, x = k²
        return math.erf(k / math.sqrt(2)) - math.sqrt(2 / math.pi) * k * math.exp(-k * k / 2)
    lo, hi = 0.0, 10.0
    for _ in range(80):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if cdf(mid) < p else (lo, mid)
    return (lo + hi) / 2


# ---------------------------------------------------------------------------- Struktur
class Structure:
    """Asymmetrische Einheit, Symmetrie, Bindungen (ASU-bezogen) und geom-Tabellen eines Blocks."""

    def __init__(self, block: Block):
        self.block = block
        self.name = block.name
        self.cell = cell_of(block)
        self.M = ortho_matrix(*self.cell)
        ops = block.column("_space_group_symop_operation_xyz", "_symmetry_equiv_pos_as_xyz") or ["x, y, z"]
        self.ops = [parse_symop(o) for o in ops]
        self.atoms = []
        aniso = {}
        for r in block.rows("_atom_site_aniso_label"):
            try:
                u = [num(r.get(f"_atom_site_aniso_u_{k}")) for k in ("11", "22", "33", "23", "13", "12")]
                if None in u:
                    bb = [num(r.get(f"_atom_site_aniso_b_{k}")) for k in ("11", "22", "33", "23", "13", "12")]
                    if None in bb:
                        continue
                    u = [v / (8 * math.pi ** 2) for v in bb]
                u11, u22, u33, u23, u13, u12 = u
                aniso[r["_atom_site_aniso_label"]] = np.array([[u11, u12, u13], [u12, u22, u23], [u13, u23, u33]])
            except (KeyError, TypeError):
                continue
        for r in block.rows("_atom_site_label"):
            lab = r["_atom_site_label"]
            xyz = [num(r.get(f"_atom_site_fract_{k}")) for k in "xyz"]
            if None in xyz:
                continue
            el = element_of(r.get("_atom_site_type_symbol"), lab)
            uiso = num(r.get("_atom_site_u_iso_or_equiv"))
            if uiso is None and num(r.get("_atom_site_b_iso_or_equiv")) is not None:
                uiso = num(r.get("_atom_site_b_iso_or_equiv")) / (8 * math.pi ** 2)
            part = r.get("_atom_site_disorder_group")
            try:
                part = int(part) if part not in (None, ".", "?") else 0
            except ValueError:
                part = 0
            occ = num(r.get("_atom_site_occupancy"))
            U = aniso.get(lab)
            self.atoms.append(dict(label=lab, el=el, frac=np.array(xyz, float), uiso=uiso or 0.02,
                                   U=u_cart(self.M, U) if U is not None else None,
                                   occ=1.0 if occ is None else occ, part=part,
                                   assembly=r.get("_atom_site_disorder_assembly") or "",
                                   calc=(r.get("_atom_site_calc_flag") or "").lower()))
        self.by_label = {a["label"]: i for i, a in enumerate(self.atoms)}
        self.geom_bond = self._geom("_geom_bond_atom_site_label_1", 2, "_geom_bond_distance", "_geom_bond")
        self.geom_angle = self._geom("_geom_angle_atom_site_label_1", 3, "_geom_angle", "_geom_angle")
        self.geom_torsion = self._geom("_geom_torsion_atom_site_label_1", 4, "_geom_torsion", "_geom_torsion")

    def _geom(self, first, n, value_tag, prefix):
        out = []
        for r in self.block.rows(first):
            labs = [r.get(f"{prefix}_atom_site_label_{k + 1}") for k in range(n)]
            val = r.get(value_tag)
            if None in labs or val in (None, "?", "."):
                continue
            codes = [r.get(f"{prefix}_site_symmetry_{k + 1}") for k in range(n)]
            if n == 3 and not codes[1]:
                codes[1] = "."
            try:
                trs = [symcode(c, self.ops) for c in codes]
            except (ValueError, IndexError):
                continue
            out.append(dict(labels=labs, trs=trs, value=val))
        return out

    def cart(self, frac):
        return self.M @ frac

    # -- Fehlordnung --------------------------------------------------------------------
    def visible_asu(self, mode="main"):
        """Indizes der gezeigten ASU-Atome: 'main' = PART 0 + je Gruppe der Teil mit der größten Besetzung."""
        if mode == "all":
            return list(range(len(self.atoms)))
        if mode == "minor":
            main = set(self.visible_asu("main"))
            return [i for i, a in enumerate(self.atoms) if a["part"] == 0 or i not in main]
        groups = {}
        for i, a in enumerate(self.atoms):
            if a["part"]:
                groups.setdefault(a["assembly"], {}).setdefault(abs(a["part"]), []).append(a["occ"])
        keep = {}
        for asm, parts in groups.items():
            keep[asm] = max(sorted(parts), key=lambda p: (np.mean(parts[p]) - 1e-9 * p))
        return [i for i, a in enumerate(self.atoms) if not a["part"] or keep.get(a["assembly"]) == abs(a["part"])]

    def parts(self):
        return sorted({a["part"] for a in self.atoms if a["part"]})

    # -- Bindungen in ASU-Form: (i, j, T) = Atom i (Identität) gebunden an T(Atom j) ----------
    def asu_bonds(self, visible):
        vis = set(visible)
        out = []
        if self.geom_bond:
            have_h = set()
            for g in self.geom_bond:
                i, j = (self.by_label.get(x) for x in g["labels"])
                if i is None or j is None or i not in vis or j not in vis:
                    continue
                if not self._part_ok(i, j):
                    continue
                Ti = g["trs"][0]
                T = compose(inverse(Ti), g["trs"][1])
                out.append((i, j, T))
                for k in (i, j):
                    if self.atoms[k]["el"] in ("H", "D"):
                        have_h.add(k)
            # H ohne Eintrag in _geom_bond (BOND $H fehlte): an das nächste Nicht-H-Atom hängen
            missing = [k for k in vis if self.atoms[k]["el"] in ("H", "D") and k not in have_h]
            if missing:
                cov = self._covalent(vis)
                out += [b for b in cov if (b[0] in missing) != (b[1] in missing)
                        and ({self.atoms[b[0]]["el"], self.atoms[b[1]]["el"]} - {"H", "D"})]
        else:
            out = self._covalent(vis)
        full = []
        for i, j, T in out:                          # beide Richtungen
            full.append((i, j, T))
            full.append((j, i, inverse(T)))
        return full

    def _part_ok(self, i, j):
        a, b = self.atoms[i], self.atoms[j]
        return not (a["part"] and b["part"] and a["part"] != b["part"] and a["assembly"] == b["assembly"])

    def _covalent(self, vis):
        idx = sorted(vis)
        if not idx:
            return []
        F = np.array([self.atoms[i]["frac"] for i in idx])
        rad = np.array([COV.get(self.atoms[i]["el"], 1.5) for i in idx])
        isH = np.array([self.atoms[i]["el"] in ("H", "D") for i in idx])
        cut = rad[:, None] + rad[None, :] + BOND_TOL
        offs = np.array([(a, b, c) for a in (-1, 0, 1) for b in (-1, 0, 1) for c in (-1, 0, 1)], float)
        out, seen = [], set()
        for op in self.ops:
            img = (op[0] @ F.T).T + op[1]                       # Bilder aller Atome j
            D = img[None, :, :] - F[:, None, :]                 # (i, j, 3)
            base = -np.rint(D)
            for o in offs:
                Dc = (D + base + o) @ self.M.T
                d = np.linalg.norm(Dc, axis=2)
                hit = np.argwhere((d > 0.4) & (d < cut) & ~(isH[:, None] & isH[None, :]))
                for a, b in hit:
                    i, j = idx[a], idx[b]
                    if not self._part_ok(i, j):
                        continue
                    T = (op[0], op[1] + base[a, b] + o)
                    key = (i, j, tr_key(T))
                    rkey = (j, i, tr_key(inverse(T)))
                    if key in seen or rkey in seen:
                        continue
                    if i == j and is_identity(T):
                        continue
                    seen.add(key)
                    out.append((i, j, T))
        return out


class Molecule:
    """Vervollständigtes Modell: gezeigte Atome (mit Transformation zur ASU), Bindungen, Fragmente."""

    def __init__(self, S: Structure, disorder="main", max_factor=8):
        self.S = S
        vis = S.visible_asu(disorder)
        nbrs = {}
        for i, j, Tr in S.asu_bonds(vis):
            nbrs.setdefault(i, []).append((j, Tr))
        self.atoms, self.bonds = [], set()
        index, by_label = {}, {}
        self.truncated = False
        limit = max_factor * max(1, len(vis)) + 50

        def place(i, T):
            frac = apply(T, S.atoms[i]["frac"])
            key = (i,) + tuple(np.round(frac, 3) + 0.0)
            if key in index:
                return index[key], False
            R = T[0]
            a = S.atoms[i]
            U = None
            if a["U"] is not None:                   # U_kart dreht mit dem kartesischen Teil der Operation
                Rc = S.M @ R @ np.linalg.inv(S.M)
                U = Rc @ a["U"] @ Rc.T
            generated = not (np.array_equal(R, np.eye(3, dtype=int)) and i not in by_label)
            k = len(self.atoms)
            self.atoms.append(dict(asu=i, label=a["label"], el=a["el"], T=T, frac=frac, xyz=S.cart(frac), U=U,
                                   uiso=a["uiso"], part=a["part"], generated=generated and not is_identity(T)))
            index[key] = k
            by_label.setdefault(i, []).append(k)
            return k, True

        for start in vis:
            if start in by_label:
                continue
            k0, _ = place(start, IDENT)
            queue = [k0]
            while queue:
                k = queue.pop(0)
                at = self.atoms[k]
                for j, Trel in nbrs.get(at["asu"], []):
                    if len(self.atoms) >= limit:
                        self.truncated = True
                        break
                    kj, new = place(j, compose(at["T"], Trel))
                    if kj != k:
                        self.bonds.add((min(k, kj), max(k, kj)))
                    if new:
                        queue.append(kj)
        # Symmetriezusätze: je unterschiedlicher Operation eine römische Ziffer
        self.ops = {}
        for at in self.atoms:
            if at["generated"]:
                key = op_string(at["T"])
                if key not in self.ops:
                    self.ops[key] = roman(len(self.ops))
                at["sym"] = self.ops[key]
                at["opkey"] = key
            else:
                at["sym"], at["opkey"] = "", ""
            at["key"] = at["label"] + ("|" + at["opkey"] if at["opkey"] else "")
        self.bonds = sorted(self.bonds)
        self._fragments()

    def _fragments(self):
        n = len(self.atoms)
        parent = list(range(n))

        def find(x):
            while parent[x] != x:
                parent[x] = parent[parent[x]]
                x = parent[x]
            return x
        for a, b in self.bonds:
            parent[find(a)] = find(b)
        comps = {}
        for k in range(n):
            comps.setdefault(find(k), []).append(k)
        info = {}
        for root, c in comps.items():
            els = [self.atoms[k]["el"] for k in c]
            info[root] = (frag_kind(hill(els), hill([e for e in els if e not in ("H", "D")])), els)
        # Hauptmolekül = größtes Fragment, das kein bekanntes Lösungsmittel/Gegenion ist (z. B. BArF⁻ > Kation)
        frags = sorted(comps.items(), key=lambda rc: (info[rc[0]][0][0] is not None,
                                                      -sum(e not in ("H", "D") for e in info[rc[0]][1]), min(rc[1])))
        self.fragments = []
        main_heavy = None
        for fi, (root, c) in enumerate(frags):
            for k in c:
                self.atoms[k]["frag"] = fi
            (kind, name), els = info[root]
            heavy = hill([e for e in els if e not in ("H", "D")])
            if fi == 0:
                kind, main_heavy = "main", heavy
            elif kind is None:
                kind = "molecule" if heavy == main_heavy else "other"
            self.fragments.append(dict(atoms=c, formula=hill(els), kind=kind, name=name))
        self._nbrs = None

    def neighbors(self):
        if self._nbrs is None:
            self._nbrs = {k: [] for k in range(len(self.atoms))}
            for a, b in self.bonds:
                self._nbrs[a].append(b)
                self._nbrs[b].append(a)
        return self._nbrs

    def substituent(self, a, b):
        """Atome hinter der Bindung a–b (die kleinere Seite, nach Nicht-H-Atomen). Liegt a–b in einem Ring: das
        Ringsystem hinter a (siehe beyond) oder None."""
        nb = self.neighbors()

        def side(start, block):
            seen, todo = {start}, [start]
            while todo:
                x = todo.pop()
                for y in nb[x]:
                    if y not in seen and not (x == start and y == block):
                        seen.add(y)
                        todo.append(y)
            return seen
        heavy = lambda c: sum(self.atoms[k]["el"] not in ("H", "D") for k in c)  # noqa: E731
        sb = side(b, a)
        if a not in sb:
            sa = side(a, b)
            return sb if heavy(sb) <= heavy(sa) else sa
        return self.beyond(a, b)

    def beyond(self, a, b):
        """Ringsystem hinter dem Verknüpfungsatom a (z. B. Spiro-C): alles, was von b aus ohne a erreichbar ist,
        sofern es höchstens die Hälfte des Fragments ist; sonst None."""
        nb = self.neighbors()
        seen, todo = {b}, [b]
        while todo:
            x = todo.pop()
            for y in nb[x]:
                if y != a and y not in seen:
                    seen.add(y)
                    todo.append(y)
        frag = [k for k, at in enumerate(self.atoms) if at.get("frag") == self.atoms[a].get("frag")]
        heavy = lambda c: sum(self.atoms[k]["el"] not in ("H", "D") for k in c)  # noqa: E731
        return seen if heavy(seen) <= (heavy(frag) - 1) / 2 else None

    def find_key(self, key):
        for k, a in enumerate(self.atoms):
            if a["key"] == key:
                return k
        return None

    def footnotes(self):
        """['i: 1−x, 1−y, 1−z', …]"""
        return [f"{r}: {op}" for op, r in self.ops.items()]


# Bekannte Lösungsmittel und Gegenionen (Hill-Formel mit H; ohne H zusätzlich über die Schweratome, weil H in
# fehlgeordneten Lösungsmitteln oft nicht verfeinert sind). Name "" = mehrdeutig (dann Formel in der Bildunterschrift).
SOLVENTS = {"C4H8O": "THF", "C4H10O": "Et₂O", "CH2Cl2": "CH₂Cl₂", "CHCl3": "CHCl₃", "CCl4": "CCl₄", "C7H8": "toluene",
            "C6H6": "benzene", "C6H14": "n-hexane", "C5H12": "n-pentane", "C7H16": "n-heptane", "C6H12": "cyclohexane",
            "C2H3N": "MeCN", "H2O": "water", "CH4O": "MeOH", "C2H6O": "EtOH", "C3H8O": "iPrOH", "C3H6O": "acetone",
            "C2H6OS": "DMSO", "C3H7NO": "DMF", "C4H8O2": "", "C4H10O2": "DME", "C2H4Cl2": "1,2-dichloroethane",
            "C6H5F": "fluorobenzene", "C6H4F2": "difluorobenzene", "C6H5Cl": "chlorobenzene",
            "C6H4Cl2": "dichlorobenzene", "C5H5N": "pyridine", "C8H10": "xylene", "C5H10O": "2-MeTHF",
            "C5H12O": "MTBE", "C3H4O3": "", "CS2": "CS₂"}
IONS = {"CF3O3S": "OTf⁻", "BF4": "BF₄⁻", "F6P": "PF₆⁻", "F6Sb": "SbF₆⁻", "AsF6": "AsF₆⁻", "ClO4": "ClO₄⁻",
        "NO3": "NO₃⁻", "C2F6NO4S2": "NTf₂⁻", "C32H12BF24": "BArF₂₄⁻", "C24BF20": "[B(C₆F₅)₄]⁻",
        "C24H20B": "BPh₄⁻", "C16AlF36O4": "[Al(OC(CF₃)₃)₄]⁻", "AlCl4": "AlCl₄⁻", "Cl4Ga": "GaCl₄⁻",
        "Cl4Fe": "FeCl₄⁻", "Cl6Sb": "SbCl₆⁻", "I3": "I₃⁻", "Cl": "Cl⁻", "Br": "Br⁻", "I": "I⁻", "F": "F⁻"}


def _heavy_map(d):
    out = {}
    for f, name in d.items():
        hv = re.sub(r"H\d*(?![a-z])", "", f)
        out.setdefault(hv, set()).add(name)
    return out


SOLVENTS_HEAVY, IONS_HEAVY = _heavy_map(SOLVENTS), _heavy_map(IONS)


def frag_kind(formula, heavy):
    """-> (Art 'solvent'|'ion'|None, Name) für ein abgetrenntes Fragment (Hill-Formeln mit und ohne H)."""
    for kind, full, hv in (("solvent", SOLVENTS, SOLVENTS_HEAVY), ("ion", IONS, IONS_HEAVY)):
        if formula in full:
            return kind, full[formula]
        if heavy in hv and heavy:
            names = hv[heavy]
            return kind, next(iter(names)) if len(names) == 1 else ""
    return None, ""


def disorder_occupancy(S):
    """'95% (A) and 5% (B)' für eine Fehlordnung mit zwei Teilen (Buchstaben aus den Atomlabels, sonst PART-Nummer);
    None ohne oder bei mehreren Fehlordnungen."""
    parts = {}
    for a in S.atoms:
        if a["part"]:
            parts.setdefault(abs(a["part"]), []).append(a)
    if len(parts) != 2 or len({a["assembly"] for v in parts.values() for a in v}) > 1:
        return None
    bits = []
    for p_, ats in sorted(parts.items()):
        occ = round(100 * float(np.mean([a["occ"] for a in ats])))
        tags = {re.sub(r"^[A-Z][a-z]?\d+", "", a["label"]) for a in ats}
        tag = next(iter(tags)) if len(tags) == 1 and next(iter(tags)) else f"PART {p_}"
        bits.append(f"{occ}% ({tag})")
    return " and ".join(bits)


def hill(els):
    cnt = {}
    for e in els:
        cnt[e] = cnt.get(e, 0) + 1
    keys = [k for k in ("C", "H") if k in cnt] + sorted(k for k in cnt if k not in ("C", "H"))
    if "C" not in cnt:
        keys = sorted(cnt)
    return "".join(f"{k}{cnt[k] if cnt[k] > 1 else ''}" for k in keys)


# ---------------------------------------------------------------------------- Geometrie
def distance(p):
    return float(np.linalg.norm(p[1] - p[0]))


def angle(p):
    u, v = p[0] - p[1], p[2] - p[1]
    return float(math.degrees(math.atan2(np.linalg.norm(np.cross(u, v)), np.dot(u, v))))


def torsion(p):
    """IUPAC-Vorzeichen (wie SHELXL): Blick entlang B->C, A nach D im Uhrzeigersinn = positiv."""
    b1, b2, b3 = p[1] - p[0], p[2] - p[1], p[3] - p[2]
    x = np.dot(np.cross(b1, b2), np.cross(b2, b3))
    y = np.linalg.norm(b2) * np.dot(b1, np.cross(b2, b3))
    return float(math.degrees(math.atan2(y, x)))


def match_geom(S: Structure, rows, atoms):
    """Eintrag aus _geom_* zu gewählten Atomen (Label + Transformation T): gesucht ist G mit T_k = G∘S_k für alle k,
    auch in umgekehrter Reihenfolge. -> Textwert aus der CIF oder None."""
    labs = [a["label"] for a in atoms]
    for g in rows:
        for order in (labs, labs[::-1]):
            if g["labels"] != order:
                continue
            ats = atoms if order is labs else atoms[::-1]
            G = compose(ats[0]["T"], inverse(g["trs"][0]))
            if all(tr_key(compose(G, s)) == tr_key(a["T"]) for s, a in zip(g["trs"], ats)):
                return g["value"]
    return None


def cif_order(mol: Molecule, keys):
    """Atomreihenfolge wie in der _geom-Tabelle der CIF (z. B. C4–C5 statt angeklickt C5–C4); sonst unverändert."""
    idx = [mol.find_key(k) for k in keys]
    if None in idx or len(idx) not in (2, 3, 4):
        return list(keys)
    ats = [mol.atoms[k] for k in idx]
    rows = {2: mol.S.geom_bond, 3: mol.S.geom_angle, 4: mol.S.geom_torsion}[len(idx)]
    labs = [a["label"] for a in ats]
    if labs == labs[::-1]:
        return list(keys)
    for g in rows:
        if g["labels"] == labs:
            return list(keys)
        if g["labels"] == labs[::-1]:
            rev = ats[::-1]
            G = compose(rev[0]["T"], inverse(g["trs"][0]))
            if all(tr_key(compose(G, t)) == tr_key(a["T"]) for t, a in zip(g["trs"], rev)):
                return list(keys)[::-1]
    return list(keys)


def measure(mol: Molecule, keys):
    """2/3/4 Atome -> dict(kind, value (float), text, esd (bool), source ('cif'|'calc'))."""
    idx = [mol.find_key(k) for k in keys]
    if None in idx or len(idx) not in (2, 3, 4) or len(set(idx)) != len(idx):
        return None
    ats = [mol.atoms[k] for k in idx]
    pts = [a["xyz"] for a in ats]
    kind = {2: "bond", 3: "angle", 4: "torsion"}[len(idx)]
    val = {2: distance, 3: angle, 4: torsion}[len(idx)](pts)
    txt = None
    if mol.S is not None:                            # CIF: Wert mit esd aus den geom-Tabellen; xyz: nur rechnen
        rows = {2: mol.S.geom_bond, 3: mol.S.geom_angle, 4: mol.S.geom_torsion}[len(idx)]
        txt = match_geom(mol.S, rows, ats)
    if txt is not None:
        return dict(kind=kind, value=val, text=txt, esd=esd_parts(txt)[1] is not None, source="cif")
    dec = 3 if kind == "bond" else 1
    txt = f"{val:.{dec}f}"
    if float(txt) == 0:
        txt = txt.lstrip("-")                        # kein „−0.0“
    return dict(kind=kind, value=val, text=txt, esd=False, source="calc")


def atom_name(a):
    return a["label"] + (a["sym"].translate(SUP) if a.get("sym") else "")


def param_name(mol, keys, sep="–"):
    names = []
    for k in keys:
        i = mol.find_key(k)
        names.append(atom_name(mol.atoms[i]) if i is not None else k.split("|")[0])
    return sep.join(names)


def param_label(keys):
    """Zeilenname unabhängig vom Modell (für die Tabelle über mehrere Verbindungen): 'C1–C2ⁱ'."""
    out = []
    for k in keys:
        lab, _, op = k.partition("|")
        out.append(lab)
    return "–".join(out)


def _fib_dirs(n):
    """n gleichmäßig verteilte Richtungen auf der oberen Halbkugel (Gegenrichtung zeigt dieselbe Überdeckung)."""
    i = np.arange(n) + 0.5
    z = 1 - i / n
    r = np.sqrt(1 - z ** 2)
    phi = i * math.pi * (3 - math.sqrt(5))
    return np.column_stack([r * np.cos(phi), r * np.sin(phi), z])


def overlap_score(P, d0=0.9):
    """Überdeckung projizierter Atome (2D, Å): Summe exp(−(d/d0)²) über alle Paare – kleiner = übersichtlicher."""
    diff = P[:, None, :] - P[None, :, :]
    d2 = np.einsum("ijk,ijk->ij", diff, diff)
    iu = np.triu_indices(len(P), 1)
    return float(np.exp(-d2[iu] / d0 ** 2).sum())


def default_view(mol: Molecule, frag=0, use=None, n_dirs=240):
    """Startansicht: Blickrichtung mit der geringsten Überdeckung der Nicht-H-Atome (Halbkugel abgetastet, dazu die
    Normale der Hauptebene), größte Ausdehnung waagerecht. use = zu berücksichtigende Atomindizes (sonst das
    Hauptfragment ohne H)."""
    idx = list(use) if use else [k for k, a in enumerate(mol.atoms) if a.get("frag") == frag
                                  and a["el"] not in ("H", "D")]
    pts = np.array([mol.atoms[k]["xyz"] for k in idx] or [a["xyz"] for a in mol.atoms] or [[0, 0, 0]], float)
    if len(pts) < 3:
        return np.eye(3)
    c = pts - pts.mean(axis=0)
    if len(c) > 160:                                   # sehr große Strukturen: Stichprobe reicht für die Wahl
        c = c[np.linspace(0, len(c) - 1, 160).astype(int)]
    w, v = np.linalg.eigh(c.T @ c)
    normal = v[:, 0]
    best = None
    for n, d in enumerate([normal] + list(_fib_dirs(n_dirs))):
        d = d / np.linalg.norm(d)
        e1 = np.cross(d, [1.0, 0, 0] if abs(d[0]) < 0.9 else [0, 1.0, 0])
        e1 /= np.linalg.norm(e1)
        e2 = np.cross(d, e1)
        P = np.column_stack([c @ e1, c @ e2])
        sc = overlap_score(P) * (0.97 if n == 0 else 1.0)      # bei Gleichstand die Hauptebene bevorzugen
        if best is None or sc < best[0]:
            best = (sc, d, e1, e2, P)
    _, d, e1, e2, P = best
    w2, v2 = np.linalg.eigh(P.T @ P)
    ax = v2[:, 1]                                     # größte Ausdehnung im Bild -> x
    x = ax[0] * e1 + ax[1] * e2
    y = np.cross(d, x)
    R = np.array([x, y, d])
    if np.linalg.det(R) < 0:
        R[1] *= -1
    return R


# ---------------------------------------------------------------------------- Ausrichten
SCREEN_DIRS = {"down": (0.0, -1.0, 0.0), "up": (0.0, 1.0, 0.0), "right": (1.0, 0.0, 0.0), "left": (-1.0, 0.0, 0.0)}


def rot_between(u, t):
    """Kleinste Drehung, die den Einheitsvektor u auf t bringt (Rodrigues)."""
    u = np.asarray(u, float) / np.linalg.norm(u)
    t = np.asarray(t, float) / np.linalg.norm(t)
    v = np.cross(u, t)
    c = float(np.dot(u, t))
    if np.linalg.norm(v) < 1e-12:
        if c > 0:
            return np.eye(3)
        a = np.cross(u, [1.0, 0, 0] if abs(u[0]) < 0.9 else [0, 1.0, 0])   # 180° um eine Senkrechte
        a /= np.linalg.norm(a)
        return 2 * np.outer(a, a) - np.eye(3)
    K = np.array([[0, -v[2], v[1]], [v[2], 0, -v[0]], [-v[1], v[0], 0]])
    return np.eye(3) + K + K @ K / (1 + c)


def align_bond(R, mol, keys, direction="down"):
    """Bindung (erstes -> zweites Atom) im Bild in die Richtung drehen (unten/oben/rechts/links); die Bindung liegt
    danach in der Bildebene. Kleinste nötige Drehung, die Blickrichtung ändert sich so wenig wie möglich."""
    a, b = (mol.atoms[mol.find_key(k)]["xyz"] for k in keys[:2])
    v = np.asarray(R) @ (b - a)
    return orthonormalize(rot_between(v, SCREEN_DIRS[direction]) @ np.asarray(R))


def align_plane(R, mol, keys):
    """Ebene durch die gewählten Atome (≥ 3) in die Bildebene: Blick senkrecht darauf, von der Seite, die der
    bisherigen Ansicht näher ist."""
    pts = np.array([mol.atoms[mol.find_key(k)]["xyz"] for k in keys])
    c = pts - pts.mean(axis=0)
    n = np.linalg.svd(c)[2][-1]
    nv = np.asarray(R) @ n
    target = (0.0, 0.0, 1.0 if nv[2] >= 0 else -1.0)
    return orthonormalize(rot_between(nv, target) @ np.asarray(R))


def align_to_template(mol, template_path, frag=0):
    """Ausrichtung wie in einer ChemDraw-Zeichnung (CDXML): Atome der Zeichnung über das Bindungsnetz den Atomen der
    Struktur zuordnen (nur Elemente und Verknüpfung, ohne Bindungsordnung und H), dann die Drehung, deren Projektion
    am besten auf die 2D-Koordinaten passt (Kabsch). -> (R, dict(n, n_template, rmsd, ambiguous)); ValueError, wenn
    keine eindeutige Zuordnung gelingt."""
    from rdkit import Chem
    from rdkit.Chem import rdFMCS
    ms = [m for m in Chem.MolsFromCDXMLFile(str(template_path), sanitize=False, removeHs=False) if m is not None]
    if not ms:
        raise ValueError("keine Struktur in der Datei")
    tmol = max(ms, key=lambda m: sum(a.GetAtomicNum() > 1 for a in m.GetAtoms()))

    def heavy_graph(n_atoms, elem, bonds):
        rw = Chem.RWMol()
        for z_ in elem:
            a = Chem.Atom(int(z_))
            a.SetNoImplicit(True)
            rw.AddAtom(a)
        for i, j in bonds:
            if rw.GetBondBetweenAtoms(i, j) is None:
                rw.AddBond(i, j, Chem.BondType.SINGLE)
        m = rw.GetMol()
        m.UpdatePropertyCache(strict=False)
        Chem.FastFindRings(m)
        return m
    t_idx = [a.GetIdx() for a in tmol.GetAtoms() if a.GetAtomicNum() > 1]
    t_map = {k: i for i, k in enumerate(t_idx)}
    conf = tmol.GetConformer()
    t_xy = np.array([[conf.GetAtomPosition(k).x, conf.GetAtomPosition(k).y] for k in t_idx])
    t_bonds = [(t_map[b.GetBeginAtomIdx()], t_map[b.GetEndAtomIdx()]) for b in tmol.GetBonds()
               if b.GetBeginAtomIdx() in t_map and b.GetEndAtomIdx() in t_map]
    tg = heavy_graph(len(t_idx), [tmol.GetAtomWithIdx(k).GetAtomicNum() for k in t_idx], t_bonds)
    pt = Chem.GetPeriodicTable()
    c_idx = [k for k, a in enumerate(mol.atoms) if a.get("frag") == frag and a["el"] not in ("H", "D")]
    c_map = {k: i for i, k in enumerate(c_idx)}
    c_xyz = np.array([mol.atoms[k]["xyz"] for k in c_idx])
    c_bonds = [(c_map[a], c_map[b]) for a, b in mol.bonds if a in c_map and b in c_map]
    cg = heavy_graph(len(c_idx), [pt.GetAtomicNumber(mol.atoms[k]["el"]) for k in c_idx], c_bonds)
    matches = []
    pairs = cg.GetSubstructMatches(tg, uniquify=False, maxMatches=2000, useChirality=False)
    if pairs:
        matches = [list(zip(range(len(t_idx)), m)) for m in pairs]
    else:                                              # Zeichnung und Struktur weichen ab: größter gemeinsamer Teil
        res = rdFMCS.FindMCS([tg, cg], timeout=5, atomCompare=rdFMCS.AtomCompare.CompareElements,
                             bondCompare=rdFMCS.BondCompare.CompareAny, ringMatchesRingOnly=False)
        if not res.numAtoms or res.numAtoms < max(3, 0.6 * len(t_idx)):
            raise ValueError("Zeichnung passt nicht zur Struktur (weniger als 60 % der Atome zuzuordnen)")
        q = Chem.MolFromSmarts(res.smartsString)
        tm_ = tg.GetSubstructMatches(q, uniquify=False, maxMatches=50)
        cm_ = cg.GetSubstructMatches(q, uniquify=False, maxMatches=2000)
        if not tm_ or not cm_:
            raise ValueError("Zuordnung fehlgeschlagen")
        matches = [list(zip(tm_[0], m)) for m in cm_]
    lt = np.mean([np.linalg.norm(t_xy[i] - t_xy[j]) for i, j in t_bonds]) if t_bonds else 1.0
    lc = np.mean([np.linalg.norm(c_xyz[i] - c_xyz[j]) for i, j in c_bonds]) if c_bonds else 1.0
    best = []
    for m in matches:
        ti = [a for a, _ in m]
        ci = [b for _, b in m]
        P = c_xyz[ci] - c_xyz[ci].mean(axis=0)
        Q = np.column_stack([t_xy[ti] * lc / lt, np.zeros(len(ti))])
        Q -= Q.mean(axis=0)
        H = P.T @ Q
        U, S, Vt = np.linalg.svd(H)
        d = np.sign(np.linalg.det(Vt.T @ U.T)) or 1.0
        Rk = Vt.T @ np.diag([1, 1, d]) @ U.T              # Rk @ P ≈ Q
        proj = (Rk @ P.T).T[:, :2]
        rmsd = float(np.sqrt(np.mean(np.sum((proj - Q[:, :2]) ** 2, axis=1))))
        best.append((rmsd, Rk))
    best.sort(key=lambda x: x[0])
    rmsd, Rk = best[0]
    # mehrdeutig: fast gleich gute Zuordnung mit deutlich anderer Ansicht (z. B. symmetrische Moleküle)
    amb = any(r2 - rmsd < 0.05 and np.degrees(np.arccos(np.clip((np.trace(R2.T @ Rk) - 1) / 2, -1, 1))) > 20
              for r2, R2 in best[1:50])
    return orthonormalize(Rk), dict(n=len(matches[0]), n_template=len(t_idx), rmsd=rmsd, ambiguous=amb)


def rot(axis, ang):
    c, s = math.cos(ang), math.sin(ang)
    if axis == "x":
        return np.array([[1, 0, 0], [0, c, -s], [0, s, c]])
    if axis == "y":
        return np.array([[c, 0, s], [0, 1, 0], [-s, 0, c]])
    return np.array([[c, -s, 0], [s, c, 0], [0, 0, 1]])


def orthonormalize(R):
    u, _, vt = np.linalg.svd(np.asarray(R, float))
    Rn = u @ vt
    if np.linalg.det(Rn) < 0:
        Rn = -Rn
    return Rn


# ---------------------------------------------------------------------------- Projektion
def project(mol: Molecule, R, k, show, h_radius=0.11, n_seg=48):
    """Szene für die Abbildung (Koordinaten in Å, Bildschirm: x rechts, y oben, z zum Betrachter).
    -> dict(atoms=[dict(k, x, y, z, w, h, ang, iso, curves=[(xs, ys)], el, r_pick)], bonds=[(k1, k2, z)])"""
    R = np.asarray(R, float)
    center = np.array([a["xyz"] for a in mol.atoms]).mean(axis=0) if mol.atoms else np.zeros(3)
    out_atoms, pos = [], {}
    t = np.linspace(0, 2 * np.pi, n_seg + 1)
    ct, st = np.cos(t), np.sin(t)
    for i, a in enumerate(mol.atoms):
        if i not in show:
            continue
        p = R @ (a["xyz"] - center)
        pos[i] = p
        item = dict(k=i, x=float(p[0]), y=float(p[1]), z=float(p[2]), el=a["el"], curves=[], npd=False)
        U = a["U"]
        if a["el"] in ("H", "D"):
            item.update(w=2 * h_radius, h=2 * h_radius, ang=0.0, iso=True, kind="H")
        elif U is not None:
            Ur = R @ U @ R.T
            w2, v2 = np.linalg.eigh(Ur[:2, :2])
            ev, evec = np.linalg.eigh(U)
            if (ev <= 0).any() or (w2 <= 0).any():
                r = k * math.sqrt(max(a["uiso"], 1e-4))
                item.update(w=2 * r, h=2 * r, ang=0.0, iso=True, kind="iso", npd=True)
            else:
                item.update(w=2 * k * math.sqrt(w2[1]), h=2 * k * math.sqrt(w2[0]),
                            ang=float(math.degrees(math.atan2(v2[1, 1], v2[0, 1]))), iso=False, kind="aniso")
                axes = [R @ (evec[:, m] * k * math.sqrt(ev[m])) for m in range(3)]
                for m1, m2 in ((0, 1), (1, 2), (0, 2)):
                    off = np.outer(ct, axes[m1]) + np.outer(st, axes[m2])
                    front = off[:, 2] >= -1e-9
                    xs = np.where(front, p[0] + off[:, 0], np.nan)
                    ys = np.where(front, p[1] + off[:, 1], np.nan)
                    item["curves"].append((xs, ys))
        else:
            r = k * math.sqrt(max(a["uiso"], 1e-4))
            item.update(w=2 * r, h=2 * r, ang=0.0, iso=True, kind="iso")
        out_atoms.append(item)
    bonds = []
    for a, b in mol.bonds:
        if a in pos and b in pos:
            bonds.append((a, b, float((pos[a][2] + pos[b][2]) / 2)))
    return dict(atoms=out_atoms, bonds=bonds, center=center)


def pick(scene, x, y, tol=0.0):
    """Vorderstes Atom, dessen projizierter Umriss (x, y) enthält (Å, Toleranz in Å) -> Atomindex oder None."""
    best, bz = None, -1e9
    for it in scene["atoms"]:
        dx, dy = x - it["x"], y - it["y"]
        th = math.radians(it["ang"])
        u = dx * math.cos(th) + dy * math.sin(th)
        v = -dx * math.sin(th) + dy * math.cos(th)
        a, b = it["w"] / 2 + tol, it["h"] / 2 + tol
        if (u / a) ** 2 + (v / b) ** 2 <= 1 and it["z"] > bz:
            best, bz = it["k"], it["z"]
    return best


# ---------------------------------------------------------------------------- Tabellen (Rich Text)
def runs(markup):
    """Mini-Auszeichnung -> [(Text, Stil)]: *kursiv*, _{tief}, ^{hoch}, ~{Überstrich}."""
    out, ital, i, n = [], False, 0, len(markup)
    buf = ""

    def flush(style_extra=""):
        nonlocal buf
        if buf:
            out.append((buf, ("i" if ital else "") + style_extra))
            buf = ""
    while i < n:
        c = markup[i]
        if c == "*":
            flush()
            ital = not ital
            i += 1
        elif c in "_^~" and i + 1 < n and markup[i + 1] == "{":
            flush()
            j = markup.index("}", i)
            inner = markup[i + 2:j]
            st = {"_": "s", "^": "p", "~": "o"}[c]
            for t, s in runs(inner):
                out.append((t, s + st + ("i" if ital and "i" not in s else "")))
            i = j + 1
        else:
            buf += c
            i += 1
    flush()
    return out


def to_html(markup):
    h = []
    for t, s in runs(markup):
        t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if "o" in s:
            t = f'<span style="text-decoration: overline">{t}</span>'
        if "i" in s:
            t = f"<i>{t}</i>"
        if "s" in s:
            t = f"<sub>{t}</sub>"
        if "p" in s:
            t = f"<sup>{t}</sup>"
        h.append(t)
    return "".join(h)


_SUB_ALL = str.maketrans("0123456789+−-=()aehijklmnoprstuvx", "₀₁₂₃₄₅₆₇₈₉₊₋₋₌₍₎ₐₑₕᵢⱼₖₗₘₙₒₚᵣₛₜᵤᵥₓ")


def to_plain(markup):
    """Für CSV/Text: Tief-/Hochstellung als Unicode, falls möglich (sonst R_int), Überstrich kombinierend."""
    out = []
    for t, s in runs(markup):
        if "s" in s:
            t2 = t.translate(_SUB_ALL)
            t = t2 if all(c != d or c in " ,." for c, d in zip(t, t2)) else t
        if "p" in s:
            t2 = t.translate(SUP)
            t = t2 if all(c != d for c, d in zip(t, t2)) else "^" + t
        if "o" in s:
            t = "".join(ch + "\u0305" for ch in t)
        out.append(t)
    return "".join(out)


def to_rich_xlsx(markup):
    """CellRichText für openpyxl (kursiv, tief-/hochgestellt; Überstrich als kombinierendes Zeichen)."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    parts = []
    for t, s in runs(markup):
        if "o" in s:
            t = "".join(ch + "\u0305" for ch in t)
        if not ("i" in s or "s" in s or "p" in s):
            parts.append(t)
            continue
        f = InlineFont(i=("i" in s) or None, vertAlign="subscript" if "s" in s else ("superscript" if "p" in s else None))
        parts.append(TextBlock(f, t))
    if all(isinstance(p, str) for p in parts):
        return "".join(parts)
    return CellRichText(parts)


def space_group_markup(hm):
    """'P 21/c' -> '*P*2_{1}/*c*', 'P -1' -> '*P*~{1}', 'Pca21' -> '*Pca*2_{1}'."""
    s = str(hm or "").strip().strip("'\"")
    if not s:
        return ""
    s = re.sub(r"\s*:\s*\w+$", "", s)                # Ursprungswahl ':1'/':H' weglassen
    toks = s.split() if " " in s else None
    if toks and len(toks) == 4 and sum(t == "1" for t in toks[1:]) == 2:    # 'P 1 21/c 1' -> P2₁/c
        toks = [toks[0]] + [t for t in toks[1:] if t != "1"]
    out = []

    def screw(d1, d2):
        return int(d2) < int(d1) and d1 in "2346"

    def fmt_token(tok, first):
        res, i = "", 0
        while i < len(tok):
            c = tok[i]
            if c.isalpha():
                res += f"*{c}*"
                i += 1
            elif c == "-" and i + 1 < len(tok) and tok[i + 1].isdigit():
                res += "~{" + tok[i + 1] + "}"
                i += 2
            elif c.isdigit():
                if i + 1 < len(tok) and tok[i + 1].isdigit() and screw(c, tok[i + 1]):
                    res += c + "_{" + tok[i + 1] + "}"
                    i += 2
                else:
                    res += c
                    i += 1
            else:
                res += c
                i += 1
        return res
    if toks:
        out = [fmt_token(t, k == 0) for k, t in enumerate(toks)]
        txt = "".join(out)
    else:
        txt = fmt_token(s, True)
    return txt.replace("**", "")


def formula_markup(sum_formula):
    """'C20 H30 N2 Si' -> 'C_{20}H_{30}N_{2}Si'."""
    s = str(sum_formula or "").strip().strip("'\"")
    out = []
    for tok in s.split():
        m = re.match(r"^([A-Z][a-z]?)(\d*\.?\d*)([+-]?\d*[+-]?)$", tok)
        if not m:
            out.append(tok)
            continue
        el, n, ch = m.groups()
        out.append(el + (f"_{{{n}}}" if n and n != "1" else "") + (f"^{{{ch}}}" if ch else ""))
    return "".join(out)


RADIATION = [(0.71073, 0.002, "MoK_{*α*}"), (1.54178, 0.0012, "CuK_{*α*}"), (0.56086, 0.002, "AgK_{*α*}"),
             (1.34139, 0.002, "GaK_{*α*}")]


def radiation_markup(b: Block):
    lam = b.get("_diffrn_radiation_wavelength")
    v = num(lam)
    if v is None:
        return ""
    name = None
    for ref, tol, nm in RADIATION:
        if abs(v - ref) < tol:
            name = nm
    if name is None:
        typ = (b.get("_diffrn_radiation_type") or "").replace("\\a", "α").replace("\\b", "β").replace(" ", "")
        name = typ or "λ"
    return f"{name} (*λ* = {lam})"


def _dec(s):
    s = str(s)
    return len(s.split(".")[1]) if "." in s else 0


CRYSTAL_ROWS = [                                   # Beschriftung wie in den SI der Arbeitsgruppe
    ("ref", "Internal reference"),
    ("ccdc", "CCDC number"),
    ("formula", "Formula"),
    ("fw", "Formula weight"),
    ("temp", "Temperature / K"),
    ("system", "Crystal system"),
    ("sg", "Space group"),
    ("a", "*a* / Å"), ("b", "*b* / Å"), ("c", "*c* / Å"),
    ("alpha", "*α* / °"), ("beta", "*β* / °"), ("gamma", "*γ* / °"),
    ("vol", "Volume / Å^{3}"),
    ("z", "*Z*"),
    ("rho", "Density (calc.) g / cm^{3}"),
    ("mu", "Absorption coefficient / mm^{-1}"),
    ("f000", "*F*(000)"),
    ("size", "Crystal size / mm"),
    ("rad", "Radiation"),
    ("theta", "2*θ* range for data collection / °"),
    ("index", "Index ranges"),
    ("refl", "Reflections collected"),
    ("indep", "Independent reflections"),
    ("method", "Refinement method"),
    ("drp", "Data / restraints / parameters"),
    ("gof", "Goodness-of-fit on *F*^{2}"),
    ("rgt", "Final *R* indices [*I*>2*σ*(I)]"),
    ("rall", "*R* indices (all data)"),
    ("diff", "Largest diff. peak and hole/e. Å^{-3}"),
]


def crystal_values(b: Block, ref="", ccdc=""):
    """Spalte der Kristalldatentabelle: Schlüssel -> Markup (Zahlen exakt wie in der CIF, echtes Minus)."""
    g = b.get
    v = {}
    v["ref"] = ref or ""
    cc = g("_database_code_depnum_ccdc_archive") or ""
    cc = re.sub(r"(?i)^\s*ccdc\s*", "", cc.strip())
    v["ccdc"] = ccdc or cc
    v["formula"] = formula_markup(g("_chemical_formula_sum"))
    v["fw"] = g("_chemical_formula_weight") or ""
    v["temp"] = g("_diffrn_ambient_temperature", "_cell_measurement_temperature") or ""
    v["system"] = g("_space_group_crystal_system", "_symmetry_cell_setting") or ""
    v["sg"] = space_group_markup(g("_space_group_name_h-m_alt", "_symmetry_space_group_name_h-m"))
    for k, tag in (("a", "_cell_length_a"), ("b", "_cell_length_b"), ("c", "_cell_length_c"),
                   ("alpha", "_cell_angle_alpha"), ("beta", "_cell_angle_beta"), ("gamma", "_cell_angle_gamma"),
                   ("vol", "_cell_volume"), ("z", "_cell_formula_units_z"), ("rho", "_exptl_crystal_density_diffrn"),
                   ("mu", "_exptl_absorpt_coefficient_mu"), ("f000", "_exptl_crystal_f_000"),
                   ("refl", "_diffrn_reflns_number"), ("gof", "_refine_ls_goodness_of_fit_ref")):
        v[k] = g(tag) or ""
    sizes = [g(f"_exptl_crystal_size_{x}") for x in ("max", "mid", "min")]
    v["size"] = " × ".join(s for s in sizes if s) if all(sizes) else ""
    v["rad"] = radiation_markup(b)
    t0, t1 = g("_diffrn_reflns_theta_min"), g("_diffrn_reflns_theta_max")
    if t0 and t1 and num(t0) is not None and num(t1) is not None:
        d0, d1 = _dec(esd_parts(t0)[0]), _dec(esd_parts(t1)[0])
        v["theta"] = f"{2 * num(t0):.{d0}f} to {2 * num(t1):.{d1}f}"
    else:
        v["theta"] = ""
    lim = [g(f"_diffrn_reflns_limit_{h}_{m}") for h in "hkl" for m in ("min", "max")]
    if all(lim):
        v["index"] = ", ".join(f"{minus(lim[2 * i])} ≤ *{h}* ≤ {minus(lim[2 * i + 1])}" for i, h in enumerate("hkl"))
    else:
        v["index"] = ""
    tot, rint, rsig = g("_reflns_number_total"), g("_diffrn_reflns_av_r_equivalents"), \
        g("_diffrn_reflns_av_unetI/netI", "_diffrn_reflns_av_sigmaI/netI")
    if tot:
        extra = []
        if rint:
            extra.append(f"*R*_{{int}} = {rint}")
        if rsig:
            extra.append(f"*R*_{{sigma}} = {rsig}")
        v["indep"] = tot + (f" [{', '.join(extra)}]" if extra else "")
    else:
        v["indep"] = ""
    coef = (g("_refine_ls_structure_factor_coef") or "").lower()
    mat = (g("_refine_ls_matrix_type") or "full").lower()
    on = "*F*^{2}" if coef.startswith("fsq") else ("*F*" if coef else "*F*^{2}")
    v["method"] = ("Full-matrix" if mat == "full" else mat.capitalize()) + f" least-squares on {on}"
    nref = g("_refine_ls_number_reflns", "_reflns_number_total")
    nres, npar = g("_refine_ls_number_restraints"), g("_refine_ls_number_parameters")
    v["drp"] = f"{nref}/{nres}/{npar}" if nref and nres and npar else ""
    r1, wr = g("_refine_ls_r_factor_gt"), g("_refine_ls_wr_factor_gt")
    v["rgt"] = f"*R*_{{1}} = {r1}, w*R*_{{2}} = {wr}" if r1 and wr else ""
    r1a, wra = g("_refine_ls_r_factor_all"), g("_refine_ls_wr_factor_ref")
    v["rall"] = f"*R*_{{1}} = {r1a}, w*R*_{{2}} = {wra}" if r1a and wra else ""
    dmax, dmin = g("_refine_diff_density_max"), g("_refine_diff_density_min")
    if dmax and dmin:
        hole = dmin.strip()
        hole = "\u2212" + hole.lstrip("-\u2212") if not hole.startswith("+") else hole
        v["diff"] = f"{dmax}/{hole}"
    else:
        v["diff"] = ""
    return v


def write_table_csv(path, header, rows):
    """rows: [[markup, …]] -> CSV (UTF-8 mit BOM, Unicode-Tief-/Hochstellung)."""
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow([to_plain(h) for h in header])
        for r in rows:
            w.writerow([to_plain(c) for c in r])
    return path


def table_html(header, rows, title=None, notes=()):
    th = "".join(f'<th style="border-bottom:1px solid #000; padding:2px 8px; text-align:left">{to_html(h)}</th>'
                 for h in header)
    body = "".join("<tr>" + "".join(f'<td style="padding:1px 8px">{to_html(c)}</td>' for c in r) + "</tr>"
                   for r in rows)
    cap = f"<p><b>{to_html(title)}</b></p>" if title else ""
    foot = "".join(f'<p style="font-size:small">{to_html(n)}</p>' for n in notes)
    return f'{cap}<table style="border-collapse:collapse">{"<tr>" + th + "</tr>"}{body}</table>{foot}'


def table_plain(header, rows):
    return "\n".join("\t".join(to_plain(c) for c in r) for r in [header] + rows)


# ---------------------------------------------------------------------------- Abbildung (matplotlib)
import matplotlib  # noqa: E402
import matplotlib.patheffects as PE  # noqa: E402
from matplotlib.collections import LineCollection, PolyCollection  # noqa: E402
from matplotlib.colors import to_rgb  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from matplotlib.patches import Ellipse  # noqa: E402
from matplotlib.transforms import Affine2D  # noqa: E402

SEL_COLOR = "#FF8C00"
LIGHT = np.array([-0.55, 0.65]) / math.hypot(0.55, 0.65)    # Licht von oben links (Bildschirm), wie Mercury
_T48 = np.linspace(0, 2 * np.pi, 49)


def _ellipse_verts(cx, cy, w, h, ang, n=None):
    t = _T48 if n is None else np.linspace(0, 2 * np.pi, n + 1)
    th = math.radians(ang)
    x, y = w / 2 * np.cos(t), h / 2 * np.sin(t)
    return np.column_stack([cx + x * math.cos(th) - y * math.sin(th), cy + x * math.sin(th) + y * math.cos(th)])


def shade_colors(color, n):
    """Mercury-ähnliche, matte Schattierung: Rand etwas dunkler, Grundfarbe dominiert, schwacher Glanz oben links."""
    base = np.array(to_rgb(color))
    out = []
    for t in np.linspace(0, 1, n):
        if t < 0.3:
            c = base * (0.62 + 0.38 * t / 0.3)
        else:
            f = (t - 0.3) / 0.7
            c = base + (1 - base) * 0.45 * f ** 1.6
        out.append(tuple(np.clip(c, 0, 1)))
    return out


STYLES = ("mercury", "flat", "ballstick", "depth", "cartoon", "iboview")
BALL_STYLES = ("ballstick", "depth", "cartoon", "iboview")   # Kugeln nach Element statt Ellipsoiden, Stäbe statt Linien
XRAY_STYLES = ("mercury", "flat")                 # Auswahl im Tab X-Ray (Nutzerwunsch: nur Mercury und Olex2)
DFT_STYLES = ("mercury", "flat", "depth", "cartoon", "iboview")   # Auswahl im Tab DFT (alter Glanz-Stil „ballstick“ entfernt)


def allowed_style(style, allowed, fallback):
    """Stil alter Projekte, der im Tab nicht mehr wählbar ist -> Ersatz (zeichnen kann XrayFigure alle)."""
    return style if style in allowed else fallback
DEPTH_A = 8.0                             # Tiefe (Å), über die der Stil „depth“ voll aufhellt
DEPTH_MAX = 0.35                          # … bis 35 % zum Weiß
FOG_STYLES = ()                           # Stile mit Aufhellung nach hinten (ORCA-Vorlage hat keine)
# Phong-Beleuchtung (Grundanteil ka + diffus kd·(n·L) + Glanz ks·(n·H)^p), angepasst an eine IboView-Abbildung
# des Nutzers (Fit an die Kohlenstoffkugel: Licht oben links 40° über der Bildebene, Grundanteil 0.31, diffus 0.48 bei
# Grau 0.55; Glanzpunkt klein und rein weiß). „depth“: matt, weicher, kaum Glanz. Schichten = Höhenlinien der
# Helligkeit (Vektor in PDF/SVG): Kugeln über contourpy, Stäbe als Zylinder je nach Lage zum Licht.
SHADE = {
    "ballstick": dict(el=40.0, ka=0.567, kd=0.874, ks=0.32, p=160, nd=30, nr=16, outline=True),
    # Kugel-Stab wie im ORCA-Handbuch (an dessen Abbildung gemessen: Licht oben links 47° über der Bildebene,
    # Mitte/Lichtseite gesättigt, Schattenseite bis ~50 %, weicher Glanzpunkt bei halbem Radius), ohne Umriss
    "depth": dict(el=47.0, ka=0.5, kd=0.97, ks=0.3, p=150, nd=24, nr=14, outline=False, soft=True, spot=0.5),
    # Strichzeichnung wie in älteren JOC-Abbildungen (Nutzervorlage): weiße Kugeln mit Umriss, hohle Stäbe
    "cartoon": dict(el=45.0, ka=1.0, kd=0.0, ks=0.0, p=1, nd=1, nr=1, outline=True),
}


def _ibo_shade():
    """Atome im Orbitalbild wie IboView (drei Lichter, orbitals.ibo_light): Schichten über die diffuse Helligkeit
    D (normiert auf ihr Maximum) und den Glanz S > 0; ka/kd/ks so, dass shade_color die Werte von D/S liefert.
    Der schwache dunkle Ring um den Glanzpunkt (S < 0) entfällt."""
    U, V = np.meshgrid(np.linspace(-1, 1, 181), np.linspace(-1, 1, 181))
    W = np.sqrt(np.clip(1 - U ** 2 - V ** 2, 0, 1))
    D, S = ORB.ibo_light(np.stack([U, V, W], -1)[(U ** 2 + V ** 2) <= 1])
    i = int(np.argmax(S))
    return dict(el=45.0, ka=0.0, kd=float(D.max()), ks=float(S.max()), p=1, nd=24, nr=14, outline=False,
                soft=True, ibo=True, d_spec=float(D[i]))


SHADE["iboview"] = _ibo_shade()
SPEC_LEVELS = tuple(np.linspace(0.04, 0.96, 12))           # weich auslaufender Glanzpunkt


def _light3(style):
    """Lichtrichtung L und Halbvektor H (3D, Bildschirm: x rechts, y oben, z zum Betrachter) für einen Stil."""
    e = math.radians(SHADE[style]["el"])
    L = np.array([LIGHT[0] * math.cos(e), LIGHT[1] * math.cos(e), math.sin(e)])
    H = L + np.array([0.0, 0.0, 1.0])
    hr = SHADE[style].get("spot")                 # Glanzpunkt bei diesem Anteil des Radius (eigenes Glanzlicht)
    if hr is not None:
        H = np.array([LIGHT[0] * hr, LIGHT[1] * hr, math.sqrt(1 - hr * hr)])
    return L, H / np.linalg.norm(H)


def _fields(style, nx, ny, nz):
    """Helligkeitsfelder für die Schichten je Normale: (diffus 0…1, Glanz 0…1) – Phong (Licht L, Halbvektor H) oder
    wie IboView (drei Lichter, auf das Maximum normiert)."""
    sp = SHADE[style]
    if sp.get("ibo"):
        D, S = ORB.ibo_light(np.stack(np.broadcast_arrays(nx, ny, nz), -1))
        return D / sp["kd"], np.clip(S, 0, None) / sp["ks"]
    L, H = _light3(style)
    ndl = np.clip(nx * L[0] + ny * L[1] + nz * L[2], 0, None)
    ndh = np.clip(nx * H[0] + ny * H[1] + nz * H[2], 0, None)
    return ndl, ndh ** sp["p"]


def spec_val(sp, lev):
    """Helligkeitszuschlag einer Glanzschicht: weich (soft, ORCA: auslaufender Lichtfleck) oder als Scheibe."""
    return sp["ks"] * lev if sp.get("soft") else sp["ks"] * (lev + (1 - lev) * 0.5)


_SPHERE_CACHE = {}


def sphere_unit_layers(style):
    """Schichten einer Kugel vom Radius 1: [(Polygon N×2, Art, Wert)], Art "d" = Helligkeit, "s" = Glanz.
    Erste Schicht = ganze Scheibe; Regionen gleicher oder größerer Helligkeit als Höhenlinien (contourpy)."""
    if style in _SPHERE_CACHE:
        return _SPHERE_CACHE[style]
    import contourpy
    sp = SHADE[style]
    g = np.linspace(-1.25, 1.25, 241)
    X, Y = np.meshgrid(g, g)
    rr = np.hypot(X, Y)
    s_ = np.where(rr > 1, 1 / np.maximum(rr, 1e-9), 1.0)       # außerhalb: Wert am Rand (für sauberes Abschneiden)
    u, v = X * s_, Y * s_
    w = np.sqrt(np.clip(1 - u ** 2 - v ** 2, 0, 1))
    ndl, ndh = _fields(style, u, v, w)
    t = np.linspace(0, 2 * np.pi, 97)
    out = [(np.column_stack([np.cos(t), np.sin(t)]), "d", sp["ka"] + sp["kd"] * 0.5 / sp["nd"])]

    def regions(field, lev):
        pts, offs = contourpy.contour_generator(X, Y, field, fill_type=contourpy.FillType.OuterOffset).filled(lev, 1e9)
        res = []
        for P, O in zip(pts, offs):
            P = P[O[0]:O[1]]
            n = np.hypot(P[:, 0], P[:, 1])
            res.append(np.where((n > 1)[:, None], P / np.maximum(n, 1e-9)[:, None], P))   # auf die Scheibe
        return res
    for j in range(1, sp["nd"]):
        for P in regions(ndl, j / sp["nd"]):
            out.append((P, "d", sp["ka"] + sp["kd"] * (j + 0.5) / sp["nd"]))
    for lev in (SPEC_LEVELS if sp["ks"] > 0 else ()):
        val = spec_val(sp, lev)
        for P in regions(ndh, lev):
            out.append((P, "s", val))
    _SPHERE_CACHE[style] = out
    return out


_SPHERE_DRAG = {}
DRAG_TOL_PX = 0.25                                         # Ziehen: Kugelumrisse höchstens so weit (Pixel) vereinfacht
DRAG_TOLS = (0.04, 0.02, 0.01, 0.005, 0.0025, 0.00125)      # Stufen (Anteil des Radius), gröbste passende gewinnt


def simplify_polyline(P, tol):
    """Douglas-Peucker: Punkte weglassen, solange die Linie höchstens tol (Einheiten von P) abweicht."""
    keep = np.zeros(len(P), bool)
    keep[0] = keep[-1] = True
    stack = [(0, len(P) - 1)]
    while stack:
        i, j = stack.pop()
        if j <= i + 1:
            continue
        a, d, Q = P[i], P[j] - P[i], P[i + 1:j]
        L = math.hypot(d[0], d[1])
        dist = (np.abs(d[0] * (Q[:, 1] - a[1]) - d[1] * (Q[:, 0] - a[0])) / L if L > 0
                else np.hypot(Q[:, 0] - a[0], Q[:, 1] - a[1]))
        m = int(dist.argmax())
        if dist[m] > tol:
            keep[i + 1 + m] = True
            stack += [(i, i + 1 + m), (i + 1 + m, j)]
    return P[keep]


def sphere_layer_verts(style, tol=None):
    """Polygone der Einheitskugel (Schichten + Umriss) für die PolyCollection; tol: vereinfacht (nur beim Ziehen,
    die Contour-Polygone haben je Schicht Hunderte Punkte – Abbildung und Export immer mit tol=None)."""
    key = (style, tol)
    if key not in _SPHERE_DRAG:
        lay = [P for P, _, _ in sphere_unit_layers(style)]
        if tol is not None:
            lay = [simplify_polyline(P, tol) for P in lay]
        _SPHERE_DRAG[key] = lay + [_ellipse_verts(0.0, 0.0, 2.0, 2.0, 0.0)]
    return _SPHERE_DRAG[key]


_ROD_U = np.linspace(-1, 1, 241)


def rod_layouts(style, dx, dy):
    """Stäbe als Zylinder (Achse in der Bildebene, Richtungen dx, dy: Arrays, ein Eintrag je Stab), vektorisiert ->
    (Breitenanteil, Versatz in Stabbreiten zur Lichtseite), je Array Stab × Schicht in der Reihenfolge von
    rod_levels(style); leere Schicht -> 0, 0. Exakt dieselbe Rechnung wie früher je Stab (kein Raster über Winkel)."""
    dx, dy = np.asarray(dx, float), np.asarray(dy, float)
    n_ = np.hypot(dx, dy)
    n_ = np.where(n_ != 0, n_, 1.0)
    nx, ny = -dy / n_, dx / n_
    flip = np.where(nx * LIGHT[0] + ny * LIGHT[1] < 0, -1.0, 1.0)
    nx, ny = nx * flip, ny * flip
    u = _ROD_U
    w = np.sqrt(1 - u ** 2)
    ndl, ndh = _fields(style, u * nx[:, None], u * ny[:, None], w)
    lev = rod_levels(style)
    td = np.array([t for k, t, _ in lev if k == "d"])
    ts = np.array([t for k, t, _ in lev if k == "s"])
    m = np.concatenate([ndl[:, None, :] >= td[None, :, None], ndh[:, None, :] >= ts[None, :, None]], axis=1)
    hit = m.any(axis=2)
    u1 = u[m.argmax(axis=2)]                                     # erste und letzte Stützstelle der Schicht
    u2 = u[len(u) - 1 - m[:, :, ::-1].argmax(axis=2)]
    return np.where(hit, (u2 - u1) / 2, 0.0), np.where(hit, (u1 + u2) / 4, 0.0)


def rod_layout(style, dx, dy):
    """Ein Stab (siehe rod_layouts): je Schicht (Breitenanteil, Versatz in Stabbreiten zur Lichtseite)."""
    fr, off = rod_layouts(style, [dx], [dy])
    return [(float(f), float(o)) for f, o in zip(fr[0], off[0])]


def rod_levels(style):
    """Schichten eines Stabs: [(Art, Schwelle, Wert)] – gleiche Helligkeitsstufen wie die Kugeln, gröber."""
    sp = SHADE[style]
    out = [("d", j / sp["nr"], sp["ka"] + sp["kd"] * (j + 0.5) / sp["nr"]) for j in range(sp["nr"])]
    out[0] = ("d", -1.0, out[0][2])                              # ganze Breite
    return out + [("s", lev, spec_val(sp, lev)) for lev in (SPEC_LEVELS if sp["ks"] > 0 else ())]


def shade_color(base, kind, val, style, fog=0.0):
    """Farbe einer Schicht: Grundfarbe × Helligkeit, Glanz = Helligkeit am Glanzpunkt + Weißanteil; dann Tiefe."""
    b = np.array(to_rgb(base))
    if kind == "d":
        c = b * val
    elif SHADE[style].get("ibo"):
        c = b * SHADE[style]["d_spec"] + val
    else:
        sp = SHADE[style]
        L, H = _light3(style)
        c = b * (sp["ka"] + sp["kd"] * float(np.dot(L, H))) + val
    c = np.clip(c, 0, 1)
    return tuple(np.clip(c + (1 - c) * fog, 0, 1))


ROD_W = {"ballstick": 0.2, "depth": 0.2, "cartoon": 0.17,      # Stabdicke (Å) wie in den Vorlagen
         "iboview": 2 * ORB.IBO_ROD_R * ORB.BOHR}
BALL_R = {"ballstick": 0.45, "depth": 0.45, "cartoon": 0.6}     # Kugelradius = Faktor × kovalenter Radius
BALL_RMIN = {"cartoon": 0.24, "depth": 0.21}                    # kleinster Radius (H), sonst 0.2 Å
BALL_SQRT = ("cartoon",)                      # Radius wächst nur mit √(kovalenter Radius): Cl, Br, Metalle nicht riesig
ROD_INNER = {"cartoon": 0.62}                                   # Kern des Stabs (Anteil der Breite), Rest = Umriss
PERSPECTIVE = {"cartoon": 0.07}                                 # Kugeln je Å näher am Betrachter größer


def ball_radius(el, style="ballstick"):
    """Kugelradius im Kugel-Stab-Modell (Å): proportional zum kovalenten Radius, H nicht zu klein."""
    if style == "iboview":
        return ORB.ibo_radius(el)
    cov = COV.get(el, 0.76)
    if style in BALL_SQRT:
        cov = 0.76 * math.sqrt(cov / 0.76)                    # C unverändert
    return max(BALL_RMIN.get(style, 0.2), BALL_R.get(style, 0.45) * cov)


def fogged(color, fog):
    c = np.array(to_rgb(color))
    return tuple(np.clip(c + (1 - c) * fog, 0, 1))


def shade_layers(it, n, shift=0.55, smin=0.25):
    """Polygone der Schattierungsschichten (äußere zuerst). Jede Schicht ist die verkleinerte Ellipse, zum Licht
    verschoben, aber immer ganz innerhalb des Umrisses (Abstand in der Metrik der Ellipse ≤ 1 − s)."""
    a, b = it["w"] / 2, it["h"] / 2
    th = math.radians(it["ang"])
    u = LIGHT[0] * math.cos(th) + LIGHT[1] * math.sin(th)
    v = -LIGHT[0] * math.sin(th) + LIGHT[1] * math.cos(th)
    nrm = math.sqrt((u / a) ** 2 + (v / b) ** 2) or 1.0
    verts = []
    for s in np.linspace(1.0, smin, n):
        d = (1 - s) * shift / nrm
        verts.append(_ellipse_verts(it["x"] + LIGHT[0] * d, it["y"] + LIGHT[1] * d, it["w"] * s, it["h"] * s,
                                    it["ang"]))
    return verts


def _radius_along(it, dx, dy):
    """Radius der projizierten Ellipse in Richtung (dx, dy)."""
    n = math.hypot(dx, dy) or 1.0
    th = math.radians(it["ang"])
    u = (dx * math.cos(th) + dy * math.sin(th)) / n
    v = (-dx * math.sin(th) + dy * math.cos(th)) / n
    a, b = it["w"] / 2, it["h"] / 2
    return 1.0 / math.sqrt((u / a) ** 2 + (v / b) ** 2)


def label_shown(a, o):
    """Label eines Atoms: Standard alle Nicht-H-Atome (H nur mit „H-Labels“), je Atom umschaltbar."""
    if not o.get("labels", True):
        return False
    over = (o.get("label_on") or {}).get(a["key"])
    if over is not None:
        return bool(over)
    if a["key"] in (o.get("wire") or ()):                     # Drahtmodell-Gruppen: standardmäßig ohne Label
        return False
    if o.get("label_mode") == "hetero" and a["el"] == "C":
        return False
    return a["el"] not in ("H", "D") or bool(o.get("h_labels"))


def wire_color(color):
    """Linienfarbe im Drahtmodell: Elementfarbe, sehr helle Farben (H) grau, damit sie auf Weiß sichtbar bleiben."""
    r, g, b = to_rgb(color)
    return "#9a9a9a" if 0.299 * r + 0.587 * g + 0.114 * b > 0.85 else color


def _half(xs, ys, half, overlap=0.0):
    """Hälfte einer Bindung (0 = zum ersten Atom, 1 = zum zweiten); overlap (Anteil der Länge) über die Mitte hinaus,
    damit zwischen zwei Stabhälften keine Naht sichtbar ist."""
    f = 0.5 + overlap if half == 0 else 0.5 - overlap
    xm, ym = xs[0] + f * (xs[1] - xs[0]), ys[0] + f * (ys[1] - ys[0])
    return ([xs[0], xm], [ys[0], ym]) if half == 0 else ([xm, xs[1]], [ym, ys[1]])


def join_en(items):
    items = [i for i in items if i]
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + " and " + items[-1]


class XrayFigure:
    """ORTEP-Abbildung eines Moleküls. Schnittstelle wie core.Figure, soweit Tab/Export sie brauchen
    (fig, ax, cfg, ann_items, axes_items, current_layout, export). Labels: Versatz in Punkten zum Atom,
    damit sie beim Drehen mitwandern; gespeicherte Versätze werden vor der automatischen Platzierung reserviert."""

    def __init__(self, cfg, mol, R, opts, layout, fig, selection=(), zoom=1.0, pan=(0.0, 0.0)):
        self.cfg, self.mol, self.opts, self.layout = cfg, mol, opts, layout or {}
        self.R = orthonormalize(R)
        self.zoom, self.pan = float(zoom), [float(pan[0]), float(pan[1])]
        self.selection = list(selection)
        fig.clear()
        fig.set_size_inches(cfg["figsize_in"], forward=False)
        self.bg_face, self.bg_transparent, self.bg_show = bg_parts(opts.get("bg", "white"))
        fig.set_facecolor(self.bg_show)                       # transparent: Vorschau weiß
        self.fig = fig
        self.ax = fig.add_axes([0, 0, 1, 1])
        self.ax.set_axis_off()
        self.axes_items, self.ann_items, self.curves = {}, {}, []
        self.k = chi2_scale(opts.get("prob", 50) / 100)
        wire = set(opts.get("wire") or ())
        self.wire = {k for k, a in enumerate(mol.atoms) if a["key"] in wire}
        self.show = self._visible()
        pts = np.array([mol.atoms[i]["xyz"] for i in self.show]) if self.show else np.zeros((1, 3))
        self.center = pts.mean(axis=0)
        self.rmax = float(np.max(np.linalg.norm(pts - self.center, axis=1))) + 1.0
        sc = cfg.get("scale", 1.0)
        self.style = opts.get("style", "mercury")
        self.ball = self.style in BALL_STYLES
        # Bindungsdicke in Å (Mercury: Zylinder); Kugel-Stab: kräftige Stäbe
        # Kugel-Stab: Stabdicke des Stils, mit „Bindungsdicke“ relativ zum Standard des Tabs skaliert
        self.bond_w = ROD_W[self.style] * float(opts.get("rod_scale", 1.0)) if self.ball \
            else float(opts.get("bond_w", 0.08))
        self.rods = None                                         # Kugel-Stab: Schichten aller Stabhälften (_build_rods)
        self.atom_rods = {}                                      # Kugel-Stab: Atom -> [LineCollection seiner Stabhälften]
        self.atom_rods_over = {}                                 # IboView: Stabhälften vor ihrem Atom
        self.sphere_tr = {}                                      # Kugel-Stab: Atom -> Transform der Einheitskugel
        self.lw_edge = max(0.4, 0.6 * sc)
        self.atom_art, self.bond_art, self.sel_art, self.bond_lines = {}, [], [], {}
        self.elem_lab = {}
        self.auto_pos = {}
        self.orb, self.orb_art = opts.get("orbital"), None      # Orbital (elucida.core.orbitals): Bild über den Atomen
        self.scene = self._project()
        # Maßstab aus der Ausdehnung in der aktuellen Ansicht (beim Drehen bleibt er fest, nach dem Loslassen
        # passt das Neuzeichnen ihn an)
        xs = [it["x"] + s_ * max(it["w"], it["h"]) / 2 for it in self.scene["atoms"] for s_ in (-1, 1)] or [-1, 1]
        ys = [it["y"] + s_ * max(it["w"], it["h"]) / 2 for it in self.scene["atoms"] for s_ in (-1, 1)] or [-1, 1]
        for x_, y_ in self.orbital_outline():                  # Orbitallappen ganz im Bild
            xs.append(x_)
            ys.append(y_)
        self.ext = (max(abs(min(xs)), abs(max(xs)), 0.5) + 0.6, max(abs(min(ys)), abs(max(ys)), 0.5) + 0.6)
        self._limits()
        self.lw_bond = max(0.6, self.bond_w * self.pt_per_A())
        self.lw_wire = max(0.7, 0.4 * self.lw_bond)
        self._build()
        if self.orb:
            self.orb_art = ORB.OrbitalImage(self.ax, self, self.orb)
            self.orb_art.set_zorder(45)                        # über allen Atomen, unter Auswahl und Labels
            self.ax.add_image(self.orb_art)
            self.order.append(self.orb_art)
        self._footnote()
        self._labels()

    # -- Orbital ----------------------------------------------------------------------------
    def orbital_outline(self):
        """Ecken der Orbitallappen (Bildkoordinaten, Å): Ausdehnung von |ψ| > iso in der aktuellen Ansicht."""
        if not self.orb:
            return []
        P = ORB.support_points(self.orb["vol"], self.orb["iso"])
        if not len(P):
            return []
        S = (P - self.center) @ self.R.T
        lo, hi = S[:, :2].min(axis=0), S[:, :2].max(axis=0)
        return [(lo[0], lo[1]), (hi[0], hi[1])]

    def occluder(self, xl, yl, W, H):
        """Tiefe der vordersten Atom-/Stabfläche je Pixel (Szene) und Tiefe der Szenenmitte relativ zu self.center."""
        sc = self.scene
        spheres = [(it["x"], it["y"], it["z"], it["w"] / 2) for it in sc["atoms"] if it.get("kind") != "wire"]
        rods = []
        if self.rods is not None and self.rods.get("ibo") and self.rods["specs"]:
            P0, P1, _ = self._ibo_seg_geometry()
            rods = [(*p0, *p1, s["r0"], s["r1"]) for p0, p1, s in zip(P0, P1, self.rods["specs"])]
        elif self.ball:
            r0 = self.bond_w / 2
            r1 = r0
            for a, b, _ in sc["bonds"]:
                A, B = sc["by_k"][a], sc["by_k"][b]
                xm, ym, zm = (A["x"] + B["x"]) / 2, (A["y"] + B["y"]) / 2, (A["z"] + B["z"]) / 2
                rods += [(A["x"], A["y"], A["z"], xm, ym, zm, r0, r1), (B["x"], B["y"], B["z"], xm, ym, zm, r0, r1)]
        zoff = float((self.R @ (self.center - sc["center"]))[2])
        return ORB.occluder_depth(spheres, rods, xl, yl, W, H), zoff

    # -- Szene ----------------------------------------------------------------------------
    def _visible(self):
        o, mol = self.opts, self.mol
        hidden = set(o.get("hidden_frags") or [])
        kinds = set(o.get("hide_kinds") or [])
        hidden_atoms = set(o.get("hidden_atoms") or [])
        nb = mol.neighbors()
        out = []
        for i, a in enumerate(mol.atoms):
            fr = a.get("frag")
            if fr in hidden or mol.fragments[fr]["kind"] in kinds or a["key"] in hidden_atoms:
                continue
            if a["el"] in ("H", "D"):
                if not o.get("show_h", True):
                    continue
                # H an Drahtmodell-Gruppen nur mit eigener Option
                if not o.get("wire_h") and nb[i] and all(j in self.wire for j in nb[i]):
                    continue
            out.append(i)
        return out

    def _limits(self):
        W, H = self.cfg["figsize_in"]
        foot = 0.06 if self.mol.ops else 0.0                        # Platz für die Symmetrie-Fußnote
        s = max(2 * self.ext[0] / W, 2 * self.ext[1] / (H * (1 - 2 * foot))) * 1.04 / self.zoom   # Å pro Zoll
        cx, cy = self.pan
        self.ax.set_xlim(cx - W * s / 2, cx + W * s / 2)
        self.ax.set_ylim(cy - H * s / 2, cy + H * s / 2)

    def pt_per_A(self):
        x0, x1 = self.ax.get_xlim()
        return 72.0 * self.cfg["figsize_in"][0] / (x1 - x0)

    def _project(self):
        scene = project(self.mol, self.R, self.k, set(self.show), h_radius=self.opts.get("h_radius", 0.09))
        # Projektion um den Schwerpunkt der gezeigten Atome (nicht aller), damit das Bild zentriert bleibt
        off = self.R @ (self.center - scene["center"])
        for it in scene["atoms"]:
            it["x"] -= off[0]
            it["y"] -= off[1]
            for c in range(len(it["curves"])):
                xs, ys = it["curves"][c]
                it["curves"][c] = (xs - off[0], ys - off[1])
        zc = float(np.mean([it["z"] for it in scene["atoms"]])) if scene["atoms"] else 0.0
        pers = PERSPECTIVE.get(self.style, 0.0)
        for it in scene["atoms"]:
            if it["k"] in self.wire:                       # Drahtmodell: kein Ellipsoid, kleiner Klickbereich
                it.update(w=0.24, h=0.24, ang=0.0, iso=True, kind="wire", curves=[])
            elif self.ball:                                # Kugel-Stab: Kugel nach Element, keine Hauptachsen
                r = ball_radius(self.mol.atoms[it["k"]]["el"], self.style) * float(self.opts.get("ball_scale", 1.0))
                r *= min(1.35, max(0.7, 1 + pers * (it["z"] - zc)))       # Perspektive (nur Cartoon)
                it.update(w=2 * r, h=2 * r, ang=0.0, iso=True, curves=[])
        scene["by_k"] = {it["k"]: it for it in scene["atoms"]}
        return scene

    def _bond_xy(self, a, b):
        A, B = self.scene["by_k"][a], self.scene["by_k"][b]
        if self.ball and a not in self.wire and b not in self.wire:
            # Kugel-Stab: Stab von Mitte zu Mitte, die Kugel verdeckt das Ende (jede Hälfte liegt direkt unter ihrem
            # Atom, siehe _zorders) – sonst stehen die Ecken dicker Stäbe über den Kugelrand
            return [A["x"], B["x"]], [A["y"], B["y"]]
        dx, dy = B["x"] - A["x"], B["y"] - A["y"]
        L = math.hypot(dx, dy)
        ra = 0.0 if A.get("kind") == "wire" else _radius_along(A, dx, dy)
        rb = 0.0 if B.get("kind") == "wire" else _radius_along(B, -dx, -dy)
        if L <= ra + rb:
            return [np.nan, np.nan], [np.nan, np.nan]
        ux, uy = dx / L, dy / L
        return [A["x"] + ux * ra, B["x"] - ux * rb], [A["y"] + uy * ra, B["y"] - uy * rb]

    def _layers(self, it):
        """Polygone eines Atoms: Schattierungsschichten (Mercury) bzw. eine Fläche (flach), zuletzt der Umriss."""
        if it.get("kind") == "wire":
            return []
        base = _ellipse_verts(it["x"], it["y"], it["w"], it["h"], it["ang"])
        if self.style == "flat":
            return [base, base]
        if self.ball:
            r = it["w"] / 2
            c = np.array([it["x"], it["y"]])
            return [c + r * P for P, _, _ in sphere_unit_layers(self.style)] + [base]
        n = 6 if it.get("kind") == "H" else 12
        return shade_layers(it, n) + [base]

    def _fog(self, z):
        """Tiefe nach hinten heller (DEPTH_A/DEPTH_MAX) – derzeit in keinem Stil (FOG_STYLES; Kugel-Stab wie ORCA
        ohne Aufhellung)."""
        if self.style == "iboview":                       # wie IboView: ab der Mitte nach hinten zu Weiß
            return float(ORB.ibo_fog(z))
        if self.style not in FOG_STYLES:
            return 0.0
        return DEPTH_MAX * min(1.0, max(0.0, (self._zmax - z) / DEPTH_A))

    def _build_rods(self, bonds, colors):
        """Kugel-Stab: alle Schichten aller Stabhälften als Segmente, je Atom eine LineCollection (Hälften und
        Schichten in derselben Reihenfolge wie früher die Einzellinien). Die Segmente sind Sichten auf einen
        gemeinsamen Puffer, den _place_rods bei jeder Drehung in einem Rechenschritt neu füllt."""
        outline = SHADE[self.style]["outline"]
        levels = ([("o", None, None)] if outline else []) + rod_levels(self.style)
        per_atom = {}
        for i, (n, a, b) in enumerate(bonds):
            fog = self._fog(self.scene["bonds"][n][2])
            for half, k in ((0, a), (1, b)):
                el = self.mol.atoms[k]["el"]
                base = "white" if self.style == "cartoon" else colors.get(el) or default_color(el, self.style)
                for j, (kind, _, val) in enumerate(levels):
                    col_ = ((0, 0, 0) if self.style == "cartoon" else (0.15, 0.15, 0.15)) if kind == "o" \
                        else shade_color(base, kind, val, self.style, fog)
                    # nur innere Schichten überlappen in der Mitte: der Außenrand der oberen Hälfte endet genau dort,
                    # sonst bliebe über den hellen Schichten der anderen Hälfte eine Naht
                    per_atom.setdefault(k, []).append((i, half, j - (1 if outline else 0), 0.0 if j == 0 else 0.02,
                                                       col_))
        segs = [t for k in per_atom for t in per_atom[k]]
        buf = np.full((len(segs), 2, 2), np.nan)
        self.rods = dict(ab=[(a, b) for _, a, b in bonds], bond=np.array([t[0] for t in segs], int),
                         half=np.array([t[1] for t in segs], int), layer=np.array([t[2] for t in segs], int),
                         ov=np.array([t[3] for t in segs]), buf=buf, frac=np.zeros(len(segs)), slices={})
        pos = 0
        for k, lst in per_atom.items():
            sl = slice(pos, pos + len(lst))
            pos += len(lst)
            lc = LineCollection(list(buf[sl]), colors=[t[4] for t in lst], linewidths=0, capstyle="butt",
                                antialiaseds=True)
            self.ax.add_collection(lc, autolim=False)
            self.atom_rods[k] = [lc]
            self.rods["slices"][k] = sl
        # Segmente teilen sich den Puffer (Path übernimmt float-Arrays ohne Kopie); sonst jedes Mal neu setzen
        first = next(iter(self.atom_rods.values()), [None])[0]
        self.rods["inplace"] = first is not None and np.shares_memory(first.get_paths()[0].vertices, buf)
        self._place_rods()

    def _ibo_rod_specs(self, bonds):
        """Stäbe wie IboView: je Bindung n = ⌈Bindungsordnung⌉ Stäbe nebeneinander (Dicke 1.4/n, Abstand 0.7 Stabradien,
        Richtung zum nächsten nicht kollinearen Nachbarn), Teilbindung (< 0.8) gestrichelt; jede Hälfte ab 0.12 × Zeichen-
        radius im Atom bis zur Mitte, verjüngt auf 72 %. -> Segmente dict(k, a, b, t0, t1, off, r0, r1, bond)."""
        atoms = self.mol.atoms
        X = np.array([a["xyz"] for a in atoms], float)
        els = [a["el"] for a in atoms]
        bo_map = self.opts.get("bond_orders") or {}
        rb = self.bond_w / 2
        thr = ORB.IBO_FULL_BOND
        specs = []
        for n, a, b in bonds:
            bo = float(bo_map.get((min(a, b), max(a, b)), 1.0))
            if abs(bo - math.floor(bo + 0.5)) < thr:
                bo = float(math.floor(bo + 0.5))
            nmb = max(1, math.ceil(bo - 1e-9))
            u = ORB.ibo_tan_u(X, els, a, b)
            L = float(np.linalg.norm(X[b] - X[a])) or 1.0
            for k, o in ((a, b), (b, a)):
                t_base = ORB.IBO_BOND_OFFS * ORB.IBO_RADII.get(els[k], 1.5) * ORB.BOHR / L
                for im in range(nmb):
                    scale = 1.4 / nmb if nmb > 1 else 1.0
                    cx = 2.0 * im / (nmb - 1) - 1.0 if nmb > 1 else 0.0
                    w = 1.0
                    if im == nmb - 1 and bo - (nmb - 1) < 1 - thr:
                        w = bo - (nmb - 1)
                        scale *= 0.25 + 0.75 * ORB.IBO_TAPER
                    off = u * (0.7 * rb * cx)
                    rr = rb * scale

                    def seg(t0, t1, r0, r1):
                        # t von Atom k aus gezählt -> Anteil a→b
                        ta, tb = (t0, t1) if k == a else (1 - t0, 1 - t1)
                        specs.append(dict(k=k, a=a, b=b, t0=ta, t1=tb, off=off, r0=r0, r1=r1, bond=n))
                    if w >= 1.0:
                        seg(t_base, 0.5, rr, rr * ORB.IBO_TAPER)
                    else:                                          # gestrichelt: Gewicht 1 − (1 − w)², Radius × √w
                        dw = 1 - (1 - w) ** 2
                        rr *= math.sqrt(w)
                        h = (0.5 - t_base) * L
                        nseg = int(math.floor(1 + h / (0.5 * ORB.BOHR)))
                        ds = 1 - dw
                        sh = h / (nseg - ds / 2)
                        for i in range(nseg):
                            s0 = t_base * L + i * sh
                            seg(s0 / L, min(0.5, (s0 + dw * sh) / L), rr, rr)
        return specs

    def _build_rods_ibo(self, bonds, colors):
        """IboView-Stäbe (_ibo_rod_specs) als Längsstreifen (ORB.IBO_ROD_STRIPS) mit der exakten Farbe des Zylinders;
        die Farben hängen von der Lage des Stabs ab und werden bei jeder Drehung neu gesetzt (_place_rods_ibo)."""
        specs = self._ibo_rod_specs(bonds)
        ns = ORB.IBO_ROD_STRIPS
        zb = {n: z for n, (_, _, z) in enumerate(self.scene["bonds"])}
        per_atom = {}
        for si, sp in enumerate(specs):
            per_atom.setdefault(sp["k"], []).append(si)
        order, slices, pos = [], {}, 0
        for k, lst in per_atom.items():
            rows = [(si, j) for si in lst for j in range(ns)]
            slices[k] = slice(pos, pos + len(rows))
            pos += len(rows)
            order += rows
            # unter dem Atom (Stab zeigt vom Betrachter weg) und darüber (zum Betrachter, ab der Kugeloberfläche)
            pcs = []
            for _ in range(2):
                pc = PolyCollection([np.zeros((4, 2))] * len(rows), edgecolors="none", linewidths=0, antialiaseds=True)
                self.ax.add_collection(pc, autolim=False)
                pcs.append(pc)
            self.atom_rods[k] = [pcs[0]]
            self.atom_rods_over[k] = pcs[1]
        base = {}
        for sp in specs:
            el = self.mol.atoms[sp["k"]]["el"]
            base[sp["k"]] = np.array(to_rgb(colors.get(el) or default_color(el, self.style)))
        self.rods = dict(ibo=True, specs=specs, seg=np.array([t[0] for t in order], int),
                         strip=np.array([t[1] for t in order], int), slices=slices,
                         base=np.array([base[sp["k"]] for sp in specs]).reshape(-1, 3),
                         fog=np.array([self._fog(zb.get(sp["bond"], 0.0)) for sp in specs]))
        self._place_rods_ibo()

    def _ibo_seg_geometry(self):
        """Segmente in der aktuellen Ansicht -> (P0, P1 (N×3, Bildschirm, z), Richtung je Segment (dx, dy, dz))."""
        sp = self.rods["specs"]
        by = self.scene["by_k"]
        A = np.array([(by[s["a"]]["x"], by[s["a"]]["y"], by[s["a"]]["z"]) for s in sp], float)
        B = np.array([(by[s["b"]]["x"], by[s["b"]]["y"], by[s["b"]]["z"]) for s in sp], float)
        off = np.array([s["off"] for s in sp], float) @ self.R.T
        t0 = np.array([s["t0"] for s in sp])[:, None]
        t1 = np.array([s["t1"] for s in sp])[:, None]
        return A + t0 * (B - A) + off, A + t1 * (B - A) + off, B - A

    def _place_rods_ibo(self):
        r = self.rods
        if not r["specs"]:
            return
        P0, P1, D = self._ibo_seg_geometry()
        sp = r["specs"]
        ns = ORB.IBO_ROD_STRIPS
        cols = np.empty((len(sp), ns, 3))
        wps, wfs = np.empty((len(sp), 3)), np.empty((len(sp), 3))
        for bkey in np.unique(r["base"], axis=0):                 # je Grundfarbe ein Rechenschritt
            m = np.all(r["base"] == bkey, axis=1)
            cols[m], wps[m], wfs[m] = ORB.ibo_rod_colors(bkey, D[m])
        cols = cols + (1 - cols) * r["fog"][:, None, None]
        si, j = r["seg"], r["strip"]
        u0 = -1 + 2 * j / ns
        u1 = np.minimum(1.0, -1 + 2 * (j + 1.6) / ns)              # leicht überlappend: keine Nähte
        r0 = np.array([s["r0"] for s in sp])[si]
        r1 = np.array([s["r1"] for s in sp])[si]
        # Zum Betrachter zeigende Stäbe liegen vor ihrer Kugel (Tiefenpuffer in IboView): Streifen erst ab der Stelle,
        # an der die Zylinderoberfläche die Kugel verlässt, und über dem Atom gezeichnet
        by = self.scene["by_k"]
        own = np.array([s["k"] for s in sp])
        Ck = np.array([(by[k]["x"], by[k]["y"], by[k]["z"]) for k in own], float)
        Rk = np.array([by[k]["w"] / 2 for k in own])
        dseg = P1 - P0
        toward = (dseg[:, 2] > 0)[si]

        def start(uu):
            n3 = uu[:, None] * wps[si] + np.sqrt(np.clip(1 - uu ** 2, 0, 1))[:, None] * wfs[si]
            Q = P0[si] + r0[:, None] * n3 - Ck[si]
            d = dseg[si]
            dd = np.maximum(np.einsum("ij,ij->i", d, d), 1e-12)
            qd = np.einsum("ij,ij->i", Q, d)
            disc = qd ** 2 - dd * (np.einsum("ij,ij->i", Q, Q) - Rk[si] ** 2)
            t = (-qd + np.sqrt(np.maximum(disc, 0))) / dd
            return np.clip(np.where(disc > 0, t, 0.0), 0.0, 1.0)
        s0, s1 = np.where(toward, start(u0), 0.0), np.where(toward, start(u1), 0.0)
        q = np.empty((len(si), 4, 2))
        for c, (ss, uu, end) in enumerate(((s0, u0, False), (s0, u0, True), (s1, u1, True), (s1, u1, False))):
            tt = np.ones_like(ss) if end else ss
            P = P0[si] + tt[:, None] * dseg[si]
            rr = r0 + tt * (r1 - r0)
            # Punkt auf der sichtbaren Zylinderoberfläche (bei schrägen Stäben entlang der Achse verschoben)
            n3 = uu[:, None] * wps[si] + np.sqrt(np.clip(1 - uu ** 2, 0, 1))[:, None] * wfs[si]
            q[:, c, :] = P[:, :2] + rr[:, None] * n3[:, :2]
        empty = np.zeros_like(q)
        fc = cols[si, j]
        for k, sl in r["slices"].items():
            tw = toward[sl][:, None, None]
            for pc, qq in ((self.atom_rods[k][0], np.where(tw, empty[sl], q[sl])),
                           (self.atom_rods_over[k], np.where(tw, q[sl], empty[sl]))):
                pc.set_verts(qq)
                pc.set_facecolors(fc[sl])

    def _place_rods(self):
        """Geometrie aller Stabschichten nach der aktuellen Ansicht (Breite/Versatz je Lage zum Licht)."""
        r = self.rods
        if r is not None and r.get("ibo"):
            return self._place_rods_ibo()
        if r is None or not len(r["bond"]):
            return
        by = self.scene["by_k"]
        P = np.array([(by[a]["x"], by[a]["y"], by[b]["x"], by[b]["y"]) for a, b in r["ab"]], float)
        ok = np.isfinite(P).all(axis=1)
        dx, dy = P[:, 2] - P[:, 0], P[:, 3] - P[:, 1]
        frac, off = rod_layouts(self.style, np.where(ok, dx, 1.0), np.where(ok, dy, 0.0))
        frac[~ok], off[~ok] = 0.0, 0.0
        inner = ROD_INNER.get(self.style, 0.95) if SHADE[self.style]["outline"] else 1.0
        bi, lj = r["bond"], r["layer"]
        inside = lj >= 0
        fr = np.where(inside, frac[bi, np.maximum(lj, 0)] * inner, 1.0)          # Umriss (-1): volle Breite
        of = np.where(inside, off[bi, np.maximum(lj, 0)] * inner, 0.0) * self.bond_w
        # Versatz senkrecht zur Bindung zur Lichtseite (Glanzstreifen oben links)
        L = np.hypot(dx, dy)
        L = np.where(L != 0, L, 1.0)
        nx, ny = -dy / L, dx / L
        flip = np.where(nx * LIGHT[0] + ny * LIGHT[1] < 0, -1.0, 1.0)
        nx, ny = (nx * flip)[bi], (ny * flip)[bi]
        x0, y0, x1, y1 = (P[bi, c] for c in range(4))
        f = np.where(r["half"] == 0, 0.5 + r["ov"], 0.5 - r["ov"])
        xm, ym = x0 + f * (x1 - x0), y0 + f * (y1 - y0)
        h0 = r["half"] == 0
        buf = r["buf"]
        buf[:, 0, 0] = np.where(h0, x0, xm) + nx * of
        buf[:, 0, 1] = np.where(h0, y0, ym) + ny * of
        buf[:, 1, 0] = np.where(h0, xm, x1) + nx * of
        buf[:, 1, 1] = np.where(h0, ym, y1) + ny * of
        r["frac"] = fr
        self._rod_widths(reset=not r["inplace"])

    def _rod_widths(self, reset=False):
        """Linienbreiten der Stabschichten (Breitenanteil × Stabdicke in pt; leere Schicht = 0, wird nicht gezeichnet)."""
        r = self.rods
        if r.get("ibo"):                              # IboView-Stäbe: Breite steckt in der Geometrie (Å)
            return
        for k, sl in r["slices"].items():
            lc = self.atom_rods[k][0]
            if reset:
                lc.set_segments(list(r["buf"][sl]))
            lc.set_linewidths(self.lw_bond * r["frac"][sl])

    def _build(self):
        ax, colors = self.ax, self.opts.get("colors") or {}
        mercury = self.style == "mercury"
        zs = [it["z"] for it in self.scene["atoms"]] or [0.0]
        self._zmax = max(zs)
        for it in self.scene["atoms"]:
            a = self.mol.atoms[it["k"]]
            col = colors.get(a["el"]) or default_color(a["el"], self.style)
            verts = self._layers(it)
            if not verts:                                  # Drahtmodell: nur die Bindungen
                p = PolyCollection([], facecolors="none", edgecolors="none")
                ax.add_collection(p)
                self.atom_art[it["k"]] = (p, [])
                continue
            nfill = len(verts) - 1
            if self.style == "iboview":                    # wie IboView: genaue Schichten der fertigen Farbe
                lay = ORB.ibo_sphere_layers(col)
                fog = self._fog(it["z"])
                p = PolyCollection([], antialiaseds=True, closed=True)
                p.set_verts_and_codes([v for v, _, _ in lay], [c for _, c, _ in lay])
                p.set_facecolors([fogged(c, fog) for _, _, c in lay])
                p.set_edgecolors("none")
                p.set_linewidths(0)
                tr = Affine2D().scale(it["w"] / 2).translate(it["x"], it["y"])
                self.sphere_tr[it["k"]] = tr
                p.set_transform(tr + ax.transData)
                ax.add_collection(p)
                self.atom_art[it["k"]] = (p, [])
                continue
            if self.ball:                                  # Einheitskugel, verschoben/skaliert über die Transform
                verts = sphere_layer_verts(self.style)
            if self.ball:                                  # wie die Vorlagen: IboView mit feinem Umriss, Tiefe ohne
                fog = self._fog(it["z"])
                faces = [shade_color(col, kd, v, self.style, fog) for _, kd, v in sphere_unit_layers(self.style)]
                edge, elw = "none", 0.0                    # wie die Vorlagen: Kugeln ohne Umriss
                if self.style == "cartoon":                # Strichzeichnung: weiß mit schwarzem Umriss
                    faces = ["white"] * len(faces)
                    edge, elw = (0, 0, 0), max(0.5, 0.75 * self.cfg.get("scale", 1.0))
            elif mercury:
                faces, edge, elw = shade_colors(col, nfill), (0.1, 0.1, 0.1), 0.8 * self.lw_edge
            else:
                faces, edge, elw = [col], (0, 0, 0), self.lw_edge
            p = PolyCollection(verts, facecolors=faces + ["none"], edgecolors=["none"] * nfill + [edge],
                               linewidths=[0] * nfill + [elw], antialiaseds=True, closed=True)
            if self.ball:
                tr = Affine2D().scale(it["w"] / 2).translate(it["x"], it["y"])
                self.sphere_tr[it["k"]] = tr
                p.set_transform(tr + ax.transData)
            ax.add_collection(p)
            lines = []
            for xs, ys in it["curves"]:
                ln = Line2D(xs, ys, color="black", lw=max(0.35, (0.7 if mercury else 0.45) * self.cfg.get("scale", 1.0)),
                            solid_capstyle="butt")
                ax.add_line(ln)
                lines.append(ln)
            self.atom_art[it["k"]] = (p, lines)
        rod_bonds = []
        for n, (a, b, _) in enumerate(self.scene["bonds"]):
            xs, ys = self._bond_xy(a, b)
            if a in self.wire or b in self.wire:             # Drahtmodell (auch die Bindung zum Kern): zwei Hälften
                self.bond_lines[n] = []
                for half, k in ((0, a), (1, b)):
                    el = self.mol.atoms[k]["el"]
                    ln = Line2D(*_half(xs, ys, half), color=wire_color(colors.get(el) or default_color(el)),
                                lw=self.lw_wire, solid_capstyle="round")
                    ax.add_line(ln)
                    self.bond_art.append((a, b, ln, half))
                    self.bond_lines[n].append(ln)
                continue
            if self.ball:                                  # Zylinder (Phong), je Hälfte in der Atomfarbe
                self.bond_lines[n] = []
                rod_bonds.append((n, a, b))
                continue
            ln = Line2D(xs, ys, color="black", lw=self.lw_bond, solid_capstyle="butt")
            ax.add_line(ln)
            self.bond_art.append((a, b, ln, None))
            self.bond_lines[n] = [ln]
        for n, k in enumerate(self.selection):
            if k in self.scene["by_k"]:
                it = self.scene["by_k"][k]
                e = Ellipse((it["x"], it["y"]), it["w"] * 1.35 + 0.12, it["h"] * 1.35 + 0.12, angle=it["ang"],
                            facecolor="none", edgecolor=SEL_COLOR, lw=2.2, zorder=50)
                ax.add_patch(e)
                self.sel_art.append((k, e))
        if rod_bonds:
            (self._build_rods_ibo if self.style == "iboview" else self._build_rods)(rod_bonds, colors)
        self._elem_labels(colors)
        self._zorders()

    def _elem_labels(self, colors):
        """IboView: Elementsymbole auf C und Heteroatomen (nicht H), Atomfarbe um 20 % dunkler mit schwarzem Rand,
        Größe 0.8 Bohr, unter dem Orbital; Aufhellung nach hinten wie die Atome."""
        self.elem_lab = {}
        if self.style != "iboview" or not self.opts.get("elem_labels", True):
            return
        import matplotlib.patheffects as PE
        fs = self._elem_fs()
        for it in self.scene["atoms"]:
            el = self.mol.atoms[it["k"]]["el"]
            if el in ("H", "D") or it.get("kind") == "wire":
                continue
            fog = self._fog(it["z"])
            col = np.array(to_rgb(colors.get(el) or default_color(el, self.style))) * (1 - ORB.IBO_LABEL_DARK)
            t = self.ax.text(it["x"], it["y"], el, ha="center", va="center", fontsize=fs, color=fogged(col, fog),
                             family=ORB.IBO_LABEL_FONTS,
                             path_effects=[PE.withStroke(linewidth=ORB.IBO_LABEL_EDGE * fs,
                                                         foreground=fogged((0, 0, 0), fog))])
            t._fog = fog
            self.elem_lab[it["k"]] = t

    def _elem_fs(self):
        return ORB.IBO_LABEL_EM * ORB.BOHR * self.pt_per_A() * float(self.opts.get("ball_scale", 1.0))

    def _zorders(self):
        """Zeichenreihenfolge nach Tiefe (hinten zuerst); Bindungen zwischen ihren Atomen."""
        items = [(it["z"], "a", it["k"]) for it in self.scene["atoms"]]
        # Kugel-Stab: Stabhälften gehören zu ihrem Atom (direkt darunter gezeichnet), Drahtmodell-Bindungen wie sonst
        items += [(z, "b", n) for n, (a, b, z) in enumerate(self.scene["bonds"]) if self.bond_lines.get(n)]
        items.sort(key=lambda t: (t[0], t[1] == "a"))
        self.order = []
        step = 40 / max(1, len(items))
        for rank, (_, kind, key) in enumerate(items):
            zo = 2 + step * rank
            if kind == "a":
                p, lines = self.atom_art[key]
                rods = self.atom_rods.get(key, ())
                for j, ln in enumerate(rods):                # zwischen dem vorigen Element und dem Atom
                    ln.set_zorder(zo - 0.5 * step + 0.4 * step * j / max(1, len(rods)))
                    self.order.append(ln)
                p.set_zorder(zo)
                self.order.append(p)
                for ln in lines:
                    ln.set_zorder(zo + 1e-4)
                    self.order.append(ln)
                over = self.atom_rods_over.get(key)
                if over is not None:                          # IboView: Stäbe zum Betrachter vor der Kugel
                    over.set_zorder(zo + 1.5e-4)
                    self.order.append(over)
                t = self.elem_lab.get(key)
                if t is not None:                             # IboView: Elementsymbol auf seinem Atom
                    t.set_zorder(zo + 2e-4)
                    self.order.append(t)
            else:
                for ln in self.bond_lines[key]:
                    ln.set_zorder(zo)
                    self.order.append(ln)
        self.order += [e for _, e in self.sel_art] + list(self.ann_items.values())
        if getattr(self, "orb_art", None) is not None:
            self.orb_art.stale = True
            self.order.append(self.orb_art)

    def set_view(self, R):
        """Nur Geometrie der vorhandenen Artists ändern (für flüssiges Drehen per Blitting)."""
        self.R = orthonormalize(R)
        self.scene = self._project()
        for it in self.scene["atoms"]:
            p, lines = self.atom_art[it["k"]]
            tr = self.sphere_tr.get(it["k"])
            if tr is not None:                             # Kugel: nur verschieben, Schichten bleiben
                tr.clear().scale(it["w"] / 2).translate(it["x"], it["y"])
                p.stale = True
            else:
                p.set_verts(self._layers(it))
            for ln, (xs, ys) in zip(lines, it["curves"]):
                ln.set_data(xs, ys)
        self._place_rods()
        for a, b, ln, half in self.bond_art:
            xs, ys = self._bond_xy(a, b)
            ln.set_data(*(_half(xs, ys, half) if half is not None else (xs, ys)))
        for k, t in getattr(self, "elem_lab", {}).items():
            it = self.scene["by_k"][k]
            t.set_position((it["x"], it["y"]))
        for k, e in self.sel_art:
            it = self.scene["by_k"][k]
            e.set_center((it["x"], it["y"]))
            e.width, e.height, e.angle = it["w"] * 1.35 + 0.12, it["h"] * 1.35 + 0.12, it["ang"]
        for key, ann in self.ann_items.items():
            ann.xy = self._anchor(ann._atom)
        self._zorders()

    def set_drag(self, on):
        """Kugel-Stab beim Ziehen: Kugelschichten mit vereinfachten Polygonen (Abweichung ≤ DRAG_TOL_PX Pixel bei der
        aktuellen Größe), danach wieder die vollen. Ändert nur die Vorschau während der Bewegung."""
        if self.orb_art is not None:                       # Orbital beim Ziehen gröber gerechnet
            self.orb_art.drag = on
            self.orb_art.stale = True
        if not self.sphere_tr:
            return
        px = 1.0 / self.data_per_px()
        if self.style == "iboview":                         # IboView-Kugeln: Konturen beim Ziehen gröber
            colors = self.opts.get("colors") or {}
            for k in self.sphere_tr:
                el = self.mol.atoms[k]["el"]
                tol = None
                if on:
                    r_px = self.scene["by_k"][k]["w"] / 2 * px
                    tol = next((t for t in DRAG_TOLS if t * r_px <= DRAG_TOL_PX), None)
                lay = ORB.ibo_sphere_layers(colors.get(el) or default_color(el, self.style), tol=tol)
                self.atom_art[k][0].set_verts_and_codes([v for v, _, _ in lay], [c for _, c, _ in lay])
            return
        for k in self.sphere_tr:
            tol = None
            if on:
                r_px = self.scene["by_k"][k]["w"] / 2 * px
                tol = next((t for t in DRAG_TOLS if t * r_px <= DRAG_TOL_PX), None)
            self.atom_art[k][0].set_verts(sphere_layer_verts(self.style, tol))

    def set_zoom_pan(self, zoom=None, pan=None):
        if zoom is not None:
            self.zoom = zoom
        if pan is not None:
            self.pan = list(pan)
        self._limits()
        self.lw_bond = max(0.6, self.bond_w * self.pt_per_A())
        self.lw_wire = max(0.7, 0.4 * self.lw_bond)
        for _, _, ln, half in self.bond_art:
            ln.set_linewidth(self.lw_bond if half is None else self.lw_wire)
        if self.rods is not None:
            self._rod_widths()
        import matplotlib.patheffects as PE
        for t in getattr(self, "elem_lab", {}).values():
            fs = self._elem_fs()
            t.set_fontsize(fs)
            t.set_path_effects([PE.withStroke(linewidth=ORB.IBO_LABEL_EDGE * fs, foreground=fogged((0, 0, 0), t._fog))])

    def artists(self):
        return self.order

    def _anchor(self, k):
        it = self.scene["by_k"][k]
        return it["x"], it["y"]

    # -- Labels -----------------------------------------------------------------------------
    def _label_text(self, a, bold=False):
        t = a["label"]
        if a.get("sym"):
            t += r"$^{\math" + ("bf" if bold else "rm") + "{" + a["sym"] + "}}$"
        return t

    def _labels(self):
        o = self.opts
        if not o.get("labels", True):
            return
        lc = self.cfg["labels"]
        fs = o.get("label_size") or lc["font_size"]           # eigene Labelgröße des Tabs (0 = Textformat)
        todo = []
        halo = [PE.withStroke(linewidth=2.2 * self.cfg.get("scale", 1.0), foreground=self.bg_show)]
        # Mercury: fett ohne Rand (wie Mercury-Abbildungen); flach (Olex2): fett mit weißem Rand; Kugel-Stab: normal
        # mit weißem Rand (Labels über glänzenden Kugeln lesbar)
        extra = {"flat": dict(fontweight="bold", path_effects=halo), "mercury": dict(fontweight="bold")}.get(
            self.style, dict(path_effects=halo))
        for it in self.scene["atoms"]:
            a = self.mol.atoms[it["k"]]
            if not label_shown(a, o):
                continue
            key = "lbl_" + a["key"]
            ann = self.ax.annotate(self._label_text(a, self.style in ("flat", "mercury")), xy=(it["x"], it["y"]), xycoords="data",
                                   xytext=(0, 0), textcoords="offset points", ha="center", va="center",
                                   fontsize=fs, zorder=60,
                                   color=label_color(a["el"], o.get("colors") or {}, o.get("label_color"), self.style)
                                   or lc.get("color") or "black",
                                   annotation_clip=False, **extra)
            ann._atom = it["k"]
            self.ann_items[key] = ann
            todo.append((key, ann, it))
        if not todo:
            return
        from matplotlib.backends.backend_agg import RendererAgg
        W, H = self.fig.get_size_inches() * self.fig.dpi
        rend = RendererAgg(int(W), int(H), self.fig.dpi)
        tr = self.ax.transData
        ppt = self.fig.dpi / 72.0
        # Belegt: Atome (Kreise um den Umriss) und Bindungen (Punkte entlang der Linie), in Pixeln
        circ = []
        for it in self.scene["atoms"]:
            cx, cy = tr.transform((it["x"], it["y"]))
            rx = abs(tr.transform((it["x"] + max(it["w"], it["h"]) / 2, it["y"]))[0] - cx)
            circ.append((cx, cy, rx))
        # Bindungen dicht abgetastet (alle ~2 px), mit halber Linienbreite + Rand: ein Label darf keine Bindung berühren
        segpts, seghw = [], []
        margin = 2.5 * ppt
        for _, _, ln, _h in self.bond_art:
            xs, ys = ln.get_data()
            if np.isfinite(xs).all():
                a_ = tr.transform((xs[0], ys[0]))
                b_ = tr.transform((xs[1], ys[1]))
                n_ = max(8, int(np.hypot(*(b_ - a_)) / 2))
                f_ = np.linspace(0, 1, n_)[:, None]
                segpts.append(a_ + f_ * (b_ - a_))
                seghw.append(np.full(n_, ln.get_linewidth() * ppt / 2 + margin))
        circ = np.array(circ)
        segpts = np.vstack(segpts) if segpts else np.zeros((0, 2))
        seghw = np.concatenate(seghw) if seghw else np.zeros(0)
        boxes = []

        def cost(bx):
            x0, y0, x1, y1 = bx
            c = 0.0
            if len(circ):
                nx = np.clip(circ[:, 0], x0, x1)
                ny = np.clip(circ[:, 1], y0, y1)
                d = np.hypot(nx - circ[:, 0], ny - circ[:, 1])
                c += float(np.sum(np.clip(circ[:, 2] - d, 0, None))) * 2
            if len(segpts):
                dx = np.maximum(np.maximum(x0 - segpts[:, 0], segpts[:, 0] - x1), 0)
                dy = np.maximum(np.maximum(y0 - segpts[:, 1], segpts[:, 1] - y1), 0)
                over = np.clip(seghw - np.hypot(dx, dy), 0, None)
                c += 40.0 * float(np.count_nonzero(over)) + 10.0 * float(over.sum())
            for b in boxes:
                ox = min(x1, b[2]) - max(x0, b[0])
                oy = min(y1, b[3]) - max(y0, b[1])
                if ox > 0 and oy > 0:
                    c += 0.5 * ox * oy
            fw, fh = W, H
            if x0 < 0 or y0 < 0 or x1 > fw or y1 > fh:
                c += 1e4
            return c
        def hits(bx):
            """Label liegt auf einer Bindung (Linie selbst, ohne Rand) oder einem Atom."""
            x0, y0, x1, y1 = bx
            if len(circ):
                d = np.hypot(np.clip(circ[:, 0], x0, x1) - circ[:, 0], np.clip(circ[:, 1], y0, y1) - circ[:, 1])
                if (circ[:, 2] - d > 1.0).any():
                    return True
            if len(segpts):
                dx = np.maximum(np.maximum(x0 - segpts[:, 0], segpts[:, 0] - x1), 0)
                dy = np.maximum(np.maximum(y0 - segpts[:, 1], segpts[:, 1] - y1), 0)
                if (seghw - margin - np.hypot(dx, dy) > 0).any():
                    return True
            return False
        sizes = {}
        for key, ann, it in todo:
            bb = ann.get_window_extent(rend)
            sizes[key] = (bb.width, bb.height)
        # 1) gespeicherte Versätze zuerst eintragen – außer sie liegen (z. B. nach dem Drehen) auf einer Bindung oder
        #    einem Atom; dann wird das Label wie ein neues platziert
        kept = set()
        for key, ann, it in todo:
            pos = self.layout.get(key)
            if pos and len(pos) == 2:
                cx, cy = tr.transform((it["x"], it["y"]))
                w, h = sizes[key]
                bx = cx + pos[0] * ppt, cy + pos[1] * ppt
                box = (bx[0] - w / 2, bx[1] - h / 2, bx[0] + w / 2, bx[1] + h / 2)
                if hits(box):
                    continue
                ann.xyann = (float(pos[0]), float(pos[1]))
                boxes.append(box)
                kept.add(key)
        # 2) übrige: beste von 16 Richtungen, Abstand = Umriss + Rand
        for key, ann, it in todo:
            if key in kept:
                continue
            cx, cy = tr.transform((it["x"], it["y"]))
            w, h = sizes[key]
            best = None
            for rr in (1.0, 1.35, 1.7, 2.1):
                for ang in np.linspace(0, 2 * np.pi, 16, endpoint=False):
                    ux, uy = math.cos(ang), math.sin(ang)
                    r_at = _radius_along(it, ux, uy)
                    rpx = abs(tr.transform((it["x"] + r_at, it["y"]))[0] - cx)
                    # Abstand so, dass die Textbox den Umriss gerade nicht berührt
                    ext = abs(ux) * w / 2 + abs(uy) * h / 2
                    dist = (rpx + ext + 3.0 * ppt) * rr
                    bx = (cx + ux * dist - w / 2, cy + uy * dist - h / 2, cx + ux * dist + w / 2, cy + uy * dist + h / 2)
                    c = cost(bx) + (rr - 1) * 40 + 0.01 * ang
                    if best is None or c < best[0]:
                        best = (c, bx, ux * dist / ppt, uy * dist / ppt)
            _, bx, ox, oy = best
            ann.xyann = (round(ox, 1), round(oy, 1))
            boxes.append(bx)
        self.order += [a for a in self.ann_items.values() if a not in self.order]
        self.auto_pos = self.current_layout()

    def _footnote(self):
        notes = []
        used = {self.mol.atoms[i].get("opkey") for i in self.show}
        for op, r in self.mol.ops.items():
            if op in used and (self.opts.get("labels", True)):
                notes.append(r"$^{\mathrm{" + r + "}}$" + op)
        self.foot = None
        if notes:
            fs = self.cfg["labels"]["font_size"] * 0.8
            self.foot = self.fig.text(0.01, 0.01, "; ".join(notes), ha="left", va="bottom", fontsize=fs,
                                      color="black", zorder=70)

    # -- Schnittstelle Export/Tab ---------------------------------------------------------
    def current_layout(self):
        out = {}
        for key, ann in self.ann_items.items():
            out[key] = [round(float(v), 1) for v in ann.xyann]
        return out

    def data_per_px(self):
        x0, x1 = self.ax.get_xlim()
        return (x1 - x0) / max(1.0, self.ax.bbox.width)

    def export(self):
        from matplotlib.transforms import Bbox
        out = Path(self.cfg["output"])
        out.parent.mkdir(parents=True, exist_ok=True)
        paths = [out] + [out.with_suffix("." + e.lstrip(".")) for e in self.cfg["extra_formats"]]
        W, H = self.fig.get_size_inches()
        kw = dict(dpi=self.cfg["dpi"], bbox_inches=Bbox.from_bounds(0, 0, W, H), facecolor=self.bg_face,
                  transparent=self.bg_transparent)
        from elucida.core import images as IMG
        for p in dict.fromkeys(paths):
            fast = {"pil_kwargs": {"compress_level": 1}} if p.suffix.lower() == ".png" and self.cfg.get("png_fast") else {}
            with matplotlib.rc_context({"svg.fonttype": "path"}):
                self.fig.savefig(p, **kw, **fast)
            if p.suffix.lower() == ".svg":
                IMG.svg_defs_first(p)
        return paths


def save_tight(F, svg):
    """Abbildung eng um das Molekül beschnitten (wie in der SI): SVG (Vektor, Text als Pfade) + PNG-Ersatz (600 dpi).
    -> (PNG, Breite cm, Höhe cm, SVG) für SI.write_docx."""
    from matplotlib.transforms import Bbox
    from elucida.core import images as IMG
    svg = Path(svg)
    svg.parent.mkdir(parents=True, exist_ok=True)
    rend = F.fig.canvas.get_renderer()
    # Atome über ihre Umrisse (get_window_extent der Polygone ist dafür nicht verlässlich), Text über seine Box
    pts = [(it["x"] + sx * max(it["w"], it["h"]) / 2, it["y"] + sy * max(it["w"], it["h"]) / 2)
           for it in F.scene["atoms"] for sx in (-1, 1) for sy in (-1, 1)] + F.orbital_outline()
    boxes = [Bbox.from_extents(*F.ax.transData.transform(pts).min(axis=0), *F.ax.transData.transform(pts).max(axis=0))
             ] if pts else []
    boxes += [a.get_window_extent(rend) for a in list(F.ann_items.values()) + ([F.foot] if F.foot else [])
              if a.get_visible()]
    boxes = [b for b in boxes if b.width > 0 and b.height > 0 and b.width < 1e5]
    bb = Bbox.union(boxes).transformed(F.fig.dpi_scale_trans.inverted()).padded(0.06)
    W, H = F.fig.get_size_inches()
    bb = Bbox.from_extents(max(0, bb.x0), max(0, bb.y0), min(W, bb.x1), min(H, bb.y1))
    with matplotlib.rc_context({"svg.fonttype": "path"}):
        kw = dict(bbox_inches=bb, facecolor=F.bg_face, transparent=F.bg_transparent)
        F.fig.savefig(svg, **kw)
        F.fig.savefig(svg.with_suffix(".png"), dpi=600, pil_kwargs={"compress_level": 1}, **kw)   # docx packt nochmal
    IMG.svg_defs_first(svg)
    return str(svg.with_suffix(".png")), bb.width * 2.54, bb.height * 2.54, str(svg)


# ---------------------------------------------------------------------------- Maus
class XrayMouse:
    """Ziehen = Drehen (Blitting: nur das Molekül neu zeichnen), rechte/mittlere Taste = Verschieben,
    Mausrad = Zoom, Klick = Atom wählen (vorderstes gewinnt), Labels ziehen = Versatz ändern,
    Shift + Rahmen ziehen = mehrere Atome für Labels an/aus (tab.atoms_boxed)."""

    CLICK_PX = 4

    def __init__(self, F, tab):
        self.F, self.tab = F, tab
        self.mode, self.start, self.moved, self._bg = None, None, False, None
        self.box = None
        for ann in F.ann_items.values():
            ann.draggable(True, use_blit=True)
        c = F.fig.canvas
        self.cids = [c.mpl_connect("button_press_event", self.press),
                     c.mpl_connect("motion_notify_event", self.move),
                     c.mpl_connect("button_release_event", self.release),
                     c.mpl_connect("scroll_event", self.scroll)]

    def disconnect(self):
        for cid in self.cids:
            self.F.fig.canvas.mpl_disconnect(cid)
        for ann in self.F.ann_items.values():
            try:
                ann.draggable(False)
            except Exception:
                pass

    def _over_label(self, ev):
        return any(a.contains(ev)[0] for a in self.F.ann_items.values())

    def press(self, ev):
        if ev.x is None:
            return
        if ev.dblclick and ev.button == 1:
            k = self.hit(ev)
            self.mode = None
            if k is not None:
                self.tab.atom_double_clicked(k)
            return
        if ev.button == 1 and "shift" in (ev.key or "") and ev.xdata is not None:
            from matplotlib.patches import Rectangle
            self.mode, self.start = "box", (ev.xdata, ev.ydata)
            self.box = Rectangle(self.start, 0, 0, fill=True, fc=(0.23, 0.51, 0.96, 0.08), ec="#3b82f6", lw=0.8,
                                 ls="--", zorder=1e6)
            self.F.ax.add_patch(self.box)
            return
        if self._over_label(ev) and ev.button == 1:
            self.mode = "label"
            return
        self.mode = "rot" if ev.button == 1 else "pan"
        self.start = (ev.x, ev.y)
        self.R0 = self.F.R.copy()
        self.pan0 = list(self.F.pan)
        self.moved = False

    def _start_blit(self):
        canvas = self.F.fig.canvas
        if not getattr(canvas, "supports_blit", False):
            return
        for a in self.F.artists():
            a.set_animated(True)
        canvas.draw()
        self._bg = canvas.copy_from_bbox(self.F.fig.bbox)

    def _blit(self):
        canvas = self.F.fig.canvas
        if self._bg is None:
            canvas.draw_idle()
            return
        canvas.restore_region(self._bg)
        for a in sorted(self.F.artists(), key=lambda a: a.get_zorder()):
            self.F.ax.draw_artist(a)
        canvas.blit(self.F.fig.bbox)

    def move(self, ev):
        if ev.x is None:
            return
        if self.mode == "box":
            if ev.xdata is not None and self.box is not None:
                self.box.set_bounds(self.start[0], self.start[1], ev.xdata - self.start[0], ev.ydata - self.start[1])
                self.F.fig.canvas.draw_idle()
            return
        if self.mode in ("rot", "pan"):
            dx, dy = ev.x - self.start[0], ev.y - self.start[1]
            if not self.moved and math.hypot(dx, dy) < self.CLICK_PX:
                return
            if not self.moved:
                self.moved = True
                self.F.set_drag(True)
                self._start_blit()
            if self.mode == "rot":
                s = 2.5 * math.pi / max(200.0, min(self.F.fig.bbox.width, self.F.fig.bbox.height))
                self.F.set_view(rot("y", dx * s) @ rot("x", -dy * s) @ self.R0)
            else:
                d = self.F.data_per_px()
                if self._bg is not None:          # Verschieben ändert die Achsengrenzen -> ohne Blitting
                    for a in self.F.artists():
                        a.set_animated(False)
                    self._bg = None
                self.F.set_zoom_pan(pan=(self.pan0[0] - dx * d, self.pan0[1] - dy * d))
            self._blit()
        elif self.mode is None:
            hit = self.hit(ev)
            self.tab.hover_atom(hit)

    def boxed(self):
        """Atome, deren Mitte im gezogenen Rahmen liegt (Bildkoordinaten der Szene)."""
        if self.box is None or self.F.scene is None:
            return []
        x0, y0 = self.box.get_x(), self.box.get_y()
        x1, y1 = x0 + self.box.get_width(), y0 + self.box.get_height()
        x0, x1, y0, y1 = min(x0, x1), max(x0, x1), min(y0, y1), max(y0, y1)
        return [it["k"] for it in self.F.scene["atoms"] if x0 <= it["x"] <= x1 and y0 <= it["y"] <= y1]

    def hit(self, ev):
        if ev.xdata is None or self.F.scene is None:
            return None
        return pick(self.F.scene, ev.xdata, ev.ydata, tol=2 * self.F.data_per_px())

    def release(self, ev):
        mode, self.mode = self.mode, None
        if mode == "box":
            ks = self.boxed()
            self.box.remove()
            self.box = None
            self.F.fig.canvas.draw_idle()
            if ks:
                QTimer.singleShot(0, lambda: self.tab.atoms_boxed(ks))
            return
        if mode == "label":
            QTimer.singleShot(0, lambda: self.tab.labels_moved(self.F.current_layout()))
            return
        if mode is None:
            return
        if self._bg is not None:
            for a in self.F.artists():
                a.set_animated(False)
            self._bg = None
        if not self.moved:
            if mode == "rot":
                self.tab.atom_clicked(self.hit(ev))
            return
        self.F.set_drag(False)
        self.F.fig.canvas.draw_idle()
        if mode == "rot":
            self.tab.view_changed(self.F.R)
        else:
            self.tab.pan_changed(self.F.pan)

    def scroll(self, ev):
        if ev.x is None:
            return
        import time
        now = time.monotonic()
        if now - getattr(self, "_last", 0) < 0.03:      # Trackpads liefern sehr viele Ereignisse
            return
        self._last = now
        up = (ev.step > 0) if ev.step else (ev.button == "up")
        z = min(20.0, max(0.2, self.F.zoom * (1.08 if up else 1 / 1.08)))
        self.F.set_zoom_pan(zoom=z)
        self.F.fig.canvas.draw_idle()
        self.tab.zoom_changed(z)


# ---------------------------------------------------------------------------- Tab
import uuid  # noqa: E402

from matplotlib.backends.backend_agg import FigureCanvasAgg  # noqa: E402
from matplotlib.figure import Figure as MplFigure  # noqa: E402
from PySide6.QtCore import Qt, QTimer  # noqa: E402
from PySide6.QtWidgets import (QAbstractItemView, QCheckBox, QComboBox, QDoubleSpinBox,  # noqa: E402
                               QFileDialog, QGridLayout, QGroupBox, QHBoxLayout, QHeaderView, QLabel, QLineEdit,
                               QListWidget, QListWidgetItem, QPushButton, QSpinBox, QTableWidget,
                               QTableWidgetItem, QTextBrowser, QVBoxLayout, QWidget)

from elucida.core import uvvis as core  # noqa: E402
from elucida.app import project as PRJ  # noqa: E402
from elucida.tabs import base as TABS  # noqa: E402
from elucida.app.i18n import T  # noqa: E402
from elucida.tabs.ir_ea import copy_rich  # noqa: E402
from elucida.app import style as STYLE  # noqa: E402
from elucida.app.style import HINT_CSS, WARN_CSS, make_form  # noqa: E402

CAPTION_MAX = 8                       # bis zu so vielen ausgewählten Werten in der Bildunterschrift, sonst Tabellen
KIND_HEAD = {"bond": "Bond lengths [Å]", "angle": "Angles [°]", "torsion": "Torsion angles [°]"}
CALC_NOTE = "calculated from the atomic coordinates; no s.u. in the CIF"


class _Vis:
    """Minimale Figur-Attrappe, um XrayFigure._visible ohne Zeichnen zu nutzen (Bildunterschrift)."""

    def __init__(self, mol, opts):
        self.mol, self.opts = mol, opts
        self.wire = {k for k, a in enumerate(mol.atoms) if a["key"] in (opts.get("wire") or ())}


def label_size_spin(st):
    """Labelgröße in pt; 0 = wie unter Format → Text (Standard)."""
    sp = QDoubleSpinBox(decimals=1, minimum=0.0, maximum=30.0, singleStep=0.5, suffix=" pt",
                        value=float(st.get("label_pt", 0.0)))
    sp.setSpecialValueText(T("xr_label_size_std"))
    sp.setToolTip(T("xr_label_size_tip"))
    return sp


def label_color_combo(st, on_change):
    """Labelfarbe: schwarz (Standard) oder wie das Atom (zu helle Farben abgedunkelt, damit lesbar)."""
    cb = QComboBox()
    for key in ("black", "atom"):
        cb.addItem(T("xr_label_color_" + key), key)
    cb.setCurrentIndex(max(0, cb.findData(st.get("label_color", "black"))))
    cb.currentIndexChanged.connect(on_change)
    return cb


def label_color(el, colors, mode, style=None):
    """Farbe eines Atomlabels: schwarz oder Atomfarbe; helle Farben (H weiß, Cl hellgrün) bis zur Lesbarkeit dunkler."""
    if mode != "atom":
        return None
    c = np.array(to_rgb(colors.get(el) or default_color(el, style)))
    lum = 0.299 * c[0] + 0.587 * c[1] + 0.114 * c[2]
    if lum > 0.45:
        c = c * (0.45 / lum)
    return tuple(np.clip(c, 0, 1))


def bg_parts(bg):
    """Hintergrund (Zustand „bg“: "white", "transparent" oder "#rrggbb") -> (Füllfarbe für savefig, transparent?,
    Farbe der Vorschau und der Label-Ränder)."""
    if bg == "transparent":
        return "none", True, "white"
    c = bg if isinstance(bg, str) and bg.startswith("#") else "white"
    return c, False, c


class BackgroundCombo(QComboBox):
    """Hintergrund der Abbildung (Vorschau, PDF/SVG/PNG, Word): Weiß, transparent oder eigene Farbe; schreibt st["bg"]
    und meldet die Änderung über on_change."""

    def __init__(self, st, on_change):
        super().__init__()
        self.st, self.on_change = st, on_change
        for key in ("white", "transparent", "color"):
            self.addItem(T("bg_" + key), key)
        self.setToolTip(T("bg_tip"))
        self.sync()
        self.activated.connect(self._chosen)

    def sync(self):
        from PySide6.QtGui import QColor, QIcon, QPixmap
        bg = self.st.get("bg", "white")
        kind = bg if bg in ("white", "transparent") else "color"
        self.setItemText(2, T("bg_color_x", c=bg) if kind == "color" else T("bg_color"))
        pm = QPixmap(12, 12)
        pm.fill(QColor(bg if kind == "color" else "#ffffff"))
        self.setItemIcon(2, QIcon(pm))
        self.blockSignals(True)
        self.setCurrentIndex(self.findData(kind))
        self.blockSignals(False)

    def _chosen(self, _i):
        kind = self.currentData()
        if kind == "color":
            from PySide6.QtGui import QColor
            from PySide6.QtWidgets import QColorDialog
            cur = self.st.get("bg", "")
            c = QColorDialog.getColor(QColor(cur if cur.startswith("#") else "#ffffff"), self, T("bg_pick"))
            if not c.isValid():
                self.sync()
                return
            kind = c.name()
        if kind != self.st.get("bg", "white"):
            self.st["bg"] = kind
            self.sync()
            self.on_change()


def reset_buttons(tab):
    """„Ansicht zurücksetzen“ + „Darstellung auf Standard“ (X-Ray und DFT)."""
    row = QHBoxLayout()
    for key, fn in (("xr_reset_view", tab.reset_view), ("xr_reset_display", tab.reset_display)):
        b = QPushButton(T(key))
        b.setToolTip(T(key + "_tip"))
        b.clicked.connect(fn)
        row.addWidget(b)
    row.addStretch(1)
    return row


def reset_display(tab, keep, entry_extra):
    """Alle Einstellungen des Tabs außer `keep` auf DEFAULT_STATE, eigene Labelgröße auf „wie Textformat“, die
    gewählte Struktur auf Startansicht (+ entry_extra). Rückgängig holt alles zurück."""
    import copy
    st = tab.st
    for k, v in tab.DEFAULT_STATE.items():
        if k not in keep:
            st[k] = copy.deepcopy(v)
    st["label_pt"] = 0.0
    e = tab.entry()
    if e is not None:
        e.update(view=None, zoom=1.0, pan=[0.0, 0.0], layout={}, **copy.deepcopy(entry_extra))
    tab.sync_widgets()
    tab.render()
    tab.main.push_undo()
    tab.main.statusBar().showMessage(T("xr_reset_display_done"), 6000)


class MeasureAlign:
    """Auswahl (bis 4 Atome), Messen und Ausrichten – gemeinsam für die Tabs X-Ray und DFT. Erwartet entry(),
    molecule(e), view_of(e, mol), render(push) und self.sel (Atom-Schlüssel, Label|Operation bzw. a<Index>)."""
    MEASURE_ESD = True          # X-Ray: Herkunft/esd des Werts melden; DFT: Werte aus der Geometrie, ohne esd

    def measure_box(self, title, take=None):
        """Gruppe mit Messwert der Auswahl, „Übernehmen“ (optional) und „Auswahl aufheben“ -> (Gruppe, Layout)."""
        g = QGroupBox(title)
        v = QVBoxLayout(g)
        self.lbl_meas = QLabel(T("xr_select_hint"))
        self.lbl_meas.setWordWrap(True)
        self.lbl_meas.setTextInteractionFlags(Qt.TextSelectableByMouse)
        v.addWidget(self.lbl_meas)
        row = QHBoxLayout()
        self.btn_take = None
        if take is not None:
            self.btn_take = QPushButton(T("xr_take"))
            self.btn_take.clicked.connect(take)
            row.addWidget(self.btn_take)
        b = QPushButton(T("xr_clear_sel"))
        b.clicked.connect(self.clear_selection)
        row.addWidget(b)
        row.addStretch(1)
        v.addLayout(row)
        return g, v

    def align_box(self):
        """Gruppe „Ausrichten“: Bindung in eine Richtung, Ebene von vorn, schrittweise drehen, wenden, ChemDraw-Vorlage."""
        g = QGroupBox(T("xr_align"))
        f = make_form(g)
        row = QHBoxLayout()
        for txt, d in (("↓", "down"), ("↑", "up"), ("←", "left"), ("→", "right")):
            b = QPushButton(txt)
            b.setToolTip(T("xr_align_bond_tip"))
            b.clicked.connect(lambda _=False, d=d: self.align_bond(d))
            row.addWidget(b)
        row.addStretch(1)
        f.addRow(T("xr_align_bond"), row)
        b = QPushButton(T("xr_align_plane"))
        b.setToolTip(T("xr_align_plane_tip"))
        b.clicked.connect(self.align_plane)
        f.addRow(b)
        row = QHBoxLayout()
        self.cmb_step = QComboBox()
        for v_ in (5, 15, 45, 90):
            self.cmb_step.addItem(f"{v_}°", v_)
        self.cmb_step.setCurrentIndex(1)
        row.addWidget(self.cmb_step)
        for txt, ax_, sg in (("x−", "x", -1), ("x+", "x", 1), ("y−", "y", -1), ("y+", "y", 1), ("z−", "z", -1),
                             ("z+", "z", 1)):
            b = QPushButton(txt)
            b.setToolTip(T("xr_rot_tip"))
            b.clicked.connect(lambda _=False, a=ax_, s_=sg: self.rotate_view(a, s_ * self.cmb_step.currentData()))
            row.addWidget(b)
        row.addStretch(1)
        f.addRow(T("xr_rotate"), row)
        b = QPushButton(T("xr_flip"))
        b.setToolTip(T("xr_flip_tip"))
        b.clicked.connect(lambda: self.rotate_view("y", 180))
        f.addRow(b)
        b = QPushButton(T("xr_template"))
        b.setToolTip(T("xr_template_tip"))
        b.clicked.connect(self.align_template)
        f.addRow(b)
        self.lbl_align = QLabel("")
        self.lbl_align.setWordWrap(True)
        self.lbl_align.setStyleSheet(HINT_CSS)
        f.addRow(self.lbl_align)
        return g

    def _take_enabled(self, on):
        if getattr(self, "btn_take", None) is not None:
            self.btn_take.setEnabled(on)

    def atoms_boxed(self, ks):
        """Shift + Rahmen: Menü für die Labels aller Atome im Rahmen."""
        from PySide6.QtGui import QCursor
        from PySide6.QtWidgets import QMenu
        m = QMenu(self)
        a_off = m.addAction(T("xr_box_labels_off", n=len(ks)))
        a_on = m.addAction(T("xr_box_labels_on", n=len(ks)))
        a = m.exec(QCursor.pos())
        if a is not None:
            self.set_labels(ks, a is a_on)

    def set_labels(self, ks, on):
        """Labels mehrerer Atome (Indizes im Molekül der Abbildung) an/aus; entspricht es der Vorgabe, fällt die
        Ausnahme weg."""
        e = self.entry()
        if e is None or self.F is None:
            return
        if on and not self.st.get("labels", True):
            self.st["labels"] = True
            self.cb_lab.setChecked(True)
        o = dict(self.opts(e), label_on={}, labels=True)
        lo = e.setdefault("label_on", {})
        for k in ks:
            a = self.F.mol.atoms[k]
            lo.pop(a["key"], None)
            if label_shown(a, o) != on:
                lo[a["key"]] = bool(on)
        if hasattr(self, "fill_atoms"):
            was, self._loading = self._loading, True
            self.fill_atoms()
            self._loading = was
        self.render()

    def toggle_atom(self, k):
        """Klick: Atom zur Auswahl (max. 4), erneuter Klick entfernt es, Klick ins Leere hebt sie auf."""
        if self.F is None:
            return
        if k is None:
            self.sel = []
        else:
            key = self.F.mol.atoms[k]["key"]
            if key in self.sel:
                self.sel.remove(key)
            else:
                if len(self.sel) >= 4:
                    self.sel = []
                self.sel.append(key)
        self.update_measure()
        self.render(push=False)

    def clear_selection(self):
        self.sel = []
        self.update_measure()
        self.render(push=False)

    def update_measure(self):
        e = self.entry()
        mol = self.molecule(e) if e else None
        self._meas = None
        if mol is None or not self.sel:
            self.lbl_meas.setText(T("xr_select_hint"))
            self.lbl_meas.setStyleSheet("")
            self._take_enabled(False)
            return
        names = param_name(mol, self.sel)
        m = measure(mol, self.sel) if len(self.sel) >= 2 else None
        self._meas = m
        self._take_enabled(m is not None)
        if m is None:
            self.lbl_meas.setText(names)
            self.lbl_meas.setStyleSheet("")
            return
        unit = " Å" if m["kind"] == "bond" else "°"
        txt = f"{T('xr_kind_' + m['kind'])}  {names} = {minus(m['text'])}{unit}"
        warn = self.MEASURE_ESD and m["source"] == "calc"
        if warn:
            txt += "\n" + T("xr_no_esd_torsion" if m["kind"] == "torsion" else "xr_no_esd")
        elif self.MEASURE_ESD and not m["esd"]:
            txt += "\n" + T("xr_cif_no_esd")
        self.lbl_meas.setText(txt)
        self.lbl_meas.setStyleSheet(WARN_CSS if warn else "")

    # -- Ausrichten -----------------------------------------------------------------------
    def _set_view(self, R, msg="", warn=False):
        e = self.entry()
        if e is None:
            return
        e["view"] = [round(float(v), 6) for v in np.asarray(orthonormalize(R)).ravel()]
        e["pan"] = [0.0, 0.0]
        self.lbl_align.setText(msg)
        self.lbl_align.setStyleSheet(WARN_CSS if warn else HINT_CSS)
        self.render()

    def current_view(self):
        e = self.entry()
        mol = self.molecule(e) if e else None
        return (self.view_of(e, mol), mol) if mol is not None else (None, None)

    def align_bond(self, direction):
        R, mol = self.current_view()
        if R is None:
            return
        if len(self.sel) < 2:
            self.lbl_align.setText(T("xr_align_need2"))
            self.lbl_align.setStyleSheet(WARN_CSS)
            return
        keys = self.sel[:2]
        self.sel = []
        self.update_measure()
        self._set_view(align_bond(R, mol, keys, direction),
                       T("xr_align_done_bond", bond=param_name(mol, keys, sep="→")))

    def align_plane(self):
        R, mol = self.current_view()
        if R is None:
            return
        if len(self.sel) < 3:
            self.lbl_align.setText(T("xr_align_need3"))
            self.lbl_align.setStyleSheet(WARN_CSS)
            return
        keys = list(self.sel)
        self.sel = []
        self.update_measure()
        self._set_view(align_plane(R, mol, keys), T("xr_align_done_plane", atoms=param_name(mol, keys, sep=", ")))

    def rotate_view(self, axis, deg):
        R, mol = self.current_view()
        if R is not None:
            self._set_view(rot(axis, math.radians(deg)) @ R)

    def template_path(self):
        """CDXML der Verbindung des Projekts (auch neben SVG/PDF/PNG), sonst None."""
        st = (self.main.compound or {}).get("structure")
        if st:
            p = Path(st)
            p = p if p.suffix.lower() == ".cdxml" else p.with_suffix(".cdxml")
            if p.exists():
                return p
        return None

    def align_template(self, path=None):
        R, mol = self.current_view()
        if R is None:
            return
        p = path or self.template_path()
        if p is None:
            p, _ = QFileDialog.getOpenFileName(self, T("xr_template"), self.main.settings.value("last_img_dir", ""),
                                               "ChemDraw (*.cdxml)")
            if not p:
                return
        try:
            Rn, info = align_to_template(mol, p)
        except Exception as ex:
            self.lbl_align.setText(T("xr_template_fail", err=str(ex)))
            self.lbl_align.setStyleSheet(WARN_CSS)
            return
        msg = T("xr_template_done", n=info["n"], nt=info["n_template"], rmsd=f"{info['rmsd']:.2f}", f=Path(p).name)
        warn = info["ambiguous"] or info["n"] < info["n_template"] or info["rmsd"] > 0.6
        if info["ambiguous"]:
            msg += "\n" + T("xr_template_amb")
        e = self.entry()
        if e is not None:
            e["layout"] = {}                              # Labels für die neue Ansicht neu setzen
        self._set_view(Rn, msg, warn)


class XrayTab(MeasureAlign, TABS.FigureTab):
    STATE_KEY = "xray"
    GUIDE, GUIDE_ACTION = "xray", ("guide_add", "add_cif")
    TEXT_OPTS = ("fonts", "labels")              # Labelgröße hat der Tab selbst
    DEFAULT_STATE = dict(size="half_a4", w_cm=16.0, h_cm=11.0, fmt={"pdf": True, "svg": True, "png": True},
                         layout={}, entries=[], current=None, prob=50, show_h=True, disorder="main", labels=True,
                         h_labels=False, colors={}, bond_w=0.12, style="mercury", x_auto=True, x_min=0.0,
                         x_max=1.0, ymax=0.0, label_mode="all", hide_solv=False, hide_ions=False, wire_h=False,
                         bg="white", label_color="black")


    def arrange(self):
        super().arrange()
        g = STYLE.split_group(self.group_titled("xr_display"), self.lst_frag, None, T("grp_atoms"), n=4)
        self.fold(g, "atoms")
        self.fold(self.group_titled("xr_align"), "align")
        self.fold(self.group_titled("xr_colors"), "colors", ("colors",))
        self.to_right(self.group_titled("xr_caption"), T("xr_caption"))
        self.to_right(self.group_titled("xr_tables"), T("xr_tables"))
    def __init__(self, main, state):
        self.sel = []                     # gewählte Atome (Schlüssel) der aktuellen Verbindung, nicht im Undo
        self._meas = None
        self._mol_cache = {}
        super().__init__(main, state)

    # -- Bedienelemente -------------------------------------------------------------------
    def build_controls(self):
        st = self.st
        g = QGroupBox(T("xr_structures"))
        v = QVBoxLayout(g)
        row = QHBoxLayout()
        for txt, fn in ((T("xr_add"), self.add_cif), ("↑", lambda: self.move_entry(-1)),
                        ("↓", lambda: self.move_entry(1)), (T("remove"), self.remove_entry)):
            b = QPushButton(txt)
            b.clicked.connect(fn)
            row.addWidget(b)
        row.addStretch(1)
        v.addLayout(row)
        self.tbl = QTableWidget(0, 3)
        self.tbl.setHorizontalHeaderLabels([T("xr_name"), T("xr_ref"), "CCDC"])
        self.tbl.verticalHeader().setVisible(False)
        self.tbl.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.tbl.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeToContents)   # „Interne Referenz“ ganz
        self.tbl.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.tbl.setSelectionMode(QAbstractItemView.SingleSelection)
        self.tbl.setMinimumHeight(110)
        self.tbl.itemChanged.connect(self.entry_edited)
        self.tbl.itemSelectionChanged.connect(self.entry_selected)
        v.addWidget(self.tbl)
        self.lbl_info = QLabel("")
        self.lbl_info.setWordWrap(True)
        self.lbl_info.setTextFormat(Qt.RichText)
        self.lbl_info.setStyleSheet(HINT_CSS)
        v.addWidget(self.lbl_info)
        self.lv.addWidget(g)

        g = QGroupBox(T("xr_display"))
        f = make_form(g)
        self.cmb_style = QComboBox()
        for key in XRAY_STYLES:
            self.cmb_style.addItem(T("xr_style_" + key), key)
        self.cmb_style.setCurrentIndex(max(0, self.cmb_style.findData(allowed_style(st.get("style", "mercury"), XRAY_STYLES,
                                                                                    "mercury"))))
        f.addRow(T("xr_style"), self.cmb_style)
        self.sp_prob = QSpinBox(minimum=10, maximum=99, singleStep=5, suffix=" %", value=int(st["prob"]))
        self.sp_prob.setToolTip(T("xr_prob_tip"))
        f.addRow(T("xr_prob"), self.sp_prob)
        self.cmb_style.currentIndexChanged.connect(self._prob_for_style)
        self.cmb_dis = QComboBox()
        for key in ("main", "all", "minor"):
            self.cmb_dis.addItem(T("xr_dis_" + key), key)
        self.cmb_dis.setCurrentIndex(max(0, self.cmb_dis.findData(st["disorder"])))
        f.addRow(T("xr_disorder"), self.cmb_dis)
        self.cb_h = QCheckBox(T("xr_show_h"))
        self.cb_h.setChecked(st["show_h"])
        self.cb_lab = QCheckBox(T("xr_labels"))
        self.cb_lab.setChecked(st["labels"])
        self.cb_hlab = QCheckBox(T("xr_h_labels"))
        self.cb_hlab.setChecked(st["h_labels"])
        self.cb_wire_h = QCheckBox(T("xr_wire_h"))
        self.cb_wire_h.setToolTip(T("xr_wire_h_tip"))
        self.cb_wire_h.setChecked(st.get("wire_h", False))
        row = QHBoxLayout()
        for w_ in (self.cb_h, self.cb_lab, self.cb_hlab, self.cb_wire_h):
            row.addWidget(w_)
        row.addStretch(1)
        f.addRow(row)
        self.cmb_labmode = QComboBox()
        for key in ("all", "hetero"):
            self.cmb_labmode.addItem(T("xr_lab_" + key), key)
        f.addRow(T("xr_label_mode"), self.cmb_labmode)
        self.sp_bond = QDoubleSpinBox(decimals=2, minimum=0.02, maximum=0.4, singleStep=0.01, suffix=" Å",
                                      value=st.get("bond_w", 0.08))
        self.sp_bond.setToolTip(T("ball_rel_tip", std=f"{self.DEFAULT_STATE['bond_w']:g}"))
        f.addRow(T("xr_bond_lw"), self.sp_bond)
        self.cmb_bg = BackgroundCombo(st, self.changed)
        f.addRow(T("bg_label"), self.cmb_bg)
        self.sp_lab = label_size_spin(st)
        f.addRow(T("xr_label_size"), self.sp_lab)
        self.cmb_lcol = label_color_combo(st, self.changed)
        f.addRow(T("xr_label_color"), self.cmb_lcol)
        self.lst_frag = QListWidget()
        self.lst_frag.setMaximumHeight(90)
        self.lst_frag.setToolTip(T("xr_frags_tip"))
        self.lst_frag.itemChanged.connect(self.frag_toggled)
        self.cb_solv = QCheckBox(T("xr_hide_solv"))
        self.cb_ions = QCheckBox(T("xr_hide_ions"))
        self.cb_solv.setToolTip(T("xr_hide_tip"))
        self.cb_ions.setToolTip(T("xr_hide_tip"))
        f.addRow(self.cb_solv)                      # untereinander: nebeneinander bestimmten sie die Spaltenbreite
        f.addRow(self.cb_ions)                      # (unter Windows ~600 px, Abbildung zu schmal)
        f.addRow(T("xr_frags"), self.lst_frag)
        self.tbl_atoms = QTableWidget(0, 4)
        self.tbl_atoms.setHorizontalHeaderLabels([T("xr_atom"), T("xr_atom_shown"), T("xr_atom_label"),
                                                  T("xr_atom_wire")])
        self.tbl_atoms.verticalHeader().setVisible(False)
        self.tbl_atoms.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.tbl_atoms.setMinimumHeight(150)
        self.tbl_atoms.setToolTip(T("xr_atoms_tip"))
        self.tbl_atoms.itemChanged.connect(self.atom_toggled)
        f.addRow(T("xr_atoms"), self.tbl_atoms)
        row = QHBoxLayout()
        for txt, fn in ((T("xr_labels_all"), lambda: self.set_all_labels(True)),
                        (T("xr_labels_none"), lambda: self.set_all_labels(False)),
                        (T("xr_atoms_reset"), self.reset_atom_overrides)):
            b = QPushButton(txt)
            b.clicked.connect(fn)
            row.addWidget(b)
        row.addStretch(1)
        f.addRow(row)
        f.addRow(reset_buttons(self))
        hint = QLabel(T("xr_mouse_hint"))
        hint.setWordWrap(True)
        hint.setStyleSheet(HINT_CSS)
        f.addRow(hint)
        self.lv.addWidget(g)

        self.lv.addWidget(self.align_box())

        g = QGroupBox(T("xr_colors"))
        v = QVBoxLayout(g)
        self.color_grid = QGridLayout()
        v.addLayout(self.color_grid)
        b = QPushButton(T("xr_colors_reset"))
        b.clicked.connect(self.reset_colors)
        v.addWidget(b)
        self.lv.addWidget(g)

        g, v = self.measure_box(T("xr_select"), take=self.take_param)
        row = QHBoxLayout()
        self.btn_wire = QPushButton(T("xr_wire"))
        self.btn_wire.setToolTip(T("xr_wire_tip"))
        self.btn_wire.clicked.connect(self.wire_selected)
        b = QPushButton(T("xr_unwire"))
        b.setToolTip(T("xr_unwire_tip"))
        b.clicked.connect(self.unwire_selected)
        row.addWidget(self.btn_wire)
        row.addWidget(b)
        row.addStretch(1)
        v.addLayout(row)
        row = QHBoxLayout()
        row.addWidget(QLabel(T("xr_wire_name")))
        self.ed_wire = QLineEdit()
        self.ed_wire.setPlaceholderText(T("xr_wire_name_ph"))
        self.ed_wire.editingFinished.connect(self.wire_name_changed)
        row.addWidget(self.ed_wire, 1)
        v.addLayout(row)
        self.lst_par = QListWidget()
        self.lst_par.setMaximumHeight(130)
        v.addWidget(self.lst_par)
        row = QHBoxLayout()
        for txt, fn in (("↑", lambda: self.move_param(-1)), ("↓", lambda: self.move_param(1)),
                        (T("remove"), self.remove_param)):
            b = QPushButton(txt)
            b.clicked.connect(fn)
            row.addWidget(b)
        row.addStretch(1)
        v.addLayout(row)
        self.lv.addWidget(g)

        g = QGroupBox(T("xr_caption"))
        v = QVBoxLayout(g)
        self.txt_cap = QTextBrowser()
        self.txt_cap.setMaximumHeight(110)
        v.addWidget(self.txt_cap)
        b = QPushButton(T("xr_caption_copy"))
        b.clicked.connect(lambda: copy_rich(*self.caption()))
        v.addWidget(b)
        self.lv.addWidget(g)

        g = QGroupBox(T("xr_tables"))
        v = QVBoxLayout(g)
        self.txt_tab = QTextBrowser()
        self.txt_tab.setMinimumHeight(200)
        v.addWidget(self.txt_tab)
        row = QHBoxLayout()
        for txt, fn in ((T("xr_copy_crystal"), lambda: copy_rich(*self.crystal_clip())),
                        (T("xr_copy_params"), lambda: copy_rich(*self.params_clip())),
                        (T("xr_excel"), self.save_excel)):
            b = QPushButton(txt)
            b.clicked.connect(fn)
            row.addWidget(b)
        row.addStretch(1)
        v.addLayout(row)
        self.lv.addWidget(g)

        self.sp_prob.valueChanged.connect(self.changed)
        self.sp_bond.valueChanged.connect(self.changed)
        self.sp_lab.valueChanged.connect(self.changed)
        self.cmb_dis.currentIndexChanged.connect(self.disorder_changed)
        self.cmb_style.currentIndexChanged.connect(self.changed)
        for w_ in (self.cb_h, self.cb_lab, self.cb_hlab, self.cb_wire_h, self.cb_solv, self.cb_ions):
            w_.toggled.connect(self.changed)
        self.cmb_labmode.currentIndexChanged.connect(self.changed)
        self.sync_widgets()

    def common_box(self):
        """Nur Größe und Formate (keine Achsen)."""
        g = QGroupBox(T("grp_export_size"))
        f = make_form(g)
        st = self.st
        self.cmb_size = QComboBox()
        for key in ("half_a4", "custom"):
            self.cmb_size.addItem(T({"half_a4": "size_half_a4", "custom": "size_custom"}[key]), key)
        self.cmb_size.setCurrentIndex(max(0, self.cmb_size.findData(st["size"])))
        f.addRow(T("size"), self.cmb_size)
        row = QHBoxLayout()
        self.sp_w = QDoubleSpinBox(decimals=1, minimum=4, maximum=40, singleStep=0.5, suffix=" cm", value=st["w_cm"])
        self.sp_h = QDoubleSpinBox(decimals=1, minimum=3, maximum=40, singleStep=0.5, suffix=" cm", value=st["h_cm"])
        row.addWidget(self.sp_w)
        row.addWidget(QLabel("×"))
        row.addWidget(self.sp_h)
        f.addRow(T("size_wh"), row)
        row = QHBoxLayout()
        self.cb_fmt = {}
        for e in ("pdf", "svg", "png"):
            cb = QCheckBox(e.upper())
            cb.setChecked(st["fmt"].get(e, True))
            self.cb_fmt[e] = cb
            row.addWidget(cb)
        row.addStretch(1)
        f.addRow(T("formats"), row)
        for w_ in (self.sp_w, self.sp_h):
            w_.valueChanged.connect(self.changed)
        self.cmb_size.currentIndexChanged.connect(self.changed)
        for cb in self.cb_fmt.values():
            cb.toggled.connect(self.changed)
        self._sync_common()
        return g

    def _sync_common(self):
        if self.cmb_size.currentData() == "half_a4":
            for sp, v in ((self.sp_w, 16.0), (self.sp_h, 11.0)):
                sp.blockSignals(True)
                sp.setValue(v)
                sp.blockSignals(False)
        custom = self.cmb_size.currentData() == "custom"
        self.sp_w.setEnabled(custom)
        self.sp_h.setEnabled(custom)

    def collect_common(self):
        self.st.update(size=self.cmb_size.currentData(), w_cm=self.sp_w.value(), h_cm=self.sp_h.value(),
                       fmt={e: cb.isChecked() for e, cb in self.cb_fmt.items()})

    def _prob_for_style(self, *_):
        """Ellipsoid-Wahrscheinlichkeit wirkt im Kugel-Stab-Stil nicht: ausgrauen und sagen, warum."""
        ball = self.cmb_style.currentData() in BALL_STYLES
        self.sp_prob.setEnabled(not ball)
        self.sp_prob.setToolTip(T("xr_prob_ball") if ball else T("xr_prob_tip"))

    def collect(self):
        self.st.update(prob=self.sp_prob.value(), show_h=self.cb_h.isChecked(), labels=self.cb_lab.isChecked(),
                       h_labels=self.cb_hlab.isChecked(), disorder=self.cmb_dis.currentData(),
                       bond_w=self.sp_bond.value(), style=self.cmb_style.currentData(), label_pt=self.sp_lab.value(),
                       label_mode=self.cmb_labmode.currentData(), wire_h=self.cb_wire_h.isChecked(),
                       label_color=self.cmb_lcol.currentData(),
                       hide_solv=self.cb_solv.isChecked(), hide_ions=self.cb_ions.isChecked())
        self.cmb_labmode.setEnabled(self.cb_lab.isChecked())

    # -- Einträge -------------------------------------------------------------------------
    def entries(self):
        return self.st["entries"]

    def entry(self):
        es = self.entries()
        cur = self.st.get("current")
        for e in es:
            if e["uid"] == cur:
                return e
        return es[0] if es else None

    def add_cif(self):
        if not self.main.ensure_project():
            return
        ps, _ = QFileDialog.getOpenFileNames(self, T("xr_add"), self.main.settings.value("last_dir", ""),
                                             f"CIF (*.cif);;{T('flt_all')}")
        for p in ps:
            self.load_cif(p)

    def load_cif(self, p):
        """Jeder data_-Block mit Atomen wird ein Eintrag (eine Spalte in den Tabellen)."""
        p = self.main.import_file(p, PRJ.data_sub(self.STATE_KEY))
        try:
            blocks = self.cache(("cif", p), lambda: structure_blocks(p))
            if not blocks:
                raise ValueError(T("xr_no_atoms"))
            for b in blocks:
                Structure(b)
        except Exception as e:
            self.main.error(f"{Path(p).name}: {e}")
            return
        last = None
        for b in blocks:
            # Name: bei einem Block der Dateiname (Datenblöcke heißen oft nach dem Messlauf, z. B. „…sadabs“)
            name = Path(p).stem if len(blocks) == 1 else b.name
            e = dict(uid=uuid.uuid4().hex[:8], file=p, block=b.name, name=name, ref="", ccdc="", view=None,
                     zoom=1.0, pan=[0.0, 0.0], layout={}, hidden_frags=[], params=[], hidden_atoms=[],
                     label_on={}, wire=[], wire_name="")
            self.entries().append(e)
            last = e["uid"]
        self.st["current"] = last
        self.sel = []
        self.sync_widgets()
        self.render()

    def remove_entry(self):
        e = self.entry()
        if e is None:
            return
        self.st["entries"] = [x for x in self.entries() if x["uid"] != e["uid"]]
        self.st["current"] = self.entries()[0]["uid"] if self.entries() else None
        self.sel = []
        self.sync_widgets()
        self.render()

    def move_entry(self, d):
        e = self.entry()
        if e is None:
            return
        es = self.entries()
        i = es.index(e)
        j = i + d
        if 0 <= j < len(es):
            es[i], es[j] = es[j], es[i]
            self.sync_widgets()
            self.render()

    def entry_edited(self, item):
        if self._loading:
            return
        es = self.entries()
        r = item.row()
        if r >= len(es):
            return
        key = ("name", "ref", "ccdc")[item.column()]
        es[r][key] = item.text().strip()
        self.update_tables()
        self.main.push_undo()

    def entry_selected(self):
        if self._loading:
            return
        rows = self.tbl.selectionModel().selectedRows()
        if not rows:
            return
        es = self.entries()
        r = rows[0].row()
        if r < len(es) and es[r]["uid"] != self.st.get("current"):
            self.st["current"] = es[r]["uid"]
            self.sel = []
            self.sync_widgets()
            self.render()

    # -- Modell ---------------------------------------------------------------------------
    def structure(self, e):
        p = e["file"]
        if not p or not Path(p).exists():
            return None
        blocks = self.cache(("cif", p), lambda: structure_blocks(p))
        for b in blocks:
            if b.name == e["block"]:
                return self.cache(("cifstruct", p, b.name), lambda: Structure(b))
        return None

    def molecule(self, e):
        S = self.structure(e)
        if S is None:
            return None
        key = (e["file"], e["block"], self.st["disorder"])
        if key not in self._mol_cache:
            self._mol_cache[key] = Molecule(S, self.st["disorder"])
        return self._mol_cache[key]

    def view_of(self, e, mol):
        v = e.get("view")
        if v and len(v) == 9:
            return np.array(v, float).reshape(3, 3)
        o = self.opts(e)
        hid, kinds, wire = set(o["hidden_frags"]), set(o["hide_kinds"]), o["wire"]
        use = [k for k, a in enumerate(mol.atoms) if a["el"] not in ("H", "D") and a["key"] not in wire
               and a["key"] not in set(o["hidden_atoms"]) and a.get("frag") not in hid
               and mol.fragments[a["frag"]]["kind"] not in kinds]
        return default_view(mol, use=use or None)        # alle sichtbaren Teile (auch Lösungsmittel/Ionen)

    # -- Oberfläche -----------------------------------------------------------------------
    def sync_widgets(self):
        self._loading = True
        st = self.st
        self.cmb_bg.sync()
        es = self.entries()
        self.tbl.setRowCount(len(es))
        cur = self.entry()
        for r, e in enumerate(es):
            for c, k in enumerate(("name", "ref", "ccdc")):
                val = e.get(k) or ""
                if k == "ccdc" and not val:
                    S = self.structure(e)
                    val = ""
                    if S is not None:
                        val = crystal_values(S.block)["ccdc"]
                it = QTableWidgetItem(val)
                if k == "name":
                    it.setToolTip(f"{Path(e['file']).name}: data_{e['block']}")
                self.tbl.setItem(r, c, it)
        if cur is not None:
            self.tbl.selectRow(es.index(cur))
        for w_, v in ((self.sp_prob, int(st["prob"])), (self.sp_bond, st.get("bond_w", 0.08)),
                      (self.sp_lab, st.get("label_pt", 0.0))):
            w_.setValue(v)
        self.cmb_dis.setCurrentIndex(max(0, self.cmb_dis.findData(st["disorder"])))
        self.cb_h.setChecked(st["show_h"])
        self.cb_lab.setChecked(st["labels"])
        self.cb_hlab.setChecked(st["h_labels"])
        self.cb_wire_h.setChecked(st.get("wire_h", False))
        self.cb_solv.setChecked(st.get("hide_solv", False))
        self.cb_ions.setChecked(st.get("hide_ions", False))
        self.cmb_labmode.setCurrentIndex(max(0, self.cmb_labmode.findData(st.get("label_mode", "all"))))
        self.cmb_labmode.setEnabled(st["labels"])
        self.cmb_lcol.setCurrentIndex(max(0, self.cmb_lcol.findData(st.get("label_color", "black"))))
        self.ed_wire.setText((cur or {}).get("wire_name") or "")
        self.cmb_style.setCurrentIndex(max(0, self.cmb_style.findData(allowed_style(st.get("style", "mercury"), XRAY_STYLES,
                                                                                    "mercury"))))
        self._prob_for_style()
        self.fill_colors()
        self.fill_atoms()
        self.fill_frags()
        self.fill_params()
        self._loading = False

    def fill_colors(self):
        while self.color_grid.count():
            w_ = self.color_grid.takeAt(0).widget()
            if w_ is not None:
                w_.deleteLater()
        els = set()
        for e in self.entries():
            S = self.structure(e)
            if S is not None:
                els |= {a["el"] for a in S.atoms}
        order = [x for x in ("C", "H") if x in els] + sorted(els - {"C", "H"})
        cols = self.st.get("colors") or {}
        for n, el in enumerate(order):
            box = QWidget()
            h = QHBoxLayout(box)
            h.setContentsMargins(0, 0, 0, 0)
            h.addWidget(QLabel(el))
            h.addWidget(TABS.color_button(cols.get(el) or default_color(el, self.st.get("style")),
                                          lambda c, el=el: self.set_color(el, c)))
            self.color_grid.addWidget(box, n // 4, n % 4)

    def set_color(self, el, c):
        self.st.setdefault("colors", {})[el] = c
        self.render()

    def reset_colors(self):
        self.st["colors"] = {}
        self.fill_colors()
        self.render()

    def fill_frags(self):
        self.lst_frag.clear()
        e = self.entry()
        mol = self.molecule(e) if e else None
        if mol is None:
            return
        hidden = set(e.get("hidden_frags") or [])
        for i, fr in enumerate(mol.fragments):
            kind = fr.get("kind")
            tag = {"main": T("xr_main_frag"), "molecule": T("xr_kind_molecule"),
                   "solvent": T("xr_kind_solvent"), "ion": T("xr_kind_ion")}.get(kind, "")
            if fr.get("name") and kind in ("solvent", "ion"):
                tag = f"{fr['name']}, {tag}"
            it = QListWidgetItem(f"{i + 1}: {fr['formula']}" + (f"  ({tag})" if tag else ""))
            it.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled)
            it.setCheckState(Qt.Unchecked if i in hidden else Qt.Checked)
            it.setData(Qt.UserRole, i)
            self.lst_frag.addItem(it)

    def frag_toggled(self, item):
        if self._loading:
            return
        e = self.entry()
        hidden = set(e.get("hidden_frags") or [])
        i = item.data(Qt.UserRole)
        hidden.discard(i) if item.checkState() == Qt.Checked else hidden.add(i)
        e["hidden_frags"] = sorted(hidden)
        self.timer.start()

    def fill_atoms(self):
        """Je Atom: sichtbar (auch einzelne H) und Label (überschreibt die Vorgabe)."""
        was = self._loading
        self._loading = True
        e = self.entry()
        mol = self.molecule(e) if e else None
        kinds = set(self.opts(e)["hide_kinds"]) if e else set()
        atoms = [] if mol is None else [a for a in mol.atoms if a.get("frag") not in set(e.get("hidden_frags") or [])
                                        and mol.fragments[a["frag"]]["kind"] not in kinds
                                        and (self.st["show_h"] or a["el"] not in ("H", "D"))]
        wire = set((e or {}).get("wire") or [])
        self.tbl_atoms.setRowCount(len(atoms))
        hidden = set((e or {}).get("hidden_atoms") or [])
        o = self.opts(e) if e else {}
        for r, a in enumerate(atoms):
            it = QTableWidgetItem(atom_name(a))
            it.setFlags(Qt.ItemIsEnabled)
            it.setData(Qt.UserRole, a["key"])
            self.tbl_atoms.setItem(r, 0, it)
            for c, on in ((1, a["key"] not in hidden), (2, label_shown(a, o)), (3, a["key"] in wire)):
                it = QTableWidgetItem()
                it.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled)
                it.setCheckState(Qt.Checked if on else Qt.Unchecked)
                it.setData(Qt.UserRole, a["key"])
                self.tbl_atoms.setItem(r, c, it)
        self.tbl_atoms.resizeColumnsToContents()
        self._loading = was

    def atom_toggled(self, item):
        if self._loading or item.column() == 0:
            return
        e = self.entry()
        if e is None:
            return
        key, on = item.data(Qt.UserRole), item.checkState() == Qt.Checked
        if item.column() == 1:
            hidden = set(e.get("hidden_atoms") or [])
            hidden.discard(key) if on else hidden.add(key)
            e["hidden_atoms"] = sorted(hidden)
            self.sel = [k for k in self.sel if k not in hidden]
        elif item.column() == 3:
            wire = set(e.get("wire") or [])
            wire.add(key) if on else wire.discard(key)
            e["wire"] = sorted(wire)
        else:
            self.set_label(key, on)
            return
        self.timer.start()

    def set_label(self, key, on):
        """Label je Atom; entspricht es der Vorgabe, wird die Ausnahme gelöscht."""
        e = self.entry()
        mol = self.molecule(e) if e else None
        if mol is None:
            return
        k = mol.find_key(key)
        lo = e.setdefault("label_on", {})
        lo.pop(key, None)
        if k is not None and label_shown(mol.atoms[k], self.opts(e)) != on:
            lo[key] = bool(on)
        self.timer.start()

    def atom_double_clicked(self, k):
        """Doppelklick: Label dieses Atoms an/aus (der erste Klick des Doppelklicks wählt nicht aus)."""
        if self.F is None:
            return
        a = self.F.mol.atoms[k]
        if a["key"] in self.sel:
            self.sel.remove(a["key"])
            self.update_measure()
        self.set_label(a["key"], not label_shown(a, self.opts(self.entry())))
        self._loading = True
        self.fill_atoms()
        self._loading = False

    def set_all_labels(self, on):
        e = self.entry()
        mol = self.molecule(e) if e else None
        if mol is None:
            return
        o = dict(self.opts(e), label_on={})
        if not self.st["labels"]:
            self.st["labels"] = True
            self.cb_lab.setChecked(True)
        e["label_on"] = {a["key"]: on for a in mol.atoms if label_shown(a, o) != on}
        self.fill_atoms()
        self.render()

    def reset_atom_overrides(self):
        e = self.entry()
        if e is not None:
            e.update(hidden_atoms=[], label_on={})
            self.fill_atoms()
            self.render()

    def disorder_changed(self, *_):
        if self._loading:
            return
        self.sel = []
        self.changed()
        self._loading = True
        self.fill_frags()
        self._loading = False

    def fill_params(self):
        self.lst_par.clear()
        e = self.entry()
        mol = self.molecule(e) if e else None
        for p in (e or {}).get("params") or []:
            m = measure(mol, p["atoms"]) if mol else None
            name = param_name(mol, p["atoms"]) if mol else param_label(p["atoms"])
            txt = f"{name}   {minus(m['text']) if m else '–'}" + ("" if (m is None or m["source"] == "cif")
                                                                    else "  *")
            self.lst_par.addItem(txt)

    def move_param(self, d):
        e = self.entry()
        i = self.lst_par.currentRow()
        ps = (e or {}).get("params") or []
        j = i + d
        if 0 <= i < len(ps) and 0 <= j < len(ps):
            ps[i], ps[j] = ps[j], ps[i]
            self.fill_params()
            self.lst_par.setCurrentRow(j)
            self.update_tables()
            self.main.push_undo()

    def remove_param(self):
        e = self.entry()
        i = self.lst_par.currentRow()
        ps = (e or {}).get("params") or []
        if 0 <= i < len(ps):
            del ps[i]
            self.fill_params()
            self.update_tables()
            self.main.push_undo()

    # -- Drahtmodell ------------------------------------------------------------------------
    def wire_selected(self):
        """Gewählte Bindung: der Substituent dahinter (kleinere Seite) wird Drahtmodell; sonst die gewählten Atome."""
        e = self.entry()
        mol = self.molecule(e) if e else None
        if mol is None or not self.sel:
            self.main.statusBar().showMessage(T("xr_wire_hint"), 6000)
            return
        idx = [mol.find_key(k) for k in self.sel]
        idx = [k for k in idx if k is not None]
        add = set(idx)
        if len(idx) == 2 and (min(idx), max(idx)) in set(mol.bonds):
            side = mol.substituent(idx[0], idx[1])
            if side is None:
                self.main.statusBar().showMessage(T("xr_wire_ring"), 8000)
                return
            add = side
        e["wire"] = sorted(set(e.get("wire") or []) | {mol.atoms[k]["key"] for k in add})
        self.sel = []
        self.update_measure()
        self.render()

    def unwire_selected(self):
        """Gewählte Atome (bzw. Substituent hinter der gewählten Bindung) wieder normal; ohne Auswahl alle."""
        e = self.entry()
        mol = self.molecule(e) if e else None
        if mol is None:
            return
        if not self.sel:
            e["wire"] = []
        else:
            idx = [k for k in (mol.find_key(x) for x in self.sel) if k is not None]
            drop = set(idx)
            if len(idx) == 2 and (min(idx), max(idx)) in set(mol.bonds):
                drop = mol.substituent(idx[0], idx[1]) or drop
            keys = {mol.atoms[k]["key"] for k in drop}
            e["wire"] = [k for k in e.get("wire") or [] if k not in keys]
        self.sel = []
        self.update_measure()
        self.render()

    def wire_name_changed(self):
        e = self.entry()
        if e is not None and self.ed_wire.text().strip() != (e.get("wire_name") or ""):
            e["wire_name"] = self.ed_wire.text().strip()
            self.update_tables()
            self.main.push_undo()

    # -- Bildunterschrift -------------------------------------------------------------------
    def caption(self, e=None):
        """Bildunterschrift im Stil der SI (ohne „Figure Sx:“) -> (Klartext, HTML). Folgt der Abbildung: Fehlordnung
        mit Besetzung, Ausgeblendetes, Drahtmodell; bis CAPTION_MAX ausgewählte Werte stehen in der Unterschrift,
        mehr kommen in eigene Tabellen (param_tables)."""
        import html as _html
        e = e or self.entry()
        mol = self.molecule(e) if e else None
        if mol is None:
            return "", ""
        st, o = self.st, self.opts(e)
        name = e.get("name") or e["block"]
        esc = _html.escape
        shown = set(XrayFigure._visible.__get__(_Vis(mol, o))())
        omit, solv, ions = [], [], []
        hid = set(o["hidden_frags"])
        for fi, fr in enumerate(mol.fragments):
            if not (fi in hid or fr["kind"] in set(o["hide_kinds"])) or fi == 0:
                continue
            lab = fr.get("name") or to_plain(formula_markup(fr["formula"]))
            if fr["kind"] == "solvent":
                solv.append(lab)
            elif fr["kind"] == "ion":
                ions.append(lab)
            elif fr["kind"] == "molecule":
                omit.append("the second independent molecule")
            else:
                omit.append("further fragments")
        heavy_hidden = [a for a in mol.atoms if a["key"] in set(o["hidden_atoms"]) and a["el"] not in ("H", "D")
                        and a.get("frag") not in hid]
        if not st["show_h"]:
            omit.insert(0, "hydrogen atoms")
        elif any(a["el"] in ("H", "D") and a["key"] in set(o["hidden_atoms"]) for a in mol.atoms):
            omit.insert(0, "selected hydrogen atoms")
        if heavy_hidden:
            omit.append("selected atoms")
        if solv:
            omit.append("co-crystallised " + join_en(sorted(set(solv))) if len(set(solv)) <= 2
                        else "co-crystallised solvent molecules")
        if ions:
            u = sorted(set(ions))
            omit.append(("the counterion " if len(u) == 1 else "the counterions ") + join_en(u))
        wire = [k for k in range(len(mol.atoms)) if mol.atoms[k]["key"] in o["wire"] and k in shown]
        occ = disorder_occupancy(mol.S)
        head = f"X-ray solid state structure of {name}"
        head_h = f"X-ray solid state structure of <b>{esc(name)}</b>"
        if occ and st["disorder"] in ("main", "minor"):
            extra = f" ({occ} occupancy; {'major' if st['disorder'] == 'main' else 'minor'} part shown)"
            head, head_h = head + extra, head_h + esc(extra)
        elif occ:
            head, head_h = head + f" ({occ} occupancy)", head_h + esc(f" ({occ} occupancy)")
        tail = f". Thermal ellipsoids are shown with {int(st['prob'])}% probability."
        clauses = []
        if omit:
            clauses.append(join_en(omit) + (" are" if len(omit) > 1 or omit[0].endswith("s") else " is") + " omitted")
        if wire:
            wn = e.get("wire_name") or "selected substituents"
            extra = " without hydrogen atoms" if st["show_h"] and not st.get("wire_h") else ""
            clauses.append(f"{wn} {'are' if wn.endswith('s') else 'is'} shown as wireframe{extra}")
        if clauses:
            sent = " and ".join(clauses) + " for clarity."
            tail += " " + sent[0].upper() + sent[1:]
        vals = self.param_values(e, mol)
        if vals and len(vals) <= CAPTION_MAX:
            kinds = [k for k in ("bond", "angle", "torsion") if any(v[0] == k for v in vals)]
            what = {("bond",): "bond lengths in [Å]", ("angle",): "angles in [°]", ("torsion",): "torsion angles in [°]",
                    ("bond", "angle"): "bond lengths and angles in [Å] and [°]",
                    ("bond", "torsion"): "bond lengths and torsion angles in [Å] and [°]",
                    ("angle", "torsion"): "angles and torsion angles in [°]",
                    ("bond", "angle", "torsion"): "bond lengths, angles and torsion angles in [Å] and [°]"}[tuple(kinds)]
            tail += f" Selected {what}: " + "; ".join(f"{n}: {v}" for _, n, v, _ in vals) + "."
        used = {mol.atoms[k].get("opkey") for k in shown if mol.atoms[k].get("opkey")}
        for pp in e.get("params") or []:
            for k in pp["atoms"]:
                i = mol.find_key(k)
                if i is not None and mol.atoms[i].get("opkey"):
                    used.add(mol.atoms[i]["opkey"])
        codes = [f"{r}: {op}" for op, r in mol.ops.items() if op in used]
        if codes:
            tail += " Symmetry code" + ("s" if len(codes) > 1 else "") + ": " + "; ".join(codes) + "."
        return minus(head + tail), head_h + esc(minus(tail))

    def param_values(self, e, mol=None):
        """Ausgewählte Parameter einer Struktur -> [(Art, Name, Wert (CIF-Text, echtes Minus), aus der CIF?)]."""
        mol = mol or self.molecule(e)
        out = []
        for kind in ("bond", "angle", "torsion"):
            for pp in e.get("params") or []:
                m = measure(mol, pp["atoms"]) if mol else None
                if m is not None and m["kind"] == kind:
                    out.append((kind, param_name(mol, cif_order(mol, pp["atoms"])), minus(m["text"]),
                                m["source"] == "cif"))
        return out

    def param_tables(self, e):
        """Eigene Tabellen wie „Table S1: Selected bond lengths for …“, nur bei mehr als CAPTION_MAX Werten:
        [(Titel-Rest, Kopf, Zeilen)], Zellen als Markup; Werte ohne esd mit Fußnote a."""
        mol = self.molecule(e)
        if mol is None:
            return []
        vals = self.param_values(e, mol)
        if len(vals) <= CAPTION_MAX:
            return []
        name = e.get("name") or e["block"]
        out = []
        for kind, title, unit, n in (("bond", "Selected bond lengths", "Length / Å", 2),
                                     ("angle", "Selected bond angles", "Angle / °", 3),
                                     ("torsion", "Selected torsion angles", "Torsion angle / °", 4)):
            rows = []
            for pp in e.get("params") or []:
                m = measure(mol, pp["atoms"])
                if m is None or m["kind"] != kind:
                    continue
                names = [atom_name(mol.atoms[mol.find_key(k)]) for k in cif_order(mol, pp["atoms"])]
                rows.append(names + [minus(m["text"]) + ("^{a}" if m["source"] == "calc" else "")])
            if rows:
                out.append((f"{title} for", name, ["Atom"] * n + [unit], rows))
        return out

    # -- Auswahl / Maus ---------------------------------------------------------------------
    def hover_atom(self, k):
        if self.F is None or k is None:
            self.main.show_hover("")
            return
        a = self.F.mol.atoms[k]
        self.main.show_hover(f"{atom_name(a)}  ({a['el']})" + (f"   {a['opkey']}" if a.get("opkey") else ""))

    def atom_clicked(self, k):
        self.toggle_atom(k)

    def take_param(self):
        e = self.entry()
        if e is None or self._meas is None:
            return
        keys = list(self.sel)
        if not any(p["atoms"] in (keys, keys[::-1]) for p in e["params"]):
            e["params"].append({"atoms": keys})
        self.sel = []
        self.update_measure()
        self.fill_params()
        self.render()

    def view_changed(self, R):
        e = self.entry()
        if e is not None:
            e["view"] = [round(float(v), 6) for v in np.asarray(R).ravel()]
            self.render()                          # nicht verschobene Labels werden neu platziert

    def pan_changed(self, pan):
        e = self.entry()
        if e is not None:
            e["pan"] = [round(float(pan[0]), 3), round(float(pan[1]), 3)]
            self.render()

    def zoom_changed(self, z):
        e = self.entry()
        if e is not None:
            e["zoom"] = round(float(z), 4)
            self.timer.start()

    def labels_moved(self, layout):
        """Nur tatsächlich verschobene Labels speichern; die übrigen bleiben automatisch platziert."""
        e = self.entry()
        if e is None or self.F is None:
            return
        lay = self.layout()
        for key, pos in layout.items():
            if key in lay or pos != self.F.auto_pos.get(key):
                lay[key] = pos
        self.main.push_undo()

    def reset_view(self):
        e = self.entry()
        if e is not None:
            e.update(view=None, zoom=1.0, pan=[0.0, 0.0], layout={})
            self.render()

    def reset_display(self):
        """Darstellung auf Standard: Stil, Ellipsoide, Bindungsdicke, Labels, Fehlordnung, Farben …; für die gewählte
        Struktur auch Ansicht, ausgeblendete Atome/Fragmente, Labels je Atom und Drahtmodell. CIFs, Namen, Referenzen,
        ausgewählte Parameter und Exportgröße bleiben."""
        reset_display(self, ("entries", "current", "size", "w_cm", "h_cm", "fmt", "layout"),
                      dict(hidden_frags=[], hidden_atoms=[], label_on={}, wire=[]))
        self.sel = []

    reset_defaults = reset_display

    def reset_layout(self):
        e = self.entry()
        if e is not None:
            e["layout"] = {}
        self.render()

    def layout(self):
        e = self.entry()
        if e is None:
            return {}
        return e.setdefault("layout", {}).setdefault(self.size_key(), {})

    # -- Abbildung ------------------------------------------------------------------------
    def opts(self, e):
        st = self.st
        return dict(prob=st["prob"], show_h=st["show_h"], labels=st["labels"], h_labels=st["h_labels"],
                    colors=st.get("colors") or {}, bond_w=st.get("bond_w", 0.08), hidden_frags=e.get("hidden_frags") or [],
                    style=allowed_style(st.get("style", "mercury"), XRAY_STYLES, "mercury"),
                    hidden_atoms=e.get("hidden_atoms") or [],
                    label_on=e.get("label_on") or {}, wire=set(e.get("wire") or []), wire_h=st.get("wire_h", False),
                    label_mode=st.get("label_mode", "all"), label_size=st.get("label_pt", 0.0),
                    label_color=st.get("label_color", "black"),
                    rod_scale=st.get("bond_w", 0.12) / self.DEFAULT_STATE["bond_w"], bg=st.get("bg", "white"),
                    hide_kinds=[k for k, on in (("solvent", st.get("hide_solv")), ("ion", st.get("hide_ions")))
                                if on])

    def build(self, fig, e=None, interactive=True):
        e = e or self.entry()
        if e is None:
            return None
        mol = self.molecule(e)
        if mol is None or not mol.atoms:
            return None
        cfg = self.base_cfg()
        core.setup_fonts(TABS.plt, cfg)
        lay = e.setdefault("layout", {}).setdefault(self.size_key(), {})
        sel = [mol.find_key(k) for k in self.sel] if interactive else []
        return XrayFigure(cfg, mol, self.view_of(e, mol), self.opts(e), lay, fig,
                          selection=[k for k in sel if k is not None], zoom=e.get("zoom", 1.0),
                          pan=e.get("pan") or (0.0, 0.0))

    def render(self, push=True):
        self.notes_begin()
        try:
            if self.dragger:
                self.dragger.disconnect()
                self.dragger = None
            cfg = self.base_cfg()
            self.host.set_figure_size(*cfg["figsize_in"])
            self.F = self.build(self.fig)
            self.preview_stack.setCurrentIndex(0 if self.F is None else 1)
            if self.F is None:
                self.fig.clear()
            else:
                self.dragger = XrayMouse(self.F, self)
            self.update_info()
            self.update_tables()
            if not self._loading:
                self.fill_atoms()
            self.canvas.draw_idle()
            self.notes_show()
            if push:
                self.main.push_undo()
        except Exception as ex:
            import traceback
            self.main.log_signal.emit(traceback.format_exc())
            self.main.statusBar().showMessage(f"{T('err_title')}: {ex}", 8000)
            self.notes_show(error=ex)

    def update_info(self):
        e = self.entry()
        mol = self.molecule(e) if e else None
        if mol is None:
            self.lbl_info.setText("")
            return
        S = mol.S
        import html as _html
        sg = to_html(space_group_markup(S.block.get("_space_group_name_h-m_alt", "_symmetry_space_group_name_h-m")))
        bits = [_html.escape(T("xr_info", n=len(S.atoms), m=len(mol.atoms), sg="\x00")).replace("\x00", sg)]
        if mol.ops:
            bits.append(T("xr_sym_info", ops="; ".join(f"{r}: {op}" for op, r in mol.ops.items())))
        if S.parts():
            bits.append(T("xr_parts_info", parts=", ".join(str(p) for p in S.parts())))
        if not S.geom_bond:
            bits.append(T("xr_no_geom"))
        if not S.geom_torsion:
            bits.append(T("xr_no_torsion"))
        if mol.truncated:
            bits.append(T("xr_truncated"))
        self.lbl_info.setText("<br>".join(b if i == 0 else _html.escape(b) for i, b in enumerate(bits)))
        self.lbl_info.setStyleSheet(WARN_CSS if mol.truncated else HINT_CSS)

    # -- Tabellen -------------------------------------------------------------------------
    def crystal_table(self):
        """-> (Kopf, Zeilen) als Markup; eine Spalte je Verbindung."""
        cols = []
        for e in self.entries():
            S = self.structure(e)
            if S is not None:
                cols.append((e, crystal_values(S.block, e.get("ref", ""), e.get("ccdc", ""))))
        header = ["Compound"] + [e.get("name") or e["block"] for e, _ in cols]
        rows = [[lab] + [v[key] for _, v in cols] for key, lab in CRYSTAL_ROWS]
        return header, rows

    def params_table(self):
        """Zeilen = gewählte Parameter (gleichnamige Atome in einer Zeile), Spalten = Verbindungen."""
        es = [e for e in self.entries() if self.structure(e) is not None]
        rows_by_kind = {"bond": [], "angle": [], "torsion": []}
        cells, calc_used, notes = {}, False, []
        for ci, e in enumerate(es):
            mol = self.molecule(e)
            used_ops = {}
            for p in e.get("params") or []:
                m = measure(mol, p["atoms"])
                if m is None:
                    continue
                label = param_label(p["atoms"])
                key = (m["kind"], label)
                if key not in rows_by_kind[m["kind"]]:
                    rev = (m["kind"], "–".join(label.split("–")[::-1]))
                    if rev in rows_by_kind[m["kind"]]:
                        key = rev
                    else:
                        rows_by_kind[m["kind"]].append(key)
                txt = minus(m["text"])
                if m["source"] == "calc":
                    txt += "^{a}"
                    calc_used = True
                syms = []
                for k in p["atoms"]:
                    i = mol.find_key(k)
                    if i is not None and mol.atoms[i].get("sym"):
                        syms.append(mol.atoms[i]["sym"])
                        used_ops[mol.atoms[i]["sym"]] = mol.atoms[i]["opkey"]
                cells[(key, ci)] = txt + ("" if not syms else " (" + ",".join("*" + s + "*" for s in syms) + ")")
            if used_ops:
                notes.append(f"{e.get('name') or e['block']}: " + "; ".join(f"{r}: {op}" for r, op in used_ops.items()))
        header = [""] + [e.get("name") or e["block"] for e in es]
        rows = []
        for kind in ("bond", "angle", "torsion"):
            if rows_by_kind[kind]:
                rows.append([KIND_HEAD[kind]] + [""] * len(es))
                for key in rows_by_kind[kind]:
                    rows.append([key[1]] + [cells.get((key, ci), "–") for ci in range(len(es))])
        if calc_used:
            notes.append("^{a} " + CALC_NOTE)
        if notes and any("(" in c for r in rows for c in r[1:]):
            notes.insert(0, "Symmetry codes:")
        return header, rows, notes

    def crystal_clip(self):
        h, r = self.crystal_table()
        return table_plain(h, r), table_html(h, r)

    def params_clip(self):
        h, r, notes = self.params_table()
        title = "Selected bond lengths [Å], angles and torsion angles [°]"
        return table_plain(h, r) + ("\n" + "\n".join(to_plain(n) for n in notes) if notes else ""), \
            table_html(h, r, title, notes)

    def update_tables(self):
        self.txt_cap.setHtml(self.caption()[1] if self.entries() else "")
        if not self.entries():
            self.txt_tab.setHtml("")
            return
        h, r = self.crystal_table()
        html = table_html(h, r, "Crystal data and structure refinement")
        h2, r2, notes = self.params_table()
        if len(r2):
            html += "<br>" + table_html(h2, r2, "Selected bond lengths [Å], angles and torsion angles [°]", notes)
        self.txt_tab.setHtml(html)

    def write_excel(self, path):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Crystal data"
        h, r = self.crystal_table()
        self._sheet(ws, h, r, [], Font)
        ws2 = wb.create_sheet("Selected parameters")
        h2, r2, notes = self.params_table()
        self._sheet(ws2, h2, r2, notes, Font)
        wb.save(path)
        return path

    @staticmethod
    def _sheet(ws, header, rows, notes, Font):
        for c, v in enumerate(header, 1):
            cell = ws.cell(row=1, column=c, value=to_rich_xlsx(v))
            cell.font = Font(bold=True)
        for ri, row in enumerate(rows, 2):
            for c, v in enumerate(row, 1):
                ws.cell(row=ri, column=c, value=to_rich_xlsx(v) if v else None)
        for k, n in enumerate(notes):
            ws.cell(row=len(rows) + 3 + k, column=1, value=to_rich_xlsx(n))
        ws.column_dimensions["A"].width = 40
        for c in range(2, len(header) + 1):
            ws.column_dimensions[chr(64 + min(c, 26))].width = 28

    def save_excel(self):
        if not self.main.ensure_project():
            return
        pdir = Path(self.main.project_dir)
        p = PRJ.export_dir(pdir, PRJ.TABLES) / (PRJ.export_stem(self.main.export_name(), self.STATE_KEY)
                                                + "_tables.xlsx")
        try:
            self.write_excel(p)
        except PermissionError:
            return self.main.error(T("report_locked", p=p))
        self.main.exported([p.parent])

    # -- Export -----------------------------------------------------------------------------
    def export_to(self, base: Path):
        """Je Verbindung eine Abbildung (eigene Ansicht), dazu Tabellen als CSV und Excel."""
        base = Path(base).resolve()
        base.parent.mkdir(parents=True, exist_ok=True)
        files = []
        es = [e for e in self.entries() if self.structure(e) is not None]
        for e in es:
            fig = MplFigure(figsize=(8, 6), dpi=TABS.PREVIEW_DPI)
            FigureCanvasAgg(fig)
            F = self.build(fig, e, interactive=False)
            if F is None:
                continue
            name = base.name + ("_" + core.safe_name(e.get("name") or e["block"]) if len(es) > 1 else "")
            stem = base.with_name(name)
            F.cfg["output"] = str(stem.with_suffix("." + F.cfg["_main_fmt"]))
            files += F.export()
        if not files and not es:
            raise RuntimeError(T("tab_empty"))
        cap = base.with_name(base.name + "_captions.txt")
        cap.write_text("\n\n".join(f"{e.get('name') or e['block']}:\n{self.caption(e)[0]}" for e in es) + "\n",
                       encoding="utf-8")
        files.append(cap)
        h, r = self.crystal_table()
        files.append(write_table_csv(base.with_name(base.name + "_crystal_data.csv"), h, r))
        h2, r2, notes = self.params_table()
        files.append(write_table_csv(base.with_name(base.name + "_parameters.csv"), h2, r2 + [[n] for n in notes]))
        try:
            files.append(self.write_excel(base.with_name(base.name + "_tables.xlsx")))
        except PermissionError:
            self.main.error(T("report_locked", p=base.with_name(base.name + "_tables.xlsx")))
        return files


# ---------------------------------------------------------------------------- Testdaten (Selbsttest)
def make_test_cif(path, blocks=("xtest_a", "xtest_b")):
    """Künstliche Struktur für den Selbsttest: zentrosymmetrischer C₆-Ring auf dem Inversionszentrum (halber Ring
    in der ASU), Cl fehlgeordnet (PART 1 0.7 / PART 2 0.3), separates Wassermolekül, aniso U, _geom_bond und
    _geom_angle mit esd, keine _geom_torsion. Kein echter Datensatz."""
    cell = (7.0, 8.0, 9.0, 90.0, 100.0, 90.0)
    M = ortho_matrix(*cell)
    Mi = np.linalg.inv(M)
    r, rcl, rh = 1.39, 1.39 + 1.74, 1.39 + 0.95
    ex, ey = np.array([0.6, 0.8, 0.0]), np.array([-0.48, 0.36, 0.8])     # Ringebene (orthonormal)

    def cart(rad, deg):
        a = math.radians(deg)
        return rad * (math.cos(a) * ex + math.sin(a) * ey)
    water = np.array([3.4, 3.9, 2.2])
    atoms = [("C1", "C", cart(r, 0), "Uani", 1, "."), ("C2", "C", cart(r, 60), "Uani", 1, "."),
             ("C3", "C", cart(r, 120), "Uani", 1, "."), ("H3", "H", cart(rh, 120), "Uiso", 1, "."),
             ("Cl1", "Cl", cart(rcl, 0), "Uani", 0.7, "1"), ("Cl1B", "Cl", cart(rcl, 60), "Uani", 0.3, "2"),
             ("O1", "O", water, "Uani", 1, "."), ("H1O", "H", water + [0.96, 0, 0], "Uiso", 1, "."),
             ("H2O", "H", water + [-0.24, 0.93, 0], "Uiso", 1, ".")]
    lines = []
    for bn in blocks:
        lines += [f"data_{bn}", "_chemical_formula_sum 'C6 H4 Cl2 O'", "_chemical_formula_weight 163.00",
                  "_space_group_crystal_system monoclinic", "_space_group_name_H-M_alt 'P 21/c'",
                  "loop_", "  _space_group_symop_operation_xyz", " 'x, y, z'", " '-x, y+1/2, -z+1/2'",
                  " '-x, -y, -z'", " 'x, -y-1/2, z-1/2'",
                  "_cell_length_a 7.0000(3)", "_cell_length_b 8.0000(4)", "_cell_length_c 9.0000(4)",
                  "_cell_angle_alpha 90", "_cell_angle_beta 100.000(2)", "_cell_angle_gamma 90",
                  "_cell_volume 496.33(4)", "_cell_formula_units_Z 2", "_exptl_crystal_density_diffrn 1.091",
                  "_exptl_absorpt_coefficient_mu 0.611", "_exptl_crystal_F_000 164",
                  "_exptl_crystal_size_max 0.30", "_exptl_crystal_size_mid 0.20", "_exptl_crystal_size_min 0.10",
                  "_diffrn_ambient_temperature 100(2)", "_diffrn_radiation_wavelength 0.71073",
                  "_diffrn_reflns_theta_min 2.500", "_diffrn_reflns_theta_max 27.500",
                  "_diffrn_reflns_limit_h_min -9", "_diffrn_reflns_limit_h_max 9", "_diffrn_reflns_limit_k_min -10",
                  "_diffrn_reflns_limit_k_max 10", "_diffrn_reflns_limit_l_min -11", "_diffrn_reflns_limit_l_max 11",
                  "_diffrn_reflns_number 5000", "_reflns_number_total 1100", "_diffrn_reflns_av_R_equivalents 0.0310",
                  "_diffrn_reflns_av_unetI/netI 0.0250", "_refine_ls_structure_factor_coef Fsqd",
                  "_refine_ls_matrix_type full", "_refine_ls_number_reflns 1100", "_refine_ls_number_restraints 3",
                  "_refine_ls_number_parameters 70", "_refine_ls_goodness_of_fit_ref 1.050",
                  "_refine_ls_R_factor_gt 0.0350", "_refine_ls_wR_factor_gt 0.0900", "_refine_ls_R_factor_all 0.0400",
                  "_refine_ls_wR_factor_ref 0.0950", "_refine_diff_density_max 0.42", "_refine_diff_density_min -0.31",
                  "loop_", "  _atom_site_label", "  _atom_site_type_symbol", "  _atom_site_fract_x",
                  "  _atom_site_fract_y", "  _atom_site_fract_z", "  _atom_site_U_iso_or_equiv", "  _atom_site_adp_type",
                  "  _atom_site_occupancy", "  _atom_site_disorder_group"]
        for lab, el, xyz, adp, occ, part in atoms:
            f = Mi @ xyz
            lines.append(f" {lab} {el} {f[0]:.6f} {f[1]:.6f} {f[2]:.6f} 0.0300(5) {adp} {occ} {part}")
        lines += ["loop_", "  _atom_site_aniso_label"] + [f"  _atom_site_aniso_U_{k}" for k in
                                                         ("11", "22", "33", "23", "13", "12")]
        for lab, el, _, adp, _, _ in atoms:
            if adp == "Uani":
                lines.append(f" {lab} 0.0250(5) 0.0300(5) 0.0400(6) 0.0030(4) 0.0050(4) -0.0020(4)")
        lines += ["loop_", "  _geom_bond_atom_site_label_1", "  _geom_bond_atom_site_label_2", "  _geom_bond_distance",
                  "  _geom_bond_site_symmetry_2", "  _geom_bond_publ_flag",
                  " C1 C2 1.390(3) . ?", " C2 C3 1.390(3) . ?", " C3 C1 1.390(3) 3_555 ?", " C1 Cl1 1.740(2) . ?",
                  " C2 Cl1B 1.740(9) . ?", " C3 H3 0.9500 . ?", " O1 H1O 0.96(2) . ?", " O1 H2O 0.96(2) . ?",
                  "loop_", "  _geom_angle_atom_site_label_1", "  _geom_angle_atom_site_label_2",
                  "  _geom_angle_atom_site_label_3", "  _geom_angle", "  _geom_angle_site_symmetry_1",
                  "  _geom_angle_site_symmetry_3", "  _geom_angle_publ_flag",
                  " C2 C1 C3 120.0(2) . 3 ?", " C1 C2 C3 120.0(2) . . ?", ""]
    Path(path).write_text("\n".join(lines), encoding="utf-8", newline="\n")
    return Path(path), {lab: xyz for lab, _, xyz, *_ in atoms}
