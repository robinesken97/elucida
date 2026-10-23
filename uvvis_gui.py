"""
uvvis_gui.py - Grafische Oberfläche der UV-Vis-Auswertung (Windows / macOS / Linux).

Start:      python uvvis_gui.py
Selbsttest: python uvvis_gui.py --selftest AUSGABEORDNER [--selftest-isnet]
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import html
import re
import traceback
from pathlib import Path

APP_NAME = "Elucida"
OLD_APP_NAMES = ("Spectravolt", "UVVisTool")   # 1.16–1.25 bzw. bis 1.15: Einstellungen beim ersten Start übernehmen
APP_VERSION = "2.0.2"
SI_PARTS = ("char", "nmr", "figs", "methods", "xray", "dft")   # Teile des SI-Exports (Auswahl im Dialog)
SETTINGS_FILE = None      # Selbsttest: eigene INI-Datei statt der echten Benutzereinstellungen

if "--selftest" in sys.argv:
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import numpy as np  # noqa: E402
import yaml  # noqa: E402
import matplotlib  # noqa: E402

matplotlib.use("QtAgg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.backends.backend_agg import FigureCanvasAgg  # noqa: E402
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg  # noqa: E402
from matplotlib.figure import Figure as MplFigure  # noqa: E402
from PySide6.QtCore import QEvent, QSettings, QSize, Qt, QThread, QTimer, Signal  # noqa: E402
from PySide6.QtCore import QUrl  # noqa: E402
from PySide6.QtGui import (QAction, QActionGroup, QColor, QDesktopServices, QFont, QIcon, QImage,  # noqa: E402
                           QGuiApplication, QKeySequence, QPixmap, QTextCursor)
from PySide6.QtWidgets import (QAbstractItemView, QAbstractSpinBox, QApplication, QCheckBox, QColorDialog, QComboBox,  # noqa: E402
                               QDialog, QDialogButtonBox, QDoubleSpinBox, QFileDialog, QFontComboBox,
                               QFormLayout, QFrame, QGroupBox, QHBoxLayout,
                               QHeaderView, QLabel, QLineEdit, QListWidget, QListWidgetItem, QMainWindow, QMessageBox,
                               QPlainTextEdit, QProgressDialog, QPushButton, QScrollArea,
                               QSpinBox, QSplitter, QStackedWidget, QStatusBar, QTableWidget, QTableWidgetItem, QTextBrowser,
                               QTabWidget, QToolButton, QVBoxLayout, QWidget)

import uvvis_core as core  # noqa: E402
import uvvis_images as IMG  # noqa: E402
from uvvis_i18n import T, get_lang, set_lang  # noqa: E402
import uvvis_tabs  # noqa: E402
import uvvis_project as PRJ  # noqa: E402
import uvvis_report as REPORT  # noqa: E402
import uvvis_update as UPD  # noqa: E402
import uvvis_color as COLOR  # noqa: E402
import uvvis_cv as CVMOD  # noqa: E402
import uvvis_sec as SECMOD  # noqa: E402
import uvvis_epr as EPRMOD  # noqa: E402
import uvvis_hrms as HRMS  # noqa: E402
import uvvis_irms as IRMS  # noqa: E402
import uvvis_xray as XRAY  # noqa: E402
import uvvis_dft as DFT  # noqa: E402
import uvvis_nmr as NMR  # noqa: E402
import uvvis_synth as SYN  # noqa: E402
import uvvis_si as SI  # noqa: E402
import uvvis_style as STYLE  # noqa: E402
from uvvis_style import HINT_CSS, WARN_CSS, align_forms, compact, make_form  # noqa: E402

PREVIEW_DPI = 90
FIGSIZE = (8.0, 6.0)


def resource_path(rel):
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base / rel


def rgba_to_pixmap(arr, max_h=110):
    arr = np.ascontiguousarray(arr)
    h, w = arr.shape[:2]
    img = QImage(arr.data, w, h, 4 * w, QImage.Format_RGBA8888).copy()
    return QPixmap.fromImage(img).scaledToHeight(max_h, Qt.SmoothTransformation)


def parse_float(txt):
    txt = (txt or "").strip().replace(",", ".")
    if not txt:
        return None
    v = float(txt)
    if v <= 0:
        raise ValueError
    return v


class Task(QThread):
    """Funktion im Hintergrund ausführen (Freistellung, CDXML), GUI bleibt bedienbar."""
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, fn, parent=None):
        super().__init__(parent)
        self.fn = fn

    def run(self):
        try:
            self.done.emit(self.fn())
        except Exception as e:
            self.failed.emit(f"{type(e).__name__}: {e}")


class GuiDragger(core.Dragger):
    """Wie core.Dragger, meldet aber jede Änderung (Ziehen von Inset/Bild/Label, Mausrad)
    an die GUI, damit Vorschau, Serienwechsel und Export dasselbe Layout verwenden."""

    def __init__(self, F, on_change):
        self.on_change = on_change
        super().__init__(F, None)

    def _report(self):
        # erst nach allen anderen Release-Handlern auslesen (Label-Drag von matplotlib)
        QTimer.singleShot(0, lambda: self.on_change(self.F.current_layout()))

    def release(self, ev):
        super().release(ev)
        self._report()

    def scroll(self, ev):
        if super().scroll(ev):
            self._report()

    def key(self, ev):
        pass


DEFAULT_GLOB = {"path_length": 1.0, "cutoff": 1.0, "min_points": 3, "bands": "",
                "baseline": "series", "bl_a": 1050.0, "bl_b": 1100.0, "ymax": 1.0,
                "r2": "legend", "show_err": False, "table": True,
                "size": "half_a4", "w_cm": 16.0, "h_cm": 11.0,
                "fmt_pdf": True, "fmt_svg": True, "fmt_png": True,
                "x_auto": True, "x_min": 200.0, "x_max": 1100.0,
                "bands_add": True, "shoulders": False, "show_excl": False, "inset": True,
                "mark_sh": False, "label_index": False, "min_r2": 0.98, "color_show": False, "color_d": 1.0}

UNITS = [("mM", "mM"), ("uM", "µM"), ("mg/mL", "mg/mL"), ("M", "M")]

def migrate_settings(old, new):
    """Einstellungen des alten Programmnamens (zuletzt geöffnet, Sprache, Theme …) einmalig übernehmen,
    solange unter dem neuen Namen noch nichts gespeichert ist. Das Alte bleibt (Rückkehr zur alten Version)."""
    if new.allKeys() or not old.allKeys():
        return False
    for k in old.allKeys():
        new.setValue(k, old.value(k))
    new.sync()
    return True


class Section(QWidget):
    """Aufklappbarer Bereich für seltener gebrauchte Optionen. Weicht ein Wert vom Standard ab, steht
    „(geändert)“ im Titel – eingeklappte Einstellungen wirken sich sonst unbemerkt aus."""

    def __init__(self, title, key, settings):
        super().__init__()
        self.title, self.key, self.settings = title, key, settings
        lay = QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(2)
        self.btn = QToolButton(checkable=True, autoRaise=True)
        self.btn.setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
        self.btn.setStyleSheet("QToolButton { border: none; padding: 2px 0; }")
        self.btn.setFont(QApplication.font())         # macOS: Werkzeugknöpfe sonst mit kleinerer Schrift
        self.body = QWidget()
        self.form = make_form(self.body)
        self.form.setContentsMargins(0, 0, 0, 0)
        lay.addWidget(self.btn)
        lay.addWidget(self.body)
        self.btn.toggled.connect(self._toggled)
        self.btn.setChecked(settings.value(f"ui/section_{key}", False, type=bool))
        self._toggled(self.btn.isChecked())
        self.set_changed(False)

    def _toggled(self, on):
        self.btn.setArrowType(Qt.DownArrow if on else Qt.RightArrow)
        self.body.setVisible(on)
        self.settings.setValue(f"ui/section_{self.key}", on)

    def set_changed(self, changed):
        self.changed = changed
        self.btn.setText(f"{self.title} ({T('adv_changed')})" if changed else self.title)
        f = self.btn.font()
        f.setBold(changed)
        self.btn.setFont(f)


class ConcDialog(QDialog):
    """Konzentrationen je Probe eintragen (wenn sie nicht im Probennamen stehen)."""

    def __init__(self, parent, samples, overrides, default_series):
        super().__init__(parent)
        self.setWindowTitle(T("conc_title"))
        self.samples = samples
        lay = QVBoxLayout(self)
        intro = QLabel(T("conc_intro"))
        intro.setWordWrap(True)
        lay.addWidget(intro)

        unit0 = None
        for s_ in samples:
            o = overrides.get(s_["key"]) or {}
            unit0 = unit0 or o.get("unit") or (s_["parsed"][2] if s_["parsed"] else None)
        row = QHBoxLayout()
        row.addWidget(QLabel(T("unit")))
        self.cmb_unit = QComboBox()
        self.cmb_unit.addItem(T("unit_choose"), None)          # keine stille Vorbelegung
        for key, label in UNITS:
            self.cmb_unit.addItem(label, key)
        self.cmb_unit.setCurrentIndex(max(0, self.cmb_unit.findData(unit0)) if unit0 else 0)
        row.addWidget(self.cmb_unit)
        row.addStretch(1)
        lay.addLayout(row)

        self.table = QTableWidget(len(samples), 3)
        self.table.setHorizontalHeaderLabels([T("col_sample"), T("col_series"), T("col_conc")])
        self.table.verticalHeader().setVisible(False)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        for i, s_ in enumerate(samples):
            o = overrides.get(s_["key"]) or {}
            it = QTableWidgetItem(s_["key"])
            it.setFlags(it.flags() & ~Qt.ItemIsEditable)
            self.table.setItem(i, 0, it)
            series = o.get("series") or (s_["parsed"][0] if s_["parsed"] else default_series)
            self.table.setItem(i, 1, QTableWidgetItem(series))
            c = o.get("conc") or (s_["parsed"][1] if s_["parsed"] and not o.get("skip") else None)
            self.table.setItem(i, 2, QTableWidgetItem(f"{c:g}" if c else ""))
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.table.resizeColumnsToContents()
        lay.addWidget(self.table)

        row = QHBoxLayout()
        row.addWidget(QLabel(T("dil_start")))
        self.sp_start = QDoubleSpinBox(decimals=4, minimum=0.0, maximum=1e6, value=0.0)
        row.addWidget(self.sp_start)
        row.addWidget(QLabel(T("dil_factor")))
        self.sp_fac = QDoubleSpinBox(decimals=3, minimum=1.0, maximum=1000, value=2.0)
        row.addWidget(self.sp_fac)
        b = QPushButton(T("dil_fill"))
        b.clicked.connect(self.fill_dilution)
        row.addWidget(b)
        row.addStretch(1)
        lay.addLayout(row)
        hint = QLabel(T("dil_hint"))
        hint.setStyleSheet(HINT_CSS)
        lay.addWidget(hint)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(self.try_accept)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)
        self.table.setMinimumHeight(36 + 30 * min(len(samples), 12))
        self.resize(600, 300 + 30 * min(len(samples), 12))

    def fill_dilution(self):
        if self.sp_start.value() <= 0:
            QMessageBox.warning(self, T("conc_title"), T("start_missing"))
            return
        rows = sorted({i.row() for i in self.table.selectedIndexes()}) or list(range(self.table.rowCount()))
        c = self.sp_start.value()
        for r in rows:
            self.table.item(r, 2).setText(f"{c:.6g}")
            c /= self.sp_fac.value()

    def try_accept(self):
        any_conc = any((self.table.item(i, 2).text() or "").strip() for i in range(self.table.rowCount()))
        if any_conc and self.cmb_unit.currentData() is None:
            QMessageBox.warning(self, T("conc_title"), T("unit_missing"))
            return
        self.accept()

    def overrides(self):
        unit = self.cmb_unit.currentData()
        out = {}
        for i, s_ in enumerate(self.samples):
            series = (self.table.item(i, 1).text() or "").strip()
            try:
                c = parse_float(self.table.item(i, 2).text())
            except ValueError:
                c = None
            out[s_["key"]] = {"conc": c, "unit": unit, "series": series} if c else {"skip": True}
        return out


class CompoundDialog(QDialog):
    """Verbindung des Projekts: Name/Probennummer, ChemDraw (Formel und Molmasse daraus), beides änderbar."""

    def __init__(self, parent, comp):
        super().__init__(parent)
        self.setWindowTitle(T("compound_title"))
        self.main = parent
        self.path = comp.get("structure")
        f = make_form(self)
        self.ed_name = QLineEdit(comp.get("name", ""))
        self.ed_name.setPlaceholderText(T("compound_name_ph"))
        f.addRow(T("compound_name"), self.ed_name)
        row = QHBoxLayout()
        self.lbl_struct = QLabel(Path(self.path).name if self.path else "–")
        b = QPushButton(T("choose"))
        b.clicked.connect(self.choose)
        row.addWidget(self.lbl_struct, 1)
        row.addWidget(b)
        f.addRow(T("structure"), row)
        self.ed_formula = QLineEdit(comp.get("formula", ""))
        self.ed_formula.setPlaceholderText("C18H20N2")
        f.addRow(T("compound_formula"), self.ed_formula)
        self.ed_mw = QLineEdit(f"{comp['mw']:g}" if comp.get("mw") else "")
        f.addRow(T("molar_mass"), self.ed_mw)
        self.ed_mp = QLineEdit(comp.get("mp", ""))
        self.ed_mp.setPlaceholderText(T("compound_mp_ph"))
        f.addRow(T("compound_mp"), self.ed_mp)
        hint = QLabel(T("compound_hint"))
        hint.setWordWrap(True)
        hint.setStyleSheet(HINT_CSS)
        f.addRow(hint)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        f.addRow(bb)
        self.resize(460, 0)

    def choose(self):
        p, _ = QFileDialog.getOpenFileName(self, T("dlg_structure"), self.main.settings.value("last_img_dir", ""),
                                           f"{T('flt_struct')};;{T('flt_all')}")
        if p:
            self.main.settings.setValue("last_img_dir", str(Path(p).parent))
            self.set_path(p)

    def set_path(self, p):
        self.path = p
        self.lbl_struct.setText(Path(p).name)
        chem = compound_from_structure(p)
        if chem.get("formula"):
            self.ed_formula.setText(chem["formula"])
        if chem.get("mw"):
            self.ed_mw.setText(f"{chem['mw']:.2f}")
        if not self.ed_name.text().strip():
            self.ed_name.setText(Path(p).stem)

    def result_compound(self):
        return {"name": self.ed_name.text().strip(), "structure": self.path,
                "formula": self.ed_formula.text().strip().replace(" ", ""), "mw": parse_float(self.ed_mw.text()),
                "mp": self.ed_mp.text().strip()}


def compound_from_structure(p):
    """Formel und Molmasse aus einer CDXML (oder der gleichnamigen CDXML neben SVG/PDF/PNG)."""
    p = Path(p)
    src = p if p.suffix.lower() == ".cdxml" else p.with_suffix(".cdxml")
    if not src.exists():
        return {}
    try:
        import uvvis_chem
        c = uvvis_chem.load_cdxml(src)
        return {"formula": c.get("formula"), "mw": c.get("mw")}
    except Exception:
        return {}


class TextFormatDialog(QDialog):
    """Schriften, Labelstil und Rahmen (Menü Format). Gilt für alle Tabs oder – wahlweise – nur für den offenen Tab;
    gezeigt wird nur, was der Tab im Bild überhaupt hat (opts)."""

    ALL_OPTS = ("fonts", "greek", "labels", "label_size", "axes_frame", "struct_frame")

    def __init__(self, parent, tf, own=None, opts=ALL_OPTS, tab_key=None, tab_name=""):
        super().__init__(parent)
        from uvvis_i18n import TEXTS
        self.setWindowTitle(T("tf_title"))
        self.glob = dict(core.TEXT_DEFAULTS, **(tf or {}))
        self.own = dict(self.glob, **own) if own else None
        lay = QVBoxLayout(self)
        self.cmb_scope = QComboBox()
        self.cmb_scope.addItem(T("tf_scope_global"), "global")
        if tab_key:
            self.cmb_scope.addItem(T("tf_scope_own", tab=tab_name), "own")
            row = QHBoxLayout()
            row.addWidget(QLabel(T("tf_scope")))
            row.addWidget(self.cmb_scope, 1)
            lay.addLayout(row)
        hint = QLabel(T("tf_scope_hint") if tab_key else T("tf_no_text", tab=tab_name))
        hint.setWordWrap(True)
        hint.setStyleSheet(HINT_CSS)
        lay.addWidget(hint)

        g = QGroupBox(T("tf_grp_fonts"))
        f = QFormLayout(g)
        self.cmb_font, self.cmb_sym = QFontComboBox(), QFontComboBox()
        f.addRow(T("tf_font"), self.cmb_font)
        self.cb_greek = QCheckBox(T("tf_greek"))
        self.cb_greek.toggled.connect(self.cmb_sym.setEnabled)
        f.addRow(T("tf_symbol_font"), self.cmb_sym)
        f.addRow(self.cb_greek)
        if "greek" not in opts:
            for w_ in (self.cmb_sym, self.cb_greek, f.labelForField(self.cmb_sym)):
                w_.setVisible(False)
        g.setVisible("fonts" in opts)
        lay.addWidget(g)

        lk = f"tf_grp_labels_{tab_key}"
        g = QGroupBox(T(lk if lk in TEXTS else "tf_grp_labels"))
        f = QFormLayout(g)
        self.sp_size = QDoubleSpinBox(decimals=1, minimum=0, maximum=40, singleStep=0.5, suffix=" pt")
        self.sp_size.setSpecialValueText(T("tf_auto"))
        f.addRow(T("tf_label_size"), self.sp_size)
        if "label_size" not in opts:
            self.sp_size.setVisible(False)
            f.labelForField(self.sp_size).setVisible(False)
        row = QHBoxLayout()
        self.cb_bold = QCheckBox(T("tf_bold"))
        self.cb_italic = QCheckBox(T("tf_italic"))
        row.addWidget(self.cb_bold)
        row.addWidget(self.cb_italic)
        row.addStretch(1)
        f.addRow("", row)
        self.btn_color = QPushButton()
        self.btn_color.clicked.connect(self.pick_color)
        f.addRow(T("tf_color"), self.btn_color)
        g.setVisible("labels" in opts)
        lay.addWidget(g)

        g = QGroupBox(T("tf_grp_frames"))
        v = QVBoxLayout(g)
        self.cb_axes = QCheckBox(T("tf_axes_frame"))
        self.cb_struct = QCheckBox(T("tf_struct_frame"))
        v.addWidget(self.cb_axes)
        v.addWidget(self.cb_struct)
        self.cb_axes.setVisible("axes_frame" in opts)
        self.cb_struct.setVisible("struct_frame" in opts)
        g.setVisible("axes_frame" in opts or "struct_frame" in opts)
        lay.addWidget(g)

        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel | QDialogButtonBox.RestoreDefaults)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        bb.button(QDialogButtonBox.RestoreDefaults).clicked.connect(self.restore_defaults)
        lay.addWidget(bb)
        self.cmb_scope.setCurrentIndex(1 if (tab_key and own) else 0)
        self._scope = self.cmb_scope.currentData()
        self.load(self.own if self._scope == "own" else self.glob)
        self.cmb_scope.currentIndexChanged.connect(self.scope_changed)

    def scope_changed(self, *_):
        """Beim Umschalten die Werte der jeweiligen Ebene zeigen (eigenes Format startet als Kopie des Standards)."""
        cur = self.result_format()
        if self._scope == "own":
            self.own = cur
        else:
            self.glob = cur
        self._scope = self.cmb_scope.currentData()
        self.load((self.own or dict(self.glob)) if self._scope == "own" else self.glob)

    def load(self, tf):
        tf = dict(core.TEXT_DEFAULTS, **tf)
        self.cmb_font.setCurrentText(tf["font"])
        self.cmb_sym.setCurrentText(tf["symbol_font"])
        self.cb_greek.setChecked(tf["greek"])
        self.cmb_sym.setEnabled(tf["greek"])
        self.sp_size.setValue(tf["label_pt"] or 0)
        self.cb_bold.setChecked(tf["bold"])
        self.cb_italic.setChecked(tf["italic"])
        self.color = tf["color"]
        self._show_color()
        self.cb_axes.setChecked(tf["axes_frame"])
        self.cb_struct.setChecked(tf["struct_frame"])

    def _show_color(self):
        self.btn_color.setText(self.color)
        self.btn_color.setStyleSheet(f"QPushButton {{ color: {self.color}; font-weight: bold; }}")

    def pick_color(self):
        c = QColorDialog.getColor(QColor(self.color), self, T("tf_color"))
        if c.isValid():
            self.color = c.name()
            self._show_color()

    def restore_defaults(self):
        self.load(core.TEXT_DEFAULTS)

    def result_format(self):
        return {"font": self.cmb_font.currentFont().family(), "symbol_font": self.cmb_sym.currentFont().family(),
                "greek": self.cb_greek.isChecked(), "label_pt": self.sp_size.value(),
                "bold": self.cb_bold.isChecked(), "italic": self.cb_italic.isChecked(), "color": self.color,
                "axes_frame": self.cb_axes.isChecked(), "struct_frame": self.cb_struct.isChecked()}

    def results(self):
        """-> (Standard für alle Tabs, eigenes Format des Tabs oder None)."""
        cur = self.result_format()
        if self.cmb_scope.currentData() == "own":
            return self.glob, cur
        return cur, None


class MainWindow(QMainWindow):
    log_signal = Signal(str)

    def __init__(self):
        super().__init__()
        self.settings = (QSettings(str(SETTINGS_FILE), QSettings.IniFormat) if SETTINGS_FILE
                         else QSettings(APP_NAME, APP_NAME))
        if not SETTINGS_FILE:
            for old in OLD_APP_NAMES:
                if migrate_settings(QSettings(old, old), self.settings):
                    break
        set_lang(self.settings.value("lang", "de"))
        SI.set_font(self.settings.value("si_font", "Arial"))
        self.csv_path = None
        self.groups, self.units = {}, {}
        self.samples, self.conc, self.csv_paths = [], {}, []          # Rohproben und Konzentrations-Overrides
        self.series = {}               # je Serie: mw, structure, photo, bg, layout, assets
        self.current = None
        self.F = None
        self.dragger = None
        self.hover = None
        self._tasks = []
        self._loading = False
        self.glob = dict(DEFAULT_GLOB)
        self.extra_cache = {}
        self.ov_state, self.fl_state, self.cv_state, self.sec_state, self.epr_state = {}, {}, {}, {}, {}  # leer
        self.irms_state = {}
        self.hrms_state = {}
        self.xray_state = {}
        self.dft_state = {}
        self.nmr_state = {}
        self.syn_state = {}
        self.compound = {}                  # Verbindung des Projekts: name, structure, formula, mw
        self.text_fmt = dict(core.TEXT_DEFAULTS)
        self.tab_fmt = {}                                # eigenes Textformat je Tab (Schlüssel = STATE_KEY / epsilon)
        self._tab_key = "syn"                            # offener Tab (Name, nicht Position)
        self.project_dir = None
        self._undo, self._redo = [], []
        self._restoring = False
        self._saved_snap = None

        self.fig = MplFigure(figsize=FIGSIZE, dpi=PREVIEW_DPI)
        self.canvas = FigureCanvasQTAgg(self.fig)
        self.canvas.setFixedSize(int(FIGSIZE[0] * PREVIEW_DPI), int(FIGSIZE[1] * PREVIEW_DPI))
        self.timer = QTimer(self, singleShot=True, interval=400)
        self.timer.timeout.connect(self.render)
        self.autosave_timer = QTimer(self, singleShot=True, interval=800)
        self.autosave_timer.timeout.connect(self.autosave)

        # Meldungen als Einblendung unten über dem Fenster statt Statusleiste: die Leiste bliebe nach der ersten
        # Meldung als leerer Streifen unter den Fußleisten stehen (Knöpfe nicht mittig), Ein-/Ausblenden ließe die
        # Abbildung bei jedem Autosave springen
        self._status = QStatusBar()                      # nur Nachrichtenträger (showMessage mit Zeitlimit), unsichtbar
        self._toast = QLabel(self, objectName="toast", alignment=Qt.AlignCenter)
        self._toast.setAttribute(Qt.WA_TransparentForMouseEvents)
        self._toast.hide()
        self._status.messageChanged.connect(self._place_toast)
        self.log_signal.connect(self._append_log)
        core.LOG = IMG.LOG = lambda m="": self.log_signal.emit(str(m))
        icon = resource_path("assets/icon.png")
        if icon.exists():
            self.setWindowIcon(QIcon(str(icon)))
        self.apply_theme()
        QGuiApplication.styleHints().colorSchemeChanged.connect(self._system_scheme_changed)
        self.build_ui()
        self.setAcceptDrops(True)
        QTimer.singleShot(0, self._init_undo)
        if os.environ.get("QT_QPA_PLATFORM") != "offscreen":
            QTimer.singleShot(3000, lambda: self.check_updates(manual=False))
        geo = self.settings.value("geometry")
        if geo:
            self.restoreGeometry(geo)
        else:
            self.resize(1320, 900)

    # ------------------------------------------------------------------ UI
    def build_ui(self):
        self.update_title()
        mb = self.menuBar()
        mb.clear()
        m = mb.addMenu(T("menu_file"))
        act_new = QAction(STYLE.icon("new"), T("act_new_project"), self, shortcut=QKeySequence.New,
                          triggered=self.new_project)
        act_open_prj = QAction(STYLE.icon("open"), T("act_open_project"), self, shortcut=QKeySequence.Open,
                               triggered=self.open_project_dialog)
        act_save = QAction(STYLE.icon("save"), T("act_save_project"), self, shortcut=QKeySequence.Save,
                           triggered=self.save_project)
        m.addAction(act_new)
        m.addAction(act_open_prj)
        self.recent_menu = m.addMenu(T("act_recent"))
        self.fill_recent_menu()
        m.addAction(act_save)
        m.addAction(QAction(T("act_save_copy"), self, shortcut=QKeySequence("Ctrl+Shift+S"),
                            triggered=self.save_project_copy))
        self.act_folder = QAction(STYLE.icon("reveal"), T("act_show_folder"), self, triggered=self.show_project_folder)
        self.act_folder.setEnabled(bool(self.project_dir))
        m.addAction(self.act_folder)
        m.addSeparator()
        # nur für den Tab UV-Vis (ε): eigenes Untermenü, damit klar ist, wofür die Befehle gelten
        m_eps = m.addMenu(T("menu_eps"))
        m_eps.addAction(QAction(T("act_open"), self, shortcut=QKeySequence("Ctrl+Shift+O"), triggered=self.choose_csv))
        m_eps.addAction(QAction(T("act_conc"), self, triggered=self.edit_concentrations))
        m_eps.addAction(QAction(T("act_restart"), self, triggered=self.restart_file))
        m.addSeparator()
        act_export = QAction(STYLE.icon("export"), T("act_export"), self, shortcut="Ctrl+E",
                             triggered=self.export_current)
        m.addAction(act_export)
        m.addAction(QAction(T("act_export_si"), self, triggered=self.export_si_dialog))
        m.addSeparator()
        m.addAction(QAction(T("act_quit"), self, shortcut=QKeySequence.Quit, triggered=self.close))
        me = mb.addMenu(T("menu_edit"))
        self.act_undo = QAction(STYLE.icon("undo"), T("act_undo"), self, triggered=self.undo)
        self.act_undo.setShortcuts([QKeySequence.Undo])
        self.act_redo = QAction(STYLE.icon("redo"), T("act_redo"), self, triggered=self.redo)
        self.act_redo.setShortcuts([QKeySequence.Redo, QKeySequence("Ctrl+Shift+Z")])
        me.addAction(self.act_undo)
        me.addAction(self.act_redo)
        self._update_undo_actions()
        # Werkzeugleiste: Projekt, Rückgängig/Wiederholen (dieselben Aktionen wie im Menü; Exportieren steht
        # in jedem Tab unten fest sichtbar)
        if getattr(self, "toolbar", None):
            self.removeToolBar(self.toolbar)
            self.toolbar.setParent(None)                # sofort lösen, gelöscht wird später
            self.toolbar.deleteLater()
        self.toolbar = self.addToolBar(T("toolbar"))
        self.toolbar.setObjectName("mainToolbar")
        self.toolbar.setMovable(False)
        self.toolbar.setIconSize(QSize(18, 18))
        for act in (act_new, act_open_prj, act_save, self.act_folder, None, self.act_undo, self.act_redo):
            if act is None:
                self.toolbar.addSeparator()
                continue
            keys = act.shortcut().toString(QKeySequence.NativeText)
            act.setToolTip(f"{act.text().rstrip('…')} ({keys})" if keys else act.text())
            self.toolbar.addAction(act)
        # Name des Dateimanagers je System (Finder, Explorer, sonst Standard-Dateimanager, z. B. Dolphin/Nautilus)
        fm = {"darwin": "fm_finder", "win32": "fm_explorer"}.get(sys.platform, "fm_linux")
        self.act_folder.setToolTip(T("act_show_folder_tip", fm=T(fm)))
        self.toolbar.addSeparator()
        self.act_compound = QAction(STYLE.icon("compound"), self.compound_label(), self,
                                    triggered=self.edit_compound)
        self.act_compound.setToolTip(T("compound_tip"))
        self.toolbar.addAction(self.act_compound)
        self.toolbar.widgetForAction(self.act_compound).setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
        self.act_si = QAction(STYLE.icon("export"), T("act_si_short"), self, triggered=self.export_si_dialog)
        self.act_si.setToolTip(T("act_si_tip"))
        self.toolbar.addAction(self.act_si)
        tb_si = self.toolbar.widgetForAction(self.act_si)          # Ziel des Programms: als Hauptknopf hervorheben
        tb_si.setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
        tb_si.setProperty("accent", True)
        tb_si.style().unpolish(tb_si)
        tb_si.style().polish(tb_si)
        self.act_si.setIcon(STYLE.icon("export", "#FFFFFF"))
        mf = mb.addMenu(T("menu_format"))
        mf.addAction(QAction(T("act_text_format"), self, triggered=self.edit_text_format))
        mt = mf.addMenu(T("menu_theme"))
        grp_t = QActionGroup(self)
        for code in ("system", "light", "dark"):
            act = QAction(T(f"theme_{code}"), self, checkable=True, checked=self._theme_mode() == code)
            act.triggered.connect(lambda _=False, c=code: self.switch_theme(c))
            grp_t.addAction(act)
            mt.addAction(act)
        ml = mb.addMenu(T("menu_lang"))
        grp = QActionGroup(self)
        for code, name in (("de", "Deutsch"), ("en", "English")):
            act = QAction(name, self, checkable=True, checked=get_lang() == code)
            act.triggered.connect(lambda _=False, c=code: self.switch_lang(c))
            grp.addAction(act)
            ml.addAction(act)
        mh = mb.addMenu(T("menu_help"))
        mh.addAction(QAction(T("act_guide_tab"), self, shortcut=QKeySequence.HelpContents, triggered=self.show_tab_guide))
        mh.addAction(QAction(T("act_quickstart"), self, triggered=self.show_quickstart))
        mh.addAction(QAction(T("act_demo"), self, triggered=self.open_demo))
        mh.addSeparator()
        mh.addAction(QAction(T("act_log"), self, triggered=self.show_log))
        mh.addSeparator()
        mh.addAction(QAction(T("act_check_updates"), self, triggered=lambda: self.check_updates(manual=True)))
        act = QAction(T("act_auto_updates"), self, checkable=True, checked=self._auto_updates())
        act.toggled.connect(lambda on: self.settings.setValue("update_auto", on))
        mh.addAction(act)
        mh.addSeparator()
        mh.addAction(QAction(T("act_about"), self, triggered=lambda: QMessageBox.about(
            self, T("act_about"), T("about", v=APP_VERSION))))

        # ---- linke Spalte
        left = QWidget()
        lv = QVBoxLayout(left)

        forms = []
        g, body = STYLE.card(T("grp_data"), "data")
        f = make_form(body)
        f.setContentsMargins(0, 0, 0, 0)
        forms.append(f)
        row = QHBoxLayout()
        self.btn_open = QPushButton(STYLE.icon("open", "#FFFFFF"), T("act_open"))
        self.btn_open.setProperty("accent", True)
        self.btn_open.clicked.connect(self.choose_csv)
        row.addWidget(self.btn_open)
        self.lbl_file = QLabel(self.files_label() if self.csv_path else T("no_file"))
        self.lbl_file.setWordWrap(True)
        row.addWidget(self.lbl_file, 1)
        f.addRow(row)
        self.cmb_series = QComboBox()
        self.cmb_series.addItems(list(self.groups))
        if self.current:
            self.cmb_series.setCurrentText(self.current)
        self.cmb_series.currentTextChanged.connect(self.series_changed)
        row = QHBoxLayout()
        row.addWidget(self.cmb_series, 1)
        b = QPushButton(T("act_conc"))
        b.clicked.connect(self.edit_concentrations)
        row.addWidget(b)
        f.addRow(T("series"), row)
        lv.addWidget(g)

        g, body = STYLE.card(T("grp_compound"), "compound")
        f = make_form(body)
        f.setContentsMargins(0, 0, 0, 0)
        forms.append(f)
        self.ed_mw = QLineEdit()
        self.ed_mw.setPlaceholderText(T("mw_hint_none"))
        self.ed_mw.editingFinished.connect(self.mw_edited)
        row = QHBoxLayout()
        row.addWidget(self.ed_mw, 1)
        row.addWidget(QLabel("g/mol"))
        f.addRow(T("molar_mass_short"), row)
        self.lbl_chem = QLabel("")
        self.lbl_chem.setWordWrap(True)
        f.addRow("", self.lbl_chem)
        row = QHBoxLayout()
        self.lbl_struct = QLabel("–")
        b1 = QPushButton(T("choose"))
        b1.clicked.connect(self.choose_structure)
        b2 = QPushButton(T("remove"))
        b2.clicked.connect(lambda: self.set_structure(None))
        row.addWidget(self.lbl_struct, 1)
        row.addWidget(b1)
        row.addWidget(b2)
        f.addRow(T("structure"), row)
        row = QHBoxLayout()
        self.lbl_photo = QLabel("–")
        b1 = QPushButton(T("choose"))
        b1.clicked.connect(self.choose_photo)
        b2 = QPushButton(T("remove"))
        b2.clicked.connect(lambda: self.set_photo(None))
        row.addWidget(self.lbl_photo, 1)
        row.addWidget(b1)
        row.addWidget(b2)
        f.addRow(T("photo"), row)
        self.sec_images = Section(T("adv_images"), "images", self.settings)
        fa = self.sec_images.form
        forms.append(fa)
        self.sp_struct_w = QDoubleSpinBox(decimals=1, minimum=3, maximum=90, singleStep=1, suffix=" %")
        self.sp_struct_w.valueChanged.connect(lambda v: self.image_size_changed("struktur", v))
        fa.addRow(T("img_size_struct"), self.sp_struct_w)
        self.sp_photo_w = QDoubleSpinBox(decimals=1, minimum=2, maximum=60, singleStep=0.5, suffix=" %")
        self.sp_photo_w.valueChanged.connect(lambda v: self.image_size_changed("kuevette", v))
        fa.addRow(T("img_size_photo"), self.sp_photo_w)
        compact(self.sp_struct_w, self.sp_photo_w)
        self.cmb_bg = QComboBox()
        for key in ("auto", "isnet", "grabcut", "border", "none"):
            self.cmb_bg.addItem(T({"auto": "bg_auto", "isnet": "bg_isnet", "grabcut": "bg_grabcut",
                                   "border": "bg_border", "none": "bg_off"}[key]), key)
        self.cmb_bg.currentIndexChanged.connect(self.bg_changed)
        self.cmb_bg.currentIndexChanged.connect(self._update_sections)
        fa.addRow(T("bg_method"), self.cmb_bg)
        f.addRow(self.sec_images)
        self.lbl_thumb = QLabel()
        self.lbl_thumb.setAlignment(Qt.AlignCenter)
        f.addRow("", self.lbl_thumb)
        # Farbe der Lösung aus dem Spektrum (nur Anzeige, nie in der Abbildung)
        self.cb_color = QCheckBox(T("color_show"))
        self.cb_color.setToolTip(T("color_tip"))
        self.cb_color.toggled.connect(self.color_settings_changed)
        f.addRow(self.cb_color)
        row = QHBoxLayout()
        self.sp_color_c = QDoubleSpinBox(decimals=4, minimum=0.0001, maximum=1e6, singleStep=0.1)
        self.sp_color_d = QDoubleSpinBox(decimals=2, minimum=0.01, maximum=100, singleStep=0.1, suffix=" cm")
        for w_ in (QLabel("c"), self.sp_color_c, QLabel("d"), self.sp_color_d):
            row.addWidget(w_)
        self.w_color_cd = QWidget()
        self.w_color_cd.setLayout(row)
        row.setContentsMargins(0, 0, 0, 0)
        f.addRow("", self.w_color_cd)
        for sp in (self.sp_color_c, self.sp_color_d):
            sp.valueChanged.connect(self.color_settings_changed)
        self.lbl_color = QLabel()
        self.lbl_color.setFixedHeight(36)
        self.lbl_color_info = QLabel()
        self.lbl_color_info.setWordWrap(True)
        self.lbl_color_info.setTextInteractionFlags(Qt.TextSelectableByMouse)
        f.addRow("", self.lbl_color)
        f.addRow("", self.lbl_color_info)
        self._color_spectra = None
        lv.addWidget(g)

        g, body = STYLE.card(T("grp_eval"), "eval")
        f = make_form(body)
        f.setContentsMargins(0, 0, 0, 0)
        forms.append(f)
        self.sp_d = QDoubleSpinBox(decimals=3, minimum=0.001, maximum=100, singleStep=0.1, suffix=" cm")
        f.addRow(T("path_length_short"), self.sp_d)
        self.ed_solvent = QLineEdit()                    # je Serie; für SI-Zeile und Bildunterschrift
        self.ed_solvent.setPlaceholderText(T("si_solvent_ph"))
        self.ed_solvent.editingFinished.connect(self.solvent_changed)
        f.addRow(T("si_solvent"), self.ed_solvent)
        self.sp_cut = QDoubleSpinBox(decimals=2, minimum=0.05, maximum=5, singleStep=0.1)
        self.sp_cut.setToolTip(T("cutoff_tip"))
        f.addRow(T("cutoff"), self.sp_cut)
        self.ed_bands = QLineEdit()
        self.ed_bands.setPlaceholderText(T("wavelengths_ph"))
        f.addRow(T("wavelengths"), self.ed_bands)
        self.cb_bands_add = QCheckBox(T("bands_add"))
        f.addRow(self.cb_bands_add)
        self.cmb_bl = QComboBox()
        for key in ("series", "simple", "manual", "off"):
            self.cmb_bl.addItem(T({"series": "bl_series", "simple": "bl_simple", "manual": "bl_manual",
                                   "off": "bl_off"}[key]), key)
        self.cmb_bl.setToolTip(T("baseline_tip"))
        f.addRow(T("baseline"), self.cmb_bl)
        row = QHBoxLayout()
        self.sp_bla = QDoubleSpinBox(decimals=0, minimum=100, maximum=3500, suffix=" nm")
        self.sp_blb = QDoubleSpinBox(decimals=0, minimum=100, maximum=3500, suffix=" nm")
        row.addWidget(self.sp_bla)
        row.addWidget(QLabel("–"))
        row.addWidget(self.sp_blb)
        row.addStretch(1)
        f.addRow("", row)
        self.sec_eval = Section(T("adv_eval"), "eval", self.settings)
        fa = self.sec_eval.form
        forms.append(fa)
        self.sp_min = QSpinBox(minimum=2, maximum=20)
        fa.addRow(T("min_points"), self.sp_min)
        self.sp_r2 = QDoubleSpinBox(decimals=3, minimum=0.0, maximum=1.0, singleStep=0.005)
        self.sp_r2.setToolTip(T("min_r2_tip"))
        fa.addRow(T("min_r2"), self.sp_r2)
        self.cb_shoulders = QCheckBox(T("find_shoulders"))
        fa.addRow(self.cb_shoulders)
        f.addRow(self.sec_eval)
        compact(self.sp_d, self.sp_cut, self.sp_min, self.sp_r2)
        lv.addWidget(g)

        g, body = STYLE.card(T("grp_plot"), "plot")
        f = make_form(body)
        f.setContentsMargins(0, 0, 0, 0)
        forms.append(f)
        self.sp_ymax = QDoubleSpinBox(decimals=2, minimum=0.05, maximum=10, singleStep=0.1)
        compact(self.sp_ymax)
        f.addRow(T("ymax"), self.sp_ymax)
        row = QHBoxLayout()
        self.cb_xauto = QCheckBox(T("xrange_auto"))
        self.sp_xmin = QDoubleSpinBox(decimals=0, minimum=100, maximum=3500, singleStep=10)
        self.sp_xmax = QDoubleSpinBox(decimals=0, minimum=100, maximum=3500, singleStep=10)
        row.addWidget(self.cb_xauto)
        row.addWidget(self.sp_xmin)
        row.addWidget(QLabel("–"))
        row.addWidget(self.sp_xmax)
        f.addRow(T("xrange"), row)
        self.cmb_size = QComboBox()
        for key in ("half_a4", "custom", "std"):
            self.cmb_size.addItem(T({"half_a4": "size_half_a4", "custom": "size_custom",
                                     "std": "size_std"}[key]), key)
        f.addRow(T("size"), self.cmb_size)
        row = QHBoxLayout()
        self.sp_w = QDoubleSpinBox(decimals=1, minimum=4, maximum=40, singleStep=0.5, suffix=" cm")
        self.sp_h = QDoubleSpinBox(decimals=1, minimum=3, maximum=40, singleStep=0.5, suffix=" cm")
        row.addWidget(self.sp_w)
        row.addWidget(QLabel("×"))
        row.addWidget(self.sp_h)
        row.addStretch(1)
        f.addRow(T("size_wh"), row)
        self.sec_plot = Section(T("adv_plot"), "plot", self.settings)
        fa = self.sec_plot.form
        forms.append(fa)
        self.cmb_r2 = QComboBox()
        for key in ("legend", "label", "off"):
            self.cmb_r2.addItem(T({"legend": "r2_legend", "label": "r2_label", "off": "r2_off"}[key]), key)
        fa.addRow(T("r2_mode"), self.cmb_r2)
        self.cb_err = QCheckBox(T("show_err"))
        self.cb_marksh = QCheckBox(T("mark_sh"))
        self.cb_lidx = QCheckBox(T("label_index"))
        self.cb_lidx.setToolTip(T("label_index_tip"))
        self.cb_inset = QCheckBox(T("inset_show"))
        self.cb_tab = QCheckBox(T("inset_table"))
        self.cb_excl = QCheckBox(T("show_excl"))
        for cb in (self.cb_err, self.cb_marksh, self.cb_lidx, self.cb_inset, self.cb_tab, self.cb_excl):
            fa.addRow(cb)
        row = QHBoxLayout()
        self.cb_pdf, self.cb_svg, self.cb_png = QCheckBox("PDF"), QCheckBox("SVG"), QCheckBox("PNG")
        for cb in (self.cb_pdf, self.cb_svg, self.cb_png):
            row.addWidget(cb)
        row.addStretch(1)
        fa.addRow(T("formats"), row)
        f.addRow(self.sec_plot)
        lv.addWidget(g)

        hint = QLabel(T("drag_hint"))
        hint.setWordWrap(True)
        hint.setStyleSheet(HINT_CSS)
        lv.addWidget(hint)
        lv.addStretch(1)
        align_forms(forms)

        lscroll = QScrollArea()
        lscroll.setWidget(left)
        lscroll.setWidgetResizable(True)
        lscroll.setFrameShape(QFrame.NoFrame)
        # Exportieren bleibt unten sichtbar, auch wenn die Spalte gescrollt wird
        row = QHBoxLayout()
        row.setContentsMargins(8, 6, 8, 6)
        row.addWidget(uvvis_tabs.reset_button(self.reset_layout, self.reset_defaults))
        b = QPushButton(T("si_copy_methods"))
        b.setToolTip(T("si_copy_methods_tip"))
        b.clicked.connect(self.copy_eps_methods)
        row.addWidget(b)
        row.addStretch(1)
        self.btn_export = QPushButton(STYLE.icon("export", "#FFFFFF"), T("btn_export"))
        self.btn_export.setProperty("accent", True)
        self.btn_export.clicked.connect(self.export_dialog)
        self.btn_export.setToolTip(T("btn_export_tip"))
        self.btn_export.setDefault(True)
        row.addWidget(self.btn_export)
        left_col = QWidget()
        cv = QVBoxLayout(left_col)
        cv.setContentsMargins(0, 0, 0, 0)
        cv.setSpacing(0)
        cv.addWidget(lscroll, 1)
        bar = QFrame(objectName="bottomBar")
        bar.setLayout(row)
        cv.addWidget(bar)
        # so breit wie der Inhalt, sonst schneidet die Spalte rechts ab (macOS blendet die Leiste aus)
        lscroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        left_col.setMinimumWidth(max(380, left.minimumSizeHint().width()
                                     + lscroll.verticalScrollBar().sizeHint().width() + 4))

        # ---- rechte Seite
        self.canvas_host = uvvis_tabs.CanvasHost(self.canvas, self.fig)
        self.tabs = QTabWidget()                      # wie die anderen Tabs: Abbildung | Ergebnisse | SI-Zeile
        self.tabs.setDocumentMode(True)
        self.table = QTableWidget(0, 7)
        self.table.setHorizontalHeaderLabels([T("col_lam"), T("col_n"), T("col_coeff"),
                                              T("col_intercept"), T("col_r2"), T("col_cut"), T("col_warn")])
        for c_ in range(6):                           # Spaltenköpfe nicht abschneiden („Achsenabschnitt“)
            self.table.horizontalHeader().setSectionResizeMode(c_, QHeaderView.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(6, QHeaderView.Stretch)
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QTableWidget.DoubleClicked | QTableWidget.EditKeyPressed
                                   | QTableWidget.SelectedClicked)
        self.table.setToolTip(T("cut_tip"))
        self.table.itemChanged.connect(self.cutoff_edited)
        self._row_keys = []
        old_log = self.log.toPlainText() if hasattr(self, "log") else ""
        self.log = QPlainTextEdit(readOnly=True)
        self.log.setPlainText(old_log)
        g_res = QGroupBox()
        QVBoxLayout(g_res).addWidget(self.table)
        si_w = QGroupBox()
        si_l = QVBoxLayout(si_w)
        self.txt_si = QTextBrowser()
        si_l.addWidget(self.txt_si)
        hint = QLabel(T("si_uv_hint"))
        hint.setWordWrap(True)
        hint.setStyleSheet(HINT_CSS)
        si_l.addWidget(hint)
        b = QPushButton(T("uv_copy"))
        b.clicked.connect(lambda: IRMS.copy_rich(*self._uv_line))
        si_l.addWidget(b)
        self._uv_line = ("", "")
        # Protokoll: Hilfe → Protokoll (eigenes Fenster), kein Reiter mehr
        # Leer-Zustand: Anleitung (was herauskommt, welche Dateien, Schritte) statt weißer Fläche
        empty = uvvis_tabs.EmptyGuide("eps", (T("act_open"), self.choose_csv))
        empty_wrap = QWidget()
        QVBoxLayout(empty_wrap).addWidget(empty)
        self.preview_stack = QStackedWidget()
        self.preview_stack.addWidget(empty_wrap)
        self.preview_stack.addWidget(self.canvas_host)
        self.preview_stack.setCurrentIndex(1 if self.current and self.current in self.groups else 0)
        self.tabs.addTab(self.preview_stack, T("tab_figure"))
        self.tabs.addTab(uvvis_tabs.result_page(g_res), T("tab_results"))
        self.tabs.addTab(uvvis_tabs.result_page(si_w), T("tab_si_line"))
        right = self.tabs
        split = QSplitter(Qt.Horizontal)
        split.addWidget(left_col)
        split.addWidget(right)
        split.setSizes([left_col.minimumWidth(), 920])
        self.column_fit = STYLE.fit_column(left, lscroll, bar, left_col, split, floor=380)   # nie abgeschnitten
        self.eps_page = split
        self.ov_tab = uvvis_tabs.OverlayTab(self, self.ov_state)
        self.fl_tab = uvvis_tabs.FluoTab(self, self.fl_state)
        self.cv_tab = CVMOD.CVTab(self, self.cv_state)
        self.sec_tab = SECMOD.SECTab(self, self.sec_state)
        self.epr_tab = EPRMOD.EPRTab(self, self.epr_state)
        self.hrms_tab = HRMS.HRMSTab(self, self.hrms_state)
        self.irms_tab = IRMS.IRMSTab(self, self.irms_state)
        self.xray_tab = XRAY.XrayTab(self, self.xray_state)
        self.syn_tab = SYN.SynthTab(self, self.syn_state)
        self.nmr_tab = NMR.NMRTab(self, self.nmr_state)
        self.dft_tab = DFT.DFTTab(self, self.dft_state)
        # Reihenfolge wie die SI: Synthese, Charakterisierung, Spektroskopie, Elektrochemie, Struktur, Rechnung
        self.main_tabs = QTabWidget()
        self._tab_keys = []
        for key, page, label in (("syn", self.syn_tab, "tab_syn"), ("nmr", self.nmr_tab, "tab_nmr"),
                                 ("irms", self.irms_tab, "tab_irms"), ("hrms", self.hrms_tab, "tab_hrms"),
                                 ("eps", split, "tab_eps"), ("overlay", self.ov_tab, "tab_overlay"),
                                 ("fluo", self.fl_tab, "tab_fluo"), ("cv", self.cv_tab, "tab_cv"),
                                 ("sec", self.sec_tab, "tab_sec"), ("epr", self.epr_tab, "tab_epr"),
                                 ("xray", self.xray_tab, "tab_xray"), ("dft", self.dft_tab, "tab_dft")):
            self.main_tabs.addTab(page, T(label))
            self._tab_keys.append(key)
        self.main_tabs.setCurrentIndex(self._tab_keys.index(self._tab_key) if self._tab_key in self._tab_keys else 0)
        self.main_tabs.currentChanged.connect(lambda i: setattr(self, "_tab_key", self._tab_keys[i]))
        self.main_tabs.currentChanged.connect(                    # Vorlagen (CIF aus X-Ray) aktuell halten
            lambda i: self.dft_tab.fill_refs() if self.main_tabs.widget(i) is self.dft_tab else None)
        for i in range(self.main_tabs.count()):
            STYLE.tidy(self.main_tabs.widget(i))
        self.center = QStackedWidget()
        self.center.addWidget(self.welcome_page())          # 0: ohne Projekt
        self.center.addWidget(self.nav_bar())               # 1: Tabs
        self.setCentralWidget(self.center)
        self.center.setCurrentIndex(1 if self.project_dir else 0)

        self.apply_globals()
        self.apply_series_widgets()
        for w in (self.sp_d, self.sp_cut, self.sp_ymax, self.sp_bla, self.sp_blb, self.sp_w, self.sp_h,
                  self.sp_xmin, self.sp_xmax):
            w.valueChanged.connect(self.changed)
        self.cb_xauto.toggled.connect(self._xauto_toggled)
        self.cmb_size.currentIndexChanged.connect(self._size_changed)
        for cb in (self.cb_pdf, self.cb_svg, self.cb_png):
            cb.toggled.connect(lambda *_: self.collect_globals())
        self.sp_min.valueChanged.connect(self.changed)
        self.sp_r2.valueChanged.connect(self.changed)
        self.ed_bands.editingFinished.connect(self.changed)
        for w in (self.cmb_bl, self.cmb_r2):
            w.currentIndexChanged.connect(self.changed)
        for w in (self.cb_err, self.cb_tab, self.cb_bands_add, self.cb_shoulders, self.cb_excl, self.cb_inset,
                  self.cb_marksh, self.cb_lidx):
            w.toggled.connect(self.changed)
        for w in (self.sp_min, self.sp_r2):
            w.valueChanged.connect(self._update_sections)
        self.cmb_r2.currentIndexChanged.connect(self._update_sections)
        for w in (self.cb_shoulders, self.cb_err, self.cb_marksh, self.cb_lidx, self.cb_inset, self.cb_tab,
                  self.cb_excl, self.cb_pdf, self.cb_svg, self.cb_png):
            w.toggled.connect(self._update_sections)
        self._update_sections()

    def _update_sections(self, *_):
        """Eingeklappte Bereiche markieren, wenn dort etwas vom Standard abweicht."""
        d = DEFAULT_GLOB
        self.sec_images.set_changed(self.cmb_bg.currentData() not in (None, "auto"))
        self.sec_eval.set_changed(self.sp_min.value() != d["min_points"]
                                  or abs(self.sp_r2.value() - d["min_r2"]) > 1e-9
                                  or self.cb_shoulders.isChecked() != d["shoulders"])
        self.sec_plot.set_changed(self.cmb_r2.currentData() != d["r2"] or any(
            cb.isChecked() != d[k] for cb, k in (
                (self.cb_err, "show_err"), (self.cb_marksh, "mark_sh"), (self.cb_lidx, "label_index"),
                (self.cb_inset, "inset"),
                (self.cb_tab, "table"), (self.cb_excl, "show_excl"), (self.cb_pdf, "fmt_pdf"),
                (self.cb_svg, "fmt_svg"), (self.cb_png, "fmt_png"))))

    def apply_globals(self):
        self._loading = True
        gl = self.glob
        self.sp_d.setValue(gl["path_length"])
        self.sp_cut.setValue(gl["cutoff"])
        self.sp_min.setValue(gl["min_points"])
        self.sp_r2.setValue(gl.get("min_r2", 0.98))
        self.ed_bands.setText(gl["bands"])
        self.cmb_bl.setCurrentIndex(self.cmb_bl.findData(gl["baseline"]))
        self.sp_bla.setValue(gl["bl_a"])
        self.sp_blb.setValue(gl["bl_b"])
        self.sp_ymax.setValue(gl["ymax"])
        self.cmb_r2.setCurrentIndex(self.cmb_r2.findData(gl["r2"]))
        self.cb_err.setChecked(gl["show_err"])
        self.cb_tab.setChecked(gl["table"])
        self.cmb_size.setCurrentIndex(max(0, self.cmb_size.findData(gl["size"])))
        self.sp_w.setValue(gl["w_cm"])
        self.sp_h.setValue(gl["h_cm"])
        self.cb_pdf.setChecked(gl["fmt_pdf"])
        self.cb_svg.setChecked(gl["fmt_svg"])
        self.cb_png.setChecked(gl["fmt_png"])
        self.cb_xauto.setChecked(gl["x_auto"])
        self.cb_bands_add.setChecked(gl.get("bands_add", True))
        self.cb_shoulders.setChecked(gl.get("shoulders", False))
        self.cb_excl.setChecked(gl.get("show_excl", False))
        self.cb_inset.setChecked(gl.get("inset", True))
        self.cb_marksh.setChecked(gl.get("mark_sh", False))
        self.cb_lidx.setChecked(gl.get("label_index", False))
        self.cb_tab.setEnabled(gl.get("inset", True))
        self.sp_xmin.setValue(gl["x_min"])
        self.sp_xmax.setValue(gl["x_max"])
        self.sp_xmin.setEnabled(not gl["x_auto"])
        self.sp_xmax.setEnabled(not gl["x_auto"])
        self._update_size_widgets()
        manual = gl["baseline"] == "manual"
        self.sp_bla.setEnabled(manual)
        self.sp_blb.setEnabled(manual)
        self.cb_color.setChecked(gl.get("color_show", False))
        self.sp_color_d.setValue(gl.get("color_d", 1.0))
        self._loading = False
        self.update_color()

    def collect_globals(self):
        self.glob.update(path_length=self.sp_d.value(), cutoff=self.sp_cut.value(),
                         min_points=self.sp_min.value(), min_r2=self.sp_r2.value(), bands=self.ed_bands.text().strip(),
                         baseline=self.cmb_bl.currentData(), bl_a=self.sp_bla.value(),
                         bl_b=self.sp_blb.value(), ymax=self.sp_ymax.value(),
                         r2=self.cmb_r2.currentData(), show_err=self.cb_err.isChecked(),
                         table=self.cb_tab.isChecked(), size=self.cmb_size.currentData(),
                         w_cm=self.sp_w.value(), h_cm=self.sp_h.value(),
                         fmt_pdf=self.cb_pdf.isChecked(), fmt_svg=self.cb_svg.isChecked(),
                         fmt_png=self.cb_png.isChecked(), x_auto=self.cb_xauto.isChecked(),
                         bands_add=self.cb_bands_add.isChecked(), shoulders=self.cb_shoulders.isChecked(),
                         show_excl=self.cb_excl.isChecked(), inset=self.cb_inset.isChecked(),
                         mark_sh=self.cb_marksh.isChecked(), label_index=self.cb_lidx.isChecked(),
                         x_min=self.sp_xmin.value(), x_max=self.sp_xmax.value())

    def _update_size_widgets(self):
        key = self.cmb_size.currentData()
        if key == "half_a4":
            self.sp_w.blockSignals(True)
            self.sp_h.blockSignals(True)
            self.sp_w.setValue(16.0)
            self.sp_h.setValue(11.0)
            self.sp_w.blockSignals(False)
            self.sp_h.blockSignals(False)
        custom = key == "custom"
        self.sp_w.setEnabled(custom)
        self.sp_h.setEnabled(custom)

    def apply_series_widgets(self):
        self._loading = True
        st = self.series.get(self.current)
        self.ed_mw.setText(f"{st['mw']:g}" if st and st["mw"] else "")
        self.lbl_struct.setText(Path(st["structure"]).name if st and st["structure"] else "–")
        self.lbl_photo.setText(Path(st["photo"]).name if st and st["photo"] else "–")
        self.cmb_bg.setCurrentIndex(self.cmb_bg.findData(st["bg"] if st else "auto"))
        info = st.get("chem_info", "") if st else ""
        unit = self.units.get(self.current) if st else None
        warn = bool(st and st.get("chem_warn"))
        if unit and unit != "mg/mL":
            info = (info + "\n" if info else "") + T("mw_unused", u="µM" if unit == "uM" else unit)
        mw_missing = unit == "mg/mL" and not (st and st["mw"])
        if mw_missing:
            info = (info + "\n" if info else "") + T("mw_needed")
            warn = True
        self.ed_mw.setStyleSheet("border: 2px solid #d9534f;" if mw_missing else "")
        self.lbl_chem.setText(info)
        self.lbl_chem.setVisible(bool(info))          # keine leere Zeile unter der Molmasse
        # fehlende Molmasse = Fehler (rot); Hinweise zur Zeichnung (aus CDXML statt ChemDraw-Bild) nur orange
        self.lbl_chem.setStyleSheet(WARN_CSS if mw_missing else STYLE.NOTE_CSS if warn else "")
        pa = st.get("photo_asset") if st else None
        self.lbl_thumb.setPixmap(rgba_to_pixmap(pa["rgba"]) if pa else QPixmap())
        self.lbl_thumb.setVisible(bool(pa))              # ohne Foto keine Leerzeile
        self.sp_struct_w.setValue(st.get("struct_w", 17.0) if st else 17.0)
        self.sp_photo_w.setValue(st.get("photo_w", 9.0) if st else 9.0)
        self.sp_struct_w.setEnabled(bool(st and st["structure"]))
        self.sp_photo_w.setEnabled(bool(st and st["photo"]))
        self._loading = False

    def _xauto_toggled(self, on):
        if self._loading:
            return
        self.sp_xmin.setEnabled(not on)
        self.sp_xmax.setEnabled(not on)
        if not on and self.current in self.groups:       # mit dem Datenbereich vorbelegen
            xs = np.concatenate([s_["x"] for s_ in self.groups[self.current]])
            self._loading = True
            if self.sp_xmin.value() <= xs.min() or self.sp_xmin.value() >= xs.max():
                self.sp_xmin.setValue(float(np.floor(xs.min())))
            if self.sp_xmax.value() >= xs.max() or self.sp_xmax.value() <= self.sp_xmin.value():
                self.sp_xmax.setValue(float(np.ceil(xs.max())))
            self._loading = False
        self.changed()

    def _size_changed(self, *_):
        if self._loading:
            return
        self._update_size_widgets()
        self.changed()

    def _theme_mode(self):
        return self.settings.value("ui/theme", "system") or "system"

    def apply_theme(self):
        STYLE.apply(self._theme_mode(), IMG.app_data_dir() / "theme")

    def switch_theme(self, mode):
        self.collect_globals()
        self.settings.setValue("ui/theme", mode)
        self.apply_theme()
        self.build_ui()                                # Symbole in der neuen Textfarbe
        self.render()

    def _system_scheme_changed(self, *_):
        if self._theme_mode() == "system" and STYLE.resolve("system") != STYLE._mode:
            self.switch_theme("system")

    def switch_lang(self, code):
        self.collect_globals()
        set_lang(code)
        self.settings.setValue("lang", code)
        # Hinweise zur Struktur werden beim Laden übersetzt und gespeichert -> in der neuen Sprache neu laden
        for g, st in self.series.items():
            if st.get("structure") and Path(st["structure"]).exists():
                self.set_structure(st["structure"], g, render=False, sync=True, autofill_mw=False)
        self.build_ui()
        self.render()

    NAV_GROUPS = (("nav_grp_syn", ("syn", "nmr", "irms", "hrms")), ("nav_grp_uv", ("eps", "overlay", "fluo")),
                  ("nav_grp_ec", ("cv", "sec", "epr")), ("nav_grp_struct", ("xray", "dft")))

    def nav_bar(self):
        """Tabs in vier Gruppen wie die SI (Leiste oben statt Reiter): kostet Höhe statt Breite – die Abbildung
        ist in der Breite begrenzt, eine Seitenleiste würde sie um ~20 % verkleinern."""
        from PySide6.QtWidgets import QButtonGroup
        bar = QFrame(objectName="navBar")
        hb = QHBoxLayout(bar)
        hb.setContentsMargins(10, 4, 10, 6)
        hb.setSpacing(4)
        grp = QButtonGroup(bar)
        self.nav_buttons = {}
        for gi, (head, keys) in enumerate(self.NAV_GROUPS):
            if gi:
                sep = QFrame(objectName="navSep")
                sep.setFixedWidth(1)
                hb.addSpacing(8)
                hb.addWidget(sep)
                hb.addSpacing(8)
            col = QVBoxLayout()
            col.setSpacing(1)
            col.addWidget(QLabel(T(head).upper(), objectName="navGroup"), 0, Qt.AlignHCenter)   # mittig über der Gruppe
            row = QHBoxLayout()
            row.setSpacing(2)
            for k in keys:
                i = self._tab_keys.index(k)
                b = QPushButton(self.main_tabs.tabText(i), objectName="navTab", checkable=True)
                b.clicked.connect(lambda _=False, i=i: self.main_tabs.setCurrentIndex(i))
                grp.addButton(b)
                row.addWidget(b)
                self.nav_buttons[i] = b
            col.addLayout(row)
            hb.addLayout(col)
        hb.addStretch(1)
        self.main_tabs.currentChanged.connect(lambda i: self.nav_buttons[i].setChecked(True))
        self.nav_buttons[self.main_tabs.currentIndex()].setChecked(True)
        self.main_tabs.tabBar().hide()
        self.main_tabs.setDocumentMode(True)
        holder = QWidget()
        v = QVBoxLayout(holder)
        v.setContentsMargins(0, 0, 0, 0)
        v.setSpacing(0)
        v.addWidget(bar)
        v.addWidget(self.main_tabs, 1)
        return holder

    def _append_log(self, msg):
        self.log.appendPlainText(msg)

    def show_log(self):
        """Protokoll aller Tabs (Meldungen, Warnungen, Fehlerdetails) in einem eigenen Fenster."""
        dlg = QDialog(self)
        dlg.setWindowTitle(T("log_title"))
        dlg.resize(820, 520)
        v = QVBoxLayout(dlg)
        txt = QPlainTextEdit(readOnly=True)
        txt.setPlainText(self.log.toPlainText())
        txt.moveCursor(QTextCursor.End)
        v.addWidget(txt)
        bb = QDialogButtonBox(QDialogButtonBox.Close)
        bb.rejected.connect(dlg.reject)
        v.addWidget(bb)
        dlg.exec()

    # ------------------------------------------------------------ Daten
    def choose_csv(self):
        if not self.ensure_project():
            return
        start = self.settings.value("last_dir", str(Path.home()))
        ps, _ = QFileDialog.getOpenFileNames(self, T("dlg_open_csv"), start,
                                             f"{T('flt_csv')};;{T('flt_all')}")
        if ps:
            self.open_files(ps)

    def open_csv(self, path, ask=True):
        return self.open_files([path], ask)

    def files_label(self):
        names = [Path(p).name for p in self.csv_paths]
        if len(names) == 1:
            return names[0]
        short = ", ".join(names[:3]) + (", …" if len(names) > 3 else "")
        return T("files_loaded", n=len(names), names=short)

    def open_files(self, paths, ask=True):
        """Eine oder mehrere Messdateien (CSV/TXT/DSW/BSW) gemeinsam öffnen,
        z. B. wenn jede Konzentration einzeln gemessen wurde."""
        if not self.ensure_project():
            return False
        try:
            paths = sorted(self.import_file(p, "data") for p in paths)   # Kopie ins Projekt
        except OSError as e:
            self.error(str(e))
            return False
        path = paths[0]
        try:
            samples = core.read_samples(paths)
        except Exception as e:
            self.error(f"{type(e).__name__}: {e}")
            return False
        if not samples:
            self.error(T("conc_none"))
            return False
        self.csv_path, self.csv_paths, self.samples = path, paths, samples
        self.settings.setValue("last_dir", str(Path(path).parent))
        self.groups, self.units, self.series, self.conc = {}, {}, {}, {}
        _, _, missing = core.group_samples(samples, self.conc, Path(path).stem)
        if missing and ask and not self.edit_concentrations(rebuild=False):
            pass
        if not self.rebuild_groups():
            self.lbl_file.setText(self.files_label())
            self.error(T("conc_none"))
            return False
        self.lbl_file.setText(self.files_label())
        self.apply_compound_defaults()
        self.apply_globals()
        self.apply_series_widgets()
        self.render()
        return True

    def edit_concentrations(self, rebuild=True):
        if not self.samples:
            self.error(T("load_first"))
            return False
        dlg = ConcDialog(self, self.samples, self.conc, Path(self.csv_path).stem)
        if dlg.exec() != QDialog.Accepted:
            return False
        self.conc = dlg.overrides()
        if rebuild:
            self.rebuild_groups()
            self.apply_series_widgets()
            self.render()
        return True

    def rebuild_groups(self):
        """Serien aus Proben + Konzentrationen neu bilden; Einstellungen je Serie bleiben."""
        try:
            groups, units, missing = core.group_samples(self.samples, self.conc, Path(self.csv_path).stem)
        except ValueError as e:
            self.error(str(e))
            return False
        for k in missing:
            self.log_signal.emit("  " + T("no_conc_in_name", name=k))
        self.groups, self.units = groups, units
        for g in groups:
            self.series.setdefault(g, {"mw": None, "structure": None, "photo": None, "bg": "auto",
                                       "layout": {}, "structure_asset": None, "photo_asset": None,
                                       "chem_info": "", "chem_warn": False,
                                       "struct_w": 17.0, "photo_w": 9.0, "band_cutoff": {},
                                       "band_hidden": []})
        if self.current not in groups:
            self.current = next(iter(groups), None)
        self.cmb_series.blockSignals(True)
        self.cmb_series.clear()
        self.cmb_series.addItems(list(groups))
        if self.current:
            self.cmb_series.setCurrentText(self.current)
        self.cmb_series.blockSignals(False)
        log_names = ", ".join(groups) or "–"
        self.log_signal.emit(T("cary_series", name=Path(self.csv_path).name, s=log_names))
        return bool(groups)

    def restart_file(self):
        """Eingaben zu den aktuellen Messdateien verwerfen (Konzentrationen, Serien-Einstellungen)."""
        if not self.csv_path:
            return self.error(T("load_first"))
        if os.environ.get("QT_QPA_PLATFORM") != "offscreen":
            ans = QMessageBox.question(self, T("act_restart").rstrip("…"),
                                       T("restart_q", name=self.files_label()))
            if ans != QMessageBox.Yes:
                return False
        paths = list(self.csv_paths)
        self.glob = dict(DEFAULT_GLOB)
        self.series, self.conc, self.current = {}, {}, None
        return self.open_files(paths)

    # ------------------------------------------------------------ Projekt
    def get_state(self):
        """Gesamter Zustand aller Tabs (serialisierbar) – für Projektdatei und Undo."""
        if hasattr(self, "sp_d"):
            self.collect_globals()
        keys = ("mw", "structure", "photo", "bg", "layout", "struct_w", "photo_w", "band_cutoff", "band_hidden",
                "color_c")
        return {"eps": {"files": list(self.csv_paths), "glob": dict(self.glob), "conc": self.conc,
                        "series": {g: {k: st.get(k) for k in keys} for g, st in self.series.items()},
                        "current": self.current},
                "overlay": self.ov_state, "fluo": self.fl_state, "cv": self.cv_state, "sec": self.sec_state, "epr": self.epr_state,
                "irms": self.irms_state, "hrms": self.hrms_state, "xray": self.xray_state,
                "dft": self.dft_state, "nmr": self.nmr_state, "syn": self.syn_state,
                "compound": dict(self.compound),
                "tab": self._tab_key,
                "text": dict(self.text_fmt), "text_tabs": {k: dict(v) for k, v in self.tab_fmt.items()}}

    def set_state(self, state):
        """Zustand übernehmen (Projekt laden, Undo/Redo)."""
        state = json.loads(json.dumps(state))          # entkoppeln
        eps = state.get("eps") or {}
        old_series = self.series
        self.glob = dict(DEFAULT_GLOB, **(eps.get("glob") or {}))
        self.conc = eps.get("conc") or {}
        self.csv_paths = [p for p in eps.get("files") or [] if Path(p).exists()]
        self.csv_path = self.csv_paths[0] if self.csv_paths else None
        self.groups, self.units, self.series, self.samples = {}, {}, {}, []
        self.ov_state.clear()
        self.ov_state.update(state.get("overlay") or {})
        self.fl_state.clear()
        self.fl_state.update(state.get("fluo") or {})
        self.cv_state.clear()
        self.cv_state.update(state.get("cv") or {})
        self.sec_state.clear()
        self.sec_state.update(state.get("sec") or {})
        self.epr_state.clear()
        self.epr_state.update(state.get("epr") or {})
        self.irms_state.clear()
        self.irms_state.update(state.get("irms") or {})
        self.hrms_state.clear()
        hr = state.get("hrms")
        if hr is None:                                 # Projekte vor 1.21: MS-Teil stand im Tab IR + MS
            ir = state.get("irms") or {}
            hr = {k: ir[k] for k in ("ms_file", "ms_raw", "ms_rt", "adduct", "found") if k in ir}
        self.hrms_state.update(hr)
        self.xray_state.clear()
        self.xray_state.update(state.get("xray") or {})
        self.dft_state.clear()
        self.dft_state.update(state.get("dft") or {})
        self.nmr_state.clear()
        self.nmr_state.update(state.get("nmr") or {})
        self.syn_state.clear()
        self.syn_state.update(state.get("syn") or {})
        self.compound = dict(state.get("compound") or {})
        self.text_fmt = dict(core.TEXT_DEFAULTS, **(state.get("text") or {}))
        self.tab_fmt = {k: dict(v) for k, v in (state.get("text_tabs") or {}).items()}
        tab = state.get("tab", "syn")
        if isinstance(tab, int):                         # bis 1.23 gespeichert als Position (alte Reihenfolge)
            old = ("eps", "overlay", "fluo", "cv", "sec", "epr", "hrms", "irms", "xray", "dft")
            tab = old[tab] if 0 <= tab < len(old) else "syn"
        self._tab_key = tab
        self.current = eps.get("current")
        self.build_ui()
        if self.csv_paths:
            key = ("samples_multi", tuple(self.csv_paths))
            if key not in self.extra_cache:
                self.extra_cache[key] = core.read_samples(self.csv_paths)
            self.samples = self.extra_cache[key]
            self.rebuild_groups()
            for g, sd in (eps.get("series") or {}).items():
                if g not in self.series:
                    continue
                st = self.series[g]
                for k, v in sd.items():
                    if v is not None and k not in ("structure", "photo"):
                        st[k] = v
                st["layout"] = st.get("layout") or {}
                old = old_series.get(g) or {}
                for kind, setter in (("structure", self.set_structure), ("photo", self.set_photo)):
                    p = sd.get(kind)
                    if p and Path(p).exists():
                        if old.get(kind) == p and old.get(kind + "_asset") is not None:
                            st[kind], st[kind + "_asset"] = p, old[kind + "_asset"]   # wiederverwenden
                            if kind == "structure":
                                st["chem_info"], st["chem_warn"] = old.get("chem_info", ""), old.get("chem_warn")
                        elif kind == "structure":
                            setter(p, g, render=False, sync=True, autofill_mw=False)
                        else:
                            setter(p, g, render=False, sync=True)
            if self.current not in self.groups:
                self.current = next(iter(self.groups), None)
            self.cmb_series.blockSignals(True)
            self.cmb_series.setCurrentText(self.current or "")
            self.cmb_series.blockSignals(False)
            self.lbl_file.setText(self.files_label())
        self.apply_globals()
        self.apply_series_widgets()
        if self.csv_paths:
            self.render()
        else:
            self.fig.clear()
            self.canvas.draw_idle()

    def _snap(self):
        return json.dumps(self.get_state(), sort_keys=True, default=str)

    def _init_undo(self):
        self._undo, self._redo = [self._snap()], []
        self._saved_snap = self._undo[0]
        self._update_undo_actions()
        self.update_title()

    def push_undo(self):
        if self._restoring or not self._undo:
            return
        snap = self._snap()
        if snap != self._undo[-1]:
            self._undo.append(snap)
            del self._undo[:-200]
            self._redo.clear()
            self._update_undo_actions()
            self.update_title()
            if self.project_dir:
                self.autosave_timer.start()

    def _restore(self, snap):
        self._restoring = True
        self.set_state(json.loads(snap))
        QTimer.singleShot(200, self._end_restore)       # erst nach dem Neuzeichnen aller Tabs

    def _end_restore(self):
        self._restoring = False
        if self._undo:
            self._undo[-1] = self._snap()              # tatsächlichen Zustand übernehmen
        self._update_undo_actions()
        self.update_title()
        if self.project_dir:
            self.autosave_timer.start()                # Undo/Redo-Stand ebenfalls speichern

    def _load_state_into_ui(self, state):
        """Projekt/neues Projekt: Zustand setzen, Undo-Verlauf nach dem Neuzeichnen neu beginnen."""
        self._restoring = True
        self.set_state(state)
        QTimer.singleShot(200, self._after_load)

    def _after_load(self):
        self._restoring = False
        self._init_undo()

    def undo(self):
        if len(self._undo) > 1:
            self._redo.append(self._undo.pop())
            self._restore(self._undo[-1])

    def redo(self):
        if self._redo:
            self._undo.append(self._redo.pop())
            self._restore(self._undo[-1])

    def _update_undo_actions(self):
        if hasattr(self, "act_undo"):
            self.act_undo.setEnabled(len(self._undo) > 1)
            self.act_redo.setEnabled(bool(self._redo))

    def is_dirty(self):
        if self._restoring:
            return False
        return bool(self._undo) and self._undo[-1] != self._saved_snap

    def welcome_page(self):
        """Startbildschirm ohne Projekt: worum es geht, Neues Projekt / Öffnen / Beispielprojekt, zuletzt geöffnet."""
        page = QWidget()
        outer = QVBoxLayout(page)
        outer.addStretch(1)
        card = QFrame(objectName="card")
        card.setMaximumWidth(620)
        v = QVBoxLayout(card)
        v.setContentsMargins(32, 28, 32, 28)
        v.setSpacing(12)
        head = QHBoxLayout()
        icon = resource_path("assets/icon.png")
        if icon.exists():
            lb = QLabel()
            lb.setPixmap(QIcon(str(icon)).pixmap(64, 64))
            head.addWidget(lb)
        t = QLabel(f"<span style='font-size:22px; font-weight:bold'>{APP_NAME}</span><br>{T('welcome_sub')}")
        t.setWordWrap(True)
        head.addWidget(t, 1)
        v.addLayout(head)
        row = QHBoxLayout()
        b_new = QPushButton(STYLE.icon("new", "#FFFFFF"), T("act_new_project"))
        b_new.setProperty("accent", True)
        b_new.clicked.connect(self.new_project)
        b_open = QPushButton(STYLE.icon("open"), T("act_open_project"))
        b_open.clicked.connect(self.open_project_dialog)
        b_demo = QPushButton(T("act_demo"))
        b_demo.setToolTip(T("act_demo_tip"))
        b_demo.clicked.connect(self.open_demo)
        for b in (b_new, b_open, b_demo):
            row.addWidget(b)
        row.addStretch(1)
        v.addLayout(row)
        self.lst_recent = QListWidget()
        self.lst_recent.setMaximumHeight(150)
        self.lst_recent.itemActivated.connect(lambda it: self.open_project(it.data(Qt.UserRole)))
        self.lst_recent.itemClicked.connect(lambda it: self.open_project(it.data(Qt.UserRole)))
        self.lbl_recent = QLabel(f"<b>{T('welcome_recent')}</b>")
        v.addWidget(self.lbl_recent)
        v.addWidget(self.lst_recent)
        hint = QLabel(T("welcome_hint", url=uvvis_tabs.README_URL + ("README.de.md" if get_lang() == "de" else "README.md")
                        + "#" + uvvis_tabs.gh_anchor(T("readme_quickstart"))))
        hint.setWordWrap(True)
        hint.setOpenExternalLinks(True)
        hint.setStyleSheet(HINT_CSS)
        v.addWidget(hint)
        row = QHBoxLayout()
        row.addStretch(1)
        row.addWidget(card, 4)
        row.addStretch(1)
        outer.addLayout(row)
        outer.addStretch(2)
        self.fill_recent_list()
        return page

    def fill_recent_list(self):
        if getattr(self, "lst_recent", None) is None:
            return
        self.lst_recent.clear()
        rec = [r for r in self.recent() if PRJ.is_project(r)]
        for r in rec:
            it = QListWidgetItem(f"{Path(r).name}  –  {Path(r).parent}")
            it.setData(Qt.UserRole, r)
            self.lst_recent.addItem(it)
        self.lst_recent.setVisible(bool(rec))
        self.lbl_recent.setVisible(bool(rec))

    def open_demo(self, folder=None):
        """Beispielprojekt mit künstlichen Daten (alle Tabs) anlegen und öffnen."""
        import uvvis_demo as DEMO
        self.maybe_save()
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            d = DEMO.open_demo(self, folder)
        except Exception as e:
            QApplication.restoreOverrideCursor()
            self.log_signal.emit(traceback.format_exc())
            return self.error(str(e))
        QApplication.restoreOverrideCursor()
        self.statusBar().showMessage(T("demo_done", p=str(d)), 10000)
        return d

    def show_tab_guide(self):
        key = self._tab_keys[self.main_tabs.currentIndex()]
        QDesktopServices.openUrl(QUrl(uvvis_tabs.guide_url(key)))

    def show_quickstart(self):
        url = (uvvis_tabs.README_URL + ("README.de.md" if get_lang() == "de" else "README.md") + "#"
               + uvvis_tabs.gh_anchor(T("readme_quickstart")))
        QDesktopServices.openUrl(QUrl(url))

    def update_title(self):
        if getattr(self, "center", None) is not None:
            self.center.setCurrentIndex(1 if self.project_dir else 0)
            self.fill_recent_list()
        if getattr(self, "act_folder", None) is not None:
            self.act_folder.setEnabled(bool(self.project_dir))
        name = Path(self.project_dir).name if self.project_dir else T("no_project_title")
        if self.compound.get("name"):
            name = f"{self.compound['name']} · {name}"
        self.setWindowTitle(f"{name} – {T('app_title')} {APP_VERSION}")

    # -- Verbindung des Projekts (Vorgabe für alle Tabs) ------------------------------------
    def compound_label(self):
        c = self.compound
        if not (c.get("name") or c.get("formula")):
            return T("compound_none")
        return T("compound_set", c=" · ".join(x for x in (c.get("name"), c.get("formula")) if x))

    def edit_compound(self):
        if not self.ensure_project():
            return
        dlg = CompoundDialog(self, self.compound)
        if dlg.exec() != QDialog.Accepted:
            return
        self.set_compound(dlg.result_compound())

    def set_compound(self, comp):
        """Verbindung übernehmen und als Vorgabe eintragen, wo noch nichts gewählt ist."""
        comp = dict(comp)
        if comp.get("structure") and self.project_dir:
            comp["structure"] = self.import_file(comp["structure"], "images")
        self.compound = comp
        self.apply_compound_defaults()
        if hasattr(self, "act_compound"):
            self.act_compound.setText(self.compound_label())
        self.update_title()
        if hasattr(self, "irms_tab"):
            self.irms_tab.compound_changed()
        if hasattr(self, "hrms_tab"):
            self.hrms_tab.compound_changed()
        self.render()
        self.push_undo()

    def apply_compound_defaults(self):
        """ε-Serien ohne Struktur/Molmasse und der Fluoreszenz-Tab bekommen die Verbindung des Projekts.
        Was schon gewählt ist, bleibt (Vergleichsprojekte, mehrere Verbindungen in einer Cary-Datei)."""
        c = self.compound
        p = c.get("structure")
        for g, st in self.series.items():
            if p and not st.get("structure") and Path(p).exists():
                self.set_structure(p, g, render=False)
            if not st.get("mw") and c.get("mw"):
                st["mw"] = c["mw"]
        if p and not self.fl_state.get("structure") and Path(p).exists():
            self.fl_state["structure"] = p
            if hasattr(self, "fl_tab"):
                self.fl_tab.timer.start()

    # -- Projektordner zuerst: Dateien werden beim Laden kopiert, Änderungen automatisch gespeichert
    def ensure_project(self):
        if self.project_dir:
            return True
        if os.environ.get("QT_QPA_PLATFORM") == "offscreen":
            return False
        if QMessageBox.question(self, T("app_title"), T("need_project_q")) != QMessageBox.Yes:
            return False
        return self.choose_project_folder()

    def choose_project_folder(self):
        d = QFileDialog.getExistingDirectory(self, T("choose_project_dir"),
                                             self.settings.value("last_project_dir", str(Path.home())))
        if not d:
            return False
        if PRJ.is_project(d):
            return self.open_project(d, ask=False)
        return self.create_project(d)

    def create_project(self, d):
        try:
            pdir = PRJ.create(d)
        except OSError as e:
            self.error(str(e))
            return False
        self.project_dir = str(pdir)
        self.settings.setValue("last_project_dir", str(pdir.parent))
        self._load_state_into_ui({})
        self.add_recent(pdir)
        self.log_signal.emit(T("project_created", p=pdir))
        return True

    def import_file(self, path, sub):
        if not self.project_dir:
            return str(Path(path).resolve())
        new = PRJ.import_file(self.project_dir, path, sub)
        if Path(new) != Path(path).resolve():
            self.log_signal.emit("  " + T("file_imported", src=Path(path).name, dst=Path(new).relative_to(
                self.project_dir).as_posix()))
        return new

    def maybe_save(self):
        """Vor Neu/Öffnen/Schließen: ausstehendes automatisches Speichern sofort ausführen."""
        if self.autosave_timer.isActive():
            self.autosave_timer.stop()
            self.autosave()
        return True

    def autosave(self):
        if not self.project_dir or self._restoring:
            return
        try:
            PRJ.write(self.project_dir, self.get_state())
            self._saved_snap = self._undo[-1] if self._undo else None
            self.statusBar().showMessage(T("autosaved"), 2000)
        except Exception as e:
            self.log_signal.emit(f"{T('err_title')}: {e}")

    def new_project(self):
        self.maybe_save()
        self.choose_project_folder()

    def open_project_dialog(self):
        self.maybe_save()
        d = QFileDialog.getExistingDirectory(self, T("act_open_project"),
                                             self.settings.value("last_project_dir", str(Path.home())))
        if not d:
            return
        if PRJ.is_project(d):
            self.open_project(d, ask=False)
        elif QMessageBox.question(self, T("app_title"), T("make_project_q", p=d)) == QMessageBox.Yes:
            self.create_project(d)

    def open_project(self, path, ask=True):
        pdir = PRJ.is_project(path)
        if pdir is None:
            self.error(T("no_project", p=path))
            return False
        self.maybe_save()
        try:
            PRJ.backup(pdir)                           # Sicherung: project.uvvis.bak
            state = PRJ.load(pdir)
        except Exception as e:
            self.error(f"{type(e).__name__}: {e}")
            return False
        for m in PRJ.missing_files(state):
            self.log_signal.emit("  " + T("missing_file", p=m))
        self.project_dir = str(pdir)
        self._load_state_into_ui(state)
        self.add_recent(pdir)
        self.settings.setValue("last_project_dir", str(pdir.parent))
        return True

    def save_project(self):
        if not self.project_dir:
            return self.ensure_project()
        self.autosave_timer.stop()
        self.autosave()
        return True

    def save_project_copy(self):
        """Projekt (inkl. aller Dateien) als Kopie in einen neuen Ordner schreiben, z. B. als Zwischenstand."""
        if not self.project_dir:
            return self.error(T("load_first"))
        d = QFileDialog.getExistingDirectory(self, T("act_save_copy"), str(Path(self.project_dir).parent))
        if d:
            return self._write_project(Path(d), switch=False)
        return False

    def _write_project(self, pdir: Path, switch=True):
        try:
            abs_state = PRJ.save(pdir, self.get_state())
        except Exception as e:
            self.error(f"{type(e).__name__}: {e}")
            return False
        if switch:
            self.project_dir = str(pdir)
            self._load_state_into_ui(abs_state)
            self.add_recent(pdir)
        self.log_signal.emit(T("project_saved", p=pdir))
        return True

    def show_project_folder(self):
        if self.project_dir:
            QDesktopServices.openUrl(QUrl.fromLocalFile(self.project_dir))

    def add_recent(self, pdir):
        rec = [str(pdir)] + [r for r in self.recent() if r != str(pdir)]
        self.settings.setValue("recent_projects", json.dumps(rec[:10]))
        self.fill_recent_menu()

    def recent(self):
        try:
            return [r for r in json.loads(self.settings.value("recent_projects", "[]") or "[]")]
        except Exception:
            return []

    def fill_recent_menu(self):
        if not hasattr(self, "recent_menu"):
            return
        self.recent_menu.clear()
        rec = [r for r in self.recent() if PRJ.is_project(r)]
        for r in rec:
            act = QAction(f"{Path(r).name}  –  {Path(r).parent}", self)
            act.triggered.connect(lambda _=False, r=r: self.open_project(r))
            self.recent_menu.addAction(act)
        self.recent_menu.setEnabled(bool(rec))

    def series_changed(self, g):
        if not g or g not in self.series:
            return
        self.current = g
        self.apply_series_widgets()
        self.render()

    def changed(self, *_):
        if self._loading:
            return
        self.collect_globals()
        manual = self.glob["baseline"] == "manual"
        self.sp_bla.setEnabled(manual)
        self.sp_blb.setEnabled(manual)
        self.cb_tab.setEnabled(self.glob.get("inset", True))
        self.timer.start()

    def mw_edited(self):
        if self._loading or not self.current:
            return
        try:
            self.series[self.current]["mw"] = parse_float(self.ed_mw.text())
        except ValueError:
            self.ed_mw.setText("")
            self.series[self.current]["mw"] = None
        self.apply_series_widgets()                 # Warnung zur Molmasse sofort aktualisieren
        self.timer.start()

    # ------------------------------------------------------- Struktur/Foto
    def choose_structure(self):
        if not self.current:
            return self.error(T("load_first"))
        p, _ = QFileDialog.getOpenFileName(self, T("dlg_structure"), self.settings.value("last_img_dir", ""),
                                           f"{T('flt_struct')};;{T('flt_all')}")
        if p:
            self.settings.setValue("last_img_dir", str(Path(p).parent))
            self.set_structure(p)

    def set_structure(self, path, g=None, render=True, sync=False, autofill_mw=True):
        g = g or self.current
        st = self.series[g]
        if path is not None and self.project_dir:
            path = self.import_file(path, "images")
        if path is None:
            st.update(structure=None, structure_asset=None, chem_info="", chem_warn=False)
            self.apply_series_widgets()
            return self.render()

        def work():
            asset = IMG.load_structure(path)
            p = Path(path)
            if asset["mw"] is None and p.suffix.lower() != ".cdxml":
                sib = p.with_suffix(".cdxml")            # gleichnamige CDXML neben SVG/PDF
                if sib.exists():
                    import uvvis_chem
                    chem = uvvis_chem.load_cdxml(sib)
                    asset.update(mw=chem["mw"], formula=chem["formula"], warnings=chem["warnings"])
            return asset

        def done(asset):
            st["structure"], st["structure_asset"] = str(path), asset
            info = []
            if asset.get("formula") and asset.get("mw"):
                info.append(T("mw_formula", f=asset["formula"], mw=asset["mw"]))
                if autofill_mw:
                    st["mw"] = round(asset["mw"], 2)
            if asset.get("drawing_from"):
                info.append(T("drawing_from", name=asset["drawing_from"]))
            info += asset.get("warnings") or []
            st["chem_info"] = "\n".join(info)
            st["chem_warn"] = bool(asset.get("warnings"))
            if g == self.current:
                self.apply_series_widgets()
            if render:
                self.render()

        self.run(work, done, sync)

    def choose_photo(self):
        if not self.current:
            return self.error(T("load_first"))
        p, _ = QFileDialog.getOpenFileName(self, T("dlg_photo"), self.settings.value("last_img_dir", ""),
                                           f"{T('flt_photo')};;{T('flt_all')}")
        if p:
            self.settings.setValue("last_img_dir", str(Path(p).parent))
            self.set_photo(p)

    def bg_changed(self):
        if self._loading or not self.current:
            return
        st = self.series[self.current]
        st["bg"] = self.cmb_bg.currentData()
        if st["photo"]:
            self.set_photo(st["photo"])

    def ensure_model(self):
        """Fragt nach dem Modell-Download. True = ISNet nutzbar."""
        if IMG.model_path(existing_only=True):
            return True
        try:
            import onnxruntime  # noqa: F401
        except ImportError:
            return False
        if os.environ.get("QT_QPA_PLATFORM") == "offscreen" and not getattr(self, "_allow_download", False):
            return False
        if not getattr(self, "_allow_download", False):
            ans = QMessageBox.question(self, T("model_q_title"), T("model_q", mb=IMG.MODEL_SIZE_MB))
            if ans != QMessageBox.Yes:
                return False
        dlg = QProgressDialog(T("downloading"), T("cancel"), 0, 100, self)
        dlg.setWindowModality(Qt.WindowModal)
        dlg.show()

        def prog(done, total):
            dlg.setValue(int(100 * done / max(total, 1)))
            QApplication.processEvents()
            return not dlg.wasCanceled()
        try:
            IMG.download_model(prog)
            return True
        except Exception as e:
            self.log_signal.emit(str(e))
            return False
        finally:
            dlg.close()

    def set_photo(self, path, g=None, render=True, sync=False):
        g = g or self.current
        st = self.series[g]
        if path is not None and self.project_dir:
            path = self.import_file(path, "images")
        if path is None:
            st.update(photo=None, photo_asset=None)
            self.apply_series_widgets()
            return self.render()
        method = st["bg"]
        if method in ("auto", "isnet") and not self.ensure_model():
            method = "grabcut" if method == "isnet" else "auto"
        cache = Path(self.csv_path).parent / ".uvvis_cache" if self.csv_path else None

        def done(asset):
            st["photo"], st["photo_asset"] = str(path), asset
            self.log_signal.emit(f"  {Path(path).name}: {asset.get('bg_method')}")
            if g == self.current:
                self.apply_series_widgets()
            if render:
                self.render()

        self.run(lambda: IMG.load_photo(path, method, cache), done, sync)

    def run(self, fn, done, sync=False):
        if sync:
            try:
                done(fn())
            except Exception as e:
                self.error(str(e))
            return
        QApplication.setOverrideCursor(Qt.WaitCursor)
        self.statusBar().showMessage(T("working"))
        t = Task(fn, self)
        self._tasks.append(t)

        def fin():
            QApplication.restoreOverrideCursor()
            self.statusBar().clearMessage()
            self._tasks.remove(t)
        t.done.connect(done)
        t.failed.connect(self.error)
        t.finished.connect(fin)
        t.start()

    # ------------------------------------------------------------ Plot
    def make_cfg(self, g):
        gl = self.glob
        st = self.series[g]
        cfg = core.deep_merge(core.DEFAULTS, {})
        cfg["_base_dir"] = Path(self.csv_path).parent
        cfg["conc_unit"] = self.units[g]
        cfg["molar_mass_g_mol"] = st["mw"]
        cfg["path_length_cm"] = gl["path_length"]
        cfg["max_abs_fit"] = gl["cutoff"]
        cfg["min_fit_points"] = gl["min_points"]
        cfg["min_r2"] = gl.get("min_r2", 0.98)
        try:
            bands = [float(x) for x in gl["bands"].replace(";", ",").split(",") if x.strip()]
        except ValueError:
            bands = []
        cfg["wavelengths_nm"] = bands or "auto"
        cfg["bands_add_auto"] = gl.get("bands_add", True)
        cfg["find_shoulders"] = gl.get("shoulders", False)
        cfg["inset"]["show_excluded"] = gl.get("show_excl", False)
        cfg["inset"]["show"] = gl.get("inset", True)
        cfg["labels"]["mark_shoulders"] = gl.get("mark_sh", False)
        cfg["labels"]["index"] = gl.get("label_index", False)
        cfg["band_cutoff"] = {int(k): float(v) for k, v in (st.get("band_cutoff") or {}).items()}
        cfg["band_hidden"] = [int(k) for k in (st.get("band_hidden") or [])]
        if gl["baseline"] == "off":
            cfg["baseline_nm"] = None
        elif gl["baseline"] == "manual":
            cfg["baseline_nm"] = [gl["bl_a"], gl["bl_b"]]
        else:
            cfg["baseline_nm"] = "auto"
            cfg["baseline_mode"] = gl["baseline"]
        cfg["ylim"] = [0, gl["ymax"]]
        if not gl["x_auto"] and gl["x_max"] > gl["x_min"]:
            cfg["xlim"] = [gl["x_min"], gl["x_max"]]
        cfg["r2_mode"] = gl["r2"]
        cfg["labels"]["show_error"] = gl["show_err"]
        cfg["inset"]["table"] = gl["table"]
        cfg["figsize_in"] = list(FIGSIZE)
        if gl["size"] != "std":
            core.apply_export_size(cfg, gl["w_cm"], gl["h_cm"])
        fmts = [e for e, on in (("pdf", gl["fmt_pdf"]), ("svg", gl["fmt_svg"]), ("png", gl["fmt_png"])) if on]
        cfg["extra_formats"] = fmts[1:]
        cfg["_main_fmt"] = fmts[0] if fmts else "pdf"
        imgs = []
        if st["structure_asset"]:
            imgs.append({"file": st["structure"], "id": "struktur", "width": st.get("struct_w", 17) / 100,
                         "prefer": "top",
                         "_asset": st["structure_asset"]})
        if st["photo_asset"]:
            imgs.append({"file": st["photo"], "id": "kuevette", "width": st.get("photo_w", 9) / 100,
                         "prefer": "right",
                         "_asset": st["photo_asset"]})
        cfg["images"] = imgs
        core.apply_text_format(cfg, self.text_format("epsilon"))
        return cfg

    # ------------------------------------------------------------ Farbe der Lösung
    def color_settings_changed(self, *_):
        if self._loading:
            return
        self.glob.update(color_show=self.cb_color.isChecked(), color_d=self.sp_color_d.value())
        st = self.series.get(self.current)
        if st is not None and self.sender() is self.sp_color_c:    # nur echte Eingabe merken
            st["color_c"] = self.sp_color_c.value()
        self.update_color()
        self.push_undo()

    def update_color(self):
        """Farbfeld aktualisieren: c (Serieneinheit) und d frei wählbar, Standard = höchste gemessene
        Konzentration in 1 cm."""
        on = bool(self.glob.get("color_show"))
        for w_ in (self.w_color_cd, self.lbl_color, self.lbl_color_info):
            w_.setVisible(on)
        if not on or not self._color_spectra or self.current not in self.series:
            return
        spectra, cfg = self._color_spectra
        if not spectra:
            return
        st = self.series[self.current]
        unit = cfg["conc_unit"].replace("uM", "µM")
        c = st.get("color_c") or max(s_["c"] for s_ in spectra)
        d = self.glob.get("color_d", 1.0)
        loading, self._loading = self._loading, True
        self.sp_color_c.setSuffix(f" {unit}")
        self.sp_color_c.setValue(c)
        self.sp_color_d.setValue(d)
        self._loading = loading
        a, covered = COLOR.solution_absorbance(spectra, c, d, cfg["path_length_cm"], cfg["max_abs_fit"])
        rgb, out_of_gamut = COLOR.absorbance_to_rgb(a)
        hx = COLOR.to_hex(rgb)
        self.lbl_color.setStyleSheet(f"background: {hx}; border: 1px solid #888; border-radius: 4px;")
        info = T("color_info", hex=hx, c=f"{c:g}", unit=unit, d=f"{d:g}")
        if covered < 0.95:
            info += "\n" + T("color_partial", p=round(100 * covered))
        if out_of_gamut:
            info += "\n" + T("color_gamut")
        self.lbl_color_info.setText(info)
        self.lbl_color_info.setStyleSheet(WARN_CSS if covered < 0.95 else HINT_CSS)

    def show_hover(self, txt):
        if txt:
            self.statusBar().showMessage(txt)
        else:
            self.statusBar().clearMessage()

    @staticmethod
    def hover_text(lam, val, name):
        return T("hover_abs", lam=lam, a=val, name=name) if val is not None else T("hover_lam", lam=lam)

    def pick_band(self, lam):
        """Doppelklick aufs Spektrum: λ als Bande hinzufügen bzw. (±2 nm) wieder entfernen."""
        try:
            bands = [float(x) for x in self.glob["bands"].replace(";", ",").split(",") if x.strip()]
        except ValueError:
            bands = []
        lam = round(lam)
        near = [b for b in bands if abs(b - lam) <= 2]
        if near:
            bands = [b for b in bands if b not in near]
            msg = T("band_removed", lam=f"{near[0]:g}")
        else:
            bands.append(lam)
            msg = T("band_added", lam=f"{lam:g}")
        self.ed_bands.setText(", ".join(f"{b:g}" for b in sorted(bands)))
        self.changed()
        QTimer.singleShot(0, lambda: self.statusBar().showMessage(msg, 6000))

    def text_format(self, key):
        """Textformat eines Tabs: eigenes Format (Menü Format, „nur dieser Tab“), sonst das für alle Tabs."""
        return {**core.TEXT_DEFAULTS, **self.text_fmt, **(self.tab_fmt.get(key) or {})}

    def text_tab(self):
        """Offener Tab -> (Schlüssel für das Textformat oder None, Name, Optionen, die der Tab im Bild hat)."""
        w_ = self.main_tabs.currentWidget()
        name = self.main_tabs.tabText(self.main_tabs.currentIndex())
        if w_ is self.eps_page:
            return "epsilon", name, TextFormatDialog.ALL_OPTS
        if isinstance(w_, uvvis_tabs.FigureTab):
            return w_.STATE_KEY, name, w_.TEXT_OPTS
        return None, name, ()

    def edit_text_format(self):
        key, name, opts = self.text_tab()
        dlg = TextFormatDialog(self, self.text_fmt, self.tab_fmt.get(key), opts or TextFormatDialog.ALL_OPTS,
                               key, name)
        if dlg.exec() != QDialog.Accepted:
            return
        self.text_fmt, own = dlg.results()
        if key:
            if own is None:
                self.tab_fmt.pop(key, None)
            else:
                self.tab_fmt[key] = own
        for tf in [self.text_fmt] + list(self.tab_fmt.values()):
            for k_ in ("font", "symbol_font"):
                if core.first_font([tf[k_]], None) is None:
                    self.log_signal.emit("  " + T("tf_missing", name=tf[k_]))
        self.render()
        for t in (self.ov_tab, self.fl_tab, self.cv_tab, self.sec_tab, self.epr_tab, self.irms_tab, self.hrms_tab,
                  self.xray_tab, self.dft_tab):
            t.render()
        self.push_undo()

    def render(self):
        has_data = bool(self.current and self.current in self.groups)
        self.preview_stack.setCurrentIndex(1 if has_data else 0)
        if not has_data:
            return
        g = self.current
        try:
            cfg = self.make_cfg(g)
            core.setup_fonts(plt, cfg)
            cfg2, spectra, results = core.prepare_series(cfg, self.groups[g], g)
            if self.dragger:
                self.dragger.disconnect()
                self.dragger = None
            if self.hover:
                self.hover.disconnect()
                self.hover = None
            self.canvas_host.set_figure_size(*cfg2["figsize_in"])
            self.F = core.build_figure(plt, cfg2, spectra, results, self.get_layout(g), fig=self.fig)
            self.dragger = GuiDragger(self.F, self.layout_changed)
            unit = cfg2["conc_unit"].replace("uM", "µM")
            curves = [dict(x=s_["x"], y=s_["y"], val=s_["y"], name=f"c = {s_['c']:g} {unit}") for s_ in spectra]
            self.hover = uvvis_tabs.HoverReadout(self.F, curves, self.show_hover, self.hover_text, self.pick_band)
            self.sync_image_sizes(self.F.current_layout())
            self.canvas.draw_idle()
            self.fill_table(results, cfg2)
            self.update_si_line(results, cfg2)
            self._color_spectra = (spectra, cfg2)
            self.update_color()
            self._last = (cfg2, results)
            self.push_undo()
        except Exception as e:
            self.log_signal.emit(traceback.format_exc())
            self.statusBar().showMessage(f"{T('err_title')}: {e}", 8000)

    def size_key(self):
        gl = self.glob
        return "std" if gl["size"] == "std" else f"{gl['w_cm']:g}x{gl['h_cm']:g}"

    def get_layout(self, g):
        lay = self.series[g]["layout"] or {}
        if lay and not all(isinstance(v, dict) for v in lay.values()):
            lay = {"std": lay}                       # altes Format (eine Größe)
            self.series[g]["layout"] = lay
        return lay.setdefault(self.size_key(), {})

    def layout_changed(self, layout):
        if not self.current or self.F is None:
            return
        stored = self.get_layout(self.current)
        if layout != stored:
            self.series[self.current]["layout"][self.size_key()] = layout
        self.sync_image_sizes(layout)
        self.push_undo()

    def sync_image_sizes(self, layout):
        """Größenfelder an die tatsächliche Bildgröße anpassen (Mausrad, automatisches Verkleinern)."""
        st = self.series.get(self.current)
        if not st:
            return
        self._loading = True
        for key, field, spin in (("image_struktur", "struct_w", self.sp_struct_w),
                                 ("image_kuevette", "photo_w", self.sp_photo_w)):
            if key in layout and len(layout[key]) == 4:
                st[field] = round(layout[key][2] * 100, 1)
                spin.setValue(st[field])
        self._loading = False

    def image_size_changed(self, iid, value):
        """Größe aus dem Zahlenfeld: Breite setzen, Mittelpunkt und Seitenverhältnis behalten."""
        if self._loading or not self.current:
            return
        st = self.series[self.current]
        st["struct_w" if iid == "struktur" else "photo_w"] = value
        lay = self.get_layout(self.current)
        key = f"image_{iid}"
        if key in lay and len(lay[key]) == 4:
            x, y, w, h = lay[key]
            nw = value / 100
            nh = h * nw / w
            nx = min(max(x + (w - nw) / 2, 0.0), max(0.0, 1 - nw))     # in der Achse halten
            ny = min(max(y + (h - nh) / 2, 0.0), max(0.0, 1 - nh))
            lay[key] = [nx, ny, nw, nh]
        self.timer.start()

    def reset_layout(self):
        if self.current:
            self.get_layout(self.current)
            self.series[self.current]["layout"][self.size_key()] = {}
            self.render()

    def reset_defaults(self):
        """ε-Tab: Einstellungen (Cutoff, Banden, Basislinie, Darstellung …) auf DEFAULT_GLOB, für die gewählte Serie
        Cutoff/Ausblenden je Bande und Layout. Schichtdicke, Exportgröße/-formate, Molmasse, Struktur, Foto,
        Lösungsmittel und Konzentrationen bleiben. Rückgängig holt alles zurück."""
        keep = ("path_length", "size", "w_cm", "h_cm", "fmt_pdf", "fmt_svg", "fmt_png")
        self.glob.update({k: v for k, v in DEFAULT_GLOB.items() if k not in keep})
        st = self.series.get(self.current)
        if st is not None:
            st.update(band_cutoff={}, band_hidden=[], layout={})
        self.apply_globals()
        self.apply_series_widgets()
        self.render()
        self.push_undo()
        self.statusBar().showMessage(T("xr_reset_display_done"), 6000)

    def update_si_line(self, results, cfg):
        st = self.series.get(self.current) or {}
        self.ed_solvent.setText(st.get("solvent", ""))
        self._uv_line = SI.uvvis_line(results, cfg, st.get("solvent", ""))
        self.txt_si.setHtml(self._uv_line[1] or T("si_uv_none"))

    def solvent_changed(self):
        st = self.series.get(self.current)
        if st is not None and self.ed_solvent.text().strip() != st.get("solvent", ""):
            st["solvent"] = self.ed_solvent.text().strip()
            self.render()

    def uvvis_si_parts(self):
        """UV-Vis-Zeile je Serie mit Ergebnissen (für den SI-Export)."""
        out = []
        for g in self.groups:
            try:
                cfg2, _, results = core.prepare_series(self.make_cfg(g), self.groups[g], g)
            except Exception:
                continue
            line = SI.uvvis_line(results, cfg2, (self.series.get(g) or {}).get("solvent", ""))
            if line[0]:
                out.append(line)
        return out

    def eps_methods_text(self, groups=None):
        """SI-Text zur ε-Bestimmung je Serie (englisch, uvvis_si.eps_methods)."""
        out = []
        for g in (groups or list(self.groups)):
            try:
                cfg2, spectra, results = core.prepare_series(self.make_cfg(g), self.groups[g], g)
            except Exception:
                continue
            txt = SI.eps_methods(cfg2, results, spectra, (self.series.get(g) or {}).get("solvent", ""))
            if txt:
                out.append((f"Series {g}: " if len(self.groups) > 1 and g else "") + txt)
        return "\n\n".join(out)

    def copy_eps_methods(self):
        self.collect_globals()
        txt = self.eps_methods_text()
        if txt:
            QApplication.clipboard().setText(txt)
        self.statusBar().showMessage(T("si_methods_copied") if txt else T("si_methods_none"), 6000)

    def si_available(self):
        """Teile der SI, für die Daten da sind."""
        xt, dt = self.xray_tab, self.dft_tab
        return dict(char=True, nmr=bool(self.nmr_tab.spectra()),
                    figs=bool(self.groups) or any(t.F is not None for t in self.figure_tabs()),
                    methods=bool(self.si_methods_blocks()),
                    xray=any(xt.structure(e) is not None for e in xt.entries()),
                    dft=any(dt.calc_of(e) is not None for e in dt.entries()))

    def export_si_dialog(self):
        """Auswahl der Teile (Charakterisierung, X-Ray, DFT), dann export_si. Die Auswahl wird gemerkt."""
        if not self.ensure_project():
            return
        avail = self.si_available()
        saved = self.settings.value("si_parts", "char,nmr,figs,methods,xray,dft")
        saved = set(str(saved).split(",")) if saved else set()
        if not self.settings.value("si_parts_figs"):         # neu in 1.24: Abbildungen einmalig mit auswählen
            saved |= {"nmr", "figs"}
            self.settings.setValue("si_parts_figs", 1)
        if not self.settings.value("si_parts_methods"):      # neu in 1.24.1: Methodentexte einmalig mit auswählen
            saved |= {"methods"}
            self.settings.setValue("si_parts_methods", 1)
        dlg = QDialog(self)
        dlg.setWindowTitle(T("act_export_si").rstrip(" …"))
        v = QVBoxLayout(dlg)
        v.addWidget(QLabel(T("si_parts_head")))
        cbs = {}
        for k in SI_PARTS:
            cb = QCheckBox(T("si_part_" + k))
            cb.setEnabled(avail[k])
            cb.setChecked(avail[k] and k in saved)
            if not avail[k]:
                cb.setToolTip(T("si_part_empty"))
            cbs[k] = cb
            v.addWidget(cb)
        row = QHBoxLayout()                                   # Schriftart aller Word-Exporte (Größe fest 11 pt)
        cmb_font = QFontComboBox()
        cmb_font.setCurrentFont(QFont(str(self.settings.value("si_font", "Arial") or "Arial")))
        cmb_font.setToolTip(T("si_font_tip"))
        row.addWidget(QLabel(T("si_font")))
        row.addWidget(cmb_font, 1)
        v.addLayout(row)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(dlg.accept)
        bb.rejected.connect(dlg.reject)
        v.addWidget(bb)
        if dlg.exec() != QDialog.Accepted:
            return
        parts = [k for k, cb in cbs.items() if cb.isChecked()]
        self.settings.setValue("si_font", cmb_font.currentFont().family())
        SI.set_font(cmb_font.currentFont().family())
        self.settings.setValue("si_parts", ",".join(k for k, cb in cbs.items() if cb.isChecked() or not avail[k]
                                                   and k in saved))
        if parts:
            self.export_si(parts)

    def export_si(self, parts=SI_PARTS):
        """SI-Abschnitt der Verbindung als Word-Datei: Charakterisierungsabsatz (m.p.; IR; HR-MS; UV-Vis;
        Elementaranalyse), Kristalldaten, ausgewählte Parameter und Bildunterschriften aus dem X-Ray-Tab, DFT-Rechnungen
        (zweispaltig) – je nach Auswahl."""
        if not self.ensure_project():
            return
        try:
            blocks, missing = self.si_blocks(parts)
            pdir = Path(self.project_dir)
            name = core.safe_name(self.compound.get("name") or pdir.name)
            p = SI.write_docx(PRJ.export_dir(pdir, "si") / f"{name}_SI.docx", blocks)
        except PermissionError as e:
            return self.error(T("report_locked", p=getattr(e, "filename", "")))
        except Exception as e:
            self.log_signal.emit(traceback.format_exc())
            return self.error(str(e))
        msg = T("si_written", p=str(p))
        if missing:
            msg += "\n" + T("si_missing", parts=", ".join(missing))
        self.info(msg)

    def si_blocks(self, parts=SI_PARTS):
        """-> (Blöcke für SI.write_docx, fehlende Teile)."""
        comp = self.compound or {}
        title = self.syn_tab.st.get("name") or comp.get("name") or Path(self.project_dir or ".").name
        blocks, missing, plain = [], [], ""
        if "char" in parts:
            blocks, missing, plain = self.si_char_blocks(comp, title)
        elif title:
            blocks = [("heading", title)]
        self._si_plain = plain
        if "nmr" in parts and self.nmr_tab.spectra() and self.project_dir:
            blocks += [("pagebreak",)] + self.nmr_tab.docx_blocks(PRJ.export_dir(Path(self.project_dir), "si"))
        if "figs" in parts and self.project_dir:
            figs = self.si_figure_blocks(PRJ.export_dir(Path(self.project_dir), "si"), title)
            if figs:
                blocks += [("pagebreak",)] + figs
        if "methods" in parts:
            mb = self.si_methods_blocks()
            if mb:
                blocks += [("gap",)] + mb
        xt = self.xray_tab
        es = [e for e in xt.entries() if xt.structure(e) is not None] if "xray" in parts else []
        if es:
            blocks += self.si_xray_blocks(es)
        dt = self.dft_tab
        if "dft" in parts and any(dt.calc_of(e) is not None for e in dt.entries()) and self.project_dir:
            blocks += [("gap",), ("columns", 2)] + dt.docx_blocks(PRJ.export_dir(Path(self.project_dir), "si"))
        return blocks, missing

    def si_methods_blocks(self):
        """Abschnitt „Data processing and fitting“: dieselben Methodentexte wie „SI-Text kopieren“ (nur benutzte
        Schritte, aktuelle Parameter, dann Ergebnisse) für alle Tabs, die fitten; je Technik ein Absatz mit fettem
        Titel. Leer, wenn nirgends gefittet wurde."""
        self.collect_globals()
        items = [("UV-Vis absorption spectroscopy", self.eps_methods_text() if self.groups else "")]
        for title, t in (("Cyclic voltammetry", self.cv_tab), ("UV-Vis spectroelectrochemistry", self.sec_tab),
                         ("EPR spectroscopy", self.epr_tab)):
            try:
                txt = t.methods_text() if t.figure_spec() is not None else ""   # Ergebnisse/Fits aktuell
            except Exception:
                self.log_signal.emit(traceback.format_exc())
                txt = ""
            items.append((title, txt))
        blocks = []
        for title, txt in items:
            paras = [p.strip() for p in (txt or "").split("\n\n") if p.strip()]
            for k, p in enumerate(paras):
                blocks.append(("char", (f"<b>{title}.</b> " if k == 0 else "") + SI.methods_html(p)))
        return [("heading", "Data processing and fitting")] + blocks if blocks else []

    def figure_tabs(self):
        """Tabs mit Abbildung in der Reihenfolge der SI (Charakterisierung, dann Spektroskopie, Elektrochemie)."""
        return [self.irms_tab, self.hrms_tab, self.ov_tab, self.fl_tab, self.cv_tab, self.sec_tab, self.epr_tab]

    def si_figure_blocks(self, out_dir, name):
        """Abbildungen aller Tabs mit Daten (ε je Serie, IR, HR-MS, Overlay, Fluoreszenz, CV, SEC, EPR), immer zwei je
        Seite (SI.two_per_page verkleinert, wo nötig), jede mit „Figure Sn: …“."""
        figs = self.si_eps_figures(out_dir, name)
        for t in self.figure_tabs():
            try:
                fb = t.si_figure(out_dir, name)
                if fb and t.si_setup_text():                 # Aufbau-Absatz der Methode (lokal hinterlegt)
                    fb = [("char", SI.chem_html(t.si_setup_text()))] + fb
                figs += fb
            except Exception:
                self.log_signal.emit(traceback.format_exc())
        return SI.two_per_page(figs)

    def si_eps_figures(self, out_dir, name):
        """ε-Abbildung je Serie wie beim Export (Größe aus dem Tab) + Unterschrift."""
        if not self.groups:
            return []
        self.collect_globals()
        out = []
        for g in self.groups:
            cfg = self.make_cfg(g)
            core.setup_fonts(plt, cfg)
            cfg2, spectra, results = core.prepare_series(cfg, self.groups[g], g)
            fig = MplFigure(figsize=cfg2["figsize_in"], dpi=PREVIEW_DPI)
            FigureCanvasAgg(fig)
            F = core.build_figure(plt, cfg2, spectra, results, self.get_layout(g), fig=fig)
            sfx = f"_{core.safe_name(g)}" if len(self.groups) > 1 else ""
            img = uvvis_tabs.si_image(F, Path(out_dir) / f"{core.safe_name(name)}_epsilon{sfx}.svg")
            solv = (self.series.get(g) or {}).get("solvent", "").strip()
            cap = (f"UV-Vis spectra of <b>{html.escape(name)}</b>"
                   + (f" in {SI.chem_html(solv)}" if solv else "")
                   + f", measured in {cfg2['path_length_cm']:g} cm quartz cuvettes.")
            if not solv:
                self.log_signal.emit(T("si_uv_no_solvent", g=g))
            out += [img, ("figcap", cap)]
        return out

    def si_char_blocks(self, comp, title):
        """„Synthesis of Name“ + Charakterisierungsabsatz -> (Blöcke, fehlende Teile, Klartext)."""
        parts, missing = [], []
        if comp.get("mp"):
            mp = re.sub(r"(?i)^m\.?\s*p\.?\s*", "", comp["mp"]).strip()
            parts.append(("m.p. " + mp, "<b>m.p.</b> " + html.escape(mp)))       # wie in der SI: „m.p.“ fett
        else:
            missing.append("m.p.")
        nmr = self.nmr_tab.si_parts()                    # wie in der SI: m.p.; ¹H NMR …; ¹³C NMR …; IR …
        if not nmr:
            missing.append("NMR")
        parts += nmr
        ir_ms_ea = self.irms_tab.si_parts() + self.hrms_tab.si_parts()
        kinds = [p_[0].split(" ")[0] for p_ in ir_ms_ea]
        uv = self.uvvis_si_parts()
        for key, lab in (("IR", "IR"), ("HR-MS-", "HR-MS"), ("Elemental", "EA")):
            if not any(k_.startswith(key) for k_ in kinds):
                missing.append(lab)
        if not uv:
            missing.append("UV-Vis")
        order = [p_ for p_ in ir_ms_ea if p_[0].startswith("IR")] + \
            [p_ for p_ in ir_ms_ea if p_[0].startswith("HR-MS")] + uv + \
            [p_ for p_ in ir_ms_ea if p_[0].startswith("Elemental")]
        plain, para = SI.join_lines(parts + order)
        blocks = [("synthesis", title)]
        if self.syn_tab.has_content() and self.project_dir:      # Überschrift, Schema, Vorschrift aus dem Tab
            blocks = self.syn_tab.si_blocks(PRJ.export_dir(Path(self.project_dir), "si"))
        if para:
            blocks.append(("char", para))
        return blocks, missing, plain

    def si_xray_blocks(self, es):
        """Abbildungen, Unterschriften, ausgewählte Parameter und Kristalldaten aus dem X-Ray-Tab."""
        xt, blocks = self.xray_tab, []
        if es:
            from matplotlib.backends.backend_agg import FigureCanvasAgg
            from matplotlib.figure import Figure as MplFigure
            out_dir = PRJ.export_dir(Path(self.project_dir), "si") if self.project_dir else None
            blocks.append(("gap",))
            for e in es:
                fname = e.get("name") or e["block"]
                if out_dir is not None:                   # Abbildung wie im X-Ray-Tab, 300 dpi
                    fig = MplFigure(figsize=(8, 6), dpi=uvvis_tabs.PREVIEW_DPI)
                    FigureCanvasAgg(fig)
                    F = xt.build(fig, e, interactive=False)
                    if F is not None:                     # wie in der SI eng um das Molekül beschnitten
                        svg = out_dir / f"{core.safe_name(fname)}_xray.svg"
                        blocks.append(("image",) + XRAY.save_tight(F, svg))
                blocks.append(("figcap", xt.caption(e)[1]))
                for t_title, t_name, t_head, t_rows in xt.param_tables(e):
                    blocks += [("tabcap", t_title, t_name), ("params", t_head, t_rows)]
                    if any("^{a}" in r_[-1] for r_ in t_rows):
                        blocks.append(("note", "^{a} " + XRAY.CALC_NOTE))
                    blocks.append(("gap",))
            h_, r_ = xt.crystal_table()
            r_ = [[c_.replace(", ", "\n") if row_[0] == "Index ranges" else c_ for c_ in row_] for row_ in r_]
            names = [e.get("name") or e["block"] for e in es]
            blocks += [("tabcap", "Crystal data and structure refinement", ", ".join(names)), ("crystal", h_, r_)]
        return blocks

    def fill_table(self, results, cfg):
        sym = "ε [M⁻¹ cm⁻¹]" if core.conc_factor(cfg) else "a [L g⁻¹ cm⁻¹]"
        self.table.blockSignals(True)
        self.table.setHorizontalHeaderLabels([T("col_lam"), T("col_n"), sym, T("col_intercept"),
                                              T("col_r2"), T("col_cut"), T("col_warn")])
        self.table.setRowCount(len(results))
        self._row_keys = []
        for i, r in enumerate(results):
            ft = r["fit"]
            lam = f"{r['lam']:g}" + (" (sh)" if r.get("shoulder") else "")
            cut = r.get("cutoff", cfg["max_abs_fit"])
            if ft:
                v, e = core.fmt_ve(r["eps"], r["eps_err"], max_decimals=core.eps_max_decimals(cfg))
                b, be = core.fmt_ve(ft["intercept"], ft["se_intercept"])
                cells = [lam, str(ft["n"]), f"{v} ± {e}" if e else v,
                         f"{b} ± {be}" if be else b, f"{ft['r2']:.5f}", f"{cut:g}", "; ".join(r["warnings"])]
            else:
                cells = [lam, str(len(r["used"])), "–", "–", "–", f"{cut:g}",
                         T("no_fit", n=len(r["used"]), cut=cut, min=cfg["min_fit_points"])]
            for j, c in enumerate(cells):
                it = QTableWidgetItem(c)
                if j < 6:
                    it.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                if j == 0:                                   # Häkchen = im Bild anzeigen
                    it.setFlags((it.flags() | Qt.ItemIsUserCheckable) & ~Qt.ItemIsEditable)
                    hidden = int(round(r["lam_requested"])) in set(cfg.get("band_hidden") or [])
                    it.setCheckState(Qt.Unchecked if hidden else Qt.Checked)
                elif j != 5:
                    it.setFlags(it.flags() & ~Qt.ItemIsEditable)
                    if j == 4 and ft and ft["r2"] < cfg["min_r2"]:   # unter R²-Schwelle -> nicht im Bild
                        it.setBackground(QColor("#f4cccc"))
                        it.setForeground(QColor("#000000"))
                        it.setToolTip(T("r2_below", min=cfg["min_r2"]))
                elif abs(cut - cfg["max_abs_fit"]) > 1e-9:
                    it.setBackground(QColor("#fff2cc"))
                    it.setForeground(QColor("#000000"))
                self.table.setItem(i, j, it)
            self._row_keys.append(int(round(r["lam_requested"])))
        self.table.resizeColumnsToContents()
        self.table.blockSignals(False)

    def cutoff_edited(self, item):
        """Cutoff pro Bande aus der Ergebnistabelle (leer = globaler Cutoff); Häkchen in Spalte λ
        = Bande im Bild anzeigen."""
        if not self.current or item.row() >= len(self._row_keys):
            return
        if item.column() == 0:
            st = self.series[self.current]
            hid = set(st.get("band_hidden") or [])
            k = self._row_keys[item.row()]
            hid.discard(k) if item.checkState() == Qt.Checked else hid.add(k)
            st["band_hidden"] = sorted(hid)
            QTimer.singleShot(0, self.render)
            return
        if item.column() != 5:
            return
        key = str(self._row_keys[item.row()])
        st = self.series[self.current]
        bc = st.setdefault("band_cutoff", {})
        try:
            v = parse_float(item.text())
        except ValueError:
            v = None
        if v is None or abs(v - self.glob["cutoff"]) < 1e-9:
            bc.pop(key, None)
        else:
            bc[key] = v
        QTimer.singleShot(0, self.render)

    # ------------------------------------------------------------ Export
    def export_dialog(self):
        if not self.csv_path:
            return self.error(T("load_first"))
        if not self.ensure_project():
            return
        out = self.export_all()
        self.exported(p.parent for p in out)

    def export_all(self, folder: Path = None):
        """Alle Serien exportieren: Abbildungen, Ergebnis-CSV und Excel-Report.
        Ohne Ordnerangabe nach <Projekt>/exports/epsilon/<Serie>/ (wird überschrieben)."""
        self.collect_globals()
        stem = Path(self.csv_path).stem
        written = []
        for g in self.groups:
            if folder is None:
                gdir = PRJ.export_dir(self.project_dir, "epsilon", core.safe_name(g))
                name = core.safe_name(g)
            else:
                gdir = Path(folder)
                name = f"{stem}_{g}" if len(self.groups) > 1 else stem
            cfg = self.make_cfg(g)
            core.setup_fonts(plt, cfg)
            cfg2, spectra, results = core.prepare_series(cfg, self.groups[g], g)
            folder_g = gdir
            cfg2["output"] = str(folder_g / f"{name}.{cfg2['_main_fmt']}")
            fig = MplFigure(figsize=cfg2["figsize_in"], dpi=PREVIEW_DPI)
            FigureCanvasAgg(fig)
            F = core.build_figure(plt, cfg2, spectra, results, self.get_layout(g), fig=fig)
            F.export()
            core.write_results(results, folder_g / f"{name}_results.csv", cfg2)
            written += [folder_g / f"{name}.{e}" for e in [cfg2["_main_fmt"]] + cfg2["extra_formats"]]
            written.append(folder_g / f"{name}_results.csv")
            meth = SI.eps_methods(cfg2, results, spectra, (self.series.get(g) or {}).get("solvent", ""))
            if meth:
                written.append(folder_g / f"{name}_methods.txt")
                written[-1].write_text(meth + "\n", encoding="utf-8")
            try:
                rep = REPORT.write_report(folder_g / f"{name}_report.xlsx", cfg2, self.groups[g], spectra,
                                          results, g, files=self.csv_paths, version=APP_VERSION)
                written.append(Path(rep))
            except PermissionError:
                self.error(T("report_locked", p=folder_g / f"{name}_report.xlsx"))
        return written

    # ------------------------------------------------------------ Drag & Drop
    def dragEnterEvent(self, ev):
        if ev.mimeData().hasUrls():
            ev.acceptProposedAction()

    def dropEvent(self, ev):
        files = [u.toLocalFile() for u in ev.mimeData().urls()]
        for p in files:
            if PRJ.is_project(p):
                self.open_project(p)
                return
        # X-Ray: CIF-Dateien (jeder data_-Block mit Atomen wird ein Eintrag)
        cifs = [p for p in files if XRAY.is_cif(p)]
        if cifs and self.ensure_project():
            for p in cifs:
                self.xray_tab.load_cif(p)
            self.main_tabs.setCurrentWidget(self.xray_tab)
            files = [p for p in files if p not in cifs]
            if not files:
                return
        # NMR: SVG-Exporte aus MestReNova; Bruker-Experimentordner gehören zum gewählten Spektrum
        svgs = [p for p in files if NMR.is_nmr_svg(p)
                or (Path(p).suffix.lower() == ".svg" and self.main_tabs.currentWidget() is self.nmr_tab)]
        exps = [p for p in files if self.main_tabs.currentWidget() is self.nmr_tab and NMR.bruker_dir(p) is not None]
        if (svgs or exps) and self.ensure_project():
            for p in svgs:
                self.nmr_tab.load_svg(p)
            for p in exps:
                if self.nmr_tab.current() is not None:
                    self.nmr_tab.set_bruker(self.nmr_tab.current(), p)
            self.main_tabs.setCurrentWidget(self.nmr_tab)
            files = [p for p in files if p not in svgs and p not in exps]
            if not files:
                return
        # DFT: xyz-Geometrien und Ausgaben von ORCA/Gaussian
        xyzs = [p for p in files if DFT.is_xyz(p) or DFT.is_calc(p)]
        if xyzs and self.ensure_project():
            for p in xyzs:
                self.dft_tab.load_xyz(p)
            self.main_tabs.setCurrentWidget(self.dft_tab)
            files = [p for p in files if p not in xyzs]
            if not files:
                return
        # HRMS: MS-Berichte (docx) und Bruker-.d-Ordner; IR + EA: OPUS-Dateien, PDF nur auf dem Tab IR + EA
        ms = [p for p in files if IRMS.is_bruker_d(p) or Path(p).suffix.lower() == ".docx"]
        irms = [p for p in files if IRMS.is_opus(p)
                or (Path(p).suffix.lower() == ".pdf" and self.main_tabs.currentWidget() is self.irms_tab)]
        if (ms or irms) and self.ensure_project():
            for p in ms:
                (self.hrms_tab.load_raw if IRMS.is_bruker_d(p) else self.hrms_tab.load_ms)(p)
            for p in irms:
                self.irms_tab.load_ir(p)
            self.main_tabs.setCurrentWidget(self.hrms_tab if ms else self.irms_tab)
            files = [p for p in files if p not in ms and p not in irms]
            if not files:
                return
        data = [p for p in files if Path(p).suffix.lower() in {".csv", ".txt"} | core.BINARY_EXT]
        if data:
            self.open_files(data)                     # mehrere Messdateien = ein Datensatz
        for p in files:
            ext = Path(p).suffix.lower()
            if p in data:
                continue
            if not self.current:
                self.error(T("load_first"))
            elif ext in IMG.CHEM_EXT | IMG.VECTOR_EXT:
                self.set_structure(p)
            elif ext in IMG.RASTER_EXT:
                self.set_photo(p)

    # ------------------------------------------------------------ Diverses
    def error(self, msg):
        self.log_signal.emit(f"{T('err_title')}: {msg}")
        if os.environ.get("QT_QPA_PLATFORM") != "offscreen":
            QMessageBox.critical(self, T("err_title"), str(msg))

    def rebuild_tabs(self):
        self.build_ui()
        self.render()

    def export_current(self):
        cur = self.main_tabs.currentWidget()
        for tab in (self.ov_tab, self.fl_tab, self.cv_tab, self.sec_tab, self.epr_tab, self.hrms_tab, self.irms_tab,
                    self.xray_tab, self.dft_tab, self.nmr_tab, self.syn_tab):
            if cur is tab:
                return tab.export_dialog()
        return self.export_dialog()

    def info(self, msg):
        self.log_signal.emit(msg)
        if os.environ.get("QT_QPA_PLATFORM") != "offscreen":
            QMessageBox.information(self, T("app_title"), msg)

    def exported(self, folders):
        """Meldung nach einem Export mit Knopf „Ordner öffnen“ (erster Zielordner im Finder/Explorer)."""
        folders = sorted({str(f) for f in folders})
        msg = T("exported_to", p="\n".join(folders))
        self.log_signal.emit(msg)
        self.last_export = folders
        if os.environ.get("QT_QPA_PLATFORM") == "offscreen" or not folders:
            return
        box = QMessageBox(QMessageBox.Information, T("app_title"), msg, QMessageBox.Ok, self)
        b_open = box.addButton(T("open_folder"), QMessageBox.ActionRole)
        box.exec()
        if box.clickedButton() is b_open:
            QDesktopServices.openUrl(QUrl.fromLocalFile(folders[0]))

    # ------------------------------------------------------------ Updates
    def _auto_updates(self):
        return str(self.settings.value("update_auto", True)).lower() not in ("false", "0")

    def check_updates(self, manual=False):
        """Neuestes Release auf GitHub abfragen. Automatisch höchstens einmal pro Tag und still bei
        Fehlern (z. B. offline); über das Hilfe-Menü mit Rückmeldung."""
        import time
        if not manual:
            if not self._auto_updates() or time.time() - float(self.settings.value("update_last", 0) or 0) < 86400:
                return
        t = Task(lambda: UPD.fetch_latest(APP_VERSION), self)
        self._tasks.append(t)

        def done(info):
            self.settings.setValue("update_last", time.time())
            self.show_update(info, manual)

        def failed(msg):
            if manual:
                QMessageBox.warning(self, T("upd_title"), T("upd_failed", err=msg))
        t.done.connect(done)
        t.failed.connect(failed)
        t.finished.connect(lambda: self._tasks.remove(t))
        t.start()

    def update_text(self, info, manual):
        """Meldungstext oder None (nichts anzeigen). Übersprungene Versionen nur bei manueller Prüfung."""
        if not UPD.is_newer(info["version"], APP_VERSION):
            return T("upd_current", cur=APP_VERSION) if manual else None
        if not manual and self.settings.value("update_skip", "") == info["version"]:
            return None
        return T("upd_available", new=info["version"], cur=APP_VERSION,
                 file=info["asset"] or T("upd_no_asset"))

    def show_update(self, info, manual):
        txt = self.update_text(info, manual)
        if txt is None:
            return
        if not UPD.is_newer(info["version"], APP_VERSION):
            QMessageBox.information(self, T("upd_title"), txt)
            return
        box = QMessageBox(QMessageBox.Information, T("upd_title"), txt, parent=self)
        b_open = box.addButton(T("upd_open"), QMessageBox.AcceptRole)
        b_skip = box.addButton(T("upd_skip"), QMessageBox.DestructiveRole)
        box.addButton(T("upd_later"), QMessageBox.RejectRole)
        box.exec()
        if box.clickedButton() is b_open:
            QDesktopServices.openUrl(QUrl(info["url"]))
        elif box.clickedButton() is b_skip:
            self.settings.setValue("update_skip", info["version"])

    def statusBar(self):
        return self._status

    def _place_toast(self, msg=None):
        """Meldung unten rechts über der Abbildung (nimmt keinen Platz im Layout ein)."""
        t = self._toast
        msg = self._status.currentMessage() if msg is None else msg
        if not msg:
            t.hide()
            return
        t.setText(msg)
        t.ensurePolished()
        w = min(t.fontMetrics().horizontalAdvance(msg) + 32, self.width() - 40)
        h = t.fontMetrics().height() + 14
        t.setGeometry(self.width() - w - 24, self.height() - h - 24, w, h)
        t.show()
        t.raise_()

    def resizeEvent(self, ev):
        super().resizeEvent(ev)
        self._place_toast()

    def closeEvent(self, ev):
        if not self.maybe_save():
            ev.ignore()
            return
        self.settings.setValue("geometry", self.saveGeometry())
        super().closeEvent(ev)


# ----------------------------------------------------------------------------
# Selbsttest (für CI und nach der Installation)
# ----------------------------------------------------------------------------
def make_selftest_inputs(d: Path):
    """Synthetische Cary-CSV (2 Serien), CDXML und Küvettenfoto (HEIC, falls möglich)."""
    from PIL import Image, ImageDraw
    d.mkdir(parents=True, exist_ok=True)
    wl = np.arange(1100.0, 199.0, -1.0)
    bands = {"TEST-A": [(262, 6.8, 20), (484, 1.6, 30), (732, 0.25, 40), (816, 0.24, 40)],
             "TEST-B": [(255, 7.3, 18), (351, 2.7, 20), (444, 3.5, 12), (719, 0.42, 35), (804, 0.41, 35)]}
    concs = [("1", 1.0), ("0p5", 0.5), ("0p25", 0.25), ("0p125", 0.125), ("0p0625", 0.0625)]
    rng = np.random.default_rng(0)
    cols, names = [], []
    for ser, bl in bands.items():
        a = sum(h * np.exp(-0.5 * ((wl - l) / w) ** 2) for l, h, w in bl)
        for tag, c in concs:
            names.append(f"{ser}-{tag}mgml_THF")
            y = np.minimum(a * c, 4.0) - 0.03 + rng.normal(0, 0.002, wl.size)
            cols.append(y)
    lines = ["".join(f"{n},," for n in names), "".join("Wavelength (nm),Abs," for _ in names)]
    for i, x in enumerate(wl):
        lines.append("".join(f"{x:.6f},{y[i]:.8f}," for y in cols))
    lines += ["", "TEST-A-1mgml_THF", "Instrument  Cary 60"]
    (d / "selftest.csv").write_text("\r\n".join(lines), encoding="utf-8", newline="")
    nodes = [(1, 100, 100, 6), (2, 112.5, 92.8, 6), (3, 125, 100, 6), (4, 125, 114.4, 6),
             (5, 112.5, 121.6, 6), (6, 100, 114.4, 7), (7, 112.5, 78.4, 8)]
    bonds = [(1, 2, 2), (2, 3, 1), (3, 4, 2), (4, 5, 1), (5, 6, 2), (6, 1, 1), (2, 7, 1)]
    xml = ['<?xml version="1.0" encoding="UTF-8" ?><CDXML BondLength="14.40"><page id="1"><fragment id="99">']
    xml += [f'<n id="{i}" p="{x} {y}"' + (f' Element="{e}"' if e != 6 else "") + "/>" for i, x, y, e in nodes]
    xml += [f'<b id="{20 + k}" B="{b}" E="{e}"' + (' Order="2"' if o == 2 else "") + "/>"
            for k, (b, e, o) in enumerate(bonds)]
    xml += ["</fragment></page></CDXML>"]
    (d / "pyridinol.cdxml").write_text("".join(xml), encoding="utf-8")
    im = Image.new("RGB", (600, 900), (190, 185, 175))
    dr = ImageDraw.Draw(im)
    dr.rectangle([240, 100, 360, 200], fill=(25, 25, 25))
    dr.rectangle([220, 200, 380, 820], fill=(215, 220, 228))
    dr.rectangle([235, 420, 365, 805], fill=(195, 90, 35))
    photo = d / "kuevette.png"
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
        im.save(d / "kuevette.heic", quality=90)
        photo = d / "kuevette.heic"
    except Exception:
        im.save(photo)
    return d / "selftest.csv", d / "pyridinol.cdxml", photo


def selftest(outdir, with_isnet=False):
    global SETTINGS_FILE
    out = Path(outdir).resolve()
    out.mkdir(parents=True, exist_ok=True)
    SETTINGS_FILE = out / "selftest_settings.ini"     # zuletzt geöffnet, Sprache usw. nicht beim Nutzer ändern
    SETTINGS_FILE.unlink(missing_ok=True)
    report = [f"{APP_NAME} {APP_VERSION} selftest", f"python {sys.version.split()[0]}, {sys.platform}"]
    ok = True
    app = QApplication.instance() or QApplication(sys.argv[:1])  # noqa: F841
    QApplication.setApplicationVersion(APP_VERSION)
    try:
        csv, cdxml, photo = make_selftest_inputs(out / "input")
        report.append(f"photo format: {photo.suffix}")
        w = MainWindow()
        w._allow_download = with_isnet
        assert w.center.currentIndex() == 0, "ohne Projekt: Startbildschirm"
        proj = out / "projekt_test"
        assert not w.open_csv(csv), "ohne Projekt darf nichts geladen werden"
        assert w.create_project(proj)
        QApplication.processEvents()
        w.open_csv(csv)
        assert set(w.groups) == {"TEST-A", "TEST-B"}, w.groups
        assert Path(w.csv_path).parent == proj / "data", w.csv_path          # ins Projekt kopiert
        w.set_structure(str(cdxml), "TEST-A", render=False, sync=True)
        st = w.series["TEST-A"]
        report.append(f"CDXML: {st['chem_info']!r}")
        assert st["mw"] and abs(st["mw"] - 95.10) < 0.05, st["mw"]      # 3-Hydroxypyridin C5H5NO
        hint_ = lambda: T("cdxml_drawn", name="pyridinol")                # noqa: E731  Hinweis folgt der Sprache
        assert hint_() in st["chem_info"], st["chem_info"]
        for lang in ("en", "de"):
            w.switch_lang(lang)
            assert hint_() in w.series["TEST-A"]["chem_info"], (lang, w.series["TEST-A"]["chem_info"])
        report.append("structure hints follow language switch: OK")
        w.series["TEST-A"]["bg"] = "isnet" if with_isnet else "grabcut"
        w.set_photo(str(photo), "TEST-A", render=False, sync=True)
        bgm = st["photo_asset"]["bg_method"]
        report.append(f"background: {bgm}")
        if with_isnet:
            assert bgm == "isnet", bgm
        w.render()
        files = w.export_all(out)
        for f in files:
            assert f.exists() and f.stat().st_size > 200, f
        import pymupdf
        page = pymupdf.open(str(out / "selftest_TEST-A.pdf"))[0]
        report.append(f"pdf: {len(page.get_drawings())} vector paths, {len(page.get_images())} images")
        assert page.get_images(), "photo missing in PDF"
        svg = (out / "selftest_TEST-A.svg").read_text(encoding="utf-8")
        assert "structure_0" in svg and "<image" in svg, "SVG ohne Struktur/Foto"
        report.append(f"pdf size: {page.rect.width / 72 * 2.54:.2f} x {page.rect.height / 72 * 2.54:.2f} cm")
        # Probennamen ohne Konzentration (z. B. c0..c4) -> Overrides wie aus dem Dialog
        txt = csv.read_text(encoding="utf-8").splitlines()
        n_samples = len([n for n in txt[0].split(",") if n])
        txt[0] = "".join(f"c{i},," for i in range(n_samples))
        plain = out / "input" / "unnamed.csv"
        plain.write_text("\r\n".join(txt), encoding="utf-8", newline="")
        assert w.open_csv(plain, ask=False) is False                 # ohne Angaben: nichts
        w.conc = {f"c{i}": {"conc": 1.0 / 2 ** i, "unit": "mg/mL", "series": "unbenannt"}
                  for i in range(5)}
        assert w.rebuild_groups() and "unbenannt" in w.groups, w.groups
        assert len(w.groups["unbenannt"]) == 5
        w.render()
        report.append("unnamed samples: OK")
        # Einzelmessungen in getrennten Dateien + Cary-Binärdatei (.DSW)
        import struct
        singles = []
        for i, s_ in enumerate(core.read_samples([csv])[:5]):
            f = out / "input" / f"single_{i}.csv"
            f.write_text("\r\n".join([f"{s_['name']},,", "Wavelength (nm),Abs,"] +
                                      [f"{a:.4f},{b:.6f}," for a, b in zip(s_["x"], s_["y"])]),
                         encoding="utf-8", newline="")
            singles.append(f)
        assert w.open_files(singles, ask=False) and len(w.groups["TEST-A"]) == 5, w.groups
        magic = b"Varian UV-VIS Spectrophotometer"
        blob = bytearray([len(magic)]) + magic
        blob += bytes(0x400 - len(blob))
        for k in range(3):
            blk = bytearray(256)
            t = f"DSW-T-{0.5 / 2 ** k:g}".replace(".", "p").encode() + b"mgml"
            blk[:len(t)] = t
            blob += blk
            for x in np.arange(800.0, 249.0, -1.0):
                blob += struct.pack("<ff", x, 0.4 / 2 ** k * np.exp(-((x - 450) / 40) ** 2))
            blob += bytes(64)
        dsw = out / "input" / "test.DSW"
        dsw.write_bytes(bytes(blob))
        assert w.open_files([dsw], ask=False) and len(w.groups.get("DSW-T", [])) == 3, w.groups
        report.append("multiple files + DSW: OK")
        # Overlay (Messung + TD-DFT) und Fluoreszenz
        orca = out / "input" / "orca.out"
        orca.write_text("\n".join([
            "-" * 77, "         ABSORPTION SPECTRUM VIA TRANSITION ELECTRIC DIPOLE MOMENTS", "-" * 77,
            "State   Energy    Wavelength  fosc         T2        TX        TY        TZ",
            "        (cm-1)      (nm)                 (au**2)    (au)      (au)      (au)", "-" * 77,
            "   1   20660.0    484.0   0.150000000   1.00000   1.00000   0.00000   0.00000",
            "   2   38168.0    262.0   0.600000000   3.00000   1.70000   0.00000   0.00000", ""]),
            encoding="utf-8")
        assert len(__import__("uvvis_extra").parse_tddft(orca)["E"]) == 2
        w.ov_state["entries"] = [
            {"kind": "exp", "file": str(csv), "sample": "TEST-A-0p25mgml_THF", "label": "exp", "color": "#000000",
             "visible": True, "own_norm": False, "nlo": 450, "nhi": 520, "conc": 0.25, "unit": "mg/mL",
             "mw": 95.1, "d": 1.0},
            {"kind": "calc", "file": str(orca), "label": "calc", "color": "#1764e8", "visible": True,
             "own_norm": False, "nlo": 450, "nhi": 520, "fwhm": 0.3, "shift": 0.0}]
        w.ov_state.update(nlo=450, nhi=520)
        for mode in ("norm", "eps"):
            w.ov_state["mode"] = mode
            files = w.ov_tab.export_to(out / f"overlay_{mode}")
            assert all(f.exists() for f in files), files
        # Hinweis im Tab statt nur im Protokoll: ε ohne Konzentration -> Messung ausgelassen, Grund unter der Vorschau
        w.ov_state["entries"][0]["conc"] = None
        w.ov_tab.render()
        assert not w.ov_tab.lbl_note.isHidden() and "„exp“" in w.ov_tab.lbl_note.text(), w.ov_tab.lbl_note.text()
        w.ov_state["entries"][0]["conc"] = 0.25
        w.ov_state["mode"] = "norm"
        w.ov_tab.render()
        assert w.ov_tab.lbl_note.isHidden(), w.ov_tab.lbl_note.text()
        report.append("overlay notes in the tab: OK")
        em = out / "input" / "emission.csv"
        xs = np.arange(400.0, 701.0, 1.0)
        em.write_text("\r\n".join(["EM ex352,,", "Wavelength (nm),Intensity (a.u.),"] +
                                    [f"{x:.1f},{500 * np.exp(-((x - 520) / 30) ** 2):.3f}," for x in xs]),
                      encoding="utf-8", newline="")
        w.fl_state.update(abs_file=str(csv), abs_sample="TEST-A-0p25mgml_THF", em_file=str(em),
                          em_sample="EM ex352", ex=352.0, alo=450, ahi=520, eps="1234",
                          photo_day=str(photo), bg_day="grabcut", photo_uv=str(photo), bg_uv="none")
        files = w.fl_tab.export_to(out / "fluorescence")
        assert all(f.exists() for f in files), files
        # Cary-Eclipse-CSV (Name in Zeile 2, λex aus den Metadaten)
        ecl = out / "input" / "eclipse.csv"
        ecl.write_text("\r\n".join(["Wavelength (nm),Intensity (a.u.),Z Axis,", ",PROBE-1-45p3µM", ",1"] +
                                     [f"{x:g},{300 * np.exp(-((x - 480) / 25) ** 2):.4f}" for x in xs] +
                                     ["", "PROBE-1-45p3µM", "Ex. Wavelength (nm)               387.00"]),
                       encoding="utf-8", newline="")
        es = core.read_samples([ecl])
        assert len(es) == 1 and es[0]["name"] == "PROBE-1-45p3µM" and es[0]["ex"] == 387.0, es[0]["name"]
        report.append("overlay + fluorescence + Eclipse CSV: OK")

        # Fluoreszenz: Stokes-Verschiebung, A(λex), λabs-Label
        w.fl_state.update(stokes_label=True)
        files = w.fl_tab.export_to(out / "fluorescence_stokes")
        assert all(f.exists() for f in files), files
        inf = w.fl_tab.info
        exp_cm = 1e7 / inf["lam_abs"] - 1e7 / inf["lam_em"]
        assert inf["lam_em"] == 520 and abs(inf["stokes_cm"] - exp_cm) < 1e-6 and inf["stokes_cm"] > 0, inf
        assert inf["a_ex"] is not None and np.isfinite(inf["a_ex"]), inf
        assert uvvis_tabs.fmt_stokes(1473.4) == "1470"
        figf = MplFigure(figsize=(8, 6), dpi=PREVIEW_DPI)
        FigureCanvasAgg(figf)
        ltxt = " ".join(a.get_text() for a in w.fl_tab.build(figf).ann_items.values())
        assert "abs" in ltxt and "\\tilde" in ltxt, ltxt
        w.fl_state.update(stokes_label=False)
        report.append("fluorescence Stokes shift + A(ex): OK")
        # Schultern, Zusatzbande, Cutoff und Ausblenden pro Bande
        w.open_csv(csv, ask=False)
        w.glob.update(shoulders=True, bands_add=True, bands="600")
        w.series["TEST-A"]["band_cutoff"] = {"484": 2.0}
        w.series["TEST-A"]["band_hidden"] = [600]
        cfg_t = w.make_cfg("TEST-A")
        _, _, res_t = core.prepare_series(cfg_t, w.groups["TEST-A"], "TEST-A")
        lams = {int(round(r["lam_requested"])): r for r in res_t}
        assert 600 in lams and lams[600].get("hidden") and lams[484]["cutoff"] == 2.0, sorted(lams)
        report.append("bands/shoulders/cutoff: OK")
        # einstellbare R²-Schwelle: 1.0 blendet alles aus, 0 zeigt alles mit Fit
        w.series["TEST-A"]["band_hidden"] = []
        for thr, expect_hidden in ((1.0, True), (0.0, False)):
            w.glob["min_r2"] = thr
            w.apply_globals()
            assert abs(w.sp_r2.value() - thr) < 1e-9
            w.collect_globals()
            _, _, res_r = core.prepare_series(w.make_cfg("TEST-A"), w.groups["TEST-A"], "TEST-A")
            fitted = [r for r in res_r if r["fit"]]
            assert fitted and all(bool(r.get("hidden")) == expect_hidden for r in fitted), (thr, res_r)
        w.glob["min_r2"] = 0.98
        w.apply_globals()
        report.append("min. R² threshold: OK")

        def pump(sec=0.4):
            import time
            t0 = time.time()
            while time.time() - t0 < sec:
                QApplication.processEvents()
                time.sleep(0.02)

        # TD-DFT mit Spin-Bahn-Kopplung (ORCA 6): Zustandsnamen wie "0-1.0A" dürfen nicht als Zahl gelten
        soc = out / "input" / "orca_soc.out"
        hdr = ["-" * 104, "      SOC CORRECTED ABSORPTION SPECTRUM VIA TRANSITION ELECTRIC DIPOLE MOMENTS", "-" * 104,
               "      Transition         Energy     Energy  Wavelength fosc(D2)      D2       |DX|      |DY|      |DZ|",
               "                          (eV)      (cm-1)    (nm)                 (au**2)    (au)      (au)      (au)",
               "-" * 104,
               "  0-1.0A  ->  1-3.0A    2.898266   23376.1   427.8   0.000000000   0.00000   0.00001   0.00000   0.00000",
               "  0-1.0A  -> 10-1.0A    4.004892   32301.6   309.6   0.120701528   1.23017   0.69013   0.85513   0.15048",
               ""]
        soc.write_text("\n".join(orca.read_text(encoding="utf-8").splitlines() + hdr), encoding="utf-8")
        import uvvis_extra as XX
        r_soc = XX.parse_tddft(soc)
        assert r_soc["variant"] == "soc" and abs(r_soc["E"][1] - 4.004892) < 1e-6 and abs(r_soc["f"][1] - 0.1207) < 1e-3, r_soc
        assert XX.parse_tddft(soc, "nosoc")["variant"] == "nosoc"
        report.append("TD-DFT SOC: OK")

        # Zuschnitt von Strukturen: senkrechte/waagerechte Endbindungen dürfen nicht abgeschnitten werden
        import pymupdf
        tdoc = pymupdf.open()
        pg = tdoc.new_page(width=200, height=200)
        pg.draw_line((50, 50), (150, 50), width=1)        # waagerecht
        pg.draw_line((100, 50), (100, 160), width=1)      # senkrecht, reicht am weitesten nach unten
        cr = IMG._content_rect(pg)
        assert cr.y1 >= 160 and cr.x0 <= 50 and cr.x1 >= 150, cr
        report.append("structure trimming: OK")

        # Einfärben von Strukturen (Overlay: Struktur pro Spektrum in Kurvenfarbe)
        ca = IMG.load_structure(cdxml, color="#e8231b")
        opaque = ca["rgba"][..., 3] > 200
        assert ca["pdf"] and abs(int(np.median(ca["rgba"][opaque][:, 0])) - 232) < 3, "Einfärben"
        report.append("structure colour: OK")

        # Neu platzierte Labels dürfen nicht auf gespeicherten landen (zwei Banden, eine gespeichert)
        cfg_l = w.make_cfg("TEST-A")
        cfg_l2, sp_l, res_l = core.prepare_series(cfg_l, w.groups["TEST-A"], "TEST-A")
        figl = MplFigure(figsize=cfg_l2["figsize_in"], dpi=PREVIEW_DPI)
        FigureCanvasAgg(figl)
        auto = core.build_figure(plt, cfg_l2, sp_l, res_l, {}, fig=figl).current_layout()
        lk = [k for k in auto if k.startswith("label_")]
        assert len(lk) >= 2, lk
        Fl = core.build_figure(plt, cfg_l2, sp_l, res_l, {lk[-1]: auto[lk[0]]}, fig=figl)
        rl = Fl._renderer()
        b0, b1 = (Fl.ann_items[k].get_window_extent(rl) for k in (lk[0], lk[-1]))
        assert not b0.overlaps(b1), "Label-Kollision"
        report.append("label collision: OK")

        # Overlay-Optionen: Skalierung, Versatz, Linienart, Strich-Modi
        w.ov_state["entries"][1].update(scale=0.8, offset=0.2, ls="dotted", lw=2.0)
        for mode in ("max", "band", "faxis"):
            w.ov_state.update(mode="norm", stick_mode=mode, sticks=True)
            files = w.ov_tab.export_to(out / f"overlay_sticks_{mode}")
            assert all(f.exists() for f in files), files
        report.append("overlay options: OK")

        # Overlay gestapelt: Ebenen in Listenreihenfolge, TD-DFT auf der Ebene der Messung, y ohne Zahlen
        ents = w.ov_state["entries"]
        ents.append(dict(ents[0], uid="stack3", label="exp 2", color="#e8231b"))
        ents[1]["stack_with_prev"] = True

        def ov_fig():
            figo = MplFigure(figsize=(8, 6), dpi=PREVIEW_DPI)
            FigureCanvasAgg(figo)
            Fo = w.ov_tab.build(figo)
            figo.canvas.draw()
            return Fo
        w.ov_state.update(mode="norm", stick_mode="faxis", stack=False, curve_labels=False, zero_lines=False)
        flat = ov_fig()
        w.ov_state.update(stack=True, stack_gap=1.2, y_axis="auto", curve_labels=True, zero_lines=True)
        stk = ov_fig()
        for i, dy in enumerate((1.2, 1.2, 0.0)):
            d = np.nanmedian(stk.curves[i]["y"] - flat.curves[i]["y"])
            assert abs(d - dy) < 1e-9, (i, d)
        assert stk.ax2 is None and not stk.ax.yaxis.get_major_ticks()[0].label1.get_visible()
        assert len(stk.curve_texts) == 2 and stk.legend is None, len(stk.curve_texts)
        files = w.ov_tab.export_to(out / "overlay_stacked")
        assert all(f.exists() for f in files), files
        ents.pop()
        ents[1].pop("stack_with_prev")
        w.ov_state.update(stack=False, curve_labels=False, zero_lines=False, y_axis="auto")
        report.append("overlay stacked: OK")

        # TD-DFT: Triplett-Zustände (ORCA 6 Bezeichnung, ORCA 4/5 eigener Block, Gaussian) und Kombinieren
        import uvvis_extra as XT
        sep = "-" * 100
        tab_head = [sep, "      ABSORPTION SPECTRUM VIA TRANSITION ELECTRIC DIPOLE MOMENTS", sep,
                    "   Transition    Energy   Energy  Wavelength fosc(D2)   D2   DX   DY   DZ",
                    "                  (eV)    (cm-1)    (nm)", sep]
        o6 = out / "input" / "orca6_trip.out"
        o6.write_text("\n".join(tab_head + [
            "  0-1A  ->  1-3A    2.500000   20164.1   495.9   0.000000000   0.00000   0.00000   0.00000   0.00000",
            "  0-1A  ->  1-1A    3.000000   24196.9   413.3   0.200000000   1.00000   1.00000   0.00000   0.00000",
            ""]), encoding="utf-8")
        o5 = out / "input" / "orca5_trip.out"
        o5.write_text("\n".join([
            "TD-DFT/TDA EXCITED STATES (TRIPLETS)", "------------------------------------", "",
            "STATE  1:  E=   0.091874 au      2.600 eV    20970.5 cm**-1 <S**2> =   2.000000", "",
            "STATE  2:  E=   0.110000 au      2.993 eV    24142.0 cm**-1 <S**2> =   2.000000", ""]
            + tab_head + ["   1   26617.0    375.7   0.300000000   1.00000   1.00000   0.00000   0.00000", ""]),
            encoding="utf-8")
        gx = out / "input" / "gauss_trip.log"
        gx.write_text("\n".join([
            " Excited State   1:     Triplet-A      2.7000 eV  459.20 nm  f=0.0000  <S**2>=2.000",
            " Excited State   2:     Singlet-A      3.2000 eV  387.45 nm  f=0.1500  <S**2>=0.000", ""]),
            encoding="utf-8")
        for pth, n_t, e_t in ((o6, 1, 2.5), (o5, 2, 2.6), (gx, 1, 2.7)):
            rt = XT.parse_tddft(pth)
            trips = [E for E, m in zip(rt["E"], rt["mult"]) if m == 3]
            assert len(trips) == n_t and abs(min(trips) - e_t) < 1e-6, (pth.name, rt)
            assert all(fv == 0 for fv, m in zip(rt["f"], rt["mult"]) if m == 3), pth.name
        n_before = len(ents)
        ents += [{"kind": "calc", "file": str(p_), "label": p_.stem, "color": "#1f9e4a", "visible": True,
                  "own_norm": False, "nlo": 450, "nhi": 520, "fwhm": 0.3, "shift": 0.0} for p_ in (o6, gx)]
        w.ov_tab.fill_table()
        for r_ in range(n_before):                      # nur die beiden neuen markieren
            ents[r_]["visible"] = r_ == 0
        w.ov_tab.table.clearSelection()
        w.ov_tab.table.setSelectionMode(QAbstractItemView.MultiSelection)
        assert not w.ov_tab.btn_combine.isEnabled()                 # nichts markiert: „Kombinieren“ aus
        for r_ in (n_before, n_before + 1):
            w.ov_tab.table.selectRow(r_)
        assert w.ov_tab.btn_combine.isEnabled()                     # zwei TD-DFT-Einträge markiert
        w.ov_tab.combine_selected()
        comb = ents[-1]
        assert len(comb["sources"]) == 2 and not ents[n_before]["visible"], comb
        E_c, f_c, m_c = w.ov_tab.calc_data(comb)
        assert len(E_c) == 4 and m_c.count(3) == 2 and abs(f_c.sum() - 0.35) < 1e-9, (E_c, f_c, m_c)
        comb["sources"][1]["weight"] = 0.5
        assert abs(w.ov_tab.calc_data(comb)[1].sum() - 0.275) < 1e-9
        w.ov_state.update(mode="norm", triplets=True, sticks=True, stick_mode="band")
        _, _, _, stx = w.ov_tab.figure_spec()
        dashed = np.concatenate([s_["x"] for s_ in stx if s_.get("ls") == "dashed"])
        assert np.allclose(sorted(dashed), sorted(1e7 / (np.array([2.5, 2.7]) * XT.EV_TO_CM))), dashed
        files = w.ov_tab.export_to(out / "overlay_triplets")
        assert all(f.exists() for f in files), files
        assert any(c_ is comb["sources"][0] for c_, _, _ in PRJ._path_slots({"overlay": w.ov_state})), \
            "kombinierte Quellen fehlen im Projekt"
        del ents[n_before:]
        ents[1]["visible"] = True
        w.ov_state.update(triplets=False)
        w.ov_tab.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        w.ov_tab.fill_table()
        e_cc = next(e_ for e_ in w.ov_state["entries"] if e_["kind"] == "calc")
        e_cc.update(color_c=1.0, color_d=1.0)                  # Lösungsfarbe aus dem TD-DFT-Spektrum (1 mM, 1 cm)
        hx1_, _ = w.ov_tab.calc_color(e_cc)
        e_cc["color_c"] = 1e4
        hx2_, _ = w.ov_tab.calc_color(e_cc)
        lum_ = lambda h_: sum(int(h_[i_:i_ + 2], 16) for i_ in (1, 3, 5))      # noqa: E731
        assert re.fullmatch(r"#[0-9A-F]{6}", hx1_) and lum_(hx2_) <= lum_(hx1_), (hx1_, hx2_)
        e_cc["color_c"] = 1.0
        w.ov_tab.update_calc_color(e_cc)
        assert hx1_ in w.ov_tab.lbl_ccol_info.text(), w.ov_tab.lbl_ccol_info.text()
        report.append("TD-DFT triplets + combine: OK")

        # Export ins Projekt inkl. Excel-Report
        w.open_csv(csv, ask=False)
        w.series["TEST-A"]["mw"] = 95.1
        w.set_structure(str(cdxml), "TEST-A", render=False, sync=True, autofill_mw=False)
        w.render()
        pump()
        files = w.export_all()
        rep = proj / "exports" / "epsilon" / "TEST-A" / "TEST-A_report.xlsx"
        assert rep in files and rep.exists(), files
        meth_ = (proj / "exports" / "epsilon" / "TEST-A" / "TEST-A_methods.txt").read_text(encoding="utf-8")
        assert "ordinary least squares" in meth_ and "95% confidence interval" in meth_ and " nm: " in meth_ \
            and "M = 95.10 g mol-1" in meth_, meth_
        assert meth_.split("\n\n")[0] in w.eps_methods_text(), "SI-Text Export/Knopf verschieden"
        from openpyxl import load_workbook
        wbk = load_workbook(rep)
        assert len(wbk.sheetnames) >= 3 and any("nm" in n for n in wbk.sheetnames), wbk.sheetnames
        report.append("export into project + Excel report: OK")

        # Textformat: λ/ε in Symbolschrift, Labels fett/farbig, Rahmen um Plot und Struktur
        w.text_fmt = dict(core.TEXT_DEFAULTS, bold=True, color="#1764e8", label_pt=12.0,
                          axes_frame=True, struct_frame=True)
        w.render()
        assert w.F.ax.spines["top"].get_visible(), "Rahmen oben fehlt"
        ann = next(a for k, a in w.F.ann_items.items() if k.startswith("label_"))
        assert "\\mathsf{λ}" in ann.get_text() and ann.get_fontweight() == "bold" and ann.get_fontsize() == 12.0
        sa = w.F.axes_items["image_struktur"]
        assert sa._frame is not None
        rf = w.F._renderer()
        fb, ib = sa._frame.get_window_extent(rf), sa.get_window_extent(rf)
        assert fb.x0 < ib.x0 and fb.y1 > ib.y1, "Rahmen liegt nicht außen um die Struktur"
        for k, b in [(k, a.get_window_extent(rf)) for k, a in w.F.ann_items.items()]:
            assert not fb.overlaps(b), f"Strukturrahmen überdeckt {k}"
        files = w.export_all(out / "textformat")
        pdf = next(f for f in files if f.suffix == ".pdf")
        fonts = " ".join(f[3] for f in pymupdf.open(str(pdf))[0].get_fonts())
        sym_font = core.first_font(core.DEFAULTS["symbol_font"], "DejaVu Serif").replace(" ", "")
        assert sym_font.lower() in fonts.lower().replace("-", ""), fonts
        svg = next(f for f in files if f.suffix == ".svg").read_text(encoding="utf-8")
        assert "structure_0" in svg, "SVG ohne Struktur"
        w.fl_tab.export_to(out / "textformat_fluo")
        w.text_fmt = dict(core.TEXT_DEFAULTS)
        w.render()
        report.append("text format + frames: OK")

        # Hover: Maus über einem Messpunkt -> λ und A in der Statusleiste; Doppelklick -> Bande
        from matplotlib.backend_bases import MouseEvent
        w.glob["bands"] = ""
        w.apply_globals()
        w.render()
        w.canvas.draw()
        hv, ax = w.hover, w.F.ax
        def peak(c):                                   # Bande bei 484 nm, nur Kurven mit Peak im Bild
            xs = np.asarray(c["x"])
            i = int(np.nanargmax(np.where((xs > 400) & (xs < 560), c["y"], -np.inf)))
            return i, float(c["y"][i])
        top = ax.get_ylim()[1]
        sp0 = max((c for c in hv.curves if peak(c)[1] < 0.9 * top), key=lambda c: peak(c)[1])
        i0 = peak(sp0)[0]
        px, py = ax.transData.transform((sp0["x"][i0], sp0["y"][i0]))
        w.canvas.callbacks.process("motion_notify_event", MouseEvent("motion_notify_event", w.canvas, px, py))
        msg = w.statusBar().currentMessage()
        # Mausereignisse haben ganzzahlige Pixel, Messpunkte liegen in der Vorschau < 1 px auseinander:
        # Treffer darf der Nachbarpunkt sein. Geprüft wird, dass die Anzeige genau den markierten Punkt zeigt.
        assert hv.hit is not None, msg
        hc, hi = hv.hit
        # erlaubt: ein Messpunkt Abstand plus ein Pixel in Datenkoordinaten (Vorschaugröße je nach System/Schrift)
        step = float(np.median(np.abs(np.diff(np.asarray(sp0["x"], float)))))
        nm_per_px = abs(ax.transData.inverted().transform((px + 1, py))[0] - sp0["x"][i0])
        assert abs(hc["x"][hi] - sp0["x"][i0]) <= step + nm_per_px + 1e-9, (msg, hc["x"][hi], step, nm_per_px)
        assert f"λ = {hc['x'][hi]:.1f} nm" in msg and f"A = {hc['y'][hi]:.4f}" in msg, msg
        assert list(hv.marker.get_xdata()) == [hc["x"][hi]], hv.marker.get_xdata()
        assert hv.marker.get_animated(), "Marker darf nicht exportiert werden"
        w.canvas.callbacks.process("button_press_event",
                                   MouseEvent("button_press_event", w.canvas, px, py, button=1, dblclick=True))
        lam0 = round(float(hc["x"][hi]))
        assert w.glob["bands"] == f"{lam0:g}", w.glob["bands"]
        w.pick_band(lam0 + 1)                       # erneut (±2 nm) -> wieder entfernt
        assert w.glob["bands"] == "", w.glob["bands"]
        w.render()
        report.append("hover readout + band picking: OK")

        # Ziehen mit Blitting: Inset verschieben, beim Loslassen wieder normal gezeichnet und gespeichert
        w.render()
        w.canvas.draw()
        key_, a_ = next((k, a) for k, a in w.F.axes_items.items() if a.get_visible())
        p0 = a_.get_position().frozen()
        cx, cy = a_.get_window_extent().x0 + 5, a_.get_window_extent().y0 + 5
        proc = w.canvas.callbacks.process
        proc("button_press_event", MouseEvent("button_press_event", w.canvas, cx, cy, button=1))
        assert a_.get_animated() and w.dragger._bg is not None, "Blitting nicht aktiv"
        proc("motion_notify_event", MouseEvent("motion_notify_event", w.canvas, cx + 30, cy + 20, button=1))
        # nach dem Loslassen muss ganz neu gezeichnet werden (Blitting-Hintergründe, auch der der Hover-Anzeige, sind
        # ohne das gezogene Element entstanden); geprüft wird die Anforderung, nicht ihr Zeitpunkt (CI ohne Fenster)
        idle_ = []
        w.canvas.draw_idle = lambda *a_k: idle_.append(1)
        proc("button_release_event", MouseEvent("button_release_event", w.canvas, cx + 30, cy + 20, button=1))
        del w.canvas.draw_idle
        assert idle_, "nach dem Ziehen kein Neuzeichnen angefordert"
        assert w.hover.bg is None, "Hover-Hintergrund aus der Zeit des Ziehens noch gültig"
        pump()
        p1 = a_.get_position()
        assert p1.x0 > p0.x0 and p1.y0 > p0.y0 and not a_.get_animated() and w.dragger._bg is None, (p0, p1)
        assert all(getattr(ann, "_draggable", None) is None or ann._draggable._use_blit
                   for ann in w.F.ann_items.values()), "Labels ohne Blitting"
        w.series[w.current]["layout"] = {}
        w.render()
        report.append("dragging with blitting: OK")
        w.glob["label_index"] = True                            # λ₁/ε₁ … aufsteigend nach λ
        w.apply_globals()
        w.render()
        lt_ = sorted((float(k_.split("_")[1]), a_.get_text()) for k_, a_ in w.F.ann_items.items())
        assert lt_ and all(t_.count(f"_{{{i_}}}") == 2 for i_, (_, t_) in enumerate(lt_, 1)), lt_
        w.glob["label_index"] = False
        w.apply_globals()
        w.render()
        assert not any("_{" in a_.get_text() for a_ in w.F.ann_items.values())
        report.append("label indices: OK")

        # Cyclovoltammetrie: Gamry-Dateien (beide Varianten), Redoxpaare, interne Referenz, Export
        def cv_sweep(e0, ip, lo, hi, step=0.002, ref=None):
            """Modell-CV: Hinlauf lo->hi, Rücklauf hi->lo; Peaks bei E0 ± 30 mV (+ optional Referenzpaar)."""
            up = np.arange(lo, hi + step / 2, step)
            E_ = np.concatenate([up, up[::-1]])
            fw = np.r_[np.ones(len(up)), np.zeros(len(up))].astype(bool)

            def wave(e_c, i_p):
                pa = i_p * np.exp(-((E_ - (e_c + 0.03)) / 0.045) ** 2) + 0.3 * i_p / (1 + np.exp(-(E_ - e_c - 0.03) / 0.02))
                pc = -0.9 * i_p * np.exp(-((E_ - (e_c - 0.03)) / 0.045) ** 2)
                return np.where(fw, pa, pc + 0.3 * i_p * np.exp(-(hi - E_) / 0.1))
            I_ = wave(e0, ip) + (wave(ref, 0.8 * ip) if ref is not None else 0)
            return E_, I_

        def write_dta(path, cycles, rate, cycle_col=False, times=None):
            head = ["EXPLAIN", "TAG\tCV", "DATE\tLABEL\t1/1/2026\tDate", "TIME\tLABEL\t12:00:00\tTime",
                    f"SCANRATE\tQUANT\t{rate:.5E}\t&Scan Rate (mV/s)", "AREA\tQUANT\t7.10000E-002\tElectrode &Area (cm^2)",
                    "CONVENTION\tIQUANT\t1\tCurrent Convention"]
            cols = "\tPt\tT\tVf\tIm\tVu" + ("\tCycle" if cycle_col else "")
            units = "\t#\ts\tV vs. Ref.\tA\tV" + ("\t#" if cycle_col else "")
            body = []
            if cycle_col:
                body += [f"CURVE\tTABLE\t{sum(len(c[0]) for c in cycles)}", cols, units]
            pt = 0
            for n, (E_, I_) in enumerate(cycles):
                if not cycle_col:
                    body += [f"CURVE{n + 1}\tTABLE", cols, units]
                for e_, i_ in zip(E_, I_):
                    t_ = times[pt] if times is not None else pt * 0.02
                    body.append(f"\t{pt}\t{t_:.2f}\t{e_:.5E}\t{i_ * 1e-6:.5E}\t0" + (f"\t{n}" if cycle_col else ""))
                    pt += 1
            path.write_text("\r\n".join(head + body) + "\r\n", encoding="latin-1", newline="")
        cvdir = out / "input" / "cv"
        cvdir.mkdir(parents=True, exist_ok=True)
        cv_files = []
        for rate, ip in ((50, 40.0), (100, 56.0), (200, 80.0)):
            p_ = cvdir / f"cv_{rate}.DTA"
            write_dta(p_, [cv_sweep(-1.02, ip, -1.4, -0.5)], rate)
            cv_files.append(p_)
        ref_file = cvdir / "cv_ref.DTA"                  # zwei Zyklen, Variante mit Spalte 'Cycle'
        write_dta(ref_file, [cv_sweep(-1.02, 80.0, -1.4, 0.9, ref=0.55)] * 2, 200, cycle_col=True)
        d_ = CVMOD.read_gamry(ref_file)
        assert len(d_["cycles"]) == 2 and d_["meta"]["scanrate"] == 200 and d_["meta"]["area"] == 0.071, d_["meta"]
        c_ = CVMOD.read_gamry(cv_files[0])["cycles"][0]
        sug = CVMOD.suggest_couples(c_["E"], c_["I"])
        assert len(sug) == 1 and sug[0][0] < -1.02 < sug[0][1], sug
        r_ = CVMOD.analyze_couple(c_["E"], c_["I"], *sug[0])
        assert abs(r_["e12"] + 1.02) < 0.003 and abs(r_["dep"] - 0.06) < 0.006 and not r_["edge"], r_
        st_cv = w.cv_state
        st_cv["files"] = [{"file": str(p_), "uid": f"cv{i}", "visible": True, "color": "#000000",
                           "label": f"{r} mV/s"} for i, (p_, r) in enumerate(zip(cv_files, (50, 100, 200)))]
        st_cv["files"].append({"file": str(ref_file), "uid": "cvref", "visible": False, "color": "#000000",
                               "label": "ref"})
        st_cv.update(couples=[list(sug[0])], ref_mode="internal", ref_uid="cvref", ref_lo=0.3, ref_hi=0.8,
                     ref_compound="ac2fc", ref_value=0.49, cycle=2)
        w.cv_tab.fill_table()
        files = w.cv_tab.export_to(out / "cv_test")
        assert all(f.exists() for f in files), files
        res = w.cv_tab.results
        assert len(res) == 3 and all(abs(r["e12"] - (-1.02 - (0.55 - 0.49))) < 0.004 for r in res), res
        assert "Ac₂Fc" in w.cv_tab.note and "+0.49 V vs. Fc" in w.cv_tab.note, w.cv_tab.note
        b_cv = w.cv_tab.table.cellWidget(0, 2)                # Farbverlauf: Feld zeigt die Farbe im Bild, gesperrt
        assert st_cv["gradient"] and not b_cv.isEnabled() and \
            w.cv_tab.grad_colors[st_cv["files"][0]["uid"]] in b_cv.styleSheet(), b_cv.styleSheet()
        rs = CVMOD.randles_sevcik([50, 100, 200], [r["ipa"] * 1e6 for r in res])
        assert rs and rs[2] > 0.99, rs
        csv_txt = files[-1].read_text(encoding="utf-8-sig")
        assert csv_txt.count("\n") == 4 and "E_half_V" in csv_txt and "charging_step_uA" in csv_txt, csv_txt
        meth = files[-2].read_text(encoding="utf-8")
        assert "Nicholson" in meth and "diacetylferrocene" in meth and "95% confidence" in meth \
            and "Couple 1: E1/2 = −1.08 V vs. Fc+/Fc" in meth, meth
        assert any(c_ is st_cv["files"][0] for c_, _, _ in PRJ._path_slots({"cv": st_cv})), "CV-Dateien fehlen im Projekt"
        # „Alles auf Standard“: Einstellungen zurück; Messungen, Redoxpaare, Referenz, Unterschrift bleiben
        snap_cv = json.loads(json.dumps({k_: st_cv[k_] for k_ in ("files", "couples", "ref_mode", "ref_value", "cycle",
                                                                  "gradient", "baseline")}))
        st_cv.update(gradient=False, baseline=False, stack=True, cap_temp="−78 °C")
        w.cv_tab.reset_defaults()
        assert st_cv["gradient"] and st_cv["baseline"] and not st_cv["stack"] and st_cv["cycle"] == 1
        assert st_cv["cap_temp"] == "−78 °C" and all(st_cv[k_] == snap_cv[k_] for k_ in ("files", "couples", "ref_mode",
                                                                                     "ref_value"))
        assert w.cv_tab.cb_gradient.isChecked() and w.main_tabs.currentWidget() is not None
        st_cv.update(cycle=snap_cv["cycle"], cap_temp="rt")
        w.cv_tab.fill_table()
        report.append("cyclic voltammetry: OK")

        # SI-Export: Abbildungen aller Tabs mit Daten (ε je Serie, CV …), zwei je Seite
        fb_ = w.si_figure_blocks(out / "si_figs", "TEST-1")
        caps_ = [b_[1] for b_ in fb_ if b_[0] == "figcap"]
        imgs_ = [b_ for b_ in fb_ if b_[0] == "image"]
        assert any(c_.startswith("UV-Vis spectra of <b>TEST-1</b>") and "cm quartz cuvettes." in c_ for c_ in caps_), caps_
        assert any(c_.startswith("Cyclic voltammograms of <b>TEST-1</b>") for c_ in caps_) and len(imgs_) == len(caps_), caps_
        assert all(Path(b_[1]).exists() and Path(b_[4]).exists() and 0 < b_[2] <= 16.0 for b_ in imgs_), imgs_
        assert fb_.count(("pagebreak",)) == (len(imgs_) - 1) // 2 and w.si_available()["figs"]
        # Methodik der Fits im Word-Export: Texte wie „SI-Text kopieren“, je Technik ein Absatz mit fettem Titel
        mb_ = w.si_methods_blocks()
        assert mb_[0] == ("heading", "Data processing and fitting") and w.si_available()["methods"], mb_[:1]
        mtx_ = [b_[1] for b_ in mb_[1:]]
        assert any(t_.startswith("<b>UV-Vis absorption spectroscopy.</b> ") for t_ in mtx_), mtx_
        assert any(t_.startswith("<b>Cyclic voltammetry.</b> ") and "Nicholson" in t_ for t_ in mtx_), mtx_
        assert any("E<sub>1/2</sub> = −1.08 V" in t_ for t_ in mtx_) and "cm<sup>-1</sup>" in "".join(mtx_), mtx_
        assert SI.methods_html("A_red, g_lit, E1/2,app, M-1 cm-1") == \
            "A<sub>red</sub>, g<sub>lit</sub>, E<sub>1/2</sub>,app, M<sup>-1</sup> cm<sup>-1</sup>"
        b_m, _ = w.si_blocks(["methods"])
        assert ("heading", "Data processing and fitting") in b_m and not any(b_[0] == "image" for b_ in b_m)
        mdoc_ = SI.write_docx(out / "si_methods.docx", b_m)
        with __import__("zipfile").ZipFile(mdoc_) as z_:
            mx_ = z_.read("word/document.xml").decode("utf-8")
            assert "Data processing and fitting" in mx_ and 'w:ascii="Arial"' in mx_
        SI.set_font("Times New Roman")                                   # Schriftart wählbar, Größe bleibt 11 pt
        with __import__("zipfile").ZipFile(SI.write_docx(out / "si_font.docx", b_m)) as z_:
            mx_ = z_.read("word/document.xml").decode("utf-8")
            assert 'w:ascii="Times New Roman"' in mx_ and 'w:ascii="Arial"' not in mx_ and 'w:sz w:val="22"' in mx_
            assert 'w:ascii="Times New Roman"' in z_.read("word/styles.xml").decode("utf-8")
        SI.set_font("Arial")
        cap2_ = "Cyclic voltammograms of <b>X</b> (1.8±0.1 mg mL-1) in MeCN (0.1 M n-Bu4NPF6) at rt; " * 2
        tp_ = SI.two_per_page([("char", "Setup text. " * 40)] + [("image", "a.png", 16.0, 11.0, None),
                                                                   ("figcap", cap2_)] * 3)
        assert [b_[0] for b_ in tp_] == ["char", "image", "figcap", "image", "figcap", "pagebreak", "image", "figcap"]
        h_ = [b_[3] for b_ in tp_ if b_[0] == "image"]
        assert h_[0] == h_[1] < h_[2] == 11.0 and abs(tp_[1][2] / tp_[1][3] - 16 / 11) < 1e-9, h_   # zwei je Seite
        w.cv_state.update(cap_conc="1.8±0.1 mg mL-1", cap_el="0.1 M n-Bu4NPF6")        # CV-Unterschrift wie in der SI
        cap_cv = w.cv_tab.si_caption("TEST-1")
        assert cap_cv.startswith("Cyclic voltammograms of <b>TEST-1</b> (1.8±0.1 mg mL<sup>-1</sup>)") \
            and "(0.1 M n-Bu<sub>4</sub>NPF<sub>6</sub>) at rt" in cap_cv and cap_cv.endswith(".") \
            and "referenced internally against diacetylferrocene" in cap_cv, cap_cv
        assert SI.chem_html("CH2Cl2, AgNO3 (0.01 M), 7.1 mm², 3 min") == \
            "CH<sub>2</sub>Cl<sub>2</sub>, AgNO<sub>3</sub> (0.01 M), 7.1 mm², 3 min"
        w.settings.setValue("si_setup/cv", "CVs were measured in CH3CN.")              # Aufbau-Absatz vor den CVs
        fb_ = w.si_figure_blocks(out / "si_figs", "TEST-1")
        assert ("char", "CVs were measured in CH<sub>3</sub>CN.") in fb_ and \
            ("figcap", cap_cv) in fb_, [b_ for b_ in fb_ if b_[0] != "image"]
        w.settings.remove("si_setup/cv")
        w.cv_state.update(cap_conc="", cap_el="")
        report.append("SI figures from all tabs (CV caption, setup text): OK")

        # Tab-Reihenfolge wie die SI; offener Tab über den Namen (alte Projekte: Position in der alten Reihenfolge)
        assert [w.main_tabs.widget(i) for i in range(5)] == [w.syn_tab, w.nmr_tab, w.irms_tab, w.hrms_tab, w.eps_page]
        st_ = w.get_state()
        st_["tab"] = 3                                                   # 1.23: Position 3 = CV
        w.set_state(st_)
        assert w.main_tabs.currentWidget() is w.cv_tab and w.get_state()["tab"] == "cv"
        # Textformat je Tab: eigenes Format nur für CV, der Rest folgt dem Standard
        w.tab_fmt["cv"] = dict(core.TEXT_DEFAULTS, bold=True)
        assert w.cv_tab.base_cfg()["labels"]["bold"] and not w.sec_tab.base_cfg()["labels"]["bold"]
        assert w.get_state()["text_tabs"]["cv"]["bold"] and w.text_tab()[0] == "cv"
        dlg_ = TextFormatDialog(w, w.text_fmt, w.tab_fmt["cv"], w.cv_tab.TEXT_OPTS, "cv", "CV")
        assert dlg_.cmb_scope.currentData() == "own" and dlg_.cb_bold.isChecked() and dlg_.cb_greek.isHidden()
        dlg_.cmb_scope.setCurrentIndex(0)                                  # zurück auf „alle Tabs“
        assert not dlg_.cb_bold.isChecked() and dlg_.results()[1] is None
        w.tab_fmt.clear()
        w.main_tabs.setCurrentWidget(w.eps_page)
        report.append("tab order + open tab by name, text format per tab: OK")

        # CV-Basislinie: Ladestrom ±C·v + linearer Untergrund dürfen i_p und i_r/i_f nicht verändern
        E_, I_ = cv_sweep(-1.02, 50.0, -1.4, -0.5)
        dEs_ = np.gradient(E_)
        I_c = I_ + np.where(dEs_ > 0, 6.0, -6.0) + 15.0 * (E_ + 1.0)
        win_ = sug[0]
        r0_ = CVMOD.analyze_couple(E_, I_, *win_, use_baseline=True)
        r1_ = CVMOD.analyze_couple(E_, I_c, *win_, use_baseline=True)
        r2_ = CVMOD.analyze_couple(E_, I_c, *win_)
        assert r0_["bl"] and r1_["bl"] and not r1_["bl_fail"], r1_
        assert abs(r1_["bl"]["step"] - 12.0) < 0.5, r1_["bl"]
        for k_ in ("ipa", "ipc", "ratio"):
            assert abs(r1_[k_] - r0_[k_]) < 0.01 * abs(r0_[k_]) + 0.05, (k_, r0_[k_], r1_[k_])
        assert abs(r2_["ipa"] - r1_["ipa"]) > 4, (r2_["ipa"], r1_["ipa"])        # ohne Basislinie verfälscht
        assert abs(r1_["e12"] - r2_["e12"]) < 1e-12                               # Potentiale unverändert
        # kathodisch zuerst (Spiegelbild): gleiche Korrektur mit umgekehrtem Vorzeichen
        r3_ = CVMOD.analyze_couple(-E_, -I_c, -win_[1], -win_[0], use_baseline=True)
        assert abs(r3_["ratio"] - r1_["ratio"]) < 1e-6 and abs(r3_["bl"]["step"] - r1_["bl"]["step"]) < 1e-6, r3_
        # zu wenig Vorlauf im Fenster: Fallback von null, markiert
        r4_ = CVMOD.analyze_couple(E_, I_c, -1.08, -0.8, use_baseline=True)
        assert r4_ is not None and r4_["bl_fail"] and r4_["ipa"] == r4_["ipa_raw"], r4_
        # breite Welle (iR-Abfall, wie THF mit ΔE_p ≈ 240 mV): Basislinie endet vor dem Anstieg, nicht im Fuß
        upw_ = np.arange(-1.6, -0.3, 0.002)
        Ew_ = np.concatenate([upw_, upw_[::-1]])
        fww_ = np.r_[np.ones(len(upw_)), np.zeros(len(upw_))].astype(bool)
        Iw_ = np.where(fww_, 100 * np.exp(-((Ew_ + 0.80) / 0.12) ** 2), -80 * np.exp(-((Ew_ + 1.04) / 0.12) ** 2))
        rw_ = CVMOD.analyze_couple(Ew_, Iw_, -1.5, -0.5, use_baseline=True)
        assert rw_["bl"] and abs(rw_["ipa"] - rw_["ipa_raw"]) < 1.0 and rw_["bl"]["e1"] < -1.0, rw_["bl"]   # alte Regel: -0.95 (21 % im Fuß)
        # Nicholson nur ohne weiteren Prozess zwischen Hinpeak und Umkehrpotential (hier: Referenzpaar bei +0.55 V)
        cr_ = CVMOD.read_gamry(ref_file)["cycles"][0]
        rr1_ = CVMOD.analyze_couple(cr_["E"], cr_["I"], *win_)
        rr2_ = CVMOD.analyze_couple(cr_["E"], cr_["I"], 0.3, 0.8)
        assert not rr1_["ratio_ok"] and not np.isfinite(rr1_["ratio"]) and rr2_["ratio_ok"], (rr1_, rr2_)
        assert r1_["ratio_ok"] and np.isfinite(r1_["ratio"])
        f_ = CVMOD.rate_fit([1, 2, 3, 4], [2.1, 3.9, 6.2, 7.9])
        # von Hand: RSS = 0.063, SE = √(0.063/2/5) = 0.079373, t(0.975; 2) = 4.302653
        assert abs(f_["slope"] - 1.97) < 1e-9 and abs(f_["ci"] - 4.302653 * 0.0793725) < 1e-5, f_
        assert CVMOD.fmt_ci(f_["slope"], f_["ci"]) == "2.0 ± 0.3", CVMOD.fmt_ci(f_["slope"], f_["ci"])
        ll_ = CVMOD.loglog_fit([50, 100, 200, 400], [10 * (v / 50) ** 0.5 for v in (50, 100, 200, 400)])
        assert abs(ll_["slope"] - 0.5) < 1e-9, ll_
        report.append("CV baseline, Nicholson with baseline, fits with CI: OK")

        # CV-Ausbau: Gruppe + eigene Referenz je Messung (zweites "Lösungsmittel"), Stapeln, E½ je Kurve,
        # Auftragungen gegen die Scanrate (ΔE_p, Randles-Ševčík, Trumpet)
        t_cv = w.cv_tab
        for e_ in st_cv["files"]:
            e_["group"] = "A"
        for r_ in (50, 100):                            # Gruppe B: E½ −0.92 V, eigenes Ac₂Fc bei +0.75 V (gemessen)
            p_ = cvdir / f"cvB_{r_}.DTA"
            write_dta(p_, [cv_sweep(-0.92, 60.0 * (r_ / 50) ** 0.5, -1.4, 1.1, ref=0.75)], r_)
            st_cv["files"].append({"file": str(p_), "uid": f"cvB{r_}", "visible": True, "color": "#000000",
                                   "label": f"B {r_} mV/s", "group": "B"})
        for e_ in st_cv["files"][-2:]:
            e_["ref"] = {"mode": "internal", "uid": e_["uid"], "lo": 0.5, "hi": 1.0, "compound": "ac2fc", "value": 0.49}
        st_cv.update(plot="cv", stack=True, e12_each=True, show_e12=True)
        t_cv.fill_table()
        spec_ = t_cv.figure_spec()
        resB = [r_ for r_ in t_cv.results if r_["group"] == "B"]
        assert len(resB) == 2 and all(abs(r_["e12"] - (-0.92 - (0.75 - 0.49))) < 0.004 for r_ in resB), resB
        resA = [r_ for r_ in t_cv.results if r_["group"] == "A"]
        assert len(resA) == 3 and all(abs(r_["e12"] - (-1.02 - (0.55 - 0.49))) < 0.004 for r_ in resA), resA
        cfg_, curves_, labels_, _ = spec_
        offs_ = [c_.get("off", 0) for c_ in curves_]
        assert offs_ == sorted(offs_, reverse=True) and offs_[0] > 0 == offs_[-1] and cfg_["y_axis"] == "no_numbers"
        assert len(labels_) == 5 and all(l_.get("y") is not None and "1/2" in l_["text"] for l_ in labels_), labels_
        t_cv.render()
        assert all(getattr(a_, "_anchor_y", None) is not None for a_ in t_cv.F.ann_items.values())
        # Peakpotentiale an einzelnen Peaks, y-Achse aus, Struktur je Messung (in Kurvenfarbe)
        st_cv.update(peak_labels="both", y_axis="off")
        st_cv["files"][0].update(structure=str(cdxml), struct_color=True)
        cfg_, curves_, labels_, _ = t_cv.figure_spec()
        assert len(labels_) == 10 and {l_["side"] for l_ in labels_} == {"above", "below"}, len(labels_)
        assert cfg_["y_axis"] == "off" and [i_["id"] for i_ in cfg_["images"]] == ["struktur_cv0"], cfg_["images"]
        t_cv.render()
        assert "image_struktur_cv0" in t_cv.F.axes_items
        assert any(c_ is st_cv["files"][0] and k_ == "structure" for c_, k_, _ in PRJ._path_slots({"cv": st_cv}))
        st_cv["files"][0].pop("structure")
        st_cv.update(peak_labels="e12", y_axis="auto")
        st_cv["show_bl"] = True                          # Basislinien zur Kontrolle: Hilfslinien, nicht in CSV/Hover
        cfg_, curves_, _, _ = t_cv.figure_spec()
        aux_ = [c_ for c_ in curves_ if c_.get("aux")]
        assert len(aux_) >= 5 * 3 and all(c_.get("label") is None for c_ in aux_), len(aux_)
        t_cv.render()
        assert len(t_cv.hover.curves) == len(curves_) - len(aux_), (len(t_cv.hover.curves), len(curves_))
        st_cv["show_bl"] = False
        st_cv.update(stack=False, e12_each=False)          # E½ ohne Stapeln: je Gruppe, nie über Gruppen gemittelt
        _, _, labels_, _ = t_cv.figure_spec()
        assert sorted(l_["text"].split(" =")[0] for l_ in labels_) == ["E$_{1/2}$ (A)", "E$_{1/2}$ (B)"], labels_
        assert {round(l_["lam"], 2) for l_ in labels_} == {-1.08, -1.18}, labels_
        st_cv.update(stack=True, e12_each=True)
        for plot_, n_series in (("dep", 2), ("rs", 4), ("loglog", 4), ("ratio", 1), ("trumpet", 4)):
            st_cv["plot"] = plot_
            cfg_, curves_, _, _ = t_cv.figure_spec()
            pts_ = [c_ for c_ in curves_ if c_.get("marker")]
            assert len(pts_) == n_series, (plot_, len(pts_))
            if plot_ == "dep":
                assert all(abs(y_ - 60) < 8 for c_ in pts_ for y_ in c_["y"]), [list(c_["y"]) for c_ in pts_]
                assert sorted(len(c_["x"]) for c_ in pts_) == [2, 3]
            if plot_ == "rs":
                assert len([c_ for c_ in curves_ if c_.get("ls") == "--"]) == 2      # Gerade nur ab 3 Scanraten (A)
            if plot_ == "trumpet":
                assert abs(min(c_["x"].min() for c_ in pts_) - np.log10(0.05)) < 1e-9
            if plot_ == "ratio":
                assert all(0.8 < y_ < 1.3 for c_ in pts_ for y_ in c_["y"]), [list(c_["y"]) for c_ in pts_]
            if plot_ == "loglog":
                fits_ = t_cv.rate_fits()
                assert list(fits_) == [("A", 1)] and abs(fits_[("A", 1)]["ll_a"]["slope"] - 0.5) < 0.05, fits_
                assert "log-log" in t_cv.lbl_res.text() or "log–log" in t_cv.lbl_res.text(), t_cv.lbl_res.text()
        files = t_cv.export_to(out / "cv_dep")
        assert all(f.exists() for f in files) and "group" in files[-1].read_text(encoding="utf-8-sig")
        assert "Couple 1 (A)" in files[-2].read_text(encoding="utf-8")
        st_cv["files"] = st_cv["files"][:1]
        assert t_cv.figure_spec() is None and "zwei Scanraten" in t_cv.note or "two scan rates" in t_cv.note
        st_cv.update(files=[], couples=[], ref_mode="none", plot="cv", stack=False, e12_each=False)
        w.cv_tab.fill_table()
        report.append("CV stacking, reference per file, rate plots: OK")

        # Spektroelektrochemie: BWSpec-Spektren + Gamry-Potential über die Uhrzeit, Nernst-Fit, Export
        import datetime as _dt
        secdir = out / "input" / "sec"
        secdir.mkdir(parents=True, exist_ok=True)
        t_start = _dt.datetime(2026, 1, 1, 12, 0, 0)
        T_ = np.arange(0.0, 1150.0, 2.0)
        E_run = -1.3 + 0.002 * T_                      # 2 mV/s
        write_dta(secdir / "sec.DTA", [(E_run, 1.0 + 0 * E_run)], 2.0, cycle_col=True, times=T_)  # Start 12:00:00
        pix = np.arange(0, 2048)
        lam_ = 200.0 + 0.35 * pix
        sec_files = []
        for j in range(0, 70):
            ts = j * 15.0 + 5.0
            e_ = -1.3 + 0.002 * ts
            ox = 1 / (1 + np.exp(-SECMOD.F_RT * (e_ + 0.80)))          # E½ = −0.80 V, n = 1
            A_ = (1 - ox) * 0.3 * np.exp(-((lam_ - 432) / 25) ** 2) + ox * 0.35 * np.exp(-((lam_ - 373) / 15) ** 2)
            hdr = ["File Version;BWSpec4.11_1", f"Date;{(t_start + _dt.timedelta(seconds=ts)):%Y-%m-%d %H:%M:%S}",
                   "coefs_a0;200", "coefs_a1;0.35", "coefs_a2;0", "coefs_a3;0", "sel_pixel_start;0",
                   "Pixel;Wavelength;Wavenumber;Raman Shift;Dark;Reference;Raw data #1;Dark Subtracted #1;%TR #1;"
                   "Absorbance #1;Irradiance (W/m^2/nm) #1"]
            rows_ = [f"{p_};   ;   ;   ;0;0;0;0;0;{a_:.5f};0;" for p_, a_ in zip(pix, A_)]
            f_ = secdir / f"s{j:03d}.txt"
            f_.write_text("\r\n".join(hdr + rows_), encoding="latin-1", newline="")
            sec_files.append(str(f_))
        (secdir / "LOG_s.txt").write_text("File Format;TXTR\r\nFile Description;Peak Monitor File\r\n",
                                         encoding="latin-1")
        assert not SECMOD.is_bwspec(secdir / "LOG_s.txt") and SECMOD.is_bwspec(sec_files[0])
        sp_ = [SECMOD.read_bwspec(f_) for f_ in sec_files]
        E_map = SECMOD.potentials(sp_, CVMOD.read_gamry(secdir / "sec.DTA"))
        assert np.allclose(E_map, -1.3 + 0.002 * (np.arange(70) * 15.0 + 5.0), atol=1e-3), E_map[:3]
        a432 = np.array([np.interp(432, s_["lam"], s_["A"]) for s_ in sp_])
        fit_ = SECMOD.nernst_fit(E_map, a432)
        assert abs(fit_["e12"] + 0.80) < 0.005 and abs(fit_["n"] - 1.0) < 0.1, fit_
        st_sec = w.sec_state
        st_sec.update(spec_files=sec_files, dta=str(secdir / "sec.DTA"), bands="373, 432", mode="spectra")
        files = w.sec_tab.export_to(out / "sec_spectra")
        assert all(f.exists() for f in files), files
        assert not any(f.name.endswith("_methods.txt") for f in files), files           # kein Fit, kein SI-Text
        st_sec.update(mode="trace", trace_lams="432")
        files = w.sec_tab.export_to(out / "sec_trace")
        assert all(f.exists() for f in files) and abs(w.sec_tab.info[432.0]["e12"] + 0.80) < 0.005
        meth_ = files[-1].read_text(encoding="utf-8")
        assert files[-1].name.endswith("_methods.txt") and "apparent" in meth_ and "grid search" in meth_ \
            and "λ = 432 nm: E1/2,app = −0.80" in meth_, meth_
        slots_ = PRJ._path_slots({"sec": st_sec})
        assert sum(1 for c_, k_, _ in slots_ if c_ is st_sec["spec_files"]) == 70 and \
            any(c_ is st_sec and k_ == "dta" for c_, k_, _ in slots_), "SEC-Dateien fehlen im Projekt"
        st_sec.update(mode="spectra")
        w.sec_state.update(cap_solvent="THF", cap_start="neutral", cap_end="dication")    # Unterschrift wie in der SI
        cap_s = w.sec_tab.si_caption("TEST-1")
        assert cap_s.startswith("UV-vis-SEC of <b>TEST-1</b> in THF.") and "Starting potential (vs. " in cap_s \
            and "End potential" in cap_s, cap_s
        # Zwischenspezies: Potential -> nächstliegendes Spektrum grün (Nr. 17: −0.81 V), über den grauen, mit Label
        assert "green" not in cap_s and cap_s.count("<b>") == 3, cap_s                      # ohne Eingabe wie bisher
        cv0_ = w.sec_tab.figure_spec()[1]
        assert [c_["color"] for c_ in cv0_ if c_["lw"] > 1] == [SECMOD.START_COLOR, SECMOD.END_COLOR]
        w.sec_state.update(mid="-0.80", cap_mid="radical cation")
        _, cv_, lb_, _ = w.sec_tab.figure_spec()
        hi_ = [c_ for c_ in cv_ if c_["lw"] > 1]
        assert [c_["color"] for c_ in hi_] == [SECMOD.START_COLOR, SECMOD.MID_COLORS[0][0], SECMOD.END_COLOR], hi_
        assert all(c_["z"] == 3 for c_ in cv_[-3:]) and cv_[-3:] == hi_, "hervorgehobene nicht oben"
        assert "0.81" in hi_[1]["label"] and np.allclose(hi_[1]["y"], np.interp(
            cv_[0]["x"], sp_[16]["lam"], sp_[16]["A"])), hi_[1]["label"]
        assert {l_["lam"]: l_["color"] for l_ in lb_} == {373.0: SECMOD.END_COLOR, 432.0: SECMOD.START_COLOR}
        cap_s = w.sec_tab.si_caption("TEST-1")
        assert ("Color coding: <b>neutral</b> (blue), <b>radical cation</b> (green), <b>dication</b> (red).") \
            in cap_s, cap_s
        w.sec_state.update(mid="-0.80, 0.30, -0.81", cap_mid="radical cation; X")      # Doppelte fallen weg
        hi_ = [c_ for c_ in w.sec_tab.figure_spec()[1] if c_["lw"] > 1]
        assert [c_["color"] for c_ in hi_].count(SECMOD.MID_COLORS[1][0]) == 1 and len(hi_) == 4
        assert "(green), <b>X</b> (orange), <b>dication</b> (red)" in w.sec_tab.si_caption("TEST-1")
        w.sec_state.update(dta=None, mid="35")                       # ohne Potentiale: Nummer des Spektrums
        hi_ = [c_ for c_ in w.sec_tab.figure_spec()[1] if c_["lw"] > 1]
        assert hi_[1]["color"] == SECMOD.MID_COLORS[0][0] and np.allclose(hi_[1]["y"], np.interp(
            hi_[1]["x"], sp_[34]["lam"], sp_[34]["A"]))
        w.sec_state.update(highlight=False)
        assert "green" not in w.sec_tab.si_caption("TEST-1")
        w.sec_state.update(dta=str(secdir / "sec.DTA"), highlight=True, mid="", cap_mid="")
        w.sec_state.update(cap_solvent="", cap_start="", cap_end="")
        report.append("spectroelectrochemistry: OK")

        # SEC im Overlay: Spektren nach Potential übernehmen, nur normiert (kein ε aus der Dünnschichtzelle)
        ov_, st_ov = w.ov_tab, w.ov_state
        saved_ov = {k: json.loads(json.dumps(st_ov[k])) for k in ("entries", "mode", "nlo", "nhi", "baseline")}
        st_ov.update(entries=[], mode="norm", nlo=340.0, nhi=460.0)
        ov_.add_sec_entries([0, 69])
        e0, e1 = st_ov["entries"]
        assert e0["kind"] == "sec" and abs(e0["E"] - (-1.3 + 0.002 * 5.0)) < 1e-3, e0
        assert "V vs. Ag/AgNO" in e0["label"] and e1["E"] > e0["E"], (e0["label"], e1["E"])
        ov_.render()
        cs_ = [c_ for c_ in ov_.F.curves if c_.get("label") in (e0["label"], e1["label"]) or "SEC" in str(c_.get("label"))]
        assert len(cs_) == 2, [c_.get("label") for c_ in ov_.F.curves]
        for c_, peak_ in zip(cs_, (432, 373)):          # reduziert: Bande bei 432 nm, oxidiert: 373 nm
            win_ = (c_["x"] >= 340) & (c_["x"] <= 460)
            assert abs(float(np.max(c_["val"][win_])) - 1.0) < 1e-6, "SEC nicht normiert"
            assert abs(float(c_["x"][win_][np.argmax(c_["val"][win_])]) - peak_) < 2, (peak_, c_["label"])
        assert c_["x"].min() >= SECMOD.LAM_RANGE[0] and c_["x"].max() <= SECMOD.LAM_RANGE[1]
        assert any(c_ is e0 and k_ == "file" for c_, k_, _ in PRJ._path_slots({"overlay": st_ov}))
        # Serie: EIN Eintrag, gemeinsamer Faktor (auf das Startspektrum) -> Verhältnisse bleiben erhalten
        st_ov.update(entries=[], baseline=False)
        ov_.add_sec_series(0, 69)
        ov_.render()
        assert len(ov_.F.curves) == 70 and len(st_ov["entries"]) == 1, len(ov_.F.curves)
        named_ = [c_["label"] for c_ in ov_.F.curves if c_.get("label")]
        assert len(named_) == 2 and all("V vs. Ag/AgNO" in n_ for n_ in named_), named_
        raw0, raw1 = (SECMOD.read_bwspec(sec_files[i_]) for i_ in (0, 69))
        start_ = next(c_ for c_ in ov_.F.curves if c_["color"] == SECMOD.START_COLOR)
        end_ = next(c_ for c_ in ov_.F.curves if c_["color"] == SECMOD.END_COLOR)
        win_ = (start_["x"] >= 340) & (start_["x"] <= 460)
        assert abs(float(start_["val"][win_].max()) - 1.0) < 1e-6
        r_raw = np.interp(373, raw1["lam"], raw1["A"]) / np.interp(432, raw0["lam"], raw0["A"])
        r_plot = np.interp(373, end_["x"], end_["val"]) / np.interp(432, start_["x"], start_["val"])
        assert abs(r_plot / r_raw - 1) < 0.02, (r_plot, r_raw)      # nicht einzeln normiert
        assert sum(1 for c_, k_, _ in PRJ._path_slots({"overlay": st_ov}) if c_ is st_ov["entries"][0]["files"]) == 70
        st_ov["entries"][0]["ends_only"] = True
        ov_.render()
        assert len(ov_.F.curves) == 2
        st_ov["mode"] = "eps"                           # ε-Modus: SEC wird ausgelassen, nicht falsch skaliert
        assert ov_.figure_spec() is None or all("SEC" not in str(c_.get("label")) for c_ in ov_.figure_spec()[1])
        st_ov.update(saved_ov)
        st_sec.update(spec_files=[], dta=None, mode="spectra", trace_lams="", bands="")
        report.append("SEC spectra in overlay: OK")

        # EPR: BES3T lesen (G -> mT), Dunkelspektrum, Basislinie, Ausblenden, Simulation anpassen, EasySpin-Parameter
        eprdir = out / "input" / "epr"
        eprdir.mkdir(parents=True, exist_ok=True)
        Bg = np.linspace(4400.0, 5100.0, 7001)                       # G
        signal = -np.gradient(np.exp(-((Bg - 4700) / 30) ** 2)) * 1000  # Ableitungslinie bei 470 mT
        drift = 0.05 * (Bg - 4400)                                     # lineare Drift der Basislinie

        def write_bes3t(stem, y, time="11:30:15", temp=100.0):
            (eprdir / f"{stem}.dsc").write_text("\n".join([
                "#DESC\t1.2", "DSRC\tEXP", "BSEQ\tLIT", "IKKF\tREAL", "XTYP\tIDX", "YTYP\tNODATA", "IRFMT\tD",
                f"XPTS\t{len(Bg)}", f"XMIN\t{Bg[0]}", f"XWID\t{Bg[-1] - Bg[0]}", "XUNI\t'G'", "#SPL\t1.2",
                "MWFQ\t9338553208.5", f"STMP\t{temp}", "B0MA\t0.00095", "MWPW\t0.01", "DATE\t08/22/25",
                f"TIME\t{time}"]),
                encoding="latin-1")
            y.astype("<f8").tofile(eprdir / f"{stem}.DTA")
        dark_bg = 3.0 * np.sin((Bg - 4400) / 200)                      # Untergrund, der auch dunkel da ist
        write_bes3t("dark", drift + dark_bg)
        write_bes3t("light", drift + dark_bg + signal)
        sim_B = np.linspace(440.0, 510.0, 1401)                        # mT, wie EasySpin
        sim_I = 0.5 * np.interp(sim_B, Bg / 10, signal)
        np.savetxt(eprdir / "sim.txt", np.column_stack([sim_B, sim_I]), delimiter="\t")
        (eprdir / "sim_params.txt").write_text("Sys.S = 1;\nSys.g = 2.0023;\nSys.D = [11100 450]; % MHz\n"
                                               "Exp.mwFreq = 9.3386;\n", encoding="utf-8")
        sp_e = EPRMOD.read_any(eprdir / "light.DTA")[0]
        assert abs(sp_e["B"][0] - 440.0) < 1e-9 and abs(sp_e["meta"]["mw_GHz"] - 9.3385535) < 1e-6, sp_e["meta"]
        assert abs(sp_e["meta"]["mod_mT"] - 0.95) < 1e-9 and abs(sp_e["meta"]["power_mW"] - 10) < 1e-9
        pr_ = EPRMOD.parse_easyspin((eprdir / "sim_params.txt").read_text(encoding="utf-8"))
        assert pr_["Sys.D"] == [11100.0, 450.0] and "0.3703 cm" in EPRMOD.zfs_summary(pr_)
        assert abs(EPRMOD.g_value(9.3386, 333.2) - 2.0024) < 1e-3
        t_e = w.epr_tab
        ed = t_e._new("exp", str(eprdir / "dark.dsc"))
        el = t_e._new("exp", str(eprdir / "light.dsc"))
        es = t_e._new("sim", str(eprdir / "sim.txt"))
        ed["visible"] = False
        el["dark"] = ed["uid"]
        es.update(fit_to=el["uid"], params=str(eprdir / "sim_params.txt"))
        w.epr_state.update(entries=[ed, el, es], bl_order=1, exclude="499-501", smooth=0.0, normalize=False,
                           stack=False, labels="Z@452, X@490", show_params=True)
        cache_ = {}
        Bp, Ip = t_e.processed(el, cache_)
        cleaned = Ip[np.isfinite(Ip)]
        ref_ = (signal / 1.0)[np.isfinite(Ip)]
        assert np.max(np.abs(cleaned - ref_)) < 0.02 * np.ptp(signal), "Dunkelspektrum/Basislinie"
        assert np.isnan(Ip[(Bp > 499.2) & (Bp < 500.8)]).all(), "Bereich nicht ausgeblendet"
        data_, _ = t_e.components([el, es], {}, [])
        Is = data_[es["uid"]][1]
        assert abs(np.nanmax(Is) / np.nanmax(Ip) - 1) < 0.05, "Simulation nicht an Messung angepasst"
        t_e.fill_table()
        files = w.epr_tab.export_to(out / "epr_test")
        assert all(f.exists() for f in files), files
        F_e = w.epr_tab.build(MplFigure(figsize=(8, 6), dpi=PREVIEW_DPI))
        assert any("cm" in a.get_text() for a in F_e.ann_items.values()), "D/E-Label fehlt"
        slots_ = PRJ._path_slots({"epr": w.epr_state})
        assert any(c_ is es and k_ == "params" for c_, k_, _ in slots_), "EPR-Parameter fehlen im Projekt"
        # Intensität gegen Zeit: Amplituden 25/50/75/100 % im Minutentakt nach dem Dunkelspektrum (11:30:15)
        series = [ed]
        for k_, amp in enumerate((0.25, 0.5, 0.75, 1.0), 1):
            write_bes3t(f"kin{k_}", drift + dark_bg + amp * signal, time=f"11:3{k_}:15")
            ek = t_e._new("exp", str(eprdir / f"kin{k_}.dsc"))
            ek["label"] = f"kin{k_}"
            series.append(ek)
        w.epr_state.update(entries=series, exclude="", labels="", show_params=False, mode="intensity",
                           int_window="455-485", int_method="double", int_x="time", normalize=True)
        t_e.fill_table()
        t_e.table.selectRow(1)
        series[1]["dark"] = ed["uid"]
        t_e.dark_for_all()
        files = w.epr_tab.export_to(out / "epr_kin")
        xs_ = [v[0] for v in t_e.series_values]
        ys_ = [v[1] for v in t_e.series_values]
        assert np.allclose(xs_, [1, 2, 3, 4]) and np.allclose(ys_, [0.25, 0.5, 0.75, 1.0], atol=0.02), (xs_, ys_)
        w.epr_state["int_method"] = "pp"
        w.epr_tab.export_to(out / "epr_kin_pp")
        assert np.allclose([v[1] for v in t_e.series_values], [0.25, 0.5, 0.75, 1.0], atol=0.02)
        txt_ = next(f for f in files if f.name.endswith("_kin4.txt"))
        arr_ = np.loadtxt(txt_, comments="%")
        assert arr_.shape[1] == 2 and abs(arr_[0, 0] - 440.0) < 1e-6 and "mwFreq" in txt_.read_text(encoding="utf-8")
        assert np.max(np.abs(arr_[:, 1] - signal)) < 0.02 * np.ptp(signal), "Export nicht verarbeitet"
        assert any(f.name.endswith("_intensity.csv") for f in files) and any(f.name.endswith("_data.csv") for f in files)
        w.epr_state.update(entries=[], labels="", exclude="", show_params=False, mode="spectra")
        report.append("EPR (BES3T, dark spectrum, simulation, intensity, export): OK")

        # EPR S = ½: TEMPO-artiges Spektrum mit Feldfehler +0.25 mT fitten -> Kalibrierung; zwei Spezies 70:30
        mw_ = 9.40
        Bt = np.linspace(330.0, 345.0, 6001)
        tempo_I = EPRMOD.iso_spectrum(Bt, mw_, 2.0055, [(1, 1.0, 1.55)], 0.15)
        noise = np.random.default_rng(0).standard_normal(len(Bt)) * 0.01 * np.ptp(tempo_I)
        np.savetxt(eprdir / "tempo-9p400000.txt", np.column_stack([Bt + 0.25, tempo_I + noise]),
                   header="B [mT]\tI [a.U.]", comments="", delimiter="\t")
        rad_I = EPRMOD.iso_spectrum(Bt, mw_, 2.0030, [(2, 0.5, 0.45)], 0.12)
        np.savetxt(eprdir / "radikal-9p400000.txt", np.column_stack([Bt + 0.25, rad_I]),
                   header="B [mT]\tI [a.U.]", comments="", delimiter="\t")
        et = t_e._new("exp", str(eprdir / "tempo-9p400000.txt"))
        er = t_e._new("exp", str(eprdir / "radikal-9p400000.txt"))
        assert abs(t_e.spectrum(et)["meta"]["mw_GHz"] - 9.4) < 1e-9
        w.epr_state.update(entries=[et, er], bl_order=-1, smooth=0.0, normalize=True, stack=False,
                           show_sum=True, show_residual=True, calib_uid=None, calib_g=2.0055)
        t_e.fill_table()
        t_e.table.selectRow(0)
        t_e.add_iso()
        ti = w.epr_state["entries"][-1]
        ti["nuclei"] = "1 14N 1.4"
        t_e.table.selectRow(2)
        t_e.fit_current()
        assert abs(ti["g"] - EPRMOD.g_value(mw_, EPRMOD.b_res(mw_, 2.0055) + 0.25)) < 2e-5, ti
        assert "1.55" in ti["nuclei"] and ti["fit"]["r2"] > 0.98, ti     # 1 % Rauschen
        w.epr_state["calib_uid"] = ti["uid"]
        t_e.fill_table()
        dB_, _ = t_e.calibration_shift()
        assert abs(dB_ - 0.25) < 0.003, dB_
        meth_ = t_e.methods_text()
        assert "Levenberg–Marquardt" in meth_ and "Field calibration: g_fit" in meth_ and "ΔB = +0.25" in meth_ \
            and "a(1 × I = 1) = 1.5" in meth_, meth_
        t_e.table.selectRow(1)
        t_e.add_iso()
        ri = w.epr_state["entries"][-1]
        ri["nuclei"] = "2 1H 0.4"
        t_e.table.selectRow(3)
        t_e.fit_current()
        assert abs(ri["g"] - 2.0030) < 5e-5 and abs(EPRMOD.parse_nuclei(ri["nuclei"])[0][2] - 0.45) < 0.005, ri
        w.epr_state.update(cap_species="TEST-1^{•+}", cap_solvent="THF", cap_conc="1 mM")   # Unterschrift wie in der SI
        cap_e = t_e.si_caption("TEST-1")
        assert cap_e.startswith("X-Band EPR spectr") and "<b>TEST-1<sup>•+</sup></b> in THF (1 mM)." in cap_e \
            and "2×<sup>1</sup>H: 12." in cap_e and "g = 2.003" in cap_e and "referenced to a standard" in cap_e, cap_e
        w.epr_state.update(cap_species="", cap_solvent="", cap_conc="")
        # zwei Spezies: Messung = 0.7·A + 0.3·B, beide als EasySpin-ASCII importiert
        A_ = EPRMOD.iso_spectrum(Bt, mw_, 2.0030, [(1, 1.0, 1.4)], 0.2)
        B_ = EPRMOD.iso_spectrum(Bt, mw_, 2.0100, [], 0.5)
        np.savetxt(eprdir / "mix-9p400000.txt", np.column_stack([Bt, 0.7 * A_ + 0.3 * B_]), delimiter="\t")
        np.savetxt(eprdir / "simA.txt", np.column_stack([Bt, A_]), delimiter="\t")
        np.savetxt(eprdir / "simB.txt", np.column_stack([Bt, B_]), delimiter="\t")
        em = t_e._new("exp", str(eprdir / "mix-9p400000.txt"))
        sa, sb = t_e._new("sim", str(eprdir / "simA.txt")), t_e._new("sim", str(eprdir / "simB.txt"))
        sa["fit_to"] = sb["fit_to"] = em["uid"]
        w.epr_state.update(entries=[em, sa, sb], calib_uid=None, bl_order=-1)
        info_ = []
        data_, extra_ = t_e.components([em, sa, sb], {}, info_)
        wa = np.nanmax(data_[sa["uid"]][1]) / np.nanmax(A_)
        wb = np.nanmax(data_[sb["uid"]][1]) / np.nanmax(B_)
        assert abs(wa - 0.7) < 1e-3 and abs(wb - 0.3) < 1e-3, (wa, wb)
        assert np.nanmax(np.abs(extra_[em["uid"]]["res"][1])) < 1e-6 * np.ptp(A_) and "sum" in extra_[em["uid"]]
        assert info_ and "%" in info_[0], info_
        files = w.epr_tab.export_to(out / "epr_species")
        assert all(f.exists() for f in files), files
        w.epr_state.update(entries=[], calib_uid=None, show_residual=False, bl_order=1)
        report.append("EPR S=1/2 fit, TEMPO calibration, multiple species: OK")

        # EPR-Zerfall bei konstanter Heizrate: Gauß-Intensität eines Merkmals, Arrhenius-Fit (E_a = 50 kJ/mol, ln A = 45)
        Ea_true, lnA_true = 50e3, 45.0
        temps = np.arange(100.0, 135.01, 0.5)
        beta_ = 0.5 / 120.0                                         # 0,5 K alle 2 min
        I_true = EPRMOD.decay_model(temps, np.exp(lnA_true), Ea_true, beta_)
        Bk = Bg                                                       # G, 440–510 mT
        kin_entries = []
        for j, (Tk, Ik) in enumerate(zip(temps, I_true)):
            y = -2000.0 * Ik * np.exp(-((Bk / 10 - 487.0) ** 2) / 9.0) / Tk + 0.3 * (Bk - 4400) / 10
            stem = f"temp{j:03d}"
            hh, mm = divmod(2 * j, 60)
            write_bes3t(stem, y, time=f"{12 + hh:02d}:{mm:02d}:00", temp=Tk)
            ek = t_e._new("exp", str(eprdir / f"{stem}.dsc"))
            kin_entries.append(ek)
        w.epr_state.update(entries=kin_entries, mode="intensity", int_method="gauss", int_window="470-510",
                           int_center=487.0, int_x="T", int_curie=True, ref_freq=0.0, T_offset=0.0, kin_fit=True,
                           kin_norm="lsq", bl_order=-1, smooth=0.0, normalize=True, calib_uid=None)
        t_e.fill_table()
        files = w.epr_tab.export_to(out / "epr_decay")
        r_ = t_e.kin_result
        assert r_ and abs(r_["Ea"] - Ea_true) / Ea_true < 0.02 and abs(r_["lnA"] - lnA_true) < 1.0, r_
        assert r_["ci_Ea"][0] < Ea_true < r_["ci_Ea"][1] and abs(r_["beta"] - beta_) < 1e-9, r_
        meth_ = next(f for f in files if f.name.endswith("_methods.txt")).read_text(encoding="utf-8")
        assert "Gaussian on a linear baseline" in meth_ and "t(0.975; n − 2)" in meth_ and "Ea = 5" in meth_ \
            and "multiplied by the temperature (I·T) before fitting" in meth_ and "± 0.01 kJ" in meth_ \
            and " × 10^" in meth_, meth_
        w.epr_state["kin_norm"] = "script"
        w.epr_tab.export_to(out / "epr_decay_script")
        assert abs(t_e.kin_result["Ea"] - Ea_true) / Ea_true < 0.02, t_e.kin_result
        assert "Arrhenius" in next(f for f in files if f.name.endswith("_intensity.csv")).read_text(encoding="utf-8-sig")
        w.epr_state.update(entries=[], mode="spectra", kin_fit=False, int_method="double", int_window="", int_center=0.0)
        report.append("EPR decay at constant heating rate (Arrhenius): OK")

        # Farbe der Lösung: Farbmetrik (Weiß, Grau, Gelb bei Blauabsorption) und Anzeige im ε-Tab
        L_ = COLOR.LAM
        assert COLOR.to_hex(COLOR.absorbance_to_rgb(0 * L_)[0]) == "#FFFFFF"
        assert COLOR.to_hex(COLOR.absorbance_to_rgb(0 * L_ + 0.3)[0]) == "#BCBCBC"     # T = 50 % -> sRGB 188
        r_, g_, b_ = COLOR.absorbance_to_rgb(np.exp(-0.5 * ((L_ - 450) / 30) ** 2))[0]
        assert r_ > 0.9 and g_ > 0.9 and b_ < 0.5, (r_, g_, b_)                        # absorbiert Blau -> Gelb
        w.glob.update(color_show=True, color_d=1.0)
        w.series["TEST-A"].pop("color_c", None)
        w.apply_globals()
        w.render()
        assert w.lbl_color.isVisibleTo(w) and "#" in w.lbl_color_info.text(), w.lbl_color_info.text()
        hx1 = w.lbl_color_info.text().split()[0]
        c_max = max(s_["c"] for s_ in w.groups["TEST-A"])
        assert f"c = {c_max:g}" in w.lbl_color_info.text(), w.lbl_color_info.text()
        w.sp_color_c.setValue(4 * c_max)                    # konzentrierter -> dunkler
        hx2 = w.lbl_color_info.text().split()[0]
        lum = lambda h: sum(int(h[i:i + 2], 16) for i in (1, 3, 5))
        assert lum(hx2) < lum(hx1) and w.series["TEST-A"]["color_c"] == 4 * c_max, (hx1, hx2)
        w.glob["color_show"] = False
        w.series["TEST-A"].pop("color_c", None)
        w.apply_globals()
        assert not w.lbl_color.isVisibleTo(w)
        report.append("solution colour: OK")

        # Layout ε-Tab: Vorschau statt Leer-Zustand, eingeklappte Bereiche melden Abweichungen vom Standard
        assert w.preview_stack.currentIndex() == 1
        w.sec_eval.btn.setChecked(False)
        assert not w.sec_eval.body.isVisibleTo(w)
        r2_old = w.sp_r2.value()
        w.sp_r2.setValue(DEFAULT_GLOB["min_r2"])
        w.sp_min.setValue(DEFAULT_GLOB["min_points"])
        w.cb_shoulders.setChecked(DEFAULT_GLOB["shoulders"])
        assert not w.sec_eval.changed
        w.sp_r2.setValue(0.95)
        assert w.sec_eval.changed and T("adv_changed") in w.sec_eval.btn.text()
        w.sp_r2.setValue(r2_old)
        w.sec_eval.btn.setChecked(True)
        assert w.sec_eval.body.isVisibleTo(w)
        labels = [lb for lb in w.eps_page.findChildren(QLabel)
                  if lb.text() in (T("molar_mass_short"), T("ymax"))]
        assert len({lb.minimumWidth() for lb in labels}) == 1 and labels[0].minimumWidth() > 0
        # „Alles auf Standard“ im ε-Tab: Einstellungen zurück, Schichtdicke und Exportgröße bleiben
        if w.current:
            w.glob.update(cutoff=2.0, ymax=2.5, path_length=0.1, shoulders=True, size="custom")
            w.series[w.current]["band_hidden"] = [600]
            w.reset_defaults()
            assert w.glob["cutoff"] == 1.0 and w.glob["ymax"] == 1.0 and not w.glob["shoulders"]
            assert w.glob["path_length"] == 0.1 and w.glob["size"] == "custom" and abs(w.sp_cut.value() - 1.0) < 1e-9
            assert w.series[w.current]["band_hidden"] == []
            w.glob.update(path_length=1.0, size="half_a4")
            w.apply_globals()
        report.append("eps tab layout (sections, empty state): OK")
        # Anleitung in leeren Tabs: Texte in beiden Sprachen, Link zeigt auf eine vorhandene README-Überschrift
        from uvvis_i18n import TEXTS as TEXTS_
        heads_ = {lang_: [ln_[3:].strip() for ln_ in (Path(__file__).parent / f_).read_text(encoding="utf-8").splitlines()
                          if ln_.startswith("## ")]
                  for lang_, f_ in (("de", "README.de.md"), ("en", "README.md"))
                  if (Path(__file__).parent / f_).exists()}           # im gebauten Programm fehlt die README
        for key_, (de_, en_) in uvvis_tabs.GUIDE_HEAD.items():
            assert not heads_ or (de_ in heads_["de"] and en_ in heads_["en"]), (key_, de_, en_)
            for part_ in ("title", "files", "steps"):
                assert T(f"guide_{key_}_{part_}") != f"guide_{key_}_{part_}", (key_, part_)
        for lang_ in heads_:                                       # Startbildschirm/Hilfe: Schnellstart-Link
            q_ = TEXTS_["readme_quickstart"][0 if lang_ == "de" else 1]
            assert q_ in heads_[lang_], (lang_, q_)
        assert uvvis_tabs.gh_anchor("IR + EA tab") == "ir--ea-tab"
        assert uvvis_tabs.gh_anchor("CV-Tab (Cyclovoltammetrie)") == "cv-tab-cyclovoltammetrie"
        guide_ = w.cv_tab.preview_stack.widget(0).findChild(uvvis_tabs.EmptyGuide)
        assert guide_ is not None and T("guide_cv_title") in guide_.label.text()
        w.resize(1440, 900)
        w.show()
        for tab_, key_ in ((w.syn_tab, "syn"), (w.nmr_tab, "nmr"), (w.cv_tab, "cv")):   # überall gleich und mittig
            w.main_tabs.setCurrentWidget(tab_)
            pump(0.1)
            assert tab_.preview_stack.currentIndex() == 0, key_
            g_ = tab_.preview_stack.widget(0)
            g_ = g_ if isinstance(g_, uvvis_tabs.EmptyGuide) else g_.findChild(uvvis_tabs.EmptyGuide)
            assert g_ is not None and T(f"guide_{key_}_title") in g_.label.text(), key_
            b_ = g_.findChild(QPushButton)
            assert b_ is not None, key_
            top_, bottom_ = g_.label.geometry().top(), g_.height() - b_.geometry().bottom()
            assert abs(top_ - bottom_) <= 12, (key_, top_, bottom_, g_.size())
            mid_ = (g_.label.geometry().center().x(), b_.geometry().center().x())
            assert all(abs(m_ - g_.width() / 2) <= 12 for m_ in mid_), (key_, mid_, g_.width())
        # Meldungen als Einblendung: keine Statusleiste unter den Fußleisten, Knöpfe oben und unten gleich weit weg
        w.statusBar().showMessage("selftest", 5000)
        assert w._toast.isVisible() and w._toast.text() == "selftest"     # sofort (andere Meldungen folgen beim pump)
        pump(0.1)
        bar_ = w.cv_tab.findChild(QFrame, "bottomBar")
        y_ = bar_.mapTo(w, bar_.rect().bottomLeft()).y()
        assert y_ >= w.height() - 3, (y_, w.height())
        btn_ = max((c_ for c_ in bar_.findChildren(QPushButton) if c_.isVisible()), key=lambda c_: c_.height())
        gap_ = (btn_.geometry().top(), bar_.height() - btn_.geometry().bottom() - 1)
        assert abs(gap_[0] - gap_[1]) <= 2, gap_
        w.statusBar().clearMessage()                                      # Signal kommt sofort (ohne pump: Autosave)
        assert not w._toast.isVisible(), (w._toast.text(), w.statusBar().currentMessage())
        report.append("empty-tab guides: OK")
        # Navigation: Gruppen oben statt Reiter; Knopf wechselt den Tab, Tabwechsel markiert den Knopf
        assert len(w.nav_buttons) == 12 and w.main_tabs.tabBar().isHidden()
        i_cv_ = w._tab_keys.index("cv")
        w.nav_buttons[i_cv_].click()
        assert w.main_tabs.currentWidget() is w.cv_tab
        w.main_tabs.setCurrentWidget(w.ov_tab)
        assert w.nav_buttons[w._tab_keys.index("overlay")].isChecked()
        # Einstellungen des gewählten Eintrags: ohne Auswahl ausgeblendet, stattdessen ein Hinweis
        w.ov_tab.table.clearSelection()
        w.ov_tab.table.setCurrentCell(-1, -1)
        w.ov_tab.show_details()
        assert w.ov_tab.det.isHidden() == (not w.ov_tab.det.isEnabled())
        assert w.toolbar.widgetForAction(w.act_si).property("accent")
        # Mausrad über einem Zahlenfeld in der linken Spalte: ohne Fokus scrollt die Spalte, der Wert bleibt
        from PySide6.QtCore import QPoint, QPointF
        from PySide6.QtGui import QWheelEvent
        sa_ = QScrollArea()
        inner_ = QWidget()
        lv_ = QVBoxLayout(inner_)
        spins_ = [QSpinBox(maximum=100, value=50) for _ in range(40)]
        for sp_ in spins_:
            lv_.addWidget(sp_)
        sa_.setWidget(inner_)
        sa_.setWidgetResizable(True)
        sa_.resize(200, 200)
        sa_.show()
        pump(0.2)
        sp0_ = spins_[2]

        def wheel_(target_):
            return QWheelEvent(QPointF(5, 5), QPointF(target_.mapToGlobal(QPoint(5, 5))), QPoint(0, 0), QPoint(0, -120),
                               Qt.NoButton, Qt.NoModifier, Qt.NoScrollPhase, False)
        QApplication.sendEvent(sp0_, wheel_(sp0_))
        assert sp0_.value() == 50 and sa_.verticalScrollBar().value() > 0, (sp0_.value(), sa_.verticalScrollBar().value())
        assert sp0_.focusPolicy() == Qt.StrongFocus
        sp0_.setFocus()
        QApplication.sendEvent(sp0_, wheel_(sp0_))
        assert sp0_.value() != 50 or not sp0_.hasFocus()      # mit Fokus wie gewohnt (offscreen evtl. ohne Fokus)
        sa_.close()
        report.append("navigation, detail hint, SI button, wheel guard: OK")
        # Synthese/NMR mit Fußleiste wie die anderen Tabs; „Aus Schema“ nur mit Schema; Exportmeldung mit Ordner
        assert w.syn_tab.btn_export.property("accent") and w.nmr_tab.btn_export.property("accent")
        assert w.syn_tab.btn_from_scheme.isEnabled() == (w.syn_tab.scheme_parts() is not None)
        w.exported([proj / "exports"])
        assert w.last_export == [str(proj / "exports")]
        # linke Spalte: Karten eingeklappt mit „(geändert)“, Ergebnisse als Reiter neben der Abbildung
        fc_ = next(f_ for f_ in w.cv_tab.folds if f_.group is w.cv_tab.grp_common)
        fc_.set_open(False, remember=False)
        assert fc_.body.isHidden() and fc_.group.title().startswith("▶")
        w.cv_state["size"] = "custom"
        fc_.refresh()
        assert T("adv_changed") in fc_.group.title(), fc_.group.title()
        w.cv_state["size"] = "half_a4"
        fc_.refresh()
        assert T("adv_changed") not in fc_.group.title()
        fc_.set_open(True, remember=False)
        assert not fc_.body.isHidden()
        assert w.cv_tab.right_tabs is not None and w.cv_tab.right_tabs.tabText(0) == T("tab_figure")
        assert [w.xray_tab.right_tabs.tabText(i_) for i_ in range(3)] == [T("tab_figure"), T("xr_caption"), T("xr_tables")]
        assert w.dft_tab.right_tabs.count() == 2 and w.hrms_tab.right_tabs is None
        # Ergebnisseiten überall gleich: ε neben der Abbildung, Karte füllt die Seite, Hauptknopf blau rechts
        assert [w.tabs.tabText(i_) for i_ in range(w.tabs.count())] == [T("tab_figure"), T("tab_results"),
                                                                       T("tab_si_line")]
        for rt_, j_ in ((w.cv_tab.right_tabs, 1), (w.xray_tab.right_tabs, 2), (w.dft_tab.right_tabs, 1), (w.tabs, 2)):
            card_ = rt_.widget(j_).widget().findChild(QGroupBox, "resultCard")
            assert card_ is not None and card_.title() == "", (rt_.tabText(j_),)
            acc_ = [b_ for b_ in card_.findChildren(QPushButton) if b_.property("accent")]
            assert len(acc_) == 1 and not acc_[0].icon().isNull(), (rt_.tabText(j_), len(acc_))
        # Raster: jedes sichtbare Zahlenfeld der linken Spalten genau STYLE.FIELD_W breit (das Stylesheet darf die
        # feste Breite nicht nachträglich aufheben – so liefen die Felder früher auseinander)
        w.show()
        for i_ in range(w.main_tabs.count()):
            w.main_tabs.setCurrentIndex(i_)
            pump(0.05)
            page_ = w.main_tabs.widget(i_)
            odd_ = [sp_.width() for sp_ in page_.findChildren(QAbstractSpinBox)
                    if sp_.isVisibleTo(page_) and sp_.property("fieldw") and sp_.width() != STYLE.FIELD_W]
            assert not odd_, (w._tab_keys[i_], odd_)
        # linke Spalte nie abgeschnitten: Inhalt passt in die sichtbare Breite, auch auf kleinerem Bildschirm
        for size_ in ((1600, 1000), (1280, 800)):
            w.resize(*size_)
            for i_ in range(w.main_tabs.count()):
                w.main_tabs.setCurrentIndex(i_)
                pump(0.1)
                page_ = w.main_tabs.widget(i_)
                cf_ = getattr(page_, "column_fit", None) or (w.column_fit if w._tab_keys[i_] == "eps" else None)
                assert cf_ is not None, w._tab_keys[i_]
                cf_.fit()
                pump(0.05)
                need_ = cf_.content.minimumSizeHint().width()
                have_ = cf_.scroll.viewport().width()
                assert need_ <= have_ + 1, (w._tab_keys[i_], size_, need_, have_)
        # Neuaufbau (Projekt öffnen): ausstehende Anpassung einer gelöschten Spalte darf nicht abstürzen
        sp_ = QSplitter()
        col_, sc_ = QWidget(), QScrollArea()
        cont_ = QWidget()
        sc_.setWidget(cont_)
        QVBoxLayout(col_).addWidget(sc_)
        sp_.addWidget(col_)
        sp_.addWidget(QWidget())
        cf_ = STYLE.fit_column(cont_, sc_, None, col_, sp_)
        QApplication.sendEvent(cont_, QEvent(QEvent.LayoutRequest))       # Anpassung steht aus
        hook_, errs_ = sys.excepthook, []
        sys.excepthook = lambda *a: errs_.append(a[1])
        try:
            import shiboken6 as _shb
            _shb.delete(sp_)                                                  # wie der Neuaufbau der Tabs
            pump(0.1)
            cf_.fit()
            cf_._moved(10, 1)
        finally:
            sys.excepthook = hook_
        assert not errs_, errs_
        report.append("left column (folds, changed mark, result tabs, field grid, never clipped): OK")

        # Erscheinungsbild: Fusion, Palette je Modus, Akzentbuttons, Symbole; Ergebnisse bleiben erhalten
        w.render()
        n_rows = w.table.rowCount()
        for mode in ("dark", "light"):
            w.switch_theme(mode)
            pal = QApplication.instance().palette()
            assert STYLE.base_style == "fusion"
            assert pal.color(pal.ColorRole.Window).name().upper() == STYLE.COLORS[mode]["window"], mode
            assert w.btn_export.property("accent") and w.btn_open.property("accent")
            assert not w.btn_export.icon().isNull(), "icon"
            assert w.table.rowCount() == n_rows, (w.table.rowCount(), n_rows)
            css_ = QApplication.instance().styleSheet()       # Checkboxen mit eigenem Rand (dunkel sonst unsichtbar)
            assert "QCheckBox::indicator" in css_ and STYLE.COLORS[mode]["check_border"] in css_, mode
        assert w.findChildren(QFrame, "card") and not [g_ for g_ in w.eps_page.findChildren(QGroupBox)
                                                        if g_.objectName() != "resultCard"]
        w.switch_theme("system")
        # Werkzeugleiste (nach Neuaufbau nur eine; Projekt, Undo, Verbindung, SI) und Listenknöpfe mit Symbol
        bars = w.findChildren(uvvis_tabs.QWidget, "mainToolbar")
        acts = [a for a in w.toolbar.actions() if not a.isSeparator()]
        assert len(bars) == 1 and len(acts) == 8 and all(not a.icon().isNull() for a in acts), (len(bars), len(acts))
        opened_ = []                                         # Projektordner im Dateimanager (hier nur abgefangen)
        orig_open_ = QDesktopServices.openUrl
        QDesktopServices.openUrl = staticmethod(lambda u: opened_.append(u.toLocalFile()) or True)
        try:
            assert w.act_folder in acts and w.act_folder.isEnabled()
            w.act_folder.trigger()
        finally:
            QDesktopServices.openUrl = orig_open_
        assert opened_ and Path(opened_[0]) == Path(w.project_dir), opened_
        arrows = [b for b in w.epr_tab.findChildren(QPushButton) if b.toolTip() == T("move_up")]
        assert arrows and arrows[0].text() == "" and not arrows[0].icon().isNull()
        report.append("appearance (Fusion, light/dark, toolbar): OK")

        # Umbenennungen UVVisTool -> Spectravolt (1.16) -> Elucida (2.0): Einstellungen/Datenordner einmalig übernehmen
        old_s = QSettings(str(out / "old_settings.ini"), QSettings.IniFormat)
        old_s.setValue("lang", "en")
        old_s.setValue("recent", [str(out / "p1"), str(out / "p2")])
        new_s = QSettings(str(out / "new_settings.ini"), QSettings.IniFormat)
        new_s.clear()
        assert migrate_settings(old_s, new_s) and new_s.value("lang") == "en" and len(new_s.value("recent")) == 2
        new_s.setValue("lang", "de")
        assert not migrate_settings(old_s, new_s) and new_s.value("lang") == "de"   # nur einmal, nichts überschreiben
        base_ = out / "appdata"
        shutil.rmtree(base_, ignore_errors=True)
        (base_ / "UVVisTool" / "models").mkdir(parents=True)
        (base_ / "UVVisTool" / "models" / "m.onnx").write_bytes(b"x")
        d_ = IMG.app_data_dir(base_)
        assert d_.name == "Elucida" and (d_ / "models" / "m.onnx").exists() and not (base_ / "UVVisTool").exists()
        shutil.rmtree(base_, ignore_errors=True)                 # Spectravolt hat Vorrang vor UVVisTool
        for n_, f_ in (("Spectravolt", "neu.onnx"), ("UVVisTool", "alt.onnx")):
            (base_ / n_ / "models").mkdir(parents=True)
            (base_ / n_ / "models" / f_).write_bytes(b"x")
        d_ = IMG.app_data_dir(base_)
        assert (d_ / "models" / "neu.onnx").exists() and not (base_ / "Spectravolt").exists(), list(d_.rglob("*"))
        assert APP_NAME == "Elucida" and OLD_APP_NAMES[0] == "Spectravolt" and UPD.REPO.endswith("/elucida")
        report.append("rename (settings + data folder migration): OK")

        # Verbindung auf Projektebene + Tab IR + MS (künstliche OPUS-, PDF- und Word-Dateien)
        import struct as _st
        import zipfile as _zf
        w.set_compound({"name": "TEST-1", "structure": str(cdxml), **compound_from_structure(cdxml)})
        assert w.compound["formula"] and w.compound["mw"] and "TEST-1" in w.windowTitle(), w.compound
        import time as _time
        t_end_ = _time.monotonic() + 10                     # Strukturen werden im Hintergrund geladen
        while not all(st_.get("structure") for st_ in w.series.values()) and _time.monotonic() < t_end_:
            pump(0.1)
        assert all(st_.get("structure") for st_ in w.series.values()), ("Verbindung nicht als Vorgabe übernommen", {g_: st_.get("structure") for g_, st_ in w.series.items()}, w.compound)
        irdir = out / "input" / "irms"
        irdir.mkdir(parents=True, exist_ok=True)
        wn_ = np.linspace(3999.0, 400.0, 1800)
        bands_ = [(2961.2, 25, 12), (1678.4, 60, 8), (1501.3, 40, 6), (1105.0, 15, 10), (758.1, 55, 6), (3350.0, 20, 160)]
        T_ = 1.0 - sum(dp / 100 * np.exp(-((wn_ - c) / wd) ** 2) for c, dp, wd in bands_)
        def opus_par(pairs):
            out_ = b""
            for nm, typ, val in pairs:
                raw = _st.pack("<i", val) if typ == 0 else _st.pack("<d", val)
                out_ += nm.encode() + b"\0" + _st.pack("<hh", typ, len(raw) // 2) + raw
            return out_ + b"END\0" + _st.pack("<hh", 0, 0)
        par = opus_par([("NPT", 0, len(wn_)), ("FXV", 1, 3999.0), ("LXV", 1, 400.0), ("CSF", 1, 1.0)])
        dat = T_.astype("<f4").tobytes()
        off_par, off_dat = 24 + 12 * 2, 24 + 12 * 2 + len(par)
        head = b"\x0a\x0a\xfe\xfe" + _st.pack("<d", 920622.0) + _st.pack("<iii", 24, 2, 2)   # Version, Verz., max, n
        dirs = _st.pack("<BBBBii", 31, 20, 0, 0, len(par) // 4, off_par) + _st.pack("<BBBBii", 15, 20, 0, 0, len(wn_), off_dat)
        (irdir / "test.0").write_bytes(head + dirs + par + dat)
        d_ir = IRMS.read_opus(irdir / "test.0")
        assert d_ir["kind"] == "T" and abs(d_ir["y"].max() - 100) < 1 and IRMS.is_opus(irdir / "test.0")
        pk_ = IRMS.pick_peaks(d_ir["wn"], d_ir["y"], "T")
        got_ = [round(p_["wn"]) for p_ in pk_]
        assert all(any(abs(g_ - c) <= 1 for g_ in got_) for c, _, _ in bands_), got_
        br_ = next(p_ for p_ in pk_ if abs(p_["wn"] - 3350) < 5)
        assert IRMS.intensity_tag(br_).startswith("br"), IRMS.intensity_tag(br_)
        doc_ = pymupdf.open()
        pg_ = doc_.new_page()
        for i_, v_ in enumerate(["2961.20", "1678.41", "1501.33", "500", "1000", "Transmittance [%]", "80"]):
            pg_.insert_text((50, 60 + 20 * i_), v_)
        doc_.save(str(irdir / "test_ir.pdf"))
        assert IRMS.read_ir_pdf(irdir / "test_ir.pdf") == [2961.2, 1678.41, 1501.33]
        def emf(texts):
            out_ = b""
            for t_ in texts:
                body = t_.encode("utf-16-le")
                size = 76 + len(body) + (-(76 + len(body)) % 4)
                rec = bytearray(size)
                _st.pack_into("<II", rec, 0, 84, size)
                _st.pack_into("<II", rec, 44, len(t_), 76)
                rec[76:76 + len(body)] = body
                out_ += bytes(rec)
            return out_
        xml_ = "".join(f"<w:p><w:r><w:t>{t_}</w:t></w:r></w:p>" for t_ in
                       ("sample name:", "TEST-1", "molecular formula:", "C18H20N2", "exact mass:", "264.1626",
                        "method of Ionisation:", "ESI ☒ APCI ☐ HESI ☐"))
        with _zf.ZipFile(irdir / "ms.docx", "w") as z_:
            z_.writestr("word/document.xml", f"<w:document><w:body>{xml_}</w:body></w:document>")
            z_.writestr("word/media/image1.emf", emf(["Chemical Formula:", "C", "18", "Exact Mass:", "264,1626"]))
            z_.writestr("word/media/image2.emf", emf(["265.1703", "266.1735", "x.d: +MS, 0.1-0.4min", "m/z"]))
        rep_ = IRMS.read_ms_docx(irdir / "ms.docx")
        assert rep_["formula"] == "C18H20N2" and rep_["mz"] == [265.1703, 266.1735] and rep_["polarity"] == "+", rep_
        plain_, html_, calc_ = IRMS.ms_line("C18H20N2", "M+H", 265.1703)
        assert plain_ == "HR-MS-ESI(+) calc. C18H21N2+ [M+H]+ 265.1699; found 265.1703." and "<sub>21</sub>" in html_
        assert abs(IRMS.ion("C18H20N2", "M+Na")[2] - 287.1519) < 1e-4 and IRMS.ion("C18H20N2", "M-H")[1] == -1
        t_ir = w.irms_tab
        t_ms = w.hrms_tab
        assert w.main_tabs.indexOf(t_ir) + 1 == w.main_tabs.indexOf(t_ms)      # IR + EA vor HRMS (wie die SI)
        t_ir.load_ir(str(irdir / "test.0"))
        t_ms.load_ms(str(irdir / "ms.docx"))
        w.compound["formula"] = "C18H20N2"
        t_ir.render()
        t_ms.render()
        assert t_ir._ir_line[0].startswith("IR [cm-1]: ṽ = 3350, 2961, 1678, 1501, 1105, 758"), t_ir._ir_line[0]
        assert t_ms._ms_line[0].endswith("265.1699; found 265.1703."), t_ms._ms_line[0]
        assert t_ms.F is None                                    # ohne Rohdaten kein Spektrum
        w.irms_state["ir_off"] = [1105]
        w.irms_state["with_int"] = True
        t_ir.render()
        assert "1105" not in t_ir._ir_line[0] and "(br " in t_ir._ir_line[0], t_ir._ir_line[0]
        w.compound["formula"] = "C18H22N2"                 # falsche Verbindung -> Warnung
        t_ms.render()
        assert "⚠" in t_ms.lbl_ms_info.text(), t_ms.lbl_ms_info.text()
        files = t_ir.export_to(out / "irms") + t_ms.export_to(out / "hrms")
        assert sum(f.name.endswith("_SI.txt") for f in files) == 2 and all(f.exists() for f in files)
        slots_ = PRJ._path_slots(w.get_state())
        assert any(k_ == "ir_file" for _, k_, _ in slots_) and any(k_ == "ms_file" for _, k_, _ in slots_)
        # ChemDraw-Objekt im Bericht: Abgleich und Übernahme der Verbindung (künstliche Struktur, kein CDX-Schreiber)
        from rdkit import Chem as _Chem
        rep_ = t_ms.ms_report()
        rep_["structure"] = {"formula": "C18H20N2", "mw": 264.37, "n_fragments": 1,
                             "cdxml": _Chem.MolToCDXMLBlock(_Chem.MolFromSmiles("CC(C)(C)c1ccc2nc(-c3ccccc3)n(C)c2c1"))
                             if hasattr(_Chem, "MolToCDXMLBlock") else None}          # macOS Intel: RDKit 2025.9
        w.compound = {}
        t_ms.render()
        assert t_ms.formula() == ("C18H20N2", "report_struct") == t_ir.formula(), t_ms.formula()
        assert not t_ms.btn_take.isHidden()
        rep_["formula"] = "C18H21N2"                       # Formelfeld vertippt -> Hinweis
        t_ms.render()
        assert "⚠" in t_ms.lbl_ms_info.text(), t_ms.lbl_ms_info.text()
        t_ms.take_compound()
        assert w.compound["formula"] == "C18H20N2" and w.compound["name"] == "TEST-1", w.compound
        if rep_["structure"]["cdxml"]:
            assert Path(w.compound["structure"]).exists() and "aus_MS-Bericht" in Path(w.compound["structure"]).name
        w.irms_state.update(ir_file=None, ir_pdf=None, ir_off=[], with_int=False)
        w.hrms_state.update(ms_file=None)
        w.compound = {}
        # alte Projekte (bis 1.20): MS-Teil im Zustand von IR + MS wird in den Tab HRMS übernommen
        st_old = w.get_state()
        st_old.pop("hrms")
        st_old["irms"].update(ms_file=str(irdir / "ms.docx"), adduct="M+Na")
        w.set_state(st_old)
        assert w.hrms_state["ms_file"] and w.hrms_state["adduct"] == "M+Na", w.hrms_state
        t_ir, t_ms = w.irms_tab, w.hrms_tab
        w.hrms_state.update(ms_file=None, adduct="M+H")
        w.irms_state.pop("ms_file", None)
        report.append("project compound + IR/HRMS lines: OK")

        # Bruker-Rohdaten (.d): künstlicher LC-Lauf, m/z über die Umkehrung der Kalibrierung in Flugzeit-Indizes
        import sqlite3 as _sq
        from matplotlib.backend_bases import MouseEvent
        ddir = irdir / "test.d"
        ddir.mkdir(exist_ok=True)
        cal_ = (24790.0, 0.2, 309.0, 158509.1, 2.54e-4, 0.0, -0.0011)   # Delay, Takt, c0, c1, c2, –, c6
        def to_idx(m):
            m1 = np.asarray(m, float) + cal_[6]
            return (cal_[2] + np.sqrt(1e12 * m1 / cal_[3]) + cal_[4] * m1 - cal_[0]) / cal_[1]
        calc_ = IRMS.ion("C18H20N2", "M+H")[2]
        rts_ = np.arange(0.0, 120.0, 10.0)
        offs_ = 1e-4 * np.array([3, -2, 4, -3, 2, 1, -1, 3, -4, 2, 0, 1])            # Streuung je Scan
        amp_ = 1e5 * np.exp(-0.5 * ((rts_ - 60) / 12) ** 2)
        baf_, rows_ = bytearray(b"\0" * 64), []
        for i_, rt_ in enumerate(rts_):
            mz_ = np.array([180.0634, calc_ + offs_[i_], calc_ + 1.00335, 400.1234])
            in_ = np.array([5e3, amp_[i_], 0.2 * amp_[i_], 8e3], "<f4")
            n_ = len(mz_)
            body = (_st.pack("<4I", n_, 16, 16 + 8 * n_, 16 + 12 * n_) + to_idx(mz_).astype("<f8").tobytes()
                    + in_.tobytes() + np.full(n_, 0.01, "<f4").tobytes())
            head_ = bytearray(52)
            _st.pack_into("<III", head_, 0, 52 + len(body), IRMS.BAF_LINE_MAGIC, 52)
            rows_.append((i_ + 1, float(rt_), 1, (0x10 << 56) | len(baf_), 0))
            baf_ += head_ + body
        rows_.append((99, 61.0, 1, (0x10 << 56) | 64, 0))                            # MS2 -> wird übersprungen
        (ddir / "analysis.baf").write_bytes(bytes(baf_))
        (ddir / "readme.txt").write_text("nicht ins Projekt")
        tf_ = _st.pack("<i7d", 10, *cal_)
        tf_ += b"\0" * (112 - len(tf_))
        (ddir / "analysis.sqlite").unlink(missing_ok=True)       # Selbsttest-Ordner kann wiederverwendet werden
        con_ = _sq.connect(ddir / "analysis.sqlite")
        con_.executescript("create table Properties (Key text, Value text); create table AcquisitionKeys "
                           "(Id integer, Polarity integer); create table Spectra (Id integer, Rt real, TransformatorId "
                           "integer, LineIndexId integer, AcquisitionKey integer); create table Steps (TargetSpectrum "
                           "integer); create table Transformators (Id integer, Blob blob);")
        con_.executemany("insert into Properties values (?, ?)", [("SchemaType", "Baf2Sql"), ("SampleName", "TEST-1"),
                                                                   ("InstrumentName", "timsTOF test")])
        con_.execute("insert into AcquisitionKeys values (0, 0)")
        con_.executemany("insert into Spectra values (?, ?, 1, ?, ?)", [(a, b, c, d) for a, b, _, c, d in rows_])
        con_.execute("insert into Steps values (99)")
        con_.execute("insert into Transformators values (1, ?)", (tf_,))
        con_.commit()
        con_.close()
        raw_ = IRMS.read_bruker_d(ddir)
        assert IRMS.is_bruker_d(ddir) and len(raw_["spectra"]) == 12 and raw_["polarity"] == "+", raw_["polarity"]
        assert np.abs(raw_["spectra"][6][1][1] - (calc_ + offs_[6])) < 1e-7          # Kalibrierung zurückgerechnet
        lc_ = IRMS.lc_found(raw_["spectra"], calc_)
        sel_ = slice(5, 8)                                                         # 50/60/70 s: >= 50 % des Maximums
        exp_ = float(np.average(calc_ + offs_[sel_], weights=amp_[sel_]))
        assert lc_["n"] == 3 and abs(lc_["mz"] - exp_) < 1e-6 and lc_["rt_lo"] == 50 and lc_["rt_hi"] == 70, lc_
        tf_bad = _st.pack("<i7d", 12, *cal_) + b"\0" * 80
        try:
            IRMS._baf_mz([1.0], tf_bad)
            raise AssertionError("unbekannte Kalibrier-Variante nicht abgelehnt")
        except ValueError:
            pass
        w.compound = {"formula": "C18H20N2"}
        t_ms.load_raw(str(ddir))
        t_ms.render()
        assert w.hrms_state["ms_raw"] and t_ms._ms_line[0].endswith(f"found {exp_:.4f}."), t_ms._ms_line[0]
        assert "mDa" in t_ms.lbl_ms_info.text() and t_ms.cmb_found.currentData() == "raw", t_ms.lbl_ms_info.text()
        assert abs(t_ms.sp_rt_lo.value() - 50 / 60) < 0.01 and abs(t_ms.sp_rt_hi.value() - 70 / 60) < 0.01
        if w.project_dir:                                  # aus dem .d nur Rohdaten + Index übernommen
            assert sorted(x.name for x in Path(w.hrms_state["ms_raw"]).iterdir()) == ["analysis.baf", "analysis.sqlite"]
        # Spektrum: gemittelt über den LC-Bereich, Isotopenmuster aus der Formel, Peak per Klick
        pat_ = HRMS.isotope_pattern("C18H21N2", 1)
        assert abs(pat_[0][0] - calc_) < 2e-5 and pat_[0][1] == 100 and 19.5 < pat_[1][1] < 22.5, pat_[:3]
        assert t_ms.F is not None and len(t_ms._sticks[0]) >= 2, t_ms._sticks
        lo_, hi_ = t_ms.view_range()
        assert lo_ < calc_ < hi_ and hi_ - lo_ < 6, (lo_, hi_)
        assert "M+1" in t_ms.lbl_ms_info.text(), t_ms.lbl_ms_info.text()
        mzs_, rel_ = t_ms._sticks
        j_ = int(np.argmin(np.abs(mzs_ - calc_)))
        px_, py_ = t_ms.F.ax.transData.transform((mzs_[j_], rel_[j_] / 2))
        t_ms.spectrum_clicked(MouseEvent("button_press_event", t_ms.canvas, px_, py_, button=1))
        pump(0.05)
        assert w.hrms_state["found"] == "pick" and abs(w.hrms_state["found_pick"] - mzs_[j_]) < 1e-5, w.hrms_state
        assert t_ms._ms_line[0].endswith(f"found {mzs_[j_]:.4f}."), t_ms._ms_line[0]
        t_ms.rt_selected(0.0, 2.0)                         # Ziehen im Chromatogramm: alle Scans mit Signal
        assert w.hrms_state["ms_rt"] == [0.0, 2.0] and t_ms.raw_found(calc_)["n"] == 12
        assert not t_ms.cb_rt_auto.isChecked()
        files = t_ms.export_to(out / "hrms" / "raw")
        assert any(f.name.endswith("_sticks.csv") for f in files) and all(f.exists() for f in files)
        w.hrms_state["found"] = None
        t_ms.load_ms(str(irdir / "ms.docx"))               # Bericht + Rohdaten: Bericht vorgewählt, Rohdaten wählbar
        t_ms.render()
        assert t_ms._ms_line[0].endswith("found 265.1703.") and t_ms.cmb_found.findData("raw") >= 0
        assert any(k_ == "ms_raw" and c_ is w.hrms_state for c_, k_, _ in PRJ._path_slots(w.get_state()))
        w.hrms_state.update(ms_file=None, ms_raw=None, ms_rt=None, found=None, found_pick=None)
        w.compound = {}
        report.append("Bruker .d raw data (LC-peak mean): OK")

        # X-Ray: künstliche CIF (Ring auf Inversionszentrum, PART 1/2, Wasser als zweites Fragment, ohne Torsionen)
        from matplotlib.backend_bases import MouseEvent
        xdir = out / "input" / "xray"
        xdir.mkdir(parents=True, exist_ok=True)
        cif_, xyz_ = XRAY.make_test_cif(xdir / "xtest.cif")
        assert abs(XRAY.chi2_scale(0.5) - 1.5382) < 1e-4, XRAY.chi2_scale(0.5)
        S_ = XRAY.Structure(XRAY.structure_blocks(cif_)[0])
        assert len(S_.ops) == 4 and len(S_.atoms) == 9 and S_.parts() == [1, 2]
        c1_ = S_.cart(S_.atoms[S_.by_label["C1"]]["frac"])
        assert np.abs(c1_ - xyz_["C1"]).max() < 1e-3, c1_                         # fraktional -> kartesisch
        mol_ = XRAY.Molecule(S_, "main")
        names_ = sorted(XRAY.atom_name(a_).encode("ascii", "replace").decode() for a_ in mol_.atoms)
        assert len(mol_.atoms) == 13 and mol_.ops == {"\u2212x, \u2212y, \u2212z": "i"}, names_
        assert mol_.find_key("Cl1B") is None and mol_.find_key("Cl1|\u2212x, \u2212y, \u2212z") is not None
        assert [f_["formula"] for f_ in mol_.fragments] == ["C6H2Cl2", "H2O"], mol_.fragments
        assert XRAY.Molecule(S_, "all").find_key("Cl1B") is not None
        mi_ = XRAY.Molecule(S_, "minor")
        assert mi_.find_key("Cl1B") is not None and mi_.find_key("Cl1") is None
        inv_ = "|\u2212x, \u2212y, \u2212z"
        m_ = XRAY.measure(mol_, ["C1", "C2"])
        d_ = float(np.linalg.norm(mol_.atoms[mol_.find_key("C2")]["xyz"] - mol_.atoms[mol_.find_key("C1")]["xyz"]))
        assert m_["text"] == "1.390(3)" and abs(d_ - 1.390) < 2e-3, (m_, d_)      # Wert aus _geom_bond
        assert XRAY.measure(mol_, ["C3", "C1" + inv_])["text"] == "1.390(3)"       # Symmetriecode 3_555
        assert XRAY.measure(mol_, ["C3" + inv_, "C1", "C2"])["text"] == "120.0(2)"  # umgekehrte Reihenfolge
        tor_ = XRAY.measure(mol_, ["C1", "C2", "C3", "C1" + inv_])
        assert tor_["source"] == "calc" and not tor_["esd"] and tor_["text"] == "0.0", tor_
        sg_ = XRAY.space_group_markup("P 21/c")
        assert sg_ == "*P*2_{1}/*c*" and "<sub>1</sub>" in XRAY.to_html(sg_), sg_
        assert XRAY.to_plain(XRAY.space_group_markup("P -1")) == "P1\u0305"
        cv_ = XRAY.crystal_values(S_.block)
        assert cv_["diff"] == "0.42/\u22120.31" and cv_["theta"] == "5.000 to 55.000", cv_["diff"]
        assert cv_["index"].startswith("\u22129 \u2264 *h*") and cv_["formula"] == "C_{6}H_{4}Cl_{2}O"
        t_x = w.xray_tab
        w.main_tabs.setCurrentWidget(t_x)
        t_x.load_cif(str(cif_))
        assert len(w.xray_state["entries"]) == 2 and Path(w.xray_state["entries"][0]["file"]).parent == proj / "data"
        assert any(k_ == "file" and c_ is w.xray_state["entries"][0] for c_, k_, _ in PRJ._path_slots(w.get_state()))
        w.xray_state["current"] = w.xray_state["entries"][0]["uid"]
        t_x.sync_widgets()
        t_x.render()
        pump()
        F_x = t_x.F
        assert F_x is not None and len(F_x.scene["atoms"]) == 13 and F_x.foot is not None
        # simulierter Klick auf C1 und C2 (Atommitte; Pixel gerundet -> nur das getroffene Atom prüfen)
        proc_ = t_x.canvas.callbacks.process

        def click_(key):
            it_ = t_x.F.scene["by_k"][t_x.F.mol.find_key(key)]
            px_, py_ = t_x.F.ax.transData.transform((it_["x"], it_["y"]))
            proc_("button_press_event", MouseEvent("button_press_event", t_x.canvas, px_, py_, button=1))
            proc_("button_release_event", MouseEvent("button_release_event", t_x.canvas, px_, py_, button=1))
            pump(0.05)
        t_x.sel = []
        click_("C1")
        assert t_x.sel == ["C1"], t_x.sel
        click_("C2")
        assert t_x.sel == ["C1", "C2"] and t_x._meas["text"] == "1.390(3)", (t_x.sel, t_x._meas)
        t_x.take_param()
        click_("C1")
        click_("C2")
        click_("C3")
        click_("C1" + inv_)
        assert t_x._meas and t_x._meas["kind"] == "torsion" and t_x._meas["source"] == "calc", t_x._meas
        t_x.take_param()
        ent_ = w.xray_state["entries"]
        ent_[1]["params"] = [{"atoms": ["C2", "C1"]}]                      # gleiche Bindung, andere Richtung
        hdr_, rows_, notes_ = t_x.params_table()
        flat_ = [r_[0] for r_ in rows_]
        assert flat_ == ["Bond lengths [\u00c5]", "C1\u2013C2", "Torsion angles [\u00b0]",
                         "C1\u2013C2\u2013C3\u2013C1"], flat_
        assert rows_[1][1:] == ["1.390(3)", "1.390(3)"] and rows_[3][1].startswith("0.0^{a}"), rows_
        assert any("no s.u." in n_ for n_ in notes_) and any(n_.startswith("Symmetry") for n_ in notes_), notes_
        hc_, rc_ = t_x.crystal_table()
        assert len(hc_) == 3 and [r_ for r_ in rc_ if r_[0] == "Space group"][0][1] == "*P*2_{1}/*c*"
        # Drehen per Ziehen: Ansicht im Zustand, Projektion ändert sich
        x0_ = [(a_["x"], a_["y"]) for a_ in t_x.F.scene["atoms"]]
        cx_, cy_ = t_x.F.ax.transData.transform((0.0, 0.0))
        cy_ += 0.3 * t_x.F.ax.bbox.height                               # neben das Molekül (kein Label, kein Atom)
        proc_("button_press_event", MouseEvent("button_press_event", t_x.canvas, cx_, cy_, button=1))
        proc_("motion_notify_event", MouseEvent("motion_notify_event", t_x.canvas, cx_ + 40, cy_ + 10, button=1))
        proc_("button_release_event", MouseEvent("button_release_event", t_x.canvas, cx_ + 40, cy_ + 10, button=1))
        pump(0.1)
        assert ent_[0]["view"] and len(ent_[0]["view"]) == 9, ent_[0].get("view")
        x1_ = [(a_["x"], a_["y"]) for a_ in t_x.F.scene["atoms"]]
        assert max(abs(a_[0] - b_[0]) + abs(a_[1] - b_[1]) for a_, b_ in zip(x0_, x1_)) > 0.05
        # Mercury-Stil (Standard): schattierte Schichten je Atom; je Atom ausblenden / Label per Doppelklick
        assert w.xray_state["style"] == "mercury"
        pc_ = t_x.F.atom_art[t_x.F.mol.find_key("C1")][0]
        assert len(pc_.get_paths()) == 13, len(pc_.get_paths())                 # 12 Schichten + Umriss
        assert "lbl_C1" in t_x.F.ann_items and t_x.F.ann_items["lbl_C1"].get_fontweight() == "bold"   # wie Mercury-Bilder
        w.xray_state["label_pt"] = 12.5                                     # Labelgröße je Tab
        t_x.render(push=False)
        assert t_x.F.ann_items["lbl_C1"].get_fontsize() == 12.5
        w.xray_state["label_pt"] = 0.0
        t_x.render(push=False)
        it_ = t_x.F.scene["by_k"][t_x.F.mol.find_key("C2")]
        px_, py_ = t_x.F.ax.transData.transform((it_["x"], it_["y"]))
        t_x.sel = []
        proc_("button_press_event", MouseEvent("button_press_event", t_x.canvas, px_, py_, button=1, dblclick=True))
        pump(0.6)
        assert ent_[0]["label_on"] == {"C2": False} and "lbl_C2" not in t_x.F.ann_items, ent_[0]["label_on"]
        assert not t_x.sel
        # Shift + Rahmen: Atome im Rahmen -> Menü (hier abgefangen), dann Labels mehrerer Atome aus und wieder an
        got_, boxed_ = [], t_x.atoms_boxed
        t_x.atoms_boxed = got_.extend
        try:
            p1_, p2_ = (t_x.F.ax.transData.transform((t_x.F.scene["by_k"][t_x.F.mol.find_key(l_)]["x"],
                                                      t_x.F.scene["by_k"][t_x.F.mol.find_key(l_)]["y"]))
                        for l_ in ("C1", "C2"))
            x0_, y0_ = min(p1_[0], p2_[0]) - 3, min(p1_[1], p2_[1]) - 3
            x1_, y1_ = max(p1_[0], p2_[0]) + 3, max(p1_[1], p2_[1]) + 3
            proc_("button_press_event", MouseEvent("button_press_event", t_x.canvas, x0_, y0_, button=1, key="shift"))
            proc_("motion_notify_event", MouseEvent("motion_notify_event", t_x.canvas, x1_, y1_, button=1,
                                                    key="shift"))
            proc_("button_release_event", MouseEvent("button_release_event", t_x.canvas, x1_, y1_, button=1,
                                                     key="shift"))
            pump(0.1)
        finally:
            t_x.atoms_boxed = boxed_
        keys_ = {t_x.F.mol.atoms[k_]["key"] for k_ in got_}
        assert {"C1", "C2"} <= keys_, keys_
        assert not any(p_.get_linestyle() == "--" for p_ in t_x.F.ax.patches), "Rahmen bleibt stehen"
        t_x.set_labels(got_, False)
        assert not any(f"lbl_{k_}" in t_x.F.ann_items for k_ in keys_), keys_
        t_x.set_labels(got_, True)
        assert all(f"lbl_{k_}" in t_x.F.ann_items for k_ in keys_ if not k_.startswith("H")), keys_
        ent_[0]["label_on"] = {}
        t_x.set_label("H3", True)                                           # H-Label einzeln an
        ent_[0]["hidden_atoms"] = ["H3" + inv_]                             # ein einzelnes H ausblenden
        t_x.render()
        assert "lbl_H3" in t_x.F.ann_items and t_x.F.mol.find_key("H3" + inv_) not in t_x.F.scene["by_k"]
        w.xray_state["style"] = "flat"
        t_x.render()
        assert len(t_x.F.atom_art[t_x.F.mol.find_key("C1")][0].get_paths()) == 2
        assert t_x.F.ann_items["lbl_C1"].get_fontweight() == "bold"
        w.xray_state["style"] = "mercury"
        ent_[0].update(label_on={}, hidden_atoms=[])
        ent_[0]["hidden_frags"] = [1]                                   # Wasser ausblenden
        w.xray_state["show_h"] = False
        t_x.sync_widgets()
        t_x.render()
        assert len(t_x.F.scene["atoms"]) == 8, len(t_x.F.scene["atoms"])
        files = t_x.export_to(out / "xray" / "xray_test")
        for f in files:
            assert f.exists() and f.stat().st_size > 200, f
        import pymupdf
        npath_ = len(pymupdf.open(str(out / "xray" / "xray_test_xtest_a.pdf"))[0].get_drawings())
        assert npath_ > 20, npath_
        csv_ = (out / "xray" / "xray_test_crystal_data.csv").read_text(encoding="utf-8-sig")
        assert "P2\u2081/c" in csv_ and "0.42/\u22120.31" in csv_
        from openpyxl import load_workbook
        wb_ = load_workbook(out / "xray" / "xray_test_tables.xlsx")
        assert wb_.sheetnames == ["Crystal data", "Selected parameters"]
        assert (out / "xray" / "xray_test_captions.txt").read_text(encoding="utf-8").count("X-ray solid state structure") == 2
        # Lösungsmittel-Schalter, Startansicht, Label-Modus, Drahtmodell, Bildunterschrift
        mol_ = t_x.molecule(ent_[0])
        assert [f_["kind"] for f_ in mol_.fragments] == ["main", "solvent"] and mol_.fragments[1]["name"] == "water"
        assert XRAY.cif_order(mol_, ["C2", "C1"]) == ["C1", "C2"]                # Reihenfolge wie in der CIF
        R_ = XRAY.default_view(mol_)
        assert np.allclose(R_ @ R_.T, np.eye(3), atol=1e-9) and abs(np.linalg.det(R_) - 1) < 1e-9
        # Ausrichten: Bindung nach unten, Ebene von vorn, Drehen, ChemDraw als Vorlage
        xyz_k = lambda k_: mol_.atoms[mol_.find_key(k_)]["xyz"]  # noqa: E731
        Rb_ = XRAY.align_bond(R_, mol_, ["C1", "Cl1"], "down")
        v_ = Rb_ @ (xyz_k("Cl1") - xyz_k("C1"))
        assert abs(v_[0]) < 1e-6 and abs(v_[2]) < 1e-6 and v_[1] < 0, v_
        Rp_ = XRAY.align_plane(R_, mol_, ["C1", "C2", "C3"])
        n_ = np.cross(xyz_k("C2") - xyz_k("C1"), xyz_k("C3") - xyz_k("C1"))
        assert abs(abs((Rp_ @ (n_ / np.linalg.norm(n_)))[2]) - 1) < 1e-6
        Rpb_ = XRAY.align_bond(Rp_, mol_, ["C1", "C2"], "right")           # Bindung in der Ebene: Ebene bleibt
        assert abs(abs((Rpb_ @ (n_ / np.linalg.norm(n_)))[2]) - 1) < 1e-6
        from rdkit import Chem as _Chem
        if hasattr(_Chem, "MolToCDXMLBlock"):
            from rdkit.Chem import AllChem as _AC
            tm_ = _Chem.MolFromSmiles("Clc1ccc(Cl)cc1")
            _AC.Compute2DCoords(tm_)
            tpl_ = xdir / "dcb.cdxml"
            tpl_.write_text(_Chem.MolToCDXMLBlock(tm_), encoding="utf-8")
            Rt_, inf_ = XRAY.align_to_template(mol_, tpl_)
            assert inf_["n"] == 8 and inf_["rmsd"] < 0.3, inf_
            cls_ = [a_.GetIdx() for a_ in tm_.GetAtoms() if a_.GetSymbol() == "Cl"]
            cf_ = tm_.GetConformer()
            dt_ = np.array(cf_.GetAtomPosition(cls_[1])) - np.array(cf_.GetAtomPosition(cls_[0]))
            dc_ = (Rt_ @ (xyz_k("Cl1|" + [o_ for o_ in mol_.ops][0]) - xyz_k("Cl1")))[:2]
            cosang_ = abs(np.dot(dt_[:2], dc_)) / np.linalg.norm(dt_[:2]) / np.linalg.norm(dc_)
            assert cosang_ > 0.98, cosang_                                   # Cl–Cl wie in der Zeichnung
        ent_[0].update(hidden_frags=[], view=None)
        w.xray_state.update(show_h=True, hide_solv=True)
        t_x.sync_widgets()
        t_x.render()
        assert t_x.F.scene["atoms"] and all(mol_.atoms[it_["k"]]["frag"] == 0 for it_ in t_x.F.scene["atoms"])
        assert "<sub>1</sub>" in t_x.lbl_info.text(), t_x.lbl_info.text()          # Raumgruppe formatiert
        cap_ = t_x.caption()[0]
        assert cap_.startswith("X-ray solid state structure of xtest_a (70% (PART 1) and 30% (B) occupancy; major part "
                               "shown). Thermal ellipsoids are shown with 50% probability."), cap_
        assert "Co-crystallised water is omitted for clarity." in cap_ and "Symmetry code: i: " in cap_, cap_
        assert "Selected bond lengths and torsion angles in [Å] and [°]: C1–C2: 1.390(3); C1–C2–C3–C1ⁱ: 0.0." in cap_, cap_
        assert "<b>xtest_a</b>" in t_x.caption()[1] and not t_x.param_tables(ent_[0])   # wenige Werte: in der Unterschrift
        c1_, cl1_ = mol_.find_key("C1"), mol_.find_key("Cl1")
        assert mol_.substituent(c1_, mol_.find_key("C2")) is None                  # Ringbindung im Kern
        assert mol_.beyond(c1_, cl1_) == {cl1_}
        assert mol_.substituent(c1_, cl1_) == {cl1_} and mol_.substituent(cl1_, c1_) == {cl1_}
        t_x.sel = ["C1", "Cl1"]
        t_x.wire_selected()
        assert ent_[0]["wire"] == ["Cl1"] and not t_x.sel, ent_[0]["wire"]
        t_x.render()                                                              # Bindung Kern–Draht dünn
        mixed_ = [ln_ for a_, b_, ln_, h_ in t_x.F.bond_art if {a_, b_} == {c1_, cl1_}]
        assert len(mixed_) == 2 and all(ln_.get_linewidth() < t_x.F.lw_bond for ln_ in mixed_), mixed_
        assert "lbl_C1" in t_x.F.ann_items and t_x.F.scene["by_k"][c1_]["kind"] != "wire"
        ent_[0]["wire"] = ["C1", "Cl1"]                                           # beide Enden -> Drahtbindung
        ent_[0]["wire_name"] = "the C-Cl group"
        t_x.render()
        assert t_x.F.scene["by_k"][cl1_]["kind"] == "wire" and "lbl_Cl1" not in t_x.F.ann_items
        halves_ = [ln_ for a_, b_, ln_, h_ in t_x.F.bond_art if {a_, b_} == {c1_, cl1_}]
        assert len(halves_) == 2 and halves_[0].get_linewidth() < t_x.F.lw_bond, halves_
        assert "the C-Cl group is shown as wireframe" in t_x.caption()[0], t_x.caption()[0]
        t_x.sel = []
        t_x.unwire_selected()
        t_x.layout()["lbl_C1"] = [0.0, 0.0]                  # gespeichert, liegt aber auf dem Atom -> neu platziert
        t_x.render()
        assert tuple(float(v_) for v_ in t_x.F.ann_items["lbl_C1"].xyann) != (0.0, 0.0)
        t_x.layout().pop("lbl_C1")
        w.xray_state["label_mode"] = "hetero"
        t_x.sync_widgets()
        t_x.render()
        assert not ent_[0]["wire"] and "lbl_Cl1" in t_x.F.ann_items and "lbl_C1" not in t_x.F.ann_items
        # Darstellung auf Standard: Einstellungen + gewählte Struktur zurück, CIFs/Namen/Parameter bleiben
        e_x = t_x.entry()
        w.xray_state.update(bond_w=0.2, prob=30, style="flat", label_pt=14.0, show_h=False, disorder="all")
        e_x.update(hidden_atoms=[e_x["hidden_atoms"] or "C2"], label_on={"C1": False}, wire=["C1"], zoom=2.0,
                   name="umbenannt")
        n_ent, params_ = len(w.xray_state["entries"]), list(e_x.get("params") or [])
        t_x.reset_display()
        d_x = XRAY.XrayTab.DEFAULT_STATE
        assert all(w.xray_state[k_] == d_x[k_] for k_ in ("bond_w", "prob", "style", "show_h", "disorder"))
        assert w.xray_state["label_pt"] == 0.0 and abs(t_x.sp_bond.value() - 0.12) < 1e-9 and t_x.sp_prob.value() == 50
        assert not e_x["hidden_atoms"] and not e_x["label_on"] and not e_x["wire"] and e_x["zoom"] == 1.0
        assert len(w.xray_state["entries"]) == n_ent and e_x["name"] == "umbenannt" and e_x["params"] == params_
        assert t_x.F is not None
        # Stile: Mercury (C weiß, schwarze Bindungen, fette Labels), Olex2, Kugel-Stab glänzend, Kugel-Stab mit Tiefe,
        # Strichzeichnung. Im Tab X-Ray wählbar nur Mercury/Olex2 (alte Projekte -> Mercury); gezeichnet werden alle
        assert [t_x.cmb_style.itemData(i_) for i_ in range(t_x.cmb_style.count())] == list(XRAY.XRAY_STYLES)
        w.xray_state["style"] = "ballstick"
        assert t_x.opts(t_x.entry())["style"] == "mercury"
        styles0_, XRAY.XRAY_STYLES = XRAY.XRAY_STYLES, XRAY.STYLES
        for sty_ in XRAY.STYLES:
            w.xray_state["style"] = sty_
            t_x.render()
            F_s = t_x.F
            assert F_s.style == sty_
            c_it = next(i_ for i_ in F_s.scene["atoms"] if i_["el"] == "C")
            if sty_ in XRAY.BALL_STYLES:
                assert (sty_ == "cartoon" or abs(c_it["w"] - 2 * XRAY.ball_radius("C", sty_)) < 1e-9) \
                    and c_it["w"] == c_it["h"] and not c_it["curves"]
                assert F_s.bond_w == XRAY.ROD_W[sty_] and len(F_s.rods["frac"]) > 10 and F_s.rods["frac"].max() == 1.0
                assert F_s.rods["inplace"], "Stabsegmente nicht im gemeinsamen Puffer"      # schnelles Drehen
                pc_s = F_s.atom_art[c_it["k"]][0]
                nv_ = [len(q_.vertices) for q_ in pc_s.get_paths()]
                F_s.set_drag(True)                                     # Ziehen: gröbere Kugelpolygone, gleiche Schichten
                nd_ = [len(q_.vertices) for q_ in pc_s.get_paths()]
                F_s.set_drag(False)
                assert len(nd_) == len(nv_) and sum(nd_) < sum(nv_), (sum(nd_), sum(nv_))
                assert [len(q_.vertices) for q_ in pc_s.get_paths()] == nv_
                fc_ = F_s.atom_art[c_it["k"]][0].get_facecolors()                         # Phong: dunkel … Glanz
                if sty_ == "cartoon":                           # weiß mit schwarzem Umriss, hohle Stäbe, Perspektive
                    assert all(tuple(f_[:3]) == (1, 1, 1) for f_ in fc_[:-1]), fc_
                    assert tuple(pc_s.get_edgecolors()[-1][:3]) == (0, 0, 0)
                    lc_ = F_s.atom_rods[c_it["k"]][0].get_colors()
                    assert {tuple(c_[:3]) for c_ in lc_} == {(0, 0, 0), (1, 1, 1)}, lc_
                    ws_ = {round(i_["w"], 6) for i_ in F_s.scene["atoms"] if i_["el"] == "C"}
                    assert len(ws_) > 1, ws_                     # vordere C größer als hintere
                    continue
                assert len(fc_) > 10 and sum(fc_[0][:3]) < sum(fc_[len(fc_) // 2][:3])    # Rand dunkler als Mitte
                if sty_ == "ballstick":
                    assert min(max(f_[:3]) for f_ in fc_[-3:-1]) > 0.95, fc_[-3:]          # Glanzpunkt fast weiß
            if sty_ == "mercury":
                assert XRAY.default_color("C", "mercury") == XRAY.OLEX2_COLORS["C"]               # C grau wie bisher
                assert all(a_.get_fontweight() == "bold" for a_ in F_s.ann_items.values())
                assert all(ln_.get_color() == "black" for _, _, ln_, h_ in F_s.bond_art if h_ is None)
            if sty_ == "depth":                                 # wie ORCA-Handbuch: ohne Aufhellung, Jmol-Farben
                assert F_s._fog(F_s._zmax - XRAY.DEPTH_A - 1) == 0.0
                assert XRAY.default_color("O", "depth") == "#FF0D0D" and XRAY.default_color("O") != "#FF0D0D"
            f_s = t_x.export_to(out / f"xray_style_{sty_}")
            assert all(f_.exists() and f_.stat().st_size > 500 for f_ in f_s if f_.suffix in (".pdf", ".svg", ".png"))
        XRAY.XRAY_STYLES = styles0_
        w.xray_state["style"] = "mercury"
        # Labelfarbe: schwarz (Standard) oder wie das Atom, helle Farben abgedunkelt (lesbar)
        w.xray_state["label_color"] = "atom"
        t_x.sync_widgets()
        t_x.render()
        assert t_x.cmb_lcol.currentData() == "atom"
        from matplotlib.colors import to_rgb
        n_lab_ = next(a_ for k_, a_ in t_x.F.ann_items.items()
                      if k_.startswith("lbl_") and t_x.F.mol.atoms[a_._atom]["el"] not in ("C", "H"))
        el_ = t_x.F.mol.atoms[n_lab_._atom]["el"]
        assert to_rgb(n_lab_.get_color()) == to_rgb(XRAY.label_color(el_, {}, "atom", "mercury")), (el_, n_lab_.get_color())
        assert to_rgb(n_lab_.get_color()) != (0, 0, 0)
        hc_ = XRAY.label_color("H", {}, "atom")
        assert 0.299 * hc_[0] + 0.587 * hc_[1] + 0.114 * hc_[2] <= 0.45 + 1e-9, hc_
        w.xray_state["label_color"] = "black"
        t_x.sync_widgets()
        t_x.render()
        assert all(to_rgb(a_.get_color()) == (0, 0, 0) for a_ in t_x.F.ann_items.values())
        # Zeichenfläche darf durch Leisten nicht schrumpfen: bei 1440 × 900 mindestens die halbe Fensterbreite
        # (gemessen ~55 %; eine Seitenleiste läge bei ~44 %). Relativ, damit andere Schriften in CI nicht stören.
        size0_ = w.size()
        w.show()
        w.resize(1440, 900)
        w.main_tabs.setCurrentWidget(t_x)
        pump(0.5)
        t_x.host.fit()
        pump(0.2)
        if t_x.canvas.width() < 0.5 * w.width():                  # Messwerte für die Fehlersuche (andere Systeme)
            left_ = t_x.lv.parentWidget()
            cards_ = sorted(((c_.minimumSizeHint().width(), getattr(c_, "title", lambda: type(c_).__name__)()[:28])
                             for c_ in left_.findChildren(QWidget) if c_.parentWidget() is left_ and c_.isVisibleTo(left_)),
                            reverse=True)[:4]
            wide_ = sorted(((x_.minimumSizeHint().width(), type(x_).__name__, (getattr(x_, "text", lambda: "")() or "")[:30])
                            for x_ in left_.findChildren(QWidget) if x_.isVisibleTo(left_)), reverse=True)[:6]
            raise AssertionError(f"Abbildung {t_x.canvas.width()} px bei Fenster {w.width()} px; linke Spalte "
                                 f"{left_.width()} px (min {left_.minimumSizeHint().width()}), Karten {cards_}, "
                                 f"breiteste Elemente {wide_}, Schrift {QApplication.font().family()} "
                                 f"{QApplication.font().pointSizeF()} pt, DPI {w.logicalDpiX()}")
        w.resize(size0_)
        report.append(f"X-Ray (CIF, symmetry, disorder, ORTEP, picking, tables, wireframe, caption): OK "
                      f"({npath_} pdf paths)")
        w.xray_state.update(entries=[], current=None, show_h=True, hide_solv=False, label_mode="all")
        t_x.sel = []
        t_x.sync_widgets()
        t_x.render()

        # DFT: xyz (zwei Bilder, CRLF), Bindungen, Umnummerieren per Klick, xyz speichern (Zeilen unverändert)
        ddir = out / "input" / "dft"
        xyzf_ = DFT.make_test_xyz(ddir / "dtest.xyz")
        fr_ = DFT.read_xyz(xyzf_)
        assert len(fr_) == 2 and fr_[-1]["els"] == ["H", "C", "Cl", "C", "H", "H", "H", "H"], fr_[-1]["els"]
        assert fr_[-1]["comment"] == "E -539.123456789 (test)" and not fr_[0]["lines"][0].endswith("\r")
        assert DFT.parse_xyz("1\n\n6 0 0 0\n")[0]["els"] == ["C"]                 # Ordnungszahl statt Symbol
        for bad_ in ("C 0 0 0\n", "3\n\nC 0 0 0\n"):
            try:
                DFT.parse_xyz(bad_)
                raise AssertionError(bad_)
            except ValueError:
                pass
        bo_ = DFT.bonds_of(fr_[-1]["els"], fr_[-1]["xyz"])
        assert sorted(bo_) == [(0, 3), (1, 2), (1, 3), (1, 4), (1, 5), (3, 6), (3, 7)], bo_
        nb_ = {k_: [] for k_ in range(8)}
        for a_, b_ in bo_:
            nb_[a_].append(b_)
            nb_[b_].append(a_)
        assert DFT.renumber(range(8), [1, 3, 2]) == [1, 3, 2, 0, 4, 5, 6, 7]
        assert DFT.renumber(range(8), [1, 3, 2], fr_[-1]["els"], nb_) == [1, 4, 5, 3, 0, 6, 7, 2]
        assert DFT.renumber([2, 1, 0, 3, 4, 5, 6, 7], [3]) == [3, 2, 1, 0, 4, 5, 6, 7]   # Rest: bisherige Reihenfolge
        t_d = w.dft_tab
        w.main_tabs.setCurrentWidget(t_d)
        e_d = t_d.load_xyz(str(xyzf_))
        assert e_d and Path(e_d["file"]).parent == proj / "data"
        assert any(k_ == "file" and c_ is e_d for c_, k_, _ in PRJ._path_slots(w.get_state()))
        pump()
        assert t_d.F is not None and len(t_d.F.scene["atoms"]) == 8 and len(t_d.F.scene["bonds"]) == 7
        assert t_d.F.ann_items["lbl_a2"].get_text() == "Cl3"
        proc_d = t_d.canvas.callbacks.process

        def click_d(k_):
            it_ = t_d.F.scene["by_k"][k_]
            px_, py_ = t_d.F.ax.transData.transform((it_["x"], it_["y"]))
            proc_d("button_press_event", MouseEvent("button_press_event", t_d.canvas, px_, py_, button=1))
            proc_d("button_release_event", MouseEvent("button_release_event", t_d.canvas, px_, py_, button=1))
            pump(0.05)
        click_d(1)
        assert t_d.picks == [], t_d.picks                                     # ohne Modus: Klick nummeriert nicht
        t_d.btn_num.setChecked(True)
        for k_ in (1, 3, 2):
            click_d(k_)
        assert t_d.picks == [1, 3, 2], t_d.picks
        click_d(3)                                                             # erneuter Klick nimmt es heraus
        assert t_d.picks == [1, 2], t_d.picks
        click_d(3)
        assert t_d.picks == [1, 2, 3], t_d.picks
        assert t_d.F.ann_items["lbl_a3"].get_text() == "C3" and t_d.F.ann_items["lbl_a3"].get_fontweight() == "bold"
        t_d.picks = [1, 3, 2]
        t_d.apply_numbering()
        assert e_d["order"] == [1, 3, 2, 0, 4, 5, 6, 7] and not t_d.numbering and not t_d.btn_num.isChecked()
        assert t_d.F.ann_items["lbl_a2"].get_text() == "Cl3" and t_d.F.ann_items["lbl_a3"].get_text() == "C2"
        w.dft_state["label_pt"] = 14.0                                       # eigene Labelgröße des Tabs
        t_d.render(push=False)
        assert t_d.F.ann_items["lbl_a1"].get_fontsize() == 14.0
        w.dft_state["label_pt"] = 0.0
        w.dft_state["zero_based"] = True
        t_d.render(push=False)
        assert t_d.F.ann_items["lbl_a1"].get_text() == "C0"
        w.dft_state["zero_based"] = False
        pump(0.3)
        w.undo()                                                               # Umnummerieren rückgängig
        pump(0.4)
        assert w.dft_state["entries"][0]["order"] is None, w.dft_state["entries"][0]["order"]
        w.redo()
        pump(0.4)
        t_d = w.dft_tab
        assert w.dft_state["entries"][0]["order"] == [1, 3, 2, 0, 4, 5, 6, 7]
        sx_ = t_d.save_xyz(out / "dft" / "dtest_new.xyz")
        fr2_ = DFT.read_xyz(sx_)
        assert len(fr2_) == 2 and fr2_[-1]["els"] == ["C", "C", "Cl", "H", "H", "H", "H", "H"]
        assert fr2_[0]["lines"] == [fr_[0]["lines"][k_] for k_ in (1, 3, 2, 0, 4, 5, 6, 7)]    # Zeilen unverändert
        src_ = Path(w.dft_state["entries"][0]["file"])
        assert t_d.save_xyz(src_) is None and DFT.read_xyz(src_)[0]["lines"] == fr_[0]["lines"]  # nie überschreiben
        files = t_d.export_to(out / "dft" / "dft_test")
        assert {f_.suffix for f_ in files} == {".pdf", ".svg", ".png", ".xyz"}, files
        assert all(f_.exists() and f_.stat().st_size > 100 for f_ in files)
        t_d.reset_order()
        assert w.dft_state["entries"][0]["order"] is None
        # Rechnungen: ORCA 6 (UTF-16/CRLF, CPCM, eine imaginäre Frequenz) und Gaussian (Format geraten); SI-Block, Word
        orf_, gf_ = DFT.make_test_orca(ddir / "otest.out"), DFT.make_test_gaussian(ddir / "gtest.log")
        assert DFT.is_calc(orf_) and DFT.is_calc(gf_) and not DFT.is_calc(xyzf_)
        co_ = DFT.read_calc(orf_)
        assert (co_["version"], co_["charge"], co_["mult"], co_["opt_converged"], co_["normal"]) == ("6.1.1", 0, 1, True, True)
        assert DFT.imaginary(co_) == ["-31.20"] and co_["thermo"]["G"] == "-539.94683332", co_["thermo"]
        assert co_["solvation"]["e_diel"] == "-0.00412345678901" and co_["solvation"]["solvent"] == "TETRAHYDROFURAN"
        assert co_["conv"][0] == ("Energy change", "-0.0000002000", "0.0000010000", "YES") and len(co_["conv"]) == 5
        cg_ = DFT.read_calc(gf_)
        assert cg_["solvation"]["model"] == "SMD" and cg_["thermo"]["H"] == "-539.916710" and len(cg_["conv"]) == 4
        assert cg_["energy"] == "-539.987654321" and not DFT.imaginary(cg_) and cg_["geometry"][1][1][0] == "Cl"
        # gleichnamige xyz (ORCA, 14 Stellen) wird mitkopiert und benutzt, wenn sie zur Endgeometrie passt
        with open(ddir / "otest.xyz", "w", encoding="utf-8", newline="\n") as fh_:
            fh_.write("8\ntest\n" + "".join(f"{el_}  {float(x_):.14f}  {float(y_):.14f}  {float(z_):.14f}\n"
                                            for el_, x_, y_, z_ in DFT._TEST_GEOM))
        fr_p, used_ = DFT.calc_frame(co_, partner=ddir / "otest.xyz")
        assert used_ and fr_p["lines"][1].startswith("Cl  0.00000000000000"), fr_p["lines"][1]
        assert not DFT.calc_frame(co_, partner=xyzf_)[1]                      # andere Geometrie -> ignoriert
        assert co_["point_group"] == "C1" and cg_["point_group"] == "CS" and DFT.pg_html("CS") == "C<sub>s</sub>"
        e_o = t_d.load_xyz(str(orf_))
        assert (Path(e_o["file"]).with_suffix(".xyz")).is_file()
        t_d.load_xyz(str(gf_))
        e_o["order"] = [1, 0, 2, 3, 4, 5, 6, 7]
        w.dft_state["current"] = e_o["uid"]
        t_d.sync_widgets()
        t_d.render()
        assert t_d.F is not None and len(t_d.F.scene["atoms"]) == 8
        lab_ = [k_ for k_ in range(8) if t_d.F.mol.atoms[k_]["el"] != "H"]    # Labels mehrerer Atome (Rahmen)
        t_d.set_labels(lab_, False)
        assert all(e_o["label_on"].get(t_d.F.mol.atoms[k_]["key"]) is False for k_ in lab_), e_o["label_on"]
        t_d.set_labels(lab_, True)
        assert not e_o["label_on"], e_o["label_on"]                            # Vorgabe: Ausnahmen fallen weg
        pl_ = DFT.si_plain(t_d.si_of(e_o))
        for frag_ in ("ORCA 6.1.1: ! B3LYP D4 def2-SVP TightOpt Freq CPCM(thf)",
                      "Charge: 0\tMultiplicity: 1\tPoint group: C1", "1 imaginary frequency (\u221231.20 cm-1)",
                      "thermochemistry at 298.15 K and 1.00 atm (quasi-RRHO)",
                      "CPCM dielectric energy (included in E(el)): \u22120.00412345678901 Eh",
                      "Energy change\t\u22120.0000002000\t0.0000010000\tyes", "\nCl  0.00000000000000",
                      "E_el [Eh]\tH [Eh]\tG [Eh]\n\u2212539.98765432\t\u2212539.91670987\t\u2212539.94683332"):
            assert frag_ in pl_, (frag_, pl_)
        assert pl_.index("Energy change") < pl_.index("Frequency calculation")      # Tabelle unter der Optimierung
        assert pl_.index("\nCl  ") < pl_.index("\nC  ")                              # Koordinaten umnummeriert
        assert T("dft_partner_used") in t_d.lbl_info.text()
        w.dft_state["si_conv"] = False                                             # Teile abschaltbar
        assert "Energy change" not in DFT.si_plain(t_d.si_of(e_o))
        w.dft_state["si_conv"] = True
        assert T("dft_w_solv_cpcm") in t_d.lbl_warn.text() and t_d.btn_si_copy.isEnabled()
        files = t_d.export_to(out / "dft" / "dft_calc")
        names_ = {f_.name for f_ in files}
        assert {"dft_calc_SI.docx", "dft_calc_SI.txt"} <= names_, names_
        import zipfile as _zf
        with _zf.ZipFile(out / "dft" / "dft_calc_SI.docx") as z_:
            doc_ = z_.read("word/document.xml").decode("utf-8")
            assert "word/media/image1.png" in z_.namelist() and "word/media/image2.png" in z_.namelist()
        for frag_ in ("Point group", "Courier New", 'w:cols w:num="2"', "Max. gradient", "otest", "Gaussian 16",
                      "quasi-RRHO", "<w:tab/>"):
            assert frag_ in doc_, frag_
        # Messen und Ausrichten wie im Tab X-Ray (gemeinsame Klasse XRAY.MeasureAlign), Werte aus der Geometrie
        t_d.set_numbering(False)
        g_d = t_d.geometry(t_d.entry())
        xyz_d = np.asarray(g_d[2], float)
        for keys_, kind_ in ((["a1", "a3"], "bond"), (["a2", "a1", "a3"], "angle"), (["a2", "a1", "a3", "a0"], "torsion")):
            t_d.sel = []
            for k_ in keys_:
                t_d.atom_clicked(int(k_[1:]))
            assert t_d.sel == keys_, t_d.sel
            m_d = t_d._meas
            idx_ = [int(k_[1:]) for k_ in keys_]
            ref_ = {"bond": XRAY.distance, "angle": XRAY.angle, "torsion": XRAY.torsion}[kind_]([xyz_d[i_] for i_ in idx_])
            assert m_d["kind"] == kind_ and abs(m_d["value"] - ref_) < 1e-9 and m_d["source"] == "calc", m_d
            assert "\n" not in t_d.lbl_meas.text() and t_d.lbl_meas.styleSheet() == "", t_d.lbl_meas.text()  # keine esd-Warnung
        t_d.sel = ["a1", "a2"]
        t_d.align_bond("down")                                 # Bindung C–Cl zeigt im Bild nach unten
        R_d = np.array(t_d.entry()["view"]).reshape(3, 3)
        v_d = R_d @ (xyz_d[2] - xyz_d[1])
        assert v_d[1] < -0.99 * np.linalg.norm(v_d) and t_d.sel == [], v_d
        t_d.sel = ["a1", "a2", "a3"]
        t_d.align_plane()
        R_d = np.array(t_d.entry()["view"]).reshape(3, 3)
        n_d = np.cross(xyz_d[2] - xyz_d[1], xyz_d[3] - xyz_d[1])
        assert abs(abs((R_d @ n_d)[2]) - np.linalg.norm(n_d)) < 1e-6           # Ebene von vorn
        t_d.sel = ["a1"]
        t_d.set_numbering(True)                                # Nummerieren hebt die Messauswahl auf
        assert t_d.sel == [] and t_d.picks == []
        t_d.set_numbering(False)
        # Kugel-Stab: Atomradius/Bindungsdicke relativ zum Standard; Hintergrund wählbar (auch im Export)
        # (alte Projekte mit „ballstick“/IboView zeichnen als Kugel-Stab „depth“; Auswahl ohne IboView, mit Strichzeichnung)
        assert [t_d.cmb_style.itemData(i_) for i_ in range(t_d.cmb_style.count())] == list(XRAY.DFT_STYLES)
        w.dft_state.update(style="ballstick", radius=0.5, bond_w=0.16, bg="#203040")
        t_d.sync_widgets()
        t_d.render()
        assert t_d.F.style == "depth" and t_d.cmb_style.currentData() == "depth"
        c_d = next(i_ for i_ in t_d.F.scene["atoms"] if i_["el"] == "C")
        assert abs(c_d["w"] - 4 * XRAY.ball_radius("C", "depth")) < 1e-9, c_d["w"]
        assert abs(t_d.F.bond_w - 2 * XRAY.ROD_W["depth"]) < 1e-9 and t_d.cmb_bg.currentData() == "color"
        from PIL import Image as PILImage
        png_bg = next(f_ for f_ in t_d.export_to(out / "dft_bg") if f_.suffix == ".png")
        with PILImage.open(png_bg) as im_:
            assert im_.convert("RGB").getpixel((1, 1)) == (0x20, 0x30, 0x40), im_.convert("RGB").getpixel((1, 1))
        w.dft_state["bg"] = "transparent"
        t_d.sync_widgets()
        png_bg = next(f_ for f_ in t_d.export_to(out / "dft_bg") if f_.suffix == ".png")
        with PILImage.open(png_bg) as im_:
            assert im_.convert("RGBA").getpixel((1, 1))[3] == 0
        w.dft_state["style"] = "mercury"
        e_dr = t_d.entry()                                     # Darstellung auf Standard, Nummerierung bleibt
        ord_ = list(e_dr.get("order") or [])
        w.dft_state.update(bond_w=0.3, radius=0.5, show_h=False)
        e_dr["zoom"] = 3.0
        t_d.reset_display()
        assert w.dft_state["bond_w"] == 0.08 and w.dft_state["radius"] == 0.25 and w.dft_state["show_h"]
        assert w.dft_state["bg"] == "white" and t_d.cmb_bg.currentData() == "white"
        assert e_dr["zoom"] == 1.0 and list(e_dr.get("order") or []) == ord_ and abs(t_d.sp_bond.value() - 0.08) < 1e-9
        report.append("DFT xyz (reader, bonds, click renumbering, save, export): OK")
        report.append("DFT outputs (ORCA 6 UTF-16, Gaussian, SI block, Word): OK")
        # Vergleichsmodus: Nummerierung einer anderen Geometrie bzw. einer CIF übernehmen
        w.dft_state.update(entries=[], current=None)
        e_ref = t_d.load_xyz(str(orf_))                                   # Vorlage: C, Cl, C, H…
        e_tg = t_d.load_xyz(str(xyzf_))                                   # gleiche Geometrie, andere Reihenfolge
        t_d.fill_refs()
        t_d.cmb_ref.setCurrentIndex(t_d.cmb_ref.findData("dft:" + e_ref["uid"]))
        t_d.apply_reference()
        assert e_tg["order"] == [1, 2, 3, 4, 5, 6, 7, 0], e_tg["order"]
        assert "0.00" in t_d.lbl_ref.text() and not e_tg["ref_labels"], t_d.lbl_ref.text()
        w.xray_tab.load_cif(str(cif_))                                    # CIF als Vorlage: Labels der Kristallstruktur
        mol_c = w.xray_tab.molecule(w.xray_state["entries"][0])
        main_c = [k_ for k_, a_ in enumerate(mol_c.atoms) if a_["frag"] == 0]
        perm_c = main_c[::-1]                                              # umgekehrte Reihenfolge, leicht verschoben
        cxyz_ = ddir / "from_cif.xyz"
        cxyz_.write_text(f"{len(perm_c)}\nfrom cif\n" + "".join(
            f"{mol_c.atoms[k_]['el']} {x_ + 0.01:.6f} {y_:.6f} {z_:.6f}\n" for k_ in perm_c
            for x_, y_, z_ in [mol_c.atoms[k_]["xyz"]]), encoding="utf-8")
        e_cx = t_d.load_xyz(str(cxyz_))
        w.main_tabs.setCurrentWidget(w.xray_tab)
        w.main_tabs.setCurrentWidget(t_d)                                 # Tabwechsel füllt die Vorlagen neu
        t_d.cmb_ref.setCurrentIndex(t_d.cmb_ref.findData("xray:" + w.xray_state["entries"][0]["uid"]))
        t_d.apply_reference()
        labs_ = e_cx["ref_labels"]
        assert sorted(labs_.values()) == sorted(XRAY.atom_name(mol_c.atoms[k_]) for k_ in main_c), labs_
        assert all(labs_[str(i_)].rstrip("\u2071").startswith(mol_c.atoms[k_]["el"]) for i_, k_ in enumerate(perm_c))
        assert w.dft_state["label_style"] == "ref" and t_d.F.ann_items["lbl_a0"].get_text() in labs_.values()
        assert "\u26a0" in t_d.lbl_ref.text()        # zentrosymmetrisch: C1/C1(i) gleichwertig -> gemeldet
        g_cx = t_d.geometry(e_cx)
        ok_ = DFT.match_reference(dict(els=g_cx[1], xyz=g_cx[2], bonds=g_cx[3], order=None),
                                  dict(els=g_cx[1], xyz=g_cx[2], bonds=g_cx[3], order=None))
        assert ok_["order"] == list(range(len(perm_c))) or ok_["ambiguous"]
        t_d.reset_order()
        assert e_cx["order"] is None and not e_cx["ref_labels"]
        w.xray_state.update(entries=[], current=None)
        w.dft_state.update(label_style="el_num")
        report.append("DFT compare mode (number like xyz / CIF template): OK")

        # NMR: Mnova-SVGs (Kern aus der Achse, sonst von Hand), Sortierung, Mnova-Text + Zuordnung, Bruker acqus, Word
        ndir = out / "input" / "nmr"
        sv_h = NMR.make_test_svg(ndir / "t-1H.svg", ("1H",))
        sv_c = NMR.make_test_svg(ndir / "t-13C.svg", ("13C",), peaks=("141.20", "77.16 CDCl3", "21.05"))
        sv_q = NMR.make_test_svg(ndir / "t-HSQC.svg", ("1H", "13C"))
        sv_f = NMR.make_test_svg(ndir / "t-fluor.svg", ("19F",), axis=False)
        assert NMR.svg_info(sv_q)["nuclei"] == ["1H", "13C"] and NMR.svg_info(sv_q)["is2d"]
        assert NMR.svg_info(sv_f)["nuclei"] == [] and NMR.is_nmr_svg(sv_h)
        t_n = w.nmr_tab
        w.main_tabs.setCurrentWidget(t_n)
        for sv_ in (sv_q, sv_f, sv_c, sv_h):
            t_n.load_svg(str(sv_))
        e_f = [s_ for s_ in w.nmr_state["spectra"] if Path(s_["file"]).stem == "t-fluor"][0]
        w.nmr_state["current"] = e_f["uid"]
        t_n.sync_widgets()
        t_n.ed_nuc.setText("19F")
        t_n.spec_edited()
        assert e_f["nuclei"] == ["19F"] and not e_f["nuc_auto"]
        assert [(s_["kind"], s_["nuclei"]) for s_ in t_n.spectra()] == [("1D", ["1H"]), ("1D", ["13C"]), ("1D", ["19F"]),
                                                                        ("HSQC", ["1H", "13C"])], t_n.spectra()
        t_n.ed_text.setPlainText("1H NMR (501 MHz, CDCl3) \u03b4 7.52 \u2013 7.45 (m, 2H), 6.91 (dd, J = 7.0, 1.5 Hz, 1H), "
                                 "2.47 (s, 3H).\n13C NMR (126 MHz, CDCl3) \u03b4 141.2, 21.1.\n"
                                 "19F NMR (471 MHz, CDCl3) \u03b4 -63.2 (s).")
        t_n.text_changed()
        assert t_n.tbl_sig.rowCount() == 6
        t_n.tbl_sig.item(1, 5).setText("Ar-H")                           # Zuordnung je Signal
        t_n.tbl_sig.item(3, 5).setText("Ar-*C*_{q}")
        ln_ = t_n.nmr_lines()
        assert ln_[0][0] == ("\u00b9H NMR (500 MHz, CDCl\u2083): \u03b4 [ppm] = 7.52 \u2013 7.45 (m, 2H), "
                             "6.91 (dd, J = 7.0, 1.5 Hz, 1H, Ar-H), 2.47 (s, 3H)"), ln_[0][0]
        assert "<i>J</i>" in ln_[0][1] and "Ar-<i>C</i><sub>q</sub>" in ln_[1][1] and "\u221263.2" in ln_[2][0]
        assert NMR.svg_peaks(sv_c) == (["141.20", "21.05"], "CDCl3"), NMR.svg_peaks(sv_c)
        w.nmr_state["current"] = [s_ for s_ in t_n.spectra() if s_["nuclei"] == ["13C"]][0]["uid"]
        t_n.peaks_from_svg()                                   # 13C-Zeile aus den Peaklabels (ersetzt die alte)
        assert "13C NMR (126 MHz, CDCl3) \u03b4 141.20, 21.05." in w.nmr_state["text"] and \
            w.nmr_state["text"].count("13C NMR") == 1, w.nmr_state["text"]
        acq_ = NMR.make_test_acqus(ndir / "exp" / "10", nuc="1H", te=300.2)
        (acq_ / "fid").write_bytes(b"\0" * 1000)                          # FID wird nicht ins Projekt kopiert
        e_h = t_n.spectra()[0]
        assert t_n.set_bruker(e_h, str(acq_ / "acqus"))
        assert Path(e_h["bruker"]).parent == proj / "data" and not (Path(e_h["bruker"]) / "fid").exists()
        assert abs(e_h["params"]["o1p"] - 6.175) < 1e-6 and e_h["params"]["d1"] == 1.0
        w.nmr_state["current"] = e_h["uid"]
        t_n.temp_from_bruker()
        assert e_h["temp"] == 300
        cap_q = NMR.caption_html(t_n.spectra()[3], t_n.blocks(), "TEST-1")
        assert cap_q == ("<sup>1</sup>H/<sup>13</sup>C HSQC (500/126 MHz, CDCl<sub>3</sub>, 298 K) of "
                         "<b>TEST-1</b>."), cap_q
        assert any(k_ == "bruker" and c_ is e_h for c_, k_, _ in PRJ._path_slots(w.get_state()))
        from rdkit import Chem as _Chem
        if hasattr(_Chem, "MolToCDXMLBlock"):                               # Struktur als Vektor ins Spektrum
            from rdkit.Chem import AllChem as _AC
            sm_ = _Chem.MolFromSmiles("Cc1ccc(C#C)nc1")             # erfundenes Testmolekül
            _AC.Compute2DCoords(sm_)
            (ndir / "struct.cdxml").write_text(_Chem.MolToCDXMLBlock(sm_), encoding="utf-8")
            w.nmr_state["structure"] = str(ndir / "struct.cdxml")
            fs_ = t_n.figure_svg(t_n.spectra()[0])
            assert fs_ != t_n.spectra()[0]["file"] and 'id="structure_0"' in Path(fs_).read_text(encoding="utf-8")
            t_n.spectra()[1]["struct_on"] = False
            assert t_n.figure_svg(t_n.spectra()[1]) == t_n.spectra()[1]["file"]
            e0_ = t_n.spectra()[0]                                         # gezogen: Lage im Dateinamen
            e0_.update(struct_pos="manual", struct_xy=[0.5, 0.3])
            assert Path(t_n.figure_svg(e0_)).stem.endswith("_500_300_200"), t_n.figure_svg(e0_)
            w.nmr_state["current"] = e0_["uid"]                            # Mausrad über der Struktur: Größe
            t_n.sync_widgets()
            t_n.update_preview()
            if t_n._drag_rect is not None:
                from PySide6.QtCore import QPoint, QPointF
                from PySide6.QtGui import QWheelEvent
                x_, y_, wr_, hr_ = t_n._drag_rect
                pw_, ph_ = t_n._pm_size
                pos_ = QPointF((t_n.lbl_prev.width() - pw_) / 2 + (x_ + wr_ / 2) * pw_,
                               (t_n.lbl_prev.height() - ph_) / 2 + (y_ + hr_ / 2) * ph_)
                ev_ = QWheelEvent(pos_, pos_, QPoint(), QPoint(0, 240), Qt.NoButton, Qt.NoModifier,
                                  Qt.NoScrollPhase, False)
                assert t_n.eventFilter(t_n.lbl_prev, ev_) and w.nmr_state["struct_w"] == 22   # gilt für alle
            e0_.update(struct_pos="auto", struct_xy=None)
            w.nmr_state["struct_w"] = 20.0
        m_ = NMR.occupancy(sv_h)                                           # freie Stelle: verdeckt nichts
        fx_, fy_ = NMR.auto_place(m_, 0.2, 0.5)
        H_, W_ = m_.shape
        x0_, y0_ = int(fx_ * W_), int(fy_ * H_)
        assert m_.any() and not m_[y0_:y0_ + int(0.1 * W_), x0_:x0_ + int(0.2 * W_)].any(), (fx_, fy_)
        nd_ = t_n.export_dialog(out / "nmr" / "nmr_test.docx")
        with _zf.ZipFile(nd_) as z_:
            xml_ = z_.read("word/document.xml").decode("utf-8")
            svgs_ = [n_ for n_ in z_.namelist() if n_.endswith(".svg")]
        assert len(svgs_) == 4 and xml_.count("<w:pageBreakBefore/>") == 1 and xml_.count("SEQ Figure") == 4, len(svgs_)
        assert "HSQC" in xml_ and ">Figure S<" in xml_ and "projekt_test" in xml_
        w.nmr_state.update(spectra=[], current=None, text="", assign={}, structure=None)
        t_n.sync_widgets()
        report.append("NMR (Mnova SVG, nucleus from axis/manual, order, Mnova text + assignment, Bruker acqus, Word): OK")

        # Synthese: Mengenprüfung mit Rundung, Reagenzientabelle, Schema als bearbeitbares ChemDraw-Objekt (CDX)
        proc_ = ("In a flask <b>S1</b> (1.03 g, 5.00 mmol, 1.0 eq.) was dissolved in MeOH (10 mL) and "
                 "K<sub>2</sub>CO<sub>3</sub> (829 mg, 6.00 mmol, 1.2 eq.) was added to give <b>S2</b> "
                 "(525 mg, 4.95 mmol, 99%).")                     # erfundene Werte
        tab_ = {"S1": {"M": 206.36}, "K2CO3": {"M": SYN.formula_mw("K2CO3")}, "S2": {"M": 106.12}}
        res_ = SYN.check_amounts(SYN.parse_amounts(proc_), tab_)
        assert [l_ for l_, _ in res_] == ["ok", "ok", "ok", "ok", "ok"], res_        # 1.03 g -> 4.99 vs. 5.00 mmol: Rundung
        bad_ = SYN.check_amounts(SYN.parse_amounts(proc_.replace("6.00", "6.60").replace("99%", "95%")), tab_)
        assert [l_ for l_, _ in bad_].count("warn") == 3, bad_
        assert abs(SYN.formula_mw("K2CO3") - 138.204) < 0.01 and SYN.molar_mass("206.36") == 206.36
        t_s = w.syn_tab
        w.main_tabs.setCurrentWidget(t_s)
        w.syn_state["reagents"] = [dict(name="S1", M="C12H18OSi", m_mg="1032", role="ref"),
                                   dict(name="K2CO3", M="K2CO3", m_mg="829", role="reagent"),
                                   dict(name="S2", M="C7H6O", m_mg="525", role="product")]
        t_s.sync_widgets()
        assert t_s.tbl.item(1, len(SYN.COLS) + 1).text() == "1.20" and t_s.tbl.item(2, len(SYN.COLS) + 1).text() == "99%"
        r0_ = w.syn_state["reagents"][0]
        assert SYN.amount_text(r0_, dict(SYN.row_amounts(r0_), eq=1.0)) == "(1.032 g, 5.00 mmol, 1.0 eq.)"
        t_s.ed_proc.setHtml(proc_)
        t_s.proc_changed()
        assert "<b>S1</b>" in w.syn_state["proc"] and "<sub>2</sub>" in w.syn_state["proc"], w.syn_state["proc"]
        assert t_s.lbl_check.text().count("\u2713") == 5, t_s.lbl_check.text()
        sdir = out / "input" / "syn"
        sdir.mkdir(parents=True, exist_ok=True)
        if hasattr(_Chem, "MolToCDXMLBlock"):
            (sdir / "scheme.cdxml").write_text(_Chem.MolToCDXMLBlock(sm_), encoding="utf-8")
            from rdkit.Chem import Draw as _Draw
            _Draw.MolToFile(sm_, str(sdir / "scheme.svg"))                   # steht für den ChemDraw-SVG-Export
            (sdir / "scheme.cdx").write_bytes(b"VjCD0100\x04\x03\x02\x01" + b"\0" * 200)   # Platzhalter-CDX
            t_s.set_scheme(str(sdir / "scheme.cdxml"))
            sp_ = t_s.scheme_parts()
            assert sp_["picture"].suffix == ".svg" and sp_["cdx"] is not None and sp_["emf"] is None
            assert Path(w.syn_state["scheme"]).parent == proj / "images"
            sb_ = t_s.si_blocks(out / "syn")
            assert sb_[0] == ("synthesis", "TEST-1" if w.compound.get("name") == "TEST-1" else t_s.compound_name())
            assert [b_[0] for b_ in sb_] == ["synthesis", "ole", "char"], [b_[0] for b_ in sb_]
            sd_ = t_s.export_dialog()
            import olefile
            with _zf.ZipFile(sd_) as z_:
                xml_ = z_.read("word/document.xml").decode("utf-8")
                ole_ = olefile.OleFileIO(z_.read("word/embeddings/oleObject1.bin"))
            assert 'ProgID="ChemDraw.Document.6.0"' in xml_ and ole_.openstream("CONTENTS").read().startswith(b"VjCD0100")
            assert any(k_ == "scheme" and c_ is w.syn_state for c_, k_, _ in PRJ._path_slots(w.get_state()))
            (Path(w.syn_state["scheme"]).with_suffix(".cdx")).unlink()           # ohne CDX: Vektorbild
            assert [b_[0] for b_ in t_s.si_blocks(out / "syn")][1] == "image"
            # Schema ohne ChemDraw-Export: selbst gezeichnet (Pfeil, Bedingungen), Reagenzien aus den Pfeiltexten
            import xml.etree.ElementTree as _ET
            import uvvis_cdxml as CDXD
            fr_ = _ET.fromstring(_Chem.MolToCDXMLBlock(sm_)).find(".//fragment")
            fr2_ = _ET.fromstring(_Chem.MolToCDXMLBlock(sm_)).find(".//fragment")
            for k_, n_ in enumerate(fr2_.iter()):                              # Produkt: 150 pt weiter rechts
                if n_.get("id"):
                    n_.set("id", str(90000 + k_))
                if n_.tag == "n":
                    x_, y_ = map(float, n_.get("p").split())
                    n_.set("p", f"{x_ + 150} {y_}")
            ids_ = {o_.get("id"): str(90000 + k_) for k_, o_ in enumerate(
                _ET.fromstring(_Chem.MolToCDXMLBlock(sm_)).find(".//fragment").iter()) if o_.get("id")}
            for b_ in fr2_.iter("b"):
                b_.set("B", ids_[b_.get("B")])
                b_.set("E", ids_[b_.get("E")])
            ys_ = [float(n_.get("p").split()[1]) for n_ in fr_.iter("n")]
            xs_ = [float(n_.get("p").split()[0]) for n_ in fr_.iter("n")]
            ym_, xr_ = sum(ys_) / len(ys_), max(xs_)
            sch_ = (f'<?xml version="1.0"?><CDXML LineWidth="0.6" BondSpacing="18" LabelSize="10"><page>'
                    + _ET.tostring(fr_, encoding="unicode") + _ET.tostring(fr2_, encoding="unicode")
                    + f'<t id="501" p="{xr_ + 30} {ym_ - 6}"><s face="96" size="10">K2CO3 (1.2 eq.)</s></t>'
                    + f'<t id="502" p="{xr_ + 30} {ym_ + 12}"><s size="9.7">MeOH, rt, 1 h</s></t>'
                    + f'<arrow id="503" Head3D="{xr_ + 120} {ym_} 0" Tail3D="{xr_ + 20} {ym_} 0"/>'
                    + f'<scheme><step ReactionStepReactants="{fr_.get("id")}" ReactionStepProducts="{fr2_.get("id")}" '
                    + 'ReactionStepObjectsAboveArrow="501" ReactionStepObjectsBelowArrow="502"/></scheme></page></CDXML>')
            (sdir / "rxn.cdxml").write_text(sch_, encoding="utf-8")
            svg_, notes_ = CDXD.render_svg(sdir / "rxn.cdxml")
            assert svg_.count(">N</text>") == 2, svg_.count(">N</text>")      # Labels auch ohne <t> (RDKit-CDXML)
            assert not notes_ and "<polygon" in svg_ and ">2</text>" in svg_ and "1.2\u00a0eq.)" in svg_, notes_
            assert CDXD.scheme_reagents(sdir / "rxn.cdxml") == [("K2CO3", "K2CO3", "reagent"), ("MeOH", "CH4O", "other")]
            t_s.set_scheme(str(sdir / "rxn.cdxml"))
            assert t_s.scheme_asset()["drawing_from"] == "cdxml" and t_s.scheme_asset()["pdf"]
            w.syn_state["reagents"] = []
            t_s.rows_from_scheme()
            assert [(r_["name"], r_["role"]) for r_ in w.syn_state["reagents"]] == \
                [("", "ref"), ("K2CO3", "reagent"), ("MeOH", "other"), ("", "product")], w.syn_state["reagents"]
        b_all, _ = w.si_blocks(["char"])
        assert b_all[0][0] == "synthesis" and any(b_[0] == "char" and "S1" in b_[1] for b_ in b_all)
        w.syn_state.update(name="", scheme=None, proc="", reagents=[])
        t_s.sync_widgets()
        report.append("Synthesis (amount check with rounding, reagent table, ChemDraw object, SI): OK")
        w.dft_state.update(entries=[], current=None)
        t_d.sync_widgets()
        t_d.render()

        # SI-Zeilen: Elementaranalyse, UV-Vis, SI-Abschnitt als Word-Datei
        comp_ = SI.ea_composition("C18H20N2")
        assert [round(comp_[e_], 2) for e_ in "CHN"] == [81.78, 7.63, 10.60], comp_
        assert SI.ea_formula_label("C18H20N2", "0.5 CH2Cl2") == "C18H20N2·0.5CH2Cl2"
        assert abs(SI.ea_composition("C18H20N2", "0.5 CH2Cl2")["C"] - 72.42) < 0.01
        pl_, ht_, nt_ = SI.ea_line("C18H20N2", "", {"C": "81.52", "H": "7.70", "N": "10.10"})
        assert pl_ == "Elemental analysis calcd (%) for C18H20N2: C 81.78, H 7.63, N 10.60; found: C 81.52, " \
                      "H 7.70, N 10.10", pl_
        assert [e_ for e_, _ in nt_] == ["N"] and "<sub>18</sub>" in ht_       # N weicht um 0.50 % ab
        cfg_uv = dict(core.DEFAULTS, conc_unit="mM")
        res_uv = [{"lam": 412.0, "fit": {"n": 5}, "eps": 12345.6, "eps_err": 120.0},
                  {"lam": 298.0, "fit": {"n": 5}, "eps": 24100.2, "eps_err": 0.4},
                  {"lam": 350.0, "fit": {"n": 5}, "eps": 5000.0, "eps_err": 50.0, "hidden": True},
                  {"lam": 500.0, "fit": None}]
        uvp_, uvh_ = SI.uvvis_line(res_uv, cfg_uv, "THF")
        assert uvp_ == "UV-Vis (THF): λmax: 298 nm (ε = 24100 cm-1 M-1), 412 nm (ε = 12350 cm-1 M-1)", uvp_
        assert "<b>λ<sub>max</sub></b>" in uvh_ and "cm<sup>-1</sup>" in uvh_, uvh_
        assert SI.join_lines([("m.p. 67 °C", "m.p. 67 °C"), (uvp_, uvh_)])[0].endswith("M-1).")
        w.compound = {"name": "TEST-1", "formula": "C18H20N2", "mp": "67 °C"}
        w.irms_state.update(ea_found={"C": "81.70", "H": "7.60", "N": "10.55"}, ea_solvate="")
        w.irms_tab.sync_widgets()
        blocks_, missing_ = w.si_blocks()
        para_ = [b_[1] for b_ in blocks_ if b_[0] == "char"][0]
        assert blocks_[0] == ("synthesis", "TEST-1"), blocks_[0]
        assert para_.startswith("<b>m.p.</b> 67 °C; ") and "<b>Elemental analysis</b>" in para_, para_
        assert "IR" in missing_ and "m.p." not in missing_, missing_
        w.export_si()
        docx_ = Path(w.project_dir) / "exports" / "si" / "TEST-1_SI.docx"
        with _zf.ZipFile(docx_) as z_:
            xml_ = z_.read("word/document.xml").decode("utf-8")
            assert 'w:val="15"' in z_.read("word/settings.xml").decode("utf-8")   # kein Kompatibilitätsmodus
        assert "Elemental analysis" in xml_ and 'w:val="subscript"' in xml_ and "TEST-1" in xml_
        assert "67\u00a0°C" in xml_ and "C\u00a081." in xml_                  # geschützte Leerzeichen
        assert SI.nbsp("298 nm (ε = 5 cm-1 M-1)") == "298\u00a0nm (ε\u00a0=\u00a05\u00a0cm-1\u00a0M-1)"
        assert "Courier New" not in xml_                                   # ohne DFT-Rechnung kein DFT-Teil
        # Teile wählbar: DFT zweispaltig hinter der Charakterisierung, ohne Charakterisierung nur Überschrift
        w.dft_tab.load_xyz(str(orf_))
        assert {k_: v_ for k_, v_ in w.si_available().items() if k_ not in ("figs", "methods")} == \
            {"char": True, "nmr": False, "xray": False, "dft": True}
        b_dft, _ = w.si_blocks(["dft"])
        assert b_dft[0] == ("heading", "TEST-1") and ("columns", 2) in b_dft and not any(b_[0] == "char" for b_ in b_dft)
        w.export_si(["char", "dft"])
        with _zf.ZipFile(docx_) as z_:
            xml_ = z_.read("word/document.xml").decode("utf-8")
        assert "Elemental analysis" in xml_ and "Courier New" in xml_ and xml_.count("<w:sectPr>") == 2
        assert xml_.index("Elemental analysis") < xml_.index('w:cols w:num="2"')   # einspaltig, dann zweispaltig
        w.dft_state.update(entries=[], current=None)
        w.dft_tab.sync_widgets()
        w.irms_state.update(ea_found={}, ea_solvate="")
        w.compound = {}
        report.append("SI lines (EA, UV-Vis) + SI section as Word (parts selectable): OK")

        # Release Notes aus den Changelogs (nur im Quellcode; das Skript liegt nicht in der gebauten App)
        rn_path = Path(__file__).resolve().parent / "release_notes.py"
        if rn_path.exists() and not getattr(sys, "frozen", False):
            import importlib.util
            spec_ = importlib.util.spec_from_file_location("release_notes", rn_path)
            RN = importlib.util.module_from_spec(spec_)
            spec_.loader.exec_module(RN)
            en_ = (rn_path.parent / "CHANGELOG.md").read_text(encoding="utf-8")
            de_ = (rn_path.parent / "CHANGELOG.de.md").read_text(encoding="utf-8")
            txt_ = RN.notes(APP_VERSION, en_, de_)
            assert txt_.startswith("## English") and "## Deutsch" in txt_ and f"{APP_NAME}-windows-x64.zip" in txt_
            assert f"## {APP_VERSION}" not in txt_ and "\n## 1." not in txt_          # nur dieser Abschnitt
            try:
                RN.notes("0.0.1", en_, de_)
                raise AssertionError("fehlende Version nicht erkannt")
            except ValueError:
                pass
            report.append("release notes from changelogs: OK")

        # Update-Prüfung (ohne Netz: Versionsvergleich, API-Antwort, Meldungstext)
        assert UPD.parse_version("v1.10.0") > UPD.parse_version("1.9.9") and UPD.parse_version("x") == ()
        assert UPD.is_newer("1.4.1", "1.4.0") and not UPD.is_newer("1.4", "1.4.0") and not UPD.is_newer("", "1.0")
        lbl = UPD.platform_label()
        fake = {"tag_name": "v99.0.0", "html_url": "https://example.invalid/r",
                "assets": [{"name": f"Elucida-{x}.zip"} for x in ("windows-x64", "macos-arm64", "macos-x64")]
                + [{"name": "Elucida-linux-x64.tar.gz"}]}
        info = UPD.parse_release(fake)
        assert info["version"] == "99.0.0" and lbl in info["asset"], info
        assert "99.0.0" in w.update_text(info, manual=False)
        assert w.update_text(dict(info, version=APP_VERSION), manual=False) is None
        assert APP_VERSION in w.update_text(dict(info, version=APP_VERSION), manual=True)
        report.append("update check: OK")

        # Projekt automatisch gespeichert -> in neuem Fenster öffnen (Dateien als Kopie, Pfade relativ)
        w.ov_state["entries"][0]["file"] = w.import_file(w.ov_state["entries"][0]["file"], "data")
        w.push_undo()
        pump(1.5)                                      # automatisches Speichern abwarten
        w.maybe_save()
        txt = (proj / PRJ.PROJECT_FILE).read_text(encoding="utf-8")
        assert str(out / "input") not in txt.replace(str(proj), ""), "absolute Pfade im Projekt"
        assert (proj / "data" / csv.name).exists() and (proj / "images" / cdxml.name).exists()
        w2 = MainWindow()
        assert w2.open_project(proj, ask=False)
        pump()
        assert set(w2.groups) == {"TEST-A", "TEST-B"} and w2.series["TEST-A"]["mw"] == 95.1, w2.groups
        assert len(w2.ov_state.get("entries", [])) == 2 and w2.fl_state.get("em_file"), "Tabs nicht geladen"
        report.append("project save/load: OK")

        # Undo/Redo
        before = w2.glob["cutoff"]
        w2.sp_cut.setValue(before + 0.5)
        pump(0.8)
        assert w2.glob["cutoff"] == before + 0.5
        w2.undo()
        pump()
        assert abs(w2.glob["cutoff"] - before) < 1e-9, w2.glob["cutoff"]
        w2.redo()
        pump()
        assert abs(w2.glob["cutoff"] - (before + 0.5)) < 1e-9, w2.glob["cutoff"]
        report.append("undo/redo: OK")
        # Startbildschirm/Tabs je nach Projekt; Beispielprojekt (künstliche Daten) füllt die Tabs
        assert w.center.currentIndex() == 1
        d_demo = w.open_demo(out / "demo")
        assert d_demo and Path(w.project_dir) == out / "demo" and w.center.currentIndex() == 1, (
            d_demo, w.project_dir, w.center.currentIndex(), w.log.toPlainText()[-3000:])
        assert w.groups and w.cv_state["files"] and w.xray_state["entries"] and w.dft_state["entries"]
        assert w.syn_state["proc"] and w.nmr_tab.spectra()
        report.append("welcome page + example project: OK")
        for lang in ("en", "de"):
            w.switch_lang(lang)
        report.append("OK")
    except Exception:
        ok = False
        report.append("FAILED\n" + traceback.format_exc())
    (out / "selftest_report.txt").write_text("\n".join(report), encoding="utf-8")
    try:                                   # Windows-Konsole (cp1252) kann z. B. kein λ: ersetzen statt abstürzen
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass
    print("\n".join(report))
    return 0 if ok else 1


# ----------------------------------------------------------------------------
def install_crash_handler():
    def hook(t, v, tb):
        msg = "".join(traceback.format_exception(t, v, tb))
        try:
            (IMG.app_data_dir() / "crash.log").write_text(msg, encoding="utf-8")
        except Exception:
            pass
        if QApplication.instance() and os.environ.get("QT_QPA_PLATFORM") != "offscreen":
            QMessageBox.critical(None, "Error", msg[-2000:])
        sys.__excepthook__(t, v, tb)
    sys.excepthook = hook


def main():
    if "--selftest" in sys.argv:
        i = sys.argv.index("--selftest")
        outdir = sys.argv[i + 1] if len(sys.argv) > i + 1 and not sys.argv[i + 1].startswith("--") \
            else "selftest_out"
        sys.exit(selftest(outdir, "--selftest-isnet" in sys.argv))
    install_crash_handler()
    QApplication.setApplicationName(APP_NAME)
    QApplication.setApplicationVersion(APP_VERSION)
    app = QApplication(sys.argv)
    w = MainWindow()
    w.show()
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    if args and Path(args[0]).exists():
        w.open_csv(args[0])          # Datei per Doppelklick / Drag auf das Programmsymbol
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
