"""uvvis_xray.py - Kristallstrukturen aus CIF: ORTEP-Abbildung, Bindungen/Winkel/Diederwinkel, SI-Tabellen.

CIF:   eigener Leser (CIF 1.1-Syntax: data_-Blöcke, loop_, Strings in '…'/"…", Textfelder ;…;), Zahlen bleiben
       zusätzlich als Text erhalten, damit Werte mit Standardabweichung exakt wie in der CIF erscheinen.
Modell: fraktionale -> kartesische Koordinaten (a entlang x, b in der xy-Ebene), U_ij -> U_kart = M·N·U·N·Mᵀ.
       Bindungen aus _geom_bond (mit Symmetriecode n_klm), sonst aus kovalenten Radien. Molekül über die Bindungen
       und die Symmetrieoperationen vervollständigen; erzeugte Atome tragen einen Zusatz (C1ⁱ) mit Fußnote.
ORTEP: Ellipsoid mit Wahrscheinlichkeit p: xᵀU⁻¹x = k², k² = χ²-Quantil (3 Freiheitsgrade); 50 % -> k = 1.5382.
       Projektion auf den Bildschirm = 2×2-Block von R·U·Rᵀ (exakter Umriss), dazu die drei Hauptachsen-Ellipsen
       (vordere Hälften) wie in Olex2.
esd:   Standardabweichungen nur aus der CIF (_geom_bond/_geom_angle/_geom_torsion), nie geschätzt.
"""
from __future__ import annotations

import csv
import math
import re
from pathlib import Path

import numpy as np

# ---------------------------------------------------------------------------- CIF lesen
_NUM = re.compile(r"^([+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?)(?:\((\d+)\))?$")


def _tokens(text):
    """-> (Token, in Anführungszeichen?) – Textfelder ;…; als ein Token."""
    lines = text.splitlines()
    i, n_lines = 0, len(lines)
    while i < n_lines:
        line = lines[i]
        if line.startswith(";"):
            buf = [line[1:]]
            i += 1
            while i < n_lines and not lines[i].startswith(";"):
                buf.append(lines[i])
                i += 1
            i += 1
            yield "\n".join(buf).strip("\n"), True
            continue
        p, n = 0, len(line)
        while p < n:
            c = line[p]
            if c.isspace():
                p += 1
                continue
            if c == "#":
                break
            if c in "'\"":
                j = p + 1
                while j < n and not (line[j] == c and (j + 1 == n or line[j + 1].isspace())):
                    j += 1
                yield line[p + 1:j], True
                p = j + 1
                continue
            j = p
            while j < n and not line[j].isspace():
                j += 1
            yield line[p:j], False
            p = j
        i += 1


class Block:
    """Ein data_-Block: Einzelwerte und Schleifen, Tags in Kleinbuchstaben."""

    def __init__(self, name):
        self.name = name
        self.items = {}          # tag -> Text
        self.loops = []          # [{tag: [Werte]}]
        self.loop_of = {}        # tag -> Schleife

    def get(self, *tags):
        """Erster vorhandener Tag (Aliase DDL1/DDL2); '?' und '.' gelten als fehlend."""
        for t in tags:
            v = self.items.get(t.lower())
            if v is None and t.lower() in self.loop_of:
                col = self.loop_of[t.lower()][t.lower()]
                v = col[0] if len(col) == 1 else None
            if v is not None and v not in ("?", "."):
                return v
        return None

    def num(self, *tags):
        return num(self.get(*tags))

    def column(self, *tags):
        for t in tags:
            lp = self.loop_of.get(t.lower())
            if lp is not None:
                return lp[t.lower()]
            if t.lower() in self.items:
                return [self.items[t.lower()]]
        return None

    def rows(self, first, *more):
        """Zeilen der Schleife mit Tag 'first' als dicts (fehlende Spalten -> None)."""
        lp = self.loop_of.get(first.lower())
        if lp is None:
            if first.lower() not in self.items:
                return []
            lp = {t.lower(): [self.items[t.lower()]] for t in (first,) + more if t.lower() in self.items}
        n = len(lp[first.lower()])
        return [{k: v[i] for k, v in lp.items()} for i in range(n)]


def parse_cif_text(text) -> list[Block]:
    toks = list(_tokens(text))
    blocks, cur, i, n = [], None, 0, len(toks)

    def is_key(t, q):
        return not q and (t.startswith("_") or t.lower() == "loop_" or t.lower().startswith(("data_", "save_",
                                                                                              "global_")))
    while i < n:
        tok, quoted = toks[i]
        low = tok.lower()
        if not quoted and low.startswith("data_"):
            cur = Block(tok[5:])
            blocks.append(cur)
            i += 1
            continue
        if cur is None:
            cur = Block("global")
            blocks.append(cur)
        if not quoted and low == "loop_":
            i += 1
            tags = []
            while i < n and not toks[i][1] and toks[i][0].startswith("_"):
                tags.append(toks[i][0].lower())
                i += 1
            vals = []
            while i < n and not is_key(*toks[i]):
                vals.append(toks[i][0])
                i += 1
            if tags:
                m = len(vals) // len(tags)               # unvollständige letzte Zeile verwerfen
                lp = {t: vals[k::len(tags)][:m] for k, t in enumerate(tags)}
                cur.loops.append(lp)
                for t in tags:
                    cur.loop_of[t] = lp
            continue
        if not quoted and tok.startswith("_"):
            if i + 1 < n and not is_key(*toks[i + 1]):
                cur.items[low] = toks[i + 1][0]
                i += 2
            else:
                i += 1
            continue
        i += 1
    return blocks


def read_cif(path) -> list[Block]:
    raw = Path(path).read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    return parse_cif_text(text)


def structure_blocks(path) -> list[Block]:
    """Nur Blöcke mit Atomen (Olex2/SHELXL: ein Block je Struktur)."""
    return [b for b in read_cif(path) if b.column("_atom_site_fract_x")]


def num(s):
    """'1.392(3)' -> 1.392; '?'/'.'/None -> None."""
    if s is None:
        return None
    m = _NUM.match(str(s).strip())
    return float(m.group(1)) if m else None


def esd_parts(s):
    """'1.392(3)' -> ('1.392', '3'); ohne esd -> ('1.392', None)."""
    m = _NUM.match(str(s).strip())
    if not m:
        return str(s), None
    return m.group(1), m.group(2)


def minus(s):
    """Echtes Minuszeichen statt Bindestrich (Tabellen, Diff.-Dichte)."""
    return re.sub(r"(^|[\s(/\[])-(?=[\d.])", "\\1\u2212", str(s))


# ---------------------------------------------------------------------------- Symmetrie
def parse_symop(s):
    """'-x+1/2, y+1/2, -z+1/2' -> (R 3×3 int, t 3)."""
    parts = [p.strip().lower().replace(" ", "") for p in str(s).split(",")]
    if len(parts) != 3:
        raise ValueError(f"Symmetrieoperation unlesbar: {s}")
    R = np.zeros((3, 3), int)
    t = np.zeros(3)
    for i, p in enumerate(parts):
        for term in re.findall(r"[+-]?[^+-]+", p):
            sign = -1 if term.startswith("-") else 1
            term = term.lstrip("+-")
            m = re.match(r"^(\d*\.?\d*(?:/\d+)?)\*?([xyz])$", term)
            if m:
                coef = m.group(1)
                c = float(eval_frac(coef)) if coef else 1.0
                R[i, "xyz".index(m.group(2))] += int(round(sign * c))
            else:
                t[i] += sign * eval_frac(term)
    return R, t


def eval_frac(s):
    if "/" in s:
        a, b = s.split("/")
        return float(a) / float(b)
    return float(s)


IDENT = (np.eye(3, dtype=int), np.zeros(3))


def tr_key(T):
    R, t = T
    return tuple(int(v) for v in R.ravel()) + tuple(round(float(v), 4) + 0.0 for v in t)


def compose(A, B):
    """(A∘B)(x) = A(B(x))."""
    return A[0] @ B[0], A[0] @ B[1] + A[1]


def inverse(A):
    Ri = np.rint(np.linalg.inv(A[0])).astype(int)
    return Ri, -Ri @ A[1]


def apply(T, frac):
    return T[0] @ frac + T[1]


def is_identity(T):
    return tr_key(T) == tr_key(IDENT)


def op_string(T):
    """(R, t) -> '1−x, 1−y, 1−z' (wie Fußnoten in SI-Tabellen)."""
    out = []
    for i in range(3):
        terms = []
        tv = float(T[1][i])
        if abs(tv) > 1e-6:
            fr = _frac_str(tv)
            terms.append(fr)
        for j, ax in enumerate("xyz"):
            c = int(T[0][i, j])
            if c:
                s = ("−" if c < 0 else ("+" if terms else "")) + (str(abs(c)) if abs(c) != 1 else "") + ax
                terms.append(s)
        out.append("".join(terms) or "0")
    return ", ".join(out)


def _frac_str(v):
    for den in (1, 2, 3, 4, 6, 8, 12):
        if abs(v * den - round(v * den)) < 1e-4:
            nume = int(round(v * den))
            s = str(abs(nume)) if den == 1 else f"{abs(nume)}/{den}"
            return ("−" if nume < 0 else "") + s
    return f"{v:g}".replace("-", "−")


def symcode(code, ops):
    """SHELXL-Code '3_665' -> Transformation (Operation 3, Translation +1,+1,0); '.' -> Identität."""
    if code in (None, ".", "?", "") or code == "1_555":
        return IDENT
    m = re.match(r"^(\d+)(?:_(\d+))?$", str(code))
    if not m:
        raise ValueError(f"Symmetriecode unlesbar: {code}")
    n, tt = int(m.group(1)), m.group(2) or "555"            # '3' = Operation 3 ohne Translation
    if len(tt) % 3:
        raise ValueError(f"Symmetriecode unlesbar: {code}")
    k = len(tt) // 3                             # 555 = keine Translation; zweistellig (z. B. 1_105555) -> 55
    shift = np.array([int(tt[i * k:(i + 1) * k]) for i in range(3)], float) - (5 if k == 1 else 55)
    R, t = ops[n - 1]
    return R, t + shift


ROMAN = ["i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x", "xi", "xii", "xiii", "xiv", "xv", "xvi"]
SUP = str.maketrans("0123456789+−-=()nivx", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁻⁼⁽⁾ⁿⁱᵛˣ")
SUB = str.maketrans("0123456789+−-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₋₌₍₎")


def roman(i):
    return ROMAN[i] if i < len(ROMAN) else str(i + 1)


# ---------------------------------------------------------------------------- Elemente
# Olex2-ähnliche Standardfarben (Beispielbilder), sonst Jmol/CPK
OLEX2_COLORS = {"C": "#8C8C8C", "H": "#FFFFFF", "N": "#1F3FE0", "O": "#E01010", "F": "#C8E632", "Cl": "#3CD23C",
                "S": "#E6A03C", "B": "#FFB5B5", "Br": "#A62929", "P": "#FF8000", "I": "#940094", "Si": "#F0C8A0"}
METAL_COLOR = "#0A7D8C"
JMOL = {"He": "#D9FFFF", "Li": "#CC80FF", "Be": "#C2FF00", "Ne": "#B3E3F5", "Na": "#AB5CF2", "Mg": "#8AFF00",
        "Al": "#BFA6A6", "Ar": "#80D1E3", "K": "#8F40D4", "Ca": "#3DFF00", "Se": "#FFA100", "Kr": "#5CB8D1",
        "Rb": "#702EB0", "Sr": "#00FF00", "Te": "#D47A00", "Xe": "#429EB0", "Cs": "#57178F", "Ba": "#00C900",
        "As": "#BD80E3", "Ge": "#668F8F", "Sb": "#9E63B5", "Ga": "#C28F8F", "In": "#A67573", "Sn": "#668080",
        "Tl": "#A6544D", "Pb": "#575961", "Bi": "#9E4FB5", "D": "#FFFFC0"}
METALS = set("Sc Ti V Cr Mn Fe Co Ni Cu Zn Y Zr Nb Mo Tc Ru Rh Pd Ag Cd Hf Ta W Re Os Ir Pt Au Hg "
             "La Ce Pr Nd Sm Eu Gd Tb Dy Ho Er Tm Yb Lu U Th".split())
# kovalente Radien (Cordero et al. 2008), Å
COV = {"H": 0.31, "D": 0.31, "He": 0.28, "Li": 1.28, "Be": 0.96, "B": 0.84, "C": 0.76, "N": 0.71, "O": 0.66,
       "F": 0.57, "Ne": 0.58, "Na": 1.66, "Mg": 1.41, "Al": 1.21, "Si": 1.11, "P": 1.07, "S": 1.05, "Cl": 1.02,
       "Ar": 1.06, "K": 2.03, "Ca": 1.76, "Sc": 1.70, "Ti": 1.60, "V": 1.53, "Cr": 1.39, "Mn": 1.39, "Fe": 1.32,
       "Co": 1.26, "Ni": 1.24, "Cu": 1.32, "Zn": 1.22, "Ga": 1.22, "Ge": 1.20, "As": 1.19, "Se": 1.20, "Br": 1.20,
       "Kr": 1.16, "Rb": 2.20, "Sr": 1.95, "Y": 1.90, "Zr": 1.75, "Nb": 1.64, "Mo": 1.54, "Tc": 1.47, "Ru": 1.46,
       "Rh": 1.42, "Pd": 1.39, "Ag": 1.45, "Cd": 1.44, "In": 1.42, "Sn": 1.39, "Sb": 1.39, "Te": 1.38, "I": 1.39,
       "Xe": 1.40, "Cs": 2.44, "Ba": 2.15, "La": 2.07, "Hf": 1.75, "Ta": 1.70, "W": 1.62, "Re": 1.51, "Os": 1.44,
       "Ir": 1.41, "Pt": 1.36, "Au": 1.36, "Hg": 1.32, "Tl": 1.45, "Pb": 1.46, "Bi": 1.48}
BOND_TOL = 0.45


def default_color(el):
    if el in OLEX2_COLORS:
        return OLEX2_COLORS[el]
    if el in METALS:
        return METAL_COLOR
    return JMOL.get(el, "#FF1493")


def element_of(type_symbol, label):
    s = type_symbol or label or "C"
    m = re.match(r"^([A-Za-z]{1,2})", s)
    el = m.group(1) if m else "C"
    el = el[0].upper() + el[1:].lower()
    if el not in COV and el[0] in COV:
        el = el[0]
    return el


# ---------------------------------------------------------------------------- Zelle / ADP
def cell_of(b: Block):
    vals = [b.num(f"_cell_length_{x}") for x in "abc"] + [b.num(f"_cell_angle_{x}") for x in ("alpha", "beta", "gamma")]
    if any(v is None for v in vals):
        raise ValueError("Zellparameter fehlen")
    return vals


def ortho_matrix(a, b, c, al, be, ga):
    """Fraktional -> kartesisch (Å), a entlang x, b in der xy-Ebene."""
    ca, cb, cg = (math.cos(math.radians(v)) for v in (al, be, ga))
    sg = math.sin(math.radians(ga))
    v = a * b * c * math.sqrt(max(1e-12, 1 - ca * ca - cb * cb - cg * cg + 2 * ca * cb * cg))
    return np.array([[a, b * cg, c * cb],
                     [0, b * sg, c * (ca - cb * cg) / sg],
                     [0, 0, v / (a * b * sg)]])


def u_cart(M, U):
    """U_ij (CIF, bezogen auf a*, b*, c*) -> kartesische U-Matrix."""
    Minv = np.linalg.inv(M)
    nstar = np.diag(np.linalg.norm(Minv, axis=1))   # Zeilen von M⁻¹ = reziproke Vektoren -> |a*|, |b*|, |c*|
    A = M @ nstar
    return A @ U @ A.T


def chi2_scale(p):
    """k mit P(χ²₃ ≤ k²) = p (Ellipsoid mit Aufenthaltswahrscheinlichkeit p); 0.5 -> 1.5382."""
    p = min(max(float(p), 0.01), 0.999)

    def cdf(k):        # χ²-Verteilung mit 3 Freiheitsgraden, x = k²
        return math.erf(k / math.sqrt(2)) - math.sqrt(2 / math.pi) * k * math.exp(-k * k / 2)
    lo, hi = 0.0, 10.0
    for _ in range(80):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if cdf(mid) < p else (lo, mid)
    return (lo + hi) / 2


# ---------------------------------------------------------------------------- Struktur
class Structure:
    """Asymmetrische Einheit, Symmetrie, Bindungen (ASU-bezogen) und geom-Tabellen eines Blocks."""

    def __init__(self, block: Block):
        self.block = block
        self.name = block.name
        self.cell = cell_of(block)
        self.M = ortho_matrix(*self.cell)
        ops = block.column("_space_group_symop_operation_xyz", "_symmetry_equiv_pos_as_xyz") or ["x, y, z"]
        self.ops = [parse_symop(o) for o in ops]
        self.atoms = []
        aniso = {}
        for r in block.rows("_atom_site_aniso_label"):
            try:
                u = [num(r.get(f"_atom_site_aniso_u_{k}")) for k in ("11", "22", "33", "23", "13", "12")]
                if None in u:
                    bb = [num(r.get(f"_atom_site_aniso_b_{k}")) for k in ("11", "22", "33", "23", "13", "12")]
                    if None in bb:
                        continue
                    u = [v / (8 * math.pi ** 2) for v in bb]
                u11, u22, u33, u23, u13, u12 = u
                aniso[r["_atom_site_aniso_label"]] = np.array([[u11, u12, u13], [u12, u22, u23], [u13, u23, u33]])
            except (KeyError, TypeError):
                continue
        for r in block.rows("_atom_site_label"):
            lab = r["_atom_site_label"]
            xyz = [num(r.get(f"_atom_site_fract_{k}")) for k in "xyz"]
            if None in xyz:
                continue
            el = element_of(r.get("_atom_site_type_symbol"), lab)
            uiso = num(r.get("_atom_site_u_iso_or_equiv"))
            if uiso is None and num(r.get("_atom_site_b_iso_or_equiv")) is not None:
                uiso = num(r.get("_atom_site_b_iso_or_equiv")) / (8 * math.pi ** 2)
            part = r.get("_atom_site_disorder_group")
            try:
                part = int(part) if part not in (None, ".", "?") else 0
            except ValueError:
                part = 0
            occ = num(r.get("_atom_site_occupancy"))
            U = aniso.get(lab)
            self.atoms.append(dict(label=lab, el=el, frac=np.array(xyz, float), uiso=uiso or 0.02,
                                   U=u_cart(self.M, U) if U is not None else None,
                                   occ=1.0 if occ is None else occ, part=part,
                                   assembly=r.get("_atom_site_disorder_assembly") or "",
                                   calc=(r.get("_atom_site_calc_flag") or "").lower()))
        self.by_label = {a["label"]: i for i, a in enumerate(self.atoms)}
        self.geom_bond = self._geom("_geom_bond_atom_site_label_1", 2, "_geom_bond_distance", "_geom_bond")
        self.geom_angle = self._geom("_geom_angle_atom_site_label_1", 3, "_geom_angle", "_geom_angle")
        self.geom_torsion = self._geom("_geom_torsion_atom_site_label_1", 4, "_geom_torsion", "_geom_torsion")

    def _geom(self, first, n, value_tag, prefix):
        out = []
        for r in self.block.rows(first):
            labs = [r.get(f"{prefix}_atom_site_label_{k + 1}") for k in range(n)]
            val = r.get(value_tag)
            if None in labs or val in (None, "?", "."):
                continue
            codes = [r.get(f"{prefix}_site_symmetry_{k + 1}") for k in range(n)]
            if n == 3 and not codes[1]:
                codes[1] = "."
            try:
                trs = [symcode(c, self.ops) for c in codes]
            except (ValueError, IndexError):
                continue
            out.append(dict(labels=labs, trs=trs, value=val))
        return out

    def cart(self, frac):
        return self.M @ frac

    # -- Fehlordnung --------------------------------------------------------------------
    def visible_asu(self, mode="main"):
        """Indizes der gezeigten ASU-Atome: 'main' = PART 0 + je Gruppe der Teil mit der größten Besetzung."""
        if mode == "all":
            return list(range(len(self.atoms)))
        if mode == "minor":
            main = set(self.visible_asu("main"))
            return [i for i, a in enumerate(self.atoms) if a["part"] == 0 or i not in main]
        groups = {}
        for i, a in enumerate(self.atoms):
            if a["part"]:
                groups.setdefault(a["assembly"], {}).setdefault(abs(a["part"]), []).append(a["occ"])
        keep = {}
        for asm, parts in groups.items():
            keep[asm] = max(sorted(parts), key=lambda p: (np.mean(parts[p]) - 1e-9 * p))
        return [i for i, a in enumerate(self.atoms) if not a["part"] or keep.get(a["assembly"]) == abs(a["part"])]

    def parts(self):
        return sorted({a["part"] for a in self.atoms if a["part"]})

    # -- Bindungen in ASU-Form: (i, j, T) = Atom i (Identität) gebunden an T(Atom j) ----------
    def asu_bonds(self, visible):
        vis = set(visible)
        out = []
        if self.geom_bond:
            have_h = set()
            for g in self.geom_bond:
                i, j = (self.by_label.get(x) for x in g["labels"])
                if i is None or j is None or i not in vis or j not in vis:
                    continue
                if not self._part_ok(i, j):
                    continue
                Ti = g["trs"][0]
                T = compose(inverse(Ti), g["trs"][1])
                out.append((i, j, T))
                for k in (i, j):
                    if self.atoms[k]["el"] in ("H", "D"):
                        have_h.add(k)
            # H ohne Eintrag in _geom_bond (BOND $H fehlte): an das nächste Nicht-H-Atom hängen
            missing = [k for k in vis if self.atoms[k]["el"] in ("H", "D") and k not in have_h]
            if missing:
                cov = self._covalent(vis)
                out += [b for b in cov if (b[0] in missing) != (b[1] in missing)
                        and ({self.atoms[b[0]]["el"], self.atoms[b[1]]["el"]} - {"H", "D"})]
        else:
            out = self._covalent(vis)
        full = []
        for i, j, T in out:                          # beide Richtungen
            full.append((i, j, T))
            full.append((j, i, inverse(T)))
        return full

    def _part_ok(self, i, j):
        a, b = self.atoms[i], self.atoms[j]
        return not (a["part"] and b["part"] and a["part"] != b["part"] and a["assembly"] == b["assembly"])

    def _covalent(self, vis):
        idx = sorted(vis)
        if not idx:
            return []
        F = np.array([self.atoms[i]["frac"] for i in idx])
        rad = np.array([COV.get(self.atoms[i]["el"], 1.5) for i in idx])
        isH = np.array([self.atoms[i]["el"] in ("H", "D") for i in idx])
        cut = rad[:, None] + rad[None, :] + BOND_TOL
        offs = np.array([(a, b, c) for a in (-1, 0, 1) for b in (-1, 0, 1) for c in (-1, 0, 1)], float)
        out, seen = [], set()
        for op in self.ops:
            img = (op[0] @ F.T).T + op[1]                       # Bilder aller Atome j
            D = img[None, :, :] - F[:, None, :]                 # (i, j, 3)
            base = -np.rint(D)
            for o in offs:
                Dc = (D + base + o) @ self.M.T
                d = np.linalg.norm(Dc, axis=2)
                hit = np.argwhere((d > 0.4) & (d < cut) & ~(isH[:, None] & isH[None, :]))
                for a, b in hit:
                    i, j = idx[a], idx[b]
                    if not self._part_ok(i, j):
                        continue
                    T = (op[0], op[1] + base[a, b] + o)
                    key = (i, j, tr_key(T))
                    rkey = (j, i, tr_key(inverse(T)))
                    if key in seen or rkey in seen:
                        continue
                    if i == j and is_identity(T):
                        continue
                    seen.add(key)
                    out.append((i, j, T))
        return out


class Molecule:
    """Vervollständigtes Modell: gezeigte Atome (mit Transformation zur ASU), Bindungen, Fragmente."""

    def __init__(self, S: Structure, disorder="main", max_factor=8):
        self.S = S
        vis = S.visible_asu(disorder)
        nbrs = {}
        for i, j, Tr in S.asu_bonds(vis):
            nbrs.setdefault(i, []).append((j, Tr))
        self.atoms, self.bonds = [], set()
        index, by_label = {}, {}
        self.truncated = False
        limit = max_factor * max(1, len(vis)) + 50

        def place(i, T):
            frac = apply(T, S.atoms[i]["frac"])
            key = (i,) + tuple(np.round(frac, 3) + 0.0)
            if key in index:
                return index[key], False
            R = T[0]
            a = S.atoms[i]
            U = None
            if a["U"] is not None:                   # U_kart dreht mit dem kartesischen Teil der Operation
                Rc = S.M @ R @ np.linalg.inv(S.M)
                U = Rc @ a["U"] @ Rc.T
            generated = not (np.array_equal(R, np.eye(3, dtype=int)) and i not in by_label)
            k = len(self.atoms)
            self.atoms.append(dict(asu=i, label=a["label"], el=a["el"], T=T, frac=frac, xyz=S.cart(frac), U=U,
                                   uiso=a["uiso"], part=a["part"], generated=generated and not is_identity(T)))
            index[key] = k
            by_label.setdefault(i, []).append(k)
            return k, True

        for start in vis:
            if start in by_label:
                continue
            k0, _ = place(start, IDENT)
            queue = [k0]
            while queue:
                k = queue.pop(0)
                at = self.atoms[k]
                for j, Trel in nbrs.get(at["asu"], []):
                    if len(self.atoms) >= limit:
                        self.truncated = True
                        break
                    kj, new = place(j, compose(at["T"], Trel))
                    if kj != k:
                        self.bonds.add((min(k, kj), max(k, kj)))
                    if new:
                        queue.append(kj)
        # Symmetriezusätze: je unterschiedlicher Operation eine römische Ziffer
        self.ops = {}
        for at in self.atoms:
            if at["generated"]:
                key = op_string(at["T"])
                if key not in self.ops:
                    self.ops[key] = roman(len(self.ops))
                at["sym"] = self.ops[key]
                at["opkey"] = key
            else:
                at["sym"], at["opkey"] = "", ""
            at["key"] = at["label"] + ("|" + at["opkey"] if at["opkey"] else "")
        self.bonds = sorted(self.bonds)
        self._fragments()

    def _fragments(self):
        n = len(self.atoms)
        parent = list(range(n))

        def find(x):
            while parent[x] != x:
                parent[x] = parent[parent[x]]
                x = parent[x]
            return x
        for a, b in self.bonds:
            parent[find(a)] = find(b)
        comps = {}
        for k in range(n):
            comps.setdefault(find(k), []).append(k)
        frags = sorted(comps.values(), key=lambda c: (-sum(self.atoms[k]["el"] not in ("H", "D") for k in c), min(c)))
        self.fragments = []
        for fi, c in enumerate(frags):
            for k in c:
                self.atoms[k]["frag"] = fi
            self.fragments.append(dict(atoms=c, formula=hill([self.atoms[k]["el"] for k in c])))

    def find_key(self, key):
        for k, a in enumerate(self.atoms):
            if a["key"] == key:
                return k
        return None

    def footnotes(self):
        """['i: 1−x, 1−y, 1−z', …]"""
        return [f"{r}: {op}" for op, r in self.ops.items()]


def hill(els):
    cnt = {}
    for e in els:
        cnt[e] = cnt.get(e, 0) + 1
    keys = [k for k in ("C", "H") if k in cnt] + sorted(k for k in cnt if k not in ("C", "H"))
    if "C" not in cnt:
        keys = sorted(cnt)
    return "".join(f"{k}{cnt[k] if cnt[k] > 1 else ''}" for k in keys)


# ---------------------------------------------------------------------------- Geometrie
def distance(p):
    return float(np.linalg.norm(p[1] - p[0]))


def angle(p):
    u, v = p[0] - p[1], p[2] - p[1]
    return float(math.degrees(math.atan2(np.linalg.norm(np.cross(u, v)), np.dot(u, v))))


def torsion(p):
    """IUPAC-Vorzeichen (wie SHELXL): Blick entlang B->C, A nach D im Uhrzeigersinn = positiv."""
    b1, b2, b3 = p[1] - p[0], p[2] - p[1], p[3] - p[2]
    x = np.dot(np.cross(b1, b2), np.cross(b2, b3))
    y = np.linalg.norm(b2) * np.dot(b1, np.cross(b2, b3))
    return float(math.degrees(math.atan2(y, x)))


def match_geom(S: Structure, rows, atoms):
    """Eintrag aus _geom_* zu gewählten Atomen (Label + Transformation T): gesucht ist G mit T_k = G∘S_k für alle k,
    auch in umgekehrter Reihenfolge. -> Textwert aus der CIF oder None."""
    labs = [a["label"] for a in atoms]
    for g in rows:
        for order in (labs, labs[::-1]):
            if g["labels"] != order:
                continue
            ats = atoms if order is labs else atoms[::-1]
            G = compose(ats[0]["T"], inverse(g["trs"][0]))
            if all(tr_key(compose(G, s)) == tr_key(a["T"]) for s, a in zip(g["trs"], ats)):
                return g["value"]
    return None


def measure(mol: Molecule, keys):
    """2/3/4 Atome -> dict(kind, value (float), text, esd (bool), source ('cif'|'calc'))."""
    idx = [mol.find_key(k) for k in keys]
    if None in idx or len(idx) not in (2, 3, 4) or len(set(idx)) != len(idx):
        return None
    ats = [mol.atoms[k] for k in idx]
    pts = [a["xyz"] for a in ats]
    kind = {2: "bond", 3: "angle", 4: "torsion"}[len(idx)]
    val = {2: distance, 3: angle, 4: torsion}[len(idx)](pts)
    rows = {2: mol.S.geom_bond, 3: mol.S.geom_angle, 4: mol.S.geom_torsion}[len(idx)]
    txt = match_geom(mol.S, rows, ats)
    if txt is not None:
        return dict(kind=kind, value=val, text=txt, esd=esd_parts(txt)[1] is not None, source="cif")
    dec = 3 if kind == "bond" else 1
    txt = f"{val:.{dec}f}"
    if float(txt) == 0:
        txt = txt.lstrip("-")                        # kein „−0.0“
    return dict(kind=kind, value=val, text=txt, esd=False, source="calc")


def atom_name(a):
    return a["label"] + (a["sym"].translate(SUP) if a.get("sym") else "")


def param_name(mol, keys, sep="–"):
    names = []
    for k in keys:
        i = mol.find_key(k)
        names.append(atom_name(mol.atoms[i]) if i is not None else k.split("|")[0])
    return sep.join(names)


def param_label(keys):
    """Zeilenname unabhängig vom Modell (für die Tabelle über mehrere Verbindungen): 'C1–C2ⁱ'."""
    out = []
    for k in keys:
        lab, _, op = k.partition("|")
        out.append(lab)
    return "–".join(out)


def default_view(mol: Molecule, frag=0):
    """Blick senkrecht auf die Ebene des größten Fragments (kleinste Hauptachse = Blickrichtung)."""
    pts = np.array([a["xyz"] for a in mol.atoms if a.get("frag") == frag and a["el"] not in ("H", "D")]
                   or [a["xyz"] for a in mol.atoms] or [[0, 0, 0]])
    if len(pts) < 3:
        return np.eye(3)
    c = pts - pts.mean(axis=0)
    w, v = np.linalg.eigh(c.T @ c)
    R = v[:, ::-1].T                     # Zeilen: x = größte, y = mittlere, z = kleinste Ausdehnung
    if np.linalg.det(R) < 0:
        R[2] *= -1
    return R


def rot(axis, ang):
    c, s = math.cos(ang), math.sin(ang)
    if axis == "x":
        return np.array([[1, 0, 0], [0, c, -s], [0, s, c]])
    if axis == "y":
        return np.array([[c, 0, s], [0, 1, 0], [-s, 0, c]])
    return np.array([[c, -s, 0], [s, c, 0], [0, 0, 1]])


def orthonormalize(R):
    u, _, vt = np.linalg.svd(np.asarray(R, float))
    Rn = u @ vt
    if np.linalg.det(Rn) < 0:
        Rn = -Rn
    return Rn


# ---------------------------------------------------------------------------- Projektion
def project(mol: Molecule, R, k, show, h_radius=0.11, n_seg=48):
    """Szene für die Abbildung (Koordinaten in Å, Bildschirm: x rechts, y oben, z zum Betrachter).
    -> dict(atoms=[dict(k, x, y, z, w, h, ang, iso, curves=[(xs, ys)], el, r_pick)], bonds=[(k1, k2, z)])"""
    R = np.asarray(R, float)
    center = np.array([a["xyz"] for a in mol.atoms]).mean(axis=0) if mol.atoms else np.zeros(3)
    out_atoms, pos = [], {}
    t = np.linspace(0, 2 * np.pi, n_seg + 1)
    ct, st = np.cos(t), np.sin(t)
    for i, a in enumerate(mol.atoms):
        if i not in show:
            continue
        p = R @ (a["xyz"] - center)
        pos[i] = p
        item = dict(k=i, x=float(p[0]), y=float(p[1]), z=float(p[2]), el=a["el"], curves=[], npd=False)
        U = a["U"]
        if a["el"] in ("H", "D"):
            item.update(w=2 * h_radius, h=2 * h_radius, ang=0.0, iso=True, kind="H")
        elif U is not None:
            Ur = R @ U @ R.T
            w2, v2 = np.linalg.eigh(Ur[:2, :2])
            ev, evec = np.linalg.eigh(U)
            if (ev <= 0).any() or (w2 <= 0).any():
                r = k * math.sqrt(max(a["uiso"], 1e-4))
                item.update(w=2 * r, h=2 * r, ang=0.0, iso=True, kind="iso", npd=True)
            else:
                item.update(w=2 * k * math.sqrt(w2[1]), h=2 * k * math.sqrt(w2[0]),
                            ang=float(math.degrees(math.atan2(v2[1, 1], v2[0, 1]))), iso=False, kind="aniso")
                axes = [R @ (evec[:, m] * k * math.sqrt(ev[m])) for m in range(3)]
                for m1, m2 in ((0, 1), (1, 2), (0, 2)):
                    off = np.outer(ct, axes[m1]) + np.outer(st, axes[m2])
                    front = off[:, 2] >= -1e-9
                    xs = np.where(front, p[0] + off[:, 0], np.nan)
                    ys = np.where(front, p[1] + off[:, 1], np.nan)
                    item["curves"].append((xs, ys))
        else:
            r = k * math.sqrt(max(a["uiso"], 1e-4))
            item.update(w=2 * r, h=2 * r, ang=0.0, iso=True, kind="iso")
        out_atoms.append(item)
    bonds = []
    for a, b in mol.bonds:
        if a in pos and b in pos:
            bonds.append((a, b, float((pos[a][2] + pos[b][2]) / 2)))
    return dict(atoms=out_atoms, bonds=bonds, center=center)


def pick(scene, x, y, tol=0.0):
    """Vorderstes Atom, dessen projizierter Umriss (x, y) enthält (Å, Toleranz in Å) -> Atomindex oder None."""
    best, bz = None, -1e9
    for it in scene["atoms"]:
        dx, dy = x - it["x"], y - it["y"]
        th = math.radians(it["ang"])
        u = dx * math.cos(th) + dy * math.sin(th)
        v = -dx * math.sin(th) + dy * math.cos(th)
        a, b = it["w"] / 2 + tol, it["h"] / 2 + tol
        if (u / a) ** 2 + (v / b) ** 2 <= 1 and it["z"] > bz:
            best, bz = it["k"], it["z"]
    return best


# ---------------------------------------------------------------------------- Tabellen (Rich Text)
def runs(markup):
    """Mini-Auszeichnung -> [(Text, Stil)]: *kursiv*, _{tief}, ^{hoch}, ~{Überstrich}."""
    out, ital, i, n = [], False, 0, len(markup)
    buf = ""

    def flush(style_extra=""):
        nonlocal buf
        if buf:
            out.append((buf, ("i" if ital else "") + style_extra))
            buf = ""
    while i < n:
        c = markup[i]
        if c == "*":
            flush()
            ital = not ital
            i += 1
        elif c in "_^~" and i + 1 < n and markup[i + 1] == "{":
            flush()
            j = markup.index("}", i)
            inner = markup[i + 2:j]
            st = {"_": "s", "^": "p", "~": "o"}[c]
            for t, s in runs(inner):
                out.append((t, s + st + ("i" if ital and "i" not in s else "")))
            i = j + 1
        else:
            buf += c
            i += 1
    flush()
    return out


def to_html(markup):
    h = []
    for t, s in runs(markup):
        t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if "o" in s:
            t = f'<span style="text-decoration: overline">{t}</span>'
        if "i" in s:
            t = f"<i>{t}</i>"
        if "s" in s:
            t = f"<sub>{t}</sub>"
        if "p" in s:
            t = f"<sup>{t}</sup>"
        h.append(t)
    return "".join(h)


_SUB_ALL = str.maketrans("0123456789+−-=()aehijklmnoprstuvx", "₀₁₂₃₄₅₆₇₈₉₊₋₋₌₍₎ₐₑₕᵢⱼₖₗₘₙₒₚᵣₛₜᵤᵥₓ")


def to_plain(markup):
    """Für CSV/Text: Tief-/Hochstellung als Unicode, falls möglich (sonst R_int), Überstrich kombinierend."""
    out = []
    for t, s in runs(markup):
        if "s" in s:
            t2 = t.translate(_SUB_ALL)
            t = t2 if all(c != d or c in " ,." for c, d in zip(t, t2)) else t
        if "p" in s:
            t2 = t.translate(SUP)
            t = t2 if all(c != d for c, d in zip(t, t2)) else "^" + t
        if "o" in s:
            t = "".join(ch + "\u0305" for ch in t)
        out.append(t)
    return "".join(out)


def to_rich_xlsx(markup):
    """CellRichText für openpyxl (kursiv, tief-/hochgestellt; Überstrich als kombinierendes Zeichen)."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    parts = []
    for t, s in runs(markup):
        if "o" in s:
            t = "".join(ch + "\u0305" for ch in t)
        if not ("i" in s or "s" in s or "p" in s):
            parts.append(t)
            continue
        f = InlineFont(i=("i" in s) or None, vertAlign="subscript" if "s" in s else ("superscript" if "p" in s else None))
        parts.append(TextBlock(f, t))
    if all(isinstance(p, str) for p in parts):
        return "".join(parts)
    return CellRichText(parts)


def space_group_markup(hm):
    """'P 21/c' -> '*P*2_{1}/*c*', 'P -1' -> '*P*~{1}', 'Pca21' -> '*Pca*2_{1}'."""
    s = str(hm or "").strip().strip("'\"")
    if not s:
        return ""
    s = re.sub(r"\s*:\s*\w+$", "", s)                # Ursprungswahl ':1'/':H' weglassen
    toks = s.split() if " " in s else None
    if toks and len(toks) == 4 and sum(t == "1" for t in toks[1:]) == 2:    # 'P 1 21/c 1' -> P2₁/c
        toks = [toks[0]] + [t for t in toks[1:] if t != "1"]
    out = []

    def screw(d1, d2):
        return int(d2) < int(d1) and d1 in "2346"

    def fmt_token(tok, first):
        res, i = "", 0
        while i < len(tok):
            c = tok[i]
            if c.isalpha():
                res += f"*{c}*"
                i += 1
            elif c == "-" and i + 1 < len(tok) and tok[i + 1].isdigit():
                res += "~{" + tok[i + 1] + "}"
                i += 2
            elif c.isdigit():
                if i + 1 < len(tok) and tok[i + 1].isdigit() and screw(c, tok[i + 1]):
                    res += c + "_{" + tok[i + 1] + "}"
                    i += 2
                else:
                    res += c
                    i += 1
            else:
                res += c
                i += 1
        return res
    if toks:
        out = [fmt_token(t, k == 0) for k, t in enumerate(toks)]
        txt = "".join(out)
    else:
        txt = fmt_token(s, True)
    return txt.replace("**", "")


def formula_markup(sum_formula):
    """'C20 H30 N2 Si' -> 'C_{20}H_{30}N_{2}Si'."""
    s = str(sum_formula or "").strip().strip("'\"")
    out = []
    for tok in s.split():
        m = re.match(r"^([A-Z][a-z]?)(\d*\.?\d*)([+-]?\d*[+-]?)$", tok)
        if not m:
            out.append(tok)
            continue
        el, n, ch = m.groups()
        out.append(el + (f"_{{{n}}}" if n and n != "1" else "") + (f"^{{{ch}}}" if ch else ""))
    return "".join(out)


RADIATION = [(0.71073, 0.002, "MoK_{α}"), (1.54178, 0.0012, "CuK_{α}"), (0.56086, 0.002, "AgK_{α}"),
             (1.34139, 0.002, "GaK_{α}")]


def radiation_markup(b: Block):
    lam = b.get("_diffrn_radiation_wavelength")
    v = num(lam)
    if v is None:
        return ""
    name = None
    for ref, tol, nm in RADIATION:
        if abs(v - ref) < tol:
            name = nm
    if name is None:
        typ = (b.get("_diffrn_radiation_type") or "").replace("\\a", "α").replace("\\b", "β").replace(" ", "")
        name = typ or "λ"
    return f"{name} (λ = {lam})"


def _dec(s):
    s = str(s)
    return len(s.split(".")[1]) if "." in s else 0


CRYSTAL_ROWS = [
    ("ref", "Internal reference"),
    ("ccdc", "CCDC number"),
    ("formula", "Formula"),
    ("fw", "Formula weight"),
    ("temp", "Temperature/K"),
    ("system", "Crystal system"),
    ("sg", "Space group"),
    ("a", "*a*/Å"), ("b", "*b*/Å"), ("c", "*c*/Å"),
    ("alpha", "α/°"), ("beta", "β/°"), ("gamma", "γ/°"),
    ("vol", "Volume/Å^{3}"),
    ("z", "*Z*"),
    ("rho", "Density (calc.)/g cm^{−3}"),
    ("mu", "Absorption coefficient/mm^{−1}"),
    ("f000", "*F*(000)"),
    ("size", "Crystal size/mm^{3}"),
    ("rad", "Radiation"),
    ("theta", "2θ range for data collection/°"),
    ("index", "Index ranges"),
    ("refl", "Reflections collected"),
    ("indep", "Independent reflections"),
    ("method", "Refinement method"),
    ("drp", "Data/restraints/parameters"),
    ("gof", "Goodness-of-fit on *F*^{2}"),
    ("rgt", "Final *R* indices [*I* > 2σ(*I*)]"),
    ("rall", "*R* indices (all data)"),
    ("diff", "Largest diff. peak and hole/e Å^{−3}"),
]


def crystal_values(b: Block, ref="", ccdc=""):
    """Spalte der Kristalldatentabelle: Schlüssel -> Markup (Zahlen exakt wie in der CIF, echtes Minus)."""
    g = b.get
    v = {}
    v["ref"] = ref or ""
    cc = g("_database_code_depnum_ccdc_archive") or ""
    cc = re.sub(r"(?i)^\s*ccdc\s*", "", cc.strip())
    v["ccdc"] = ccdc or cc
    v["formula"] = formula_markup(g("_chemical_formula_sum"))
    v["fw"] = g("_chemical_formula_weight") or ""
    v["temp"] = g("_diffrn_ambient_temperature", "_cell_measurement_temperature") or ""
    v["system"] = g("_space_group_crystal_system", "_symmetry_cell_setting") or ""
    v["sg"] = space_group_markup(g("_space_group_name_h-m_alt", "_symmetry_space_group_name_h-m"))
    for k, tag in (("a", "_cell_length_a"), ("b", "_cell_length_b"), ("c", "_cell_length_c"),
                   ("alpha", "_cell_angle_alpha"), ("beta", "_cell_angle_beta"), ("gamma", "_cell_angle_gamma"),
                   ("vol", "_cell_volume"), ("z", "_cell_formula_units_z"), ("rho", "_exptl_crystal_density_diffrn"),
                   ("mu", "_exptl_absorpt_coefficient_mu"), ("f000", "_exptl_crystal_f_000"),
                   ("refl", "_diffrn_reflns_number"), ("gof", "_refine_ls_goodness_of_fit_ref")):
        v[k] = g(tag) or ""
    sizes = [g(f"_exptl_crystal_size_{x}") for x in ("max", "mid", "min")]
    v["size"] = " × ".join(s for s in sizes if s) if all(sizes) else ""
    v["rad"] = radiation_markup(b)
    t0, t1 = g("_diffrn_reflns_theta_min"), g("_diffrn_reflns_theta_max")
    if t0 and t1 and num(t0) is not None and num(t1) is not None:
        d0, d1 = _dec(esd_parts(t0)[0]), _dec(esd_parts(t1)[0])
        v["theta"] = f"{2 * num(t0):.{d0}f} to {2 * num(t1):.{d1}f}"
    else:
        v["theta"] = ""
    lim = [g(f"_diffrn_reflns_limit_{h}_{m}") for h in "hkl" for m in ("min", "max")]
    if all(lim):
        v["index"] = ", ".join(f"{minus(lim[2 * i])} ≤ *{h}* ≤ {minus(lim[2 * i + 1])}" for i, h in enumerate("hkl"))
    else:
        v["index"] = ""
    tot, rint, rsig = g("_reflns_number_total"), g("_diffrn_reflns_av_r_equivalents"), \
        g("_diffrn_reflns_av_unetI/netI", "_diffrn_reflns_av_sigmaI/netI")
    if tot:
        extra = []
        if rint:
            extra.append(f"*R*_{{int}} = {rint}")
        if rsig:
            extra.append(f"*R*_{{sigma}} = {rsig}")
        v["indep"] = tot + (f" [{', '.join(extra)}]" if extra else "")
    else:
        v["indep"] = ""
    coef = (g("_refine_ls_structure_factor_coef") or "").lower()
    mat = (g("_refine_ls_matrix_type") or "full").lower()
    on = "*F*^{2}" if coef.startswith("fsq") else ("*F*" if coef else "*F*^{2}")
    v["method"] = ("Full-matrix" if mat == "full" else mat.capitalize()) + f" least-squares on {on}"
    nref = g("_refine_ls_number_reflns", "_reflns_number_total")
    nres, npar = g("_refine_ls_number_restraints"), g("_refine_ls_number_parameters")
    v["drp"] = f"{nref}/{nres}/{npar}" if nref and nres and npar else ""
    r1, wr = g("_refine_ls_r_factor_gt"), g("_refine_ls_wr_factor_gt")
    v["rgt"] = f"*R*_{{1}} = {r1}, *wR*_{{2}} = {wr}" if r1 and wr else ""
    r1a, wra = g("_refine_ls_r_factor_all"), g("_refine_ls_wr_factor_ref")
    v["rall"] = f"*R*_{{1}} = {r1a}, *wR*_{{2}} = {wra}" if r1a and wra else ""
    dmax, dmin = g("_refine_diff_density_max"), g("_refine_diff_density_min")
    if dmax and dmin:
        hole = dmin.strip()
        hole = "\u2212" + hole.lstrip("-\u2212") if not hole.startswith("+") else hole
        v["diff"] = f"{dmax}/{hole}"
    else:
        v["diff"] = ""
    return v


def write_table_csv(path, header, rows):
    """rows: [[markup, …]] -> CSV (UTF-8 mit BOM, Unicode-Tief-/Hochstellung)."""
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow([to_plain(h) for h in header])
        for r in rows:
            w.writerow([to_plain(c) for c in r])
    return path


def table_html(header, rows, title=None, notes=()):
    th = "".join(f'<th style="border-bottom:1px solid #000; padding:2px 8px; text-align:left">{to_html(h)}</th>'
                 for h in header)
    body = "".join("<tr>" + "".join(f'<td style="padding:1px 8px">{to_html(c)}</td>' for c in r) + "</tr>"
                   for r in rows)
    cap = f"<p><b>{to_html(title)}</b></p>" if title else ""
    foot = "".join(f'<p style="font-size:small">{to_html(n)}</p>' for n in notes)
    return f'{cap}<table style="border-collapse:collapse">{"<tr>" + th + "</tr>"}{body}</table>{foot}'


def table_plain(header, rows):
    return "\n".join("\t".join(to_plain(c) for c in r) for r in [header] + rows)


# ---------------------------------------------------------------------------- Abbildung (matplotlib)
import matplotlib  # noqa: E402
import matplotlib.patheffects as PE  # noqa: E402
from matplotlib.collections import PolyCollection  # noqa: E402
from matplotlib.colors import to_rgb  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from matplotlib.patches import Ellipse  # noqa: E402

SEL_COLOR = "#FF8C00"
LIGHT = np.array([-0.55, 0.65]) / math.hypot(0.55, 0.65)    # Licht von oben links (Bildschirm), wie Mercury
_T48 = np.linspace(0, 2 * np.pi, 49)


def _ellipse_verts(cx, cy, w, h, ang, n=None):
    t = _T48 if n is None else np.linspace(0, 2 * np.pi, n + 1)
    th = math.radians(ang)
    x, y = w / 2 * np.cos(t), h / 2 * np.sin(t)
    return np.column_stack([cx + x * math.cos(th) - y * math.sin(th), cy + x * math.sin(th) + y * math.cos(th)])


def shade_colors(color, n):
    """Mercury-ähnliche Schattierung: Rand dunkel, Mitte Grundfarbe, Glanzpunkt fast weiß."""
    base = np.array(to_rgb(color))
    out = []
    for t in np.linspace(0, 1, n):
        if t < 0.45:
            c = base * (0.5 + 0.5 * t / 0.45)
        else:
            f = (t - 0.45) / 0.55
            c = base + (1 - base) * 0.8 * f ** 1.4
        out.append(tuple(np.clip(c, 0, 1)))
    return out


def shade_layers(it, n):
    """Polygone der Schattierungsschichten (äußere zuerst). Jede Schicht ist die verkleinerte Ellipse, zum Licht
    verschoben, aber immer ganz innerhalb des Umrisses (Abstand in der Metrik der Ellipse ≤ 1 − s)."""
    a, b = it["w"] / 2, it["h"] / 2
    th = math.radians(it["ang"])
    u = LIGHT[0] * math.cos(th) + LIGHT[1] * math.sin(th)
    v = -LIGHT[0] * math.sin(th) + LIGHT[1] * math.cos(th)
    nrm = math.sqrt((u / a) ** 2 + (v / b) ** 2) or 1.0
    verts = []
    for s in np.linspace(1.0, 0.12, n):
        d = (1 - s) * 0.6 / nrm
        verts.append(_ellipse_verts(it["x"] + LIGHT[0] * d, it["y"] + LIGHT[1] * d, it["w"] * s, it["h"] * s,
                                    it["ang"]))
    return verts


def _radius_along(it, dx, dy):
    """Radius der projizierten Ellipse in Richtung (dx, dy)."""
    n = math.hypot(dx, dy) or 1.0
    th = math.radians(it["ang"])
    u = (dx * math.cos(th) + dy * math.sin(th)) / n
    v = (-dx * math.sin(th) + dy * math.cos(th)) / n
    a, b = it["w"] / 2, it["h"] / 2
    return 1.0 / math.sqrt((u / a) ** 2 + (v / b) ** 2)


def label_shown(a, o):
    """Label eines Atoms: Standard alle Nicht-H-Atome (H nur mit „H-Labels“), je Atom umschaltbar."""
    if not o.get("labels", True):
        return False
    over = (o.get("label_on") or {}).get(a["key"])
    if over is not None:
        return bool(over)
    return a["el"] not in ("H", "D") or bool(o.get("h_labels"))


class XrayFigure:
    """ORTEP-Abbildung eines Moleküls. Schnittstelle wie core.Figure, soweit Tab/Export sie brauchen
    (fig, ax, cfg, ann_items, axes_items, current_layout, export). Labels: Versatz in Punkten zum Atom,
    damit sie beim Drehen mitwandern; gespeicherte Versätze werden vor der automatischen Platzierung reserviert."""

    def __init__(self, cfg, mol, R, opts, layout, fig, selection=(), zoom=1.0, pan=(0.0, 0.0)):
        self.cfg, self.mol, self.opts, self.layout = cfg, mol, opts, layout or {}
        self.R = orthonormalize(R)
        self.zoom, self.pan = float(zoom), [float(pan[0]), float(pan[1])]
        self.selection = list(selection)
        fig.clear()
        fig.set_size_inches(cfg["figsize_in"], forward=False)
        self.fig = fig
        self.ax = fig.add_axes([0, 0, 1, 1])
        self.ax.set_axis_off()
        self.axes_items, self.ann_items, self.curves = {}, {}, []
        self.k = chi2_scale(opts.get("prob", 50) / 100)
        self.show = self._visible()
        pts = np.array([mol.atoms[i]["xyz"] for i in self.show]) if self.show else np.zeros((1, 3))
        self.center = pts.mean(axis=0)
        self.rmax = float(np.max(np.linalg.norm(pts - self.center, axis=1))) + 1.0
        sc = cfg.get("scale", 1.0)
        self.lw_bond = opts.get("bond_lw", 1.8) * sc
        self.lw_edge = max(0.4, 0.6 * sc)
        self.atom_art, self.bond_art, self.sel_art = {}, [], []
        self.auto_pos = {}
        self.scene = self._project()
        # Maßstab aus der Ausdehnung in der aktuellen Ansicht (beim Drehen bleibt er fest, nach dem Loslassen
        # passt das Neuzeichnen ihn an)
        xs = [it["x"] + s_ * max(it["w"], it["h"]) / 2 for it in self.scene["atoms"] for s_ in (-1, 1)] or [-1, 1]
        ys = [it["y"] + s_ * max(it["w"], it["h"]) / 2 for it in self.scene["atoms"] for s_ in (-1, 1)] or [-1, 1]
        self.ext = (max(abs(min(xs)), abs(max(xs)), 0.5) + 0.6, max(abs(min(ys)), abs(max(ys)), 0.5) + 0.6)
        self._limits()
        self._build()
        self._footnote()
        self._labels()

    # -- Szene ----------------------------------------------------------------------------
    def _visible(self):
        o, mol = self.opts, self.mol
        hidden = set(o.get("hidden_frags") or [])
        hidden_atoms = set(o.get("hidden_atoms") or [])
        return [i for i, a in enumerate(mol.atoms)
                if a.get("frag") not in hidden and (o.get("show_h", True) or a["el"] not in ("H", "D"))
                and a["key"] not in hidden_atoms]

    def _limits(self):
        W, H = self.cfg["figsize_in"]
        foot = 0.06 if self.mol.ops else 0.0                        # Platz für die Symmetrie-Fußnote
        s = max(2 * self.ext[0] / W, 2 * self.ext[1] / (H * (1 - 2 * foot))) * 1.04 / self.zoom   # Å pro Zoll
        cx, cy = self.pan
        self.ax.set_xlim(cx - W * s / 2, cx + W * s / 2)
        self.ax.set_ylim(cy - H * s / 2, cy + H * s / 2)

    def _project(self):
        scene = project(self.mol, self.R, self.k, set(self.show), h_radius=self.opts.get("h_radius", 0.13))
        # Projektion um den Schwerpunkt der gezeigten Atome (nicht aller), damit das Bild zentriert bleibt
        off = self.R @ (self.center - scene["center"])
        for it in scene["atoms"]:
            it["x"] -= off[0]
            it["y"] -= off[1]
            for c in range(len(it["curves"])):
                xs, ys = it["curves"][c]
                it["curves"][c] = (xs - off[0], ys - off[1])
        scene["by_k"] = {it["k"]: it for it in scene["atoms"]}
        return scene

    def _bond_xy(self, a, b):
        A, B = self.scene["by_k"][a], self.scene["by_k"][b]
        dx, dy = B["x"] - A["x"], B["y"] - A["y"]
        L = math.hypot(dx, dy)
        ra, rb = _radius_along(A, dx, dy), _radius_along(B, -dx, -dy)
        if L <= ra + rb:
            return [np.nan, np.nan], [np.nan, np.nan]
        ux, uy = dx / L, dy / L
        return [A["x"] + ux * ra, B["x"] - ux * rb], [A["y"] + uy * ra, B["y"] - uy * rb]

    def _layers(self, it):
        """Polygone eines Atoms: Schattierungsschichten (Mercury) bzw. eine Fläche (flach), zuletzt der Umriss."""
        base = _ellipse_verts(it["x"], it["y"], it["w"], it["h"], it["ang"])
        if self.opts.get("style", "mercury") == "flat":
            return [base, base]
        n = 6 if it.get("kind") == "H" else 12
        return shade_layers(it, n) + [base]

    def _build(self):
        ax, colors = self.ax, self.opts.get("colors") or {}
        mercury = self.opts.get("style", "mercury") != "flat"
        for it in self.scene["atoms"]:
            a = self.mol.atoms[it["k"]]
            col = colors.get(a["el"]) or default_color(a["el"])
            verts = self._layers(it)
            nfill = len(verts) - 1
            faces = shade_colors(col, nfill) if mercury else [col]
            edge = tuple(np.array(to_rgb(col)) * 0.35) if mercury else (0, 0, 0)
            p = PolyCollection(verts, facecolors=faces + ["none"], edgecolors=["none"] * nfill + [edge],
                               linewidths=[0] * nfill + [self.lw_edge * (0.8 if mercury else 1.0)],
                               antialiaseds=True, closed=True)
            ax.add_collection(p)
            lines = []
            for xs, ys in it["curves"]:
                ln = Line2D(xs, ys, color="#1a1a1a" if mercury else "black",
                            lw=max(0.3, 0.45 * self.cfg.get("scale", 1.0)), solid_capstyle="butt")
                ax.add_line(ln)
                lines.append(ln)
            self.atom_art[it["k"]] = (p, lines)
        for a, b, _ in self.scene["bonds"]:
            xs, ys = self._bond_xy(a, b)
            ln = Line2D(xs, ys, color="#262626" if mercury else "black", lw=self.lw_bond, solid_capstyle="butt")
            ax.add_line(ln)
            self.bond_art.append((a, b, ln))
        for n, k in enumerate(self.selection):
            if k in self.scene["by_k"]:
                it = self.scene["by_k"][k]
                e = Ellipse((it["x"], it["y"]), it["w"] * 1.35 + 0.12, it["h"] * 1.35 + 0.12, angle=it["ang"],
                            facecolor="none", edgecolor=SEL_COLOR, lw=2.2, zorder=50)
                ax.add_patch(e)
                self.sel_art.append((k, e))
        self._zorders()

    def _zorders(self):
        """Zeichenreihenfolge nach Tiefe (hinten zuerst); Bindungen zwischen ihren Atomen."""
        items = [(it["z"], "a", it["k"]) for it in self.scene["atoms"]]
        items += [(z, "b", n) for n, (_, _, z) in enumerate(self.scene["bonds"])]
        items.sort(key=lambda t: (t[0], t[1] == "a"))
        self.order = []
        for rank, (_, kind, key) in enumerate(items):
            zo = 2 + 40 * rank / max(1, len(items))
            if kind == "a":
                p, lines = self.atom_art[key]
                p.set_zorder(zo)
                self.order.append(p)
                for ln in lines:
                    ln.set_zorder(zo + 1e-4)
                    self.order.append(ln)
            else:
                ln = self.bond_art[key][2]
                ln.set_zorder(zo)
                self.order.append(ln)
        self.order += [e for _, e in self.sel_art] + list(self.ann_items.values())

    def set_view(self, R):
        """Nur Geometrie der vorhandenen Artists ändern (für flüssiges Drehen per Blitting)."""
        self.R = orthonormalize(R)
        self.scene = self._project()
        for it in self.scene["atoms"]:
            p, lines = self.atom_art[it["k"]]
            p.set_verts(self._layers(it))
            for ln, (xs, ys) in zip(lines, it["curves"]):
                ln.set_data(xs, ys)
        for a, b, ln in self.bond_art:
            ln.set_data(*self._bond_xy(a, b))
        for k, e in self.sel_art:
            it = self.scene["by_k"][k]
            e.set_center((it["x"], it["y"]))
            e.width, e.height, e.angle = it["w"] * 1.35 + 0.12, it["h"] * 1.35 + 0.12, it["ang"]
        for key, ann in self.ann_items.items():
            ann.xy = self._anchor(ann._atom)
        self._zorders()

    def set_zoom_pan(self, zoom=None, pan=None):
        if zoom is not None:
            self.zoom = zoom
        if pan is not None:
            self.pan = list(pan)
        self._limits()

    def artists(self):
        return self.order

    def _anchor(self, k):
        it = self.scene["by_k"][k]
        return it["x"], it["y"]

    # -- Labels -----------------------------------------------------------------------------
    def _label_text(self, a, bold=False):
        t = a["label"]
        if a.get("sym"):
            t += r"$^{\math" + ("bf" if bold else "rm") + "{" + a["sym"] + "}}$"
        return t

    def _labels(self):
        o = self.opts
        if not o.get("labels", True):
            return
        lc = self.cfg["labels"]
        fs = lc["font_size"]
        todo = []
        flat = o.get("style", "mercury") == "flat"
        # Mercury: normale Schrift ohne Rand; flach (Olex2): fett mit weißem Rand
        extra = dict(fontweight="bold", path_effects=[PE.withStroke(linewidth=2.2 * self.cfg.get("scale", 1.0),
                                                                    foreground="white")]) if flat else {}
        for it in self.scene["atoms"]:
            a = self.mol.atoms[it["k"]]
            if not label_shown(a, o):
                continue
            key = "lbl_" + a["key"]
            ann = self.ax.annotate(self._label_text(a, flat), xy=(it["x"], it["y"]), xycoords="data",
                                   xytext=(0, 0), textcoords="offset points", ha="center", va="center",
                                   fontsize=fs, color=lc.get("color") or "black", zorder=60,
                                   annotation_clip=False, **extra)
            ann._atom = it["k"]
            self.ann_items[key] = ann
            todo.append((key, ann, it))
        if not todo:
            return
        from matplotlib.backends.backend_agg import RendererAgg
        W, H = self.fig.get_size_inches() * self.fig.dpi
        rend = RendererAgg(int(W), int(H), self.fig.dpi)
        tr = self.ax.transData
        ppt = self.fig.dpi / 72.0
        # Belegt: Atome (Kreise um den Umriss) und Bindungen (Punkte entlang der Linie), in Pixeln
        circ = []
        for it in self.scene["atoms"]:
            cx, cy = tr.transform((it["x"], it["y"]))
            rx = abs(tr.transform((it["x"] + max(it["w"], it["h"]) / 2, it["y"]))[0] - cx)
            circ.append((cx, cy, rx))
        segpts = []
        for _, _, ln in self.bond_art:
            xs, ys = ln.get_data()
            if np.isfinite(xs).all():
                for f in np.linspace(0, 1, 7):
                    segpts.append(tr.transform((xs[0] + f * (xs[1] - xs[0]), ys[0] + f * (ys[1] - ys[0]))))
        circ = np.array(circ)
        segpts = np.array(segpts) if segpts else np.zeros((0, 2))
        boxes = []

        def cost(bx):
            x0, y0, x1, y1 = bx
            c = 0.0
            if len(circ):
                nx = np.clip(circ[:, 0], x0, x1)
                ny = np.clip(circ[:, 1], y0, y1)
                d = np.hypot(nx - circ[:, 0], ny - circ[:, 1])
                c += float(np.sum(np.clip(circ[:, 2] - d, 0, None))) * 2
            if len(segpts):
                inside = (segpts[:, 0] > x0) & (segpts[:, 0] < x1) & (segpts[:, 1] > y0) & (segpts[:, 1] < y1)
                c += 6.0 * int(inside.sum())
            for b in boxes:
                ox = min(x1, b[2]) - max(x0, b[0])
                oy = min(y1, b[3]) - max(y0, b[1])
                if ox > 0 and oy > 0:
                    c += 0.5 * ox * oy
            fw, fh = W, H
            if x0 < 0 or y0 < 0 or x1 > fw or y1 > fh:
                c += 1e4
            return c
        sizes = {}
        for key, ann, it in todo:
            bb = ann.get_window_extent(rend)
            sizes[key] = (bb.width, bb.height)
        # 1) gespeicherte Versätze zuerst eintragen
        for key, ann, it in todo:
            pos = self.layout.get(key)
            if pos and len(pos) == 2:
                ann.xyann = (float(pos[0]), float(pos[1]))
                cx, cy = tr.transform((it["x"], it["y"]))
                w, h = sizes[key]
                bx = cx + pos[0] * ppt, cy + pos[1] * ppt
                boxes.append((bx[0] - w / 2, bx[1] - h / 2, bx[0] + w / 2, bx[1] + h / 2))
        # 2) übrige: beste von 16 Richtungen, Abstand = Umriss + Rand
        for key, ann, it in todo:
            if self.layout.get(key):
                continue
            cx, cy = tr.transform((it["x"], it["y"]))
            w, h = sizes[key]
            best = None
            for rr in (1.0, 1.5):
                for ang in np.linspace(0, 2 * np.pi, 16, endpoint=False):
                    ux, uy = math.cos(ang), math.sin(ang)
                    r_at = _radius_along(it, ux, uy)
                    rpx = abs(tr.transform((it["x"] + r_at, it["y"]))[0] - cx)
                    # Abstand so, dass die Textbox den Umriss gerade nicht berührt
                    ext = abs(ux) * w / 2 + abs(uy) * h / 2
                    dist = (rpx + ext + 3.0 * ppt) * rr
                    bx = (cx + ux * dist - w / 2, cy + uy * dist - h / 2, cx + ux * dist + w / 2, cy + uy * dist + h / 2)
                    c = cost(bx) + (rr - 1) * 40 + 0.01 * ang
                    if best is None or c < best[0]:
                        best = (c, bx, ux * dist / ppt, uy * dist / ppt)
            _, bx, ox, oy = best
            ann.xyann = (round(ox, 1), round(oy, 1))
            boxes.append(bx)
        self.order += [a for a in self.ann_items.values() if a not in self.order]
        self.auto_pos = self.current_layout()

    def _footnote(self):
        notes = []
        used = {self.mol.atoms[i].get("opkey") for i in self.show}
        for op, r in self.mol.ops.items():
            if op in used and (self.opts.get("labels", True)):
                notes.append(r"$^{\mathrm{" + r + "}}$" + op)
        self.foot = None
        if notes:
            fs = self.cfg["labels"]["font_size"] * 0.8
            self.foot = self.fig.text(0.01, 0.01, "; ".join(notes), ha="left", va="bottom", fontsize=fs,
                                      color="black", zorder=70)

    # -- Schnittstelle Export/Tab ---------------------------------------------------------
    def current_layout(self):
        out = {}
        for key, ann in self.ann_items.items():
            out[key] = [round(float(v), 1) for v in ann.xyann]
        return out

    def data_per_px(self):
        x0, x1 = self.ax.get_xlim()
        return (x1 - x0) / max(1.0, self.ax.bbox.width)

    def export(self):
        from matplotlib.transforms import Bbox
        out = Path(self.cfg["output"])
        out.parent.mkdir(parents=True, exist_ok=True)
        paths = [out] + [out.with_suffix("." + e.lstrip(".")) for e in self.cfg["extra_formats"]]
        W, H = self.fig.get_size_inches()
        kw = dict(dpi=self.cfg["dpi"], bbox_inches=Bbox.from_bounds(0, 0, W, H), facecolor="white")
        import uvvis_images as IMG
        for p in dict.fromkeys(paths):
            with matplotlib.rc_context({"svg.fonttype": "path"}):
                self.fig.savefig(p, **kw)
            if p.suffix.lower() == ".svg":
                IMG.svg_defs_first(p)
        return paths


# ---------------------------------------------------------------------------- Maus
class XrayMouse:
    """Ziehen = Drehen (Blitting: nur das Molekül neu zeichnen), rechte/mittlere Taste = Verschieben,
    Mausrad = Zoom, Klick = Atom wählen (vorderstes gewinnt), Labels ziehen = Versatz ändern."""

    CLICK_PX = 4

    def __init__(self, F, tab):
        self.F, self.tab = F, tab
        self.mode, self.start, self.moved, self._bg = None, None, False, None
        for ann in F.ann_items.values():
            ann.draggable(True, use_blit=True)
        c = F.fig.canvas
        self.cids = [c.mpl_connect("button_press_event", self.press),
                     c.mpl_connect("motion_notify_event", self.move),
                     c.mpl_connect("button_release_event", self.release),
                     c.mpl_connect("scroll_event", self.scroll)]

    def disconnect(self):
        for cid in self.cids:
            self.F.fig.canvas.mpl_disconnect(cid)
        for ann in self.F.ann_items.values():
            try:
                ann.draggable(False)
            except Exception:
                pass

    def _over_label(self, ev):
        return any(a.contains(ev)[0] for a in self.F.ann_items.values())

    def press(self, ev):
        if ev.x is None:
            return
        if ev.dblclick and ev.button == 1:
            k = self.hit(ev)
            self.mode = None
            if k is not None:
                self.tab.atom_double_clicked(k)
            return
        if self._over_label(ev) and ev.button == 1:
            self.mode = "label"
            return
        self.mode = "rot" if ev.button == 1 else "pan"
        self.start = (ev.x, ev.y)
        self.R0 = self.F.R.copy()
        self.pan0 = list(self.F.pan)
        self.moved = False

    def _start_blit(self):
        canvas = self.F.fig.canvas
        if not getattr(canvas, "supports_blit", False):
            return
        for a in self.F.artists():
            a.set_animated(True)
        canvas.draw()
        self._bg = canvas.copy_from_bbox(self.F.fig.bbox)

    def _blit(self):
        canvas = self.F.fig.canvas
        if self._bg is None:
            canvas.draw_idle()
            return
        canvas.restore_region(self._bg)
        for a in sorted(self.F.artists(), key=lambda a: a.get_zorder()):
            self.F.ax.draw_artist(a)
        canvas.blit(self.F.fig.bbox)

    def move(self, ev):
        if ev.x is None:
            return
        if self.mode in ("rot", "pan"):
            dx, dy = ev.x - self.start[0], ev.y - self.start[1]
            if not self.moved and math.hypot(dx, dy) < self.CLICK_PX:
                return
            if not self.moved:
                self.moved = True
                self._start_blit()
            if self.mode == "rot":
                s = 2.5 * math.pi / max(200.0, min(self.F.fig.bbox.width, self.F.fig.bbox.height))
                self.F.set_view(rot("y", dx * s) @ rot("x", -dy * s) @ self.R0)
            else:
                d = self.F.data_per_px()
                if self._bg is not None:          # Verschieben ändert die Achsengrenzen -> ohne Blitting
                    for a in self.F.artists():
                        a.set_animated(False)
                    self._bg = None
                self.F.set_zoom_pan(pan=(self.pan0[0] - dx * d, self.pan0[1] - dy * d))
            self._blit()
        elif self.mode is None:
            hit = self.hit(ev)
            self.tab.hover_atom(hit)

    def hit(self, ev):
        if ev.xdata is None or self.F.scene is None:
            return None
        return pick(self.F.scene, ev.xdata, ev.ydata, tol=2 * self.F.data_per_px())

    def release(self, ev):
        mode, self.mode = self.mode, None
        if mode == "label":
            QTimer.singleShot(0, lambda: self.tab.labels_moved(self.F.current_layout()))
            return
        if mode is None:
            return
        if self._bg is not None:
            for a in self.F.artists():
                a.set_animated(False)
            self._bg = None
        if not self.moved:
            if mode == "rot":
                self.tab.atom_clicked(self.hit(ev))
            return
        if mode == "rot":
            self.tab.view_changed(self.F.R)
        else:
            self.tab.pan_changed(self.F.pan)

    def scroll(self, ev):
        if ev.x is None:
            return
        import time
        now = time.monotonic()
        if now - getattr(self, "_last", 0) < 0.03:      # Trackpads liefern sehr viele Ereignisse
            return
        self._last = now
        up = (ev.step > 0) if ev.step else (ev.button == "up")
        z = min(20.0, max(0.2, self.F.zoom * (1.08 if up else 1 / 1.08)))
        self.F.set_zoom_pan(zoom=z)
        self.F.fig.canvas.draw_idle()
        self.tab.zoom_changed(z)


# ---------------------------------------------------------------------------- Tab
import uuid  # noqa: E402

from matplotlib.backends.backend_agg import FigureCanvasAgg  # noqa: E402
from matplotlib.figure import Figure as MplFigure  # noqa: E402
from PySide6.QtCore import Qt, QTimer  # noqa: E402
from PySide6.QtWidgets import (QAbstractItemView, QCheckBox, QComboBox, QDoubleSpinBox,  # noqa: E402
                               QFileDialog, QGridLayout, QGroupBox, QHBoxLayout, QHeaderView, QLabel,
                               QListWidget, QListWidgetItem, QPushButton, QSpinBox, QTableWidget,
                               QTableWidgetItem, QTextBrowser, QVBoxLayout, QWidget)

import uvvis_core as core  # noqa: E402
import uvvis_project as PRJ  # noqa: E402
import uvvis_tabs as TABS  # noqa: E402
from uvvis_i18n import T  # noqa: E402
from uvvis_irms import copy_rich  # noqa: E402
from uvvis_style import HINT_CSS, WARN_CSS, make_form  # noqa: E402

KIND_HEAD = {"bond": "Bond lengths [Å]", "angle": "Angles [°]", "torsion": "Torsion angles [°]"}
CALC_NOTE = "calculated from the atomic coordinates; no s.u. in the CIF"


def is_cif(path):
    return Path(path).suffix.lower() == ".cif" and Path(path).is_file()


class XrayTab(TABS.FigureTab):
    STATE_KEY = "xray"
    DEFAULT_STATE = dict(size="half_a4", w_cm=16.0, h_cm=11.0, fmt={"pdf": True, "svg": True, "png": True},
                         layout={}, entries=[], current=None, prob=50, show_h=True, disorder="main", labels=True,
                         h_labels=False, colors={}, bond_lw=1.8, style="mercury", x_auto=True, x_min=0.0,
                         x_max=1.0, ymax=0.0)

    def __init__(self, main, state):
        self.sel = []                     # gewählte Atome (Schlüssel) der aktuellen Verbindung, nicht im Undo
        self._meas = None
        self._mol_cache = {}
        super().__init__(main, state)

    # -- Bedienelemente -------------------------------------------------------------------
    def build_controls(self):
        st = self.st
        g = QGroupBox(T("xr_structures"))
        v = QVBoxLayout(g)
        row = QHBoxLayout()
        for txt, fn in ((T("xr_add"), self.add_cif), ("↑", lambda: self.move_entry(-1)),
                        ("↓", lambda: self.move_entry(1)), (T("remove"), self.remove_entry)):
            b = QPushButton(txt)
            b.clicked.connect(fn)
            row.addWidget(b)
        row.addStretch(1)
        v.addLayout(row)
        self.tbl = QTableWidget(0, 3)
        self.tbl.setHorizontalHeaderLabels([T("xr_name"), T("xr_ref"), "CCDC"])
        self.tbl.verticalHeader().setVisible(False)
        self.tbl.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.tbl.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.tbl.setSelectionMode(QAbstractItemView.SingleSelection)
        self.tbl.setMinimumHeight(110)
        self.tbl.itemChanged.connect(self.entry_edited)
        self.tbl.itemSelectionChanged.connect(self.entry_selected)
        v.addWidget(self.tbl)
        self.lbl_info = QLabel("")
        self.lbl_info.setWordWrap(True)
        self.lbl_info.setStyleSheet(HINT_CSS)
        v.addWidget(self.lbl_info)
        self.lv.addWidget(g)

        g = QGroupBox(T("xr_display"))
        f = make_form(g)
        self.cmb_style = QComboBox()
        for key in ("mercury", "flat"):
            self.cmb_style.addItem(T("xr_style_" + key), key)
        self.cmb_style.setCurrentIndex(max(0, self.cmb_style.findData(st.get("style", "mercury"))))
        f.addRow(T("xr_style"), self.cmb_style)
        self.sp_prob = QSpinBox(minimum=10, maximum=99, singleStep=5, suffix=" %", value=int(st["prob"]))
        self.sp_prob.setToolTip(T("xr_prob_tip"))
        f.addRow(T("xr_prob"), self.sp_prob)
        self.cmb_dis = QComboBox()
        for key in ("main", "all", "minor"):
            self.cmb_dis.addItem(T("xr_dis_" + key), key)
        self.cmb_dis.setCurrentIndex(max(0, self.cmb_dis.findData(st["disorder"])))
        f.addRow(T("xr_disorder"), self.cmb_dis)
        self.cb_h = QCheckBox(T("xr_show_h"))
        self.cb_h.setChecked(st["show_h"])
        self.cb_lab = QCheckBox(T("xr_labels"))
        self.cb_lab.setChecked(st["labels"])
        self.cb_hlab = QCheckBox(T("xr_h_labels"))
        self.cb_hlab.setChecked(st["h_labels"])
        row = QHBoxLayout()
        for w_ in (self.cb_h, self.cb_lab, self.cb_hlab):
            row.addWidget(w_)
        row.addStretch(1)
        f.addRow(row)
        self.sp_bond = QDoubleSpinBox(decimals=1, minimum=0.5, maximum=5.0, singleStep=0.2, suffix=" pt",
                                      value=st["bond_lw"])
        f.addRow(T("xr_bond_lw"), self.sp_bond)
        self.lst_frag = QListWidget()
        self.lst_frag.setMaximumHeight(90)
        self.lst_frag.setToolTip(T("xr_frags_tip"))
        self.lst_frag.itemChanged.connect(self.frag_toggled)
        f.addRow(T("xr_frags"), self.lst_frag)
        self.tbl_atoms = QTableWidget(0, 3)
        self.tbl_atoms.setHorizontalHeaderLabels([T("xr_atom"), T("xr_atom_shown"), T("xr_atom_label")])
        self.tbl_atoms.verticalHeader().setVisible(False)
        self.tbl_atoms.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.tbl_atoms.setMinimumHeight(150)
        self.tbl_atoms.setToolTip(T("xr_atoms_tip"))
        self.tbl_atoms.itemChanged.connect(self.atom_toggled)
        f.addRow(T("xr_atoms"), self.tbl_atoms)
        row = QHBoxLayout()
        for txt, fn in ((T("xr_labels_all"), lambda: self.set_all_labels(True)),
                        (T("xr_labels_none"), lambda: self.set_all_labels(False)),
                        (T("xr_atoms_reset"), self.reset_atom_overrides)):
            b = QPushButton(txt)
            b.clicked.connect(fn)
            row.addWidget(b)
        row.addStretch(1)
        f.addRow(row)
        b = QPushButton(T("xr_reset_view"))
        b.clicked.connect(self.reset_view)
        f.addRow(b)
        hint = QLabel(T("xr_mouse_hint"))
        hint.setWordWrap(True)
        hint.setStyleSheet(HINT_CSS)
        f.addRow(hint)
        self.lv.addWidget(g)

        g = QGroupBox(T("xr_colors"))
        v = QVBoxLayout(g)
        self.color_grid = QGridLayout()
        v.addLayout(self.color_grid)
        b = QPushButton(T("xr_colors_reset"))
        b.clicked.connect(self.reset_colors)
        v.addWidget(b)
        self.lv.addWidget(g)

        g = QGroupBox(T("xr_select"))
        v = QVBoxLayout(g)
        self.lbl_meas = QLabel(T("xr_select_hint"))
        self.lbl_meas.setWordWrap(True)
        v.addWidget(self.lbl_meas)
        row = QHBoxLayout()
        self.btn_take = QPushButton(T("xr_take"))
        self.btn_take.clicked.connect(self.take_param)
        b = QPushButton(T("xr_clear_sel"))
        b.clicked.connect(self.clear_selection)
        row.addWidget(self.btn_take)
        row.addWidget(b)
        row.addStretch(1)
        v.addLayout(row)
        self.lst_par = QListWidget()
        self.lst_par.setMaximumHeight(130)
        v.addWidget(self.lst_par)
        row = QHBoxLayout()
        for txt, fn in (("↑", lambda: self.move_param(-1)), ("↓", lambda: self.move_param(1)),
                        (T("remove"), self.remove_param)):
            b = QPushButton(txt)
            b.clicked.connect(fn)
            row.addWidget(b)
        row.addStretch(1)
        v.addLayout(row)
        self.lv.addWidget(g)

        g = QGroupBox(T("xr_tables"))
        v = QVBoxLayout(g)
        self.txt_tab = QTextBrowser()
        self.txt_tab.setMinimumHeight(200)
        v.addWidget(self.txt_tab)
        row = QHBoxLayout()
        for txt, fn in ((T("xr_copy_crystal"), lambda: copy_rich(*self.crystal_clip())),
                        (T("xr_copy_params"), lambda: copy_rich(*self.params_clip())),
                        (T("xr_excel"), self.save_excel)):
            b = QPushButton(txt)
            b.clicked.connect(fn)
            row.addWidget(b)
        row.addStretch(1)
        v.addLayout(row)
        self.lv.addWidget(g)

        self.sp_prob.valueChanged.connect(self.changed)
        self.sp_bond.valueChanged.connect(self.changed)
        self.cmb_dis.currentIndexChanged.connect(self.disorder_changed)
        self.cmb_style.currentIndexChanged.connect(self.changed)
        for w_ in (self.cb_h, self.cb_lab, self.cb_hlab):
            w_.toggled.connect(self.changed)
        self.sync_widgets()

    def common_box(self):
        """Nur Größe und Formate (keine Achsen)."""
        g = QGroupBox(T("grp_plot"))
        f = make_form(g)
        st = self.st
        self.cmb_size = QComboBox()
        for key in ("half_a4", "custom"):
            self.cmb_size.addItem(T({"half_a4": "size_half_a4", "custom": "size_custom"}[key]), key)
        self.cmb_size.setCurrentIndex(max(0, self.cmb_size.findData(st["size"])))
        f.addRow(T("size"), self.cmb_size)
        row = QHBoxLayout()
        self.sp_w = QDoubleSpinBox(decimals=1, minimum=4, maximum=40, singleStep=0.5, suffix=" cm", value=st["w_cm"])
        self.sp_h = QDoubleSpinBox(decimals=1, minimum=3, maximum=40, singleStep=0.5, suffix=" cm", value=st["h_cm"])
        row.addWidget(self.sp_w)
        row.addWidget(QLabel("×"))
        row.addWidget(self.sp_h)
        f.addRow(T("size_wh"), row)
        row = QHBoxLayout()
        self.cb_fmt = {}
        for e in ("pdf", "svg", "png"):
            cb = QCheckBox(e.upper())
            cb.setChecked(st["fmt"].get(e, True))
            self.cb_fmt[e] = cb
            row.addWidget(cb)
        row.addStretch(1)
        f.addRow(T("formats"), row)
        for w_ in (self.sp_w, self.sp_h):
            w_.valueChanged.connect(self.changed)
        self.cmb_size.currentIndexChanged.connect(self.changed)
        for cb in self.cb_fmt.values():
            cb.toggled.connect(self.changed)
        self._sync_common()
        return g

    def _sync_common(self):
        if self.cmb_size.currentData() == "half_a4":
            for sp, v in ((self.sp_w, 16.0), (self.sp_h, 11.0)):
                sp.blockSignals(True)
                sp.setValue(v)
                sp.blockSignals(False)
        custom = self.cmb_size.currentData() == "custom"
        self.sp_w.setEnabled(custom)
        self.sp_h.setEnabled(custom)

    def collect_common(self):
        self.st.update(size=self.cmb_size.currentData(), w_cm=self.sp_w.value(), h_cm=self.sp_h.value(),
                       fmt={e: cb.isChecked() for e, cb in self.cb_fmt.items()})

    def collect(self):
        self.st.update(prob=self.sp_prob.value(), show_h=self.cb_h.isChecked(), labels=self.cb_lab.isChecked(),
                       h_labels=self.cb_hlab.isChecked(), disorder=self.cmb_dis.currentData(),
                       bond_lw=self.sp_bond.value(), style=self.cmb_style.currentData())

    # -- Einträge -------------------------------------------------------------------------
    def entries(self):
        return self.st["entries"]

    def entry(self):
        es = self.entries()
        cur = self.st.get("current")
        for e in es:
            if e["uid"] == cur:
                return e
        return es[0] if es else None

    def add_cif(self):
        if not self.main.ensure_project():
            return
        ps, _ = QFileDialog.getOpenFileNames(self, T("xr_add"), self.main.settings.value("last_dir", ""),
                                             f"CIF (*.cif);;{T('flt_all')}")
        for p in ps:
            self.load_cif(p)

    def load_cif(self, p):
        """Jeder data_-Block mit Atomen wird ein Eintrag (eine Spalte in den Tabellen)."""
        p = self.main.import_file(p, "data")
        try:
            blocks = self.cache(("cif", p), lambda: structure_blocks(p))
            if not blocks:
                raise ValueError(T("xr_no_atoms"))
            for b in blocks:
                Structure(b)
        except Exception as e:
            self.main.error(f"{Path(p).name}: {e}")
            return
        last = None
        for b in blocks:
            e = dict(uid=uuid.uuid4().hex[:8], file=p, block=b.name, name=b.name, ref="", ccdc="", view=None,
                     zoom=1.0, pan=[0.0, 0.0], layout={}, hidden_frags=[], params=[], hidden_atoms=[],
                     label_on={})
            self.entries().append(e)
            last = e["uid"]
        self.st["current"] = last
        self.sel = []
        self.sync_widgets()
        self.render()

    def remove_entry(self):
        e = self.entry()
        if e is None:
            return
        self.st["entries"] = [x for x in self.entries() if x["uid"] != e["uid"]]
        self.st["current"] = self.entries()[0]["uid"] if self.entries() else None
        self.sel = []
        self.sync_widgets()
        self.render()

    def move_entry(self, d):
        e = self.entry()
        if e is None:
            return
        es = self.entries()
        i = es.index(e)
        j = i + d
        if 0 <= j < len(es):
            es[i], es[j] = es[j], es[i]
            self.sync_widgets()
            self.render()

    def entry_edited(self, item):
        if self._loading:
            return
        es = self.entries()
        r = item.row()
        if r >= len(es):
            return
        key = ("name", "ref", "ccdc")[item.column()]
        es[r][key] = item.text().strip()
        self.update_tables()
        self.main.push_undo()

    def entry_selected(self):
        if self._loading:
            return
        rows = self.tbl.selectionModel().selectedRows()
        if not rows:
            return
        es = self.entries()
        r = rows[0].row()
        if r < len(es) and es[r]["uid"] != self.st.get("current"):
            self.st["current"] = es[r]["uid"]
            self.sel = []
            self.sync_widgets()
            self.render()

    # -- Modell ---------------------------------------------------------------------------
    def structure(self, e):
        p = e["file"]
        if not p or not Path(p).exists():
            return None
        blocks = self.cache(("cif", p), lambda: structure_blocks(p))
        for b in blocks:
            if b.name == e["block"]:
                return self.cache(("cifstruct", p, b.name), lambda: Structure(b))
        return None

    def molecule(self, e):
        S = self.structure(e)
        if S is None:
            return None
        key = (e["file"], e["block"], self.st["disorder"])
        if key not in self._mol_cache:
            self._mol_cache[key] = Molecule(S, self.st["disorder"])
        return self._mol_cache[key]

    def view_of(self, e, mol):
        v = e.get("view")
        if v and len(v) == 9:
            return np.array(v, float).reshape(3, 3)
        return default_view(mol)

    # -- Oberfläche -----------------------------------------------------------------------
    def sync_widgets(self):
        self._loading = True
        st = self.st
        es = self.entries()
        self.tbl.setRowCount(len(es))
        cur = self.entry()
        for r, e in enumerate(es):
            for c, k in enumerate(("name", "ref", "ccdc")):
                val = e.get(k) or ""
                if k == "ccdc" and not val:
                    S = self.structure(e)
                    val = ""
                    if S is not None:
                        val = crystal_values(S.block)["ccdc"]
                it = QTableWidgetItem(val)
                if k == "name":
                    it.setToolTip(f"{Path(e['file']).name}: data_{e['block']}")
                self.tbl.setItem(r, c, it)
        if cur is not None:
            self.tbl.selectRow(es.index(cur))
        for w_, v in ((self.sp_prob, int(st["prob"])), (self.sp_bond, st["bond_lw"])):
            w_.setValue(v)
        self.cmb_dis.setCurrentIndex(max(0, self.cmb_dis.findData(st["disorder"])))
        self.cb_h.setChecked(st["show_h"])
        self.cb_lab.setChecked(st["labels"])
        self.cb_hlab.setChecked(st["h_labels"])
        self.cmb_style.setCurrentIndex(max(0, self.cmb_style.findData(st.get("style", "mercury"))))
        self.fill_colors()
        self.fill_atoms()
        self.fill_frags()
        self.fill_params()
        self._loading = False

    def fill_colors(self):
        while self.color_grid.count():
            w_ = self.color_grid.takeAt(0).widget()
            if w_ is not None:
                w_.deleteLater()
        els = set()
        for e in self.entries():
            S = self.structure(e)
            if S is not None:
                els |= {a["el"] for a in S.atoms}
        order = [x for x in ("C", "H") if x in els] + sorted(els - {"C", "H"})
        cols = self.st.get("colors") or {}
        for n, el in enumerate(order):
            box = QWidget()
            h = QHBoxLayout(box)
            h.setContentsMargins(0, 0, 0, 0)
            h.addWidget(QLabel(el))
            h.addWidget(TABS.color_button(cols.get(el) or default_color(el), lambda c, el=el: self.set_color(el, c)))
            self.color_grid.addWidget(box, n // 4, n % 4)

    def set_color(self, el, c):
        self.st.setdefault("colors", {})[el] = c
        self.render()

    def reset_colors(self):
        self.st["colors"] = {}
        self.fill_colors()
        self.render()

    def fill_frags(self):
        self.lst_frag.clear()
        e = self.entry()
        mol = self.molecule(e) if e else None
        if mol is None:
            return
        hidden = set(e.get("hidden_frags") or [])
        for i, fr in enumerate(mol.fragments):
            it = QListWidgetItem(f"{i + 1}: {fr['formula']}" + ("" if i else f"  ({T('xr_main_frag')})"))
            it.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled)
            it.setCheckState(Qt.Unchecked if i in hidden else Qt.Checked)
            it.setData(Qt.UserRole, i)
            self.lst_frag.addItem(it)

    def frag_toggled(self, item):
        if self._loading:
            return
        e = self.entry()
        hidden = set(e.get("hidden_frags") or [])
        i = item.data(Qt.UserRole)
        hidden.discard(i) if item.checkState() == Qt.Checked else hidden.add(i)
        e["hidden_frags"] = sorted(hidden)
        self.timer.start()

    def fill_atoms(self):
        """Je Atom: sichtbar (auch einzelne H) und Label (überschreibt die Vorgabe)."""
        was = self._loading
        self._loading = True
        e = self.entry()
        mol = self.molecule(e) if e else None
        atoms = [] if mol is None else [a for a in mol.atoms if a.get("frag") not in set(e.get("hidden_frags") or [])
                                        and (self.st["show_h"] or a["el"] not in ("H", "D"))]
        self.tbl_atoms.setRowCount(len(atoms))
        hidden = set((e or {}).get("hidden_atoms") or [])
        o = self.opts(e) if e else {}
        for r, a in enumerate(atoms):
            it = QTableWidgetItem(atom_name(a))
            it.setFlags(Qt.ItemIsEnabled)
            it.setData(Qt.UserRole, a["key"])
            self.tbl_atoms.setItem(r, 0, it)
            for c, on in ((1, a["key"] not in hidden), (2, label_shown(a, o))):
                it = QTableWidgetItem()
                it.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled)
                it.setCheckState(Qt.Checked if on else Qt.Unchecked)
                it.setData(Qt.UserRole, a["key"])
                self.tbl_atoms.setItem(r, c, it)
        self.tbl_atoms.resizeColumnsToContents()
        self._loading = was

    def atom_toggled(self, item):
        if self._loading or item.column() == 0:
            return
        e = self.entry()
        if e is None:
            return
        key, on = item.data(Qt.UserRole), item.checkState() == Qt.Checked
        if item.column() == 1:
            hidden = set(e.get("hidden_atoms") or [])
            hidden.discard(key) if on else hidden.add(key)
            e["hidden_atoms"] = sorted(hidden)
            self.sel = [k for k in self.sel if k not in hidden]
        else:
            self.set_label(key, on)
            return
        self.timer.start()

    def set_label(self, key, on):
        """Label je Atom; entspricht es der Vorgabe, wird die Ausnahme gelöscht."""
        e = self.entry()
        mol = self.molecule(e) if e else None
        if mol is None:
            return
        k = mol.find_key(key)
        lo = e.setdefault("label_on", {})
        lo.pop(key, None)
        if k is not None and label_shown(mol.atoms[k], self.opts(e)) != on:
            lo[key] = bool(on)
        self.timer.start()

    def atom_double_clicked(self, k):
        """Doppelklick: Label dieses Atoms an/aus (der erste Klick des Doppelklicks wählt nicht aus)."""
        if self.F is None:
            return
        a = self.F.mol.atoms[k]
        if a["key"] in self.sel:
            self.sel.remove(a["key"])
            self.update_measure()
        self.set_label(a["key"], not label_shown(a, self.opts(self.entry())))
        self._loading = True
        self.fill_atoms()
        self._loading = False

    def set_all_labels(self, on):
        e = self.entry()
        mol = self.molecule(e) if e else None
        if mol is None:
            return
        o = dict(self.opts(e), label_on={})
        if not self.st["labels"]:
            self.st["labels"] = True
            self.cb_lab.setChecked(True)
        e["label_on"] = {a["key"]: on for a in mol.atoms if label_shown(a, o) != on}
        self.fill_atoms()
        self.render()

    def reset_atom_overrides(self):
        e = self.entry()
        if e is not None:
            e.update(hidden_atoms=[], label_on={})
            self.fill_atoms()
            self.render()

    def disorder_changed(self, *_):
        if self._loading:
            return
        self.sel = []
        self.changed()
        self._loading = True
        self.fill_frags()
        self._loading = False

    def fill_params(self):
        self.lst_par.clear()
        e = self.entry()
        mol = self.molecule(e) if e else None
        for p in (e or {}).get("params") or []:
            m = measure(mol, p["atoms"]) if mol else None
            name = param_name(mol, p["atoms"]) if mol else param_label(p["atoms"])
            txt = f"{name}   {minus(m['text']) if m else '–'}" + ("" if (m is None or m["source"] == "cif")
                                                                    else "  *")
            self.lst_par.addItem(txt)

    def move_param(self, d):
        e = self.entry()
        i = self.lst_par.currentRow()
        ps = (e or {}).get("params") or []
        j = i + d
        if 0 <= i < len(ps) and 0 <= j < len(ps):
            ps[i], ps[j] = ps[j], ps[i]
            self.fill_params()
            self.lst_par.setCurrentRow(j)
            self.update_tables()
            self.main.push_undo()

    def remove_param(self):
        e = self.entry()
        i = self.lst_par.currentRow()
        ps = (e or {}).get("params") or []
        if 0 <= i < len(ps):
            del ps[i]
            self.fill_params()
            self.update_tables()
            self.main.push_undo()

    # -- Auswahl / Maus ---------------------------------------------------------------------
    def hover_atom(self, k):
        if self.F is None or k is None:
            self.main.show_hover("")
            return
        a = self.F.mol.atoms[k]
        self.main.show_hover(f"{atom_name(a)}  ({a['el']})" + (f"   {a['opkey']}" if a.get("opkey") else ""))

    def atom_clicked(self, k):
        """Klick: Atom zur Auswahl (max. 4), erneuter Klick entfernt es, Klick ins Leere hebt sie auf."""
        if self.F is None:
            return
        if k is None:
            self.sel = []
        else:
            key = self.F.mol.atoms[k]["key"]
            if key in self.sel:
                self.sel.remove(key)
            else:
                if len(self.sel) >= 4:
                    self.sel = []
                self.sel.append(key)
        self.update_measure()
        self.render(push=False)

    def clear_selection(self):
        self.sel = []
        self.update_measure()
        self.render(push=False)

    def update_measure(self):
        e = self.entry()
        mol = self.molecule(e) if e else None
        self._meas = None
        if mol is None or not self.sel:
            self.lbl_meas.setText(T("xr_select_hint"))
            self.lbl_meas.setStyleSheet("")
            self.btn_take.setEnabled(False)
            return
        names = param_name(mol, self.sel)
        m = measure(mol, self.sel) if len(self.sel) >= 2 else None
        self._meas = m
        self.btn_take.setEnabled(m is not None)
        if m is None:
            self.lbl_meas.setText(names)
            self.lbl_meas.setStyleSheet("")
            return
        unit = " Å" if m["kind"] == "bond" else "°"
        txt = f"{T('xr_kind_' + m['kind'])}  {names} = {minus(m['text'])}{unit}"
        if m["source"] == "calc":
            txt += "\n" + T("xr_no_esd_torsion" if m["kind"] == "torsion" else "xr_no_esd")
        elif not m["esd"]:
            txt += "\n" + T("xr_cif_no_esd")
        self.lbl_meas.setText(txt)
        self.lbl_meas.setStyleSheet(WARN_CSS if m["source"] == "calc" else "")

    def take_param(self):
        e = self.entry()
        if e is None or self._meas is None:
            return
        keys = list(self.sel)
        if not any(p["atoms"] in (keys, keys[::-1]) for p in e["params"]):
            e["params"].append({"atoms": keys})
        self.sel = []
        self.update_measure()
        self.fill_params()
        self.render()

    def view_changed(self, R):
        e = self.entry()
        if e is not None:
            e["view"] = [round(float(v), 6) for v in np.asarray(R).ravel()]
            self.render()                          # nicht verschobene Labels werden neu platziert

    def pan_changed(self, pan):
        e = self.entry()
        if e is not None:
            e["pan"] = [round(float(pan[0]), 3), round(float(pan[1]), 3)]
            self.render()

    def zoom_changed(self, z):
        e = self.entry()
        if e is not None:
            e["zoom"] = round(float(z), 4)
            self.timer.start()

    def labels_moved(self, layout):
        """Nur tatsächlich verschobene Labels speichern; die übrigen bleiben automatisch platziert."""
        e = self.entry()
        if e is None or self.F is None:
            return
        lay = self.layout()
        for key, pos in layout.items():
            if key in lay or pos != self.F.auto_pos.get(key):
                lay[key] = pos
        self.main.push_undo()

    def reset_view(self):
        e = self.entry()
        if e is not None:
            e.update(view=None, zoom=1.0, pan=[0.0, 0.0], layout={})
            self.render()

    def reset_layout(self):
        e = self.entry()
        if e is not None:
            e["layout"] = {}
        self.render()

    def layout(self):
        e = self.entry()
        if e is None:
            return {}
        return e.setdefault("layout", {}).setdefault(self.size_key(), {})

    # -- Abbildung ------------------------------------------------------------------------
    def opts(self, e):
        st = self.st
        return dict(prob=st["prob"], show_h=st["show_h"], labels=st["labels"], h_labels=st["h_labels"],
                    colors=st.get("colors") or {}, bond_lw=st["bond_lw"], hidden_frags=e.get("hidden_frags") or [],
                    style=st.get("style", "mercury"), hidden_atoms=e.get("hidden_atoms") or [],
                    label_on=e.get("label_on") or {})

    def build(self, fig, e=None, interactive=True):
        e = e or self.entry()
        if e is None:
            return None
        mol = self.molecule(e)
        if mol is None or not mol.atoms:
            return None
        cfg = self.base_cfg()
        core.setup_fonts(TABS.plt, cfg)
        lay = e.setdefault("layout", {}).setdefault(self.size_key(), {})
        sel = [mol.find_key(k) for k in self.sel] if interactive else []
        return XrayFigure(cfg, mol, self.view_of(e, mol), self.opts(e), lay, fig,
                          selection=[k for k in sel if k is not None], zoom=e.get("zoom", 1.0),
                          pan=e.get("pan") or (0.0, 0.0))

    def render(self, push=True):
        try:
            if self.dragger:
                self.dragger.disconnect()
                self.dragger = None
            cfg = self.base_cfg()
            self.host.set_figure_size(*cfg["figsize_in"])
            self.F = self.build(self.fig)
            self.preview_stack.setCurrentIndex(0 if self.F is None else 1)
            if self.F is None:
                self.fig.clear()
            else:
                self.dragger = XrayMouse(self.F, self)
            self.update_info()
            self.update_tables()
            if not self._loading:
                self.fill_atoms()
            self.canvas.draw_idle()
            if push:
                self.main.push_undo()
        except Exception as ex:
            import traceback
            self.main.log_signal.emit(traceback.format_exc())
            self.main.statusBar().showMessage(f"{T('err_title')}: {ex}", 8000)

    def update_info(self):
        e = self.entry()
        mol = self.molecule(e) if e else None
        if mol is None:
            self.lbl_info.setText("")
            return
        S = mol.S
        bits = [T("xr_info", n=len(S.atoms), m=len(mol.atoms), sg=to_plain(space_group_markup(
            S.block.get("_space_group_name_h-m_alt", "_symmetry_space_group_name_h-m"))))]
        if mol.ops:
            bits.append(T("xr_sym_info", ops="; ".join(f"{r}: {op}" for op, r in mol.ops.items())))
        if S.parts():
            bits.append(T("xr_parts_info", parts=", ".join(str(p) for p in S.parts())))
        if not S.geom_bond:
            bits.append(T("xr_no_geom"))
        if not S.geom_torsion:
            bits.append(T("xr_no_torsion"))
        if mol.truncated:
            bits.append(T("xr_truncated"))
        self.lbl_info.setText("\n".join(bits))
        self.lbl_info.setStyleSheet(WARN_CSS if mol.truncated else HINT_CSS)

    # -- Tabellen -------------------------------------------------------------------------
    def crystal_table(self):
        """-> (Kopf, Zeilen) als Markup; eine Spalte je Verbindung."""
        cols = []
        for e in self.entries():
            S = self.structure(e)
            if S is not None:
                cols.append((e, crystal_values(S.block, e.get("ref", ""), e.get("ccdc", ""))))
        header = ["Compound"] + [e.get("name") or e["block"] for e, _ in cols]
        rows = [[lab] + [v[key] for _, v in cols] for key, lab in CRYSTAL_ROWS]
        return header, rows

    def params_table(self):
        """Zeilen = gewählte Parameter (gleichnamige Atome in einer Zeile), Spalten = Verbindungen."""
        es = [e for e in self.entries() if self.structure(e) is not None]
        rows_by_kind = {"bond": [], "angle": [], "torsion": []}
        cells, calc_used, notes = {}, False, []
        for ci, e in enumerate(es):
            mol = self.molecule(e)
            used_ops = {}
            for p in e.get("params") or []:
                m = measure(mol, p["atoms"])
                if m is None:
                    continue
                label = param_label(p["atoms"])
                key = (m["kind"], label)
                if key not in rows_by_kind[m["kind"]]:
                    rev = (m["kind"], "–".join(label.split("–")[::-1]))
                    if rev in rows_by_kind[m["kind"]]:
                        key = rev
                    else:
                        rows_by_kind[m["kind"]].append(key)
                txt = minus(m["text"])
                if m["source"] == "calc":
                    txt += "^{a}"
                    calc_used = True
                syms = []
                for k in p["atoms"]:
                    i = mol.find_key(k)
                    if i is not None and mol.atoms[i].get("sym"):
                        syms.append(mol.atoms[i]["sym"])
                        used_ops[mol.atoms[i]["sym"]] = mol.atoms[i]["opkey"]
                cells[(key, ci)] = txt + ("" if not syms else " (" + ",".join("*" + s + "*" for s in syms) + ")")
            if used_ops:
                notes.append(f"{e.get('name') or e['block']}: " + "; ".join(f"{r}: {op}" for r, op in used_ops.items()))
        header = [""] + [e.get("name") or e["block"] for e in es]
        rows = []
        for kind in ("bond", "angle", "torsion"):
            if rows_by_kind[kind]:
                rows.append([KIND_HEAD[kind]] + [""] * len(es))
                for key in rows_by_kind[kind]:
                    rows.append([key[1]] + [cells.get((key, ci), "–") for ci in range(len(es))])
        if calc_used:
            notes.append("^{a} " + CALC_NOTE)
        if notes and any("(" in c for r in rows for c in r[1:]):
            notes.insert(0, "Symmetry codes:")
        return header, rows, notes

    def crystal_clip(self):
        h, r = self.crystal_table()
        return table_plain(h, r), table_html(h, r)

    def params_clip(self):
        h, r, notes = self.params_table()
        title = "Selected bond lengths [Å], angles and torsion angles [°]"
        return table_plain(h, r) + ("\n" + "\n".join(to_plain(n) for n in notes) if notes else ""), \
            table_html(h, r, title, notes)

    def update_tables(self):
        if not self.entries():
            self.txt_tab.setHtml("")
            return
        h, r = self.crystal_table()
        html = table_html(h, r, "Crystal data and structure refinement")
        h2, r2, notes = self.params_table()
        if len(r2):
            html += "<br>" + table_html(h2, r2, "Selected bond lengths [Å], angles and torsion angles [°]", notes)
        self.txt_tab.setHtml(html)

    def write_excel(self, path):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Crystal data"
        h, r = self.crystal_table()
        self._sheet(ws, h, r, [], Font)
        ws2 = wb.create_sheet("Selected parameters")
        h2, r2, notes = self.params_table()
        self._sheet(ws2, h2, r2, notes, Font)
        wb.save(path)
        return path

    @staticmethod
    def _sheet(ws, header, rows, notes, Font):
        for c, v in enumerate(header, 1):
            cell = ws.cell(row=1, column=c, value=to_rich_xlsx(v))
            cell.font = Font(bold=True)
        for ri, row in enumerate(rows, 2):
            for c, v in enumerate(row, 1):
                ws.cell(row=ri, column=c, value=to_rich_xlsx(v) if v else None)
        for k, n in enumerate(notes):
            ws.cell(row=len(rows) + 3 + k, column=1, value=to_rich_xlsx(n))
        ws.column_dimensions["A"].width = 40
        for c in range(2, len(header) + 1):
            ws.column_dimensions[chr(64 + min(c, 26))].width = 28

    def save_excel(self):
        if not self.main.ensure_project():
            return
        pdir = Path(self.main.project_dir)
        p = PRJ.export_dir(pdir, self.STATE_KEY) / f"{pdir.name}_xray_tables.xlsx"
        try:
            self.write_excel(p)
        except PermissionError:
            return self.main.error(T("report_locked", p=p))
        self.main.info(T("exported_to", p=str(p.parent)))

    # -- Export -----------------------------------------------------------------------------
    def export_to(self, base: Path):
        """Je Verbindung eine Abbildung (eigene Ansicht), dazu Tabellen als CSV und Excel."""
        base = Path(base).resolve()
        base.parent.mkdir(parents=True, exist_ok=True)
        files = []
        es = [e for e in self.entries() if self.structure(e) is not None]
        for e in es:
            fig = MplFigure(figsize=(8, 6), dpi=TABS.PREVIEW_DPI)
            FigureCanvasAgg(fig)
            F = self.build(fig, e, interactive=False)
            if F is None:
                continue
            name = base.name + ("_" + core.safe_name(e.get("name") or e["block"]) if len(es) > 1 else "")
            stem = base.with_name(name)
            F.cfg["output"] = str(stem.with_suffix("." + F.cfg["_main_fmt"]))
            files += F.export()
        if not files and not es:
            raise RuntimeError(T("tab_empty"))
        h, r = self.crystal_table()
        files.append(write_table_csv(base.with_name(base.name + "_crystal_data.csv"), h, r))
        h2, r2, notes = self.params_table()
        files.append(write_table_csv(base.with_name(base.name + "_parameters.csv"), h2, r2 + [[n] for n in notes]))
        try:
            files.append(self.write_excel(base.with_name(base.name + "_tables.xlsx")))
        except PermissionError:
            self.main.error(T("report_locked", p=base.with_name(base.name + "_tables.xlsx")))
        return files


# ---------------------------------------------------------------------------- Testdaten (Selbsttest)
def make_test_cif(path, blocks=("xtest_a", "xtest_b")):
    """Künstliche Struktur für den Selbsttest: zentrosymmetrischer C₆-Ring auf dem Inversionszentrum (halber Ring
    in der ASU), Cl fehlgeordnet (PART 1 0.7 / PART 2 0.3), separates Wassermolekül, aniso U, _geom_bond und
    _geom_angle mit esd, keine _geom_torsion. Kein echter Datensatz."""
    cell = (7.0, 8.0, 9.0, 90.0, 100.0, 90.0)
    M = ortho_matrix(*cell)
    Mi = np.linalg.inv(M)
    r, rcl, rh = 1.39, 1.39 + 1.74, 1.39 + 0.95
    ex, ey = np.array([0.6, 0.8, 0.0]), np.array([-0.48, 0.36, 0.8])     # Ringebene (orthonormal)

    def cart(rad, deg):
        a = math.radians(deg)
        return rad * (math.cos(a) * ex + math.sin(a) * ey)
    water = np.array([3.4, 3.9, 2.2])
    atoms = [("C1", "C", cart(r, 0), "Uani", 1, "."), ("C2", "C", cart(r, 60), "Uani", 1, "."),
             ("C3", "C", cart(r, 120), "Uani", 1, "."), ("H3", "H", cart(rh, 120), "Uiso", 1, "."),
             ("Cl1", "Cl", cart(rcl, 0), "Uani", 0.7, "1"), ("Cl1B", "Cl", cart(rcl, 60), "Uani", 0.3, "2"),
             ("O1", "O", water, "Uani", 1, "."), ("H1O", "H", water + [0.96, 0, 0], "Uiso", 1, "."),
             ("H2O", "H", water + [-0.24, 0.93, 0], "Uiso", 1, ".")]
    lines = []
    for bn in blocks:
        lines += [f"data_{bn}", "_chemical_formula_sum 'C6 H4 Cl2 O'", "_chemical_formula_weight 163.00",
                  "_space_group_crystal_system monoclinic", "_space_group_name_H-M_alt 'P 21/c'",
                  "loop_", "  _space_group_symop_operation_xyz", " 'x, y, z'", " '-x, y+1/2, -z+1/2'",
                  " '-x, -y, -z'", " 'x, -y-1/2, z-1/2'",
                  "_cell_length_a 7.0000(3)", "_cell_length_b 8.0000(4)", "_cell_length_c 9.0000(4)",
                  "_cell_angle_alpha 90", "_cell_angle_beta 100.000(2)", "_cell_angle_gamma 90",
                  "_cell_volume 496.33(4)", "_cell_formula_units_Z 2", "_exptl_crystal_density_diffrn 1.091",
                  "_exptl_absorpt_coefficient_mu 0.611", "_exptl_crystal_F_000 164",
                  "_exptl_crystal_size_max 0.30", "_exptl_crystal_size_mid 0.20", "_exptl_crystal_size_min 0.10",
                  "_diffrn_ambient_temperature 100(2)", "_diffrn_radiation_wavelength 0.71073",
                  "_diffrn_reflns_theta_min 2.500", "_diffrn_reflns_theta_max 27.500",
                  "_diffrn_reflns_limit_h_min -9", "_diffrn_reflns_limit_h_max 9", "_diffrn_reflns_limit_k_min -10",
                  "_diffrn_reflns_limit_k_max 10", "_diffrn_reflns_limit_l_min -11", "_diffrn_reflns_limit_l_max 11",
                  "_diffrn_reflns_number 5000", "_reflns_number_total 1100", "_diffrn_reflns_av_R_equivalents 0.0310",
                  "_diffrn_reflns_av_unetI/netI 0.0250", "_refine_ls_structure_factor_coef Fsqd",
                  "_refine_ls_matrix_type full", "_refine_ls_number_reflns 1100", "_refine_ls_number_restraints 3",
                  "_refine_ls_number_parameters 70", "_refine_ls_goodness_of_fit_ref 1.050",
                  "_refine_ls_R_factor_gt 0.0350", "_refine_ls_wR_factor_gt 0.0900", "_refine_ls_R_factor_all 0.0400",
                  "_refine_ls_wR_factor_ref 0.0950", "_refine_diff_density_max 0.42", "_refine_diff_density_min -0.31",
                  "loop_", "  _atom_site_label", "  _atom_site_type_symbol", "  _atom_site_fract_x",
                  "  _atom_site_fract_y", "  _atom_site_fract_z", "  _atom_site_U_iso_or_equiv", "  _atom_site_adp_type",
                  "  _atom_site_occupancy", "  _atom_site_disorder_group"]
        for lab, el, xyz, adp, occ, part in atoms:
            f = Mi @ xyz
            lines.append(f" {lab} {el} {f[0]:.6f} {f[1]:.6f} {f[2]:.6f} 0.0300(5) {adp} {occ} {part}")
        lines += ["loop_", "  _atom_site_aniso_label"] + [f"  _atom_site_aniso_U_{k}" for k in
                                                         ("11", "22", "33", "23", "13", "12")]
        for lab, el, _, adp, _, _ in atoms:
            if adp == "Uani":
                lines.append(f" {lab} 0.0250(5) 0.0300(5) 0.0400(6) 0.0030(4) 0.0050(4) -0.0020(4)")
        lines += ["loop_", "  _geom_bond_atom_site_label_1", "  _geom_bond_atom_site_label_2", "  _geom_bond_distance",
                  "  _geom_bond_site_symmetry_2", "  _geom_bond_publ_flag",
                  " C1 C2 1.390(3) . ?", " C2 C3 1.390(3) . ?", " C3 C1 1.390(3) 3_555 ?", " C1 Cl1 1.740(2) . ?",
                  " C2 Cl1B 1.740(9) . ?", " C3 H3 0.9500 . ?", " O1 H1O 0.96(2) . ?", " O1 H2O 0.96(2) . ?",
                  "loop_", "  _geom_angle_atom_site_label_1", "  _geom_angle_atom_site_label_2",
                  "  _geom_angle_atom_site_label_3", "  _geom_angle", "  _geom_angle_site_symmetry_1",
                  "  _geom_angle_site_symmetry_3", "  _geom_angle_publ_flag",
                  " C2 C1 C3 120.0(2) . 3 ?", " C1 C2 C3 120.0(2) . . ?", ""]
    Path(path).write_text("\n".join(lines), encoding="utf-8", newline="\n")
    return Path(path), {lab: xyz for lab, _, xyz, *_ in atoms}
